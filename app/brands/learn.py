"""Arculat tanulása mintákból: képek, PDF, pptx, docx → paletta, betűk, logó, elrendezési szabályok.

Lánc mintánként:
  • pptx: a téma színei és betűi PONTOSAN (style_ref._pptx_theme), a mesterdia/elrendezések kis, sarokban ülő képe = logó-jelölt;
  • docx: a dokumentum és a téma színei/betűi (brand_extract._scan_docx), a fejléc képe = logó-jelölt;
  • pdf: a betűnevek (pdffonts), az első oldal sarkában ülő kis kép = logó-jelölt;
  • kép: átlátszó vagy kevés színű kis kép = maga a logó; különben mintaoldal.
Minden minta első oldalának képéből k-means paletta (mért színek), majd EGY látómodell-hívás (DeepSeek-V4.1-Flash) a
mintaoldalakkal: a színek szerepe (fő/kiemelő/háttér/szöveg), a betűk, a logó helye, ha nincs jelölt, és az elrendezési
szokások (archetípusok, igazítás, kis/nagybetű, díszítés). A modell hibája nem állítja meg: a mért adatokból áll össze.
A betűneveket a design_auto.font_family az elérhető családra képezi (a minta „Montserrat”-ja → Lato); az eredeti név megmarad.
"""
from __future__ import annotations

import asyncio
import collections
import json
import logging
import re
from pathlib import Path
from typing import Any, Optional

from app.config import settings

log = logging.getLogger("aice.brand_learn")
_HEX = re.compile(r"^#[0-9A-Fa-f]{6}$")
RULE_KEYS = ("layouts", "align", "title_case", "decoration", "corners", "photos", "mood")


# ---------- logó-jelöltek a fájlokból ----------

def _pptx_logo(src: Path) -> Optional[bytes]:
    """A mesterdián/elrendezéseken ülő legkisebb, sarokhoz közeli kép (a logó szokásos helye)."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(src))
    W, H = prs.slide_width or 1, prs.slide_height or 1
    best: tuple[float, Optional[bytes]] = (9.0, None)
    shapes = []
    for m in prs.slide_masters:
        shapes += list(m.shapes)
        for lay in m.slide_layouts:
            shapes += list(lay.shapes)
    for sh in shapes:
        if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        area = (sh.width * sh.height) / (W * H)
        if area > 0.12 or area < 0.0005:
            continue
        cx, cy = (sh.left + sh.width / 2) / W, (sh.top + sh.height / 2) / H
        corner = min(abs(cx - 0.0) + abs(cy - 0.0), abs(cx - 1) + abs(cy - 0), abs(cx - 0) + abs(cy - 1), abs(cx - 1) + abs(cy - 1))
        score = corner + area * 3
        if score < best[0]:
            best = (score, sh.image.blob)
    return best[1]


def _docx_logo(src: Path) -> Optional[bytes]:
    from docx import Document
    doc = Document(str(src))
    for sec in doc.sections[:1]:
        for part in (sec.header, sec.first_page_header):
            try:
                for rel in part.part.rels.values():
                    if "image" in rel.reltype:
                        return rel.target_part.blob
            except Exception:
                continue
    return None


def _pdf_logo(src: Path) -> Optional[bytes]:
    import fitz
    with fitz.open(str(src)) as doc:
        if not doc.page_count:
            return None
        page = doc[0]
        W, H = page.rect.width, page.rect.height
        best: tuple[float, Optional[int]] = (9.0, None)
        for img in page.get_images(full=True):
            xref = img[0]
            for r in page.get_image_rects(xref):
                area = (r.width * r.height) / (W * H)
                if area > 0.08 or area < 0.0005:
                    continue
                cx, cy = (r.x0 + r.x1) / 2 / W, (r.y0 + r.y1) / 2 / H
                score = min(cx, 1 - cx) + min(cy, 1 - cy) + area * 3
                if cy < 0.25 or cy > 0.85:
                    score -= 0.2                       # fejléc/lábléc sávja
                if score < best[0]:
                    best = (score, xref)
        if best[1] is None:
            return None
        pix = fitz.Pixmap(doc, best[1])
        if pix.n - pix.alpha >= 4:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        return pix.tobytes("png")


def _image_is_logo(src: Path) -> bool:
    """Átlátszó, vagy kicsi és kevés színű kép → maga a logó (nem fotó)."""
    from PIL import Image
    with Image.open(src) as im:
        if im.mode in ("RGBA", "LA", "P") and "A" in im.convert("RGBA").getbands():
            a = im.convert("RGBA").getchannel("A")
            if a.getextrema()[0] < 200:
                return True
        small = im.convert("RGB").resize((64, 64))
        colors = len(set(small.getdata()))
        return max(im.size) <= 1200 and colors < 400


def _clean_logo(data: bytes) -> bytes:
    """A logó egyszínű hátterét átlátszóvá teszi (a sarkok színéből), és a látható részre vágja."""
    import io
    import numpy as np
    from PIL import Image
    im = Image.open(io.BytesIO(data)).convert("RGBA")
    a = np.asarray(im).astype(np.int16)
    corners = np.array([a[0, 0, :3], a[0, -1, :3], a[-1, 0, :3], a[-1, -1, :3]])
    if a[..., 3].min() > 200 and np.ptp(corners, axis=0).max() < 24:      # nincs átlátszóság, egyszínű a háttér
        bg = corners.mean(0)
        d = np.abs(a[..., :3] - bg).max(-1)
        alpha = np.clip((d - 12) * 255 / 30, 0, 255)
        a[..., 3] = np.minimum(a[..., 3], alpha)
    out = Image.fromarray(a.astype(np.uint8), "RGBA")
    bb = out.getchannel("A").point(lambda v: 255 if v > 20 else 0).getbbox()
    if bb:
        out = out.crop(bb)
    out.thumbnail((1200, 1200), Image.Resampling.LANCZOS)
    buf = io.BytesIO()
    out.save(buf, "PNG")
    return buf.getvalue()


# ---------- mintaoldalak ----------

def _render(src: Path, dst: Path) -> list[str]:
    from app.tools import style_ref
    return style_ref._first_page_png(src, dst)


def _colors_of(png: Path, k: int = 8) -> list[tuple[str, float]]:
    """(szín, arány) a mintaoldalon, gyakoriság szerint."""
    import numpy as np
    from PIL import Image
    from app.tools.design_auto import _hex
    a = np.asarray(Image.open(png).convert("RGB").resize((96, 96)), dtype=np.float32).reshape(-1, 3)
    rng = np.random.default_rng(3)
    c = a[rng.choice(len(a), k, replace=False)]
    for _ in range(14):
        lab = ((a[:, None, :] - c[None]) ** 2).sum(-1).argmin(1)
        for i in range(k):
            sel = a[lab == i]
            if len(sel):
                c[i] = sel.mean(0)
    counts = np.bincount(lab, minlength=k) / len(a)
    return [(_hex(c[i]), float(counts[i])) for i in counts.argsort()[::-1] if counts[i] > 0.004]


def _roles(measured: list[tuple[str, float]], exact: list[str]) -> dict[str, Any]:
    """Mért (és pontos) színekből szerepek: háttér = a leggyakoribb világos/sötét, fő = a leggyakoribb színes-sötét,
    kiemelő = a legtelítettebb, az előzőektől eltérő árnyalat."""
    from app.tools.design_auto import _lum, _sat, readable
    pool = [c for c, _ in measured] + [c for c in exact if _HEX.match(c)]
    if not pool:
        return {"primary": "#14213D", "accent": "#FCA311", "bg": "#FFFFFF", "text": "#1E2430", "extra": []}
    bg = measured[0][0] if measured else "#FFFFFF"
    cand = [c for c in pool if c != bg]
    primary = next((c for c in cand if _sat(c) > 0.25 and _lum(c) < 0.35), None) or min(cand or [bg], key=_lum)
    accent = max((c for c in cand if c != primary), key=lambda c: _sat(c) * (0.4 + min(_lum(c), 0.6)), default=primary)
    text = readable(min(pool, key=_lum), bg)
    extra = [c for c in dict.fromkeys(cand) if c not in (primary, accent)][:4]
    return {"primary": primary.upper(), "accent": accent.upper(), "bg": bg.upper(), "text": text.upper(), "extra": [e.upper() for e in extra]}


_VISION = """These are examples of one brand's designs (posts, slides, documents). Describe the brand's visual identity
as ONE JSON object:
{"name": "the brand / organisation name if visible, else null",
 "colors": {"primary": "#hex main brand colour", "accent": "#hex highlight colour", "bg": "#hex usual background",
            "text": "#hex main text colour", "extra": ["#hex", "..."]},
 "fonts": {"heading": "the headline typeface name, your best guess", "body": "the body typeface name"},
 "logo": {"image": 0-based index of the example that shows the logo most clearly, "box": [x0, y0, x1, y1] as 0-1 fractions} or null,
 "rules": {"layouts": ["2-3 of: full_bleed, split, typographic, framed, band, cutout"], "align": "left|center",
           "title_case": "upper|as_is", "decoration": "rule|frame|blocks|none", "corners": "square|rounded",
           "photos": "full-bleed|framed|cut-out|rare", "mood": "one sentence on tone and layout habits"}}
Measured colours (most used first): MEASURED. Fonts named in the files: FONTS. Use the measured colours where they fit.
Only the JSON."""


async def _vision(pngs: list[Path], measured: list[str], fonts: list[str]) -> dict[str, Any]:
    from app.llm import vision
    prompt = _VISION.replace("MEASURED", ", ".join(measured[:10]) or "none").replace("FONTS", ", ".join(fonts[:8]) or "none")
    content: list[dict[str, Any]] = [{"type": "text", "text": prompt}]
    for p in pngs[:3]:
        content.append({"type": "image_url", "image_url": {"url": vision.image_to_data_url(p, long_edge=1024)}})
    last: Exception = RuntimeError("no model")
    for model in (settings.LLM_MODEL_VISION, settings.LLM_MODEL_VISION_FALLBACK):
        try:
            raw = await vision._call([{"role": "user", "content": content}], max_tokens=900, model=model)
            return json.loads(raw[raw.find("{"):raw.rfind("}") + 1])
        except Exception as e:
            last = e
            log.warning("arculat-látás %s: %s", model, e)
    raise last


async def learn(file_ids: list[str], name: Optional[str] = None, progress=None) -> dict[str, Any]:
    """Minták → arculat-spec (a logó külön fájlként mentve): {name, colors, fonts, logo_file_id, rules, sources}."""
    from app.storage import output_path, temp_path, write_meta
    from app.storage.files import new_file_id
    from app.tools import crafts, design_auto, style_ref
    from app.brands.extract import _scan_docx
    if not file_ids:
        raise ValueError("Add a few examples of the brand — posts, slides, a PDF or a logo.")
    pngs: list[Path] = []
    exact: list[str] = []
    fonts: collections.Counter = collections.Counter()
    logo: Optional[bytes] = None
    theme_fonts: dict[str, str] = {}
    tmp: list[Path] = []
    try:
        for i, fid in enumerate(file_ids[:8]):
            if progress:
                progress(i, min(len(file_ids), 8), f"reading example {i + 1}")
            src = crafts.existing(fid)
            ext = src.suffix.lower()
            try:
                if ext in (".pptx", ".potx"):
                    t = await asyncio.to_thread(style_ref._pptx_theme, src)
                    exact += [t[k] for k in ("primary", "accent", "bg", "text") if t.get(k)]
                    theme_fonts.setdefault("heading", t.get("font_heading"))
                    theme_fonts.setdefault("body", t.get("font_body"))
                    logo = logo or await asyncio.to_thread(_pptx_logo, src)
                elif ext in (".docx", ".dotx"):
                    sc = await asyncio.to_thread(_scan_docx, src)
                    exact += [c["hex"] for c in sc.get("top_colors", [])[:4]] + sc.get("theme_colors", [])[:4]
                    for f in sc.get("top_fonts", []):
                        fonts[f["name"]] += f["count"]
                    for f in sc.get("theme_fonts", [])[:2]:
                        fonts[f] += 5
                    logo = logo or await asyncio.to_thread(_docx_logo, src)
                elif ext == ".pdf":
                    logo = logo or await asyncio.to_thread(_pdf_logo, src)
                elif ext in (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff", ".svg"):
                    if ext != ".svg" and not logo and await asyncio.to_thread(_image_is_logo, src):
                        logo = src.read_bytes()
                        continue                           # a logó nem mintaoldal (a színeit a többi adja)
            except Exception as e:
                log.info("arculat-minta %s: %s", fid, e)
            png = temp_path(new_file_id("png"))
            try:
                found = await asyncio.to_thread(_render, src, png)
                for f in found:
                    fonts[f] += 3
                pngs.append(png)
                tmp.append(png)
            except Exception as e:
                log.info("arculat-minta képe %s: %s", fid, e)
        measured: list[tuple[str, float]] = []
        for p in pngs:
            for c, share in await asyncio.to_thread(_colors_of, p):
                measured.append((c, share / max(1, len(pngs))))
        measured.sort(key=lambda t: -t[1])
        if logo and not pngs:                              # csak logó: a színei
            lp = temp_path(new_file_id("png"))
            lp.write_bytes(await asyncio.to_thread(_clean_logo, logo))
            tmp.append(lp)
            measured = [(c, s) for c, s in await asyncio.to_thread(_colors_of, lp)]
        roles = _roles(measured, exact)
        declared = [f for f, _ in fonts.most_common(6)]
        vis: dict[str, Any] = {}
        if pngs:
            if progress:
                progress(len(file_ids), len(file_ids) + 1, "looking at the examples")
            try:
                vis = await _vision(pngs, [c for c, _ in measured] + exact, declared + [v for v in theme_fonts.values() if v])
            except Exception as e:
                log.warning("arculat: a látómodell nem válaszolt (%s) — a mért adatokból", e)
        vc = vis.get("colors") if isinstance(vis.get("colors"), dict) else {}
        colors = {k: (str(vc[k]).upper() if _HEX.match(str(vc.get(k, ""))) else roles[k]) for k in ("primary", "accent", "bg", "text")}
        def far(c, others, d=48):
            from app.tools.design_auto import _rgb
            return all(sum((p - q) ** 2 for p, q in zip(_rgb(c), _rgb(o))) ** 0.5 > d for o in others)
        extra: list[str] = []
        for c in [str(c).upper() for c in (vc.get("extra") or []) if _HEX.match(str(c))] + roles["extra"]:
            # a k-means élsimításból jövő szürke keverékek nem arculati színek: csak a valóban színes extrák maradnak
            if far(c, [colors[k] for k in ("primary", "accent", "bg", "text")] + extra) and design_auto._sat(c) > 0.25:
                extra.append(c)
        colors["extra"] = extra[:4]
        colors["text"] = design_auto.readable(colors["text"], colors["bg"])
        vf = vis.get("fonts") if isinstance(vis.get("fonts"), dict) else {}
        h_decl = theme_fonts.get("heading") or vf.get("heading") or (declared[0] if declared else None)
        b_decl = theme_fonts.get("body") or vf.get("body") or (declared[1] if len(declared) > 1 else h_decl)
        font_spec = {"heading": design_auto.font_family(h_decl) if h_decl else "Noto Sans Display",
                     "body": design_auto.font_family(b_decl, body=True) if b_decl else "Source Sans 3",
                     "heading_declared": h_decl, "body_declared": b_decl}
        # logó: a fájlból kiolvasott jelölt, különben a modell által megjelölt dobozból kivágva
        if not logo and isinstance(vis.get("logo"), dict) and pngs:
            try:
                idx = max(0, min(len(pngs) - 1, int(vis["logo"].get("image") or 0)))
                x0, y0, x1, y1 = [float(v) for v in vis["logo"]["box"]]
                if 0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1 and (x1 - x0) * (y1 - y0) < 0.25:
                    from PIL import Image
                    import io
                    with Image.open(pngs[idx]) as im:
                        W, H = im.size
                        pad = 0.01
                        box = (int(max(0, x0 - pad) * W), int(max(0, y0 - pad) * H), int(min(1, x1 + pad) * W), int(min(1, y1 + pad) * H))
                        buf = io.BytesIO()
                        im.convert("RGB").crop(box).save(buf, "PNG")
                        logo = buf.getvalue()
            except Exception as e:
                log.info("logó-kivágás: %s", e)
        bname = (name or vis.get("name") or "").strip() or "My brand"
        logo_fid = None
        if logo:
            logo_fid = new_file_id("png")
            output_path(logo_fid).write_bytes(await asyncio.to_thread(_clean_logo, logo))
            write_meta(logo_fid, label=f"{bname[:60]} logo.png", operation="brand logo")
        vr = vis.get("rules") if isinstance(vis.get("rules"), dict) else {}
        rules = {k: vr[k] for k in RULE_KEYS if vr.get(k)}
        if isinstance(rules.get("layouts"), list):
            rules["layouts"] = [a for a in rules["layouts"] if a in design_auto.ARCHETYPES][:3]
        return {"name": bname[:80], "colors": colors, "fonts": font_spec, "logo_file_id": logo_fid, "rules": rules,
                "sources": list(file_ids[:8])}
    finally:
        for p in tmp:
            p.unlink(missing_ok=True)
