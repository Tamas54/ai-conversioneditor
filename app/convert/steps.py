"""A konverziós gráf éleinek végrehajtása — motoronként egy függvény.

Minden lépés: `async def step(src: Path, fmt: str, ctx: Ctx) -> list[Path]` — egy bemenetből egy vagy több
kimenet (pl. egy többoldalas PDF → oldalanként egy JPG). A következő lépés minden kimenetre lefut.
A lépések csak validált opciókat kapnak (options.validate), a parancssorok argumentumlisták (nincs shell).
"""
from __future__ import annotations

import asyncio
import bz2
import csv
import gzip
import io
import json
import lzma
import os
import re
import shutil
import tarfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Awaitable, Callable, Optional

from app.convert import advanced
from app.convert.engines import EngineError, available, run
from app.convert.options import OptionError, parse_pages

# a formátumkulcs ≠ fájlkiterjesztés
EXT_FILE = {"alac": "m4a", "pdfa": "pdf"}
VIDEO_TIMEOUT = float(os.getenv("VIDEO_TIMEOUT_S", "1800"))
MAX_EXTRACT_BYTES = int(os.getenv("ARCHIVE_MAX_EXTRACT_MB", "4096")) * 1024 * 1024
MAX_EXTRACT_FILES = int(os.getenv("ARCHIVE_MAX_FILES", "20000"))


def file_ext(fmt: str) -> str:
    return EXT_FILE.get(fmt, fmt)


@dataclass
class Ctx:
    work: Path
    opts: dict[str, Any]
    final: bool = True
    source_fmt: str = ""
    resolve: Optional[Callable[[str], Path]] = None     # file_id → útvonal (borító, felirat, vízjel-kép)
    _n: int = 0
    notes: list[str] = field(default_factory=list)

    def out(self, src: Path, fmt: str, suffix: str = "") -> Path:
        self._n += 1
        stem = re.sub(r"[^\w\-]+", "_", _stem(src))[:60] or "file"
        d = self.work / f"s{self._n}"
        d.mkdir(parents=True, exist_ok=True)
        return d / f"{stem}{suffix}.{file_ext(fmt)}"

    def sub(self, name: str) -> Path:
        self._n += 1
        d = self.work / f"{name}{self._n}"
        d.mkdir(parents=True, exist_ok=True)
        return d

    def file(self, key: str) -> Optional[Path]:
        fid = self.opts.get(key)
        if not fid:
            return None
        if self.resolve is None:
            raise OptionError(f"{key}: files can't be referenced here")
        return self.resolve(fid)


def _stem(p: Path) -> str:
    name = p.name
    for multi in ("tar.gz", "tar.bz2", "tar.bz", "tar.xz", "tar.zst", "tar.7z", "tar.lzo", "tar.z"):
        if name.lower().endswith("." + multi):
            return name[: -len(multi) - 1]
    return p.stem


def _fmt_of(p: Path) -> str:
    from app.catalog.formats import from_filename
    return from_filename(p.name)


def _ok(paths: list[Path]) -> list[Path]:
    paths = [p for p in paths if p.exists() and p.stat().st_size > 0]
    if not paths:
        raise EngineError("the conversion produced an empty file")
    return paths


# ======================================================================
# LibreOffice
# ======================================================================

_LO_FILTER = {
    "writer": {"pdf": "writer_pdf_Export", "docx": "MS Word 2007 XML", "doc": "MS Word 97", "odt": "writer8",
               "ott": "writer8_template", "fodt": "OpenDocument Text Flat XML", "rtf": "Rich Text Format",
               "txt": "Text (encoded):UTF8", "html": "HTML (StarWriter)", "dotx": "MS Word 2007 XML Template",
               "dot": "MS Word 97 Vorlage", "docm": "MS Word 2007 XML VBA", "epub": "EPUB"},
    "calc": {"xlsx": "Calc MS Excel 2007 XML", "xls": "MS Excel 97", "xltx": "Calc MS Excel 2007 XML Template",
             "xlt": "MS Excel 97 Vorlage/Template", "xlsm": "Calc MS Excel 2007 VBA XML", "ods": "calc8",
             "fods": "OpenDocument Spreadsheet Flat XML", "csv": "Text - txt - csv (StarCalc):44,34,76,1",
             "pdf": "calc_pdf_Export", "html": "HTML (StarCalc)", "dif": "DIF:76", "slk": "SYLK"},
    "impress": {"pptx": "Impress MS PowerPoint 2007 XML", "ppt": "MS PowerPoint 97", "pptm": "Impress MS PowerPoint 2007 XML VBA",
                "pps": "MS PowerPoint 97 AutoPlay", "ppsx": "Impress MS PowerPoint 2007 XML AutoPlay",
                "pot": "MS PowerPoint 97 Vorlage", "potx": "Impress MS PowerPoint 2007 XML Template", "odp": "impress8",
                "fodp": "OpenDocument Presentation Flat XML", "pdf": "impress_pdf_Export", "html": "impress_html_Export"},
    "draw": {"pdf": "draw_pdf_Export", "odg": "draw8", "fodg": "OpenDocument Drawing Flat XML", "svg": "draw_svg_Export",
             "png": "draw_png_Export", "emf": "draw_emf_Export", "wmf": "draw_wmf_Export"},
}


def _lo_class(fmt: str) -> str:
    from app.convert import graph
    if fmt in graph.LO_CALC_IN:
        return "calc"
    if fmt in graph.LO_IMPRESS_IN:
        return "impress"
    if fmt in graph.LO_DRAW_IN:
        return "draw"
    return "writer"


def _sniff_delimiter(p: Path) -> str:
    head = p.read_bytes()[:65536].decode("utf-8-sig", "replace")
    try:
        return csv.Sniffer().sniff(head, delimiters=",;\t|").delimiter
    except csv.Error:
        return "\t" if p.suffix.lower() == ".tsv" else ","


async def wordcraft(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    """WordCraft (Rust Word-motor, a meglévő CLI): docx/odt/rtf/html/md/tex/txt → pdf/odt/rtf/html/md/tex/txt, md → docx."""
    from app.tools import wordcraft as wc
    dst = ctx.out(src, fmt)
    try:
        await wc.convert_file(src, dst)
    except RuntimeError as e:
        raise EngineError(str(e))
    if not dst.exists() or dst.stat().st_size == 0:
        raise EngineError("WordCraft produced no output")
    return [dst]


async def _craft(app: str, *args: str) -> None:
    from app.tools import crafts
    cli = {"deck": lambda: crafts.office_cli("deck"), "sheet": lambda: crafts.office_cli("sheet"),
           "photo": crafts.photo_cli, "vector": crafts.vector_cli}[app]()
    try:
        await crafts.run_cli(cli, f"{app}craft-cli", *args)
    except Exception as e:                       # CliError / RuntimeError → a gráf motorhibája
        raise EngineError(str(e)[:500])


async def deckcraft(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    """DeckCraft (Rust PowerPoint-motor): pptx → pdf (convert) | png diánként (render --all)."""
    if fmt == "pdf":
        dst = ctx.out(src, "pdf")
        await _craft("deck", "convert", str(src), str(dst))
        return _ok([dst])
    work = ctx.sub("deck")
    scale = max(0.5, min(4.0, float(ctx.opts.get("dpi") or 144) / 96))
    await _craft("deck", "render", str(src), "--all", "--scale", f"{scale:.2f}", str(work / "slides.png"))
    pages = sorted(p for p in work.rglob("*.png") if p.is_file())       # a kimenet egy „slides.png” mappa
    if not pages:
        raise EngineError("DeckCraft rendered no slides")
    stem = re.sub(r"[^\w\-]+", "_", _stem(src))[:60] or "slide"
    out = []
    for i, p in enumerate(pages, 1):
        dst = ctx.out(src, "png", suffix=f"-{i:03d}" if len(pages) > 1 else "")
        shutil.move(str(p), dst)
        out.append(dst)
    return _ok(out)


async def gridcraft(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    """GridCraft (Rust táblázatmotor): xlsx/csv/tsv → xlsx/csv/tsv/html (a CSV elválasztóját maga ismeri fel)."""
    dst = ctx.out(src, fmt)
    await _craft("sheet", "convert", str(src), str(dst))
    return _ok([dst])


async def photocraft(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    """PhotoCraft (Rust képmotor): webp/gif cél, rétegzett PSD."""
    dst = ctx.out(src, fmt)
    await _craft("photo", "convert", str(src), str(dst))
    return _ok([dst])


async def vectorcraft(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    """VectorCraft (Rust vektormotor): svg/svgz/dxf → svg/pdf/png/jpg/webp/eps/dxf/tiff."""
    dst = ctx.out(src, fmt)
    await _craft("vector", "convert", str(src), str(dst))
    return _ok([dst])


async def libreoffice(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    sfmt = _fmt_of(src)
    cls = _lo_class(sfmt)
    filt = _LO_FILTER[cls].get(fmt)
    if not filt:
        raise EngineError(f"LibreOffice can't write .{fmt} from this file")
    work = ctx.sub("lo")
    inp = work / f"in.{src.suffix.lstrip('.') or sfmt}"
    shutil.copy(src, inp)
    args = ["soffice", f"-env:UserInstallation=file://{work}/profile", "--headless", "--norestore"]
    if sfmt in ("csv", "tsv"):
        d = ctx.opts.get("delimiter", "auto")
        d = {"tab": "\t", "auto": _sniff_delimiter(src)}.get(d, d)
        args.append(f"--infilter=CSV:{ord(d)},34,76,1")
    if cls == "calc" and fmt == "csv" and ctx.opts.get("delimiter") not in (None, "auto"):
        d = {"tab": "\t"}.get(ctx.opts["delimiter"], ctx.opts["delimiter"])
        filt = f"Text - txt - csv (StarCalc):{ord(d)},34,76,1"
    if sfmt in ("html", "htm", "xhtml") and fmt == "ott":
        args.append("--infilter=HTML (StarWriter)")
    out_ext = fmt
    args += ["--convert-to", f"{out_ext}:{filt}", "--outdir", str(work / "out"), str(inp)]
    await run(*args, timeout=600)
    outs = sorted((work / "out").glob("*"))
    if not outs:
        raise EngineError("LibreOffice produced nothing (the file may be damaged or password-protected)")
    dst = ctx.out(src, fmt)
    shutil.move(str(outs[0]), dst)
    return _ok([dst])


# ======================================================================
# Pandoc, WeasyPrint, pdf2docx
# ======================================================================

_PANDOC_R = {"md": "markdown", "rst": "rst", "tex": "latex", "html": "html", "docx": "docx", "odt": "odt",
             "epub": "epub", "fb2": "fb2", "org": "org", "textile": "textile", "wiki": "mediawiki", "ipynb": "ipynb"}
_PANDOC_W = {"md": "gfm", "rst": "rst", "tex": "latex", "html": "html5", "docx": "docx", "odt": "odt", "epub": "epub3",
             "fb2": "fb2", "org": "org", "adoc": "asciidoc", "textile": "textile", "wiki": "mediawiki",
             "ipynb": "ipynb", "txt": "plain", "rtf": "rtf"}


async def pandoc(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    sfmt = _fmt_of(src)
    dst = ctx.out(src, fmt)
    args = ["pandoc", "-f", _PANDOC_R[sfmt], "-t", _PANDOC_W[fmt], "-o", str(dst)]
    if fmt in ("html", "tex", "rtf", "epub", "fb2"):
        args.append("-s")
    if fmt == "html":
        args.append("--embed-resources")
    # a cím csak metaadat legyen, ne kerüljön új címsorként a szövegbe (a fájlnév nem a dokumentum része)
    if ctx.opts.get("title") or fmt in ("epub", "fb2"):
        args += ["--metadata", f"title={ctx.opts.get('title') or _stem(src)}"]
    elif fmt == "html":
        args += ["--metadata", f"pagetitle={_stem(src)}"]
    if ctx.opts.get("author"):
        args += ["--metadata", f"author={ctx.opts['author']}"]
    media = ctx.sub("media")
    if sfmt in ("docx", "odt", "epub"):
        args += ["--extract-media", str(media)]
    args.append(str(src))
    await run(*args, timeout=300, cwd=src.parent)
    if fmt == "docx":                           # a pandoc üres táblarácsot ír: Wordben csak az első oszlop látszana
        from app.tools import ooxml_safe
        await asyncio.to_thread(ooxml_safe.fix, dst)
    return _ok([dst])


_PRINT_CSS = "@page { size: A4; margin: 18mm } body { font-family: 'DejaVu Sans', 'Liberation Sans', sans-serif; }"


async def weasyprint(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    dst = ctx.out(src, "pdf")

    def go():
        from weasyprint import CSS, HTML
        HTML(filename=str(src), base_url=str(src.parent)).write_pdf(str(dst), stylesheets=[CSS(string=_PRINT_CSS)])
    await asyncio.to_thread(go)
    return _ok([dst])


async def pdf2docx(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    dst = ctx.out(src, "docx")

    def go():
        from pdf2docx import Converter
        cv = Converter(str(src))
        try:
            cv.convert(str(dst))
        finally:
            cv.close()
    await asyncio.to_thread(go)
    return _ok([dst])


# ======================================================================
# PDF: Poppler, MuPDF, Ghostscript, DjVuLibre
# ======================================================================

def _pdf_pages(src: Path) -> int:
    import pymupdf
    with pymupdf.open(str(src)) as d:
        return d.page_count


async def poppler(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    if fmt == "txt":
        dst = ctx.out(src, "txt")
        await run("pdftotext", "-layout", "-enc", "UTF-8", str(src), str(dst))
        return _ok([dst])
    work = ctx.sub("pg")
    dpi = str(ctx.opts.get("dpi") or 150)
    flag = {"png": "-png", "jpg": "-jpeg", "tiff": "-tiff"}[fmt]
    args = ["pdftoppm", "-r", dpi, flag]
    if fmt == "jpg":
        args += ["-jpegopt", f"quality={ctx.opts.get('quality') or 88}"]
    await run(*args, str(src), str(work / "p"), timeout=900)
    pages = sorted(work.glob("p-*"), key=lambda p: int(re.findall(r"(\d+)", p.stem)[-1]))
    outs = []
    for i, p in enumerate(pages, 1):
        dst = ctx.out(src, fmt, f"-{i:03d}" if len(pages) > 1 else "")
        shutil.move(str(p), dst)
        outs.append(dst)
    return _ok(outs)


async def mupdf(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    sfmt = _fmt_of(src)
    if sfmt == "xps":
        dst = ctx.out(src, "pdf")
        await run("mutool", "convert", "-o", str(dst), str(src))
        return _ok([dst])
    work = ctx.sub("mu")
    if fmt == "svg":
        await run("mutool", "convert", "-F", "svg", "-o", str(work / "p%d.svg"), str(src), timeout=600)
        pages = sorted(work.glob("p*.svg"), key=lambda p: int(re.findall(r"\d+", p.stem)[0]))
        outs = []
        for i, p in enumerate(pages, 1):
            dst = ctx.out(src, "svg", f"-{i:03d}" if len(pages) > 1 else "")
            shutil.move(str(p), dst)
            outs.append(dst)
        return _ok(outs)
    dst = ctx.out(src, "html")
    await run("mutool", "convert", "-F", "html", "-o", str(dst), str(src), timeout=600)
    return _ok([dst])


def _icc() -> Optional[Path]:
    for p in sorted(Path("/usr/share/ghostscript").glob("*/iccprofiles/default_rgb.icc"), reverse=True):
        return p
    return None


async def ghostscript(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    sfmt = _fmt_of(src)
    dst = ctx.out(src, fmt)
    base = ["gs", "-dBATCH", "-dNOPAUSE", "-dSAFER", "-dQUIET"]
    if fmt == "pdfa":
        level = str(ctx.opts.get("pdfa_level", "2b"))[0]
        icc = _icc()
        work = ctx.sub("gs")
        defs = work / "PDFA_def.ps"
        if icc:
            defs.write_text(
                "%!\n[/_objdef {icc_PDFA} /type /stream /OBJ pdfmark\n"
                f"[{{icc_PDFA}} <</N 3>> /PUT pdfmark\n[{{icc_PDFA}} ({icc}) (r) file /PUT pdfmark\n"
                "[/_objdef {OutputIntent_PDFA} /type /dict /OBJ pdfmark\n"
                "[{OutputIntent_PDFA} <</Type /OutputIntent /S /GTS_PDFA1 /DestOutputProfile {icc_PDFA} "
                "/OutputConditionIdentifier (sRGB)>> /PUT pdfmark\n"
                "[{Catalog} <</OutputIntents [ {OutputIntent_PDFA} ]>> /PUT pdfmark\n")
        args = base + [f"-dPDFA={level}", "-dPDFACompatibilityPolicy=1", "-sColorConversionStrategy=RGB",
                       "-sProcessColorModel=DeviceRGB", "-sDEVICE=pdfwrite", "-dNOOUTERSAVE"]
        if icc:
            args += [f"--permit-file-read={icc}", f"-sOutputICCProfile={icc}"]
        args += [f"-sOutputFile={dst}"] + ([str(defs)] if icc else []) + [str(src)]
        await run(*args, timeout=900)
    elif fmt == "pdf":
        await run(*base, "-sDEVICE=pdfwrite", "-dEPSCrop", f"-sOutputFile={dst}", str(src), timeout=900)
    elif fmt == "ps":
        await run(*base, "-sDEVICE=ps2write", f"-sOutputFile={dst}", str(src), timeout=900)
    elif fmt == "eps":
        await run(*base, "-sDEVICE=eps2write", "-dFirstPage=1", "-dLastPage=1", f"-sOutputFile={dst}", str(src), timeout=600)
    else:
        raise EngineError(f"Ghostscript can't write .{fmt}")
    return _ok([dst])


async def djvulibre(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    dst = ctx.out(src, fmt)
    if fmt == "txt":
        await run("djvutxt", str(src), str(dst), timeout=600)
    else:
        await run("ddjvu", f"-format={'pdf' if fmt == 'pdf' else 'tiff'}", "-quality=85", str(src), str(dst), timeout=900)
    return _ok([dst])


# ======================================================================
# Calibre (e-könyvek)
# ======================================================================

async def calibre(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    real = {"azw": "mobi", "prc": "mobi"}.get(fmt, fmt)
    work = ctx.sub("cal")
    inp = src
    if _fmt_of(src) == "md":                      # a Calibre a .md-t txt+markdown-ként olvassa
        inp = work / f"{_stem(src)}.txt"
        shutil.copy(src, inp)
    tmp = work / f"{_stem(src)}.{real}"
    args = ["ebook-convert", str(inp), str(tmp)]
    if _fmt_of(src) in ("md", "txt"):
        args += ["--formatting-type", "markdown" if _fmt_of(src) == "md" else "auto"]
    if ctx.opts.get("title"):
        args += ["--title", ctx.opts["title"]]
    if ctx.opts.get("author"):
        args += ["--authors", ctx.opts["author"]]
    cover = ctx.file("cover")
    if cover:
        args += ["--cover", str(cover)]
    prof = ctx.opts.get("output_profile")
    if prof and prof != "default":
        args += ["--output-profile", prof]
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    os.environ.setdefault("QTWEBENGINE_CHROMIUM_FLAGS", "--no-sandbox")
    await run(*args, timeout=1200, cwd=work)
    dst = ctx.out(src, fmt)
    shutil.move(str(tmp), dst)
    return _ok([dst])


# ======================================================================
# Képek: Pillow, ImageMagick, libheif, JPEG XL, LibRaw, Potrace
# ======================================================================

def _register_plugins() -> None:
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except Exception:
        pass
    try:
        import pillow_jxl  # noqa: F401
    except Exception:
        pass


def _apply_image_opts(im, opts: dict[str, Any], fmt: str):
    from PIL import Image, ImageOps
    w, h = opts.get("width"), opts.get("height")
    if w or h:
        fit = opts.get("fit", "inside")
        W, H = w or round(im.width * h / im.height), h or round(im.height * w / im.width)
        if fit == "fill":
            im = im.resize((W, H), Image.LANCZOS)
        elif fit == "cover":
            im = ImageOps.fit(im, (W, H), Image.LANCZOS)
        elif fit == "contain":
            bg = opts.get("background") or ("#00000000" if im.mode in ("RGBA", "LA") else "#ffffff")
            im = ImageOps.pad(im.convert("RGBA") if "A" in im.mode else im, (W, H), Image.LANCZOS, color=_rgba(bg))
        else:                                       # inside: a dobozba illesztve (nagyítva is)
            r = min(W / im.width, H / im.height)
            im = im.resize((max(1, round(im.width * r)), max(1, round(im.height * r))), Image.LANCZOS)
    if opts.get("rotate"):
        im = im.rotate(-int(opts["rotate"]), expand=True)
    if opts.get("flip") == "horizontal":
        im = ImageOps.mirror(im)
    elif opts.get("flip") == "vertical":
        im = ImageOps.flip(im)
    if opts.get("grayscale"):
        im = ImageOps.grayscale(im) if "A" not in im.mode else im.convert("LA")
    return im


def _rgba(color: str) -> tuple:
    from PIL import ImageColor
    if color == "transparent":
        return (0, 0, 0, 0)
    c = ImageColor.getrgb(color)
    return c if len(c) == 4 else (*c, 255)


def _flatten(im, bg: Optional[str]):
    from PIL import Image
    if im.mode in ("RGBA", "LA", "P", "PA"):
        im = im.convert("RGBA")
        base = Image.new("RGBA", im.size, _rgba(bg or "#ffffff"))
        base.alpha_composite(im)
        return base.convert("RGB")
    return im.convert("RGB") if im.mode not in ("RGB", "L") else im


_PIL_NAME = {"jpg": "JPEG", "jfif": "JPEG", "png": "PNG", "apng": "PNG", "gif": "GIF", "bmp": "BMP", "tiff": "TIFF",
             "webp": "WEBP", "avif": "AVIF", "ico": "ICO", "icns": "ICNS", "tga": "TGA", "jp2": "JPEG2000",
             "pcx": "PCX", "ppm": "PPM", "pgm": "PPM", "pbm": "PPM", "pnm": "PPM", "sgi": "SGI", "xbm": "XBM",
             "dds": "DDS", "pdf": "PDF", "eps": "EPS", "heic": "HEIF", "heif": "HEIF", "jxl": "JXL"}


def _pil_save(im, dst: Path, fmt: str, opts: dict[str, Any], frames: Optional[list] = None, info: Optional[dict] = None) -> None:
    from PIL import Image
    q = opts.get("quality")
    kw: dict[str, Any] = {}
    keep_meta = not opts.get("strip_metadata")
    if opts.get("dpi"):
        kw["dpi"] = (opts["dpi"], opts["dpi"])
    if fmt in ("jpg", "jfif"):
        im = _flatten(im, opts.get("background"))
        kw.update(quality=q or 88, optimize=True, progressive=True)
        if keep_meta and info and info.get("exif"):
            kw["exif"] = info["exif"]
        if keep_meta and info and info.get("icc_profile"):
            kw["icc_profile"] = info["icc_profile"]
    elif fmt == "webp":
        kw.update(quality=q or 85, method=4, lossless=bool(opts.get("lossless")))
        if keep_meta and info and info.get("exif"):
            kw["exif"] = info["exif"]
    elif fmt == "avif":
        kw.update(quality=q or 70)
    elif fmt in ("heic", "heif"):
        kw.update(quality=q or 80)
        im = im.convert("RGBA") if "A" in im.mode else im.convert("RGB")
    elif fmt == "jxl":
        kw.update(quality=q or 85, lossless=bool(opts.get("lossless")))
    elif fmt in ("png", "apng"):
        kw.update(optimize=True)
        if keep_meta and info and info.get("icc_profile"):
            kw["icc_profile"] = info["icc_profile"]
    elif fmt == "tiff":
        kw.update(compression="tiff_lzw")
    elif fmt == "gif":
        im = im.convert("RGBA").convert("P", palette=Image.ADAPTIVE) if im.mode not in ("P", "L") else im
    elif fmt == "ico":
        im = im.convert("RGBA")
        kw["sizes"] = [s for s in [(16, 16), (24, 24), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)]
                       if s[0] <= max(im.size)] or [(16, 16)]
    elif fmt == "icns":
        im = im.convert("RGBA")
        side = max(im.size)
        if side < 1024 and (im.width != im.height):
            from PIL import ImageOps
            im = ImageOps.pad(im, (side, side), color=(0, 0, 0, 0))
    elif fmt in ("pbm", "xbm"):
        im = im.convert("L").point(lambda v: 255 if v >= opts.get("threshold", 128) else 0, "1")
    elif fmt == "pgm":
        im = im.convert("L")
    elif fmt in ("ppm", "pnm", "pcx", "sgi", "eps", "jp2"):
        im = _flatten(im, opts.get("background")) if fmt in ("eps", "ppm", "pnm", "pcx") else im
        if fmt in ("pcx", "sgi", "jp2") and im.mode not in ("RGB", "RGBA", "L"):
            im = im.convert("RGB")
    elif fmt == "pdf":
        im = _flatten(im, opts.get("background"))
        kw.setdefault("resolution", float(opts.get("dpi") or 150))
    elif fmt == "tga":
        im = im.convert("RGBA") if "A" in im.mode else im.convert("RGB")
    elif fmt == "dds":
        im = im.convert("RGBA")
    if frames and fmt in ("gif", "webp", "apng", "png") and len(frames) > 1:
        kw.update(save_all=True, append_images=frames[1:], loop=0 if opts.get("loop", True) else 1,
                  duration=(info or {}).get("duration", 100))
        im = frames[0]
    im.save(dst, _PIL_NAME[fmt], **kw)


async def pillow(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    dst = ctx.out(src, fmt)

    def go():
        from PIL import Image, ImageOps, ImageSequence
        _register_plugins()
        Image.MAX_IMAGE_PIXELS = 400_000_000
        with Image.open(src) as im0:
            info = dict(im0.info)
            animated = getattr(im0, "is_animated", False) and fmt in ("gif", "webp", "apng") and ctx.final
            if animated:
                frames = [_apply_image_opts(ImageOps.exif_transpose(f.convert("RGBA")), ctx.opts, fmt)
                          for f in ImageSequence.Iterator(im0)]
                _pil_save(frames[0], dst, fmt, ctx.opts, frames, info)
                return
            if getattr(im0, "n_frames", 1) > 1:
                im0.seek(min(int(ctx.opts.get("frame") or 0), im0.n_frames - 1))
            im = ImageOps.exif_transpose(im0)
            im.load()
            if ctx.final:
                im = _apply_image_opts(im, ctx.opts, fmt)
            _pil_save(im, dst, fmt, ctx.opts if ctx.final else {}, None, info)
    await asyncio.to_thread(go)
    return _ok([dst])


def _im() -> str:
    return "magick" if shutil.which("magick") else "convert"


async def imagemagick(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    dst = ctx.out(src, fmt)
    args = [_im(), f"{src}[{int(ctx.opts.get('frame') or 0)}]", "-auto-orient"]
    o = ctx.opts if ctx.final else {}
    if o.get("width") or o.get("height"):
        args += ["-resize", f"{o.get('width') or ''}x{o.get('height') or ''}"]
    if o.get("rotate"):
        args += ["-rotate", str(o["rotate"])]
    if o.get("flip") == "horizontal":
        args.append("-flop")
    elif o.get("flip") == "vertical":
        args.append("-flip")
    if o.get("grayscale"):
        args += ["-colorspace", "Gray"]
    if o.get("quality"):
        args += ["-quality", str(o["quality"])]
    if o.get("strip_metadata"):
        args.append("-strip")
    if o.get("dpi"):
        args += ["-density", str(o["dpi"])]
    if fmt in ("cur", "ico"):
        args += ["-resize", "256x256>"]
    if fmt in ("jpg", "pcx", "ppm") or o.get("background"):
        args += ["-background", o.get("background") or "white", "-flatten"]
    if o.get("advanced"):
        args += advanced.check(o["advanced"], "imagemagick")
    args.append(f"{ {'jpg': 'jpeg'}.get(fmt, fmt) }:{dst}")
    await run(*args, timeout=600)
    if fmt == "cur":
        _ico_as_cur(dst)
    return _ok([dst])


def _ico_as_cur(p: Path) -> None:
    """Az ImageMagick a .cur-t ikonként (ICO, típus 1) írja — a kurzorhoz típus 2 és forrópont kell (bal felső sarok);
    a képadat ugyanaz."""
    b = bytearray(p.read_bytes())
    if len(b) < 6 or b[:4] != b"\0\0\1\0":
        return
    b[2] = 2
    for k in range(int.from_bytes(b[4:6], "little")):
        e = 6 + 16 * k
        if e + 16 > len(b):
            break
        b[e + 4:e + 8] = b"\0\0\0\0"          # színsíkok/bitmélység helyén: forrópont x, y
    p.write_bytes(bytes(b))


async def heif(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    return await pillow(src, fmt, ctx)


async def jxl(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    return await pillow(src, fmt, ctx)


async def libraw(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    work = ctx.sub("raw")
    inp = work / src.name
    shutil.copy(src, inp)
    out, _ = await run("dcraw_emu", "-T", "-w", "-o", "1", "-q", "3", "-Z", "-", str(inp), timeout=600)
    dst = ctx.out(src, "tiff")
    dst.write_bytes(out)
    return _ok([dst])


async def potrace(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    work = ctx.sub("trace")
    bmp = work / "in.bmp"
    thr = int(ctx.opts.get("threshold") or 128)

    def prep():
        from PIL import Image, ImageOps
        _register_plugins()
        with Image.open(src) as im:
            g = _flatten(ImageOps.exif_transpose(im), "#ffffff").convert("L")
            g.point(lambda v: 255 if v >= thr else 0, "1").save(bmp)
    await asyncio.to_thread(prep)
    dst = ctx.out(src, fmt)
    await run("potrace", "-b", "svg" if fmt == "svg" else "eps", "-o", str(dst), str(bmp), timeout=300)
    return _ok([dst])


# ======================================================================
# Vektor: Inkscape, rsvg
# ======================================================================

_INK_TYPE = {"ai": "pdf", "plt": "hpgl"}


async def inkscape(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    work = ctx.sub("ink")
    kind = _INK_TYPE.get(fmt, fmt)
    tmp = work / f"out.{kind}"
    inp = work / f"in.{src.suffix.lstrip('.')}"
    shutil.copy(src, inp)
    args = ["inkscape", str(inp), f"--export-type={kind}", f"--export-filename={tmp}"]
    o = ctx.opts if ctx.final else {}
    if kind == "png":
        if o.get("width"):
            args.append(f"--export-width={o['width']}")
        elif o.get("height"):
            args.append(f"--export-height={o['height']}")
        else:
            args.append(f"--export-dpi={o.get('dpi') or 150}")
        if o.get("background"):
            args += [f"--export-background={o['background']}", "--export-background-opacity=1"]
    if o.get("text_to_path") and kind in ("pdf", "eps", "ps", "svg"):
        args.append("--export-text-to-path")
    if _fmt_of(src) in ("pdf", "ai"):
        args.append("--pdf-poppler")
    os.environ.setdefault("HOME", str(work))
    await run(*args, timeout=600, cwd=work)
    dst = ctx.out(src, fmt)
    shutil.move(str(tmp), dst)
    return _ok([dst])


async def rsvg(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    dst = ctx.out(src, fmt)
    args = ["rsvg-convert", "-f", fmt, "-o", str(dst)]
    o = ctx.opts if ctx.final else {}
    if fmt == "png":
        if o.get("width"):
            args += ["-w", str(o["width"]), "-a"]
        elif o.get("height"):
            args += ["-h", str(o["height"]), "-a"]
        if o.get("dpi"):
            args += ["-d", str(o["dpi"]), "-p", str(o["dpi"])]
        if o.get("background"):
            args += ["-b", o["background"]]
    args.append(str(src))
    await run(*args, timeout=300)
    return _ok([dst])


# ======================================================================
# FFmpeg (videó, hang, animált kép)
# ======================================================================

_V_CODEC = {"h264": ["-c:v", "libx264"], "h265": ["-c:v", "libx265", "-tag:v", "hvc1"], "vp9": ["-c:v", "libvpx-vp9"],
            "av1": ["-c:v", "libaom-av1", "-cpu-used", "6", "-row-mt", "1"], "mpeg4": ["-c:v", "mpeg4", "-q:v", "3"],
            "mpeg2": ["-c:v", "mpeg2video", "-q:v", "3"], "prores": ["-c:v", "prores_ks", "-profile:v", "2"],
            "copy": ["-c:v", "copy"]}
_A_CODEC = {"aac": ["-c:a", "aac"], "mp3": ["-c:a", "libmp3lame"], "opus": ["-c:a", "libopus"],
            "vorbis": ["-c:a", "libvorbis"], "flac": ["-c:a", "flac"], "pcm": ["-c:a", "pcm_s16le"],
            "ac3": ["-c:a", "ac3"], "copy": ["-c:a", "copy"], "none": ["-an"]}
# konténer → (alap videókodek, alap hangkodek, muxer, extra)
_VIDEO_DEF: dict[str, tuple[str, str, Optional[str], list[str]]] = {
    "mp4": ("h264", "aac", "mp4", ["-movflags", "+faststart"]), "m4v": ("h264", "aac", "mp4", ["-movflags", "+faststart"]),
    "mov": ("h264", "aac", "mov", []), "mkv": ("h264", "aac", "matroska", []), "webm": ("vp9", "opus", "webm", ["-row-mt", "1", "-b:v", "0"]),
    "avi": ("mpeg4", "mp3", "avi", []), "divx": ("mpeg4", "mp3", "avi", ["-vtag", "DIVX"]),
    "flv": ("h264", "aac", "flv", []), "f4v": ("h264", "aac", "mp4", []),
    "wmv": ("wmv2", "wmav2", "asf", []), "asf": ("wmv2", "wmav2", "asf", []), "wtv": ("mpeg2", "mp2", "wtv", []),
    "mpeg": ("mpeg2", "mp2", "mpeg", []), "mpg": ("mpeg2", "mp2", "mpeg", []), "vob": ("mpeg2", "ac3", "vob", []),
    "ts": ("h264", "aac", "mpegts", []), "mts": ("h264", "aac", "mpegts", ["-mpegts_m2ts_mode", "1"]),
    "m2ts": ("h264", "aac", "mpegts", ["-mpegts_m2ts_mode", "1"]), "3gp": ("h264", "aac", "3gp", ["-profile:v", "baseline"]),
    "3g2": ("h264", "aac", "3g2", ["-profile:v", "baseline"]), "ogv": ("theora", "vorbis", "ogg", ["-q:v", "7"]),
    "mxf": ("mpeg2", "pcm", "mxf", ["-r", "25", "-ar", "48000"]), "dv": ("dv", "pcm", "dv", []),
    "hevc": ("h265", "none", "hevc", []), "swf": ("flv1", "mp3", "swf", ["-ar", "44100"]),
    "mjpeg": ("mjpeg", "none", "mjpeg", ["-q:v", "3"]),
}
_SPECIAL_V = {"wmv2": ["-c:v", "wmv2", "-q:v", "4"], "theora": ["-c:v", "libtheora"], "dv": ["-c:v", "dvvideo"],
              "flv1": ["-c:v", "flv1", "-q:v", "4"], "mjpeg": ["-c:v", "mjpeg"]}
_SPECIAL_A = {"wmav2": ["-c:a", "wmav2"], "mp2": ["-c:a", "mp2"]}
# hang-konténer → (kodek-argumentumok, muxer, alap bitráta)
_AUDIO_DEF: dict[str, tuple[list[str], str, Optional[str]]] = {
    "mp3": (["-c:a", "libmp3lame"], "mp3", "192k"), "wav": (["-c:a", "pcm_s16le"], "wav", None),
    "aac": (["-c:a", "aac"], "adts", "192k"), "m4a": (["-c:a", "aac"], "ipod", "192k"),
    "ogg": (["-c:a", "libvorbis", "-q:a", "5"], "ogg", None), "flac": (["-c:a", "flac"], "flac", None),
    "opus": (["-c:a", "libopus"], "opus", "128k"), "wma": (["-c:a", "wmav2"], "asf", "192k"),
    "aiff": (["-c:a", "pcm_s16be"], "aiff", None), "alac": (["-c:a", "alac"], "ipod", None),
    "m4b": (["-c:a", "aac"], "ipod", "96k"), "m4r": (["-c:a", "aac"], "ipod", "192k"),
    "amr": (["-c:a", "libopencore_amrnb", "-ar", "8000", "-ac", "1"], "amr", "12.2k"),
    "ac3": (["-c:a", "ac3"], "ac3", "384k"), "dts": (["-c:a", "dca", "-strict", "-2"], "dts", "768k"),
    "wv": (["-c:a", "wavpack"], "wv", None), "mka": (["-c:a", "flac"], "matroska", None),
    "au": (["-c:a", "pcm_s16be"], "au", None), "caf": (["-c:a", "pcm_s16le"], "caf", None),
    "mp2": (["-c:a", "mp2"], "mp2", "192k"), "weba": (["-c:a", "libopus"], "webm", "128k"),
    "gsm": (["-c:a", "libgsm", "-ar", "8000", "-ac", "1"], "gsm", None), "voc": (["-c:a", "pcm_u8"], "voc", None),
}
_RES = {"4k": "3840:2160", "2160p": "3840:2160", "1440p": "2560:1440", "1080p": "1920:1080", "720p": "1280:720",
        "480p": "854:480", "360p": "640:360", "240p": "426:240"}


async def probe(src: Path) -> dict[str, Any]:
    out, _ = await run("ffprobe", "-v", "error", "-print_format", "json", "-show_format", "-show_streams", str(src), timeout=60)
    return json.loads(out or b"{}")


def _has_audio(info: dict) -> bool:
    return any(s.get("codec_type") == "audio" for s in info.get("streams", []))


def _has_video(info: dict) -> bool:
    return any(s.get("codec_type") == "video" and (s.get("disposition") or {}).get("attached_pic") != 1
               for s in info.get("streams", []))


def _trim_args(o: dict) -> list[str]:
    a = []
    if o.get("start"):
        a += ["-ss", str(o["start"])]
    if o.get("end"):
        a += ["-to", str(o["end"])]
    return a


def _vfilters(o: dict, fmt: str, sub: Optional[Path]) -> list[str]:
    vf = []
    res = o.get("resolution")
    if res:
        wh = _RES.get(res, res.replace("x", ":"))
        vf.append(f"scale={wh}:force_original_aspect_ratio=decrease,pad={wh}:(ow-iw)/2:(oh-ih)/2" if o.get("aspect_mode", "pad") == "pad"
                  else f"scale={wh}")
    elif o.get("width") or o.get("height"):
        vf.append(f"scale={o.get('width') or -2}:{o.get('height') or -2}")
    if o.get("aspect"):
        a, b = (int(x) for x in o["aspect"].split(":"))
        mode = o.get("aspect_mode", "pad")
        if mode == "crop":
            vf.append(f"crop='min(iw,ih*{a}/{b})':'min(ih,iw*{b}/{a})'")
        elif mode == "pad":
            vf.append(f"pad='max(iw,ih*{a}/{b})':'max(ih,iw*{b}/{a})':(ow-iw)/2:(oh-ih)/2")
        else:
            vf.append(f"scale='iw':'iw*{b}/{a}'")
        vf.append("setsar=1")
    rot = o.get("rotate")
    if rot == 90:
        vf.append("transpose=1")
    elif rot == 180:
        vf.append("transpose=1,transpose=1")
    elif rot == 270:
        vf.append("transpose=2")
    if o.get("flip") == "horizontal":
        vf.append("hflip")
    elif o.get("flip") == "vertical":
        vf.append("vflip")
    if o.get("fps") and fmt not in ("gif",):
        vf.append(f"fps={o['fps']}")
    if sub is not None:
        vf.append(f"subtitles={sub.name}")
    # páros méret (x264/x265 yuv420p) — a végén
    if fmt in ("mp4", "m4v", "mov", "mkv", "flv", "f4v", "ts", "mts", "m2ts", "3gp", "3g2", "hevc", "webm"):
        vf.append("scale=trunc(iw/2)*2:trunc(ih/2)*2")
    return vf


async def ffmpeg(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    from app.catalog.formats import FORMATS
    cat = FORMATS[fmt].category if fmt in FORMATS else ""
    sfmt = _fmt_of(src)
    scat = FORMATS[sfmt].category if sfmt in FORMATS else ""
    o = ctx.opts if ctx.final else {}
    dst = ctx.out(src, fmt)
    work = dst.parent
    if scat == "audio" and cat == "video":
        return await _audio_to_video(src, fmt, ctx, dst)
    if scat == "image" and cat == "image" and fmt not in ("gif", "apng", "webp"):      # pl. OpenEXR
        pix = {"exr": ["-pix_fmt", "gbrpf32le"]}.get(fmt, [])
        await run("ffmpeg", "-y", "-hide_banner", "-loglevel", "error", "-i", str(src), "-frames:v", "1", *pix, str(dst), timeout=300)
        return _ok([dst])
    if fmt in ("jpg", "png") and scat == "video":
        at = o.get("at") or "1"
        await run("ffmpeg", "-y", "-ss", str(at), "-i", str(src), "-frames:v", "1", *(["-q:v", "2"] if fmt == "jpg" else []),
                  str(dst), timeout=120)
        if not dst.exists():               # rövidebb, mint `at`
            await run("ffmpeg", "-y", "-i", str(src), "-frames:v", "1", str(dst), timeout=120)
        return _ok([dst])
    if fmt in ("gif", "apng", "webp") and scat in ("video", "image"):
        return await _animated(src, fmt, ctx, dst)
    info = await probe(src)
    args = ["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", *_trim_args(o), "-i", str(src)]
    sub = None
    subtitle_src = ctx.file("subtitles") if cat == "video" and o.get("subtitles") else None
    if subtitle_src is not None:
        sub = work / f"subs{subtitle_src.suffix.lower()}"
        shutil.copy(subtitle_src, sub)
        if o.get("subtitles_mode") == "soft":
            args += ["-i", str(sub)]
            sub = None
    if cat == "audio":
        codec, mux, br = _AUDIO_DEF[fmt]
        a_codec = o.get("audio_codec", "auto")
        args += ["-vn", "-map", "0:a:0?"] + (_A_CODEC[a_codec] if a_codec not in ("auto", None) else codec)
        if br and not o.get("audio_bitrate") and a_codec in ("auto", None):
            args += ["-b:a", br]
        args += _audio_opts(o, fmt)
        if fmt == "mp2" and not o.get("sample_rate"):
            args += ["-ar", "44100"]
        if fmt == "m4r" and not o.get("end"):
            args += ["-t", "40"]
        args += ["-map_metadata", "-1"] if o.get("strip_metadata") else []
        args += ["-f", mux, str(dst)]
        await run(*args, timeout=VIDEO_TIMEOUT)
        return _ok([dst])
    vdef, adef, mux, extra = _VIDEO_DEF[fmt]
    vc = o.get("video_codec", "auto")
    ac = o.get("audio_codec", "auto")
    vargs = _V_CODEC.get(vc) if vc not in ("auto", None) else (_V_CODEC.get(vdef) or _SPECIAL_V[vdef])
    if vdef == "dv":
        args += ["-target", "pal-dv"]
        vargs, extra = [], []
    vf = [] if vc == "copy" else _vfilters(o, fmt, sub)
    args += vargs
    if vf:
        args += ["-vf", ",".join(vf)]
    eff_v = vc if vc not in ("auto", None) else vdef
    if eff_v in ("h264", "h265", "vp9", "av1"):
        if o.get("crf") is not None:
            args += ["-crf", str(o["crf"])]
        elif not o.get("video_bitrate"):
            args += ["-crf", {"h264": "23", "h265": "28", "vp9": "32", "av1": "34"}[eff_v]]
        if eff_v in ("h264", "h265"):
            args += ["-preset", o.get("preset", "medium"), "-pix_fmt", "yuv420p"]
    if o.get("video_bitrate") and "-b:v" in extra:
        k = extra.index("-b:v")
        extra = extra[:k] + extra[k + 2:]
    if o.get("video_bitrate"):
        args += ["-b:v", o["video_bitrate"]]
    if o.get("remove_audio") or adef == "none" or not _has_audio(info):
        args.append("-an")
    else:
        args += _A_CODEC.get(ac) if ac not in ("auto", None) else (_A_CODEC.get(adef) or _SPECIAL_A[adef])
        args += _audio_opts(o, fmt)
    if o.get("subtitles") and o.get("subtitles_mode") == "soft":
        args += ["-map", "0:v:0", "-map", "0:a?", "-map", "1:0", "-c:s",
                 "mov_text" if fmt in ("mp4", "m4v", "mov") else "srt" if fmt == "mkv" else "webvtt"]
    if o.get("advanced"):
        args += advanced.check(o["advanced"], "ffmpeg")
    args += extra
    args += ["-map_metadata", "-1"] if o.get("strip_metadata") else []
    args += ["-f", mux, str(dst)]
    await run(*args, timeout=VIDEO_TIMEOUT, cwd=work)
    return _ok([dst])


def _audio_opts(o: dict, fmt: str) -> list[str]:
    a = []
    if o.get("audio_bitrate"):
        a += ["-b:a", o["audio_bitrate"]]
    if o.get("sample_rate"):
        a += ["-ar", str(o["sample_rate"])]
    if o.get("channels"):
        a += ["-ac", str(o["channels"])]
    if o.get("bit_depth") and fmt in ("wav", "aiff", "flac"):
        a += ["-sample_fmt", {16: "s16", 24: "s32", 32: "s32"}[int(o["bit_depth"])]]
        if fmt == "wav" and int(o["bit_depth"]) in (24, 32):
            a += ["-c:a", f"pcm_s{o['bit_depth']}le"]
    if o.get("volume"):
        a += ["-af", f"volume={o['volume']}"]
    return a


async def _animated(src: Path, fmt: str, ctx: Ctx, dst: Path) -> list[Path]:
    o = ctx.opts if ctx.final else {}
    fps = o.get("fps") or 12
    width = o.get("width") or o.get("max_width") or 480
    scale = f"fps={fps},scale={width}:-1:flags=lanczos"
    base = ["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", *_trim_args(o), "-i", str(src)]
    if fmt == "gif":
        await run(*base, "-vf", f"{scale},split[a][b];[a]palettegen=stats_mode=diff[p];[b][p]paletteuse=dither=bayer:bayer_scale=5",
                  "-loop", "0" if o.get("loop", True) else "-1", str(dst), timeout=VIDEO_TIMEOUT)
    elif fmt == "webp":
        await run(*base, "-vf", scale, "-c:v", "libwebp", "-lossless", "0", "-q:v", str(o.get("quality") or 70),
                  "-loop", "0", "-an", str(dst), timeout=VIDEO_TIMEOUT)
    else:
        await run(*base, "-vf", scale, "-plays", "0", "-f", "apng", str(dst), timeout=VIDEO_TIMEOUT)
    return _ok([dst])


async def _audio_to_video(src: Path, fmt: str, ctx: Ctx, dst: Path) -> list[Path]:
    o = ctx.opts
    size = (o.get("resolution") and _RES.get(o["resolution"], o["resolution"].replace("x", ":"))) or "1280:720"
    w, h = size.split(":")
    cover = ctx.file("cover") if o.get("cover") else None
    vin = ["-loop", "1", "-i", str(cover)] if cover else ["-f", "lavfi", "-i", f"color=c=black:s={w}x{h}:r=2"]
    vf = f"scale={w}:{h}:force_original_aspect_ratio=decrease,pad={w}:{h}:(ow-iw)/2:(oh-ih)/2,format=yuv420p"
    vc = ["-c:v", "libvpx-vp9", "-b:v", "0", "-crf", "40"] if fmt == "webm" else ["-c:v", "libx264", "-tune", "stillimage", "-preset", "veryfast"]
    ac = ["-c:a", "libopus"] if fmt == "webm" else ["-c:a", "aac", "-b:a", "192k"]
    await run("ffmpeg", "-y", "-hide_banner", "-loglevel", "error", *vin, *_trim_args(o), "-i", str(src), "-vf", vf,
              *vc, *ac, "-r", "2", "-shortest", "-map", "0:v", "-map", "1:a", str(dst), timeout=VIDEO_TIMEOUT)
    return _ok([dst])


async def slideshow(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    """Prezentáció → videó/GIF: diák PDF-be (LibreOffice), képpé (pdftoppm), majd diánként N másodperc."""
    sub = Ctx(ctx.sub("slides"), {}, final=False, source_fmt=ctx.source_fmt)
    from app.convert.engines import available
    if _fmt_of(src) == "pptx" and available("deckcraft"):     # Rust: diánként közvetlenül, LibreOffice nélkül
        pages = await deckcraft(src, "png", Ctx(sub.work, {"dpi": 144}, final=False))
    else:
        pdf = (await libreoffice(src, "pdf", sub))[0]
        pages = await poppler(pdf, "png", Ctx(sub.work, {"dpi": 110}, final=False))
    secs = float(ctx.opts.get("seconds_per_slide") or 3.0)
    lst = sub.work / "list.txt"
    lines = []
    for p in pages:
        lines += [f"file '{p.as_posix()}'", f"duration {secs}"]
    lines.append(f"file '{pages[-1].as_posix()}'")
    lst.write_text("\n".join(lines), "utf-8")
    dst = ctx.out(src, fmt)
    w = ctx.opts.get("width") or (720 if fmt == "gif" else 1920)
    if fmt == "gif":
        await run("ffmpeg", "-y", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", str(lst), "-vf",
                  f"fps=2,scale={w}:-1:flags=lanczos,split[a][b];[a]palettegen[p];[b][p]paletteuse", "-loop", "0", str(dst),
                  timeout=VIDEO_TIMEOUT)
    else:
        vc = ["-c:v", "libvpx-vp9", "-b:v", "0", "-crf", "34"] if fmt == "webm" else ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-tune", "stillimage"]
        await run("ffmpeg", "-y", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", str(lst), "-vf",
                  f"scale={w}:-2:force_original_aspect_ratio=decrease,pad=ceil(iw/2)*2:ceil(ih/2)*2,fps=25", *vc, str(dst),
                  timeout=VIDEO_TIMEOUT)
    return _ok([dst])


# ======================================================================
# Archívumok
# ======================================================================

_TAR_MODE = {"tar": "w", "tar.gz": "w:gz", "tgz": "w:gz", "tar.bz2": "w:bz2", "tbz2": "w:bz2", "tbz": "w:bz2",
             "tar.bz": "w:bz2", "tar.xz": "w:xz", "txz": "w:xz"}
_SINGLE = {"gz", "bz2", "bz", "xz", "zst", "lz", "lzma", "lzo", "z"}


def _safe_member(name: str) -> bool:
    n = name.replace("\\", "/")
    return not (n.startswith("/") or n.startswith("../") or "/../" in n or n == ".." or re.match(r"^[A-Za-z]:", n))


def _check_tree(root: Path) -> None:
    total, count = 0, 0
    for p in root.rglob("*"):
        if p.is_symlink():
            p.unlink()
            continue
        if p.is_file():
            count += 1
            total += p.stat().st_size
            if total > MAX_EXTRACT_BYTES or count > MAX_EXTRACT_FILES:
                raise EngineError("the archive is too large to unpack here")
        try:
            p.resolve().relative_to(root.resolve())
        except ValueError:
            raise EngineError("unsafe path inside the archive")


async def extract(src: Path, dest: Path, password: Optional[str] = None) -> Path:
    """Bármely archívum kibontása `dest`-be (zip-slip és zip-bomb ellen védve)."""
    fmt = _fmt_of(src)
    dest.mkdir(parents=True, exist_ok=True)

    def py_extract() -> bool:
        if fmt in ("zip", "jar", "war", "apk", "cbz") and not password:
            with zipfile.ZipFile(src) as z:
                if sum(i.file_size for i in z.infolist()) > MAX_EXTRACT_BYTES or len(z.infolist()) > MAX_EXTRACT_FILES:
                    raise EngineError("the archive is too large to unpack here")
                for i in z.infolist():
                    if not _safe_member(i.filename):
                        raise EngineError("unsafe path inside the archive")
                    name = i.filename
                    if not (i.flag_bits & 0x800):          # régi zip: cp437-ből vissza (a magyar ékezet így épül)
                        try:
                            name = name.encode("cp437").decode("utf-8")
                        except (UnicodeEncodeError, UnicodeDecodeError):
                            pass
                    target = dest / name
                    if i.is_dir():
                        target.mkdir(parents=True, exist_ok=True)
                        continue
                    target.parent.mkdir(parents=True, exist_ok=True)
                    with z.open(i) as r, open(target, "wb") as w:
                        shutil.copyfileobj(r, w)
            return True
        if fmt in ("tar", "tar.gz", "tgz", "tar.bz2", "tbz2", "tbz", "tar.bz", "tar.xz", "txz"):
            with tarfile.open(src) as t:
                members = []
                for m in t.getmembers():
                    if not _safe_member(m.name) or m.issym() or m.islnk() or m.isdev():
                        continue
                    members.append(m)
                t.extractall(dest, members=members)
            return True
        if fmt in ("gz", "bz2", "bz", "xz", "lzma"):
            opener = {"gz": gzip.open, "bz2": bz2.open, "bz": bz2.open, "xz": lzma.open, "lzma": lzma.open}[fmt]
            name = src.name[: -len(fmt) - 1] if src.name.lower().endswith("." + fmt) else src.stem
            with opener(src, "rb") as r, open(dest / (name or "file"), "wb") as w:
                n = 0
                while chunk := r.read(1 << 20):
                    n += len(chunk)
                    if n > MAX_EXTRACT_BYTES:
                        raise EngineError("the archive is too large to unpack here")
                    w.write(chunk)
            return True
        if fmt == "7z" and not shutil.which("7z"):
            import py7zr
            with py7zr.SevenZipFile(src, password=password) as z:
                if any(not _safe_member(n) for n in z.getnames()):
                    raise EngineError("unsafe path inside the archive")
                z.extractall(dest)
            return True
        return False

    if not await asyncio.to_thread(py_extract):
        if fmt in ("zst", "tar.zst"):
            mid = dest.parent / (_stem(src) + (".tar" if fmt == "tar.zst" else ""))
            await run("zstd", "-d", "-f", "-o", str(mid), str(src), timeout=600)
            if fmt == "tar.zst":
                await extract(mid, dest)
            else:
                shutil.move(str(mid), dest / mid.name)
        elif fmt in ("lz",):
            out, _ = await run("lzip", "-d", "-c", str(src), timeout=600)
            (dest / _stem(src)).write_bytes(out)
        elif fmt in ("lzo", "tar.lzo", "tzo"):
            mid = dest.parent / (_stem(src) + (".tar" if fmt != "lzo" else ""))
            out, _ = await run("lzop", "-d", "-c", str(src), timeout=600)
            mid.write_bytes(out)
            if fmt != "lzo":
                await extract(mid, dest)
            else:
                shutil.move(str(mid), dest / mid.name)
        elif fmt in ("tar.z", "tz"):
            mid = dest.parent / "tarz"
            await extract_7z(src, mid, None)
            inner = [p for p in mid.rglob("*") if p.is_file()]
            if not inner:
                raise EngineError("empty .tar.Z")
            tar = inner[0].with_name(inner[0].name + ".tar")
            inner[0].rename(tar)
            await extract(tar, dest)
        elif fmt == "tar.7z":
            mid = dest.parent / "tar7z"
            await extract_7z(src, mid, password)
            tars = list(mid.rglob("*.tar"))
            if not tars:
                raise EngineError("no tar inside the .tar.7z")
            await extract(tars[0], dest)
        elif shutil.which("7z") and fmt not in ("ace", "alz", "arc"):
            await extract_7z(src, dest, password)
        elif shutil.which("unar"):
            args = ["unar", "-q", "-f", "-D", "-o", str(dest)]
            if password:
                args += ["-p", password]
            await run(*args, str(src), timeout=1200)
        elif shutil.which("bsdtar"):
            await run("bsdtar", "-xf", str(src), "-C", str(dest), timeout=1200)
        else:
            raise EngineError(f"no unpacker for .{fmt} on this server")
    await asyncio.to_thread(_check_tree, dest)
    return dest


async def extract_7z(src: Path, dest: Path, password: Optional[str]) -> None:
    dest.mkdir(parents=True, exist_ok=True)
    out, _ = await run("7z", "l", "-slt", "-p" + (password or ""), str(src), timeout=300, ok_codes=(0, 2))
    names = [l[7:] for l in out.decode("utf-8", "replace").splitlines() if l.startswith("Path = ")][1:]
    if any(not _safe_member(n) for n in names):
        raise EngineError("unsafe path inside the archive")
    await run("7z", "x", "-y", "-p" + (password or ""), f"-o{dest}", str(src), timeout=1800)


async def pack(root: Path, dst: Path, fmt: str, password: Optional[str] = None, level: Optional[int] = None) -> Path:
    """Egy mappa tartalma → archívum (a mappa NEM kerül bele, csak a tartalma)."""
    files = sorted(p for p in root.rglob("*") if p.is_file())

    def py_pack() -> bool:
        if fmt in ("zip", "jar", "war") and not password:
            comp = zipfile.ZIP_DEFLATED
            with zipfile.ZipFile(dst, "w", comp, compresslevel=level if level is not None else 6) as z:
                for f in files:
                    z.write(f, f.relative_to(root).as_posix())
            return True
        if fmt in _TAR_MODE:
            kw = {"compresslevel": level} if level is not None and fmt not in ("tar", "tar.xz", "txz") else {}
            with tarfile.open(dst, _TAR_MODE[fmt], **kw) as t:
                for f in files:
                    t.add(f, f.relative_to(root).as_posix())
            return True
        if fmt in ("gz", "bz2", "bz", "xz", "lzma"):
            if len(files) != 1:
                raise EngineError(f".{fmt} holds a single file — choose .tar.{fmt if fmt != 'lzma' else 'xz'} for several")
            opener = {"gz": gzip.open, "bz2": bz2.open, "bz": bz2.open, "xz": lzma.open,
                      "lzma": lambda p, m: lzma.open(p, m, format=lzma.FORMAT_ALONE)}[fmt]
            with open(files[0], "rb") as r, opener(dst, "wb") as w:
                shutil.copyfileobj(r, w)
            return True
        if fmt == "7z" and not shutil.which("7z"):
            import py7zr
            with py7zr.SevenZipFile(dst, "w", password=password, header_encryption=bool(password)) as z:
                for f in files:
                    z.write(f, f.relative_to(root).as_posix())
            return True
        return False

    if await asyncio.to_thread(py_pack):
        return dst
    if fmt in ("zip", "7z", "wim", "jar", "war"):
        t = {"jar": "zip", "war": "zip"}.get(fmt, fmt)
        args = ["7z", "a", f"-t{t}", "-y"]
        if password and t in ("zip", "7z"):
            args += [f"-p{password}"] + (["-mhe=on"] if t == "7z" else ["-mem=AES256"])
        if level is not None:
            args.append(f"-mx={level}")
        await run(*args, str(dst), ".", cwd=root, timeout=1800)
    elif fmt in ("tar.zst", "zst", "tar.7z", "tar.lzo", "tzo", "lzo", "lz"):
        single = fmt in ("zst", "lzo", "lz")
        if single and len(files) != 1:
            raise EngineError(f".{fmt} holds a single file — choose a tar-based format for several")
        mid = files[0] if single else dst.parent / "pack.tar"
        if not single:
            with tarfile.open(mid, "w") as t:
                for f in files:
                    t.add(f, f.relative_to(root).as_posix())
        if fmt in ("tar.zst", "zst"):
            await run("zstd", "-q", "-f", f"-{level or 3}", "-o", str(dst), str(mid), timeout=1800)
        elif fmt in ("tar.lzo", "tzo", "lzo"):
            out, _ = await run("lzop", "-c", str(mid), timeout=1800)
            dst.write_bytes(out)
        elif fmt == "lz":
            out, _ = await run("lzip", "-c", str(mid), timeout=1800)
            dst.write_bytes(out)
        else:
            await run("7z", "a", "-t7z", "-y", str(dst), str(mid), timeout=1800)
    elif fmt in ("cpio", "iso"):
        await run("bsdtar", "--format", "cpio" if fmt == "cpio" else "iso9660", "-cf", str(dst), "-C", str(root), ".", timeout=1800)
    elif fmt in ("tar.z", "tz", "z"):
        if fmt == "z" and len(files) != 1:
            raise EngineError(".Z holds a single file — choose .tar.Z for several")
        mid = files[0] if fmt == "z" else dst.parent / "pack.tar"
        if fmt != "z":
            with tarfile.open(mid, "w") as t:
                for f in files:
                    t.add(f, f.relative_to(root).as_posix())
        await _compress_single(mid, dst, "z")
    else:
        raise EngineError(f"can't write .{fmt} archives")
    return dst


async def archive(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    tmp = ctx.sub("x")
    await extract(src, tmp / "c", ctx.opts.get("password"))
    n = sum(1 for p in (tmp / "c").rglob("*") if p.is_file())
    if fmt in _SINGLE and n != 1:
        # egyfájlos tömörítő több fájlra: előbb tar, és a név .tar.<ext> lesz (ahogy a 7-Zip és a tar is teszi)
        dst = ctx.out(src, fmt).with_name(_stem(src) + f".tar.{fmt}")
        tar = tmp / "pack.tar"
        await pack(tmp / "c", tar, "tar")
        await _compress_single(tar, dst, fmt, ctx.opts.get("level"))
        return _ok([dst])
    dst = ctx.out(src, fmt)
    await pack(tmp / "c", dst, fmt, ctx.opts.get("password") if ctx.final else None, ctx.opts.get("level"))
    return _ok([dst])


async def _compress_single(src: Path, dst: Path, fmt: str, level: Optional[int] = None) -> None:
    """Egyetlen fájl tömörítése: gz, bz2/bz, xz, lzma, zst, lz, lzo, z."""
    if fmt in ("gz", "bz2", "bz", "xz", "lzma"):
        def go():
            opener = {"gz": gzip.open, "bz2": bz2.open, "bz": bz2.open, "xz": lzma.open,
                      "lzma": lambda p, m: lzma.open(p, m, format=lzma.FORMAT_ALONE)}[fmt]
            with open(src, "rb") as r, opener(dst, "wb") as w:
                shutil.copyfileobj(r, w)
        await asyncio.to_thread(go)
    elif fmt == "z":
        from app.convert.pyformats import compress_z
        data = await asyncio.to_thread(src.read_bytes)
        await asyncio.to_thread(dst.write_bytes, compress_z(data))
    elif fmt == "zst":
        await run("zstd", "-q", "-f", f"-{level or 3}", "-o", str(dst), str(src), timeout=1800)
    elif fmt in ("lzo", "lz"):
        out, _ = await run("lzop" if fmt == "lzo" else "lzip", "-c", str(src), timeout=1800)
        dst.write_bytes(out)
    else:
        raise EngineError(f"can't write .{fmt}")


# ======================================================================
# Betűk: fontTools (+ saját EOT), FontForge
# ======================================================================

def _ttf_to_otf(font) -> None:
    """Kvadratikus (glyf) → köbös (CFF) körvonalak."""
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.t2CharStringPen import T2CharStringPen
    gs = font.getGlyphSet()
    charstrings = {}
    for name in font.getGlyphOrder():
        pen = T2CharStringPen(font["hmtx"][name][0], gs)
        gs[name].draw(pen)
        charstrings[name] = pen.getCharString()
    names = {n.nameID: n.toUnicode() for n in font["name"].names if n.platformID in (3, 1)}
    ps_name = re.sub(r"[^A-Za-z0-9\-]", "", names.get(6) or names.get(4) or "Font") or "Font"
    for t in ("glyf", "loca", "fpgm", "prep", "cvt ", "gasp", "hdmx", "LTSH", "VDMX"):
        if t in font:
            del font[t]
    fb = FontBuilder(font=font, isTTF=False)
    fb.setupCFF(ps_name, {"FullName": names.get(4, ps_name), "FamilyName": names.get(1, ps_name)}, charstrings, {})
    font.sfntVersion = "OTTO"
    font["maxp"].tableVersion = 0x00005000


def _otf_to_ttf(font) -> None:
    """Köbös (CFF) → kvadratikus (glyf) körvonalak (cu2qu)."""
    from fontTools.pens.cu2quPen import Cu2QuPen
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    from fontTools.ttLib import newTable
    gs = font.getGlyphSet()
    order = font.getGlyphOrder()
    glyphs = {}
    for name in order:
        pen = TTGlyphPen(gs)
        gs[name].draw(Cu2QuPen(pen, 1.0, reverse_direction=True))
        glyphs[name] = pen.glyph()
    for t in ("CFF ", "CFF2", "VORG"):
        if t in font:
            del font[t]
    glyf = newTable("glyf")
    glyf.glyphOrder = order
    glyf.glyphs = glyphs
    font["glyf"] = glyf
    font["loca"] = newTable("loca")
    maxp = font["maxp"]
    maxp.tableVersion = 0x00010000
    for attr, v in (("maxZones", 1), ("maxTwilightPoints", 0), ("maxStorage", 0), ("maxFunctionDefs", 0),
                    ("maxInstructionDefs", 0), ("maxStackElements", 0), ("maxSizeOfInstructions", 0),
                    ("maxComponentElements", 0), ("maxPoints", 0), ("maxContours", 0), ("maxCompositePoints", 0),
                    ("maxCompositeContours", 0), ("maxComponentDepth", 0)):
        setattr(maxp, attr, v)
    post = font["post"]
    post.formatType = 2.0
    post.extraNames = []
    post.mapping = {}
    font["head"].glyphDataFormat = 0
    font.sfntVersion = "\x00\x01\x00\x00"


def eot_wrap(ttf: bytes) -> bytes:
    """TTF → EOT (2.1-es fejléc, tömörítetlen, MTX nélkül) — az IE ≤ 8 webbetű-formátuma."""
    import struct
    from fontTools.ttLib import TTFont
    f = TTFont(io.BytesIO(ttf))
    os2, head, name = f["OS/2"], f["head"], f["name"]

    def nm(i):
        n = name.getName(i, 3, 1, 0x409) or name.getName(i, 1, 0, 0)
        s = (n.toUnicode() if n else "").encode("utf-16-le")
        return struct.pack("<H", len(s)) + s

    panose = bytes([os2.panose.bFamilyType, os2.panose.bSerifStyle, os2.panose.bWeight, os2.panose.bProportion,
                    os2.panose.bContrast, os2.panose.bStrokeVariation, os2.panose.bArmStyle, os2.panose.bLetterForm,
                    os2.panose.bMidline, os2.panose.bXHeight])
    body = (panose + struct.pack("<BBLHH", 1, 1 if os2.fsSelection & 1 else 0, os2.usWeightClass, os2.fsType, 0x504C)
            + struct.pack("<LLLL", os2.ulUnicodeRange1, os2.ulUnicodeRange2, os2.ulUnicodeRange3, os2.ulUnicodeRange4)
            + struct.pack("<LLL", os2.ulCodePageRange1, getattr(os2, "ulCodePageRange2", 0), head.checkSumAdjustment)
            + b"\x00" * 16 + struct.pack("<H", 0) + nm(1) + struct.pack("<H", 0) + nm(2) + struct.pack("<H", 0) + nm(5)
            + struct.pack("<H", 0) + nm(4) + struct.pack("<HH", 0, 0) + struct.pack("<LL", 0, 0))
    total = 4 + 4 + 4 + 4 + len(body) + len(ttf)
    return struct.pack("<LLLL", total, len(ttf), 0x00020001, 0) + body + ttf


def eot_unwrap(eot: bytes) -> bytes:
    import struct
    total, size, version, flags = struct.unpack("<LLLL", eot[:16])
    if flags & 0x4:
        raise EngineError("this EOT is MicroType-compressed; only uncompressed EOT can be converted")
    if flags & 0x10000000:
        raise EngineError("this EOT is obfuscated (XOR) and can't be converted")
    return eot[total - size: total]


async def fonttools(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    dst = ctx.out(src, fmt)
    sfmt = _fmt_of(src)

    def go():
        from fontTools.ttLib import TTCollection, TTFont
        if sfmt == "eot":
            font = TTFont(io.BytesIO(eot_unwrap(src.read_bytes())))
        elif sfmt == "ttc":
            font = TTCollection(str(src)).fonts[int(ctx.opts.get("font_index") or 0)]
        else:
            font = TTFont(str(src))
        font.flavor = None
        is_cff = "CFF " in font or "CFF2" in font
        if ctx.opts.get("subset_text"):
            from fontTools import subset
            s = subset.Subsetter()
            s.populate(text=ctx.opts["subset_text"])
            s.subset(font)
        if fmt in ("ttf", "eot") and is_cff:
            _otf_to_ttf(font)
        elif fmt == "otf" and not is_cff:
            _ttf_to_otf(font)
        if fmt in ("woff", "woff2"):
            font.flavor = fmt
            font.save(str(dst))
        elif fmt == "eot":
            buf = io.BytesIO()
            font.save(buf)
            dst.write_bytes(eot_wrap(buf.getvalue()))
        elif fmt == "ttc":
            c = TTCollection()
            c.fonts = [font]
            c.save(str(dst))
        else:
            font.save(str(dst))
    await asyncio.to_thread(go)
    return _ok([dst])


_FF_SCRIPT = ("import fontforge,sys\n"
              "f=fontforge.open(sys.argv[1])\n"
              "o=sys.argv[2]\n"
              "if o.endswith('.bdf'):\n"
              "    f.bitmapSizes=((16,),)\n"
              "    f.regenBitmaps((16,))\n"
              "    f.generate(o[:-4]+'.', bitmap_type='bdf')\n"
              "elif o.endswith('.sfd'):\n"
              "    f.save(o)\n"
              "else:\n"
              "    f.generate(o)\n")


async def fontforge(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    work = ctx.sub("ff")
    script = work / "conv.py"
    script.write_text(_FF_SCRIPT)
    tmp = work / f"out.{fmt}"
    await run("fontforge", "-lang=py", "-script", str(script), str(src), str(tmp), timeout=300, cwd=work)
    produced = tmp if tmp.exists() else next(iter(sorted(work.glob("out*.bdf"))), None) if fmt == "bdf" else None
    if not produced:
        raise EngineError("FontForge produced nothing")
    dst = ctx.out(src, fmt)
    shutil.move(str(produced), dst)
    return _ok([dst])


# ======================================================================
# CAD: LibreDWG, ezdxf
# ======================================================================

async def libredwg(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    dst = ctx.out(src, fmt)
    exe = "dwg2dxf" if fmt == "dxf" else "dxf2dwg"
    await run(exe, "-y", "-o", str(dst), str(src), timeout=600, ok_codes=(0, 1))
    return _ok([dst])


import threading  # noqa: E402

_EZDXF_LOCK = threading.Lock()          # az ezdxf betűkezelője és a matplotlib nem szálbiztos


async def ezdxf_render(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    dst = ctx.out(src, fmt)
    o = ctx.opts if ctx.final else {}

    def go():
        with _EZDXF_LOCK:
            _draw()

    def _draw():
        import matplotlib
        matplotlib.use("Agg")
        import ezdxf
        from ezdxf import recover
        from ezdxf.addons.drawing import matplotlib as mpl
        try:
            doc = ezdxf.readfile(str(src))
        except Exception:
            doc, _ = recover.readfile(str(src))
        bg = o.get("background") or "#FFFFFF"
        dpi = int(o.get("dpi") or 200)
        mpl.qsave(doc.modelspace(), str(dst), bg=bg if bg != "transparent" else "#FFFFFF00", dpi=dpi)
    await asyncio.to_thread(go)
    return _ok([dst])


# ======================================================================
# Kis Python-átalakítók: tömörített változatok, e-mail, HWPX, adatformátumok, képregény
# ======================================================================

async def python_step(src: Path, fmt: str, ctx: Ctx) -> list[Path]:
    sfmt = _fmt_of(src)
    dst = ctx.out(src, fmt)
    if sfmt in ("cbz", "cbr", "cb7") and fmt == "pdf":
        tmp = await extract(src, ctx.sub("comic"))
        await asyncio.to_thread(_images_to_pdf, tmp, dst)
        return _ok([dst])
    await asyncio.to_thread(_py_convert, src, sfmt, dst, fmt, ctx)
    return _ok([dst])


def _py_convert(src: Path, sfmt: str, dst: Path, fmt: str, ctx: Ctx) -> None:
    if (sfmt, fmt) in (("zabw", "abw"), ("svgz", "svg"), ("emz", "emf"), ("wmz", "wmf")):
        with gzip.open(src, "rb") as r:
            dst.write_bytes(r.read(MAX_EXTRACT_BYTES))
        return
    if (sfmt, fmt) in (("svg", "svgz"), ("emf", "emz"), ("wmf", "wmz")):
        with gzip.open(dst, "wb") as w:
            w.write(src.read_bytes())
        return
    if (sfmt, fmt) in (("pdfa", "pdf"), ("pdf", "ai"), ("ai", "pdf")):
        shutil.copy(src, dst)              # az .ai PDF-kompatibilis: ugyanaz a tartalom
        return
    if (sfmt, fmt) == ("plt", "svg"):
        from app.convert.pyformats import hpgl_to_svg
        return hpgl_to_svg(src, dst)
    if (sfmt, fmt) == ("pdf", "plt"):
        from app.convert.pyformats import pdf_to_hpgl
        return pdf_to_hpgl(src, dst)
    if sfmt == "eml":
        return _eml(src, dst, fmt)
    if sfmt == "hwpx":
        return _hwpx(src, dst, fmt)
    if sfmt == "txt" and fmt in ("md", "html"):
        text = src.read_text("utf-8", "replace")
        if fmt == "md":
            dst.write_text(text, "utf-8")
        else:
            import html as H
            body = "".join(f"<p>{H.escape(p).replace(chr(10), '<br>')}</p>\n" for p in re.split(r"\n\s*\n", text) if p.strip())
            dst.write_text(f"<!doctype html><html><head><meta charset=\"utf-8\"><title>{H.escape(_stem(src))}</title></head>"
                           f"<body>\n{body}</body></html>", "utf-8")
        return
    if sfmt == "pdf" and fmt == "cbz":
        import pymupdf
        with pymupdf.open(str(src)) as d, zipfile.ZipFile(dst, "w") as z:
            for i, page in enumerate(d, 1):
                z.writestr(f"{i:04d}.jpg", page.get_pixmap(dpi=int(ctx.opts.get("dpi") or 150)).tobytes("jpeg"))
        return
    if sfmt == "csv" and fmt == "tsv" or sfmt == "tsv" and fmt == "csv":
        rows = _read_table(src, sfmt, ctx.opts)
        _write_rows(rows, dst, fmt)
        return
    # adatformátumok
    data = _read_data(src, sfmt, ctx.opts)
    _write_data(data, dst, fmt, _stem(src))


def _eml(src: Path, dst: Path, fmt: str) -> None:
    import email
    import html as H
    from email import policy
    msg = email.message_from_bytes(src.read_bytes(), policy=policy.default)
    head = [(k, str(msg.get(k, ""))) for k in ("From", "To", "Cc", "Date", "Subject") if msg.get(k)]
    body = msg.get_body(preferencelist=("html", "plain"))
    content = body.get_content() if body else ""
    is_html = body is not None and body.get_content_type() == "text/html"
    atts = [p.get_filename() for p in msg.iter_attachments() if p.get_filename()]
    if fmt == "txt":
        text = content
        if is_html:
            from bs4 import BeautifulSoup
            text = BeautifulSoup(content, "lxml").get_text("\n")
        dst.write_text("\n".join(f"{k}: {v}" for k, v in head) + "\n\n" + text.strip() + "\n" +
                       ("\nAttachments: " + ", ".join(atts) + "\n" if atts else ""), "utf-8")
        return
    if not is_html:
        content = "<pre style=\"white-space:pre-wrap;font-family:inherit\">" + H.escape(content) + "</pre>"
    else:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(content, "lxml")
        for t in soup(["script", "iframe", "object", "embed"]):
            t.decompose()
        content = str(soup.body or soup)
    rows = "".join(f"<tr><th>{H.escape(k)}</th><td>{H.escape(v)}</td></tr>" for k, v in head)
    att = ("<p><b>Attachments:</b> " + H.escape(", ".join(atts)) + "</p>") if atts else ""
    dst.write_text(f"<!doctype html><html><head><meta charset=\"utf-8\"><title>{H.escape(str(msg.get('Subject', 'E-mail')))}</title>"
                   "<style>table{border-collapse:collapse;margin-bottom:1em}th{text-align:left;padding:2px 12px 2px 0;color:#555}"
                   "</style></head><body><table>" + rows + "</table>" + att + "<hr>" + content + "</body></html>", "utf-8")


def _hwpx(src: Path, dst: Path, fmt: str) -> None:
    from lxml import etree
    paras: list[str] = []
    with zipfile.ZipFile(src) as z:
        secs = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n))
        for s in secs:
            root = etree.fromstring(z.read(s))
            for p in root.iter():
                if etree.QName(p).localname == "p":
                    txt = "".join(t.text or "" for t in p.iter() if etree.QName(t).localname == "t")
                    paras.append(txt)
    if fmt == "txt":
        dst.write_text("\n".join(paras) + "\n", "utf-8")
    elif fmt == "html":
        import html as H
        dst.write_text("<!doctype html><html><head><meta charset=\"utf-8\"></head><body>" +
                       "".join(f"<p>{H.escape(p)}</p>" for p in paras) + "</body></html>", "utf-8")
    else:
        from docx import Document
        d = Document()
        for p in paras:
            d.add_paragraph(p)
        d.save(str(dst))


def _images_to_pdf(folder: Path, dst: Path) -> None:
    from PIL import Image
    imgs = sorted(p for p in folder.rglob("*") if p.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp"))
    if not imgs:
        raise EngineError("no pictures inside the comic archive")
    pages = [Image.open(p).convert("RGB") for p in imgs]
    pages[0].save(dst, "PDF", save_all=True, append_images=pages[1:], resolution=150)


def _read_table(src: Path, sfmt: str, opts: dict) -> list[list[str]]:
    raw = src.read_bytes().decode("utf-8-sig", "replace")
    d = opts.get("delimiter", "auto")
    d = {"tab": "\t", "auto": None}.get(d, d) or ("\t" if sfmt == "tsv" else _sniff_delimiter(src))
    return [r for r in csv.reader(io.StringIO(raw), delimiter=d)]


def _write_rows(rows: list[list[str]], dst: Path, fmt: str) -> None:
    with open(dst, "w", newline="", encoding="utf-8") as f:
        csv.writer(f, delimiter="\t" if fmt == "tsv" else ",").writerows(rows)


def _read_data(src: Path, sfmt: str, opts: dict) -> Any:
    if sfmt == "json":
        return json.loads(src.read_text("utf-8-sig"))
    if sfmt == "yaml":
        import yaml
        return yaml.safe_load(src.read_text("utf-8-sig"))
    if sfmt == "xml":
        from lxml import etree
        parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
        root = etree.parse(str(src), parser).getroot()
        return {etree.QName(root).localname: _xml_to_obj(root)}
    if sfmt in ("csv", "tsv"):
        rows = _read_table(src, sfmt, opts)
        if not rows:
            return []
        head = [h.strip() or f"col{i + 1}" for i, h in enumerate(rows[0])]
        return [dict(zip(head, r)) for r in rows[1:] if any(c.strip() for c in r)]
    if sfmt == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(src), read_only=True, data_only=True)
        ws = wb[opts["sheet"]] if opts.get("sheet") in wb.sheetnames else wb.active
        rows = [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
        if not rows:
            return []
        head = [str(h).strip() or f"col{i + 1}" for i, h in enumerate(rows[0])]
        return [{h: (v.isoformat() if hasattr(v, "isoformat") else v) for h, v in zip(head, r)} for r in rows[1:]
                if any(c not in ("", None) for c in r)]
    raise EngineError(f"can't read .{sfmt} as data")


def _xml_to_obj(el) -> Any:
    from lxml import etree
    kids = [c for c in el if isinstance(c.tag, str)]
    attrs = {f"@{etree.QName(k).localname}": v for k, v in el.attrib.items()}
    text = (el.text or "").strip()
    if not kids and not attrs:
        return text
    out: dict[str, Any] = dict(attrs)
    for c in kids:
        k = etree.QName(c).localname
        v = _xml_to_obj(c)
        if k in out:
            out[k] = out[k] if isinstance(out[k], list) else [out[k]]
            out[k].append(v)
        else:
            out[k] = v
    if text:
        out["#text"] = text
    return out


def _records(data: Any) -> list[dict]:
    if isinstance(data, dict):
        lists = [v for v in data.values() if isinstance(v, list)]
        if len(data) == 1 and isinstance(next(iter(data.values())), (dict, list)):
            return _records(next(iter(data.values())))
        if lists:
            return _records(lists[0])
        return [data]
    if isinstance(data, list):
        return [r if isinstance(r, dict) else {"value": r} for r in data]
    return [{"value": data}]


def _flat(d: dict, pre: str = "") -> dict:
    out = {}
    for k, v in d.items():
        key = f"{pre}{k}"
        if isinstance(v, dict):
            out.update(_flat(v, key + "."))
        elif isinstance(v, list):
            out[key] = json.dumps(v, ensure_ascii=False)
        else:
            out[key] = v
    return out


def _write_data(data: Any, dst: Path, fmt: str, title: str) -> None:
    if fmt == "json":
        dst.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), "utf-8")
    elif fmt == "yaml":
        import yaml
        dst.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), "utf-8")
    elif fmt == "xml":
        from lxml import etree

        def build(parent, key, v):
            key = re.sub(r"[^\w.\-]", "_", str(key)) or "item"
            if key[0].isdigit() or key.startswith("-"):
                key = "_" + key
            if isinstance(v, list):
                for x in v:
                    build(parent, key, x)
                return
            el = etree.SubElement(parent, key.lstrip("@"))
            if isinstance(v, dict):
                for k2, v2 in v.items():
                    if k2 == "#text":
                        el.text = str(v2)
                    else:
                        build(el, k2, v2)
            elif v is not None:
                el.text = str(v)
        if isinstance(data, dict) and len(data) == 1 and isinstance(next(iter(data.values())), (dict, list)):
            k, v = next(iter(data.items()))
            root = etree.Element(re.sub(r"[^\w.\-]", "_", k) or "root")
            for k2, v2 in (v.items() if isinstance(v, dict) else [("item", v)]):
                build(root, k2, v2)
        else:
            root = etree.Element("root")
            for x in (data if isinstance(data, list) else [data]):
                build(root, "item", x)
        dst.write_bytes(etree.tostring(root, pretty_print=True, xml_declaration=True, encoding="utf-8"))
    elif fmt in ("csv", "tsv", "xlsx", "md", "html", "txt"):
        recs = [_flat(r) for r in _records(data)]
        head: list[str] = []
        for r in recs:
            for k in r:
                if k not in head:
                    head.append(k)
        rows = [[("" if r.get(h) is None else r.get(h)) for h in head] for r in recs]
        if fmt in ("csv", "tsv"):
            _write_rows([head] + [[str(c) for c in r] for r in rows], dst, fmt)
        elif fmt == "xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.append(head)
            for r in rows:
                ws.append([c if isinstance(c, (int, float)) else str(c) for c in r])
            wb.save(str(dst))
        elif fmt == "md":
            esc = lambda c: str(c).replace("|", "\\|").replace("\n", " ")
            lines = ["| " + " | ".join(esc(h) for h in head) + " |", "|" + "---|" * len(head)]
            lines += ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows]
            dst.write_text("\n".join(lines) + "\n", "utf-8")
        elif fmt == "html":
            import html as H
            th = "".join(f"<th>{H.escape(str(h))}</th>" for h in head)
            trs = "".join("<tr>" + "".join(f"<td>{H.escape(str(c))}</td>" for c in r) + "</tr>" for r in rows)
            dst.write_text(f"<!doctype html><html><head><meta charset=\"utf-8\"><title>{H.escape(title)}</title>"
                           "<style>table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:4px 8px}</style></head>"
                           f"<body><table><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table></body></html>", "utf-8")
        else:
            dst.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str) if not recs else
                           "\n".join("\t".join(str(c) for c in r) for r in [head] + rows) + "\n", "utf-8")
    else:
        raise EngineError(f"can't write data as .{fmt}")


STEPS: dict[str, Callable[[Path, str, Ctx], Awaitable[list[Path]]]] = {
    "wordcraft": wordcraft, "deckcraft": deckcraft, "gridcraft": gridcraft, "photocraft": photocraft,
    "vectorcraft": vectorcraft, "libreoffice": libreoffice, "pandoc": pandoc, "weasyprint": weasyprint, "pdf2docx": pdf2docx, "poppler": poppler,
    "mupdf": mupdf, "ghostscript": ghostscript, "djvulibre": djvulibre, "calibre": calibre, "pillow": pillow,
    "imagemagick": imagemagick, "heif": heif, "jxl": jxl, "libraw": libraw, "potrace": potrace, "inkscape": inkscape,
    "rsvg": rsvg, "ffmpeg": ffmpeg, "slideshow": slideshow, "archive": archive, "fonttools": fonttools,
    "fontforge": fontforge, "libredwg": libredwg, "ezdxf": ezdxf_render, "python": python_step,
}
