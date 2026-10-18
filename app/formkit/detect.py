"""Űrlap-felismerés bármiből → egységes mezőséma (schema.py).

    kitölthető PDF   a valódi mezők (név, típus, opciók, érték), címke a mező bal oldaláról / fölüle
    lapos PDF        vonalak és aláhúzás-sorok ("Név: ______"), üres keretek, jelölőnégyzet-keretek + a mellettük
                     álló felirat; a betűméretből számolt tűrésekkel (nincs AI)
    DOCX             "____", "……", "[   ]", ☐ minták, {{mező}} helyőrzők, tartalomvezérlők (SDT), táblázatban a
                     felirat melletti üres cella, illetve fejléc + üres sorok (table_fields — szöveges cella soha)
    XLSX             felirat-cella melletti üres cella (űrlap-oszlop), "____" cellák, legördülő lista opciói
    fotó / szken     ha van látó modell (OCR_MODEL): a meglévő forms.read_layout; különben helyben: Tesseract-szavak
                     + OpenCV-vonalak és -négyzetek, ugyanazzal a címke-logikával, mint a lapos PDF
"""
from __future__ import annotations

import asyncio
import re
import statistics
from pathlib import Path
from typing import Any, Optional

from app.catalog.formats import FORMATS, from_filename
from app.formkit import schema as S

BLANK_RE = re.compile(r"_{3,}|\.{5,}|…{2,}|\[\s*\]|\(\s{2,}\)|☐|□|\{\{\s*([^{}]{1,60}?)\s*\}\}")


# ---------------------------------------------------------------- közös: címke-keresés elrendezésből

def _lines_of(words: list[tuple]) -> dict[Any, list[tuple]]:
    lines: dict[Any, list[tuple]] = {}
    for w in words:
        lines.setdefault(w[5], []).append(w)
    for k in lines:
        lines[k].sort(key=lambda w: w[0])
    return lines


def _label_near(words: list[tuple], box: tuple[float, float, float, float], h: float, prefer_right: bool = False,
                stop_x: Optional[float] = None) -> str:
    """A doboz felirata: ugyanabban a sávban balra a legközelebbi sor szavai (a doboz előtt), ennek híján fölötte."""
    x0, y0, x1, y1 = box
    cy = (y0 + y1) / 2
    band = [w for w in words if abs((w[1] + w[3]) / 2 - cy) <= 0.9 * h and not _is_blank_word(w[4])]
    left = [w for w in band if w[2] <= x0 + 1 and (stop_x is None or w[0] >= stop_x)]
    right = [w for w in band if w[0] >= x1 - 1]
    sides = (right, left) if prefer_right else (left, right)
    for side in sides:
        if not side:
            continue
        near = min(side, key=lambda w: (x0 - w[2]) if w[2] <= x0 + 1 else (w[0] - x1))
        same = [w for w in side if w[5] == near[5]]
        if side is left:
            # csak a doboz előtti, összefüggő szakasz (egy sorban több mező is lehet)
            same = [w for w in same if w[2] <= x0 + 1]
            same.sort(key=lambda w: w[0])
            chunk: list[tuple] = []
            for w in reversed(same):
                if chunk and chunk[0][0] - w[2] > 3.5 * h:
                    break
                chunk.insert(0, w)
            same = chunk
        else:
            same = [w for w in same if w[0] - x1 < 25 * h][:6]
        text = S.clean_label(" ".join(w[4] for w in same))
        if text and (side is left and x0 - same[-1][2] < 30 * h or side is right):
            return text
    above = [w for w in words if 0 <= y0 - w[3] <= 2.4 * h and w[0] < x1 and w[2] > x0 and not _is_blank_word(w[4])]
    if above:
        key = max(above, key=lambda w: w[3])[5]
        return S.clean_label(" ".join(w[4] for w in sorted([w for w in above if w[5] == key], key=lambda w: w[0])))
    return ""


def _is_blank_word(t: str) -> bool:
    return bool(re.fullmatch(r"[_.…\-–]{3,}|\[\s*\]|☐|□", t.strip()))


def _occupied(words: list[tuple], x0: float, x1: float, y: float, h: float) -> bool:
    """Van-e már írás a vonalon (kitöltött vagy aláhúzott szöveg) — akkor nem üres mező."""
    cover = 0.0
    for w in words:
        if _is_blank_word(w[4]):
            continue
        if -0.3 * h <= y - w[3] <= 1.2 * h and w[0] < x1 and w[2] > x0:
            cover += min(w[2], x1) - max(w[0], x0)
    return cover > 0.25 * (x1 - x0)


# ---------------------------------------------------------------- közös: táblázatok

numeric = S.numeric_label


def table_fields(rows: list[list[Optional[str]]]) -> list[tuple[int, int, str]]:
    """Egy táblázat mezői: CSAK üres cella lehet mező, szöveges cella soha. → [(sor, oszlop, felirat)].

    rows: cellaszövegek soronként; None = egyesített cella folytatása (nem önálló cella). Két űrlap-alak:
      felirat + üres cella mellette   „Név | ⎵” — a bal cella felirat (nem szám), a jobb üres
      fejléc + üres sorok             az első sor felirat-fejléc, alatta üres (vagy csak az első oszlopban kitöltött)
                                      sorok — pl. jelenléti ív napi sorai; a felirat: fejléc (+ a sor első cellája)
    A kitöltött adattábla (minden sorban adat) üres cellája sem mező."""
    def txt(r: int, c: int) -> str:
        v = rows[r][c] if c < len(rows[r]) else None
        return (v or "").strip()

    def empty(r: int, c: int) -> bool:
        return c < len(rows[r]) and rows[r][c] is not None and not BLANK_RE.sub("", txt(r, c)).strip()

    out: list[tuple[int, int, str]] = []
    seen: set[tuple[int, int]] = set()
    n_cols = max((len(r) for r in rows), default=0)
    header = [txt(0, c) for c in range(n_cols)] if rows else []
    body = range(1, len(rows))
    # fejléc: az első oszlop után minden oszlopnak van (nem szám) felirata — a „Név | ⎵” kétoszlopos tábla nem ilyen
    gridlike = len(rows) >= 2 and n_cols >= 2 and all(header[c] and not numeric(header[c]) for c in range(1, n_cols)) and any(
        all(empty(r, c) or (c == 0 and txt(r, 0)) for c in range(n_cols)) and any(empty(r, c) for c in range(1, n_cols))
        for r in body)
    if gridlike:
        for r in body:
            lab0 = txt(r, 0)
            for c in range(n_cols):
                if empty(r, c) and header[c] and not numeric(header[c]) and (r, c) not in seen:
                    seen.add((r, c))
                    out.append((r, c, f"{header[c]} – {lab0}" if lab0 and c else f"{header[c]} {r}"))
        return out
    for r in range(len(rows)):
        for c in range(n_cols - 1):
            lab = S.clean_label(txt(r, c))
            if lab and len(lab) <= 80 and not numeric(lab) and not BLANK_RE.search(txt(r, c)) and empty(r, c + 1):
                out.append((r, c + 1, lab))
    return out


def _merge_segs(segs: list[tuple], tol: float) -> list[list[float]]:
    """Egy egyenesbe eső, összeérő szakaszok (a0, a1, c) összevonva."""
    out: list[list[float]] = []
    for a0, a1, c in sorted(segs, key=lambda s: (round(s[2] / tol), s[0])):
        if out and abs(out[-1][2] - c) <= tol and a0 <= out[-1][1] + tol:
            out[-1][1] = max(out[-1][1], a1)
        else:
            out.append([a0, a1, c])
    return out


def _cluster(vals: list[float], tol: float) -> list[float]:
    out: list[float] = []
    for v in sorted(vals):
        if not out or v - out[-1] > tol:
            out.append(v)
    return out


def find_tables(hlines: list[tuple], vlines: list[tuple], tol: float = 2.5) -> list[dict[str, Any]]:
    """Rács a vonalakból: egymást metsző vízszintes (x0, x1, y) és függőleges (y0, y1, x) vonalak, ≥ 2 cella.
    → [{"xs": oszlophatárok, "ys": sorhatárok, "box": (x0, y0, x1, y1)}]"""
    H, V = _merge_segs(hlines, tol), _merge_segs(vlines, tol)
    if len(H) < 2 or len(V) < 2:
        return []
    parent = list(range(len(H) + len(V)))

    def root(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
    for i, (x0, x1, y) in enumerate(H):
        for j, (y0, y1, x) in enumerate(V):
            if x0 - tol <= x <= x1 + tol and y0 - tol <= y <= y1 + tol:
                parent[root(i)] = root(len(H) + j)
    comps: dict[int, tuple[list, list]] = {}
    for i, h in enumerate(H):
        comps.setdefault(root(i), ([], []))[0].append(h)
    for j, v in enumerate(V):
        comps.setdefault(root(len(H) + j), ([], []))[1].append(v)
    out = []
    for hs, vs in comps.values():
        ys, xs = _cluster([h[2] for h in hs], tol), _cluster([v[2] for v in vs], tol)
        if len(ys) >= 2 and len(xs) >= 2 and (len(ys) - 1) * (len(xs) - 1) >= 2:
            out.append({"xs": xs, "ys": ys, "box": (xs[0], ys[0], xs[-1], ys[-1])})
    return out


def _inside(box: tuple, t: dict[str, Any], tol: float = 3.0) -> bool:
    x0, y0, x1, y1 = t["box"]
    return box[0] >= x0 - tol and box[2] <= x1 + tol and box[1] >= y0 - tol and box[3] <= y1 + tol


def fields_from_layout(words: list[tuple], hlines: list[tuple], squares: list[tuple], boxes: list[tuple],
                       page: int, kind_anchor: str, vlines: Optional[list[tuple]] = None) -> list[dict[str, Any]]:
    """words: (x0, y0, x1, y1, szöveg, sorkulcs); hlines: (x0, x1, y); vlines: (y0, y1, x); squares/boxes: (x0, y0, x1, y1).
    A vonalakból összeálló táblázat cellái a table_fields szabálya szerint lesznek mezők (szöveges cella soha);
    a rácsvonalak nem „aláhúzások”, a cellakeretek nem „üres keretek”."""
    hs = [w[3] - w[1] for w in words if w[4].strip() and not _is_blank_word(w[4])]
    h = statistics.median(hs) if hs else 12.0
    out: list[dict[str, Any]] = []
    taken: list[tuple] = []
    tables = find_tables(hlines + [(b[0], b[2], y) for b in boxes for y in (b[1], b[3])],
                         (vlines or []) + [(b[1], b[3], x) for b in boxes for x in (b[0], b[2])])

    def add(label: str, typ: str, box: tuple, extra: Optional[dict] = None):
        for t in taken:      # ugyanaz a hely kétszer (aláhúzás-szó + vonal)
            if abs(t[0] - box[0]) < h and abs(t[3] - box[3]) < h * 0.8 and abs(t[2] - box[2]) < 2 * h:
                return
        taken.append(box)
        num = numeric(label)
        if num:                                  # „120” mellett üres hely: adat, nem felirat — a minőség sem „good”
            label = ""
        label = label or f"Field {len(out) + 1}"
        out.append(S.field(S.slug(label), label, typ, page=page, box=list(box),
                           anchor={"kind": kind_anchor, "page": page, "rect": [round(v, 2) for v in box], **(extra or {})}))
        if num:
            out[-1]["_numeric_label"] = True

    # 0) táblázatok: csak az űrlapjellegű tábla üres cellái
    for ti, t in enumerate(tables):
        xs, ys = t["xs"], t["ys"]
        grid: list[list[Optional[str]]] = [[None] * (len(xs) - 1) for _ in range(len(ys) - 1)]
        for r in range(len(ys) - 1):
            for c in range(len(xs) - 1):
                inside = [w for w in words if xs[c] <= (w[0] + w[2]) / 2 < xs[c + 1] and ys[r] <= (w[1] + w[3]) / 2 < ys[r + 1]]
                grid[r][c] = " ".join(w[4] for w in sorted(inside, key=lambda w: (w[5], w[0])))
        for r, c, lab in table_fields(grid):
            box = (xs[c] + 1.5, ys[r] + 1.5, xs[c + 1] - 1.5, ys[r + 1] - 1.5)
            add(lab, S.guess_type(lab, "multiline" if box[3] - box[1] > 3 * h else "text"), box,
                {"table": ti, "row": r, "col": c})
    hlines = [l for l in hlines if not any(_inside((l[0], l[2], l[1], l[2]), t) for t in tables)]
    boxes = [b for b in boxes if not any(_inside(b, t) for t in tables)]

    # 1) aláhúzás- / pont-sorok a szövegben ("Név: ________")
    for w in words:
        if re.fullmatch(r"_{3,}|\.{5,}|…{2,}", w[4].strip()):
            box = (w[0], w[3] - 1.5 * h, w[2], w[3])
            lab = _label_near(words, box, h)
            add(lab, S.guess_type(lab), box)
        elif re.fullmatch(r"☐|□", w[4].strip()):
            box = (w[0], w[1], w[2], w[3])
            lab = _label_near(words, box, h, prefer_right=True)
            add(lab, "checkbox", box)
    # 2) vízszintes vonalak
    for x0, x1, y in sorted(hlines, key=lambda l: (l[2], l[0])):
        if x1 - x0 < 2.5 * h or _occupied(words, x0, x1, y, h):
            continue
        box = (x0, y - 1.5 * h, x1, y - 0.5)
        lab = _label_near(words, box, h)
        add(lab, S.guess_type(lab), box)
    # 3) jelölőnégyzet-keretek
    for b in squares:
        if any(w[0] >= b[0] - 1 and w[2] <= b[2] + 1 and w[1] >= b[1] - 1 and w[3] <= b[3] + 1 and not _is_blank_word(w[4]) for w in words):
            continue
        lab = _label_near(words, b, h, prefer_right=True)
        add(lab, "checkbox", b)
    # 4) üres keretek (szövegdoboz)
    for b in boxes:
        if any(w[0] < b[2] and w[2] > b[0] and w[1] < b[3] and w[3] > b[1] for w in words):
            continue
        lab = _label_near(words, b, h)
        typ = S.guess_type(lab, "multiline" if b[3] - b[1] > 3 * h else "text")
        add(lab, typ, b)
    out.sort(key=lambda f: (f["box"][1] // max(h, 1), f["box"][0]))
    return out


# ---------------------------------------------------------------- PDF

def _pdf_widgets(doc) -> list[dict[str, Any]]:
    import pymupdf
    types = {pymupdf.PDF_WIDGET_TYPE_CHECKBOX: "checkbox", pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON: "choice",
             pymupdf.PDF_WIDGET_TYPE_COMBOBOX: "choice", pymupdf.PDF_WIDGET_TYPE_LISTBOX: "choice",
             pymupdf.PDF_WIDGET_TYPE_SIGNATURE: "signature", pymupdf.PDF_WIDGET_TYPE_TEXT: "text"}
    out: list[dict[str, Any]] = []
    seen: set[str] = set()
    for pi, page in enumerate(doc, 1):
        words = [(*w[:5], (w[5], w[6])) for w in page.get_text("words")]
        hs = [w[3] - w[1] for w in words]
        h = statistics.median(hs) if hs else 12.0
        for wd in page.widgets() or []:
            name = wd.field_name or f"field_{len(out) + 1}"
            if name in seen and wd.field_type != pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON:
                continue
            r = wd.rect
            label = (wd.field_label or "").strip() or _label_near(words, (r.x0, r.y0, r.x1, r.y1), h,
                                                                   prefer_right=wd.field_type == pymupdf.PDF_WIDGET_TYPE_CHECKBOX)
            typ = types.get(wd.field_type, "text")
            if typ == "text":
                typ = "multiline" if wd.field_flags & pymupdf.PDF_TX_FIELD_IS_MULTILINE else S.guess_type(label or name)
                if typ == "checkbox":
                    typ = "text"
            if name in seen:      # rádiógomb: a csoport egy mező, az opciók a gombok értékei
                f = next(x for x in out if x["name"] == name)
                on = [v for v in (wd.button_states() or {}).get("normal", []) if v != "Off"]
                f["options"] = sorted(set(f["options"]) | set(on))
                continue
            seen.add(name)
            opts = list(wd.choice_values or []) if typ == "choice" else []
            if wd.field_type == pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON:
                opts = [v for v in (wd.button_states() or {}).get("normal", []) if v != "Off"]
            val = wd.field_value
            if typ == "checkbox":
                val = val not in (None, False, "", "Off", "off")
            out.append(S.field(name, label or name, typ, page=pi, box=[r.x0, r.y0, r.x1, r.y1], options=opts,
                               value=val if val not in ("", None) else None, required=bool(wd.field_flags & 2),
                               anchor={"kind": "acroform", "field": name}))
    return out


def _pdf_flat(doc) -> tuple[list[dict[str, Any]], int]:
    """Vissza: (mezők, a „Felirat:” formák száma — a minőség-becsléshez)."""
    out: list[dict[str, Any]] = []
    hints = 0
    for pi, page in enumerate(doc, 1):
        words = [(*w[:5], (w[5], w[6])) for w in page.get_text("words")]
        hlines, vlines, squares, boxes = [], [], [], []
        for d in page.get_drawings():
            for it in d["items"]:
                if it[0] == "l":
                    p1, p2 = it[1], it[2]
                    if abs(p1.y - p2.y) < 1.5 and abs(p2.x - p1.x) >= 20:
                        hlines.append((min(p1.x, p2.x), max(p1.x, p2.x), (p1.y + p2.y) / 2))
                    elif abs(p1.x - p2.x) < 1.5 and abs(p2.y - p1.y) >= 6:
                        vlines.append((min(p1.y, p2.y), max(p1.y, p2.y), (p1.x + p2.x) / 2))
                elif it[0] in ("re", "qu"):
                    r = it[1] if it[0] == "re" else it[1].rect
                    if r.height < 3 and r.width >= 20:
                        hlines.append((r.x0, r.x1, r.y1))
                    elif r.width < 3 and r.height >= 6:
                        vlines.append((r.y0, r.y1, (r.x0 + r.x1) / 2))
                    elif 5 <= r.width <= 24 and 5 <= r.height <= 24 and abs(r.width - r.height) < 4:
                        squares.append((r.x0, r.y0, r.x1, r.y1))
                    elif r.width >= 40 and 10 <= r.height <= 200 and r.width < page.rect.width * 0.95:
                        boxes.append((r.x0, r.y0, r.x1, r.y1))
        out += fields_from_layout(words, hlines, squares, boxes, pi, "xy", vlines)
        hints += _label_hints(words)
    return out, hints


def recognize_pdf(path: Path) -> dict[str, Any]:
    import pymupdf
    with pymupdf.open(str(path)) as doc:
        if doc.needs_pass:
            raise ValueError("The PDF is password-protected — unlock it first.")
        pages = [{"page": i + 1, "width": round(p.rect.width, 2), "height": round(p.rect.height, 2)} for i, p in enumerate(doc)]
        fields = _pdf_widgets(doc)
        if fields:
            # az aláírás helye gyakran csak egy vonal ("Aláírás: ____") — azt is felvesszük, ha nincs ott mező
            for f in _pdf_flat(doc)[0]:
                f.pop("_numeric_label", None)
                if f["type"] == "signature" and not any(
                        g["page"] == f["page"] and g["box"][0] < f["box"][2] and g["box"][2] > f["box"][0]
                        and g["box"][1] < f["box"][3] and g["box"][3] > f["box"][1] for g in fields):
                    fields.append(f)
            return S.make("pdf", "acroform", pages, fields)
        has_text = any(p.get_text().strip() for p in doc)
        fields, hints = _pdf_flat(doc)
    return S.make("pdf", "flat_pdf", pages, fields, label_hints=hints, numeric_labels=_strip_numeric(fields),
                  **({} if has_text else {"scanned": True}))


def _strip_numeric(fields: list[dict[str, Any]]) -> int:
    """Hány mező mellett állt felirat helyett csak szám (a minőség-becsléshez); a belső jelölő kikerül."""
    return sum(1 for f in fields if f.pop("_numeric_label", False))


# ---------------------------------------------------------------- DOCX

def _docx_paragraphs(doc):
    """(hely, bekezdés) párok: törzs-bekezdések és táblázatcellák bekezdései (egyesített cella egyszer)."""
    for i, p in enumerate(doc.paragraphs):
        yield {"p": i}, p
    for ti, t in enumerate(doc.tables):
        seen = set()
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                for pi, p in enumerate(cell.paragraphs):
                    yield {"t": ti, "r": ri, "c": ci, "p": pi}, p


def _sdts(doc) -> list:
    from docx.oxml.ns import qn
    return list(doc.element.body.iter(qn("w:sdt")))


def recognize_docx(path: Path) -> dict[str, Any]:
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(str(path))
    fields: list[dict[str, Any]] = []
    # tartalomvezérlők
    for i, sdt in enumerate(_sdts(doc)):
        pr = sdt.find(qn("w:sdtPr"))
        if pr is None:
            continue
        alias = pr.find(qn("w:alias"))
        tag = pr.find(qn("w:tag"))
        label = (alias.get(qn("w:val")) if alias is not None else None) or (tag.get(qn("w:val")) if tag is not None else None)
        if not label:
            continue
        typ, opts = S.guess_type(label), []
        if any(c.tag.endswith("}checkbox") for c in pr):
            typ = "checkbox"
        elif pr.find(qn("w:date")) is not None:
            typ = "date"
        elif pr.find(qn("w:dropDownList")) is not None or pr.find(qn("w:comboBox")) is not None:
            typ = "choice"
            lst = pr.find(qn("w:dropDownList")) if pr.find(qn("w:dropDownList")) is not None else pr.find(qn("w:comboBox"))
            opts = [li.get(qn("w:displayText")) or li.get(qn("w:value")) for li in lst.findall(qn("w:listItem"))]
        fields.append(S.field(S.slug(label), label, typ, options=opts, anchor={"kind": "docx_sdt", "index": i}))
    # minták a bekezdésekben
    prev_text = ""
    for loc, p in _docx_paragraphs(doc):
        text = p.text
        last = 0
        for n, m in enumerate(BLANK_RE.finditer(text)):
            tok = m.group(0)
            before = S.clean_label(text[last:m.start()])
            if m.group(1):                                  # {{név}}
                label, typ = m.group(1).strip(), S.guess_type(m.group(1))
            elif tok in ("☐", "□"):
                after = S.clean_label(BLANK_RE.split(text[m.end():])[0])
                label, typ = (before or after or prev_text), "checkbox"
            else:
                label = before or S.clean_label(prev_text)
                typ = S.guess_type(label)
            last = m.end()
            fields.append(S.field(S.slug(label or f"field {len(fields) + 1}"), label or f"Field {len(fields) + 1}", typ,
                                  anchor={"kind": "docx_text", "loc": loc, "token": tok, "nth": n}))
        if text.strip():
            prev_text = text
    # táblázat: csak üres cella, és csak űrlapjellegű táblában (felirat + üres cella, vagy fejléc + üres sorok)
    for ti, t in enumerate(doc.tables):
        grid: list[list[Optional[str]]] = []
        for row in t.rows:
            cells = row.cells
            grid.append([None if ci and cells[ci]._tc is cells[ci - 1]._tc else cells[ci].text for ci in range(len(cells))])
        for r, c, lab in table_fields(grid):
            fields.append(S.field(S.slug(lab), lab, S.guess_type(lab), anchor={"kind": "docx_cell", "t": ti, "r": r, "c": c}))
    return S.make("docx", "docx", [], fields)


# ---------------------------------------------------------------- XLSX

def recognize_xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    fields: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        merged_tl = {}
        for rng in ws.merged_cells.ranges:
            for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row, min_col=rng.min_col, max_col=rng.max_col):
                for c in row:
                    merged_tl[c.coordinate] = ws.cell(rng.min_row, rng.min_col).coordinate
        dv_opts: dict[str, list[str]] = {}
        for dv in ws.data_validations.dataValidation:
            if dv.type == "list" and dv.formula1 and dv.formula1.startswith('"'):
                opts = [o.strip() for o in dv.formula1.strip('"').split(",")]
                for rng in dv.sqref.ranges:
                    for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row, min_col=rng.min_col, max_col=rng.max_col):
                        for c in row:
                            dv_opts[c.coordinate] = opts
        cands: list[tuple] = []
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and re.fullmatch(r"\s*_{3,}\s*", v):
                    left = ws.cell(c.row, c.column - 1).value if c.column > 1 else None
                    cands.append((str(left or f"Field {len(cands) + 1}"), c.coordinate, c.column - 1, True))
                    continue
                if not isinstance(v, str) or not v.strip() or len(v) > 60 or re.fullmatch(r"[\d\s.,%-]+", v):
                    continue
                right = ws.cell(c.row, c.column + 1)
                coord = merged_tl.get(right.coordinate, right.coordinate)
                if coord == merged_tl.get(c.coordinate, c.coordinate):
                    continue
                if ws[coord].value in (None, ""):
                    cands.append((v, coord, c.column, v.strip().endswith(":") or coord in dv_opts))
        col_count: dict[int, int] = {}
        for _, _, col, _ in cands:
            col_count[col] = col_count.get(col, 0) + 1
        for label, coord, col, strong in cands:
            if not strong and col_count.get(col, 0) < 2:
                continue
            lab = S.clean_label(label)
            typ = "choice" if coord in dv_opts else S.guess_type(lab)
            fields.append(S.field(S.slug(lab), lab, typ, options=dv_opts.get(coord, []),
                                  anchor={"kind": "xlsx_cell", "sheet": ws.title, "cell": coord}))
    return S.make("xlsx", "xlsx", [{"page": i + 1, "sheet": ws.title} for i, ws in enumerate(wb.worksheets)], fields)


# ---------------------------------------------------------------- fotó, szken

async def _tesseract_words(img: Path, lang: str) -> tuple[list[tuple], Optional[float]]:
    from app.convert.engines import run
    from app.convert.service import ocr_languages
    have = ocr_languages()
    langs = "+".join(l for l in lang.split("+") if l in have) or ("eng" if "eng" in have else "")
    out, _ = await run("tesseract", str(img), "stdout", *(["-l", langs] if langs else []), "--psm", "11", "tsv", timeout=300)
    words, confs = [], []
    for ln in out.decode("utf-8", "replace").splitlines()[1:]:
        p = ln.split("\t")
        if len(p) < 12 or not p[11].strip() or float(p[10]) < 0:
            continue
        confs.append(float(p[10]))
        if float(p[10]) < 30:
            continue
        x, y, w, h = (int(p[i]) for i in (6, 7, 8, 9))
        words.append((x, y, x + w, y + h, p[11], (p[2], p[3], p[4])))
    return words, (sum(confs) / len(confs) if confs else None)


def _cv_shapes(img: Path) -> tuple[list, list, list, list]:
    import cv2
    import numpy as np
    g = cv2.imdecode(np.frombuffer(img.read_bytes(), np.uint8), cv2.IMREAD_GRAYSCALE)
    H, W = g.shape
    bw = cv2.adaptiveThreshold(g, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 25, 15)
    horiz = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (max(20, W // 25), 1)))
    hlines = []
    for c in cv2.findContours(horiz, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)[0]:
        x, y, w, h = cv2.boundingRect(c)
        if h <= max(6, H // 300) and w >= W // 20:
            hlines.append((x, x + w, y + h / 2))
    vert = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(15, H // 60))))
    vlines = []
    for c in cv2.findContours(vert, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)[0]:
        x, y, w, h = cv2.boundingRect(c)
        if w <= max(6, W // 300) and h >= H // 60:
            vlines.append((y, y + h, x + w / 2))
    squares, boxes = [], []
    for c in cv2.findContours(bw, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)[0]:
        x, y, w, h = cv2.boundingRect(c)
        ap = cv2.approxPolyDP(c, 0.04 * cv2.arcLength(c, True), True)
        if len(ap) != 4:
            continue
        fill = cv2.countNonZero(bw[y + 3:y + h - 3, x + 3:x + w - 3]) / max(1, (w - 6) * (h - 6))
        if W / 120 <= w <= W / 18 and abs(w - h) <= 0.25 * w and fill < 0.15:
            squares.append((x, y, x + w, y + h))
        elif w >= W / 8 and H / 60 <= h <= H / 4 and fill < 0.05 and w < 0.95 * W:
            boxes.append((x, y, x + w, y + h))
    return hlines, squares, boxes, vlines


async def recognize_image(path: Path, language: str = "hun+eng") -> dict[str, Any]:
    """Fotó/szken helyben, ingyen: Tesseract-szavak + OpenCV-vonalak/-négyzetek (a látó modell: deep.py)."""
    from PIL import Image, ImageOps
    from app.convert.steps import _register_plugins
    _register_plugins()
    work = path.with_name(path.stem + "-form.png")
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im).convert("RGB")
        im.save(work)
        W, H = im.size
    try:
        words, conf = await _tesseract_words(work, language)
        hlines, squares, boxes, vlines = await asyncio.to_thread(_cv_shapes, work)
    finally:
        work.unlink(missing_ok=True)
    fields = fields_from_layout(words, hlines, squares, boxes, 1, "xy", vlines)
    return S.make("image", "image", [{"page": 1, "width": W, "height": H}], fields, engine="offline",
                  ocr_confidence=round(conf, 1) if conf is not None else None, label_hints=_label_hints(words),
                  numeric_labels=_strip_numeric(fields))


def _label_hints(words: list[tuple]) -> int:
    """Hány „Felirat:” forma van a lapon (a minőség-becsléshez: ennyi körüli mezőt várunk)."""
    return sum(1 for w in words if len(w[4]) > 2 and w[4].endswith(":") and not _is_blank_word(w[4]))


# ----------------------------------------------------------------

async def recognize(path: Path, language: str = "hun+eng", deep: str = "auto") -> dict[str, Any]:
    """deep: "auto" (gyenge vagy összetett eredménynél, ha jár: fizetős látó modell), "on" (kérve), "off"."""
    fmt = from_filename(path.name) or path.suffix.lstrip(".").lower()
    cat = FORMATS[fmt].category if fmt in FORMATS else ""
    if fmt == "pdf":
        sch = await asyncio.to_thread(recognize_pdf, path)
        if sch.get("kind") == "flat_pdf" and sch.get("scanned"):
            sch = await _scanned_pdf(path, sch, language)
        return await _maybe_deep(sch, path, language, deep)
    if fmt in ("docx", "docm", "dotx"):
        return await asyncio.to_thread(recognize_docx, path)
    if fmt in ("xlsx", "xlsm"):
        return await asyncio.to_thread(recognize_xlsx, path)
    if cat == "image":
        return await _maybe_deep(await recognize_image(path, language), path, language, deep)
    if cat in ("document", "spreadsheet"):
        raise ValueError(f".{fmt}: convert it to {'DOCX' if cat == 'document' else 'XLSX'} or PDF first, then recognize")
    raise ValueError(f".{fmt} files can't hold a form")


async def _scanned_pdf(path: Path, sch: dict[str, Any], language: str) -> dict[str, Any]:
    """Képből álló PDF: oldalanként (legfeljebb 10) helyi felismerés, a dobozok pontra váltva."""
    import pymupdf
    from app.formkit.deep import MAX_PAGES
    fields, confs, hints, nums = [], [], 0, 0
    with pymupdf.open(str(path)) as d:
        n = min(d.page_count, MAX_PAGES)
        for i in range(n):
            img = path.with_name(f"{path.stem}-p{i + 1}.png")
            d[i].get_pixmap(dpi=150).save(str(img))
            try:
                im = await recognize_image(img, language)
            finally:
                img.unlink(missing_ok=True)
            scale = sch["pages"][i]["width"] / im["pages"][0]["width"]
            for f in im["fields"]:
                f["box"] = [round(v * scale, 2) for v in f["box"]]
                f["page"] = i + 1
                f["anchor"] = {"kind": "xy", "page": i + 1, "rect": f["box"]}
                fields.append(f)
            if im.get("ocr_confidence") is not None:
                confs.append(im["ocr_confidence"])
            hints += im.get("label_hints") or 0
            nums += im.get("numeric_labels") or 0
    sch["fields"] = S.unique_names(fields)
    sch["numeric_labels"] = nums
    sch["engine"] = "offline"
    sch["ocr_confidence"] = round(sum(confs) / len(confs), 1) if confs else None
    sch["label_hints"] = hints
    return sch


async def _maybe_deep(sch: dict[str, Any], src: Path, language: str, deep: str) -> dict[str, Any]:
    """Gyenge vagy összetett eredmény → fizetős, több lépéses látó modell (deep.py); MCP-n a kliens ügynöke."""
    import shutil
    from app.formkit import deep as D
    if sch["kind"] not in ("flat_pdf", "image"):
        return sch
    q = D.quality(sch)
    sch["quality"] = q
    if deep == "off" or not (deep == "on" or q["level"] == "low" or q["complex"]):
        return sch
    ok, why = await D.eligibility()
    if not ok:
        if deep == "on":
            raise D.DeepUnavailable(D.why_text(why))
        sch["deep"] = {"used": False, "reason": why, "hint": D.why_text(why)}
        return sch
    from app.storage import temp_path
    from app.storage.files import new_file_id
    work = temp_path(new_file_id("deep"))
    work.mkdir(parents=True)
    try:
        pages: list[tuple[Path, float, float]] = []
        if sch["kind"] == "image":
            pages.append((src, sch["pages"][0]["width"], sch["pages"][0]["height"]))
        else:
            import pymupdf
            with pymupdf.open(str(src)) as d:
                for i in range(min(d.page_count, D.MAX_PAGES)):
                    img = work / f"page{i + 1}.png"
                    d[i].get_pixmap(dpi=150).save(str(img))
                    pages.append((img, d[i].rect.width, d[i].rect.height))
        hint = {"hun": "Hungarian", "eng": "English", "deu": "German"}.get(language.split("+")[0])
        out = await D.recognize(sch, pages, work, hint)
        out["quality"] = D.quality(out)
        return out
    finally:
        shutil.rmtree(work, ignore_errors=True)
