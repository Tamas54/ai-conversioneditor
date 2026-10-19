"""Kitöltés az eredeti formátumban, a séma (schema.py) horgonyai szerint.

    acroform   valódi PDF-mezők (PyMuPDF), a megjelenés Unicode-betűvel (ő, ű); a mező betűje utána is a
               beágyazott DejaVu (forms.unicode_fields), így a néző újrarajzolva sem veszít ékezetet;
               flatten → forms._bake (nyomtatható, nem szerkeszthető)
    xy         lapos PDF / szken / fotó: a szöveg a mező dobozába kerül (beférő betűmérettel), jelölő → ✓,
               aláírás → kép
    docx       a minta ("____", "[   ]", ☐, {{név}}) helyére írunk a futás (run) formázásával; cella; tartalomvezérlő
    xlsx       a cella értéke (dátum dátumként, szám számként), a cella formázása marad
    image      keep_look=True: a fotóra írunk (ugyanaz a formátum); False: tiszta PDF (a kép + kereshető szöveg)
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Any

from app.formkit import schema as S
from app.formkit import validate as V

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
TRUE = ("1", "true", "yes", "igen", "x", "✓", "✔", "on", "y", "i")


def truthy(v: Any) -> bool:
    return v if isinstance(v, bool) else str(v).strip().lower() in TRUE


def match_values(sch: dict[str, Any], values: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    """A kulcs lehet mezőnév vagy címke, kis-nagybetű és ékezet nélkül is (forms._coerce-szal egyezően)."""
    norm = lambda s: re.sub(r"[^a-z0-9]+", "", S.ascii_fold(s))
    by: dict[str, str] = {}
    for f in sch["fields"]:
        by.setdefault(norm(f["name"]), f["name"])
        by.setdefault(norm(f.get("label") or ""), f["name"])
    out, skipped = {}, []
    for k, v in (values or {}).items():
        name = by.get(norm(k))
        if name is None and len(norm(k)) >= 4:      # rövidített felirat: egyértelmű előtag-egyezés
            hits = {n for key, n in by.items() if key and (key.startswith(norm(k)) or norm(k).startswith(key) and len(key) >= 4)}
            name = hits.pop() if len(hits) == 1 else None
        if name is None:
            skipped.append(str(k))
        elif v is not None and v != "":
            out[name] = v
    return out, skipped


def prepare(sch: dict[str, Any], values: dict[str, Any], date_style: str = "auto") -> tuple[dict[str, Any], dict[str, str], list[str]]:
    """Értékek → (egységesített értékek, hibák mezőnként, ismeretlen kulcsok)."""
    vals, skipped = match_values(sch, values)
    errors: dict[str, str] = {}
    fields = {f["name"]: f for f in sch["fields"]}
    hu = date_style == "hu" or (date_style == "auto" and sch.get("language", "hu").startswith("hu"))
    for name, v in list(vals.items()):
        f = fields[name]
        if f["type"] == "checkbox":
            vals[name] = truthy(v)
            continue
        if f["type"] == "signature" or isinstance(v, dict):
            continue
        if f["type"] == "choice" and f.get("options") and str(v) not in f["options"]:
            m = [o for o in f["options"] if S.ascii_fold(o) == S.ascii_fold(v)]
            if m:
                vals[name] = m[0]
            else:
                errors[name] = f"choose one of: {', '.join(f['options'][:20])}"
                vals.pop(name)
            continue
        kind = f.get("check")
        if kind == "date":
            ok, nv, msg = V.check_date(v, "hu" if hu else "iso")
        else:
            ok, nv, msg = V.check(kind, v)
        if ok:
            vals[name] = nv
        else:                       # hibás értéket nem írunk be: a hibát mezőnként visszaadjuk
            errors[name] = msg or "invalid"
            vals.pop(name)
    return vals, errors, skipped


# ---------------------------------------------------------------- PDF: valódi mezők

def _fill_acroform(src: Path, dst: Path, sch: dict[str, Any], vals: dict[str, Any], flatten: bool) -> list[str]:
    import pikepdf
    import pymupdf
    from app.tools import forms
    work = dst.with_suffix(".work.pdf")
    shutil.copy(src, work)
    try:
        forms.unicode_fields(work)
    except Exception:
        pass
    by_anchor = {f["anchor"].get("field"): f for f in sch["fields"] if f["anchor"].get("kind") == "acroform"}
    filled: list[str] = []
    sizes: dict[str, float] = {}
    sig_xy: list[dict[str, Any]] = []
    with pymupdf.open(str(work)) as d:
        for page in d:
            for w in page.widgets() or []:
                f = by_anchor.get(w.field_name)
                if not f or f["name"] not in vals:
                    continue
                v = vals[f["name"]]
                if f["type"] == "signature":
                    if isinstance(v, dict) and v.get("image"):
                        sig_xy.append({"page": page.number + 1, "box": list(w.rect), "image": v["image"]})
                        filled.append(f["name"])
                    continue
                if w.field_type == pymupdf.PDF_WIDGET_TYPE_CHECKBOX:
                    w.field_value = w.on_state() if truthy(v) else "Off"
                elif w.field_type == pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON:
                    if str(v) not in (w.button_states() or {}).get("normal", []):
                        continue
                    w.field_value = str(v)
                else:
                    if not w.text_fontsize:
                        multi = w.field_flags & pymupdf.PDF_TX_FIELD_IS_MULTILINE
                        w.text_fontsize = 10 if multi else max(6.0, min(11.0, w.rect.height * 0.7))
                    sizes[w.field_name] = w.text_fontsize
                    w.field_value = str(v)
                w.update()
                filled.append(f["name"])
        for s in sig_xy:
            from app.formkit.signature import stamp_pdf_page
            stamp_pdf_page(d[s["page"] - 1], s["box"], Path(s["image"]))
        d.save(str(dst), garbage=3, deflate=True)
    work.unlink(missing_ok=True)
    # a mező betűje vissza a beágyazott Unicode-betűre (a PyMuPDF Helv-re írja át), és a megjelenést is mi
    # rajzoljuk vele: a MuPDF a nem latin-1 betűkhöz tartalékbetűt vesz, és az alapvonalat a dobozon kívülre teszi
    with pikepdf.open(str(dst), allow_overwriting_input=True) as pdf:
        af = pdf.Root.get("/AcroForm")
        fonts = af.get("/DR", {}).get("/Font", {}) if af is not None else {}
        if af is not None and "/UniF" in fonts:
            # a glifazonosítók a TÉNYLEGESEN beágyazott mezőbetűből (a gépelési részhalmaz más sorrendű, mint a teljes
            # DejaVu — azzal mérve a megjelenés eltolt betűket mutatna); ha nem olvasható, a teljes betű
            try:
                font = pymupdf.Font(fontbuffer=bytes(fonts.UniF.DescendantFonts[0].FontDescriptor.FontFile2.read_bytes()))
            except Exception:
                font = pymupdf.Font(fontfile=FONT)

            def fix(fld, name=None):
                name = str(fld.get("/T", name or ""))
                if name in sizes and (fld.get("/FT") == pikepdf.Name.Tx or "/FT" not in fld):
                    fld.DA = pikepdf.String(f"/UniF {sizes[name]:g} Tf 0 g")
                    if "/Rect" in fld:
                        _text_ap(pdf, fld, fonts.UniF, font, str(vals_by_field.get(name, "")), sizes[name])
                for k in fld.get("/Kids", []):
                    fix(k, name)
            vals_by_field = {f["anchor"].get("field"): vals[f["name"]] for f in sch["fields"]
                             if f["anchor"].get("kind") == "acroform" and f["name"] in vals}
            for fld in af.get("/Fields", []):
                fix(fld)
            af.NeedAppearances = False
        pdf.save(str(dst))
    if flatten:
        forms._bake(dst)
    return filled


def _text_ap(pdf, widget, font_ref, font, text: str, size: float) -> None:
    """Szövegmező megjelenése a beágyazott (Identity-H) DejaVu-val: glifazonosítók, sortörés többsoros mezőben."""
    import pikepdf
    r = [float(v) for v in widget.Rect]
    w, h = abs(r[2] - r[0]), abs(r[3] - r[1])
    multi = int(widget.get("/Ff", 0)) & 4096
    if not multi:
        while size > 5 and font.text_length(text, fontsize=size) > w - 4:
            size -= 0.5
        lines = [text]
    else:
        lines = []
        for para in text.splitlines() or [""]:
            cur = ""
            for word in para.split(" "):
                t = (cur + " " + word).strip()
                if cur and font.text_length(t, fontsize=size) > w - 4:
                    lines.append(cur)
                    cur = word
                else:
                    cur = t
            lines.append(cur)
    gid = lambda ch: font.has_glyph(ord(ch)) or font.has_glyph(ord("?"))
    ops = ["/Tx BMC", "q", f"1 1 {w - 2:.2f} {h - 2:.2f} re W n", "BT", f"/UniF {size:g} Tf", "0 g"]
    lead = size * 1.2
    y = (h - size) / 2 + 0.22 * size if not multi else h - 2 - size
    ops.append(f"2 {y:.2f} Td")
    for i, ln in enumerate(lines):
        if i:
            ops.append(f"0 {-lead:.2f} Td")
        ops.append("<" + "".join(f"{gid(c):04x}" for c in ln) + "> Tj")
    ops += ["ET", "Q", "EMC"]
    stream = pikepdf.Stream(pdf, "\n".join(ops).encode("ascii"))
    stream.Type, stream.Subtype = pikepdf.Name.XObject, pikepdf.Name.Form
    stream.BBox = [0, 0, w, h]
    stream.Resources = pikepdf.Dictionary(Font=pikepdf.Dictionary(UniF=font_ref))
    widget.AP = pikepdf.Dictionary(N=stream)


# ---------------------------------------------------------------- PDF/kép: dobozba írás

def _fit_size(font, text: str, w: float, h: float, max_size: float) -> float:
    size = max(5.0, min(max_size, h * 0.75))
    while size > 5 and font.text_length(text, fontsize=size) > w - 4:
        size -= 0.5
    return size


def _overlay_pdf(src: Path, dst: Path, sch: dict[str, Any], vals: dict[str, Any], scale: float = 1.0) -> list[str]:
    import pymupdf
    from app.tools import pdffonts
    buf = pdffonts.subset_for("".join(str(v) for v in vals.values() if not isinstance(v, dict)), FONT)   # csak a beírt jelek
    font = pymupdf.Font(fontbuffer=buf)
    filled = []
    with pymupdf.open(str(src)) as d:
        for f in sch["fields"]:
            if f["name"] not in vals or not f.get("box"):
                continue
            v = vals[f["name"]]
            page = d[(f.get("page") or 1) - 1]
            x0, y0, x1, y1 = (c * scale for c in f["box"])
            r = pymupdf.Rect(x0, y0, x1, y1)
            if f["type"] == "signature":
                if isinstance(v, dict) and v.get("image"):
                    from app.formkit.signature import stamp_pdf_page
                    stamp_pdf_page(page, list(r), Path(v["image"]))
                    filled.append(f["name"])
                continue
            page.insert_font(fontname="UniX", fontbuffer=buf)
            if f["type"] == "checkbox":
                if truthy(v):
                    s = min(r.width, r.height)
                    page.insert_text(pymupdf.Point(r.x0 + s * 0.12, r.y1 - s * 0.15), "✓", fontsize=s * 0.95, fontname="UniX")
                filled.append(f["name"])
                continue
            text = str(v)
            if f["type"] == "multiline" and r.height > 24:
                page.insert_textbox(r + (2, 2, -2, -2), text, fontsize=10, fontname="UniX")
            else:
                size = _fit_size(font, text, r.width, max(r.height, 12), 11)
                page.insert_text(pymupdf.Point(r.x0 + 2, r.y1 - max(2.0, r.height * 0.18)), text, fontsize=size, fontname="UniX")
            filled.append(f["name"])
        d.save(str(dst), garbage=3, deflate=True)
    return filled


def _overlay_image(src: Path, dst: Path, sch: dict[str, Any], vals: dict[str, Any]) -> list[str]:
    from PIL import Image, ImageDraw, ImageFont, ImageOps
    from app.convert.steps import _register_plugins
    _register_plugins()
    with Image.open(src) as im0:
        im = ImageOps.exif_transpose(im0).convert("RGB")
    dr = ImageDraw.Draw(im)
    ink = (20, 30, 90)
    filled = []
    for f in sch["fields"]:
        if f["name"] not in vals or not f.get("box"):
            continue
        v = vals[f["name"]]
        x0, y0, x1, y1 = f["box"]
        if f["type"] == "signature":
            if isinstance(v, dict) and v.get("image"):
                with Image.open(v["image"]) as sig:
                    sig = sig.convert("RGBA")
                    sig.thumbnail((int(x1 - x0), int(max(y1 - y0, 20) * 1.6)))
                    im.paste(sig, (int(x0), int(y1 - sig.height)), sig)
                filled.append(f["name"])
            continue
        if f["type"] == "checkbox":
            if truthy(v):
                s = int(min(x1 - x0, y1 - y0))
                fnt = ImageFont.truetype(FONT, max(8, int(s * 1.0)))
                dr.text((x0 + s * 0.08, y0 - s * 0.12), "✓", font=fnt, fill=ink)
            filled.append(f["name"])
            continue
        text = str(v)
        size = int(max(10, (y1 - y0) * 0.8))
        fnt = ImageFont.truetype(FONT, size)
        while size > 8 and dr.textlength(text, font=fnt) > (x1 - x0) - 6:
            size -= 1
            fnt = ImageFont.truetype(FONT, size)
        dr.text((x0 + 3, y1 - size * 1.1), text, font=fnt, fill=ink)
        filled.append(f["name"])
    fmt = (dst.suffix.lstrip(".") or "png").lower()
    im.save(dst, **({"quality": 92} if fmt in ("jpg", "jpeg") else {}))
    return filled


def _image_to_pdf(src: Path, dst: Path, dpi: int = 150) -> float:
    """Kép → egyoldalas PDF; vissza: pont/pixel arány (a dobozok átváltásához)."""
    import pymupdf
    from PIL import Image
    with Image.open(src) as im:
        w, h = im.size
    scale = 72 / dpi
    with pymupdf.open() as d:
        page = d.new_page(width=w * scale, height=h * scale)
        page.insert_image(page.rect, filename=str(src))
        d.save(str(dst))
    return scale


# ---------------------------------------------------------------- DOCX

def _docx_para(doc, loc: dict[str, Any]):
    if "t" in loc:
        return doc.tables[loc["t"]].cell(loc["r"], loc["c"]).paragraphs[loc["p"]]
    return doc.paragraphs[loc["p"]]


def _replace_in_runs(p, start: int, end: int, new: str) -> None:
    """A bekezdés [start, end) szakaszának cseréje úgy, hogy az első érintett futás formázása marad."""
    pos = 0
    first = True
    for run in p.runs:
        t = run.text
        a, b = pos, pos + len(t)
        pos = b
        if b <= start or a >= end:
            continue
        s0, s1 = max(start, a) - a, min(end, b) - a
        if first:
            run.text = t[:s0] + new + t[s1:]
            first = False
        else:
            run.text = t[:s0] + t[s1:]


def _fill_docx(src: Path, dst: Path, sch: dict[str, Any], vals: dict[str, Any]) -> list[str]:
    from docx import Document
    from docx.oxml.ns import qn
    from app.formkit.detect import BLANK_RE, _sdts
    doc = Document(str(src))
    filled = []
    # ugyanabban a bekezdésben hátulról előre, hogy a korábbi minták helye ne csússzon el
    texts = sorted((f for f in sch["fields"] if f["anchor"].get("kind") == "docx_text" and f["name"] in vals),
                   key=lambda f: (str(f["anchor"]["loc"]), -f["anchor"]["nth"]))
    for f in texts:
        a, v = f["anchor"], vals[f["name"]]
        p = _docx_para(doc, a["loc"])
        ms = list(BLANK_RE.finditer(p.text))
        if a["nth"] >= len(ms):
            continue
        m = ms[a["nth"]]
        if f["type"] == "signature" and isinstance(v, dict) and v.get("image"):
            from app.formkit.signature import docx_picture_at
            docx_picture_at(p, m.start(), m.end(), Path(v["image"]))
        elif f["type"] == "checkbox":
            _replace_in_runs(p, m.start(), m.end(), "☒" if truthy(v) else "☐")
        else:
            _replace_in_runs(p, m.start(), m.end(), str(v))
        filled.append(f["name"])
    for f in sch["fields"]:
        a = f["anchor"]
        if f["name"] not in vals:
            continue
        v = vals[f["name"]]
        if a.get("kind") == "docx_cell":
            cell = doc.tables[a["t"]].cell(a["r"], a["c"])
            p = cell.paragraphs[0]
            if f["type"] == "signature" and isinstance(v, dict) and v.get("image"):
                from app.formkit.signature import docx_picture_at
                docx_picture_at(p, len(p.text), len(p.text), Path(v["image"]))
            else:
                txt = ("☒" if truthy(v) else "☐") if f["type"] == "checkbox" else str(v)
                for r in p.runs[1:]:
                    r.text = ""
                if p.runs:
                    p.runs[0].text = txt
                else:
                    p.add_run(txt)
            filled.append(f["name"])
        elif a.get("kind") == "docx_after":
            from app.formkit.detect import _docx_paragraphs
            key = str(a.get("label") or "").lower()
            p = next((p for _, p in _docx_paragraphs(doc) if key and key in p.text.lower()), None)
            if p is None:
                continue
            if f["type"] == "signature" and isinstance(v, dict) and v.get("image"):
                from app.formkit.signature import docx_picture_at
                docx_picture_at(p, len(p.text), len(p.text), Path(v["image"]))
            else:
                p.add_run(" " + (("☒" if truthy(v) else "☐") if f["type"] == "checkbox" else str(v)))
            filled.append(f["name"])
        elif a.get("kind") == "docx_sdt":
            sdts = _sdts(doc)
            if a["index"] >= len(sdts):
                continue
            sdt = sdts[a["index"]]
            pr, content = sdt.find(qn("w:sdtPr")), sdt.find(qn("w:sdtContent"))
            ph = pr.find(qn("w:showingPlcHdr"))
            if ph is not None:
                pr.remove(ph)
            ts = list(content.iter(qn("w:t")))
            if f["type"] == "checkbox":
                for c in pr:
                    if c.tag.endswith("}checkbox"):
                        for ch in c:
                            if ch.tag.endswith("}checked"):
                                ch.set(ch.tag.replace("checked", "val"), "1" if truthy(v) else "0")
                if ts:
                    ts[0].text = "☒" if truthy(v) else "☐"
            elif ts:
                ts[0].text = str(v)
                for t in ts[1:]:
                    t.text = ""
            filled.append(f["name"])
    doc.save(str(dst))
    return filled


# ---------------------------------------------------------------- XLSX

def _fill_xlsx(src: Path, dst: Path, sch: dict[str, Any], vals: dict[str, Any]) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(str(src))
    filled = []
    for f in sch["fields"]:
        a = f["anchor"]
        if a.get("kind") != "xlsx_cell" or f["name"] not in vals:
            continue
        v = vals[f["name"]]
        c = wb[a["sheet"]][a["cell"]]
        if f["type"] == "date" and V.parse_date(v):
            c.value = V.parse_date(v)
            c.number_format = "yyyy.mm.dd."
        elif f["type"] == "number":
            try:
                c.value = float(str(v).replace(" ", "").replace(",", "."))
            except ValueError:
                c.value = str(v)
        elif f["type"] == "checkbox":
            c.value = "☒" if truthy(v) else "☐"
        elif isinstance(v, dict):
            continue
        else:
            c.value = str(v)
        filled.append(f["name"])
    wb.save(str(dst))
    return filled


# ----------------------------------------------------------------

def fill(src: Path, sch: dict[str, Any], vals: dict[str, Any], dst_dir: Path, stem: str, *, flatten: bool = False,
         keep_look: bool = True) -> tuple[Path, list[str]]:
    kind = sch["kind"]
    if kind == "acroform":
        dst = dst_dir / f"{stem}.pdf"
        filled = _fill_acroform(src, dst, sch, vals, flatten)
        xy = [f for f in sch["fields"] if f["anchor"].get("kind") == "xy" and f["name"] in vals]
        if xy:      # a mezők mellett: aláírás-vonal, kézzel felvett doboz
            tmp = dst.with_suffix(".xy.pdf")
            filled += _overlay_pdf(dst, tmp, {"fields": xy}, vals)
            tmp.replace(dst)
        return dst, filled
    if kind == "flat_pdf":
        dst = dst_dir / f"{stem}.pdf"
        return dst, _overlay_pdf(src, dst, sch, vals)
    if kind == "docx":
        dst = dst_dir / f"{stem}.docx"
        return dst, _fill_docx(src, dst, sch, vals)
    if kind == "xlsx":
        dst = dst_dir / f"{stem}.xlsx"
        return dst, _fill_xlsx(src, dst, sch, vals)
    if kind == "image":
        if keep_look:
            ext = src.suffix.lower() if src.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp") else ".png"
            dst = dst_dir / f"{stem}{ext}"
            return dst, _overlay_image(src, dst, sch, vals)
        base = dst_dir / f"{stem}.base.pdf"
        flat = src
        if src.suffix.lower() not in (".jpg", ".jpeg", ".png"):
            from PIL import Image
            flat = dst_dir / f"{stem}.src.png"
            with Image.open(src) as im:
                im.convert("RGB").save(flat)
        scale = _image_to_pdf(flat, base)
        dst = dst_dir / f"{stem}.pdf"
        filled = _overlay_pdf(base, dst, sch, vals, scale)
        base.unlink(missing_ok=True)
        return dst, filled
    raise ValueError(f"can't fill a {kind} form")


__all__ = ["fill", "prepare", "match_values", "truthy"]
