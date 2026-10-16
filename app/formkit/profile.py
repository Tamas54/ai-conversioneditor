"""Ügynöki kitöltés: a felhasználó adatprofiljából (titkosítva tárolva) és a feltöltött iratokból tölt ki,
a hiányzót kérdésként adja vissza, mindent validál.

Sorrend mezőnként: kifejezett érték → válasz egy korábbi kérdésre → profil (címke-szabályok) → iratokból
kiolvasva ("Név: …" párok, IBAN/adóazonosító/TAJ/e-mail minták ellenőrzőszámmal) → AI (ha van kulcs és
engedélyezett) → kérdés. Minden értéknél megmondjuk, honnan jött (provenance).
"""
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path
from typing import Any, Optional

from app.formkit import schema as S
from app.formkit import validate as V

PROFILE_KEYS = ("full_name", "family_name", "given_name", "birth_name", "mother_name", "birth_date", "birth_place",
                "nationality", "address", "postcode", "city", "street", "mailing_address", "email", "phone", "tax_id",
                "taj", "id_card_number", "passport_number", "iban", "bank_account", "company_name",
                "company_tax_number", "company_registration", "company_address", "job_title", "signature_id")

# címke (ékezet nélkül) → profilkulcs; az első találat nyer (a konkrétabb elöl)
RULES: list[tuple[str, str]] = [
    (r"anyja (szuletesi )?neve|mother'?s (maiden )?name", "mother_name"),
    (r"szuletesi nev|birth name|maiden name", "birth_name"),
    (r"szuletesi hely|place of birth|birth ?place", "birth_place"),
    (r"szuletesi (ido|datum)|date of birth|birth ?date|\bdob\b", "birth_date"),
    (r"cegjegyzek|company reg|registration (no|number)", "company_registration"),
    (r"\badoszam|vat (no|number)|\bvat\b", "company_tax_number"),
    (r"(ceg|company|munkaltato|employer).{0,12}(cim|address|szekhely)|szekhely", "company_address"),
    (r"ceg ?nev|cegnev|company( name)?|employer|munkaltato( neve)?", "company_name"),
    (r"adoazonosito|tax ?id|\btin\b", "tax_id"),
    (r"\btaj\b|social security", "taj"),
    (r"szemelyi igazolvany|szemelyazonosito igazolvany|id card|\bszig", "id_card_number"),
    (r"utlevel|passport", "passport_number"),
    (r"\biban\b", "iban"),
    (r"bankszamla|szamlaszam|bank account|account (no|number)", "bank_account"),
    (r"e-?mail", "email"),
    (r"telefon|phone|mobil", "phone"),
    (r"iranyitoszam|post ?code|\bzip\b", "postcode"),
    (r"levelezesi cim|mailing address|postal address", "mailing_address"),
    (r"telepules|\bvaros\b|\bcity\b|\btown\b", "city"),
    (r"\butca\b|street", "street"),
    (r"lakcim|lakohely|allando cim|\bcim\b|\baddress\b|residence", "address"),
    (r"allampolgarsag|nationality|citizenship", "nationality"),
    (r"vezeteknev|csaladi nev|family name|surname|last name", "family_name"),
    (r"keresztnev|utonev|given name|first name", "given_name"),
    (r"beosztas|munkakor|job title|position", "job_title"),
    (r"\bneve?\b|\bname\b|alulirott", "full_name"),
]
TODAY_RE = r"^(datum|date|kelt|keltezes|today)$|^kelt\b|^dated?\b"


def key_for(label: str) -> Optional[str]:
    a = S.ascii_fold(label)
    for pat, key in RULES:
        if re.search(pat, a):
            return key
    return None


def hu_iban(bank: str) -> Optional[str]:
    d = re.sub(r"\D", "", bank or "")
    if len(d) == 16:
        d += "0" * 8
    if len(d) != 24:
        return None
    num = "".join(str(int(c, 36)) for c in d + "HU00")
    chk = 98 - int(num) % 97
    return V.check_iban(f"HU{chk:02d}{d}")[1]


def derive(p: dict[str, Any]) -> dict[str, Any]:
    """Hiányzó, de kikövetkeztethető profilmezők (név részei, cím, születési dátum az adóazonosítóból, IBAN)."""
    p = {k: v for k, v in (p or {}).items() if v not in (None, "")}
    if not p.get("full_name") and p.get("family_name") and p.get("given_name"):
        p["full_name"] = f"{p['family_name']} {p['given_name']}"
    if p.get("full_name") and not (p.get("family_name") and p.get("given_name")):
        parts = str(p["full_name"]).split()
        if len(parts) >= 2:
            p.setdefault("family_name", parts[0])
            p.setdefault("given_name", " ".join(parts[1:]))
    if not p.get("address") and p.get("city") and p.get("street"):
        p["address"] = f"{p.get('postcode', '')} {p['city']}, {p['street']}".strip()
    if not p.get("birth_date") and p.get("tax_id") and V.check_hu_tax_id(p["tax_id"])[0]:
        p["birth_date"] = V.hu_tax_id_birth(p["tax_id"]).isoformat()
    if not p.get("iban") and p.get("bank_account"):
        ib = hu_iban(str(p["bank_account"]))
        if ib:
            p["iban"] = ib
    if not p.get("bank_account") and str(p.get("iban", "")).replace(" ", "").upper().startswith("HU"):
        p["bank_account"] = V.check_hu_bank(re.sub(r"\D", "", str(p["iban"]))[2:])[1]
    return p


# ---------------------------------------------------------------- iratokból

_PATTERNS = {
    "iban": r"\b[A-Z]{2}\d{2}(?:\s?[A-Z0-9]{4}){3,7}(?:\s?[A-Z0-9]{1,4})?\b",
    "hu_tax_id": r"\b8\d{9}\b|\b8\d{3}\s?\d{3}\s?\d{3}\b",
    "taj": r"\b\d{3}[\s-]?\d{3}[\s-]?\d{3}\b",
    "email": r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+",
    "hu_bank": r"\b\d{8}-\d{8}(?:-\d{8})?\b",
}


def from_sources(sch: dict[str, Any], text: str) -> dict[str, Any]:
    """Iratszövegből: "Felirat: érték" párok a mezők címkéihez, és ellenőrzőszámos minták a típusos mezőkhöz."""
    out: dict[str, Any] = {}
    pairs: dict[str, str] = {}
    for line in text.splitlines():
        m = re.match(r"\s*([^:\n]{2,50}?)\s*[:：]\s*(.{1,200}?)\s*$", line)
        if m and not re.fullmatch(r"[_.\s]*", m.group(2)):
            pairs.setdefault(S.ascii_fold(S.clean_label(m.group(1))), m.group(2).strip())
    for f in sch["fields"]:
        lab = S.ascii_fold(f.get("label") or f["name"])
        if lab in pairs:
            out[f["name"]] = pairs[lab]
            continue
        pat = _PATTERNS.get(f["type"])
        if pat:
            for m in re.finditer(pat, text):
                ok, val, _ = V.check(f["type"], m.group(0))
                if ok:
                    out[f["name"]] = val
                    break
    return out


async def source_text(paths: list[Path], max_chars: int = 20000) -> str:
    from app.convert.engines import available, run
    chunks = []
    for p in paths:
        ext = p.suffix.lower()
        try:
            if ext == ".pdf":
                import pymupdf
                with pymupdf.open(str(p)) as d:
                    t = "\n".join(pg.get_text() for pg in d)
                if not t.strip() and available("tesseract"):
                    with pymupdf.open(str(p)) as d:
                        img = p.with_name(p.stem + "-ocr.png")
                        d[0].get_pixmap(dpi=200).save(str(img))
                    out, _ = await run("tesseract", str(img), "stdout", "-l", "hun+eng", timeout=300)
                    img.unlink(missing_ok=True)
                    t = out.decode("utf-8", "replace")
            elif ext in (".txt", ".md", ".csv", ".json"):
                t = p.read_text("utf-8", errors="replace")
            elif ext == ".docx":
                from docx import Document
                d = Document(str(p))
                t = "\n".join([x.text for x in d.paragraphs] + [" : ".join(c.text for c in r.cells) for tb in d.tables for r in tb.rows])
            elif ext in (".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".heic", ".bmp") and available("tesseract"):
                out, _ = await run("tesseract", str(p), "stdout", "-l", "hun+eng", timeout=300)
                t = out.decode("utf-8", "replace")
            else:
                from app.convert.service import convert_path
                outs, _ = await convert_path(p, "txt", {}, p.parent / f"_txt_{p.stem}")
                t = outs[0].read_text("utf-8", errors="replace")
        except Exception:
            t = ""
        chunks.append(t)
    return "\n\n".join(chunks)[:max_chars]


# ---------------------------------------------------------------- AI-megfeleltetés

async def ai_values(sch: dict[str, Any], missing: list[dict[str, Any]], profile: dict[str, Any], text: str) -> dict[str, Any]:
    from app.config import settings
    from app.llm import siliconflow
    if not settings.SILICONFLOW_API_KEY or not missing:
        return {}
    want = [{"name": f["name"], "label": f.get("label"), "type": f["type"], "options": f.get("options") or None} for f in missing]
    prompt = ("Fill in form fields from the person's profile and their documents. Use ONLY facts that are stated "
              "there; if a field's value isn't known, leave it out. Keep the language and spelling of the source "
              "(Hungarian accents included). Dates as YYYY-MM-DD. Checkboxes: true/false only if clearly implied.\n"
              f"FIELDS: {json.dumps(want, ensure_ascii=False)}\nPROFILE: {json.dumps(profile, ensure_ascii=False)}\n"
              f"DOCUMENTS:\n{text[:12000]}\n\nReply with one JSON object {{field name: value}}.")
    try:
        raw = await siliconflow._call_model(settings.LLM_MODEL_PRIMARY, [{"role": "user", "content": prompt}], thinking=False)
        data = json.loads(siliconflow._strip_json_fences(raw))
    except Exception:
        return {}
    names = {f["name"] for f in missing}
    return {k: v for k, v in (data or {}).items() if k in names and v not in (None, "")}


# ---------------------------------------------------------------- összerakva

async def plan(sch: dict[str, Any], *, values: Optional[dict] = None, answers: Optional[dict] = None,
               profile: Optional[dict] = None, sources: str = "", use_ai: bool = True) -> dict[str, Any]:
    from app.formkit.fill import match_values, prepare
    prof = derive(profile or {})
    got: dict[str, Any] = {}
    origin: dict[str, str] = {}

    def put(vals: dict[str, Any], src: str):
        for k, v in vals.items():
            if k not in got and v not in (None, ""):
                got[k], origin[k] = v, src
    put(match_values(sch, values or {})[0], "value")
    put(match_values(sch, answers or {})[0], "answer")
    for f in sch["fields"]:
        if f["name"] in got or f["type"] == "signature":
            continue
        lab = f.get("label") or f["name"]
        if re.search(TODAY_RE, S.ascii_fold(lab)):
            got[f["name"]], origin[f["name"]] = date.today().isoformat(), "today"
            continue
        key = key_for(lab)
        if key == "full_name" and f["type"] != "text":
            key = None
        if key and prof.get(key) not in (None, ""):
            got[f["name"]], origin[f["name"]] = prof[key], f"profile:{key}"
    if sources:
        put(from_sources(sch, sources), "document")
    left = [f for f in sch["fields"] if f["name"] not in got and f["type"] not in ("signature", "checkbox")]
    if use_ai and left:
        put(await ai_values(sch, left, prof, sources), "ai")
    vals, errors, _ = prepare(sch, got)
    for k in errors:
        vals.pop(k, None)
    sig_fields = [f for f in sch["fields"] if f["type"] == "signature"]
    questions = []
    for f in sch["fields"]:
        if f["name"] in vals or f["type"] == "signature":
            continue
        if f["type"] == "checkbox" and not f.get("required"):
            continue
        q = {"name": f["name"], "label": f.get("label") or f["name"], "type": f["type"], "options": f.get("options") or [],
             "required": bool(f.get("required"))}
        if f["name"] in errors:
            q["error"] = errors[f["name"]]
            q["given"] = got.get(f["name"])
        questions.append(q)
    return {"values": vals, "provenance": {k: origin.get(k, "value") for k in vals}, "questions": questions,
            "errors": errors, "signature_fields": [f["name"] for f in sig_fields], "profile_used": sorted(prof)}


# ---------------------------------------------------------------- profil tárolása (titkosítva)

def _fs(uid: str) -> Path:
    import hashlib
    from app.config import settings
    d = Path(settings.DATA_DIR) / "secure" / "profiles"
    d.mkdir(parents=True, exist_ok=True)
    return d / (hashlib.sha256(uid.encode()).hexdigest()[:24] + ".bin")


def clean_profile(p: dict[str, Any]) -> dict[str, Any]:
    out = {}
    for k, v in (p or {}).items():
        if k not in PROFILE_KEYS and not re.fullmatch(r"custom_[a-z0-9_]{1,40}", str(k)):
            raise ValueError(f"unknown profile field {k}. Fields: {', '.join(PROFILE_KEYS)} (or custom_<name>)")
        if v in (None, ""):
            continue
        s = str(v).strip()
        if len(s) > 500:
            raise ValueError(f"{k}: at most 500 characters")
        out[k] = s
    return out


async def load_profile(uid: str) -> dict[str, Any]:
    from app import db
    from app.formkit import crypto
    if db.enabled():
        row = await db.profile_get(uid)
        if not row:
            return {}
        raw = crypto.open_(uid, "profile", bytes(row["nonce"]), bytes(row["data"]))
    else:
        p = _fs(uid)
        if not p.exists():
            return {}
        b = p.read_bytes()
        raw = crypto.open_(uid, "profile", b[:12], b[12:])
    return json.loads(raw.decode("utf-8"))


async def save_profile(uid: str, data: dict[str, Any], merge: bool = True) -> dict[str, Any]:
    from app import db
    from app.formkit import crypto
    data = clean_profile(data)
    if merge:
        data = {**(await load_profile(uid)), **data}
    nonce, enc = crypto.seal(uid, "profile", json.dumps(data, ensure_ascii=False).encode("utf-8"))
    if db.enabled():
        await db.profile_put(uid, nonce, enc)
    else:
        _fs(uid).write_bytes(nonce + enc)
    return data


async def delete_profile(uid: str) -> bool:
    from app import db
    if db.enabled():
        return await db.profile_delete(uid)
    p = _fs(uid)
    ok = p.exists()
    p.unlink(missing_ok=True)
    return ok
