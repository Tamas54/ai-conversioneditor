"""Ismétlődő XLSX-sablon (havi jelenléti ív, munkaidő-igazoló lap) kitöltése EGY hívásból, determinisztikusan.

  register  a sablon (.xlsx, vagy régi .xls — a LibreOffice képletekkel, formázással xlsx-be teszi) + a profil, a
            felhasználó fiókjában (Postgres, titkosítva: a fejléc állandó értékei személyes adatok)
  learn     üres sablon + egy korábbi, kitöltött hónap → javasolt profil a különbségből; a felhasználó hagyja jóvá
            (register a javaslattal, igazítva)
  fill      template_id + period (ÉÉÉÉ-HH) + alapértékek (8:00–16:00) + kivételek (dátum + fajta) → kitöltött XLSX
            (kérésre XLS is) és PDF, és a check

Mi kerülhet a lapra: csak a kapott vagy a munkanaptárból (app/tools/workcalendar) levezethető adat. Munkanapra (az
áthelyezett szombati munkanap is) az alapérték, kivételnapra a kivétel szabálya; hétvége, munkaszüneti nap,
áthelyezett pihenőnap üres marad — semmit nem találunk ki. Az idő időként, a dátum dátumként, a szám számként kerül a
cellába (szövegként a sablon képletei nem számolnának); írás után a LibreOffice MINDEN képletet újraszámol, és csak
utána készül a PDF. A kész xlsx-ben a számolt értékek is benne vannak, és az Excel megnyitáskor újraszámol.

Profil (register / learn):
  {"sheet"?: "Jelenléti ív",
   "header": {"name": "B3", "job_title": "B4", "code": "F4", "month": "F3"},     szerep → cella
   "month_format"?: "date" | "{year}. {month}" | "{year}-{mm}" | "{month} {year}"  (month = magyar hónapnév)
   "days": {"first_row": 8, "step"?: 1},          a d. nap sora: first_row + (d − 1) · step
   "columns": {"arrive": "B", "leave": "C", "absence": "D"},                    szerep → oszlop
   "rules": {"workday": {"arrive": "{start}", "leave": "{end}"},
             "kinds": {"szabadság": {"leave": "SZABADSÁG", "absence": "+"}, "betegség": {…}}},
   "totals"?: {"hours": "E39", "absence_days": "D39", "worked_days": "B39"},     a check ezeket olvassa vissza
   "header_values"?: {"name": "…", "job_title": "…", "code": "…"},              havonta ugyanaz
   "defaults"?: {"start": "8:00", "end": "16:00"}}
"""
from __future__ import annotations

import asyncio
import base64
import datetime as dt
import hashlib
import io
import json
import re
import secrets
import shutil
import unicodedata
import zipfile
from pathlib import Path
from typing import Any, Optional

from app.tools import workcalendar as wc

KINDS_HINT = "szabadság, betegség, … (the profile's rules.kinds)"
_ERRORS = ("#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")
_CELL = re.compile(r"[A-Z]{1,3}[1-9]\d{0,6}")
_COL = re.compile(r"[A-Z]{1,3}")
_TIME = re.compile(r"(\d{1,2})[:.](\d{2})")


class TemplateError(ValueError):
    pass


# ---------------------------------------------------------------- profil

def _fold(s: str) -> str:
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().casefold().strip()


def validate_profile(p: Any) -> dict[str, Any]:
    """A profil alakja (hibánál pontosan megmondja, mi a baj). Vissza: a normalizált profil."""
    if not isinstance(p, dict):
        raise TemplateError("profile must be an object (see the forms tool: header, days, columns, rules)")
    out = dict(p)
    hdr = p.get("header") or {}
    if not isinstance(hdr, dict) or any(not isinstance(v, str) or not _CELL.fullmatch(v.upper()) for v in hdr.values()):
        raise TemplateError('profile.header: {"name": "B3", "month": "F3", …} — role → cell')
    out["header"] = {k: v.upper() for k, v in hdr.items()}
    days = p.get("days") or {}
    try:
        first, step = int(days["first_row"]), int(days.get("step", 1))
        assert first >= 1 and step >= 1
    except (KeyError, TypeError, ValueError, AssertionError):
        raise TemplateError('profile.days: {"first_row": 8, "step": 1} — the row of day d is first_row + (d − 1) · step')
    out["days"] = {"first_row": first, "step": step}
    cols = p.get("columns") or {}
    if not cols or not isinstance(cols, dict) or any(not isinstance(v, str) or not _COL.fullmatch(v.upper()) for v in cols.values()):
        raise TemplateError('profile.columns: {"arrive": "B", "leave": "C", "absence": "D"} — role → column letter')
    out["columns"] = {k: v.upper() for k, v in cols.items()}
    rules = p.get("rules") or {}
    work = rules.get("workday")
    kinds = rules.get("kinds") or {}
    if not isinstance(work, dict) or not work or not isinstance(kinds, dict):
        raise TemplateError('profile.rules: {"workday": {"arrive": "{start}", "leave": "{end}"}, "kinds": {"szabadság": '
                            '{"leave": "SZABADSÁG", "absence": "+"}}}')
    for name, r in [("workday", work), *[(f"kinds.{k}", v) for k, v in kinds.items()]]:
        if not isinstance(r, dict):
            raise TemplateError(f"profile.rules.{name} must be an object role → value")
        bad = [k for k in r if k not in out["columns"]]
        if bad:
            raise TemplateError(f"profile.rules.{name}: {', '.join(bad)} isn't a column role ({', '.join(out['columns'])})")
    out["rules"] = {"workday": work, "kinds": {str(k).strip().lower(): v for k, v in kinds.items()}}
    fmt = p.get("month_format") or "date"
    if fmt != "date" and "{year}" not in fmt:
        raise TemplateError('profile.month_format: "date", or a pattern with {year} and {month} / {mm}, e.g. "{year}. {month}"')
    out["month_format"] = fmt
    tot = p.get("totals") or {}
    if any(not isinstance(v, str) or not _CELL.fullmatch(v.upper()) for v in tot.values()):
        raise TemplateError('profile.totals: {"hours": "E39", "absence_days": "D39"} — role → cell')
    out["totals"] = {k: v.upper() for k, v in tot.items()}
    out["header_values"] = {k: v for k, v in (p.get("header_values") or {}).items() if v not in (None, "")}
    out["defaults"] = dict(p.get("defaults") or {})
    return out


# ---------------------------------------------------------------- tárolás (a felhasználó fiókjában, titkosítva)

def _fs_dir() -> Path:
    from app.config import settings
    d = Path(settings.DATA_DIR) / "secure" / "sheet_templates"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _fs_path(uid: str, tid: str) -> Path:
    return _fs_dir() / (hashlib.sha256(f"{uid}:{tid}".encode()).hexdigest()[:32] + ".bin")


async def _to_xlsx(src: Path) -> bytes:
    ext = src.suffix.lstrip(".").lower()
    if ext == "xlsx":
        return src.read_bytes()
    if ext not in ("xls", "ods", "xlsm"):
        raise TemplateError(f"the template must be an .xlsx (or .xls) spreadsheet, not .{ext}")
    from app.storage import temp_path
    from app.storage.files import new_file_id
    work = temp_path(new_file_id("tpl"))
    work.mkdir(parents=True)
    try:
        out = await _soffice(src, work, "xlsx:Calc MS Excel 2007 XML")
        return out.read_bytes()
    finally:
        shutil.rmtree(work, ignore_errors=True)


async def register(uid: str, template: Path, profile: dict[str, Any], name: Optional[str] = None) -> dict[str, Any]:
    from app import db
    from app.formkit import crypto
    prof = validate_profile(profile)
    xlsx = await _to_xlsx(template)
    _open_sheet(xlsx, prof)                          # a profil cellái a sablonon is értelmesek (lap, sorok)
    tid = "tpl_" + secrets.token_hex(6)
    name = (name or Path(template).stem)[:120]
    blob = json.dumps({"profile": prof, "xlsx": base64.b64encode(xlsx).decode(), "name": name}, ensure_ascii=False).encode()
    nonce, enc = crypto.seal(uid, "sheet_template", blob)
    if db.enabled():
        await db.sheet_template_put(uid, tid, name, nonce, enc)
    else:
        _fs_path(uid, tid).write_bytes(nonce + enc)
        idx = _fs_dir() / (hashlib.sha256(uid.encode()).hexdigest()[:24] + ".idx")
        items = json.loads(idx.read_text("utf-8")) if idx.exists() else []
        idx.write_text(json.dumps(items + [{"id": tid, "name": name, "created_at": dt.datetime.now().isoformat()}]), "utf-8")
    return {"template_id": tid, "name": name, "profile": prof}


async def load(uid: str, tid: str) -> tuple[str, dict[str, Any], bytes]:
    from app import db
    from app.formkit import crypto
    if db.enabled():
        row = await db.sheet_template_get(uid, tid)
        if not row:
            raise TemplateError(f"no such template: {tid}")
        raw = crypto.open_(uid, "sheet_template", bytes(row["nonce"]), bytes(row["data"]))
    else:
        p = _fs_path(uid, tid)
        if not p.exists():
            raise TemplateError(f"no such template: {tid}")
        b = p.read_bytes()
        raw = crypto.open_(uid, "sheet_template", b[:12], b[12:])
    d = json.loads(raw.decode("utf-8"))
    return d["name"], d["profile"], base64.b64decode(d["xlsx"])


async def list_(uid: str) -> list[dict[str, Any]]:
    from app import db
    if db.enabled():
        return [{"template_id": r["id"], "name": r["name"], "created_at": str(r["created_at"])}
                for r in await db.sheet_template_list(uid)]
    idx = _fs_dir() / (hashlib.sha256(uid.encode()).hexdigest()[:24] + ".idx")
    items = json.loads(idx.read_text("utf-8")) if idx.exists() else []
    return [{"template_id": i["id"], "name": i["name"], "created_at": i["created_at"]} for i in items
            if _fs_path(uid, i["id"]).exists()]


async def delete(uid: str, tid: str) -> bool:
    from app import db
    if db.enabled():
        return await db.sheet_template_delete(uid, tid)
    p = _fs_path(uid, tid)
    ok = p.exists()
    p.unlink(missing_ok=True)
    return ok


# ---------------------------------------------------------------- munkalap

def _open_sheet(xlsx: bytes, prof: dict[str, Any], data_only: bool = False):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx), data_only=data_only)
    if prof.get("sheet"):
        if prof["sheet"] not in wb.sheetnames:
            raise TemplateError(f"the template has no sheet {prof['sheet']!r} (sheets: {', '.join(wb.sheetnames)})")
        ws = wb[prof["sheet"]]
    else:
        ws = wb.worksheets[0]
    last = prof["days"]["first_row"] + 30 * prof["days"]["step"]
    if last > ws.max_row + 40:
        raise TemplateError(f"profile.days: day 31 would be row {last}, but the sheet ends at row {ws.max_row}")
    return wb, ws


def _target(ws, coord: str):
    """Egyesített cellába a bal felső cellába írunk (a többi csak olvasható)."""
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            return ws.cell(rng.min_row, rng.min_col)
    return ws[coord]


def parse_time(v: Any) -> dt.time:
    if isinstance(v, dt.time):
        return v
    m = _TIME.fullmatch(str(v).strip())
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        raise TemplateError(f"{v!r} isn't a time — write it as 8:00")
    return dt.time(int(m.group(1)), int(m.group(2)))


def _value(v: Any, ctx: dict[str, Any]) -> Any:
    """Egy szabály-érték a cellába: {start}/{end} → idő, „8:00” → idő, szám → szám, a többi szöveg (betűre)."""
    if isinstance(v, str):
        s = v.strip()
        m = re.fullmatch(r"\{(\w+)\}", s)
        if m:
            key = m.group(1)
            if key not in ctx:
                raise TemplateError(f"the rule needs {{{key}}}, but no such value was given (defaults: start, end)")
            return ctx[key]
        if _TIME.fullmatch(s):
            return parse_time(s)
        return v
    return v


def _put(cell, v: Any) -> None:
    """Típusosan: idő időként (h:mm), dátum dátumként, szám számként — a sablon számformátuma marad, ha van."""
    cell.value = v
    if isinstance(v, dt.time) and cell.number_format in ("General", "@"):
        cell.number_format = "h:mm"
    elif isinstance(v, dt.date) and cell.number_format in ("General", "@"):
        cell.number_format = "yyyy.mm.dd"


def _month_value(fmt: str, y: int, m: int) -> Any:
    if fmt == "date":
        return dt.date(y, m, 1)
    return fmt.format(year=y, month=wc.MONTHS[m - 1], mm=f"{m:02d}", m=m)


def parse_period(period: str) -> tuple[int, int]:
    m = re.fullmatch(r"(\d{4})-(\d{2})", str(period or "").strip())
    if not m or not 1 <= int(m.group(2)) <= 12:
        raise TemplateError("period: the month to fill, as YYYY-MM (e.g. 2025-12)")
    return int(m.group(1)), int(m.group(2))


def _exceptions(items: Any, y: int, m: int, kinds: dict[str, Any], days: dict[dt.date, wc.Day]) -> dict[dt.date, str]:
    out: dict[dt.date, str] = {}
    for it in items or []:
        if not isinstance(it, dict) or not it.get("kind"):
            raise TemplateError('exceptions: [{"date": "2025-12-01", "kind": "szabadság"}] or {"from": …, "to": …, "kind": …}')
        kind = str(it["kind"]).strip().lower()
        if kind not in kinds:
            match = next((k for k in kinds if _fold(k) == _fold(kind)), None)
            if not match:
                raise TemplateError(f"exception kind {it['kind']!r} isn't in the template's rules — known: {', '.join(kinds) or 'none'}")
            kind = match
        try:
            a = dt.date.fromisoformat(str(it.get("date") or it.get("from")))
            b = dt.date.fromisoformat(str(it.get("to") or it.get("date") or it.get("from")))
        except ValueError:
            raise TemplateError(f"exception date {it.get('date') or it.get('from')!r}: write it as YYYY-MM-DD")
        d = a
        while d <= b:
            if (d.year, d.month) != (y, m):
                raise TemplateError(f"{d.isoformat()} isn't in {y}-{m:02d}")
            if d in days and days[d].is_workday:
                out[d] = kind
            elif not it.get("to"):
                k = days[d]
                raise TemplateError(f"{d.isoformat()} ({k.weekday}) is not a workday ({k.name or k.kind}) — an exception "
                                    "only goes on a workday")
            d += dt.timedelta(days=1)
    return out


# ---------------------------------------------------------------- LibreOffice: teljes újraszámolás, PDF, XLS

_RECALC_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


async def _soffice(src: Path, work: Path, target: str) -> Path:
    """LibreOffice saját profillal, amelyben a betöltéskori újraszámolás „mindig” (OOXMLRecalcMode = 0)."""
    from app.convert.engines import run
    prof = work / "profile"
    (prof / "user").mkdir(parents=True, exist_ok=True)
    (prof / "user" / "registrymodifications.xcu").write_text(_RECALC_XCU, "utf-8")
    out = work / f"out-{target.split(':')[0]}-{secrets.token_hex(3)}"
    out.mkdir()
    await run("soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--norestore", "--convert-to", target,
              "--outdir", str(out), str(src), timeout=180)
    got = sorted(out.glob("*"))
    if not got:
        raise RuntimeError("LibreOffice produced nothing")
    return got[0]


def _calc_on_load(path: Path) -> None:
    """Az Excel is számoljon újra megnyitáskor (a LibreOffice által számolt értékek a fájlban maradnak)."""
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "xl/workbook.xml":
                x = data.decode("utf-8")
                if "<calcPr" in x:
                    x = re.sub(r"<calcPr\b([^>]*?)(/?)>", lambda m: "<calcPr" + re.sub(r'\s+fullCalcOnLoad="[^"]*"', "", m.group(1))
                               + ' fullCalcOnLoad="1"' + m.group(2) + ">", x, count=1)
                else:
                    x = x.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')
                data = x.encode("utf-8")
            zo.writestr(item, data)
    tmp.replace(path)


# ---------------------------------------------------------------- kitöltés

def _hhmm(hours: float) -> str:
    mins = int(round(hours * 60))
    return f"{mins // 60}:{mins % 60:02d}"


def _hours(cell) -> Optional[float]:
    """Egy összesítő cella értéke órában. Az idő-formátumú szám napban van (az openpyxl dátummá/időtartammá alakítja —
    a to_excel adja vissza a sorozatszámot, az 1900-as naptár szökőnap-hibáját is helyesen); a sima szám már óra."""
    from openpyxl.utils.datetime import to_excel
    v = cell.value
    if isinstance(v, (dt.datetime, dt.date, dt.time, dt.timedelta)):
        return float(to_excel(v)) * 24
    if isinstance(v, (int, float)):
        return float(v) * 24 if re.search(r"[hH]", cell.number_format or "") else float(v)
    return None


async def fill(uid: str, template_id: str, period: str, defaults: Optional[dict[str, Any]] = None,
               exceptions: Optional[list] = None, formats: Optional[list[str]] = None) -> dict[str, Any]:
    from app.storage import output_path, public_url, temp_path, write_meta
    from app.storage.files import new_file_id
    name, prof, xlsx = await load(uid, template_id)
    y, m = parse_period(period)
    try:
        cal = wc.month(y, m)
    except wc.UnknownYear as e:
        raise TemplateError(str(e))
    days = {d.date: d for d in cal}
    given = {**prof.get("defaults", {}), **{k: v for k, v in (defaults or {}).items() if v not in (None, "")}}
    ctx: dict[str, Any] = {}
    for k in ("start", "end"):
        if k in given:
            ctx[k] = parse_time(given[k])
    head = {**prof.get("header_values", {}), **{k: v for k, v in given.items() if k in prof["header"] and k != "month"}}
    exc = _exceptions(exceptions, y, m, prof["rules"]["kinds"], days)
    wb, ws = _open_sheet(xlsx, prof)
    for role, coord in prof["header"].items():
        if role == "month":
            _put(_target(ws, coord), _month_value(prof["month_format"], y, m))
        elif role in head:
            _put(_target(ws, coord), head[role])
    cols, first, step = prof["columns"], prof["days"]["first_row"], prof["days"]["step"]
    written: dict[dt.date, dict[str, Any]] = {}
    for d in cal:
        rule = prof["rules"]["kinds"][exc[d.date]] if d.date in exc else prof["rules"]["workday"] if d.is_workday else None
        if rule is None:
            continue                                      # hétvége, munkaszüneti nap, pihenőnap: üres marad
        row = first + (d.date.day - 1) * step
        vals = {role: _value(v, ctx) for role, v in rule.items()}
        for role, v in vals.items():
            _put(_target(ws, f"{cols[role]}{row}"), v)
        written[d.date] = vals
    work = temp_path(new_file_id("tsh"))
    work.mkdir(parents=True)
    try:
        raw = work / f"{_safe(name)} {y}-{m:02d}.xlsx"
        wb.save(str(raw))
        calc = await _soffice(raw, work, "xlsx:Calc MS Excel 2007 XML")       # teljes újraszámolás
        _calc_on_load(calc)
        check, summary = verify(calc.read_bytes(), prof, cal, exc, written, ctx)
        label = f"{name} {y}-{m:02d}"
        files = []
        wanted = [f.lower().lstrip(".") for f in (formats or ["xlsx", "pdf"])]
        for fmt in wanted:
            if fmt == "xlsx":
                src = calc
            elif fmt == "pdf":
                src = await _soffice(calc, work, "pdf:calc_pdf_Export")
            elif fmt == "xls":
                src = await _soffice(calc, work, "xls:MS Excel 97")
            else:
                raise TemplateError(f"formats: xlsx, pdf, xls — not {fmt}")
            fid = new_file_id(fmt)
            shutil.copy(src, output_path(fid))
            write_meta(fid, label=f"{label}.{fmt}", operation="sheet template fill",
                       extra={"template_id": template_id, "period": f"{y}-{m:02d}"})
            files.append({"file_id": fid, "url": public_url(fid), "label": f"{label}.{fmt}", "format": fmt})
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return {"files": files, "file_id": files[0]["file_id"] if files else None, "check": check, "summary": summary,
            "period": f"{y}-{m:02d}", "calendar_sources": wc.sources(y)}


def _safe(name: str) -> str:
    return re.sub(r"[\\/:*?\"<>|\x00-\x1f]+", "_", name).strip() or "sheet"


# ---------------------------------------------------------------- ellenőrzés

def verify(xlsx: bytes, prof: dict[str, Any], cal: list[wc.Day], exc: dict[dt.date, str],
           written: dict[dt.date, dict[str, Any]], ctx: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    """A LibreOffice által újraszámolt munkafüzet: napok, órák, hibás képlet, minden munkanapnak van bejegyzése.
    → (check a közös alakban, összesítés: munkanapok, kivételnapok, órák)"""
    from app.tools.checks import make
    wb, ws = _open_sheet(xlsx, prof, data_only=True)
    findings, verified = [], []
    workdays = [d for d in cal if d.is_workday]
    worked = [d for d in workdays if d.date not in exc]
    errs = [f"{c.coordinate}: {c.value}" for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith(_ERRORS)]
    if errs:
        findings.append(f"Formula errors in the sheet: {', '.join(errs[:8])}{'…' if len(errs) > 8 else ''}.")
    else:
        verified.append("No #VALUE!, #REF! or other formula error anywhere on the sheet.")
    cols, first, step = prof["columns"], prof["days"]["first_row"], prof["days"]["step"]
    missing = []
    for d in workdays:
        row = first + (d.date.day - 1) * step
        want = written.get(d.date) or {}
        if not want or any(ws[f"{cols[r]}{row}"].value in (None, "") for r in want):
            missing.append(d.date.isoformat())
    if missing:
        findings.append(f"Workdays without an entry: {', '.join(missing)}.")
    else:
        verified.append(f"Every workday has its entry: {len(workdays)} workdays ({len(worked)} worked, "
                        f"{len(exc)} exception day{'s' if len(exc) != 1 else ''}).")
    summary: dict[str, Any] = {"workdays": len(workdays), "worked_days": len(worked), "exception_days": len(exc),
                               "transferred_workdays": [d.date.isoformat() for d in cal if d.kind == "transferred_workday"],
                               "rest_days": [d.date.isoformat() for d in cal if d.kind in ("holiday", "restday")]}
    if "start" in ctx and "end" in ctx:
        per = (dt.datetime.combine(dt.date.min, ctx["end"]) - dt.datetime.combine(dt.date.min, ctx["start"])).total_seconds() / 3600
        expected = per * len(worked)
        summary["expected_hours"] = _hhmm(expected)
        cell = prof["totals"].get("hours")
        if cell:
            h = _hours(ws[cell])
            if h is None:
                findings.append(f"The total hours cell {cell} shows {ws[cell].value!r}, not a number.")
            else:
                summary["hours"] = _hhmm(h)
                if abs(h - expected) > 1 / 120:
                    findings.append(f"The sheet's total ({cell}) is {_hhmm(h)}, but {len(worked)} worked days × "
                                    f"{_hhmm(per)} = {_hhmm(expected)}.")
                else:
                    verified.append(f"Total hours ({cell}, the sheet's own formula): {_hhmm(h)} = {len(worked)} × {_hhmm(per)}.")
    for role, n in (("absence_days", len(exc)), ("worked_days", len(worked))):
        cell = prof["totals"].get(role)
        if cell:
            v = ws[cell].value
            if not isinstance(v, (int, float)) or int(v) != n:
                findings.append(f"{cell} ({role.replace('_', ' ')}) shows {v!r}, expected {n}.")
            else:
                verified.append(f"{role.replace('_', ' ').capitalize()} ({cell}): {n}.")
    return make("failed" if findings else "ok", findings, verified), summary


# ---------------------------------------------------------------- profil tanulása: üres sablon + egy kitöltött hónap

_ROLE_WORDS = {"name": ("nev", "name", "dolgozo"), "job_title": ("munkakor", "beosztas", "job", "position"),
               "code": ("kod", "azonosito", "torzsszam", "code", "id"), "month": ("honap", "idoszak", "month", "period", "ev/ho")}


def _is_time(v: Any) -> Optional[dt.time]:
    if isinstance(v, dt.time):
        return v
    if isinstance(v, dt.datetime) and v.date() == dt.date(1899, 12, 30):
        return v.time()
    if isinstance(v, str) and _TIME.fullmatch(v.strip()):
        try:
            return parse_time(v)
        except TemplateError:
            return None
    return None


def learn(template: bytes, sample: bytes, period: Optional[str] = None) -> dict[str, Any]:
    """A különbségből javasolt profil + indoklás. A javaslatot a felhasználó hagyja jóvá (register)."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    tw, sw = load_workbook(io.BytesIO(template)), load_workbook(io.BytesIO(sample))
    best = None
    for name in tw.sheetnames:
        if name not in sw.sheetnames:
            continue
        a, b = tw[name], sw[name]
        diff = {(c.row, c.column): c.value for row in b.iter_rows() for c in row
                if c.value not in (None, "") and a.cell(c.row, c.column).value in (None, "")}
        if best is None or len(diff) > len(best[1]):
            best = (name, diff, a)
    if not best or not best[1]:
        raise TemplateError("the filled sample has no cells that the empty template leaves blank — nothing to learn from")
    sheet, diff, ws = best
    evidence: list[str] = []
    # napi sorok: ahol idő került a cellákba
    times = {k: t for k, v in diff.items() if (t := _is_time(v))}
    rows = sorted({r for r, _ in times})
    if len(rows) < 3:
        raise TemplateError("couldn't find the daily rows: the sample has fewer than 3 rows with times")
    tcols = sorted({c for _, c in times}, key=lambda c: sum(1 for r, cc in times if cc == c), reverse=True)[:2]
    tcols.sort(key=lambda c: statistics_mean([times[(r, c)].hour * 60 + times[(r, c)].minute for r in rows if (r, c) in times]))
    arrive, leave = tcols[0], tcols[-1]
    gaps = [b - a for a, b in zip(rows, rows[1:])]
    step = min(g for g in gaps if g > 0) if gaps else 1
    first = _first_row(ws, rows, step, period)
    evidence.append(f"Daily rows from row {first} (day d at row {first} + (d − 1)·{step}); times in "
                    f"{get_column_letter(arrive)} (arrival) and {get_column_letter(leave)} (leave).")
    last = first + 30 * step
    columns = {"arrive": get_column_letter(arrive), "leave": get_column_letter(leave)}
    kinds: dict[str, dict[str, Any]] = {}
    for r in range(first, last + 1, step):
        cells = {c: v for (rr, c), v in diff.items() if rr == r}
        if not cells or any(_is_time(v) for v in cells.values()):
            continue
        texts = [str(v).strip() for v in cells.values() if isinstance(v, str) and len(str(v).strip()) > 1]
        kind = (texts[0] if texts else "távollét").lower()
        rule = {}
        for c, v in cells.items():
            role = "arrive" if c == arrive else "leave" if c == leave else next(
                (k for k, cc in columns.items() if cc == get_column_letter(c)), None)
            if role is None:
                role = "absence" if "absence" not in columns else f"col_{get_column_letter(c).lower()}"
                columns[role] = get_column_letter(c)
            rule[role] = v
        kinds.setdefault(kind, rule)
    if kinds:
        evidence.append("Exception rules from the sample: " + "; ".join(
            f"{k}: " + ", ".join(f"{columns[r]} = {v!r}" for r, v in rule.items()) for k, rule in kinds.items()) + ".")
    header, header_values, month_format = {}, {}, "date"
    for (r, c), v in diff.items():
        if first <= r <= last + step:
            continue
        label = next((ws.cell(r, cc).value for cc in range(c - 1, 0, -1) if isinstance(ws.cell(r, cc).value, str)), None)
        if label is None and r > 1:
            label = ws.cell(r - 1, c).value if isinstance(ws.cell(r - 1, c).value, str) else None
        if not label:
            continue
        f = _fold(label)
        role = next((k for k, words in _ROLE_WORDS.items() if any(w in f for w in words)), None)
        if not role or role in header:
            continue
        header[role] = f"{get_column_letter(c)}{r}"
        if role == "month":
            month_format = "date" if isinstance(v, (dt.date, dt.datetime)) else _month_pattern(str(v))
        else:
            header_values[role] = v
    if header:
        evidence.append("Header cells: " + ", ".join(f"{k} = {v}" for k, v in header.items()) + ".")
    totals = {}
    for row in ws.iter_rows(min_row=last + 1, max_row=last + 6 * step + 6):
        for c in row:
            f = c.value if isinstance(c.value, str) and c.value.startswith("=") else ""
            if not f:
                continue
            if re.search(r"COUNTIF\(", f, re.I) and get_column_letter(c.column) in columns.values():
                totals.setdefault("absence_days", c.coordinate)
            elif re.search(r"SUM\(", f, re.I):
                totals.setdefault("hours", c.coordinate)
    if totals:
        evidence.append("Totals (formulas below the days): " + ", ".join(f"{k} = {v}" for k, v in totals.items()) + ".")
    starts = [times[(r, arrive)] for r in rows if (r, arrive) in times]
    ends = [times[(r, leave)] for r in rows if (r, leave) in times]
    defaults = {}
    if starts and ends:
        defaults = {"start": _mode(starts).strftime("%-H:%M"), "end": _mode(ends).strftime("%-H:%M")}
        evidence.append(f"Usual times in the sample: {defaults['start']}–{defaults['end']}.")
    prof = {"sheet": sheet, "header": header, "month_format": month_format, "days": {"first_row": first, "step": step},
            "columns": columns, "rules": {"workday": {"arrive": "{start}", "leave": "{end}"}, "kinds": kinds},
            "totals": totals, "header_values": header_values, "defaults": defaults}
    return {"profile": validate_profile(prof), "evidence": evidence,
            "note": "A suggestion from the difference — check it (above all the exception rules and the header roles), "
                    "then register the template with this profile."}


def statistics_mean(xs: list[int]) -> float:
    return sum(xs) / len(xs) if xs else 0.0


def _mode(xs: list[dt.time]) -> dt.time:
    return max(set(xs), key=xs.count)


def _month_pattern(s: str) -> str:
    """„2025. október” → "{year}. {month}"; „2025-10” → "{year}-{mm}"; ami más, az dátumként íródik."""
    f = s.strip()
    y = re.search(r"\d{4}", f)
    low = f.lower()
    name = next((n for n in wc.MONTHS if n in low), None)
    if y and name:
        i = low.index(name)
        return (f[:i] + "{month}" + f[i + len(name):]).replace(y.group(0), "{year}")
    m = re.fullmatch(r"(\d{4})([-./]\s?)(\d{1,2})(\.?)", f)
    return "{year}" + m.group(2) + "{mm}" + m.group(4) if m else "date"


def _first_row(ws, rows: list[int], step: int, period: Optional[str]) -> int:
    """A d. nap sora. 1) a sablonban van napszám-oszlop (1, 2, 3 … egymás alatti sorokban) → onnan; 2) különben a
    kitöltött hónap (period) munkanapjaihoz igazítva: az a kezdősor, amelynél minden időt tartalmazó sor munkanap."""
    for col in range(1, min(ws.max_column, 8) + 1):
        seq = [(r, ws.cell(r, col).value) for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, col).value, (int, float))]
        for i, (r, v) in enumerate(seq):
            if v == 1 and all(any(rr == r + k * step and vv == 1 + k for rr, vv in seq) for k in range(1, 5)):
                if r <= rows[0] and rows[-1] <= r + 30 * step:
                    return r
    if not period:
        raise TemplateError("the template has no day-number column — give period (YYYY-MM of the filled sample) so "
                            "the days can be matched")
    y, m = parse_period(period)
    cal = wc.month(y, m)
    for start in range(rows[0] - 30 * step, rows[0] + 1):
        if start < 1:
            continue
        ok = all(0 <= (r - start) // step < len(cal) and (r - start) % step == 0 and cal[(r - start) // step].is_workday
                 for r in rows)
        if ok and sum(1 for d in cal if d.is_workday) >= len(rows):
            return start
    raise TemplateError(f"couldn't match the sample's rows to the workdays of {period}")
