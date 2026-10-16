"""Aláírás és kézjegy: rajzolt, fotózott, szkennelt vagy feltöltött kép → tiszta, átlátszó PNG.

Tisztítás: háttér-becslés (nagy medián) és -kiosztás → az árnyékos/szürke papír eltűnik, Otsu-küszöb, élsimított
alfa, apró pöttyök törlése, körbevágás, egységes tintaszín (a toll színe, vagy amit kérnek).
Tárolás: felhasználónként titkosítva (crypto.py, AES-256-GCM) a Postgresben; DB nélkül a DATA_DIR alatt.
Elhelyezés: PDF (aláírásmező, „Aláírás:” vonal, felirat melletti hely, vagy megadott koordináta) és DOCX
(minta, táblázatcella, felirat utáni hely) — az eredeti formátumban.
"""
from __future__ import annotations

import io
import json
import re
import secrets
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

ANCHOR_WORDS = ("aláírás", "aláírása", "signature", "signed", "kézjegy", "unterschrift", "podpis")


# ---------------------------------------------------------------- tisztítás

def _load_rgba(src: Path):
    from PIL import Image, ImageOps
    from app.convert.steps import _register_plugins
    _register_plugins()
    if src.suffix.lower() == ".pdf":
        import pymupdf
        with pymupdf.open(str(src)) as d:
            pix = d[0].get_pixmap(dpi=200)
            return Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGBA")
    if src.suffix.lower() in (".svg", ".svgz"):
        import subprocess
        png = subprocess.run(["rsvg-convert", "-w", "1200", str(src)], capture_output=True, check=True, timeout=60).stdout
        return Image.open(io.BytesIO(png)).convert("RGBA")
    with Image.open(src) as im:
        im = ImageOps.exif_transpose(im).convert("RGBA")
    im.thumbnail((2000, 2000))                       # telefonfotó: elég 2000 px a tisztításhoz
    return im


def clean(src: Path, color: Optional[str] = None, max_width: int = 1200):
    """→ PIL RGBA kép (átlátszó háttér, egyszínű tinta, körbevágva)."""
    import numpy as np
    from PIL import Image, ImageColor, ImageFilter
    im = _load_rgba(src)
    arr = np.asarray(im).astype(np.float32)
    alpha_in = arr[..., 3]
    rgb = arr[..., :3]
    if alpha_in.min() < 250 and (alpha_in < 10).mean() > 0.2:
        # már átlátszó (rajzolt, vásznon): az alfa a tinta
        ink = alpha_in / 255.0
    else:
        g = np.asarray(im.convert("L")).astype(np.float32)
        k = max(15, (min(g.shape) // 12) | 1)
        bg = np.asarray(Image.fromarray(g.astype(np.uint8)).filter(ImageFilter.MaxFilter(9)).filter(ImageFilter.MedianFilter(min(k, 31 | 1)))).astype(np.float32)
        norm = np.clip(g / np.maximum(bg, 1) * 255.0, 0, 255)           # árnyék, papírszín ki
        hist = np.histogram(norm, bins=256, range=(0, 255))[0].astype(np.float64)
        t = _otsu(hist)
        t = min(t, 200)
        ink = np.clip((t + 25 - norm) / (t + 25 - max(0, t - 60)), 0, 1)
    mask = ink > 0.35
    if mask.sum() < 30:
        raise ValueError("No signature found in the image — use a dark pen on light paper, or draw it.")
    mask = _drop_specks(mask)
    ink = ink * _dilate(mask)
    ys, xs = np.nonzero(ink > 0.15)
    pad = max(4, int(0.04 * max(np.ptp(xs), np.ptp(ys))))
    y0, y1 = max(0, ys.min() - pad), min(ink.shape[0], ys.max() + pad + 1)
    x0, x1 = max(0, xs.min() - pad), min(ink.shape[1], xs.max() + pad + 1)
    ink = ink[y0:y1, x0:x1]
    if color:
        c = ImageColor.getrgb(color)[:3]
    else:
        sel = (ink > 0.6)
        src_rgb = rgb[y0:y1, x0:x1][sel]
        c = tuple(int(v) for v in src_rgb.mean(axis=0)) if len(src_rgb) else (20, 25, 60)
        if max(c) - min(c) < 18:                                  # szürke ceruza / szken → mélyfekete-kék
            c = (22, 24, 38)
        else:                                                     # a toll színe, telítve-sötétítve
            m = max(c) or 1
            c = tuple(int(v * min(1.0, 110 / m)) for v in c)
    out = np.zeros((*ink.shape, 4), dtype=np.uint8)
    out[..., 0], out[..., 1], out[..., 2] = c
    out[..., 3] = (ink * 255).astype(np.uint8)
    res = Image.fromarray(out, "RGBA")
    if res.width > max_width:
        res = res.resize((max_width, round(res.height * max_width / res.width)), Image.LANCZOS)
    return res


def _otsu(hist) -> float:
    import numpy as np
    total = hist.sum()
    sum_all = (np.arange(256) * hist).sum()
    wb = sb = 0.0
    best, thr = 0.0, 127.0
    for i in range(256):
        wb += hist[i]
        if wb == 0:
            continue
        wf = total - wb
        if wf == 0:
            break
        sb += i * hist[i]
        mb, mf = sb / wb, (sum_all - sb) / wf
        between = wb * wf * (mb - mf) ** 2
        if between > best:
            best, thr = between, float(i)
    return thr


def _drop_specks(mask):
    import cv2
    import numpy as np
    n, lab, stats, _ = cv2.connectedComponentsWithStats(mask.astype(np.uint8), 8)
    if n <= 1:
        return mask
    areas = stats[1:, cv2.CC_STAT_AREA]
    keep_min = max(12, int(areas.max() * 0.02))
    keep = np.zeros(n, dtype=bool)
    keep[1:] = areas >= keep_min
    return keep[lab]


def _dilate(mask):
    import cv2
    import numpy as np
    return cv2.dilate(mask.astype(np.uint8), np.ones((3, 3), np.uint8)).astype(np.float32)


def to_png(im) -> bytes:
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)
    return buf.getvalue()


# ---------------------------------------------------------------- tárolás (titkosítva)

def _fs_dir(uid: str) -> Path:
    import hashlib
    from app.config import settings
    d = Path(settings.DATA_DIR) / "secure" / "signatures" / hashlib.sha256(uid.encode()).hexdigest()[:24]
    d.mkdir(parents=True, exist_ok=True)
    return d


async def save(uid: str, png: bytes, name: str, kind: str = "signature") -> dict[str, Any]:
    from PIL import Image
    from app import db
    from app.formkit import crypto
    sid = "sig_" + secrets.token_urlsafe(12)
    with Image.open(io.BytesIO(png)) as im:
        w, h = im.size
    nonce, data = crypto.seal(uid, "signature", png, sid.encode())
    meta = {"id": sid, "name": name[:80], "kind": kind, "width": w, "height": h,
            "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds")}
    if db.enabled():
        await db.signature_put(uid, sid, meta["name"], kind, nonce, data, w, h)
    else:
        d = _fs_dir(uid)
        (d / f"{sid}.bin").write_bytes(nonce + data)
        (d / f"{sid}.json").write_text(json.dumps(meta, ensure_ascii=False), "utf-8")
    return meta


async def list_(uid: str) -> list[dict[str, Any]]:
    from app import db
    if db.enabled():
        return await db.signature_list(uid)
    d = _fs_dir(uid)
    items = [json.loads(p.read_text("utf-8")) for p in d.glob("*.json")]
    return sorted(items, key=lambda m: m["created_at"], reverse=True)


async def load(uid: str, sid: str) -> bytes:
    from app import db
    from app.formkit import crypto
    if not re.fullmatch(r"sig_[A-Za-z0-9_-]{8,40}", sid or ""):
        raise KeyError(sid)
    if db.enabled():
        row = await db.signature_get(uid, sid)
        if not row:
            raise KeyError(sid)
        nonce, data = bytes(row["nonce"]), bytes(row["data"])
    else:
        p = _fs_dir(uid) / f"{sid}.bin"
        if not p.exists():
            raise KeyError(sid)
        raw = p.read_bytes()
        nonce, data = raw[:12], raw[12:]
    return crypto.open_(uid, "signature", nonce, data, sid.encode())


async def delete(uid: str, sid: str) -> bool:
    from app import db
    if db.enabled():
        return await db.signature_delete(uid, sid)
    d = _fs_dir(uid)
    ok = (d / f"{sid}.bin").exists()
    for ext in (".bin", ".json"):
        (d / f"{sid}{ext}").unlink(missing_ok=True)
    return ok


# ---------------------------------------------------------------- elhelyezés: PDF

def stamp_pdf_page(page, box: list[float], png: Path) -> list[float]:
    """Az aláírás a dobozba / vonalra: szélességben a dobozba fér, magasságban a vonal fölé nyúlhat
    (egy aláírásmező vagy egy „______” sor alacsony), az alja a doboz aljára ül."""
    import pymupdf
    from PIL import Image
    with Image.open(png) as im:
        iw, ih = im.size
    x0, y0, x1, y1 = box
    bw, bh = x1 - x0, y1 - y0
    max_h = max(bh * 2.2, 30.0)
    s = min(bw / iw, max_h / ih)
    w, h = iw * s, ih * s
    rect = pymupdf.Rect(x0 + 2, y1 - h + min(4.0, bh * 0.2), x0 + 2 + w, y1 + min(4.0, bh * 0.2))
    page.insert_image(rect, filename=str(png), keep_proportion=True, overlay=True)
    return [rect.x0, rect.y0, rect.x1, rect.y1]


def _anchor_box_pdf(doc, words: tuple[str, ...]) -> Optional[tuple[int, list[float]]]:
    for pno in range(doc.page_count - 1, -1, -1):          # az aláírás jellemzően a végén van
        page = doc[pno]
        for w in words:
            hits = page.search_for(w)
            if hits:
                r = hits[-1]
                return pno + 1, [r.x1 + 6, r.y0 - 2, min(page.rect.width - 36, r.x1 + 186), r.y1 + 2]
    return None


def place_pdf(src: Path, dst: Path, png: Path, *, placement: str = "auto", page: Optional[int] = None,
              x: Optional[float] = None, y: Optional[float] = None, width: Optional[float] = None,
              field: Optional[str] = None) -> dict[str, Any]:
    import pymupdf
    from app.formkit import detect, schema as S
    where: Optional[tuple[int, list[float]]] = None
    how = placement
    if placement in ("auto", "field"):
        sch = detect.recognize_pdf(src)
        sigs = [f for f in sch["fields"] if f["type"] == "signature" and f.get("box")]
        if field:
            key = S.ascii_fold(field)
            sigs = [f for f in sch["fields"] if f.get("box") and key in (S.ascii_fold(f["name"]), S.ascii_fold(f.get("label") or ""))] or \
                   [f for f in sigs if key in S.ascii_fold(f.get("label") or "")]
        if sigs:
            f = sigs[-1] if not field else sigs[0]
            where, how = (f["page"], f["box"]), "field"
        elif placement == "field":
            raise ValueError("No signature field or signature line found — give x and y, or a field name.")
    with pymupdf.open(str(src)) as d:
        if where is None and placement in ("auto", "anchor"):
            words = (field,) if (field and placement == "anchor") else ANCHOR_WORDS + tuple(w.capitalize() for w in ANCHOR_WORDS)
            where = _anchor_box_pdf(d, words)
            how = "anchor"
            if where is None and placement == "anchor":
                raise ValueError(f"Couldn't find {field or 'a signature label'} in the document.")
        if where is None:
            pno = (page or d.page_count)
            if not 1 <= pno <= d.page_count:
                raise ValueError(f"page {pno} doesn't exist")
            pg = d[pno - 1]
            wdt = width or 160.0
            if x is None or y is None:              # automatikus: utolsó oldal jobb alsó része
                x, y = pg.rect.width - wdt - 56, pg.rect.height - 140
                how = "bottom-right"
            else:
                how = "xy"
            where = (pno, [x, y, x + wdt, y + wdt * 0.32])
        elif width:
            b = where[1]
            where = (where[0], [b[0], b[1], b[0] + width, b[3]])
        rect = stamp_pdf_page(d[where[0] - 1], where[1], png)
        d.save(str(dst), garbage=3, deflate=True)
    return {"page": where[0], "rect": [round(v, 1) for v in rect], "placement": how}


# ---------------------------------------------------------------- elhelyezés: DOCX

def docx_picture_at(p, start: int, end: int, png: Path, width_cm: float = 4.5) -> None:
    """A bekezdés [start, end) szakasza helyére egy beágyazott kép (az aláírás)."""
    from docx.shared import Cm
    from app.formkit.fill import _replace_in_runs
    target = None
    pos = 0
    for run in p.runs:
        if pos + len(run.text) >= start:
            target = run
            break
        pos += len(run.text)
    _replace_in_runs(p, start, end, "")
    new = p.add_run()
    new.add_picture(str(png), width=Cm(width_cm))
    if target is not None and target._r is not new._r:
        target._r.addnext(new._r)


def place_docx(src: Path, dst: Path, png: Path, *, field: Optional[str] = None, width_cm: Optional[float] = None) -> dict[str, Any]:
    from docx import Document
    from docx.shared import Cm
    from app.formkit import detect, schema as S
    from app.formkit.detect import BLANK_RE
    sch = detect.recognize_docx(src)
    sigs = [f for f in sch["fields"] if f["type"] == "signature"]
    if field:
        key = S.ascii_fold(field)
        sigs = [f for f in sch["fields"] if key in (S.ascii_fold(f["name"]), S.ascii_fold(f.get("label") or ""))] or sigs
    doc = Document(str(src))
    wcm = width_cm or 4.5
    if sigs:
        f = sigs[0]
        a = f["anchor"]
        if a["kind"] == "docx_text":
            p = detect_para(doc, a["loc"])
            m = list(BLANK_RE.finditer(p.text))[a["nth"]]
            docx_picture_at(p, m.start(), m.end(), png, wcm)
        elif a["kind"] == "docx_cell":
            p = doc.tables[a["t"]].cell(a["r"], a["c"]).paragraphs[0]
            p.add_run().add_picture(str(png), width=Cm(wcm))
        else:
            sigs = []
        if sigs:
            doc.save(str(dst))
            return {"placement": "field", "field": f["name"]}
    for p in reversed(doc.paragraphs):
        if any(w in p.text.lower() for w in ANCHOR_WORDS):
            p.add_run(" ")
            p.add_run().add_picture(str(png), width=Cm(wcm))
            doc.save(str(dst))
            return {"placement": "anchor", "text": p.text.strip()[:60]}
    doc.add_paragraph().add_run().add_picture(str(png), width=Cm(wcm))
    doc.save(str(dst))
    return {"placement": "end"}


def detect_para(doc, loc):
    from app.formkit.fill import _docx_para
    return _docx_para(doc, loc)
