"""Kézi szerkesztők — a Craft-motorok WASM-buildje (változatlan upstream release-ek).

    /editor/{app}/            — az app index.html-je, beszúrt aice_bridge.js-sel
                                (app: wordcraft | photocraft | vectorcraft | deckcraft | pdfcraft | designcraft | gridcraft)
    /editor/{app}/{asset}     — wasm/js (előtömörített .gz, ha a kliens kéri)
    POST /api/editor/save     — a szerkesztő Mentés/Exportálás bájtjai → új output file_id

A dashboard iframe-ben nyitja: /editor/<app>/?src=<fájl-URL>&name=<fájlnév>&save=/api/editor/save?source=<file_id>
"""
from __future__ import annotations

import asyncio
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse

from app import events
from app.config import settings
from app.storage import input_path, output_path, public_url, save_bytes, write_meta
from app.storage.files import new_file_id
from app.tools import crafts
from app.tools.wordcraft import SAVE_FORMATS as WORD_SAVE

router = APIRouter()
log = logging.getLogger("aice.editor")

BRIDGE_JS = Path(__file__).parent.parent / "static" / "editor" / "aice_bridge.js"
_ASSET_RE = re.compile(r"^[A-Za-z0-9._-]+$")
_MEDIA = {".wasm": "application/wasm", ".js": "text/javascript", ".html": "text/html"}
_IMMUTABLE = "public, max-age=31536000, immutable"
# Amit a szerkesztők menteni/exportálni tudnak (a mentés-végpont ezt engedi be).
_SAVE_EXTS = WORD_SAVE | crafts.PHOTO_SAVE | crafts.VECTOR_SAVE | crafts.OFFICE_SAVE
# natív mentés → (motor, eredeti formátumok, amelyekre visszaalakítjuk)
_OFFICE_ROUNDTRIP = {"deckcraft": ("deck", {"pptx"})}


def _web_dirs() -> dict[str, Path]:
    return {
        "wordcraft": settings.WORDCRAFT_WEB_DIR,
        "photocraft": settings.PHOTOCRAFT_WEB_DIR,
        "vectorcraft": settings.VECTORCRAFT_WEB_DIR,
        **{spec["exe"]: crafts.office_web_dir(spec["exe"]) for spec in crafts.OFFICE.values()},
    }


def _web_dir(app: str) -> Path:
    d = _web_dirs().get(app)
    if d is None:
        raise HTTPException(404, f"Unknown editor: {app}")
    if not (d / "index.html").is_file():
        raise HTTPException(503, f"The {app} editor isn't installed on this server.")
    return d



def editor_available() -> bool:
    return (settings.WORDCRAFT_WEB_DIR / "index.html").is_file()


def crafts_status() -> dict[str, Any]:
    from app.tools import cutout, wordcraft
    clis = {"wordcraft": wordcraft.available(), "photocraft": crafts.photo_cli() is not None,
            "vectorcraft": crafts.vector_cli() is not None,
            **{spec["exe"]: crafts.office_cli(a) is not None for a, spec in crafts.OFFICE.items()}}
    st: dict[str, Any] = {app: {"cli": clis[app], "editor": (d / "index.html").is_file()}
                          for app, d in _web_dirs().items()}
    st["cutout"] = cutout.available()
    return st


@router.get("/editor", include_in_schema=False)
@router.get("/editor/", include_in_schema=False)
async def editor_legacy(request: Request):
    # Régi cím (csak WordCraft volt): /editor/?… → /editor/wordcraft/?…
    qs = request.url.query
    return RedirectResponse("/editor/wordcraft/" + (f"?{qs}" if qs else ""))


@router.get("/editor/{app}", include_in_schema=False)
async def editor_slash(app: str, request: Request):
    qs = request.url.query
    return RedirectResponse(f"/editor/{app}/" + (f"?{qs}" if qs else ""))


@router.get("/editor/{app}/", include_in_schema=False)
async def editor_index(app: str) -> HTMLResponse:
    html = (_web_dir(app) / "index.html").read_text("utf-8")
    tag = '<script src="./aice_bridge.js"></script>\n'
    # A bridge-nek a WASM-modul ELŐTT kell futnia (TrunkApplicationStarted-figyelő).
    marker = '<script type="module">'
    html = html.replace(marker, tag + marker, 1) if marker in html else html.replace("</head>", tag + "</head>", 1)
    return HTMLResponse(html, headers={"Cache-Control": "no-cache"})


@router.get("/editor/{app}/{asset}", include_in_schema=False)
async def editor_asset(app: str, asset: str, request: Request):
    if asset == "aice_bridge.js":
        return FileResponse(BRIDGE_JS, media_type="text/javascript", headers={"Cache-Control": "no-cache"})
    if not _ASSET_RE.match(asset) or asset.endswith(".gz"):
        raise HTTPException(404)
    d = _web_dir(app)
    p = d / asset
    if not p.is_file():
        raise HTTPException(404)
    suffix = p.suffix.lower()
    media = _MEDIA.get(suffix)
    headers = {"Cache-Control": _IMMUTABLE if suffix in (".wasm", ".js") else "no-cache", "Vary": "Accept-Encoding"}
    gz = d / f"{asset}.gz"
    if gz.is_file() and "gzip" in request.headers.get("accept-encoding", ""):
        return FileResponse(gz, media_type=media, headers={**headers, "Content-Encoding": "gzip"})
    return FileResponse(p, media_type=media, headers=headers)


_PLACEHOLDER_STEMS = {"untitled", "document", "letter", "report", "cv", "sample", "presentation", "your title here", "your headline here"}


def _title_from(path: Path, ext: str) -> str:
    """Az első címsor / az első nem üres sor (docx), ill. az első dia címe (pptx) — fájlnévnek."""
    try:
        if ext == "docx":
            from docx import Document
            paras = [p for p in Document(str(path)).paragraphs if p.text.strip()]
            head = next((p for p in paras if (p.style.name or "").lower().startswith(("title", "heading"))), paras[0] if paras else None)
            text = head.text if head else ""
        elif ext == "pptx":
            from pptx import Presentation
            slides = list(Presentation(str(path)).slides)
            text = slides[0].shapes.title.text if slides and slides[0].shapes.title is not None else ""
        else:
            return ""
    except Exception as e:
        log.debug("cím kiolvasása (%s): %s", path.name, e)
        return ""
    text = re.sub(r"[\\/:*?\"<>|\s]+", " ", text).strip(" .")
    return text[:60].rsplit(" ", 1)[0] if len(text) > 60 else text


@router.post("/api/editor/save")
async def editor_save(
    request: Request,
    name: str = Query(..., description="file name given by the editor"),
    source: Optional[str] = Query(None, description="a megnyitott eredeti file_id"),
) -> dict[str, Any]:
    data = await request.body()
    if not data:
        raise HTTPException(400, "Empty file.")
    if len(data) > settings.EDITOR_MAX_SAVE_MB * 1024 * 1024:
        raise HTTPException(413, f"nagyobb, mint {settings.EDITOR_MAX_SAVE_MB} MB")
    fname = Path(name).name or "document.docx"
    ext = Path(fname).suffix.lstrip(".").lower()
    if ext not in _SAVE_EXTS:
        raise HTTPException(400, f"Unexpected file type: .{ext}")
    if source and not (input_path(source).exists() or output_path(source).exists()):
        source = None

    fid = new_file_id(ext)
    await save_bytes(output_path(fid), data)
    # A natív formátumú mentést (Ctrl+S) visszaalakítjuk az eredeti Office-formátumra, ha abból nyitották:
    # a felhasználó pptx-et nyitott, pptx-et vár vissza (a .deckcraft csak a mi motorunknak olvasható).
    want = _OFFICE_ROUNDTRIP.get(ext)
    src_ext = source.rsplit(".", 1)[-1].lower() if source else ""
    if want and src_ext in want[1]:
        native, fid = output_path(fid), new_file_id(src_ext)
        try:
            await crafts.office_convert(want[0], native, output_path(fid))
            ext, fname = src_ext, f"{Path(fname).stem}.{src_ext}"
            data = output_path(fid).read_bytes()
        except Exception as e:  # marad a natív mentés — az is megnyitható nálunk
            log.warning("visszaalakítás %s → %s nem sikerült: %s", native.name, src_ext, e)
            fid = native.name
        else:
            native.unlink(missing_ok=True)
    # Változatlan mentés (pl. a „Kérj változtatást” előtti automatikus mentés egy puszta kattintás után):
    # NINCS új verzió — a forrás marad, a metájával (pakli-leírás, lánc) együtt. Kétes esetben új verzió lesz.
    if source and ext in ("pptx", "docx") and source.rsplit(".", 1)[-1].lower() == ext:
        from app.storage.metadata import read_meta
        from app.tools.content_digest import same_content
        src_path = output_path(source) if output_path(source).exists() else input_path(source)
        if await asyncio.to_thread(same_content, src_path, output_path(fid)):
            output_path(fid).unlink(missing_ok=True)
            meta = read_meta(source) or {}
            return {"file_id": source, "url": public_url(source), "size_kb": src_path.stat().st_size // 1024,
                    "label": meta.get("label") or fname, "unchanged": True}
    stamp = datetime.utcnow().strftime("%H:%M")
    stem = re.sub(r"( \(edited [^)]*\))+$", "", Path(fname).stem)
    if stem.lower() in _PLACEHOLDER_STEMS:      # „Untitled.docx” helyett a dokumentum saját címe (a Files-ban megkülönböztethető)
        stem = _title_from(output_path(fid), ext) or stem
    label = f"{stem} (edited {stamp} UTC).{ext}"
    write_meta(fid, label=label, source_file_id=source, operation="manual edit", original_filename=fname)
    result = {"file_id": fid, "url": public_url(fid), "size_kb": len(data) // 1024, "label": label}
    events.emit("tool_done", tool="editor_save", params={"source": source, "name": fname}, result=result)
    return result
