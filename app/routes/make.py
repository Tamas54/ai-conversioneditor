"""A termék „Make” végpontjai — egy rövid űrlapból kész fájl (a webes app csempéi mögött).

    POST /api/make/deck     {brief, slides?, images?: stock|ai|none, accent?}   → PPTX (a tartalmat a GLM írja)
    POST /api/make/social   {headline, subtitle?, handle?, sizes[], image}        → PNG-k (+PDF) méretenként
    POST /api/make/design   {template, data, size?}                               → PDF + PNG
    POST /api/make/photo    {file_id, action, background?, blur?}                 → új kép
    POST /api/make/book     {file_id, title?, author?, cover?}                    → könyv-PDF (borító + fejezetek)
    POST /api/make/pdf      {action, file_id | file_ids, …}                      → PDF-műveletek (PdfCraft)

Képforrás (image/background/cover): {"file_id"} | {"stock": "words"} | {"ai": "prompt"}. Az AI-kép kreditbe kerül
(vision_ops.generate_image vonja le); a pakli előre ellenőrzi, hogy elég-e az egyenleg.
"""
from __future__ import annotations

import json
import re
import logging
from pathlib import Path
from typing import Any, Optional

from fastapi import APIRouter, HTTPException, Request

from app import events, tenancy
from app.config import settings

log = logging.getLogger("aice.make")
router = APIRouter()

SOCIAL_SIZES = {"landscape", "square", "story"}


def _image_spec(src: Optional[dict]) -> Optional[Any]:
    """A felület képválasztója → images.resolve formátum."""
    if not src:
        return None
    if src.get("file_id"):
        return src["file_id"]
    if src.get("url"):                       # a képválasztó rácsából kiválasztott stockkép (forrásmegjelöléssel)
        return {"url": src["url"], **({"credit": src["credit"]} if src.get("credit") else {})}
    if src.get("stock"):
        return {"stock": src["stock"]}
    if src.get("ai"):
        return {"generate": src["ai"]}
    return None


async def _materialize(spec: Any) -> Optional[str]:
    """Stock/AI képet EGYSZER letölt/generál és fájlként ment — több méret/felhasználás ugyanazt kapja
    (különben méretenként újra generálna és újra kreditet vonna le)."""
    if spec is None or isinstance(spec, str):
        return spec
    from app.storage import output_path, write_meta
    from app.storage.files import new_file_id
    from app.tools import images
    r = await images.resolve(spec, aspect=None, max_px=2400)
    try:
        if r.source_file_id:          # generált kép: már fájl
            return r.source_file_id
        fid = new_file_id("jpg")
        output_path(fid).write_bytes(r.path.read_bytes())
        write_meta(fid, operation="stock photo", label="stock.jpg", extra={"credit": r.attribution})
        return fid
    finally:
        r.path.unlink(missing_ok=True)


async def _credits_left() -> Optional[int]:
    from app import db
    uid = tenancy.current_user_id()
    if uid == tenancy.ADMIN_ID or not db.enabled():
        return None
    u = await db.user_get(uid)
    return (u or {}).get("credits", 0)


def _fail(e: Exception) -> HTTPException:
    from app import db
    if isinstance(e, db.InsufficientCredits) and tenancy.is_guest():
        from app.auth import ACCOUNT_NEEDED
        return HTTPException(401, ACCOUNT_NEEDED)
    if isinstance(e, db.InsufficientCredits):
        return HTTPException(402, "Not enough credits — this step uses AI. Add credits, or use stock photos and the steps without AI.")
    log.exception("make failed")
    return HTTPException(500, str(e)[:300])


async def _llm_json(system: str, user: str) -> dict[str, Any]:
    from app.llm.siliconflow import _call_model, _strip_json_fences
    last: Exception = RuntimeError("no model")
    for model in (settings.LLM_MODEL_PRIMARY, settings.LLM_MODEL_FALLBACK):
        try:
            raw = await _call_model(model, [{"role": "system", "content": system}, {"role": "user", "content": user}],
                                    thinking=False)
            return json.loads(_strip_json_fences(raw))
        except Exception as e:  # következő modell
            last = e
            log.warning("deck-vázlat %s modellel nem sikerült: %s", model, e)
    raise last


_DECK_SYSTEM = """You write presentation decks as JSON for a deck builder. Reply with ONE JSON object:
{"title": str, "footer": str, "slides": [ ... ]}
Each slide has "layout" and fields:
  cover {title, subtitle, kicker}            section {title, subtitle, number}
  bullets {title, bullets[3-5], image?}      image_left|image_right {title, bullets[2-4] or text, image}
  full_image {title, caption, image}         quote {quote, author}
  stats {title, stats:[{value, label}] 2-4}  chart {title, chart_type: column|bar|line|pie, categories[], series:[{name, values[]}]}
  table {title, headers[], rows[][]}         closing {title, subtitle}
Rules: first slide cover, last slide closing; vary layouts; bullets ≤ 70 characters, max 5 per slide;
only use numbers the brief gives or well-known facts; write in the language of the brief.
For each slide that should have a picture set "image": {"query": "2-4 plain English words"} — a stock-photo
search, concrete and visual (e.g. "wind turbines", "team meeting"). Cover and full_image always get an image."""


_FROM_DOC = """
The user gives a DOCUMENT to turn into this deck. Follow its structure — its main headings become sections or slides,
in order — and keep its key facts, names, numbers and wording. Use ONLY what the document says: no invented facts.
Write in the language of the document (unless the user's instructions ask for another).
The document's pictures are marked [picture DOCIMG#k]. Put a picture on the slide where it belongs by setting
"image": "DOCIMG#k" (each at most once; prefer them to stock photos). Tables and figures may become table or chart slides."""


async def _deck_style(body: dict[str, Any], ai_ok: bool) -> tuple[Optional[str], Optional[dict[str, Any]]]:
    """A stílusminta: (teljes sablon file_id, vagy téma). pptx-sablon / pptx-téma AI nélkül; PDF/kép témája AI-val."""
    sid = body.get("style_file_id")
    if not sid:
        return None, None
    is_pptx = str(sid).lower().endswith((".pptx", ".potx"))
    if is_pptx and body.get("use_template", True):
        from app.tools import pptx_template
        from app.tools.crafts import existing as _ex
        if pptx_template.usable(_ex(sid)):          # dizájn a mesterdián → teljes sablon; különben téma
            return sid, None
    if not is_pptx and not ai_ok:
        return None, None
    from app.tools import style_ref
    try:
        return None, await style_ref.theme_from(sid)
    except Exception as e:
        log.warning("stílusminta: %s", e)
        return None, None


def _deck_look(deck: dict[str, Any], template_id: Optional[str], style: Optional[dict[str, Any]], accent: Optional[str]) -> None:
    if template_id:                    # pptx-minta = teljes sablon (elrendezés, háttér, logó, betű, szín)
        deck["template_file_id"] = template_id
    elif style:
        deck["theme"] = {k: v for k, v in style.items() if k != "mood"}
    elif accent:
        deck["theme"] = {"accent": accent}


@router.post("/api/make/deck")
async def make_deck(body: dict[str, Any]) -> dict[str, Any]:
    """Három mód: blank (üres pakli, AI nélkül → rögtön a szerkesztő), brief (leírásból), document (egy meglévő
    dokumentumból — szerkezet, tények és képek a dokumentumból)."""
    from app.tools import pptx_deck
    mode = body.get("mode") or ("document" if body.get("source_file_id") else "brief")
    if mode == "blank":
        template_id, style = await _deck_style(body, ai_ok=not tenancy.is_guest())
        title = (body.get("title") or "").strip()[:120] or "Presentation title"
        deck = {"title": title, "slides": [
            {"layout": "cover", "title": title, "subtitle": "Subtitle"},
            {"layout": "bullets", "title": "Slide title", "bullets": ["First point", "Second point", "Third point"]}]}
        _deck_look(deck, template_id, style, body.get("accent"))
        try:
            out = await pptx_deck.build(deck)
        except Exception as e:
            raise _fail(e)
        events.emit("tool_done", tool="make_deck", result={"file_id": out["file_id"]})
        return out
    if tenancy.is_guest():
        from app.auth import ACCOUNT_NEEDED
        raise HTTPException(401, ACCOUNT_NEEDED)
    brief = (body.get("brief") or "").strip()
    source = None
    if mode == "document":
        if not body.get("source_file_id"):
            raise HTTPException(400, "Choose the document to make slides from.")
        from app.tools import doc_source
        try:
            source = await doc_source.read(body["source_file_id"])
        except (ValueError, FileNotFoundError) as e:
            raise HTTPException(400, str(e))
    elif len(brief) < 10:
        raise HTTPException(400, "Describe the presentation in a sentence or two.")
    auto = source is not None and not body.get("slides")
    n = max(4, min(15, int(body.get("slides") or 8)))
    from app import billing
    await billing.use("deck")
    pics = body.get("images") or "stock"          # stock | ai | doc (csak a dokumentum képei) | none
    template_id, style = await _deck_style(body, ai_ok=True)
    if pics == "ai":  # az LLM-vázlat ELŐTT: kredit nélkül ne várjon fél percet a semmire
        left = await _credits_left()
        if left is not None and left < settings.IMAGE_CREDIT_COST:
            raise HTTPException(402, f"AI images need credits (you have {left}). Choose stock photos or add credits.")
    try:
        style_note = f"\n\nMatch this visual style and tone: {style['mood']}" if style and style.get("mood") else ""
        count = "Choose the number of slides (5 to 15) to fit the document." if auto else f"Make exactly {n} slides."
        if source:
            user = ((f"Instructions from the user: {brief}\n\n" if brief else "") + count + style_note
                    + f"\n\nDOCUMENT “{source['title']}”" + (" (long — only its first part is given)" if source["truncated"] else "")
                    + f":\n{source['text']}")
            spec = await _llm_json(_DECK_SYSTEM + _FROM_DOC, user)
        else:
            spec = await _llm_json(_DECK_SYSTEM, f"Brief:\n{brief}\n\n{count}{style_note}")
        slides = [s for s in spec.get("slides") or [] if isinstance(s, dict) and s.get("layout") in pptx_deck.LAYOUTS]
        if not slides:
            raise HTTPException(502, "The assistant returned no usable slides — try again or rephrase the brief.")
        doc_pics = {p["token"]: p["file_id"] for p in (source or {}).get("pictures", [])}
        used: set[str] = set()

        def doc_pic(v: Any) -> Optional[str]:
            tok = v if isinstance(v, str) else (v.get("doc") or v.get("query")) if isinstance(v, dict) else None
            tok = (re.search(r"DOCIMG#\d+", str(tok or "")) or [None])[0]
            if tok in doc_pics and tok not in used:
                used.add(tok)
                return doc_pics[tok]
            return None
        for s in slides:                   # a dokumentum képei: a helyükre (galériában is)
            if isinstance(s.get("images"), list):
                s["images"] = [{"file_id": f} for f in (doc_pic(x) for x in s["images"]) if f]
        wanted = [s for s in slides if isinstance(s.get("image"), dict) and not doc_pic_peek(s.get("image"), doc_pics)]
        if pics == "ai":
            left = await _credits_left()
            if left is not None and left < len(wanted) * settings.IMAGE_CREDIT_COST:
                raise HTTPException(402, f"This deck needs {len(wanted) * settings.IMAGE_CREDIT_COST} credits for AI images; "
                                         f"you have {left}. Choose stock photos or add credits.")
        for s in slides:
            img = s.get("image")
            fid = doc_pic(img) if img is not None else None
            if fid:
                s["image"] = {"file_id": fid}
                continue
            q = img.get("query") if isinstance(img, dict) else None
            if pics in ("none", "doc") or not q or "DOCIMG#" in str(q):
                s.pop("image", None)
            elif pics == "ai":
                s["image"] = {"generate": f"{q}, editorial photograph, natural light, no text"}
            else:
                s["image"] = {"stock": q}
        deck = {"title": spec.get("title") or slides[0].get("title"), "footer": spec.get("footer"), "slides": slides}
        _deck_look(deck, template_id, style, body.get("accent"))
        out = await pptx_deck.build(deck)
    except HTTPException:
        raise
    except Exception as e:
        raise _fail(e)
    if source:
        from app.storage import write_meta
        write_meta(out["file_id"], source_file_id=body["source_file_id"], operation="slides from document")
    events.emit("tool_done", tool="make_deck", result={"file_id": out["file_id"]})
    return out


def doc_pic_peek(v: Any, doc_pics: dict[str, str]) -> bool:
    """Dokumentum-kép-e a kép-spec (a kreditszámláláshoz — nem fogyasztja el a jelölőt)."""
    return bool(re.search(r"DOCIMG#\d+", str(v or ""))) and any(t in str(v) for t in doc_pics)


def _label_outputs(res: dict[str, Any], stem: str, group: Optional[str] = None, variant: Optional[str] = None) -> None:
    """A dizájn-kimenetek neve a címből (a motor belső „dizajn” neve helyett) — png, pdf, szerkeszthető .vectorcraft.
    group/variant: egy készítés testvér-fájljai (méretek × formátumok) — a felület ebből csoportosít és vált köztük."""
    from app.storage import write_meta
    for f in res.get("files") or []:
        ext = f["file_id"].rsplit(".", 1)[-1]
        f["label"] = f"{stem}.{ext}"
        write_meta(f["file_id"], label=f["label"], extra={k: v for k, v in (("group", group), ("variant", variant)) if v})


def _new_group() -> str:
    """Egy készítés azonosítója (a testvér-fájlok közös metája)."""
    import uuid
    return uuid.uuid4().hex[:12]


@router.post("/api/make/social")
async def make_social(body: dict[str, Any]) -> dict[str, Any]:
    from app.tools import design
    headline = (body.get("headline") or "").strip()
    if not headline:
        raise HTTPException(400, "Add a headline.")
    sizes = [s for s in body.get("sizes") or ["square"] if s in SOCIAL_SIZES] or ["square"]
    files, group = [], _new_group()
    try:
        from app.brands import user as user_brands
        brand = await user_brands.resolve(body.get("brand"))
        image = await _materialize(_image_spec(body.get("image")))
        if image and body.get("remove_background") and body.get("background"):
            # alany kivágása + új háttér, egyszer — minden méret ezt használja
            from app.tools import crafts
            bg_id = await _materialize(_image_spec(body["background"]))
            comp = await crafts.photo_run(file_id=image, remove_background=True, background_file_id=bg_id,
                                          background_blur=float(body.get("blur") or 12), save_format="jpg", quality=90)
            image = comp["file_id"]
        for size in sizes:
            res = await design.render("social", {"title": headline, "subtitle": body.get("subtitle"),
                                                 "handle": body.get("handle"), "image": image},
                                      size=size, brand=brand,
                                      theme={"accent": body["accent"]} if body.get("accent") and not brand else None,
                                      export=["png", "pdf", *(["vectorcraft"] if body.get("editable") else [])])
            _label_outputs(res, f"{headline[:60]} ({size})", group=group, variant=size)
            files.append({"size": size, **res})
    except HTTPException:
        raise
    except Exception as e:
        raise _fail(e)
    events.emit("tool_done", tool="make_social", result={"files": [f["file_id"] for f in files]})
    return {"posts": files}


@router.post("/api/make/blank")
async def make_blank(body: dict[str, Any]) -> dict[str, Any]:
    """„New”: üres fájl a látható szerkesztőbe — {kind: social|print|photo|book|pdf, size, orientation?, background?,
    width?, height?}. AI és feltöltés nélkül (vendégnek is)."""
    from app.tools import blank
    try:
        out = await blank.make(body)
    except ValueError as e:
        raise HTTPException(400, str(e))
    except Exception as e:
        raise _fail(e)
    events.emit("tool_done", tool="make_blank", params={"kind": body.get("kind")}, result={"file_id": out["file_id"]})
    return out


@router.post("/api/make/design")
async def make_design(body: dict[str, Any]) -> dict[str, Any]:
    from app.tools import design
    data = dict(body.get("data") or {})
    for k in ("image", "cover_image"):
        if isinstance(data.get(k), dict):
            data[k] = _image_spec(data[k])
    try:
        from app.brands import user as user_brands
        brand = await user_brands.resolve(body.get("brand"))     # a saját arculat magától (kikapcsolható: "none")
        out = await design.render(body.get("template") or "poster", data, size=body.get("size"), brand=brand,
                                  theme={"accent": body["accent"]} if body.get("accent") and not brand else None,
                                  export=["pdf", "png", *(["vectorcraft"] if body.get("editable") else [])])
    except Exception as e:
        raise _fail(e)
    _label_outputs(out, (data.get("title") or data.get("recipient") or str(body.get("template") or "design"))[:60], group=_new_group())
    events.emit("tool_done", tool="make_design", result={"file_id": out["file_id"]})
    return out


@router.post("/api/make/photo")
async def make_photo(body: dict[str, Any]) -> dict[str, Any]:
    from app.tools import crafts
    fid = body.get("file_id")
    if not fid:
        raise HTTPException(400, "Choose a photo first.")
    action = body.get("action") or "remove_background"
    try:
        if action == "remove_background":
            out = await crafts.photo_run(file_id=fid, remove_background=True, save_format="png")
        elif action == "replace_background":
            bg_id = await _materialize(_image_spec(body.get("background")))
            if not bg_id:
                raise HTTPException(400, "Pick a background: a stock search, an AI prompt or one of your images.")
            out = await crafts.photo_run(file_id=fid, remove_background=True, background_file_id=bg_id,
                                         background_blur=float(body.get("blur") or 12), save_format="jpg", quality=90)
        elif action == "enhance":
            out = await crafts.photo_run(file_id=fid, commands=[{"command": "image.autoTone"}, {"command": "image.autoColor"}])
        else:
            raise HTTPException(400, f"unknown action: {action}")
    except HTTPException:
        raise
    except Exception as e:
        raise _fail(e)
    # a név mondja meg, melyik a javított változat („urge (enhanced).jpg”), és a verziólánc is erre épül
    from app.storage import display_label, write_meta
    label = f"{Path(display_label(fid)).stem} ({PHOTO_LABELS[action]}).{out['file_id'].rsplit('.', 1)[1]}"
    write_meta(out["file_id"], source_file_id=fid, operation=f"photo {action}", label=label)
    out["label"] = label
    events.emit("tool_done", tool="make_photo", result={"file_id": out["file_id"]})
    return out


PHOTO_LABELS = {"enhance": "enhanced", "remove_background": "cut out", "replace_background": "new background"}


def _book_front(path: Path, title: str, author: Optional[str], cover: Optional[Path], size: str, ink: bool = True) -> None:
    """Könyv-docx eleje: címoldal (borítókép, cím, szerző), utána új oldalon a szöveg; A4-nél A4-es oldal.
    ink: a címoldal fekete betűs (a saját könyvsablonnál — a pandoc Title/Subtitle stílusa kék)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, RGBColor
    doc = Document(str(path))
    if size == "A4":
        for sec in doc.sections:
            sec.page_width, sec.page_height = Mm(210), Mm(297)
            sec.left_margin = sec.right_margin = Mm(25)
            sec.top_margin = sec.bottom_margin = Mm(22)
    sec = doc.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    styles = {s.name for s in doc.styles}
    new = []
    if cover and cover.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(cover), width=width)
        new.append(p)
    p = doc.add_paragraph(title, style="Title" if "Title" in styles else None); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    new.append(p)
    if author:
        p = doc.add_paragraph(author, style="Subtitle" if "Subtitle" in styles else None); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        new.append(p)
    if ink:
        for p in new:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    body = doc.element.body
    for k, p in enumerate(new):                     # a dokumentum ELEJÉRE
        body.remove(p._p)
        body.insert(k, p._p)
    # a szöveg új oldalon kezdődjön — de ha az első bekezdés stílusa eleve új oldalon kezd (fejezetcím), ne legyen üres oldal
    first = next((p for p in doc.paragraphs[len(new):] if p.text.strip()), None)
    if first is not None:
        st, pbb = first.style, first.paragraph_format.page_break_before
        while pbb is None and st is not None:
            pbb, st = st.paragraph_format.page_break_before, st.base_style
        if not pbb:
            first.paragraph_format.page_break_before = True
    doc.save(str(path))


@router.post("/api/make/book")
async def make_book(body: dict[str, Any]) -> dict[str, Any]:
    """Kézirat (DOCX) → KÖNYV a Word-szerkesztőben: könyvszedés (A5/A4, könyvbetű, címoldal borítóval) docx-ben,
    amit kiadás előtt át lehet nézni és szerkeszteni, + PDF mellé. layout="booklet": a régi dizájnos füzet-PDF."""
    from docx import Document
    from app.tools import design
    from app.tools.crafts import existing
    fid = body.get("file_id")
    if not fid:
        raise HTTPException(400, "Upload your manuscript (.docx) first.")
    if body.get("layout") != "booklet":
        import asyncio
        from app.storage import display_label, output_path, public_url, write_meta
        from app.storage.files import new_file_id
        from app.tools import digitize, doc_style, wordcraft
        if not str(fid).lower().endswith(".docx"):
            raise HTTPException(400, "The manuscript must be a .docx file.")
        title = (body.get("title") or "").strip() or Path(display_label(fid)).stem
        try:
            ref = (await doc_style.reference_for(body["style_file_id"]) if body.get("style_file_id")
                   else await asyncio.to_thread(digitize._book_reference))
            styled = new_file_id("docx")
            await asyncio.to_thread(doc_style.apply_style, existing(fid), ref, output_path(styled))
            cover_id = await _materialize(_image_spec(body.get("cover")))
            await asyncio.to_thread(_book_front, output_path(styled), title, (body.get("author") or "").strip() or None,
                                    existing(cover_id) if cover_id else None, body.get("size") or "A5",
                                    not body.get("style_file_id"))
            pdf = new_file_id("pdf")
            await wordcraft.convert_file(output_path(styled), output_path(pdf))
        except HTTPException:
            raise
        except Exception as e:
            raise _fail(e)
        group = {"group": _new_group()}             # a Word-könyv és a PDF-je testvérek: a nézetben egymásra váltanak
        write_meta(styled, source_file_id=fid, operation="book", label=f"{title} (book).docx", extra=group)
        write_meta(pdf, source_file_id=styled, operation="book", label=f"{title}.pdf", extra=group)
        events.emit("tool_done", tool="make_book", result={"file_id": styled})
        return {"file_id": styled, "url": public_url(styled), "pdf": pdf, "pdf_url": public_url(pdf),
                "size_kb": output_path(styled).stat().st_size // 1024}
    if body.get("style_file_id"):          # „olyan legyen, mint ez” — a kézirat a minta stílusaival, WordCraft-tal szedve
        import asyncio
        from pathlib import Path as _P
        from app.storage import display_label, output_path, public_url, write_meta
        from app.storage.files import new_file_id
        from app.tools import doc_style, wordcraft
        from app.tools.crafts import existing as _existing
        try:
            ref = await doc_style.reference_for(body["style_file_id"])
            styled = new_file_id("docx")
            await asyncio.to_thread(doc_style.apply_style, _existing(fid), ref, output_path(styled))
            write_meta(styled, source_file_id=fid, operation="restyle", label=f"{_P(display_label(fid)).stem} (styled).docx")
            pdf = new_file_id("pdf")
            await wordcraft.convert_file(output_path(styled), output_path(pdf))
        except HTTPException:
            raise
        except Exception as e:
            raise _fail(e)
        write_meta(pdf, source_file_id=styled, operation="book", label=f"{body.get('title') or _P(display_label(fid)).stem}.pdf")
        return {"file_id": pdf, "url": public_url(pdf), "size_kb": output_path(pdf).stat().st_size // 1024, "docx": styled}
    try:
        doc = Document(str(existing(fid)))
    except Exception:
        raise HTTPException(400, "The manuscript must be a .docx file.")
    sections: list[dict[str, Any]] = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if p.style.name.lower().startswith(("heading 1", "heading 2", "title")) or (not sections and len(txt) < 80):
            sections.append({"heading": txt, "text": ""})
        elif sections:
            sections[-1]["text"] += ("\n\n" if sections[-1]["text"] else "") + txt
    sections = [s for s in sections if s["text"]][:40]
    if not sections:
        raise HTTPException(400, "No chapters found — use Heading 1 for chapter titles.")
    data = {"title": body.get("title") or sections[0]["heading"], "subtitle": body.get("subtitle"),
            "author": body.get("author"), "organization": body.get("author"), "sections": sections,
            "cover_image": _image_spec(body.get("cover"))}
    try:
        out = await design.render("booklet", data, size=body.get("size") or "A5", export=["pdf"])
    except Exception as e:
        raise _fail(e)
    _label_outputs(out, f"{str(data['title'])[:60]} (booklet)")     # a szerkesztő címsorában ne „design.pdf” álljon
    events.emit("tool_done", tool="make_book", result={"file_id": out["file_id"]})
    return out


# PDF-műveletek (PdfCraft): action → (lépések, kimeneti formátum, mentsen-e a végén doc 1-et)
def _pdf_plan(action: str, body: dict[str, Any]) -> tuple[list[dict[str, Any]], str, bool]:
    if action == "merge":
        return [{"tool": "doc_combine", "args": {"paths": "$paths", "out": "$out"}}], "pdf", False
    if action == "compress":
        return [{"tool": "doc_reduce", "args": {"path": "$out"}}], "pdf", False
    if action == "to_word":
        return [{"tool": "doc_export_office", "args": {"path": "$out"}}], "docx", False
    if action == "page_numbers":
        pos = body.get("position") or "center"
        return [{"tool": "doc_header_footer", "args": {f"footer_{pos}": body.get("format") or "<<1 of n>>"}}], "pdf", True
    if action == "watermark":
        text = (body.get("text") or "").strip()
        if not text:
            raise HTTPException(400, "Type the watermark text.")
        return [{"tool": "doc_watermark", "args": {"text": text, "opacity": 0.25}}], "pdf", True
    if action == "redact":
        words = [w.strip() for w in body.get("words") or [] if w and w.strip()]
        patterns = [p for p in body.get("patterns") or [] if p in ("email", "phone", "credit-card", "ssn", "date")]
        if not words and not patterns:
            raise HTTPException(400, "Add words to black out, or pick a pattern (emails, phone numbers…).")
        steps = ([{"tool": "redact_mark", "args": {"words": words}}] if words else []) + \
                [{"tool": "redact_mark", "args": {"pattern": p}} for p in patterns]
        return steps + [{"tool": "redact_apply", "args": {}}, {"tool": "doc_remove_hidden", "args": {}}], "pdf", True
    if action == "ocr":
        return [{"tool": "ocr_recognize", "args": {"language": "en"}}], "pdf", True
    if action == "protect":
        pw = body.get("password") or ""
        if len(pw) < 4:
            raise HTTPException(400, "Choose a password of at least 4 characters.")
        return [{"tool": "doc_protect", "args": {"open_password": pw}}], "pdf", True
    if action == "compare":
        return [{"tool": "doc_compare_report", "args": {"doc": 2, "other": 1, "path": "$out"}}], "pdf", False
    if action == "make_fillable":   # digitális (nem fotózott) nyomtatvány: vonalak/dobozok → valódi mezők
        return [{"tool": "form_detect_fields", "args": {"add": True}}], "pdf", True
    raise HTTPException(400, f"unknown action: {action}")


PDF_LABELS = {"merge": "merged", "compress": "compressed", "to_word": "", "page_numbers": "numbered",
              "watermark": "watermarked", "redact": "redacted", "ocr": "searchable", "protect": "protected",
              "compare": "comparison", "make_fillable": "fillable"}


@router.post("/api/make/pdf")
async def make_pdf(body: dict[str, Any]) -> dict[str, Any]:
    from pathlib import Path
    from app.storage import display_label, output_path, public_url, write_meta
    from app.storage.files import new_file_id
    from app.tools import crafts
    action = body.get("action") or ""
    ids = [i for i in (body.get("file_ids") or ([body["file_id"]] if body.get("file_id") else [])) if i]
    if not ids:
        raise HTTPException(400, "Choose a PDF first.")
    need = 2 if action in ("merge", "compare") else 1
    if len(ids) < need:
        raise HTTPException(400, "Choose at least two PDFs." if action == "merge" else "Choose the old and the new version.")
    if any(not i.lower().endswith(".pdf") for i in ids):
        raise HTTPException(400, "Only PDF files here — convert other files to PDF first.")
    steps, fmt, save = _pdf_plan(action, body)
    fid = new_file_id(fmt)
    out = output_path(fid)
    try:
        if action == "merge":
            steps[0]["args"]["paths"] = [str(crafts.existing(i)) for i in ids]
            await crafts.pdf_script([], steps, out, save=False)
        elif action == "redact":
            # a nem illeszkedő szó/minta a PdfCraftban hiba („nothing matched”) → előbb egyenként kipróbáljuk
            marks = []
            for st in steps[:-2]:
                try:
                    await crafts.pdf_script(ids[:1], [st], out, save=False)
                    marks.append(st)
                except RuntimeError as e:
                    if "nothing matched" not in str(e):
                        raise
            if not marks:
                raise HTTPException(422, "Nothing to black out — none of those words or patterns appear in this PDF.")
            await crafts.pdf_script(ids[:1], marks + steps[-2:], out, save=True)
        else:
            await crafts.pdf_script(ids[:need], steps, out, save=save)
        if not out.exists():
            raise RuntimeError("no output")
    except HTTPException:
        raise
    except Exception as e:
        raise _fail(e)
    stem = Path(display_label(ids[-1] if action == "compare" else ids[0])).stem
    tag = PDF_LABELS.get(action, action)
    label = f"{stem}{' (' + tag + ')' if tag else ''}.{fmt}"
    write_meta(fid, source_file_id=ids[0], operation=f"pdf {action}", label=label)
    res = {"file_id": fid, "url": public_url(fid), "size_kb": out.stat().st_size // 1024, "label": label}
    events.emit("tool_done", tool="make_pdf", params={"action": action}, result=res)
    return res


# ---------- Convert (bármi → bármi) ----------

@router.post("/api/make/handout")
async def make_handout(body: dict[str, Any]) -> dict[str, Any]:
    """Pakli → Word-kiosztmány (diakép + szöveg + jegyzetek), AI nélkül."""
    from app.tools import deck_handout
    if not body.get("file_id"):
        raise HTTPException(400, "Choose a presentation.")
    try:
        out = await deck_handout.build(body["file_id"], slide_images=body.get("slide_images", True) is not False)
    except (ValueError, FileNotFoundError) as e:
        raise HTTPException(400, str(e))
    except Exception as e:
        raise _fail(e)
    events.emit("tool_done", tool="make_handout", result={"file_id": out["file_id"]})
    return out


@router.get("/api/convert/targets")
async def convert_targets(ext: str) -> dict[str, Any]:
    from app.tools import anyconvert
    return {"targets": [{"fmt": t, "label": anyconvert.LABELS.get(t, t.upper())} for t in anyconvert.targets_for(ext)]}


@router.post("/api/make/convert")
async def make_convert(request: Request, body: dict[str, Any]) -> dict[str, Any]:
    from app.tools import anyconvert
    fid, target = body.get("file_id"), (body.get("target") or "").strip()
    if not fid or not target:
        raise HTTPException(400, "Choose a file and a format.")
    try:
        out = await anyconvert.convert_any(getattr(request.app.state, "pool", None), fid, target)
    except (ValueError, FileNotFoundError) as e:
        raise HTTPException(400, str(e))
    except Exception as e:
        raise _fail(e)
    events.emit("tool_done", tool="make_convert", params={"target": target}, result={"file_id": out["file_id"]})
    return out


# ---------- Digitalizálás: régi könyv / fotózott oldal → szöveg, Word, új PDF, kereshető PDF ----------

@router.post("/api/make/digitize")
async def make_digitize(body: dict[str, Any]) -> dict[str, Any]:
    from app import jobs
    from app.tools import digitize
    ids = [i for i in body.get("file_ids") or [] if i]
    if not ids:
        raise HTTPException(400, "Add photos or scans of the pages.")
    outputs = [o for o in body.get("outputs") or ["docx"] if o in digitize.OUTPUTS]
    if not outputs:
        raise HTTPException(400, "Pick at least one result: Word, new PDF, searchable PDF or text.")
    lang = (body.get("language") or "").strip() or None
    style_id = body.get("style_file_id")
    limit = 1000 if tenancy.current_user_id() == tenancy.ADMIN_ID else settings.OCR_MAX_PAGES
    if body.get("max_pages"):            # „csak az első N oldal” (próba, mintafejezet)
        limit = max(1, min(limit, int(body["max_pages"])))

    async def on_pages(n: int, form_pages: int) -> None:     # a keret / kredit, amint kiderül, hány oldal
        from app import billing
        if form_pages:
            await billing.use("form_pages", form_pages)
        if n:
            await billing.scan_pages(n)

    async def work(progress):
        try:
            res = await digitize.digitize(ids, outputs, lang, limit, progress, on_pages=on_pages, style_file_id=style_id)
        except Exception as e:
            raise _fail(e)
        events.emit("tool_done", tool="make_digitize", result={"files": [f["file_id"] for f in res["files"]],
                                                               "pages": res["pages"]})
        return res
    return jobs.start("digitize", work)


# ---------- Űrlapok: mezők, kitöltés (egyenként vagy táblázatból) ----------

@router.get("/api/forms/fields")
async def form_fields(file_id: str) -> dict[str, Any]:
    from app.tools import forms
    if not file_id.lower().endswith(".pdf"):
        return {"fields": []}
    try:
        return {"fields": await forms.fields(file_id)}
    except FileNotFoundError:
        raise HTTPException(404, "File not found.")
    except Exception as e:
        log.warning("form_fields %s: %s", file_id, e)
        return {"fields": []}


@router.post("/api/make/fill_form")
async def make_fill_form(body: dict[str, Any]) -> dict[str, Any]:
    """{file_id, values: {mező: érték}} vagy {file_id, data_file_id: csv/xlsx} (soronként egy példány); flatten?"""
    from pathlib import Path
    from app.storage import display_label, output_path, public_url, write_meta
    from app.storage.files import new_file_id
    from app.tools import crafts, forms
    fid = body.get("file_id")
    if not fid:
        raise HTTPException(400, "Choose the form (a PDF with fields).")
    flatten = bool(body.get("flatten"))
    out_id = new_file_id("pdf")
    from app.storage.naming import result_stem
    try:
        if body.get("data_file_id"):
            rows = forms.read_rows(crafts.existing(body["data_file_id"]))
            info = await forms.fill_batch(fid, rows, flatten, output_path(out_id))
            label = result_stem(fid, f"{info['copies']} filled copies") + ".pdf"
        else:
            info = await forms.fill(fid, dict(body.get("values") or {}), flatten, output_path(out_id))
            label = result_stem(fid, "filled") + ".pdf"
    except (ValueError, FileNotFoundError) as e:
        raise HTTPException(400, str(e))
    except Exception as e:
        raise _fail(e)
    write_meta(out_id, source_file_id=fid, operation="fill form", label=label)
    res = {"file_id": out_id, "url": public_url(out_id), "label": label, "size_kb": output_path(out_id).stat().st_size // 1024, **info}
    events.emit("tool_done", tool="make_fill_form", result={"file_id": out_id})
    return res


# ---------- Szavas, de pontos fotószerkesztés (rétegzett PSD; tömegesen egy terv minden képre) ----------

@router.post("/api/make/photo_edit")
async def make_photo_edit(body: dict[str, Any]) -> dict[str, Any]:
    """{file_ids[], instruction, logo_file_id?, formats?: ["psd","jpg"], plan?: [lépések] (külső AI saját terve)}"""
    from app import jobs
    from app.tools import photo_edit
    ids = [i for i in body.get("file_ids") or ([body["file_id"]] if body.get("file_id") else []) if i]
    if not ids:
        raise HTTPException(400, "Add a photo.")
    instruction = (body.get("instruction") or "").strip()
    steps = body.get("plan")
    if not instruction and not steps:
        raise HTTPException(400, "Describe the edit, e.g. “crop to 4:5, warmer, headline ‘Summer sale’ at the top, logo bottom right”.")
    formats = [f for f in body.get("formats") or ["psd", "jpg"] if f in ("psd", "jpg", "png", "webp")] or ["psd", "jpg"]
    from app import billing
    plan = await billing.plan_of(tenancy.current_user_id())
    if len(ids) > billing.batch_limit(plan):
        raise HTTPException(402, f"Up to {billing.batch_limit(plan)} photos at once on the {billing.PLANS[plan]['name']} plan.")
    await billing.use("photo_edit", len(ids))

    async def work(progress):
        try:
            res = await photo_edit.run(ids, instruction, body.get("logo_file_id"), formats, steps=steps, progress=progress)
        except Exception as e:
            raise _fail(e)
        events.emit("tool_done", tool="make_photo_edit", result={"files": [f["file_id"] for f in res["files"]]})
        return res
    return jobs.start("photo_edit", work)


# ---------- Régi fotók helyreállítása (album): rétegzett PSD + JPG, opcionálisan színezve ----------

@router.post("/api/make/restore")
async def make_restore(body: dict[str, Any]) -> dict[str, Any]:
    """{file_ids[], colorize?: true|false|null (null = fekete-fehérnél igen), split_prints?: bool}"""
    from app import jobs
    from app.tools import restore
    ids = [i for i in body.get("file_ids") or [] if i]
    if not ids:
        raise HTTPException(400, "Add photos of the old pictures (or scans, or a whole album page).")
    colorize = body.get("colorize")
    split = bool(body.get("split_prints", True))
    left = await _credits_left()
    if left is not None and left < 2:
        raise HTTPException(402, f"Restoring needs credits (2 per photo, +1 to colourise); you have {left}. See Plans.")

    async def work(progress):
        try:
            res = await restore.run(ids[:50], colorize, split, progress)
        except Exception as e:
            raise _fail(e)
        events.emit("tool_done", tool="make_restore", result={"files": [f["file_id"] for f in res["files"]]})
        return res
    return jobs.start("restore", work)


# ---------- „Kérj változtatást”: iteratív szerkesztés, minden kérés új verzió ----------

@router.post("/api/make/revise")
async def make_revise(request: Request, body: dict[str, Any]) -> dict[str, Any]:
    """{file_id, instruction, style_file_id?} → job; az eredmény új verzió (meta: instruction, summary)."""
    from app import jobs
    from app.tools import revise
    fid = body.get("file_id")
    if not fid:
        raise HTTPException(400, "Choose a file.")
    instruction = (body.get("instruction") or "").strip()
    if not instruction and not body.get("style_file_id"):
        raise HTTPException(400, "Say what should change.")
    pool = getattr(request.app.state, "pool", None)
    from app import billing
    await billing.use("revise")

    async def work(progress):
        progress(0, 1, "working")
        try:
            res = await revise.revise(pool, fid, instruction, body.get("style_file_id"))
        except Exception as e:
            raise _fail(e)
        if str(body.get("branch") or "").strip():          # új ág: a változat ág-nevet kap (a fa és az összefésülés ezt mutatja)
            from app.tools import branches
            branches.set_branch(res["file_id"], str(body["branch"]))
            res["branch"] = str(body["branch"]).strip()[:60]
        events.emit("tool_done", tool="make_revise", result={"file_id": res["file_id"]})
        return res
    return jobs.start("revise", work)


@router.get("/api/revisions")
async def revisions(file_id: str) -> dict[str, Any]:
    from app.tools import revise
    return {"history": revise.history(file_id)}


# ---------- Dokumentum (Word-szerű): sablonból, vagy pár mondatból (AI) ----------

_DOC_SYSTEM = """You write documents. Reply with the document only, in Markdown: "# " title, "## " headings, normal
paragraphs, "- " lists, Markdown tables where useful. No preamble, no closing remarks, no code fences.
Write in the language of the brief. For a letter: sender block, date, recipient, salutation, body, closing, name.
Use only the facts in the brief — never invent specifics (what work is done, prices, names, dates, places). Where the
brief is silent, stay general or leave a [placeholder] the writer fills in."""
DOC_KINDS = {"letter": "a letter", "report": "a report", "article": "an article", "cv": "a CV / résumé", "memo": "a memo",
             "proposal": "a proposal", "other": "a document"}
_DOC_REF: Optional[Path] = None


def _doc_reference() -> Path:
    """Pandoc referencia-docx általános dokumentumhoz: A4, Carlito (Calibri-metrika) 11 pt, sötét címsorok —
    a pandoc alapja kék, Letter méretű."""
    global _DOC_REF
    if _DOC_REF and _DOC_REF.exists():
        return _DOC_REF
    import subprocess
    import tempfile
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    path = Path(tempfile.gettempdir()) / "aice_doc_reference.docx"
    subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"], capture_output=True, check=True)
    doc = Document(str(path))
    for sec in doc.sections:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(25)
        sec.top_margin, sec.bottom_margin = Mm(22), Mm(22)
    st = {x.name: x for x in doc.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in st:
            st[name].font.name, st[name].font.size = "Carlito", Pt(11)
            st[name].font.color.rgb = RGBColor(0x1E, 0x24, 0x30)
            st[name].paragraph_format.space_after = Pt(6)
    for lvl, size in (("Title", 24), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        if lvl in st:
            st[lvl].font.name, st[lvl].font.size, st[lvl].font.bold = "Caladea", Pt(size), True
            st[lvl].font.color.rgb = RGBColor(0x1E, 0x44, 0x36)
    for el in doc.styles.element.iter(qn("w:rFonts")):
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            el.attrib.pop(qn(a), None)
    doc.save(str(path))
    _DOC_REF = path
    return path


@router.post("/api/make/document")
async def make_document(body: dict[str, Any]) -> dict[str, Any]:
    """{mode: "template", template: blank|letter|report|resume|sample} vagy {mode: "brief", brief, kind?}."""
    import asyncio
    import pypandoc
    from app.storage import output_path, public_url, temp_path, write_meta
    from app.storage.files import new_file_id
    from app.tools import wordcraft
    mode = body.get("mode") or "template"
    if mode == "template":
        tpl = body.get("template") or "blank"
        if tpl not in wordcraft.TEMPLATES:
            raise HTTPException(400, f"Unknown template: {tpl}")
        try:
            out = await wordcraft.run(template=tpl)
            if body.get("style_file_id"):
                from app.tools import doc_style
                ref = await doc_style.reference_for(body["style_file_id"])
                await asyncio.to_thread(doc_style.apply_style, output_path(out["file_id"]), ref, output_path(out["file_id"]))
        except Exception as e:
            raise _fail(e)
        names = {"blank": "Untitled", "letter": "Letter", "report": "Report", "resume": "CV", "sample": "Sample"}
        write_meta(out["file_id"], operation="new document", label=f"{names.get(tpl, tpl)}.docx")
        return out
    if tenancy.is_guest():
        from app.auth import ACCOUNT_NEEDED
        raise HTTPException(401, ACCOUNT_NEEDED)
    brief = (body.get("brief") or "").strip()
    if len(brief) < 10:
        raise HTTPException(400, "Describe the document in a sentence or two.")
    from app import billing
    await billing.use("document")
    kind = DOC_KINDS.get(body.get("kind") or "other", "a document")
    from app.llm.siliconflow import _call_model
    try:
        md = ""
        for model in (settings.AGENT_MODEL, settings.AGENT_MODEL_FALLBACK):
            try:
                md = await _call_model(model, [{"role": "system", "content": _DOC_SYSTEM},
                                               {"role": "user", "content": f"Write {kind}.\n\nBrief:\n{brief}"}], thinking=False,
                                         json_mode=False)
                if md.strip():
                    break
            except Exception as e:
                log.warning("dokumentum-írás (%s): %s", model, e)
        md = md.strip().removeprefix("```markdown").removeprefix("```").removesuffix("```").strip()
        if not md:
            raise RuntimeError("The assistant returned nothing — try again.")
        title = next((l[2:].strip() for l in md.splitlines() if l.startswith("# ")), kind.split(" ", 1)[1].capitalize())
        src = temp_path(new_file_id("md"))
        src.write_text(md, "utf-8")
        fid = new_file_id("docx")
        if body.get("style_file_id"):
            from app.tools import doc_style
            ref = await doc_style.reference_for(body["style_file_id"])
        else:
            ref = await asyncio.to_thread(_doc_reference)
        await asyncio.to_thread(pypandoc.convert_file, str(src), "docx", outputfile=str(output_path(fid)),
                                extra_args=["--wrap=none", f"--reference-doc={ref}"])
        src.unlink(missing_ok=True)
    except HTTPException:
        raise
    except Exception as e:
        raise _fail(e)
    write_meta(fid, operation="new document (AI)", label=f"{title[:60]}.docx", extra={"brief": brief[:500]})
    events.emit("tool_done", tool="make_document", result={"file_id": fid})
    return {"file_id": fid, "url": public_url(fid), "size_kb": output_path(fid).stat().st_size // 1024}
