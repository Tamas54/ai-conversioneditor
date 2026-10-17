"""Eredmény-ellenőrzés: a TECHNIKAI siker (lett fájl) és a FELADAT teljesülése (az lett-e, amit kértek) külön.

A ChatGPT-teszt (10-12) kritikája: a háttércsere „sikeres" volt, közben eltűnt a gitárnyak — az eszköz nem
mondta, hogy nézd meg. Ezért minden MCP-fájleredmény kap egy `check` mezőt, KÖZÖS SZERZŐDÉS a sávok között:

    {"status": "ok" | "needs_check" | "failed", "findings": [str], "verified": [str]}

  ok           — egy ellenőrző TÉNYLEG igazolt valamit a kérésből (pl. a kitakarandó e-mailek 6 → 0 a szövegben)
  needs_check  — nincs gépi igazolás a célra, vagy gyanús jel van: nézd meg a preview_url-t
  failed       — az ellenőrző hibát talált (maradt kitakarandó adat, nem változott semmi, üres a kimenet)

Ha az eszköz maga ad `check`-et (pl. a kivágás a megőrzendő alakokról), az marad — itt csak alakot igazítunk.
Az ellenőrzők gyorsak (fitz, python-pptx, python-docx, Pillow), és SOSEM dobnak: hiba esetén nincs check.
"""
from __future__ import annotations

import asyncio
import logging
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Any, Optional

log = logging.getLogger("aice.checks")

STATUSES = ("ok", "needs_check", "failed")
LOOK = "Look at preview_url before telling the person it's done."


def make(status: str, findings: Optional[list[str]] = None, verified: Optional[list[str]] = None) -> dict[str, Any]:
    return {"status": status if status in STATUSES else "needs_check",
            "findings": [f for f in (findings or []) if f], "verified": [v for v in (verified or []) if v]}


def normalize(check: Any) -> Optional[dict[str, Any]]:
    """Egy eszköz saját check-je → a közös alak (ismeretlen státusz → needs_check)."""
    if not isinstance(check, dict):
        return None
    f, v = check.get("findings") or [], check.get("verified") or []
    return make(str(check.get("status") or "needs_check"),
                [str(x) for x in (f if isinstance(f, list) else [f])], [str(x) for x in (v if isinstance(v, list) else [v])])


# ---------- segédek ----------

def _path(file_id: Optional[str]) -> Optional[Path]:
    if not file_id or not isinstance(file_id, str):
        return None
    from app.tools.crafts import existing
    try:
        return existing(file_id)
    except (FileNotFoundError, ValueError):
        return None


def _ext(p: Optional[Path]) -> str:
    return p.suffix.lstrip(".").lower() if p else ""


def _fold(s: str) -> str:
    """Kisbetű, ékezet nélkül, összevont szóközök — a PDF-szövegkinyerés apró eltéréseihez."""
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def pdf_pages_text(p: Path) -> list[str]:
    import fitz
    with fitz.open(str(p)) as d:
        return [pg.get_text() for pg in d]


def _plural(n: int, word: str) -> str:
    return f"{n} {word}{'' if n == 1 else 's'}"


# ---------- minták (a PdfCraft redact_mark pattern-jei) ----------
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_PHONE = re.compile(r"(?<![\w@./-])\+?\(?\d[\d\s().\-/]{5,}\d(?![\w@])")
_DATEISH = re.compile(r"\s*(\d{4}[-./]\d{1,2}[-./]\d{1,2}|\d{1,2}[-./]\d{1,2}[-./]\d{2,4})\s*")
_SSN = re.compile(r"\b\d{3}-\d{2}-\d{4}\b")
_CARD = re.compile(r"\b(?:\d[ -]?){13,19}\b")
_DATE = re.compile(r"\b(\d{4}[-./]\d{1,2}[-./]\d{1,2}|\d{1,2}[-./]\d{1,2}[-./]\d{2,4})\b")


def _luhn(digits: str) -> bool:
    s, alt = 0, False
    for d in reversed(digits):
        n = int(d) * (2 if alt else 1)
        s += n - 9 if n > 9 else n
        alt = not alt
    return s % 10 == 0


def find_pattern(kind: str, text: str) -> list[str]:
    kind = (kind or "").lower().replace("_", "-")
    if kind == "email":
        return _EMAIL.findall(text)
    if kind == "phone":
        out = []
        for m in _PHONE.findall(text):
            digits = re.sub(r"\D", "", m)
            if 7 <= len(digits) <= 15 and not _DATEISH.fullmatch(m) and not _SSN.fullmatch(m.strip()):
                out.append(m.strip())
        return out
    if kind == "ssn":
        return _SSN.findall(text)
    if kind in ("credit-card", "card", "creditcard"):
        return [m for m in _CARD.findall(text) if _luhn(re.sub(r"\D", "", m))]
    if kind == "date":
        return [m if isinstance(m, str) else m[0] for m in _DATE.findall(text)]
    return []


# ---------- PDF: kitakarás ----------

def check_redaction(src: Optional[Path], out: Path, steps: list[dict[str, Any]]) -> Optional[dict[str, Any]]:
    """A kitakarandó minták/szavak a KIMENET szövegrétegében már nem szerepelhetnek (nem elég a fekete doboz)."""
    marks = [s.get("args") or {} for s in steps if s.get("tool") == "redact_mark"]
    if not marks:
        return None
    after_pages = pdf_pages_text(out)
    after = "\n".join(after_pages)
    before = "\n".join(pdf_pages_text(src)) if src and _ext(src) == "pdf" else ""
    findings, verified, unverifiable = [], [], []
    all_before: list[str] = []
    for m in marks:
        targets: list[tuple[str, list[str], list[str]]] = []      # (felirat, előtte, utána)
        if m.get("pattern"):
            kind = str(m["pattern"])
            targets.append((f"{kind} pattern", find_pattern(kind, before), find_pattern(kind, after)))
        words = ([m["find"]] if isinstance(m.get("find"), str) else []) + [w for w in (m.get("words") or []) if isinstance(w, str)]
        for w in words:
            fw = _fold(w)
            if not fw:
                continue
            targets.append((f"“{w}”", [w] * _fold(before).count(fw), [w] * _fold(after).count(fw)))
        if m.get("rect"):
            import fitz
            page = int(m.get("page") or 1)
            with fitz.open(str(out)) as d:
                left = d[page - 1].get_textbox(fitz.Rect(*m["rect"])).strip() if 0 < page <= d.page_count else ""
            if left:
                findings.append(f"Text is still readable inside the redaction box on page {page}: “{left[:60]}”.")
            else:
                verified.append(f"No text left inside the redaction box on page {page}.")
        if m.get("whole_pages"):
            pages = m.get("pages") or list(range(1, len(after_pages) + 1))
            left_pages = [p for p in pages if isinstance(p, int) and 0 < p <= len(after_pages) and after_pages[p - 1].strip()]
            (findings.append(f"Pages {left_pages} still have text after whole-page redaction.") if left_pages
             else verified.append("Whole-page redactions leave no text."))
        for label, b, a in targets:
            all_before += b
            if a:
                sample = ", ".join(sorted(set(a))[:6])
                findings.append(f"{label}: {len(a)} still in the text after redaction ({sample}) — redact them too "
                                f"(e.g. redact_mark with find).")
            elif b:
                verified.append(f"{label}: {len(b)} in the original → 0 left in the text layer.")
            elif src:
                unverifiable.append(f"{label}: nothing like it was found in the original text (scanned page or already absent?).")
    if before.strip() and not after.strip():
        findings.append("The result has no text left at all — check that only the intended parts were blacked out.")
    if findings:
        return make("failed", findings + unverifiable, verified)
    # a doboz kicsit rövidebb a találatnál → a széléről 1–3 karakter megmarad (mért: „+36 30 123 4567" → „7")
    removed = [m for m in all_before if m]
    src_tokens = set(before.split())
    frags = sorted({t for t in after.split() if len(t) <= 3 and t not in src_tokens
                    and any(m.endswith(t) or m.startswith(t) for m in removed)})
    if frags:
        return make("needs_check", [f"Bits of the redacted text are still there: {', '.join(repr(f) for f in frags[:10])} — the "
                                    "black boxes are a little short. Look at preview_url and redact those spots too "
                                    "(redact_mark with rect a few points wider, then redact_apply)."] + unverifiable, verified)
    if verified:
        return make("ok", unverifiable, verified + ["Checked the extracted text of the result, not just the black boxes."])
    return make("needs_check", unverifiable or ["Nothing textual to verify (image-only PDF?). " + LOOK], verified)


# ---------- PDF: űrlap ----------

def _widgets(p: Path) -> list[tuple[str, str, str]]:
    import fitz
    out = []
    with fitz.open(str(p)) as d:
        for pg in d:
            for w in pg.widgets() or []:
                out.append((w.field_name or "", w.field_type_string or "", str(w.field_value or "").strip()))
    return out


def _filled(value: str) -> bool:
    return value not in ("", "Off", "No", "False", "false", "0")


def check_form_fill(out: Path, res: dict[str, Any], values: Optional[dict] = None, data_rows: Optional[list[dict]] = None,
                    flatten: bool = False) -> dict[str, Any]:
    if flatten:
        text = _fold("\n".join(pdf_pages_text(out)))
        expected = [str(v) for row in (data_rows or [values or {}]) for v in row.values() if isinstance(v, (str, int, float)) and str(v).strip()]
        missing = [v for v in expected if _fold(v) not in text]
        if missing:
            return make("needs_check", [f"{_plural(len(missing), 'value')} not found in the flattened text: {', '.join(missing[:8])}"],
                        [f"{len(expected) - len(missing)} of {len(expected)} values are printed on the page."])
        return make("ok", [], [f"All {len(expected)} values are printed on the page (flattened)."])
    ws = _widgets(out)
    if not ws:
        return make("failed", ["The result has no form fields left."], [])
    filled = [w for w in ws if _filled(w[2])]
    findings, verified = [], [f"{len(filled)} of {len(ws)} fields have a value."]
    vals = {_fold(w[2]) for w in filled}
    rows = data_rows if data_rows is not None else [values or {}]
    skipped = {str(x) for x in res.get("skipped") or []}
    expected = [(k, str(v)) for row in rows for k, v in row.items()
                if v not in (None, "", False) and not isinstance(v, bool) and str(k) not in skipped]
    missing = [(k, v) for k, v in expected if _fold(v) not in vals]
    if data_rows is not None:
        copies = res.get("copies")
        if copies is not None and copies != len(data_rows):
            findings.append(f"{copies} copies made for {len(data_rows)} rows.")
        else:
            verified.append(f"{len(data_rows)} rows → {len(data_rows)} filled copies.")
    if res.get("skipped"):
        findings.append(f"No field matched: {', '.join(map(str, res['skipped'][:10]))} — list the fields (pdf_form with only file_id).")
    if missing:
        findings.append(f"{_plural(len(missing), 'value')} not found in any field: "
                        + ", ".join(f"{k}={v}" for k, v in missing[:8]))
    elif expected:
        verified.append(f"Every given value ({len(expected)}) is in a field.")
    if not filled:
        return make("failed", findings or ["No field got a value."], verified)
    return make("needs_check" if findings else "ok", findings, verified)


def check_fillable(out: Path) -> dict[str, Any]:
    ws = _widgets(out)
    if not ws:
        return make("failed", ["The PDF has no form fields — it is not fillable."], [])
    kinds: dict[str, int] = {}
    for _, t, _v in ws:
        kinds[t or "field"] = kinds.get(t or "field", 0) + 1
    return make("ok", [], [f"{len(ws)} fillable fields (" + ", ".join(f"{n} {k.lower()}" for k, n in sorted(kinds.items())) + ")."])


# ---------- PPTX ----------

def _pptx_slide_sigs(p: Path) -> Optional[list[str]]:
    import hashlib
    from app.tools import content_digest
    try:
        data = content_digest._pptx(p)
    except Exception:
        return None
    return [hashlib.sha1(repr(s).encode("utf-8")).hexdigest() for s in data[1]]


def _emu_pt(v) -> float:
    return (v or 0) / 12700


def pptx_problems(p: Path) -> tuple[list[str], dict[str, Any]]:
    """Gyors, gépi hibakeresés egy paklin: dián kívüli elem, (becsült) szöveg-túlcsordulás, üres helyőrző."""
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    prs = Presentation(str(p))
    sw, sh = _emu_pt(prs.slide_width), _emu_pt(prs.slide_height)
    out, stats = [], {"slides": len(prs.slides), "pictures": 0}
    for i, s in enumerate(prs.slides, 1):
        for shp in s.shapes:
            name = shp.name or "shape"
            if getattr(shp, "image", None) is not None:
                stats["pictures"] += 1
            if shp.left is not None and shp.width is not None:
                l, t, w, h = _emu_pt(shp.left), _emu_pt(shp.top), _emu_pt(shp.width), _emu_pt(shp.height)
                if l < -0.02 * sw or t < -0.02 * sh or l + w > 1.02 * sw or t + h > 1.02 * sh:
                    out.append(f"Slide {i}: “{name}” reaches outside the slide.")
            if shp.is_placeholder and shp.has_text_frame and not shp.text_frame.text.strip() \
                    and getattr(shp, "image", None) is None:
                out.append(f"Slide {i}: empty placeholder “{name}” (shows “Click to add text”).")
            if shp.has_text_frame and shp.text_frame.text.strip() and shp.width and shp.height:
                tf = shp.text_frame
                if tf.auto_size in (MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE, MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT):
                    continue
                bw = _emu_pt(shp.width) - 14.4
                bh = _emu_pt(shp.height) - 7.2
                if bw <= 20 or bh <= 10:
                    continue
                need = 0.0
                for para in tf.paragraphs:
                    size = next((r.font.size.pt for r in para.runs if r.font.size), None) or (40 if "itle" in name else 20)
                    chars = len(para.text) or 1
                    lines = max(1, -(-int(chars * size * 0.5) // int(bw)))
                    need += lines * size * 1.2
                if need > bh * 1.35:
                    out.append(f"Slide {i}: text in “{name}” likely overflows its box (needs ~{int(need)} pt, box {int(bh)} pt).")
    return out, stats


async def check_deck(out: Path, src: Optional[Path], tool: str, args: dict[str, Any]) -> tuple[dict[str, Any], Optional[int]]:
    """A pakli ellenőrzése a készítés/szerkesztés végén: UGYANAZ a renderelt elemző, mint a validate_document-é
    (a szerkesztő és a PowerPoint/LibreOffice tördelése) — a ChatGPT-teszt 10-17: a készítő „rendben”-t mondott,
    a külön validátor túllógó szöveget talált. A két út ítélete így nem térhet el; a PDF-ek gyorsítótárban
    maradnak, a rákövetkező validate_document nem rendereli újra."""
    from app.tools import validate
    rendered, note = None, None
    try:
        # a teljes ellenőrzés kerete 45 s (run) — a renderelés legfeljebb 30 s, utána becslés, de legyen check
        layouts, notes = await asyncio.wait_for(validate._layouts(out, "deck", want_file=True), timeout=30)
        try:
            if layouts:
                rendered = await asyncio.to_thread(validate.deck_issues, out, layouts)
                note = f"Rendered and checked like validate_document: " + "; ".join(
                    f"{validate.view_name(L)} — {_plural(L.pages, 'slide')}" for L in layouts) + "."
            notes and log.info("pakli-ellenőrzés: %s", notes)
        finally:
            for L in layouts:
                L.close()
    except Exception as e:
        log.info("renderelt pakli-ellenőrzés kimaradt: %s", e)
        why = "it took too long" if isinstance(e, asyncio.TimeoutError) else str(e)[:100]
        note = (f"The rendered check could not run ({why}) — text overflow was only estimated; "
                "validate_document(file_id) renders and checks every slide.")
    return await asyncio.to_thread(check_pptx, out, src, tool, args, rendered, note)


def _deck_finding(it: dict[str, Any]) -> str:
    fix = f" Fix: {it['fix']}" if it.get("fix") else ""
    return f"{it['id']} ({it['severity']}) — {it['message']}{fix}"


def check_pptx(out: Path, src: Optional[Path], tool: str, args: dict[str, Any], rendered: Optional[list[dict[str, Any]]] = None,
               note: Optional[str] = None) -> tuple[dict[str, Any], Optional[int]]:
    from app.tools.validate import ERROR
    if rendered is not None:
        _, stats = pptx_problems(out)
        problems = [_deck_finding(i) for i in rendered]
        errors = any(i["severity"] == ERROR for i in rendered)
    else:
        problems, stats = pptx_problems(out)
        errors = False
    verified, findings, first_changed = [f"{_plural(stats['slides'], 'slide')}, {_plural(stats['pictures'], 'picture')}."], [], None
    if note:
        verified.append(note)
    if stats["slides"] == 0:
        return make("failed", ["The deck has no slides."], []), None
    if src is not None and _ext(src) == "pptx":
        a, b = _pptx_slide_sigs(src), _pptx_slide_sigs(out)
        if a is not None and b is not None:
            import difflib
            new_s, changed_s, removed = [], [], 0
            for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
                if op == "insert":
                    new_s += range(j1 + 1, j2 + 1)
                elif op == "delete":
                    removed += i2 - i1
                elif op == "replace":
                    changed_s += range(j1 + 1, j1 + 1 + min(i2 - i1, j2 - j1))
                    new_s += range(j1 + 1 + min(i2 - i1, j2 - j1), j2 + 1)
                    removed += max(0, (i2 - i1) - (j2 - j1))
            if not (new_s or changed_s or removed):
                return make("failed", ["Nothing changed in the deck — the commands had no visible effect."], verified), None
            first_changed = (sorted(new_s + changed_s) or [None])[0]
            parts = ([f"new slide{'s' if len(new_s) > 1 else ''} {', '.join(map(str, new_s[:12]))}"] if new_s else []) + \
                    ([f"changed slide{'s' if len(changed_s) > 1 else ''} {', '.join(map(str, changed_s[:12]))}"] if changed_s else []) + \
                    ([f"{_plural(removed, 'slide')} removed"] if removed else [])
            verified.append("Compared with the original: " + "; ".join(parts) + ".")
    if tool == "find_replace":
        pat = str(args.get("pattern") or "")
        texts = _fold(" ".join(_pptx_texts(out)))
        if pat and not args.get("regex") and _fold(pat) in texts:
            findings.append(f"“{pat}” still appears in the deck.")
        elif pat and not args.get("regex"):
            verified.append(f"“{pat}” no longer appears in the deck.")
    findings += problems[:12] + ([f"…and {len(problems) - 12} more — validate_document(file_id) lists them all."]
                                 if len(problems) > 12 else [])
    if findings:
        # ugyanaz az ítélet, mint a validate_document-é: hiba → failed (javítani kell), figyelmeztetés → needs_check
        return make("failed" if errors else "needs_check", findings + [LOOK], verified), first_changed
    if src is None or first_changed:
        return make("ok", [], verified + ["No slide has text overflowing, text drawn over other text, anything off the slide, or empty placeholders."]), first_changed
    return make("needs_check", [LOOK], verified), first_changed


def _pptx_texts(p: Path) -> list[str]:
    from pptx import Presentation
    return [shp.text_frame.text for s in Presentation(str(p)).slides for shp in s.shapes if shp.has_text_frame]


# ---------- DOCX ----------

def _docx_review_counts(p: Path) -> tuple[int, int, list[str]]:
    """Követett változások (w:ins + w:del) és megjegyzések száma, a szerzők."""
    with zipfile.ZipFile(p) as z:
        doc = z.read("word/document.xml").decode("utf-8", "replace")
        com = z.read("word/comments.xml").decode("utf-8", "replace") if "word/comments.xml" in z.namelist() else ""
    changes = len(re.findall(r"<w:(?:ins|del)\b", doc))
    comments = len(re.findall(r"<w:comment\b", com))
    authors = sorted(set(re.findall(r'w:author="([^"]*)"', doc + com)))
    return changes, comments, authors


def _docx_text(p: Path) -> str:
    from docx import Document
    return "\n".join(par.text for par in Document(str(p)).paragraphs)


def check_docx(out: Path, src: Optional[Path], tool: str, args: dict[str, Any]) -> dict[str, Any]:
    from app.tools import content_digest
    findings, verified = [], []
    text = _docx_text(out)
    verified.append(f"{_plural(len(text.split()), 'word')}.")
    if src is not None and _ext(src) == "docx":
        if content_digest.same_content(src, out):
            return make("failed", ["Nothing changed in the document — the edit had no visible effect."], verified)
        verified.append("The content changed.")
    changes, comments, authors = _docx_review_counts(out)
    cmds = " ".join(str((c or {}).get("command") if isinstance(c, dict) else c) for c in (args.get("commands") or []))
    if changes:
        verified.append(f"{_plural(changes, 'tracked change')}" + (f" by {', '.join(authors)}" if authors else "") + ".")
    elif "trackChanges" in cmds:
        findings.append("Track changes was asked for, but the result has no tracked changes.")
    if comments:
        verified.append(f"{_plural(comments, 'comment')}.")
    elif "comment" in cmds.lower():
        findings.append("A comment was asked for, but the result has none.")
    if args.get("author") and (changes or comments) and args["author"] not in authors:
        findings.append(f"The changes are not by “{args['author']}” (authors: {', '.join(authors) or 'none'}).")
    if tool == "find_replace" and not args.get("regex"):
        pat = str(args.get("pattern") or "")
        if pat and pat in text:
            findings.append(f"“{pat}” still appears in the document.")
        elif pat:
            verified.append(f"“{pat}” no longer appears.")
    if tool == "replace_heading" and args.get("new_title"):
        (verified if args["new_title"] in text else findings).append(
            f"Heading “{args['new_title']}” is there." if args["new_title"] in text else f"“{args['new_title']}” is not in the document.")
    if findings:
        return make("needs_check", findings, verified)
    specific = len(verified) > (2 if src is not None else 1)
    return make("ok" if specific else "needs_check", [] if specific else [LOOK], verified)


async def with_layout(check: dict[str, Any], out: Path) -> dict[str, Any]:
    """A mentett Word-fájl tördelése is (a szerkesztő motorján, kép nélkül): a sikeres mentés nem igazolja a jó
    tördelést (ChatGPT-teszt 10-13). Hiba vagy figyelmeztetés → needs_check, a validate_document mutatja meg képpel."""
    from app.tools import validate
    try:
        lay = await asyncio.wait_for(validate.quick_layout(out), timeout=25)
    except Exception as e:
        log.info("tördelés-ellenőrzés kimaradt: %s", e)
        lay = None
    if not lay:
        return make(check["status"], check["findings"] + ["Layout not checked — call validate_document."], check["verified"])
    found = lay["errors"] + lay["warnings"]
    if not found:
        return make(check["status"], check["findings"], check["verified"] + [
            f"Layout checked in the editor's engine ({_plural(lay['pages'], 'page')}): no empty list items, broken numbering, "
            "split table rows, headings left at a page bottom or overflow. validate_document also checks Word's layout "
            "and shows the pages."])
    lines = [f"Page {i['page']}: {i['message']}" if i.get("page") else i["message"] for i in found[:6]]
    more = f" (+{len(found) - 6} more)" if len(found) > 6 else ""
    return make("failed" if check["status"] == "failed" else "needs_check",
                check["findings"] + [f"Layout: {validate.counts(found)}{more} — call validate_document(file_id) to see them "
                                     "marked on the pages, fix the errors, then validate again."] + lines, check["verified"])


# ---------- képek ----------

def image_change(src: Path, out: Path) -> Optional[float]:
    """A képpontok hány %-a változott érezhetően (kicsinyítve összevetve); None, ha nem összevethető."""
    from PIL import Image, ImageChops, ImageOps
    try:
        with Image.open(src) as a, Image.open(out) as b:
            a = ImageOps.exif_transpose(a).convert("RGB").resize((256, 256))
            b = ImageOps.exif_transpose(b).convert("RGB").resize((256, 256))
            diff = ImageChops.difference(a, b).convert("L")
            return round(100 * sum(1 for v in diff.getdata() if v > 24) / (256 * 256), 1)
    except Exception:
        return None


def check_photo(out: Path, src: Optional[Path], res: dict[str, Any]) -> dict[str, Any]:
    verified, findings = [], []
    steps = res.get("steps") or []
    if steps:
        verified.append(f"{_plural(len(steps), 'command')} ran.")
    if src is not None and _ext(src) in ("png", "jpg", "jpeg", "webp", "tif", "tiff", "bmp", "gif") \
            and _ext(out) in ("png", "jpg", "jpeg", "webp", "tif", "tiff", "bmp", "gif"):
        pct = image_change(src, out)
        if pct is not None:
            if pct < 0.1 and not res.get("removed_background"):
                return make("failed", ["The picture looks the same as before — the edit had no visible effect."], verified)
            verified.append(f"{pct}% of the picture changed visibly.")
    ba = res.get("before_after") or {}
    if ba.get("before") and ba.get("after"):
        verified.append(f"Layers {ba['before'].get('layers', '?')} → {ba['after'].get('layers', '?')}.")
    return make("needs_check", findings + [LOOK], verified)


# ---------- restore (album) ----------

def check_restore(res: dict[str, Any]) -> dict[str, Any]:
    from app.tools.restore import trim_border
    files = res.get("files") or []
    n = res.get("photos") or 0
    if not files or not n:
        return make("failed", ["No photo came out."], [])
    verified = [f"{_plural(n, 'photo')} found and restored; each has a layered PSD (original, restored"
                + (", colour" if any(f.get("colorized") for f in files) else "") + ") and a JPG."]
    findings = []
    import cv2
    for f in files:
        if not str(f.get("file_id", "")).endswith(".jpg"):
            continue
        p = _path(f["file_id"])
        img = cv2.imread(str(p)) if p else None
        if img is None:
            continue
        t = trim_border(img)
        lost = 1 - (t.shape[0] * t.shape[1]) / (img.shape[0] * img.shape[1])
        if lost > 0.03:
            findings.append(f"“{f.get('label', f['file_id'])}” may still have a paper border ({round(lost * 100)}% of the picture).")
    if findings:
        return make("needs_check", findings + [LOOK], verified)
    verified.append("No paper border left on the restored pictures.")
    return make("needs_check", ["Faces are partly rebuilt by AI on blurry originals — " + LOOK], verified)


# ---------- egyéb ----------

async def _pages(p: Path) -> Optional[int]:
    from app.tools import previews
    try:
        return await previews.page_count(p)
    except Exception:
        return None


async def check_convert(out: Path, src: Optional[Path]) -> dict[str, Any]:
    n = await _pages(out)
    if not n:
        return make("failed", ["The converted file is empty or can't be opened."], [])
    verified = [f"The {_ext(out).upper()} opens: {_plural(n, 'page' if _ext(out) not in ('pptx',) else 'slide')}."]
    if _ext(out) == "pdf" and src is not None and _ext(src) in ("docx", "pptx", "md", "txt", "html", "odt", "rtf"):
        text = await asyncio.to_thread(lambda: "".join(pdf_pages_text(out)))
        if not text.strip():
            return make("needs_check", ["The PDF has no text layer — fonts may not have embedded. " + LOOK], verified)
        verified.append("It has a text layer.")
    return make("ok", [], verified)


async def check_pages_math(out: Path, expected: int, what: str) -> dict[str, Any]:
    n = await _pages(out)
    if n is None:
        return make("needs_check", ["Couldn't count the pages. " + LOOK], [])
    if n != expected:
        return make("failed", [f"{what}: expected {expected} pages, the result has {n}."], [])
    return make("ok", [], [f"{what}: {_plural(n, 'page')}, as expected."])


async def generic(out: Path) -> dict[str, Any]:
    n = await _pages(out)
    v = [f"It opens: {_plural(n, 'page')}."] if n else []
    return make("needs_check", ["No automatic check for this kind of change. " + LOOK], v)


_LENGTH_BY = {"narration": "narration mode: the voice sets the length", "fixed": "fixed mode",
              "footage": "the footage sets the length"}


def check_video(p: Path, want_s: Optional[float] = None, want_wh: Optional[tuple[int, int]] = None,
                draft: bool = False, requested: Optional[float] = None, mode: Optional[str] = None,
                why: Optional[str] = None) -> dict[str, Any]:
    """Kész videó: dekódolható-e (5 mintakocka), a renderelt hossz az idővonalé-e (want_s), a KÉRT hossz (requested)
    teljesült-e — eltérésnél kimondva, miért (mode / why) —, a mérete a vártnak megfelelő-e, nem fekete-e."""
    import cv2
    c = cv2.VideoCapture(str(p))
    try:
        n, fps = c.get(cv2.CAP_PROP_FRAME_COUNT), c.get(cv2.CAP_PROP_FPS) or 0
        w, h = int(c.get(cv2.CAP_PROP_FRAME_WIDTH)), int(c.get(cv2.CAP_PROP_FRAME_HEIGHT))
        good, dark = 0, 0
        for f in (0.05, 0.25, 0.5, 0.75, 0.95):
            c.set(cv2.CAP_PROP_POS_FRAMES, max(0, int(n * f) - 1))
            ok, fr = c.read()
            if ok and fr is not None:
                good += 1
                dark += int(fr.mean() < 8)
    finally:
        c.release()
    if not n or not fps or not good:
        return make("failed", ["The video doesn't play: no frame could be decoded."], [])
    dur = n / fps
    findings, verified = [], [f"Plays: {good} of 5 sample frames decode; {dur:.1f} s at {fps:.0f} fps, {w}×{h}."]
    if want_s and abs(dur - want_s) > max(1.0, 0.1 * want_s):
        findings.append(f"It is {dur:.1f} s long, but its timeline is {want_s:.1f} s — the render may have stopped early; "
                        "render it again (video action=render).")
    if requested:                     # a kért hossz: eltérés soha nem „teljesült” (10-17: 24 s kérve, 17 s lett, mégis ok)
        if abs(dur - requested) > max(1.0, 0.05 * requested):
            reason = why or _LENGTH_BY.get(mode or "", "")
            findings.append(f"Asked for {requested:.0f} s, it is {dur:.0f} s" + (f" ({reason})" if reason else "") + ".")
        else:
            verified.append(f"Length as asked: {requested:.0f} s (it is {dur:.1f} s" + (f", {mode} mode" if mode else "") + ").")
    elif mode == "footage":
        verified.append(f"The length is the footage's: {dur:.1f} s.")
    elif mode:
        verified.append(f"No length was asked: the narration (or, without a voice, the scenes' reading time) sets it — {dur:.1f} s.")
    if want_wh and not draft and (w, h) != tuple(want_wh):
        findings.append(f"It is {w}×{h} — {want_wh[0]}×{want_wh[1]} was expected.")
    if good < 5:
        findings.append(f"Only {good} of 5 sample frames decode.")
    if dark >= 3:
        findings.append(f"{dark} of 5 sample frames are black.")
    return make("needs_check" if findings else "ok", findings, verified)


# ---------- belépési pont ----------

_NO_CHECK = {"files", "history", "extract", "ocr_image", "describe_image", "extract_table_from_image", "generate_image",
             "upload_from_url", "upload_file", "upload_chunk", "validate_document", "chart_redraw",
             "generate_chart_image", "image_to_xlsx"}


def main_file(res: dict[str, Any]) -> Optional[str]:
    """Az eredmény fő fájlja (file_id, vagy az első a files-ban)."""
    fid = res.get("file_id")
    if isinstance(fid, str) and "." in fid:
        return fid
    for f in res.get("files") or []:
        if isinstance(f, dict) and isinstance(f.get("file_id"), str):
            return f["file_id"]
    return None


async def run(tool: str, args: dict[str, Any], res: dict[str, Any]) -> tuple[Optional[dict[str, Any]], Optional[int]]:
    """→ (check, az érintett oldal a preview_url-hez | None). Sosem dob."""
    try:
        return await asyncio.wait_for(_run(tool, args, res), timeout=45)
    except Exception as e:
        log.info("check %s nem futott: %s", tool, e)
        return None, None


def merge_checks(a: dict[str, Any], b: dict[str, Any]) -> dict[str, Any]:
    """Két ellenőrzés egyben: a rosszabb státusz, mindkét lista (ismétlés nélkül)."""
    rank = {"ok": 0, "needs_check": 1, "failed": 2}
    status = max(a["status"], b["status"], key=lambda x: rank.get(x, 1))
    uniq = lambda xs: list(dict.fromkeys(xs))  # noqa: E731
    findings = uniq([f for f in a["findings"] + b["findings"] if f != LOOK]) + ([LOOK] if LOOK in a["findings"] + b["findings"] else [])
    return make(status, findings, uniq(a["verified"] + b["verified"]))


async def _run(tool: str, args: dict[str, Any], res: dict[str, Any]) -> tuple[Optional[dict[str, Any]], Optional[int]]:
    own = normalize(res.get("check"))
    if own is not None:
        out = _path(main_file(res))
        if out is not None and _ext(out) == "pptx" and not res.get("read_only"):
            # saját ellenőrzés + a renderelt pakli-ellenőrzés: minden pptx-kimenet ugyanazon az elemzőn megy át
            deck, page = await check_deck(out, None, tool, args)
            return merge_checks(own, deck), page
        return own, None
    if tool in _NO_CHECK or args.get("read") or res.get("read_only"):    # olvasás (bármely motor): nincs új fájl, nincs mit ellenőrizni
        return None, None
    # háttérmunka: a kész eredményen (digitize / restore / photo_edit — a hívás job_id-val jön vissza)
    if res.get("status") in ("running", "failed", "cancelled") and "kind" in res:
        return None, None
    if res.get("status") == "done" and isinstance(res.get("result"), dict):
        kind, inner = res.get("kind"), res["result"]
        own = normalize(inner.get("check"))
        if own is not None:
            return own, None
        if kind == "restore":
            return await asyncio.to_thread(check_restore, inner), None
        if kind == "digitize":
            return await _check_digitize(inner), None
        fid = main_file(inner)
        return (await generic(_path(fid)) if _path(fid) else None), None
    fid = main_file(res)
    out = _path(fid)
    if out is None:
        return None, None
    src = _path(args.get("file_id"))
    ext = _ext(out)
    if tool == "office" and args.get("app") == "pdf":
        steps = [c for c in args.get("commands") or [] if isinstance(c, dict)]
        red = await asyncio.to_thread(check_redaction, src, out, steps)
        if red is not None:
            return red, None
        if any(s.get("tool") == "form_fill" for s in steps):
            vals = {}
            for s in steps:
                if s.get("tool") == "form_fill":
                    vals.update((s.get("args") or {}).get("values") or {})
            return await asyncio.to_thread(check_form_fill, out, res, vals), None
    if tool == "pdf_form" and ext == "pdf":
        rows = None
        if args.get("data_file_id"):
            from app.tools import forms
            dp = _path(args["data_file_id"])
            rows = await asyncio.to_thread(forms.read_rows, dp) if dp else None
        return await asyncio.to_thread(check_form_fill, out, res, args.get("values"), rows, bool(args.get("flatten"))), None
    if tool == "convert":
        return await check_convert(out, src), None
    if tool == "merge" and args.get("file_ids"):
        total = 0
        for f in args["file_ids"]:
            p = _path(f)
            n = await _pages(p) if p else None
            if n is None:
                return await generic(out), None
            total += n
        return await check_pages_math(out, total, f"{len(args['file_ids'])} files merged"), None
    if tool == "delete_pages" and src is not None and ext == "pdf":
        from app.tools.edits import _parse_page_range
        before = await _pages(src) or 0
        removed = len(_parse_page_range(str(args.get("page_range") or ""), before))
        return await check_pages_math(out, before - removed, f"{removed} of {before} pages removed"), None
    if ext == "pptx":
        return await check_deck(out, src, tool, args)
    if ext == "docx":
        return await with_layout(await asyncio.to_thread(check_docx, out, src, tool, args), out), None
    if tool in ("photocraft", "photo_edit") or (ext in ("psd", "png", "jpg", "jpeg", "webp") and tool not in ("design", "vectorcraft")):
        return await asyncio.to_thread(check_photo, out, src, res), None
    return await generic(out), None


async def _check_digitize(inner: dict[str, Any]) -> dict[str, Any]:
    verified, findings = [], []
    for f in inner.get("files") or []:
        p = _path(f.get("file_id"))
        if p is None:
            continue
        if "fillable" in str(f.get("label", "")).lower() or (inner.get("fields") and _ext(p) == "pdf"):
            c = await asyncio.to_thread(check_fillable, p)
            (verified if c["status"] == "ok" else findings).extend(c["verified"] + c["findings"])
        elif _ext(p) == "docx":
            words = len((await asyncio.to_thread(_docx_text, p)).split())
            (verified if words else findings).append(f"{f.get('label', 'Document')}: {_plural(words, 'word')} read." if words
                                                      else f"{f.get('label', 'Document')}: no text was read.")
    if findings:
        return make("failed" if not verified else "needs_check", findings, verified)
    r = inner.get("readings") or {}
    doubtful = int(r.get("uncertain") or 0) + int(r.get("illegible") or 0)
    if verified and doubtful:
        # a kétes olvasat nem hiba, de ember nézze meg: hol vannak, és hogy a fájlokban jelölve vannak
        pages = ", ".join(map(str, (r.get("pages") or [])[:15])) + ("…" if len(r.get("pages") or []) > 15 else "")
        verified.append(f"{_plural(int(inner.get('pages') or 0), 'page')} read.")
        return make("needs_check", [f"{_plural(int(r.get('uncertain') or 0), 'uncertain reading')} and "
                                    f"{_plural(int(r.get('illegible') or 0), 'illegible word')} on page{'s' if len(r.get('pages') or []) > 1 else ''} "
                                    f"{pages} — marked in the files (Word: highlighted, each with a comment; txt/md: [reading?], [?]; "
                                    "review_pdf: highlighted next to the page image). `readings.items` lists them with their context: "
                                    "check them against the page images."], verified)
    if verified:
        verified.append(f"{_plural(int(inner.get('pages') or 0), 'page')} read; the reader marked no doubtful readings.")
        return make("ok", ["Spelling is read by AI — skim the text for misreads."], verified)
    return make("needs_check", [LOOK], [])
