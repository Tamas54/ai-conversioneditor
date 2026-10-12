"""Két fájl (vagy két változat) összevetése: oldalanként/diánként mi változott, és eltűnt-e valami.

A ChatGPT kérése (10-12): „vizuális és tartalmi különbségnézet — mely oldalon mi változott, illetve eltűnt-e
valami váratlanul". Mindkét fájl PDF-alakján (a pptx/docx/kép is oda konvertálva, temp-cache) oldalpáronként:
  • képpont-eltérés % (kicsinyítve, szürkén) → a változott oldalak, előnézet-párral (preview_url a és b)
  • szöveg: hozzáadott / eltűnt sorok (a máshova költözött sor nem „eltűnt")
  • képek száma oldalanként; pptx-nél diánként az eltűnt alakzatok (kép, szövegdoboz, diagram) neve
"""
from __future__ import annotations

import asyncio
import difflib
from pathlib import Path
from typing import Any, Optional

_W = 480            # az összevetés felbontása (px szélesség)
_CHANGED = 0.4      # ennyi % érezhetően eltérő képpont fölött „változott" az oldal


def _raster(pdf: Path, n: int):
    import fitz
    import numpy as np
    out = []
    with fitz.open(str(pdf)) as d:
        for i in range(min(n, d.page_count)):
            pg = d[i]
            z = _W / max(pg.rect.width, 1)
            pix = pg.get_pixmap(matrix=fitz.Matrix(z, z), colorspace=fitz.csGRAY)
            out.append(np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width))
    return out


def _pages_text(pdf: Path) -> list[list[str]]:
    import fitz
    with fitz.open(str(pdf)) as d:
        return [[ln.strip() for ln in pg.get_text().splitlines() if ln.strip()] for pg in d]


def _images_per_page(pdf: Path) -> list[int]:
    import fitz
    with fitz.open(str(pdf)) as d:
        return [len(pg.get_images(full=True)) for pg in d]


def _pixel_change(a, b) -> float:
    import cv2
    import numpy as np
    if a.shape != b.shape:
        b = cv2.resize(b, (a.shape[1], a.shape[0]), interpolation=cv2.INTER_AREA)
    return round(100 * float((np.abs(a.astype(np.int16) - b.astype(np.int16)) > 32).mean()), 2)


def _pptx_shapes(p: Path) -> list[list[tuple[str, str]]]:
    from pptx import Presentation
    slides = []
    for s in Presentation(str(p)).slides:
        items = []
        for sh in s.shapes:
            kind = ("picture" if getattr(sh, "image", None) is not None else "chart" if getattr(sh, "has_chart", False) and sh.has_chart
                    else "table" if getattr(sh, "has_table", False) and sh.has_table
                    else "text" if sh.has_text_frame and sh.text_frame.text.strip() else "shape")
            label = sh.text_frame.text.strip()[:40] if kind == "text" else (sh.name or kind)
            items.append((kind, label))
        slides.append(items)
    return slides


def _align(sig_a: list[str], sig_b: list[str]) -> list[tuple[Optional[int], Optional[int]]]:
    """Oldalpárok tartalom szerint (egy beszúrt dia ne tegye „változottá" az összes utána jövőt)."""
    pairs: list[tuple[Optional[int], Optional[int]]] = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, sig_a, sig_b, autojunk=False).get_opcodes():
        if op == "equal":
            pairs += list(zip(range(i1, i2), range(j1, j2)))
        elif op == "replace":
            for k in range(max(i2 - i1, j2 - j1)):
                pairs.append((i1 + k if i1 + k < i2 else None, j1 + k if j1 + k < j2 else None))
        elif op == "delete":
            pairs += [(i, None) for i in range(i1, i2)]
        else:
            pairs += [(None, j) for j in range(j1, j2)]
    return pairs


def _compare_sync(pa: Path, pb: Path, ra: Path, rb: Path, max_pages: int) -> dict[str, Any]:
    ta, tb = _pages_text(pa), _pages_text(pb)
    na, nb = len(ta), len(tb)
    ia, ib = _images_per_page(pa), _images_per_page(pb)
    rasters_a, rasters_b = _raster(pa, max_pages), _raster(pb, max_pages)
    rows, missing = [], []
    for a, b in _align(["\n".join(x) for x in ta], ["\n".join(x) for x in tb]):
        if (a is not None and a >= max_pages) or (b is not None and b >= max_pages):
            continue
        row: dict[str, Any] = {"page_a": None if a is None else a + 1, "page_b": None if b is None else b + 1}
        if a is None:
            row["note"] = "new in the second file"
            row["change_pct"] = 100.0
        elif b is None:
            row["note"] = "gone from the second file"
            row["change_pct"] = 100.0
            missing.append(f"Page {a + 1} is gone (“{(ta[a][0][:50] if ta[a] else 'no text')}”).")
        else:
            row["change_pct"] = _pixel_change(rasters_a[a], rasters_b[b]) if a < len(rasters_a) and b < len(rasters_b) else None
            gone = [ln for ln in ta[a] if ln not in tb[b]]
            new = [ln for ln in tb[b] if ln not in ta[a]]
            if gone or new:
                row["text_removed"], row["text_added"] = gone[:8], new[:8]
            if ia[a] > ib[b]:
                missing.append(f"Page {a + 1} → {b + 1}: {ia[a]} picture(s) → {ib[b]}.")
        rows.append(row)
    all_b = {ln for pg in tb for ln in pg}
    gone_lines = [ln for pg in ta for ln in pg if ln not in all_b]
    sa, sb = [ln for pg in ta for ln in pg], [ln for pg in tb for ln in pg]
    sm = difflib.SequenceMatcher(None, sa, sb, autojunk=False)
    added = sum(j2 - j1 for op, _, _, j1, j2 in sm.get_opcodes() if op in ("insert", "replace"))
    removed = sum(i2 - i1 for op, i1, i2, _, _ in sm.get_opcodes() if op in ("delete", "replace"))
    if ra.suffix.lower() == ".pptx" and rb.suffix.lower() == ".pptx":
        try:
            xa, xb = _pptx_shapes(ra), _pptx_shapes(rb)
            rest = [it for sl in xb for it in sl]          # bárhol megvan a második fájlban → nem tűnt el
            for i, shapes in enumerate(xa[:max_pages]):
                for kind, label in shapes:
                    if (kind, label) in rest:
                        rest.remove((kind, label))
                    elif kind in ("picture", "chart", "table") or (kind == "text" and label):
                        missing.append(f"Slide {i + 1}: {kind} “{label}” is gone.")
        except Exception:
            pass
    if gone_lines:
        missing.append(f"{len(gone_lines)} line(s) of text no longer appear anywhere, e.g. “{gone_lines[0][:60]}”.")
    changed = [r for r in rows if (r.get("change_pct") or 0) > _CHANGED or r.get("text_added") or r.get("text_removed") or r.get("note")]
    return {"pages_a": na, "pages_b": nb, "changes": changed,
            "text": {"lines_added": added, "lines_removed": removed, "removed_sample": gone_lines[:10]},
            "missing": missing[:30]}


async def compare(file_a: str, file_b: str, max_pages: int = 40) -> dict[str, Any]:
    from app.storage import display_label
    from app.tools import previews
    from app.tools.crafts import existing
    ra, rb = existing(file_a), existing(file_b)
    pa, pb = await previews.as_pdf(ra), await previews.as_pdf(rb)
    res = await asyncio.to_thread(_compare_sync, pa, pb, ra, rb, max(1, min(200, int(max_pages))))
    for r in res["changes"]:
        if r["page_a"]:
            r["preview_a"] = previews.preview_url(file_a, r["page_a"])
        if r["page_b"]:
            r["preview_b"] = previews.preview_url(file_b, r["page_b"])
    n = len(res["changes"])
    res["changed_pages"] = [r["page_b"] or r["page_a"] for r in res["changes"]]
    res["a"], res["b"] = {"file_id": file_a, "name": display_label(file_a)}, {"file_id": file_b, "name": display_label(file_b)}
    t = res["text"]
    res["summary"] = ("No visible difference." if not n and not res["missing"] else
                      f"{n} page(s) differ (first file {res['pages_a']} pages, second {res['pages_b']})"
                      + (f"; {t['lines_removed']} text line(s) removed, {t['lines_added']} added" if t["lines_removed"] or t["lines_added"] else "")
                      + (f"; {len(res['missing'])} thing(s) gone — see missing." if res["missing"] else "; nothing went missing."))
    return res


def latest_source(file_id: str) -> Optional[str]:
    """Ha nincs második fájl: a változat forrása (amiből készült) — „mi változott az előzőhöz képest"."""
    from app.storage import read_meta
    return (read_meta(file_id) or {}).get("source_file_id")
