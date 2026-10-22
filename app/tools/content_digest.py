"""Tartalmi ujjlenyomat (pptx, docx): ugyanaz-e két fájl a felhasználó szemében.

A szerkesztők (DeckCraft, WordCraft) mentéskor másképp sorosítanak (más részfájlok, más XML), ezért bájt-
összevetés nem jó. Itt ami LÁTSZIK: diák/bekezdések sorrendje, alakzatok helye és mérete, szöveg, betű-
formázás (félkövér, dőlt, aláhúzás, méret, szín), képek tartalma, táblázatok szövege, oldalbeállítás.
DOCX-ben a szerkezet is (10-22, ügynök-jelentés: a valódi lábjegyzet és az oldaltörés-beállítás után „Nothing
changed” jött, a szerkesztő pedig az ilyen mentést változatlannak vette és eldobta): lábjegyzet / végjegyzet,
megjegyzés, követett beszúrás / törlés, sor- és oldaltörés, mező, a bekezdés tördelési és térköz-beállításai,
számozás, élőfej / élőláb, szakaszbeállítás, a stílusok definíciója.
Kétes esetben „más” (None vagy eltérő érték) — így legfeljebb fölösleges új verzió lesz, adat nem vész el.
"""
from __future__ import annotations

import hashlib
import logging
import re
from pathlib import Path
from typing import Any, Optional

log = logging.getLogger("aice.digest")
_GRID = 12700   # 1 pt EMU-ban: a szerkesztők kerekítése ne számítson változásnak


def _q(v: Optional[int]) -> Optional[int]:
    return None if v is None else round(v / _GRID)


def _color(font) -> Optional[str]:
    try:
        return str(font.color.rgb) if font.color is not None and font.color.type is not None else None
    except Exception:
        return None


def _runs(paragraphs) -> list:
    out = []
    for p in paragraphs:
        runs = [(r.text, r.font.bold, r.font.italic, r.font.underline,
                 r.font.size.pt if r.font.size else None, _color(r.font)) for r in p.runs if r.text]
        # a szomszédos, azonos formázású futamok összevonva (a szerkesztők másképp darabolnak)
        merged: list = []
        for t, *fmt in runs:
            if merged and merged[-1][1:] == fmt:
                merged[-1][0] += t
            else:
                merged.append([t, *fmt])
        out.append((getattr(p, "level", 0), [tuple(m) for m in merged]))
    return out


def _pptx(path: Path) -> list:
    from pptx import Presentation
    from pptx.util import Emu  # noqa: F401
    prs = Presentation(str(path))
    slides = []

    def shapes(coll) -> list:
        items = []
        for sh in coll:
            d: dict[str, Any] = {"t": str(sh.shape_type), "g": (_q(sh.left), _q(sh.top), _q(sh.width), _q(sh.height))}
            if sh.has_text_frame:
                d["x"] = _runs(sh.text_frame.paragraphs)
            if getattr(sh, "image", None) is not None:
                try:
                    d["img"] = hashlib.sha1(sh.image.blob).hexdigest()
                except Exception:
                    pass
            if getattr(sh, "has_table", False) and sh.has_table:
                d["tbl"] = [[c.text for c in r.cells] for r in sh.table.rows]
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                d["ch"] = (str(ch.chart_type), [(s.name, list(s.values)) for s in ch.plots[0].series] if ch.plots else [])
            if str(sh.shape_type).startswith("GROUP"):
                d["grp"] = shapes(sh.shapes)
            items.append(d)
        return sorted(items, key=repr)       # a z-sorrend helyett a tartalom számít

    for s in prs.slides:
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        slides.append({"layout": s.slide_layout.name, "shapes": shapes(s.shapes), "notes": notes.strip()})
    return [(_q(prs.slide_width), _q(prs.slide_height)), slides]


def _docx(path: Path) -> list:
    from docx import Document
    doc = Document(str(path))
    body = []
    for p in doc.paragraphs:
        body.append(("p", p.style.name if p.style is not None else None, p.paragraph_format.alignment,
                     _runs([p])[0][1]))
    tables = [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
    imgs = sorted(hashlib.sha1(rel.target_part.blob).hexdigest() for rel in doc.part.rels.values()
                  if "image" in rel.reltype and not rel.is_external)
    secs = [(_q(s.page_width), _q(s.page_height), _q(s.left_margin), _q(s.right_margin), _q(s.top_margin),
             _q(s.bottom_margin)) for s in doc.sections]
    return [body, tables, imgs, secs, _docx_structure(path)]


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_OFF = ("0", "false", "off", "none")
# a bekezdés tördelési / térköz / számozási beállításai, amik a lapon látszanak (a szerkesztők alapértékei nélkül)
_PPR_FLAGS = ("keepNext", "keepLines", "pageBreakBefore", "widowControl", "suppressAutoHyphens", "contextualSpacing")
_PPR_VALS = {"spacing": ("before", "after", "line", "lineRule", "beforeLines", "afterLines"),
             "ind": ("left", "start", "right", "end", "firstLine", "hanging"), "jc": ("val",), "outlineLvl": ("val",),
             "numPr": ()}


def _attrs(el, names) -> tuple:
    return tuple(el.get(_W + n) for n in names)


def _ppr(p) -> tuple:
    ppr = p.find(_W + "pPr")
    if ppr is None:
        return ()
    out = []
    for f in _PPR_FLAGS:
        e = ppr.find(_W + f)
        if e is not None:
            out.append((f, (e.get(_W + "val") or "1").lower() not in _OFF))
    for tag, names in _PPR_VALS.items():
        e = ppr.find(_W + tag)
        if e is None:
            continue
        if tag == "numPr":
            il, ni = e.find(_W + "ilvl"), e.find(_W + "numId")
            out.append((tag, il.get(_W + "val") if il is not None else None, ni.get(_W + "val") if ni is not None else None))
        else:
            out.append((tag,) + _attrs(e, names))
    if ppr.find(_W + "sectPr") is not None:
        out.append(("sect", _sect(ppr.find(_W + "sectPr"))))
    return tuple(out)


def _inline(el) -> str:
    """A bekezdés tartalma jelölésekkel: szöveg (a követett beszúrás is), ⌫törölt, törések, jegyzet- és megjegyzés-
    hivatkozás, mezőkód, kép — dokumentum-sorrendben. A rsid-ek, a nyelvi hibajelek, a renderelt oldaltörés-jel és a
    könyvjelzők nem számítanak (a szerkesztők maguktól is írják / eldobják őket)."""
    out = []
    for e in el.iter():
        t = e.tag
        if t == _W + "t":
            out.append(e.text or "")
        elif t == _W + "delText":
            out.append(f"⌫{e.text or ''}⌫")
        elif t == _W + "br":
            out.append({"page": "⏎P", "column": "⏎C"}.get(e.get(_W + "type") or "", "\n"))
        elif t == _W + "tab" and e.getparent() is not None and e.getparent().tag == _W + "r":
            out.append("\t")
        elif t in (_W + "footnoteReference", _W + "endnoteReference"):
            out.append("⟨fn⟩" if t.endswith("footnoteReference") else "⟨en⟩")
        elif t == _W + "commentReference":
            out.append("⟨c⟩")
        elif t == _W + "instrText":
            out.append(f"⟨fld {' '.join((e.text or '').split())}⟩")
        elif t == _W + "fldSimple":
            out.append(f"⟨fld {' '.join((e.get(_W + 'instr') or '').split())}⟩")
        elif t.endswith("}drawing") or t.endswith("}pict"):
            out.append("⟨img⟩")
        elif t in (_W + "ins", _W + "del", _W + "moveFrom", _W + "moveTo"):
            out.append(f"⟨{t.rsplit('}', 1)[1]}⟩")
    return "".join(out)


def _sect(sp) -> tuple:
    def a(tag, names):
        e = sp.find(_W + tag)
        return None if e is None else tuple(
            (_q(int(v) * 635) if v and v.lstrip("-").isdigit() and n not in ("orient", "fmt", "start", "num", "val") else v)
            for v, n in zip(_attrs(e, names), names))
    cols, typ = a("cols", ("num", "space")), a("type", ("val",))
    cols = (cols[0] or "1", cols[1]) if cols else ("1", None)                  # az alapértékek kiírva vagy elhagyva: ugyanaz
    typ = None if typ in (None, ("nextPage",)) else typ
    return (a("pgSz", ("w", "h", "orient")), a("pgMar", ("top", "right", "bottom", "left", "header", "footer", "gutter")),
            cols, a("pgNumType", ("fmt", "start")), typ, a("vAlign", ("val",)),
            sp.find(_W + "titlePg") is not None,
            tuple(sorted((r.tag.rsplit("}", 1)[1], r.get(_W + "type")) for r in sp if r.tag in (_W + "headerReference", _W + "footerReference"))))


# a stílus-definícióból csak ami a lapon látszik, a szerkesztők sorosítási szokásai nélkül (a WordCraft a téma-betűt
# konkrét betűre oldja, az üres elemeket elhagyja — 10-22 mérés): méret, vastag, dőlt, aláhúzott, szín, igazítás,
# térköz, behúzás, tördelés
_STYLE_KEEP = {"sz": ("val",), "b": ("val",), "i": ("val",), "u": ("val",), "color": ("val",), "caps": ("val",),
               "jc": ("val",), "spacing": ("before", "after", "line"), "ind": ("left", "start", "firstLine", "hanging"),
               "keepNext": ("val",), "keepLines": ("val",), "pageBreakBefore": ("val",), "outlineLvl": ("val",)}


def _coarse(color: Optional[str]) -> Optional[str]:
    """A téma-szín feloldása motoronként egy-két egységgel eltér (10-22: a címsor accent1 + 75% Wordben 365F91, a
    WordCraft 0.9-ben 376092) — szemre ugyanaz. Csatornánként 16-os lépcsőre kerekítve hasonlítunk."""
    if not color or not re.fullmatch(r"[0-9A-Fa-f]{6}", color):
        return color
    return "".join(f"{min(255, round(int(color[k:k + 2], 16) / 16) * 16):02X}" for k in (0, 2, 4))


def _style_props(s_) -> tuple:
    out = []
    for sect in ("pPr", "rPr"):
        e = s_.find(_W + sect)
        if e is None:
            continue
        for c in e:
            tag = c.tag.rsplit("}", 1)[1]
            if tag in _STYLE_KEEP:
                vals = _attrs(c, _STYLE_KEEP[tag])
                if tag in ("b", "i", "keepNext", "keepLines", "pageBreakBefore", "caps"):
                    vals = ((vals[0] or "1").lower() not in _OFF,)
                elif tag == "color":
                    vals = (_coarse(vals[0]),)
                out.append((tag, vals))
    return tuple(sorted(out, key=repr))


def _part_paragraphs(root) -> list:
    return [(_ppr(p), _inline(p)) for p in root.iter(_W + "p")]


def _docx_structure(path: Path) -> list:
    """Ami a python-docx bekezdés-futamaiból kimarad, de a lapon látszik: lásd a modul leírását."""
    import zipfile
    from lxml import etree
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())

        def root(name):
            return etree.fromstring(z.read(name)) if name in names else None
        doc = root("word/document.xml")
        body = doc.find(_W + "body")
        out: list = [_part_paragraphs(body), [_sect(sp) for sp in body.iter(_W + "sectPr")]]
        for part, tag in (("word/footnotes.xml", "footnote"), ("word/endnotes.xml", "endnote")):
            r = root(part)
            out.append([] if r is None else [_part_paragraphs(n) for n in r.iter(_W + tag)
                                             if (n.get(_W + "type") or "normal") == "normal"])
        com = root("word/comments.xml")
        out.append([] if com is None else [(c.get(_W + "author"), _inline(c)) for c in com.iter(_W + "comment")])
        out.append(sorted(_inline(root(n)) for n in names if re.match(r"word/(header|footer)\d*\.xml$", n)))
        st = root("word/styles.xml")
        used = {e.get(_W + "val") for part in [doc] for e in part.iter() if e.tag in (_W + "pStyle", _W + "rStyle", _W + "tblStyle")}
        styles = []
        if st is not None:
            for s_ in st.iter(_W + "style"):
                if s_.get(_W + "styleId") not in used and s_.get(_W + "default") not in ("1", "true"):
                    continue
                styles.append((s_.get(_W + "styleId"), _style_props(s_)))
        out.append(sorted(styles, key=repr))
    return out


def digest(path: Path) -> Optional[str]:
    """sha1 a látható tartalomról; None, ha a formátum nem támogatott vagy nem olvasható."""
    ext = path.suffix.lower().lstrip(".")
    try:
        data = _pptx(path) if ext == "pptx" else _docx(path) if ext == "docx" else None
    except Exception as e:
        log.info("ujjlenyomat nem készült (%s): %s", path.name, e)
        return None
    return None if data is None else hashlib.sha1(repr(data).encode("utf-8")).hexdigest()


def same_content(a: Path, b: Path) -> bool:
    da, db = digest(a), digest(b)
    return da is not None and da == db
