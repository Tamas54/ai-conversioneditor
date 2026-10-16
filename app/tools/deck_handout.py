"""Pakli → Word-kiosztmány (handout), AI nélkül: diánként a dia képe, címe, szövege (szintekkel), táblázatai és a
előadói jegyzetek. A diaképeket a DeckCraft rendereli; a szöveget a python-pptx olvassa (így bármilyen pptx-ből megy,
nem csak a sajátunkból). Szedés: pandoc + az általános A4-es referencia-docx.
"""
from __future__ import annotations

import asyncio
import re
import shutil
from pathlib import Path
from typing import Any


def _md(t: str) -> str:
    """Szöveg markdownba szó szerint: a „1.”-gyel kezdődő sor ne legyen számozott lista, a * _ # [ ] ne formázzon."""
    t = re.sub(r"([\\`*_\[\]<>])", r"\\\1", str(t))
    t = re.sub(r"^(\s*)(\d+)([.)])", r"\1\2\\\3", t)
    return re.sub(r"^(\s*)([#+-])", r"\1\\\2", t)


_PAGE_NO = re.compile(r"^\s*(\d+\s*(/|of)\s*\d+|\d{1,3})\s*$")


def _spec_text(s: dict[str, Any]) -> tuple[str, list[str]]:
    """A saját paklijaink leírásából (deck_spec) — tiszta szöveg, láblécek és díszítő számok nélkül."""
    lay, lines = s.get("layout"), []
    title = s.get("title") or ""
    for key in ("kicker", "subtitle", "text", "caption"):      # a szakasz „01”-es sorszáma díszítés, kimarad
        if s.get(key):
            lines += [_md(s[key]), ""]
    lines += [f"- {_md(b)}" for b in s.get("bullets") or []]
    if lay == "quote":
        title = title or "Quote"
        lines += [f"> “{s.get('quote', '')}”" + (f" — {s['author']}" if s.get("author") else ""), ""]
    for st in s.get("stats") or []:
        lines.append(f"- **{st.get('value', '')}** {st.get('label', '')}")
    if lay == "chart":
        cats = s.get("categories") or []
        for se in s.get("series") or []:
            lines.append(f"- {se.get('name', '')}: " + ", ".join(f"{c}: {v}" for c, v in zip(cats, se.get("values") or [])))
    if s.get("headers") or s.get("rows"):
        hdr = s.get("headers") or [""] * len((s.get("rows") or [[]])[0])
        lines += ["", "| " + " | ".join(map(str, hdr)) + " |", "|" + "---|" * len(hdr)]
        lines += ["| " + " | ".join(str(c).replace("|", "/") for c in r) + " |" for r in s.get("rows") or []]
    return title, lines + [""]


def _repeated(prs) -> set[str]:
    """A diák legalább felén ismétlődő szövegek (lábléc, cégnév, dátum) — ezek nem a dia tartalma."""
    from collections import Counter
    c: Counter = Counter()
    for slide in prs.slides:
        c.update({sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()})
    n = len(prs.slides)
    return {t for t, k in c.items() if n >= 3 and k >= max(3, n / 2)}


def _slide_text(slide, skip: set[str], SH: int) -> tuple[str, list[str]]:
    """(cím, markdown-sorok) egy idegen pptx diájáról: felsorolások szintekkel, táblázatok, diagram-adatok; a lábléc,
    az oldalszám és a sok dián ismétlődő szöveg kimarad. Cím nélküli diánál a legnagyobb betűs rövid szöveg a cím."""
    from pptx.enum.shapes import PP_PLACEHOLDER as T
    title, lines = "", []
    shapes = sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0)))

    def noise(sh) -> bool:
        t = sh.text_frame.text.strip()
        return (not t or t in skip or bool(_PAGE_NO.match(t))
                or (sh.is_placeholder and sh.placeholder_format.type in (T.SLIDE_NUMBER, T.FOOTER, T.DATE))
                or ((sh.top or 0) > SH * 0.9 and len(t) < 80))

    def biggest(sh) -> float:
        sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
        return max(sizes) if sizes else 0
    one_line = lambda sh: " ".join(sh.text_frame.text.split())
    texts = [sh for sh in shapes if sh.has_text_frame and not noise(sh)]
    for sh in texts:
        if sh.is_placeholder and sh.placeholder_format.type in (T.TITLE, T.CENTER_TITLE, T.VERTICAL_TITLE):
            title = one_line(sh)
    if not title:                       # cím-helyőrző nélkül: a legnagyobb betűs rövid szöveg, egyenlőségnél a legfelső
        cands = [sh for sh in texts if len(one_line(sh)) <= 120]
        if cands:
            best = max(cands, key=lambda sh: (biggest(sh), -(sh.top or 0)))
            title = one_line(best) if biggest(best) >= 20 else one_line(cands[0])
    for sh in shapes:
        if sh.has_text_frame:
            if noise(sh) or one_line(sh) == title:
                continue
            paras = [(p.level, "".join(r.text for r in p.runs).strip()) for p in sh.text_frame.paragraphs]
            paras = [(lv, t) for lv, t in paras if re.search(r"\w", t)]       # csak betű/szám — üres pont, díszjel nem
            if not paras:
                continue
            bullet = len(paras) > 1 or (sh.is_placeholder and sh.placeholder_format.type in (T.BODY, T.OBJECT))
            for lv, t in paras:
                lines.append(("  " * lv + "- " + _md(t)) if bullet else _md(t))
            lines.append("")
        elif sh.has_table:
            rows = [[c.text.replace("|", "/").replace("\n", " ").strip() for c in r.cells] for r in sh.table.rows]
            if rows:
                lines += ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
                lines += ["| " + " | ".join(r) + " |" for r in rows[1:]] + [""]
        elif sh.has_chart:
            ch = sh.chart
            try:
                cats = [str(c) for c in ch.plots[0].categories]
                for se in ch.plots[0].series:
                    vals = ", ".join(f"{c}: {v:g}" for c, v in zip(cats, se.values) if v is not None)
                    lines.append(f"- {se.name}: {vals}")
                lines.append("")
            except Exception:
                pass
    return title, lines


async def build(file_id: str, slide_images: bool = True) -> dict[str, Any]:
    import pypandoc
    from pptx import Presentation
    from app.storage import display_label, output_path, public_url, temp_path, write_meta
    from app.storage.files import new_file_id
    from app.tools import crafts
    src = crafts.existing(file_id)
    if src.suffix.lower() != ".pptx":
        raise ValueError("A handout can be made from a PowerPoint (.pptx) file.")
    stem = Path(display_label(file_id)).stem
    work = temp_path(new_file_id("ho"))
    work.mkdir(parents=True)
    try:
        prs = await asyncio.to_thread(Presentation, str(src))
        pngs: list[Path] = []
        if slide_images:
            try:
                await crafts.run_cli(crafts.office_cli("deck"), "deckcraft-cli", "render", str(src), "--all", "--scale", "1",
                                     str(work / "s.png"))
                d = work / "s.png"
                pngs = sorted(d.glob("*.png")) if d.is_dir() else sorted(work.glob("*.png"))
            except Exception:
                pngs = []
        md = [f"# {stem}", ""]
        from app.storage import read_meta
        spec_slides = ((read_meta(file_id) or {}).get("deck_spec") or {}).get("slides") or []
        own = len(spec_slides) == len(prs.slides)      # a saját paklink: a leírásból jön a tiszta szöveg
        skip = set() if own else await asyncio.to_thread(_repeated, prs)
        for i, slide in enumerate(prs.slides, 1):
            title, lines = _spec_text(spec_slides[i - 1]) if own else _slide_text(slide, skip, prs.slide_height)
            md += [f"## {i}. {_md(title) if title else 'Slide ' + str(i)}", ""]
            if i <= len(pngs):
                md += [f"![]({pngs[i - 1]}){{width=100%}}", ""]
            md += lines
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    md += ["**Notes:** " + notes.replace("\n", " "), ""]
        from app.routes.make import _doc_reference
        ref = await asyncio.to_thread(_doc_reference)
        fid = new_file_id("docx")
        await asyncio.to_thread(pypandoc.convert_text, "\n".join(md), "docx", format="markdown",
                                outputfile=str(output_path(fid)), extra_args=[f"--reference-doc={ref}", f"--resource-path={work}"])
        from app.tools import ooxml_safe
        await asyncio.to_thread(ooxml_safe.fix, output_path(fid))     # iPhone/Pages: sémahelyes docx
    finally:
        shutil.rmtree(work, ignore_errors=True)
    label = f"{stem} (handout).docx"
    write_meta(fid, source_file_id=file_id, operation="handout from slides", label=label)
    return {"file_id": fid, "url": public_url(fid), "label": label, "slides": len(prs.slides),
            "size_kb": output_path(fid).stat().st_size // 1024}
