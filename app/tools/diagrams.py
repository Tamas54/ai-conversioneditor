"""Szakmai magyarázó ábrák: döntési fa, sávos (szereplőnkénti) folyamat, kapcsolati háló — EGY fordító a videónak
és a paklinak.

    from_plan()   az ügynök ábra-leírása → ellenőrzött, tárolható leírás (azonosítók, fajták, a számok a forrásból)
    merge()       célzott javítás: csak a megadott csúcsok / élek / beállítások változnak (id, illetve from+to szerint)
    layout()      determinisztikus elrendezés em-egységben (rendezett fa, sávrács, gyűrűs háló), tiszta Python
    to_graphic()  → a videó MEGLÉVŐ „graphic” elemei (rect/text/line/arrow/path/circle, 0–1 koordináták, anim/at/dur)
    to_pptx()     → natív, szerkeszthető PowerPoint-alakzatok (lekerekített téglalap, szöveg, összekötő, jelmagyarázat)

Szemantikus nyilak: minden élfajtának EGY rögzített stílusa van (szín, vastagság, nyílhegy), a bizonyosság a vonal
mintája (szaggatott = bizonytalan, pontozott + „?” = vitatott). Az együttmozgás (correlation) SOHA nem néz ki
oksági nyílnak: nincs nyílhegye, a két végén pötty, középről kifelé rajzolódik, a jelmagyarázat kimondja.
Szereplő-regiszter: ugyanaz a szereplő (intézmény, mutató) a projektben mindenhol ugyanazt a színt kapja — videóban a
forgatókönyvben (story["actors"]) tárolva, így a későbbi jelenetek és szerkesztések is ezt használják.
"""
from __future__ import annotations

import colorsys
import copy
import math
import re
from typing import Any, Callable, Optional

TYPES = ("decision_tree", "swimlane", "network")
ALIASES = {"tree": "decision_tree", "decision": "decision_tree", "decisiontree": "decision_tree",
           "decision-tree": "decision_tree", "flowchart": "swimlane", "swim_lane": "swimlane", "swim-lane": "swimlane",
           "lanes": "swimlane", "graph": "network", "relations": "network"}
EDGE_KINDS = ("sequence", "hierarchy", "money", "cause", "correlation", "link")
_KIND_ALIAS = {"then": "sequence", "next": "sequence", "flow": "sequence", "step": "sequence", "reports_to": "hierarchy",
               "subordinate": "hierarchy", "oversees": "hierarchy", "owns": "hierarchy", "ownership": "hierarchy",
               "controls": "hierarchy", "pays": "money", "payment": "money", "funds": "money", "goods": "money",
               "causes": "cause", "causal": "cause", "leads_to": "cause", "correlates": "correlation",
               "co-moves": "correlation", "comoves": "correlation", "moves_together": "correlation", "related": "link",
               "relation": "link", "plain": "link"}
CERTAINTY = ("certain", "uncertain", "disputed")
_CERT_ALIAS = {"sure": "certain", "known": "certain", "unsure": "uncertain", "likely": "uncertain", "possible": "uncertain",
               "unproven": "uncertain", "contested": "disputed"}
NODE_KINDS = {"decision_tree": ("question", "condition", "outcome"), "swimlane": ("start", "step", "decision", "end"),
              "network": ("node",)}
LIMITS = {"decision_tree": 15, "swimlane": 14, "network": 14}
MAX_EDGES = 24

# minden élfajtának EGY rögzített stílusa (sötét videóháttér / világos pakli) — a jelmagyarázat is ebből készül
STYLE = {
    "sequence":    {"dark": "#E4E7EE", "light": "#3B4252", "width": 4, "head": "triangle", "sym": False},
    "hierarchy":   {"dark": "#8DB6FF", "light": "#2C5BC0", "width": 5, "head": "chevron", "sym": False},
    "money":       {"dark": "#5BD08F", "light": "#1D8A55", "width": 9, "head": "triangle", "sym": False},
    "cause":       {"dark": "#FF7A5C", "light": "#C2412A", "width": 5, "head": "triangle", "sym": False},
    "correlation": {"dark": "#C7A2FF", "light": "#7445C2", "width": 4, "head": "dots", "sym": True},
    "link":        {"dark": "#A3A9B6", "light": "#7A808C", "width": 3, "head": None, "sym": True},
}
DASH = {"uncertain": (4.0, 2.8), "disputed": (0.12, 2.6)}       # a vonalvastagság szorzói: [vonás, köz]
BADGE = "#F2C14E"
LEGEND = {
    "en": {"sequence": "then", "hierarchy": "is above (hierarchy)", "money": "money / goods flow", "cause": "causes",
           "correlation": "moves together, not shown to cause", "link": "related", "uncertain": "dashed: uncertain",
           "disputed": "dotted + ?: disputed"},
    "hu": {"sequence": "ezután", "hierarchy": "fölérendelt (hierarchia)", "money": "pénz- / áruáramlás", "cause": "okozza",
           "correlation": "együtt mozog, az okozás nem igazolt", "link": "kapcsolódik", "uncertain": "szaggatott: bizonytalan",
           "disputed": "pontozott + ?: vitatott"},
}
_NAMED = {"red": "#E5484D", "green": "#30A46C", "blue": "#3E7BFA", "yellow": "#F5C542", "orange": "#F76B15",
          "purple": "#8E4EC6", "violet": "#8E4EC6", "teal": "#12A594", "pink": "#E93D82", "grey": "#8B8D98",
          "gray": "#8B8D98", "white": "#F4F5F8", "black": "#16181D"}

PX, PY, LH = 0.65, 0.5, 1.22          # csúcs belső margója és sormagasság (em) — a sormagasság a video._fit_size-é
LAB, LEG, HEAD = 0.8, 0.78, 0.95     # élcímke, jelmagyarázat, sávfejléc betűmérete a csúcsbetűhöz képest


def _id(x: Any) -> str:
    return re.sub(r"[^\w.-]", "", str(x or "").strip())[:24]


def _txt(x: Any, n: int, what: str, clean: Optional[Callable[..., str]]) -> str:
    if clean:
        return clean(x, n, what)
    s = re.sub(r"\s+", " ", str(x or "")).strip()
    return s if len(s) <= n else s[:n - 1].rsplit(" ", 1)[0] + "…"


def _hex(v: Any) -> Optional[str]:
    if not isinstance(v, str):
        return None
    v = _NAMED.get(v.strip().lower(), v.strip())
    if re.fullmatch(r"#?[0-9a-fA-F]{6}", v):
        return ("#" + v.lstrip("#")).upper()
    return None


def _kind(v: Any, where: str) -> str:
    k = str(v or "").strip().lower().replace(" ", "_")
    k = _KIND_ALIAS.get(k, k)
    if k not in EDGE_KINDS:
        raise ValueError(f"{where}: kind “{v}” is not an edge kind — use {' | '.join(EDGE_KINDS)} "
                         "(correlation = moves together, without claiming a cause).")
    return k


def _cert(v: Any, where: str) -> str:
    c = str(v or "certain").strip().lower()
    c = _CERT_ALIAS.get(c, c)
    if c not in CERTAINTY:
        raise ValueError(f"{where}: certainty “{v}” — use {' | '.join(CERTAINTY)}.")
    return c


# ---------- leírás: ellenőrzés, normalizálás, célzott javítás ----------

def _expand_branches(parent: str, branches: Any, nodes: list[dict], edges: list[dict], used: set[str]) -> None:
    """Fa-rövidítés: {"branches": [{"label", "text", "kind"?, "branches"?}]} → csúcsok + élek (az id-k a szülőé + .1, .2…)."""
    for k, b in enumerate(branches if isinstance(branches, list) else [], 1):
        if isinstance(b, str):
            b = {"text": b}
        if not isinstance(b, dict):
            continue
        nid = _id(b.get("id")) or f"{parent}.{k}"
        while nid in used:
            nid += "b"
        used.add(nid)
        node = {kk: v for kk, v in b.items() if kk not in ("branches", "label", "edge", "certainty", "kind_of_edge")}
        node["id"] = nid
        nodes.append(node)
        e = {"from": parent, "to": nid}
        for src, dst in (("label", "label"), ("certainty", "certainty"), ("edge", "kind"), ("kind_of_edge", "kind")):
            if b.get(src) not in (None, ""):
                e[dst] = b[src]
        edges.append(e)
        if b.get("branches"):
            _expand_branches(nid, b["branches"], nodes, edges, used)


def from_plan(c: dict[str, Any], bad: Optional[Callable[[str], list]] = None,
              clean: Optional[Callable[..., str]] = None) -> tuple[Optional[dict[str, Any]], list[str]]:
    """Az ábra leírása → (tárolható leírás | None, megjegyzések). Szerkezeti hiba → ValueError, ami megmondja, mit kell
    javítani. A számok ugyanazon a próbán mennek át, mint a többi ábra szövegei (bad: a forrásban nem szereplő számok):
    az élcímke kimarad (az él marad), a csúcs szövegénél az egész ábra kimarad — a megjegyzés megmondja, miért."""
    typ = ALIASES.get(str(c.get("type") or "").lower(), str(c.get("type") or "").lower())
    if typ not in TYPES:
        raise ValueError(f"diagram type: {' | '.join(TYPES)}")
    title = _txt(c.get("title"), 60, "Chart title", clean)
    out: dict[str, Any] = {"type": typ, "title": title}
    for k in ("direction", "reveal", "focus", "pace", "backdrop", "center"):
        if c.get(k) not in (None, ""):
            out[k] = c[k]
    out["direction"] = str(out.get("direction") or "auto").lower()
    if out["direction"] not in ("auto", "down", "right", "outline"):
        raise ValueError(f"diagram “{title}”: direction is auto | down | right (the way the flow goes) | outline (a tree "
                         f"as indented rows), not {out['direction']!r}.")
    out["reveal"] = str(out.get("reveal") or "steps").lower()
    if out["reveal"] not in ("steps", "all"):
        raise ValueError(f"diagram “{title}”: reveal is steps (builds up step by step — default) | all.")
    if "pace" in out:
        try:
            out["pace"] = round(min(4.0, max(0.3, float(out["pace"]))), 2)
        except (TypeError, ValueError):
            raise ValueError(f"diagram “{title}”: pace is seconds per reveal step (0.3–4), e.g. 1.5 for a slower build-up.")
    if "focus" in out:
        out["focus"] = out["focus"] not in (False, "false", "off", "no", 0)
    if "backdrop" in out:
        try:
            out["backdrop"] = max(0.0, min(1.0, float(out["backdrop"])))
        except (TypeError, ValueError):
            raise ValueError(f"diagram “{title}”: backdrop is 0–1 (the dark veil over the picture; default 0.82).")
    else:
        out["backdrop"] = 0.82                      # sűrű ábra: a kép csak halvány textúra mögötte
    acts = c.get("actors") if isinstance(c.get("actors"), dict) else {}
    out["actors"] = {}
    for name, a in acts.items():
        a = a if isinstance(a, dict) else ({"color": a} if isinstance(a, str) else {})
        nm = _txt(name, 40, "Actor", clean)
        if nm:
            out["actors"][nm] = {k: v for k, v in {"color": _hex(a.get("color")), "short": _txt(a.get("short"), 14, "Actor short name", clean) or None}.items() if v}
    if isinstance(c.get("legend"), dict):
        out["legend"] = {k: _txt(v, 44, "Legend text", clean) for k, v in c["legend"].items() if k in LEGEND["en"] and v}

    raw_nodes = c.get("nodes") if c.get("nodes") is not None else c.get("steps")
    raw_nodes = list(raw_nodes) if isinstance(raw_nodes, list) else []
    raw_edges = [e for e in (c.get("edges") or []) if isinstance(e, dict)]
    used = {_id(n.get("id")) for n in raw_nodes if isinstance(n, dict) and n.get("id")}
    if typ == "decision_tree" and (c.get("root") or c.get("question") or c.get("branches")) and not raw_nodes:
        root = c.get("root") or c.get("question") or ""
        root = root if isinstance(root, dict) else {"text": root}
        rid = _id(root.get("id")) or "n1"
        used.add(rid)
        raw_nodes = [{**{k: v for k, v in root.items() if k != "branches"}, "id": rid}]
        _expand_branches(rid, c.get("branches") or root.get("branches") or [], raw_nodes, raw_edges, used)
    nodes: list[dict[str, Any]] = []
    k_auto = 0
    for i, n in enumerate(raw_nodes):
        if isinstance(n, str):
            n = {"text": n}
        if not isinstance(n, dict):
            continue
        nid = _id(n.get("id"))
        if not nid:
            k_auto += 1
            while f"n{k_auto}" in used:
                k_auto += 1
            nid = f"n{k_auto}"
            used.add(nid)
        if any(x["id"] == nid for x in nodes):
            raise ValueError(f"diagram “{title}”: two nodes have the id “{nid}” — give each node its own id.")
        node: dict[str, Any] = {"id": nid, "text": _txt(n.get("text") or n.get("label") or n.get("name"), 70, "Diagram text", clean)}
        if not node["text"]:
            raise ValueError(f"diagram “{title}”: node “{nid}” has no text — every box needs text (≤ 70 characters).")
        kd = str(n.get("kind") or "").lower()
        if kd:
            if kd not in NODE_KINDS[typ]:
                raise ValueError(f"diagram “{title}”: node “{nid}” kind “{kd}” — for a {typ} use {' | '.join(NODE_KINDS[typ])}.")
            node["kind"] = kd
        actor = n.get("actor") or n.get("lane") or n.get("who")
        if actor:
            node["actor"] = _txt(actor, 40, "Actor", clean)
        if n.get("group"):
            node["group"] = _txt(n["group"], 40, "Group", clean)
        for key in ("at", "order", "size"):
            if isinstance(n.get(key), (int, float)) and not isinstance(n.get(key), bool):
                node[key] = float(n[key]) if key != "size" else int(max(1, min(3, n[key])))
        if "at" in node:
            node["at"] = round(max(0.0, min(60.0, node["at"])), 2)
        nodes.append(node)
        for t in ([n["next"]] if isinstance(n.get("next"), (str, dict)) else n.get("next") or []):   # „next" rövidítés
            raw_edges.append({"from": nid, **({"to": t} if isinstance(t, str) else t)})
        if n.get("branches"):                          # csúcson belüli ágak (a javításban is: „bontsd két esetre")
            sub_n: list[dict] = []
            _expand_branches(nid, n["branches"], sub_n, raw_edges, used | {x["id"] for x in nodes})
            raw_nodes[i + 1:i + 1] = sub_n
            used |= {x["id"] for x in sub_n}
    if typ == "network" and not nodes:                 # csak élek: a végpontok nevei a csúcsok
        for e in raw_edges:
            for end in (e.get("from"), e.get("to")):
                if end and not any(x["text"].lower() == str(end).strip().lower() for x in nodes):
                    nodes.append({"id": f"n{len(nodes) + 1}", "text": _txt(end, 70, "Diagram text", clean)})
    if not nodes:
        raise ValueError(f"diagram “{title}”: no nodes — give nodes: [{{id, text}}] and edges: [{{from, to}}]"
                         + (", or root + branches" if typ == "decision_tree" else ", or steps" if typ == "swimlane" else "") + ".")
    if len(nodes) > LIMITS[typ]:
        raise ValueError(f"diagram “{title}”: {len(nodes)} nodes — at most {LIMITS[typ]} stay readable in one scene. "
                         "Split it into two scenes (e.g. one per main branch).")
    ids = [n["id"] for n in nodes]

    def ref(x: Any, where: str) -> str:
        r = _id(x)
        if r in ids:
            return r
        hit = [n["id"] for n in nodes if n["text"].lower() == str(x or "").strip().lower()]
        if hit:
            return hit[0]
        raise ValueError(f"diagram “{title}”: {where} points to “{x}”, which is not a node (node ids: {', '.join(ids)}) — "
                         "fix the id or add the node.")
    edges: list[dict[str, Any]] = []
    for j, e in enumerate(raw_edges, 1):
        if e.get("remove"):
            continue
        where = f"edge {j}"
        a, b = ref(e.get("from"), f"{where} (from)"), ref(e.get("to"), f"{where} (to)")
        if a == b:
            raise ValueError(f"diagram “{title}”: {where} goes from “{a}” to itself — an edge joins two different nodes.")
        ed: dict[str, Any] = {"from": a, "to": b}
        if e.get("id"):
            ed["id"] = _id(e["id"])
        if e.get("kind") not in (None, ""):
            ed["kind"] = _kind(e["kind"], where)
        if e.get("certainty") not in (None, "", "certain"):
            ed["certainty"] = _cert(e["certainty"], where)
        lab = _txt(e.get("label") or e.get("amount"), 30, "Edge label", clean)
        if lab:
            ed["label"] = lab
        edges.append(ed)
    if typ == "swimlane" and not edges and len(nodes) > 1:     # nincs él: a lépések sorrendje a folyamat
        seq = sorted(range(len(nodes)), key=lambda i: (nodes[i].get("order", i), i))
        edges = [{"from": nodes[seq[i]]["id"], "to": nodes[seq[i + 1]]["id"]} for i in range(len(seq) - 1)]
    if len(edges) > MAX_EDGES:
        raise ValueError(f"diagram “{title}”: {len(edges)} edges — at most {MAX_EDGES} stay readable in one scene; split it.")
    if typ == "swimlane":
        lanes = []
        for ln in c.get("lanes") or []:
            nm = ln.get("name") if isinstance(ln, dict) else ln
            nm = _txt(nm, 40, "Lane", clean)
            if nm and nm not in lanes:
                lanes.append(nm)
                if isinstance(ln, dict) and (ln.get("color") or ln.get("short")):
                    out["actors"].setdefault(nm, {}).update({k: v for k, v in {"color": _hex(ln.get("color")),
                                                                              "short": _txt(ln.get("short"), 14, "Actor short name", clean)}.items() if v})
        miss = [n["id"] for n in nodes if not n.get("actor")]
        if miss:
            raise ValueError(f"diagram “{title}”: step(s) {', '.join(miss)} have no actor — every step needs actor "
                             f"(the lane it belongs to{': ' + ', '.join(lanes) if lanes else ''}).")
        for n in nodes:
            if n["actor"] not in lanes:
                lanes.append(n["actor"])
        if len(lanes) > 6:
            raise ValueError(f"diagram “{title}”: {len(lanes)} lanes — at most 6 fit; merge actors or split the process.")
        out["lanes"] = lanes
    if typ == "decision_tree":
        _check_tree(title, ids, edges)
    if typ == "network" and out.get("center") and _id(out["center"]) not in ids:
        out["center"] = ref(out["center"], "center")
    if not out["actors"]:
        out.pop("actors")
    out["nodes"], out["edges"] = nodes, edges

    notes: list[str] = []
    if bad is not None:                                 # a számok a forrásból (ugyanaz a próba, mint a grafikán)
        wrong = [n for n in nodes if bad(n["text"])] + [{"text": ln} for ln in out.get("lanes") or [] if bad(ln)]
        if wrong:
            return None, [f"⚠ The diagram “{title}” was left out: “{wrong[0]['text'][:60]}” has a number that isn't in the "
                          "source — put the figure in the source text (or the caption), or leave it out of the box."]
        dropped = [e for e in edges if e.get("label") and bad(e["label"])]
        for e in dropped:
            e.pop("label")
        if dropped:
            notes.append(f"⚠ Diagram “{title}”: left out {len(dropped)} edge label(s) with a number not in the source, "
                         f"e.g. on {dropped[0]['from']} → {dropped[0]['to']}.")
        for k, v in list((out.get("legend") or {}).items()):
            if bad(v):
                out["legend"].pop(k)
    kinds = sorted({e.get("kind") or default_kind(typ) for e in edges}, key=EDGE_KINDS.index)
    what = {"decision_tree": "Decision tree", "swimlane": "Swimlane process", "network": "Network"}[typ]
    certs = [e["certainty"] for e in edges if e.get("certainty")]
    marked = "".join(f"; {certs.count(c)} {c}" for c in ("uncertain", "disputed") if c in certs)
    notes.insert(0, f"{what} “{title or 'diagram'}”: {len(nodes)} boxes, {len(edges)} arrows ({', '.join(kinds) or 'none'}{marked})"
                 + (", no number that isn't in the source." if bad is not None else "."))
    return out, notes


def _check_tree(title: str, ids: list[str], edges: list[dict[str, Any]]) -> None:
    parents: dict[str, list[str]] = {i: [] for i in ids}
    for e in edges:
        parents[e["to"]].append(e["from"])
    two = [i for i, p in parents.items() if len(p) > 1]
    if two:
        raise ValueError(f"diagram “{title}”: node “{two[0]}” has two parents ({', '.join(parents[two[0]])}) — a decision "
                         "tree only branches; repeat the outcome as two nodes, or use a swimlane for a process that joins.")
    roots = [i for i, p in parents.items() if not p]
    if len(roots) != 1:
        raise ValueError(f"diagram “{title}”: a decision tree has exactly one root (a node nothing points to); here "
                         f"{len(roots)}: {', '.join(roots) or 'none (there is a loop)'} — connect them with edges.")
    seen, todo = set(), [roots[0]]
    kids = children(ids, edges)
    while todo:
        x = todo.pop()
        if x in seen:
            raise ValueError(f"diagram “{title}”: the branches loop back at “{x}” — a decision tree has no loops.")
        seen.add(x)
        todo += kids.get(x, [])
    lost = [i for i in ids if i not in seen]
    if lost:
        raise ValueError(f"diagram “{title}”: node(s) {', '.join(lost)} are not reachable from the root — add an edge to them.")


def children(ids: list[str], edges: list[dict[str, Any]]) -> dict[str, list[str]]:
    kids: dict[str, list[str]] = {i: [] for i in ids}
    for e in edges:
        kids.setdefault(e["from"], []).append(e["to"])
    return kids


def default_kind(typ: str) -> str:
    return "link" if typ == "network" else "sequence"


def merge(base: dict[str, Any], patch: dict[str, Any]) -> dict[str, Any]:
    """Célzott javítás egy meglévő ábrán: a csúcsok id szerint, az élek id vagy from+to szerint frissülnek
    ({"remove": true} töröl; az új id / új pár hozzáadódik); a többi kulcs felülír, None = vissza az alapra."""
    out = copy.deepcopy(base)
    for k, v in patch.items():
        if k in ("nodes", "steps", "edges"):
            continue
        if k == "actors" and isinstance(v, dict):
            out["actors"] = {**(out.get("actors") or {}), **v}
            continue
        if v is None:
            out.pop(k, None)
        else:
            out[k] = v
    nodes = [dict(n) for n in out.get("nodes") or []]
    gone: set[str] = set()
    for p in patch.get("nodes") or patch.get("steps") or []:
        if not isinstance(p, dict):
            continue
        hit = next((n for n in nodes if p.get("id") and n.get("id") == _id(p["id"])), None)
        if hit is None:
            nodes.append(dict(p))
        elif p.get("remove"):
            nodes.remove(hit)
            gone.add(hit["id"])
        else:
            for kk, vv in p.items():
                if vv is None:
                    hit.pop(kk, None)
                else:
                    hit[kk] = vv
    edges = [dict(e) for e in out.get("edges") or [] if e.get("from") not in gone and e.get("to") not in gone]
    for p in patch.get("edges") or []:
        if not isinstance(p, dict):
            continue
        hit = next((e for e in edges if (p.get("id") and e.get("id") == _id(p["id"]))
                    or (not p.get("id") and p.get("from") is not None and _id(p.get("from")) == e.get("from")
                        and _id(p.get("to")) == e.get("to"))), None)
        if hit is None:
            if not p.get("remove"):
                edges.append(dict(p))
        elif p.get("remove"):
            edges.remove(hit)
        else:
            for kk, vv in p.items():
                if vv is None:
                    hit.pop(kk, None)
                else:
                    hit[kk] = vv
    out["nodes"], out["edges"] = nodes, edges
    out.pop("steps", None)
    return out


def texts(spec: dict[str, Any]) -> list[str]:
    """Az ábra képernyőn látható szövegei (a lefedettség-számításhoz)."""
    return [n.get("text") or "" for n in spec.get("nodes") or []] + [e.get("label") or "" for e in spec.get("edges") or []] \
        + list(spec.get("lanes") or [])


def describe(spec: dict[str, Any]) -> str:
    """Rövid, id-s leírás a „kérj változtatást” modellnek (hogy célzottan javíthasson)."""
    ns = "; ".join(f"{n['id']} “{n['text'][:40]}”" + (f" ({n['actor']})" if n.get("actor") else "") for n in spec.get("nodes") or [])
    es = "; ".join(f"{e['from']}→{e['to']}" + (f" {e['kind']}" if e.get("kind") else "") + (f" {e['certainty']}" if e.get("certainty") else "")
                   + (f" “{e['label']}”" if e.get("label") else "") for e in spec.get("edges") or [])
    return f"{spec['type']} — nodes: {ns} — edges: {es}"


# ---------- kibontás: lépések és idő ----------

def _rings(spec: dict[str, Any]) -> tuple[str, dict[str, int]]:
    """Háló: a középpont (megadott, különben a legtöbb kapcsolatú) és a gyűrűk (távolság tőle, irány nélkül)."""
    ids = [n["id"] for n in spec["nodes"]]
    deg = {i: 0 for i in ids}
    adj: dict[str, list[str]] = {i: [] for i in ids}
    for e in spec["edges"]:
        deg[e["from"]] += 1
        deg[e["to"]] += 1
        adj[e["from"]].append(e["to"])
        adj[e["to"]].append(e["from"])
    center = spec.get("center") if spec.get("center") in ids else max(ids, key=lambda i: (deg[i], -ids.index(i)))
    ring, todo = {center: 0}, [center]
    while todo:
        x = todo.pop(0)
        for y in adj[x]:
            if y not in ring:
                ring[y] = ring[x] + 1
                todo.append(y)
    last = max(ring.values()) + 1
    for i in ids:
        ring.setdefault(i, last)                      # leszakadt csúcs: a külső gyűrűn
    return center, ring


def _order(spec: dict[str, Any]) -> dict[str, int]:
    """Csúcs → kibontási lépés: fa = gyökér, majd ágról ágra a kimenetelig (mélységi bejárás); sávos folyamat = a
    lépések sorrendje (azonos sorrend = együtt); háló = középpont, majd gyűrűnként."""
    typ, nodes = spec["type"], spec["nodes"]
    ids = [n["id"] for n in nodes]
    if typ == "decision_tree":
        kids = children(ids, spec["edges"])
        root = next(i for i in ids if all(e["to"] != i for e in spec["edges"]))
        out, todo = {}, [root]
        while todo:
            x = todo.pop()
            out[x] = len(out)
            todo += reversed(kids.get(x, []))
        return out
    if typ == "swimlane":
        cols = _columns(spec)
        order = sorted(set(cols.values()))
        return {i: order.index(cols[i]) for i in ids}
    return _rings(spec)[1]


def _columns(spec: dict[str, Any]) -> dict[str, float]:
    return {n["id"]: n.get("order", float(k)) for k, n in enumerate(spec["nodes"])}


def steps_of(spec: dict[str, Any]) -> int:
    return (max(_order(spec).values()) + 1) if spec.get("nodes") else 0


T_START, HOLD = 0.6, 1.8


def seconds_needed(spec: Optional[dict[str, Any]]) -> Optional[float]:
    """A kifejezett tempóhoz (pace) kellő jelenethossz: kezdés + lépések × tempó + a teljes kép nézése. Tempó nélkül
    None — az ábra a jelenet hosszába illeszkedik."""
    if not isinstance(spec, dict) or spec.get("type") not in TYPES or not spec.get("pace") or spec.get("reveal") == "all":
        return None
    return round(T_START + steps_of(spec) * float(spec["pace"]) + HOLD, 1)


def duration_note(spec: dict[str, Any], dur: Any) -> Optional[str]:
    need = seconds_needed(spec)
    if need and isinstance(dur, (int, float)) and dur < need - 0.05:
        pace = max(0.3, (float(dur) - T_START - HOLD) / max(1, steps_of(spec)))
        return (f"⚠ Diagram “{spec.get('title') or ''}”: the reveal at {spec['pace']:g} s per step needs ~{need:g} s, the "
                f"scene is {float(dur):g} s, so it builds up faster ({pace:.1f} s per step) — give the scene a duration "
                f"of {math.ceil(need)} or more, or leave the duration out.")
    return None


def timing(spec: dict[str, Any], dur: Optional[float]) -> tuple[dict[str, float], float, float]:
    """→ (csúcs → a megjelenés ideje, tempó, a kibontás vége). Csúcsonkénti `at` felülír."""
    order = _order(spec)
    n = max(order.values()) + 1
    if spec.get("reveal") == "all":
        return {i: 0.5 for i in order}, 0.0, 0.5
    avail = (float(dur) - T_START - HOLD) if dur else None
    if spec.get("pace"):
        pace = float(spec["pace"])
        if avail is not None and n * pace > avail:
            pace = max(0.3, avail / n)
    else:
        pace = min(1.3, max(0.45, avail / n)) if avail is not None else 0.9
    at = {i: round(T_START + order[i] * pace, 3) for i in order}
    for nd in spec["nodes"]:
        if nd.get("at") is not None:
            at[nd["id"]] = float(nd["at"])
    end = max(at.values()) + max(pace, 0.9)
    return at, pace, round(end, 3)


# ---------- szereplő-regiszter ----------

def _rgb(h: str) -> tuple[float, float, float]:
    h = h.lstrip("#")
    return int(h[0:2], 16) / 255, int(h[2:4], 16) / 255, int(h[4:6], 16) / 255


def _to_hex(c: tuple[float, ...]) -> str:
    return "#" + "".join(f"{max(0, min(255, round(v * 255))):02X}" for v in c[:3])


def mix(a: str, b: str, t: float) -> str:
    """a → b t arányban (t = 0: a, t = 1: b)."""
    x, y = _rgb(a), _rgb(b)
    return _to_hex(tuple(x[i] * (1 - t) + y[i] * t for i in range(3)))


def lum(h: str) -> float:
    c = [v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4 for v in _rgb(h)]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def contrast(a: str, b: str) -> float:
    la, lb = sorted((lum(a), lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def readable(color: str, bg: str, need: float = 4.5) -> str:
    """Ugyanaz az árnyalat, de olvasható a háttéren (szövegnek): világosít / sötétít, amíg a kontraszt megvan."""
    c, toward = color, ("#FFFFFF" if lum(bg) < 0.35 else "#000000")
    for t in (0, 0.15, 0.3, 0.45, 0.6, 0.75):
        c = mix(color, toward, t)
        if contrast(c, bg) >= need:
            return c
    return c


def palette(seed: Optional[str], n: int) -> list[str]:
    """Szereplőszínek az arculat akcentusából: az első az akcentus árnyalata, a többi aranymetszés-lépésnyire a
    színkörön, egységes telítettséggel és világossággal (sötét és világos háttéren is jól elválnak)."""
    h, _l, s = colorsys.rgb_to_hls(*_rgb(_hex(seed) or "#FF5A36"))
    return [_to_hex(colorsys.hls_to_rgb((h + i * 0.381966) % 1.0, 0.58, max(0.55, min(0.78, s)))) for i in range(n)]


def actor_names(spec: dict[str, Any]) -> list[str]:
    names = list(spec.get("lanes") or []) + list((spec.get("actors") or {}).keys())
    for n in spec.get("nodes") or []:
        names += [x for x in (n.get("actor"), n.get("group")) if x]
    return list(dict.fromkeys(names))


def register(spec: dict[str, Any], reg: dict[str, dict[str, Any]], seed: Optional[str]) -> dict[str, dict[str, Any]]:
    """A projekt regiszterébe az ábra szereplői: a megadott szín nyer (és onnantól a projekt színe), az új szereplő a
    következő szabad palettaszínt kapja, a régi megtartja a magáét. → a frissített regiszter."""
    reg = {k: dict(v) for k, v in (reg or {}).items() if isinstance(v, dict)}
    own = spec.get("actors") or {}
    for name in actor_names(spec):
        mine = own.get(name) or {}
        ent = reg.setdefault(name, {})
        if mine.get("color"):
            ent["color"] = mine["color"]
        if mine.get("short"):
            ent["short"] = mine["short"]
        if not ent.get("color"):
            taken = {v.get("color") for v in reg.values()}
            ent["color"] = next(c for c in palette(seed, len(reg) + 12) if c not in taken)
    return reg


def prepare(story: dict[str, Any], accent: Optional[str]) -> None:
    """Videó: a forgatókönyv ábráinak szereplői a projekt regiszterébe (story["actors"], a spec-ben tárolva) — minden
    ábra a regiszter színeit kapja ("_actors", futásidejű kulcs)."""
    reg = story.get("actors") or {}
    found = False
    for sc in story.get("scenes") or []:
        ch = (sc or {}).get("chart") or {}
        if ch.get("type") in TYPES:
            found = True
            reg = register(ch, reg, accent)
            ch["_actors"] = {k: reg[k] for k in actor_names(ch)}
    if found and reg:
        story["actors"] = reg


# ---------- elrendezés (em-egységben; a betűméret szorzója a végén) ----------

def _wrap(text: str, max_em: float, cw: float) -> list[str]:
    """Szótördelés karakterszámra (a video._fit_size mértékével); a sornál hosszabb szó a kötőjelénél is törhet."""
    maxc = max(1, int(max_em / cw))
    toks: list[tuple[str, str]] = []                  # (darab, elválasztó utána: " " | "")
    for w in str(text or "").split():
        parts = re.split(r"(?<=-)(?=\w)", w) if len(w) > maxc else [w]
        toks += [(x, "") for x in parts[:-1]] + [(parts[-1], " ")]
    lines, cur, sep = [], "", ""
    for tok, after in toks:
        if not cur:
            cur = tok
        elif len(cur) + len(sep) + len(tok) <= maxc:
            cur += sep + tok
        else:
            lines.append(cur)
            cur = tok
        sep = after
    return lines + [cur] if cur else (lines or [""])


def _measure(text: str, wrap_em: float, cw: float, font: float = 1.0, max_lines: int = 3) -> tuple[list[str], float, float]:
    """Szöveg → (sorok, szélesség em, magasság em) a font szorzójával; ha a sorok száma több a megengedettnél, szélesebben."""
    lines = _wrap(text, wrap_em / font, cw)
    for k in (1.45, 2.0, 2.8):
        if len(lines) <= max_lines:
            break
        lines = _wrap(text, wrap_em * k / font, cw)
    tw = max(len(x) for x in lines) * cw * font
    return lines, tw, len(lines) * LH * font


def _pill(text: str, wrap_em: float, cw: float) -> dict[str, Any]:
    lines, tw, th = _measure(text, wrap_em, cw, LAB, 2)
    return {"lines": lines, "w": tw + 0.9, "h": th + 0.44}


def _side(b: dict[str, Any], side: str) -> list[float]:
    return {"top": [b["x"], b["y"] - b["h"] / 2], "bottom": [b["x"], b["y"] + b["h"] / 2],
            "left": [b["x"] - b["w"] / 2, b["y"]], "right": [b["x"] + b["w"] / 2, b["y"]]}[side]


def _along(pts: list[list[float]], t: float) -> list[float]:
    """A törött vonal ívhossz szerinti t (0–1) pontja."""
    segs = [math.dist(pts[i], pts[i + 1]) for i in range(len(pts) - 1)]
    goal, acc = sum(segs) * t, 0.0
    for i, L in enumerate(segs):
        if acc + L >= goal - 1e-9:
            f = (goal - acc) / L if L else 0
            return [pts[i][0] + (pts[i + 1][0] - pts[i][0]) * f, pts[i][1] + (pts[i + 1][1] - pts[i][1]) * f]
        acc += L
    return list(pts[-1])


def _overlap(a: tuple, b: tuple, pad: float = 0.0) -> float:
    """Két doboz (cx, cy, w, h) átfedésének területe (pad: ennyivel nagyobbnak véve)."""
    ox = (a[2] + b[2]) / 2 + pad - abs(a[0] - b[0])
    oy = (a[3] + b[3]) / 2 + pad - abs(a[1] - b[1])
    return ox * oy if ox > 0 and oy > 0 else 0.0


def _seg_hits(p: list[float], q: list[float], b: dict[str, Any], pad: float = 0.25) -> bool:
    """Átmegy-e a p–q szakasz a dobozon (Liang–Barsky)?"""
    x0, x1 = b["x"] - b["w"] / 2 - pad, b["x"] + b["w"] / 2 + pad
    y0, y1 = b["y"] - b["h"] / 2 - pad, b["y"] + b["h"] / 2 + pad
    dx, dy = q[0] - p[0], q[1] - p[1]
    t0, t1 = 0.0, 1.0
    for pp, qq in ((-dx, p[0] - x0), (dx, x1 - p[0]), (-dy, p[1] - y0), (dy, y1 - p[1])):
        if pp == 0:
            if qq < 0:
                return False
            continue
        r = qq / pp
        if pp < 0:
            t0 = max(t0, r)
        else:
            t1 = min(t1, r)
        if t0 > t1:
            return False
    return t1 - t0 > 0.02


def _node_style(spec: dict[str, Any], n: dict[str, Any], kids: dict[str, list[str]], center: Optional[str]) -> tuple[str, bool, float]:
    """→ (alak: round | pill | hex, félkövér-e, betűszorzó) a csúcs fajtájából."""
    typ = spec["type"]
    kd = n.get("kind")
    if typ == "decision_tree":
        kd = kd or ("question" if all(e["to"] != n["id"] for e in spec["edges"]) else "outcome" if not kids.get(n["id"]) else "condition")
        return ("pill" if kd == "outcome" else "round"), kd in ("question", "outcome"), 1.0
    if typ == "swimlane":
        kd = kd or "step"
        return {"start": "pill", "end": "pill", "decision": "hex", "step": "round"}[kd], kd in ("start", "end", "decision"), 1.0
    size = {1: 0.9, 2: 1.0, 3: 1.15}.get(n.get("size") or (3 if n["id"] == center else 2), 1.0)
    return "round", n["id"] == center or (n.get("size") or 0) >= 3, size


def _kind_of(spec: dict[str, Any], n: dict[str, Any]) -> str:
    if spec["type"] == "decision_tree":
        if n.get("kind"):
            return n["kind"]
        if all(e["to"] != n["id"] for e in spec["edges"]):
            return "question"
        return "condition" if any(e["from"] == n["id"] for e in spec["edges"]) else "outcome"
    return n.get("kind") or ("step" if spec["type"] == "swimlane" else "node")


def _boxes(spec: dict[str, Any], wrap_em: float, cw: float, center: Optional[str] = None) -> dict[str, dict[str, Any]]:
    kids = children([n["id"] for n in spec["nodes"]], spec["edges"])
    out = {}
    for n in spec["nodes"]:
        shape, bold, font = _node_style(spec, n, kids, center)
        lines, tw, th = _measure(n["text"], wrap_em, cw * (1.06 if bold else 1.0), font, 4)
        out[n["id"]] = {"id": n["id"], "lines": lines, "tw": tw, "w": tw + 2 * PX + (1.2 if shape == "hex" else 0),
                        "h": th + 2 * PY, "shape": shape, "bold": bold, "font": font, "kind": _kind_of(spec, n),
                        "actor": n.get("group") or n.get("actor")}
    return out


def _rewrap(B: dict[str, dict[str, Any]], W: float, cw: float) -> None:
    """Egységes dobozszélességnél a szöveg a teljes belső szélességet kapja (kevesebb sor, alacsonyabb doboz)."""
    for b in B.values():
        cwb = cw * (1.06 if b["bold"] else 1.0)
        lines = _wrap(" ".join(b["lines"]), (W - 2 * PX - (1.2 if b["shape"] == "hex" else 0)) / b["font"], cwb)
        if len(lines) <= len(b["lines"]):
            b["lines"] = lines
            b["tw"] = max(len(x) for x in lines) * cwb * b["font"]
            b["h"] = len(lines) * LH * b["font"] + 2 * PY
        b["w"] = W


def _edge(spec: dict[str, Any], e: dict[str, Any], pts: list[list[float]], pill: Optional[dict[str, Any]],
          lab_at: Optional[list[float]]) -> dict[str, Any]:
    kind = e.get("kind") or default_kind(spec["type"])
    out = {"from": e["from"], "to": e["to"], "kind": kind, "certainty": e.get("certainty") or "certain", "pts": pts}
    if pill and lab_at:
        out["label"] = {**pill, "x": lab_at[0], "y": lab_at[1]}
    return out


def _tree(spec: dict[str, Any], horizontal: bool, wrap_em: float, cw: float) -> dict[str, Any]:
    B = _boxes(spec, wrap_em, cw)
    ids = [n["id"] for n in spec["nodes"]]
    kids = children(ids, spec["edges"])
    root = next(i for i in ids if all(e["to"] != i for e in spec["edges"]))
    depth, leaves, todo = {root: 0}, [], [root]
    while todo:
        x = todo.pop()
        if not kids[x]:
            leaves.append(x)
        for k in reversed(kids[x]):
            depth[k] = depth[x] + 1
            todo.append(k)
    W = max(max(b["w"] for b in B.values()), 5.5)
    _rewrap(B, W, cw)
    gap_s = 1.1
    pitch = W + gap_s if not horizontal else max(b["h"] for b in B.values()) + 0.9
    pills = {(e["from"], e["to"]): _pill(e["label"], (pitch - 1.0) if not horizontal else 10.0, cw)
             for e in spec["edges"] if e.get("label")}
    lab_h = max([p["h"] for p in pills.values()] + [0])
    lab_w = max([p["w"] for p in pills.values()] + [0])
    pos: dict[str, float] = {}
    for i, lf in enumerate(leaves):
        pos[lf] = i * pitch

    def place(x: str) -> float:
        if x not in pos:
            ks = [place(k) for k in kids[x]]
            pos[x] = (ks[0] + ks[-1]) / 2
        return pos[x]
    place(root)
    levels = max(depth.values()) + 1
    if not horizontal:
        lvl_h = [max(B[i]["h"] for i in ids if depth[i] == d) for d in range(levels)]
        gap_y = max(2.4, lab_h + 1.7)
        ys = [sum(lvl_h[:d]) + gap_y * d + lvl_h[d] / 2 for d in range(levels)]
        for i in ids:
            B[i]["x"], B[i]["y"] = pos[i], ys[depth[i]]
    else:
        gap_x = max(3.2, lab_w + 1.8)
        for i in ids:
            B[i]["x"], B[i]["y"] = depth[i] * (W + gap_x), pos[i]
    edges = []
    for e in spec["edges"]:
        a, b = B[e["from"]], B[e["to"]]
        p, q = (_side(a, "bottom"), _side(b, "top")) if not horizontal else (_side(a, "right"), _side(b, "left"))
        pill, at = pills.get((e["from"], e["to"])), None
        if pill:                                        # a címke a gyerek közelében, a vonalon
            if not horizontal:
                yy = q[1] - 0.35 - pill["h"] / 2
                f = (yy - p[1]) / ((q[1] - p[1]) or 1)
                at = [p[0] + (q[0] - p[0]) * f, yy]
            else:
                xx = q[0] - 0.35 - pill["w"] / 2
                f = (xx - p[0]) / ((q[0] - p[0]) or 1)
                at = [xx, p[1] + (q[1] - p[1]) * f]
        edges.append(_edge(spec, e, [p, q], pill, at))
    return {"nodes": B, "edges": edges, "lanes": [], "dir": "right" if horizontal else "down"}


def _outline(spec: dict[str, Any], wrap_em: float, cw: float) -> dict[str, Any]:
    """Vázlat-fa (álló képkockára): minden csúcs a saját sorában, mélységi sorrendben, a mélységgel beljebb; a szülő
    bal oldalán lefutó gerincből jobbra ágazik az él, a címke a bevezető szakaszon ül."""
    B = _boxes(spec, wrap_em, cw)
    ids = [n["id"] for n in spec["nodes"]]
    kids = children(ids, spec["edges"])
    root = next(i for i in ids if all(e["to"] != i for e in spec["edges"]))
    rows, depth, todo = [], {root: 0}, [root]
    while todo:
        x = todo.pop()
        rows.append(x)
        for k in reversed(kids[x]):
            depth[k] = depth[x] + 1
            todo.append(k)
    W = max(max(b["w"] for b in B.values()), 5.5)
    _rewrap(B, W, cw)
    pills = {(e["from"], e["to"]): _pill(e["label"], 7.0, cw) for e in spec["edges"] if e.get("label")}
    levels = max(depth.values()) + 1
    stub = [0.0] + [max([pills[(e["from"], e["to"])]["w"] + 1.0 for e in spec["edges"]
                         if (e["from"], e["to"]) in pills and depth[e["to"]] == d] + [1.6]) for d in range(1, levels)]
    left = [0.0]
    for d in range(1, levels):
        left.append(left[-1] + 1.2 + stub[d])
    y = 0.0
    for r in rows:
        b = B[r]
        pl = next((pills[(e["from"], e["to"])] for e in spec["edges"] if e["to"] == r and (e["from"], e["to"]) in pills), None)
        hh = max(b["h"], pl["h"] if pl else 0)
        b["x"], b["y"] = left[depth[r]] + W / 2, y + hh / 2
        y += hh + 0.9
    edges = []
    for e in spec["edges"]:
        a, c = B[e["from"]], B[e["to"]]
        xs = a["x"] - W / 2 + 1.2
        pts = [[xs, a["y"] + a["h"] / 2], [xs, c["y"]], [c["x"] - W / 2, c["y"]]]
        pill = pills.get((e["from"], e["to"]))
        at = [(xs + c["x"] - W / 2) / 2, c["y"]] if pill else None
        edges.append(_edge(spec, e, pts, pill, at))
    return {"nodes": B, "edges": edges, "lanes": [], "dir": "outline"}


def _swimlane(spec: dict[str, Any], horizontal: bool, wrap_em: float, cw: float, actors: dict[str, Any]) -> dict[str, Any]:
    """Sávos folyamat: szereplőnként egy sáv; az idő balról jobbra (a sávok sorok) vagy fentről le (a sávok oszlopok).
    Azonos sorrendű lépések egy oszlopban; ha egy sávban több is van, egymás után kapnak helyet."""
    B = _boxes(spec, wrap_em, cw)
    lanes = spec["lanes"]
    cols = _columns(spec)
    order = sorted(set(cols.values()))
    span = {c: max([sum(1 for n in spec["nodes"] if cols[n["id"]] == c and n["actor"] == ln) for ln in lanes] + [1]) for c in order}
    base, acc = {}, 0
    for c in order:
        base[c], acc = acc, acc + span[c]
    slot, seen = {}, {}
    for n in spec["nodes"]:
        key = (cols[n["id"]], n["actor"])
        slot[n["id"]] = base[cols[n["id"]]] + seen.get(key, 0)
        seen[key] = seen.get(key, 0) + 1
    nslots = acc
    W = max(max(b["w"] for b in B.values()), 5.5)
    _rewrap(B, W, cw)
    H = max(b["h"] for b in B.values())
    for b in B.values():
        b["h"] = H
    lane_of = {n["id"]: lanes.index(n["actor"]) for n in spec["nodes"]}
    pills = {(e["from"], e["to"]): _pill(e["label"], 6.0, cw) for e in spec["edges"] if e.get("label")}
    lab_w = max([p["w"] for p in pills.values()] + [0])
    lab_h = max([p["h"] for p in pills.values()] + [0])
    heads = []
    for ln in lanes:
        nm = ln
        lines, tw, th = _measure(nm, 7.0 if horizontal else W + 0.6, cw * 1.06, HEAD, 2)
        if len(_wrap(nm, (7.0 if horizontal else W + 0.6) / HEAD, cw * 1.06)) > 2 and (actors.get(ln) or {}).get("short"):
            lines, tw, th = _measure(actors[ln]["short"], 7.0 if horizontal else W + 0.6, cw * 1.06, HEAD, 2)
        heads.append((lines, tw, th))
    lanes_out = []
    # visszalépés címkével (fekvőben a sáv alján fut vissza): a sávok alján ennyi hely kell neki
    back_lab = max([pills[(e["from"], e["to"])]["h"] for e in spec["edges"]
                    if (e["from"], e["to"]) in pills and slot[e["to"]] < slot[e["from"]]] + [0.0])
    if horizontal:
        head_w = max(max(h[1] for h in heads) + 1.6, 4.0)
        # résenként annyi hely, amennyi az oda kerülő címkének kell (a sávváltó él címkéje a cél előtti résben, a
        # sávon belüli szomszédoké a kettejük közöttiben) — a címke nélküli rések keskenyek maradnak
        gaps = [1.9] * max(0, nslots - 1)
        for e in spec["edges"]:
            pl = pills.get((e["from"], e["to"]))
            sa_, sb_ = slot[e["from"]], slot[e["to"]]
            if pl and sb_ > sa_ and (lane_of[e["from"]] != lane_of[e["to"]] or sb_ == sa_ + 1):
                gaps[sb_ - 1] = max(gaps[sb_ - 1], pl["w"] + 1.2)
        gap = max(gaps + [1.9])
        extra = back_lab + 0.2 if back_lab else 0.0
        lane_h = H + 1.5 + extra
        x0 = head_w + 0.9
        xs = [x0 + W / 2]
        for g in gaps:
            xs.append(xs[-1] + W + g)
        for i in B:
            B[i]["x"], B[i]["y"] = xs[slot[i]], lane_of[i] * lane_h + 0.75 + H / 2
        total_w = xs[-1] + W / 2 + 0.9
        for k, ln in enumerate(lanes):
            lanes_out.append({"name": ln, "lines": heads[k][0], "band": [total_w / 2, k * lane_h + lane_h / 2, total_w, lane_h],
                              "head": [head_w / 2, k * lane_h + lane_h / 2, head_w, lane_h]})
    else:
        lane_w = W + 1.4
        head_h = max(h[2] for h in heads) + 1.0
        gap = max(1.8, lab_h + 1.1)
        pitch = H + gap
        y0 = head_h + 0.8
        for i in B:
            B[i]["x"], B[i]["y"] = lane_of[i] * lane_w + lane_w / 2, y0 + slot[i] * pitch + H / 2
        total_h = y0 + nslots * pitch - gap + 0.8
        for k, ln in enumerate(lanes):
            lanes_out.append({"name": ln, "lines": heads[k][0], "band": [k * lane_w + lane_w / 2, total_h / 2, lane_w, total_h],
                              "head": [k * lane_w + lane_w / 2, head_h / 2, lane_w, head_h]})

    def blocked(pts: list[list[float]], a: str, b: str) -> bool:
        return any(_seg_hits(pts[j], pts[j + 1], B[o]) for j in range(len(pts) - 1) for o in B if o not in (a, b))
    edges = []
    for e in spec["edges"]:
        a, b = e["from"], e["to"]
        A, Bb = B[a], B[b]
        sa, sb, la, lb = slot[a], slot[b], lane_of[a], lane_of[b]
        if horizontal:
            fwd, back, along, across = "right", "left", 0, 1
        else:
            fwd, back, along, across = "bottom", "top", 1, 0
        cands: list[list[list[float]]] = []
        if sb > sa:
            p, q = _side(A, fwd), _side(Bb, back)
            if la == lb:
                cands.append([p, q])
                off = (H / 2 + 0.45) if horizontal else (W / 2 + 0.45)    # kerülő a sáv szélén, ha útban van egy lépés
                side_a, side_b = (_side(A, "top"), _side(Bb, "top")) if horizontal else (_side(A, "right"), _side(Bb, "right"))
                if horizontal:
                    cands.append([side_a, [A["x"], A["y"] - off], [Bb["x"], Bb["y"] - off], side_b])
                else:
                    cands.append([side_a, [A["x"] + off, A["y"]], [Bb["x"] + off, Bb["y"]], side_b])
            else:
                g_b = (gaps[sb - 1] if horizontal else gap) / 2
                g_a = (gaps[sa] if horizontal else gap) / 2
                for m in (q[along] - g_b, p[along] + g_a):
                    if horizontal:
                        cands.append([p, [m, p[1]], [m, q[1]], q])
                    else:
                        cands.append([p, [p[0], m], [q[0], m], q])
        elif sb == sa:                                  # párhuzamos lépések (azonos idő, más sáv)
            if horizontal:
                cands.append([_side(A, "bottom" if lb > la else "top"), _side(Bb, "top" if lb > la else "bottom")])
            else:
                cands.append([_side(A, "right" if lb > la else "left"), _side(Bb, "left" if lb > la else "right")])
        else:                                           # visszalépés (pl. „visszaküldi”): a sávok szélén vissza
            if horizontal:                              # alulról be a cél aljába; ha útban van valami, oldalról
                yb = A["y"] + H / 2 + 0.45 + back_lab / 2
                xg = Bb["x"] - W / 2 - (gaps[sb - 1] if sb > 0 else 2.2) * 0.35
                cands.append([_side(A, "bottom"), [A["x"], yb], [Bb["x"], yb], _side(Bb, "bottom" if Bb["y"] < yb else "top")])
                cands.append([_side(A, "bottom"), [A["x"], yb], [xg, yb], [xg, Bb["y"]], _side(Bb, "left")])
            else:
                xr = A["x"] + W / 2 + 0.45
                yg = Bb["y"] - H / 2 - gap * 0.4
                cands.append([_side(A, "right"), [xr, A["y"]], [xr, Bb["y"]], _side(Bb, "right" if Bb["x"] < xr else "left")])
                cands.append([_side(A, "right"), [xr, A["y"]], [xr, yg], [Bb["x"], yg], _side(Bb, "top")])
        pts = next((c for c in cands if not blocked(c, a, b)), cands[0])
        pill = pills.get((a, b))
        at = None
        if pill and len(pts) == 2 and la == lb:         # sávon belüli egyenes: a vonal MELLÉ, a testvér-élekkel szemközt
            mid = [(pts[0][0] + pts[1][0]) / 2, (pts[0][1] + pts[1][1]) / 2]
            sib = [lane_of[x["to"]] for x in spec["edges"] if x["from"] == a and lane_of[x["to"]] != la]
            flip = bool(sib) and all(v < la for v in sib)          # a testvérek fölfelé / balra ágaznak: mi a másik oldalra
            if horizontal:
                at = [mid[0], mid[1] + (1 if flip else -1) * (0.2 + pill["h"] / 2)]
            else:
                at = [mid[0] + (-1 if (sib and not flip) else 1) * (0.25 + pill["w"] / 2), mid[1]]
        elif pill:                                      # a címke a leghosszabb szakasz közepén, a vonalon
            j = max(range(len(pts) - 1), key=lambda i: math.dist(pts[i], pts[i + 1]) + (0.01 * i))
            if sb < sa:                                 # visszalépés: a visszafutó szakaszon (ott van neki hely)
                j = 1 if horizontal else 2
            at = [(pts[j][0] + pts[j + 1][0]) / 2, (pts[j][1] + pts[j + 1][1]) / 2]
        edges.append(_edge(spec, e, pts, pill, at))
    return {"nodes": B, "edges": edges, "lanes": lanes_out, "dir": "right" if horizontal else "down"}


def _network(spec: dict[str, Any], aspect: float, wrap_em: float, cw: float) -> dict[str, Any]:
    """Kapcsolati háló: a középpont középen, körülötte gyűrűk (ellipszis a képkocka arányában); a gyűrű csúcsai a belső
    szomszédjaik irányában; a végén az átfedések szétlökése — mind determinisztikus (nincs véletlen)."""
    center, ring = _rings(spec)
    B = _boxes(spec, wrap_em, cw, center)
    ids = [n["id"] for n in spec["nodes"]]
    adj: dict[str, list[str]] = {i: [] for i in ids}
    for e in spec["edges"]:
        adj[e["from"]].append(e["to"])
        adj[e["to"]].append(e["from"])
    a = max(0.7, min(2.6, aspect))
    ang: dict[str, float] = {center: 0.0}
    B[center]["x"] = B[center]["y"] = 0.0
    rx_prev = ry_prev = 0.0
    prev = [center]
    for k in range(1, max(ring.values()) + 1):
        ring_k = [i for i in ids if ring[i] == k]
        if not ring_k:
            continue
        wmax_prev, hmax_prev = max(B[i]["w"] for i in prev), max(B[i]["h"] for i in prev)
        wmax, hmax = max(B[i]["w"] for i in ring_k), max(B[i]["h"] for i in ring_k)
        need = sum(math.hypot(B[i]["w"], B[i]["h"]) * 0.8 + 1.4 for i in ring_k)
        ry = max(ry_prev + (hmax_prev + hmax) / 2 + 1.6, (rx_prev + (wmax_prev + wmax) / 2 + 2.0) / a)
        while True:
            rx = a * ry
            per = math.pi * (3 * (rx + ry) - math.sqrt((3 * rx + ry) * (rx + 3 * ry)))
            if per >= need:
                break
            ry *= 1.08
        start = 0.0 if a > 1.15 else -math.pi / 2
        want = {}
        for j, i in enumerate(ring_k):
            inner = [ang[x] for x in adj[i] if x in ang and ring[x] < k]
            if inner and k > 1:
                want[i] = math.atan2(sum(math.sin(t) for t in inner), sum(math.cos(t) for t in inner))
            else:
                want[i] = start + 2 * math.pi * j / len(ring_k)
        seq = sorted(ring_k, key=lambda i: (want[i] % (2 * math.pi), ids.index(i)))
        th = [want[i] for i in seq]
        n_k = len(seq)
        sep = 2 * math.pi / n_k * (0.92 if k == 1 else 0.85)
        if k > 1:
            sep = min(sep, max(0.38, max(math.hypot(B[i]["w"], B[i]["h"]) for i in seq) / max(rx, ry) * 1.1))
        for _ in range(80):                            # szomszédos csúcsok szögtávolsága legalább sep
            moved = False
            for j in range(n_k if n_k > 1 else 0):
                j2 = (j + 1) % n_k
                gap = (th[j2] - th[j]) % (2 * math.pi) if n_k > 1 else 2 * math.pi
                if gap < sep - 1e-6:
                    d = (sep - gap) / 2
                    th[j] -= d
                    th[j2] += d
                    moved = True
            if not moved:
                break
        for i, t in zip(seq, th):
            ang[i] = t
            B[i]["x"], B[i]["y"] = rx * math.cos(t), ry * math.sin(t)
        rx_prev, ry_prev, prev = rx, ry, ring_k
    room = {}
    for e in spec["edges"]:                          # címkés él: a két doboz között férjen el a címke
        if e.get("label"):
            pl = _pill(e["label"], 8.0, cw)
            key = tuple(sorted((e["from"], e["to"])))
            room[key] = (max(room.get(key, (0, 0))[0], pl["w"] + 2.2), max(room.get(key, (0, 0))[1], pl["h"] + 1.6))
    _separate(B, ids, room=room)
    return {"nodes": B, "edges": _net_edges(spec, B, cw), "lanes": [], "dir": "ring"}


def _separate(B: dict[str, dict[str, Any]], ids: list[str], mx: float = 1.4, my: float = 1.0,
              room: Optional[dict[tuple, tuple]] = None) -> None:
    """Átfedő dobozok szétlökése a kisebb behatolás irányában (determinisztikus); room: páronkénti nagyobb távolság."""
    room = room or {}
    for _ in range(300):
        moved = False
        for i in range(len(ids)):
            for j in range(i + 1, len(ids)):
                a, b = B[ids[i]], B[ids[j]]
                dx, dy = b["x"] - a["x"], b["y"] - a["y"]
                rx_, ry_ = room.get(tuple(sorted((ids[i], ids[j]))), (mx, my))
                ox = (a["w"] + b["w"]) / 2 + max(mx, rx_) - abs(dx)
                oy = (a["h"] + b["h"]) / 2 + max(my, ry_) - abs(dy)
                if ox > 0 and oy > 0:
                    moved = True
                    if ox / (a["w"] + b["w"]) <= oy / (a["h"] + b["h"]):
                        s = 1 if dx >= 0 else -1
                        a["x"] -= s * ox / 2
                        b["x"] += s * ox / 2
                    else:
                        s = 1 if dy >= 0 else -1
                        a["y"] -= s * oy / 2
                        b["y"] += s * oy / 2
        if not moved:
            return


def _net_edges(spec: dict[str, Any], B: dict[str, dict[str, Any]], cw: float) -> list[dict[str, Any]]:
    """Háló élei: a két doboz egymás felé néző oldalának közepétől; ha egy harmadik dobozon menne át, megtörik."""
    edges, labels = [], []
    obst = [(b["x"], b["y"], b["w"], b["h"]) for b in B.values()]
    for e in spec["edges"]:
        A, Bb = B[e["from"]], B[e["to"]]
        dx, dy = Bb["x"] - A["x"], Bb["y"] - A["y"]
        if abs(dx) / ((A["w"] + Bb["w"]) / 2) >= abs(dy) / ((A["h"] + Bb["h"]) / 2):
            p, q = (_side(A, "right"), _side(Bb, "left")) if dx > 0 else (_side(A, "left"), _side(Bb, "right"))
        else:
            p, q = (_side(A, "bottom"), _side(Bb, "top")) if dy > 0 else (_side(A, "top"), _side(Bb, "bottom"))
        others = [o for k, o in B.items() if k not in (e["from"], e["to"])]
        pts = [p, q]
        if any(_seg_hits(p, q, o) for o in others):
            L = math.dist(p, q) or 1
            nx, ny = -(q[1] - p[1]) / L, (q[0] - p[0]) / L
            mid = [(p[0] + q[0]) / 2, (p[1] + q[1]) / 2]
            for off in (2.0, -2.0, 3.5, -3.5, 5.0, -5.0):
                c = [mid[0] + nx * off, mid[1] + ny * off]
                if not any(_seg_hits(p, c, o) or _seg_hits(c, q, o) for o in others):
                    pts = [p, c, q]
                    break
        pill, at = None, None
        if e.get("label"):
            pill = _pill(e["label"], 8.0, cw)
            at = _label_spot(pts, pill, obst, labels)
            labels.append((at[0], at[1], pill["w"], pill["h"]))
        edges.append(_edge(spec, e, pts, pill, at))
    return edges


def _label_spot(pts: list[list[float]], pill: dict[str, Any], obst: list[tuple], labels: list[tuple]) -> list[float]:
    """A címke helye az élen: a vonalon, vagy közvetlenül mellette — ahol a legkevésbé takar dobozt, másik címkét
    és a két végpontot (nyílhegy, pötty)."""
    ends = [(p[0], p[1], 1.6, 1.6) for p in (pts[0], pts[-1])]
    best = None
    for t in (0.5, 0.42, 0.58, 0.34, 0.66):
        c = _along(pts, t)
        a, b = _along(pts, max(0.0, t - 0.02)), _along(pts, min(1.0, t + 0.02))
        L = math.dist(a, b) or 1
        nx, ny = -(b[1] - a[1]) / L, (b[0] - a[0]) / L
        reach = abs(nx) * pill["w"] / 2 + abs(ny) * pill["h"] / 2 + 0.25
        for off in (0.0, reach, -reach):
            q = [c[0] + nx * off, c[1] + ny * off]
            box = (q[0], q[1], pill["w"], pill["h"])
            cost = (10 * sum(_overlap(box, o, 0.2) for o in obst) + 3 * sum(_overlap(box, lb, 0.2) for lb in labels)
                    + 4 * sum(_overlap(box, z) for z in ends) + abs(t - 0.5) + (0.15 if off else 0.0))
            if best is None or cost < best[0] - 1e-9:
                best = (cost, q)
    return best[1]


def _legend_items(spec: dict[str, Any], lay: dict[str, Any], lang: str, actors: dict[str, Any]) -> list[dict[str, Any]]:
    """A használt élfajták (és a bizonytalan / vitatott jelölés) — csak a magától értetődő egyfajta „ezután” nyíl
    nincs a jelmagyarázatban. A háló csoportjai színmintával."""
    words = {**LEGEND["hu" if lang == "hu" else "en"], **(spec.get("legend") or {})}
    kinds = [k for k in EDGE_KINDS if any(e["kind"] == k for e in lay["edges"])]
    certs = [c for c in ("uncertain", "disputed") if any(e["certainty"] == c for e in lay["edges"])]
    items = []
    if not (kinds == ["sequence"] and not certs):
        items += [{"what": "kind", "key": k, "text": words[k]} for k in kinds]
    items += [{"what": "cert", "key": c, "text": words[c]} for c in certs]
    groups = list(dict.fromkeys(n["group"] for n in spec["nodes"] if n.get("group")))
    items += [{"what": "group", "key": g, "text": g, "color": (actors.get(g) or {}).get("color")} for g in groups]
    return items


def _bbox(lay: dict[str, Any]) -> list[float]:
    xs, ys = [], []
    for b in lay["nodes"].values():
        xs += [b["x"] - b["w"] / 2, b["x"] + b["w"] / 2]
        ys += [b["y"] - b["h"] / 2, b["y"] + b["h"] / 2]
    for e in lay["edges"]:
        xs += [p[0] for p in e["pts"]]
        ys += [p[1] for p in e["pts"]]
        if e.get("label"):
            lb = e["label"]
            xs += [lb["x"] - lb["w"] / 2, lb["x"] + lb["w"] / 2]
            ys += [lb["y"] - lb["h"] / 2, lb["y"] + lb["h"] / 2]
    for ln in lay["lanes"]:
        for bx in (ln["band"], ln["head"]):
            xs += [bx[0] - bx[2] / 2, bx[0] + bx[2] / 2]
            ys += [bx[1] - bx[3] / 2, bx[1] + bx[3] / 2]
    return [min(xs), min(ys), max(xs), max(ys)]


def _legend_rows(items: list[dict[str, Any]], max_w: float, cw: float) -> tuple[list[list[dict[str, Any]]], float]:
    """Jelmagyarázat-sorok (em): minta (2,6 em) + szöveg; annyi sor, amennyi a szélességbe kell."""
    rows: list[list[dict[str, Any]]] = [[]]
    cur = 0.0
    for it in items:
        it["w"] = (1.2 if it["what"] == "group" else 2.6) + 0.5 + len(it["text"]) * cw * LEG
        add = it["w"] + (1.6 if rows[-1] else 0)
        if rows[-1] and cur + add > max_w:
            rows.append([])
            cur, add = 0.0, it["w"]
        rows[-1].append(it)
        cur += add
    rows = [r for r in rows if r]
    widest = max([sum(i["w"] for i in r) + 1.6 * (len(r) - 1) for r in rows] + [0])
    return rows, widest


def layout(spec: dict[str, Any], region: tuple[float, float], *, cw: float, s_max: float, s_min: float,
           lang: str = "en", actors: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    """A legjobb elrendezés a régióba (w, h — képpont vagy pont): irány és tördelési szélesség jelöltjei közül az,
    amelyik a legnagyobb betűt engedi. → {nodes, edges, lanes, legend, s (egység/em), ox, oy, dir}, minden
    koordináta a régió egységében (a régió bal felső sarka a 0, 0)."""
    rw, rh = region
    actors = actors or {}
    aspect = rw / rh
    dirs: list[Any] = {"down": [False], "right": [True], "outline": ["outline"], "auto": [False, True]}[spec.get("direction") or "auto"]
    if spec["type"] == "decision_tree" and (spec.get("direction") or "auto") == "auto":
        dirs.append("outline")                        # álló képkockán a vázlat-fa fér el a legjobban
    if spec["type"] == "network":
        dirs = [None]
    best = None
    for horiz in dirs:
        for wrap_em in (5.5, 7.0, 8.5, 11.0, 14.0):
            if spec["type"] == "decision_tree":
                lay = _outline(spec, wrap_em, cw) if horiz == "outline" else _tree(spec, bool(horiz), wrap_em, cw)
            elif spec["type"] == "swimlane":
                lay = _swimlane(spec, bool(horiz), wrap_em, cw, actors)
            else:
                lay = _network(spec, aspect, wrap_em, cw)
            bb = _bbox(lay)
            bw, bh = bb[2] - bb[0], bb[3] - bb[1]
            items = _legend_items(spec, lay, lang, actors)
            s = min(s_max, rw / bw, rh / bh)
            leg_rows, leg_h = [], 0.0
            if items:
                for _ in range(3):                   # a jelmagyarázat sorai a betűmérettől függnek — két-három kör elég
                    leg_rows, widest = _legend_rows(items, rw / s, cw)
                    leg_h = len(leg_rows) * LH * LEG * 1.45 + 1.2
                    s = min(s_max, rw / bw, rw / max(widest, 1), rh / (bh + leg_h))
            pref = (horiz is None) or (horiz == "outline" and aspect < 0.85) or (horiz is not None and horiz != "outline" and bool(horiz) == (aspect > 1.15))
            score = (round(s, 2), pref, -wrap_em)
            if best is None or score > best[0]:
                best = (score, lay, bb, s, leg_rows, leg_h)
    _, lay, bb, s, leg_rows, leg_h = best
    bw, bh = bb[2] - bb[0], bb[3] - bb[1]
    total_h = (bh + leg_h) * s
    ox = (rw - bw * s) / 2 - bb[0] * s
    oy = (rh - total_h) / 2 - bb[1] * s
    P = lambda p: [round(ox + p[0] * s, 2), round(oy + p[1] * s, 2)]   # noqa: E731
    for b in lay["nodes"].values():
        b["x"], b["y"] = P([b["x"], b["y"]])
        b["w"], b["h"], b["tw"] = round(b["w"] * s, 2), round(b["h"] * s, 2), round(b["tw"] * s, 2)
    for e in lay["edges"]:
        e["pts"] = [P(p) for p in e["pts"]]
        if e.get("label"):
            lb = e["label"]
            lb["x"], lb["y"] = P([lb["x"], lb["y"]])
            lb["w"], lb["h"] = round(lb["w"] * s, 2), round(lb["h"] * s, 2)
    for ln in lay["lanes"]:
        for key in ("band", "head"):
            x, y, w, h = ln[key]
            ln[key] = [*P([x, y]), round(w * s, 2), round(h * s, 2)]
    legend = []
    y = oy + (bb[3] + 1.2) * s + LH * LEG * s * 0.75
    for row in leg_rows:
        row_w = (sum(i["w"] for i in row) + 1.6 * (len(row) - 1)) * s
        x = (rw - row_w) / 2
        for it in row:
            legend.append({**it, "x": round(x, 2), "y": round(y, 2), "w": round(it["w"] * s, 2)})
            x += (it["w"] + 1.6) * s
        y += LH * LEG * 1.45 * s
    return {**lay, "legend": legend, "s": round(s, 3), "small": s < s_min, "bbox": [round(ox + bb[0] * s, 1), round(oy + bb[1] * s, 1),
                                                                               round(ox + bb[2] * s, 1), round(oy + bb[3] * s, 1)]}


def lang_of(spec: dict[str, Any]) -> str:
    from app.tools.infographic import _HU
    return "hu" if _HU.search(" ".join([spec.get("title") or ""] + texts(spec))) else "en"


# ---------- videó: a meglévő „graphic” elemek ----------

_DARK_BG = "#0E1016"
_PANEL = "#151821"


def _node_colors(b: dict[str, Any], actors: dict[str, Any], accent_hex: str, bg: str, light: bool) -> tuple[str, str, str]:
    """→ (kitöltés, keret, szöveg) a csúcs fajtájából és szereplőjéből."""
    col = (actors.get(b.get("actor")) or {}).get("color") if b.get("actor") else None
    base = col or (accent_hex if b["kind"] in ("outcome", "start", "end") else None)
    text = "#1E2430" if light else "#F4F5F8"
    if b["kind"] == "question" and not col:            # a kiinduló kérdés: semleges panel, erős világos keret
        return (mix("#3B4252", bg, 0.9), "#3B4252", text) if light else (mix("#E4E7EE", bg, 0.88), "#E4E7EE", text)
    if light:
        if base:
            return mix(base, bg, 0.84 if b["kind"] != "outcome" else 0.72), base, readable(text, mix(base, bg, 0.8))
        return mix("#8A909C", bg, 0.88), "#8A909C", text
    if base:
        return mix(base, bg, 0.80 if b["kind"] != "outcome" else 0.62), base, text
    return mix("#9AA0AE", bg, 0.86), "#7D8494", text


def to_graphic(spec: dict[str, Any], W: int, H: int, accent: Optional[list[float]] = None, dur: Optional[float] = None,
               lang: Optional[str] = None) -> list[dict[str, Any]]:
    """Az ábra → a videó „graphic” elemei (ugyanaz a formátum, amit a _graphic_from_plan ad, és a _graphic_elements
    rajzol): 0–1 koordináták, méretek 1080-as rövid oldalra, anim/at/dur, és a két kiegészítés: dash (szaggatott /
    pontozott vonal) és dim (halványítás, amíg egy másik lépést mutatunk). A felirat- és címsáv szabadon marad."""
    m = min(W, H)
    k = m / 1080
    acc = _to_hex(tuple(accent)) if accent else "#FF5A36"
    actors = spec.get("_actors") or register(spec, {}, acc)
    lang = lang or lang_of(spec)
    x0, y0, x1, y1 = W * 0.05, H * 0.17, W * 0.95, H * 0.80
    lay = layout(spec, (x1 - x0, y1 - y0), cw=0.56, s_max=30 * k, s_min=22 * k, lang=lang, actors=actors)
    s = lay["s"]
    if lay["small"]:
        try:                                            # a check megemlíti (mint a dobozába kisebb betűvel férő szöveget)
            from app.tools.video import _FIT
            notes = _FIT.get()
            if notes is not None:
                notes.append(f"diagram “{spec.get('title') or ''}” {round(22)}→{round(s / k)} px")
        except Exception:
            pass
    fx = lambda x: round((x0 + x) / W, 5)            # noqa: E731
    fy = lambda y: round((y0 + y) / H, 5)            # noqa: E731
    pt = lambda p: [fx(p[0]), fy(p[1])]              # noqa: E731
    u = lambda px: round(px / k, 2)                   # képpont → 1080-as egység   # noqa: E731
    at, pace, t_end = timing(spec, dur)
    steps = sorted(set(at.values()))
    focus = spec.get("focus", True) and spec.get("reveal") != "all" and len(steps) > 1

    def dim(t0: float) -> dict[str, Any]:
        nxt = next((t for t in steps if t > t0 + 1e-6), None)
        return {"dim": [[round(nxt, 3), round(t_end, 3)]]} if focus and nxt is not None and t_end > nxt + 0.3 else {}
    out: list[dict[str, Any]] = []
    # sávok: háttérsáv a szereplő árnyalatával, fejléc a nevével — elsőként, a többi alatt
    for i, ln in enumerate(lay["lanes"]):
        col = (actors.get(ln["name"]) or {}).get("color") or acc
        bx, by, bw, bh = ln["band"]
        out.append({"shape": "rect", "x": fx(bx), "y": fy(by), "w": round((bw - 2) / W, 5), "h": round((bh - 2) / H, 5),
                    "fill": mix(col, _DARK_BG, 0.86 if i % 2 == 0 else 0.9), "anim": "appear", "at": 0.25, "dur": 0.5})
        hx, hy, hw, hh = ln["head"]
        fs = s * HEAD
        n_l = len(ln["lines"])
        bar_w = max(4.0, 0.22 * s)
        if lay["dir"] == "right":
            out.append({"shape": "rect", "x": fx(hx - hw / 2 + bar_w / 2), "y": fy(hy), "w": round(bar_w / W, 5), "h": round((hh - 2) / H, 5),
                        "fill": col, "anim": "grow", "at": 0.3, "dur": 0.4})
        else:
            out.append({"shape": "rect", "x": fx(hx), "y": fy(hy - hh / 2 + bar_w / 2), "w": round((hw - 2) / W, 5), "h": round(bar_w / H, 5),
                        "fill": col, "anim": "grow", "at": 0.3, "dur": 0.4})
        out.append({"shape": "text", "x": fx(hx + (bar_w / 2 if lay["dir"] == "right" else 0)), "y": fy(hy - (n_l - 1) * 0.59 * fs + (bar_w / 2 if lay["dir"] != "right" else 0)),
                    "w": round((hw - bar_w - 0.4 * s) / W, 5), "text": "\n".join(ln["lines"]), "size": u(fs), "weight": "bold",
                    "align": "center", "color": readable(col, _DARK_BG), "anim": "appear", "at": 0.35, "dur": 0.4})
    # élek: a cél csúcs lépésében rajzolódnak ki (a háló a későbbi végpont lépésében); a címkéjük a dobozok fölé
    first_use: dict[str, float] = {}
    top: list[dict[str, Any]] = []
    for e in lay["edges"]:
        st = STYLE[e["kind"]]
        col = st["dark"]
        wpx = st["width"] * k
        t_e = max(at[e["to"]], at[e["from"]]) + (0.25 if spec["type"] == "network" else 0.0)
        for key in (e["kind"], e["certainty"]):
            first_use[key] = min(first_use.get(key, 99.0), t_e)
        d_ = dim(t_e)
        dash = {"dash": [round(DASH[e["certainty"]][0] * st["width"], 2), round(DASH[e["certainty"]][1] * st["width"], 2)]} \
            if e["certainty"] in DASH else {}
        line = {"stroke": col, "width": st["width"], "opacity": 100, **dash, **d_}
        pts = e["pts"]
        draw = 0.45
        if st["sym"]:                                   # irány nélküli kapcsolat: középről kifelé, két fél vonal
            mid_t = 0.5
            segs = [math.dist(pts[i], pts[i + 1]) for i in range(len(pts) - 1)]
            goal, acc_l, cut = sum(segs) * mid_t, 0.0, 0
            for i, L in enumerate(segs):
                if acc_l + L >= goal:
                    cut = i
                    break
                acc_l += L
            mid = _along(pts, mid_t)
            halves = [[mid] + list(reversed(pts[:cut + 1])), [mid] + pts[cut + 1:]]
            for hp in halves:
                if len(hp) == 2:
                    out.append({"shape": "line", "from": pt(hp[0]), "to": pt(hp[1]), **line, "anim": "draw", "at": round(t_e, 3), "dur": draw})
                else:
                    out.append({"shape": "path", "points": [pt(p) for p in hp], **line, "anim": "draw", "at": round(t_e, 3), "dur": draw})
            if st["head"] == "dots":
                for p in (pts[0], pts[-1]):
                    out.append({"shape": "circle", "x": fx(p[0]), "y": fy(p[1]), "r": round(wpx * 1.5 / m, 5), "fill": col,
                                "opacity": 100, "anim": "pop", "at": round(t_e + draw * 0.9, 3), "dur": 0.25, **d_})
        else:
            segs = [math.dist(pts[i], pts[i + 1]) for i in range(len(pts) - 1)]
            tot = sum(segs) or 1
            t_ = t_e
            if len(pts) > 2:                            # a törött vonal eleje path, az utolsó szakasza nyíl
                d0 = draw * sum(segs[:-1]) / tot
                out.append({"shape": "path", "points": [pt(p) for p in pts[:-1]], **line, "anim": "draw", "at": round(t_, 3), "dur": round(max(0.1, d0), 3)})
                t_ += d0
            d1 = max(0.12, draw * segs[-1] / tot)
            shape = "arrow" if st["head"] == "triangle" else "line"
            out.append({"shape": shape, "from": pt(pts[-2]), "to": pt(pts[-1]), **line, "anim": "draw", "at": round(t_, 3), "dur": round(d1, 3)})
            if st["head"] == "chevron":                 # nyitott „V” hegy (alárendeltség) — nem tömör nyíl
                (ax, ay), (bx_, by_) = pts[-2], pts[-1]
                ang = math.atan2(by_ - ay, bx_ - ax)
                L_ = max(16.0 * k, wpx * 3.0)
                lft = [bx_ - L_ * math.cos(ang) + L_ * 0.7 * math.sin(ang), by_ - L_ * math.sin(ang) - L_ * 0.7 * math.cos(ang)]
                rgt = [bx_ - L_ * math.cos(ang) - L_ * 0.7 * math.sin(ang), by_ - L_ * math.sin(ang) + L_ * 0.7 * math.cos(ang)]
                out.append({"shape": "path", "points": [pt(lft), pt([bx_, by_]), pt(rgt)], "stroke": col, "width": st["width"],
                            "opacity": 100, "anim": "appear", "at": round(t_ + d1 * 0.85, 3), "dur": 0.2, **d_})
        if e["certainty"] == "disputed" or e.get("label"):
            lb = e.get("label")
            fs = s * LAB
            if lb:
                top.append({"shape": "rect", "x": fx(lb["x"]), "y": fy(lb["y"]), "w": round(lb["w"] / W, 5), "h": round(lb["h"] / H, 5),
                            "fill": _PANEL, "stroke": col, "width": 2, "radius": u(lb["h"] / 2), "opacity": 100,
                            "anim": "pop", "at": round(t_e + 0.2, 3), "dur": 0.3, **d_})
                top.append({"shape": "text", "x": fx(lb["x"]), "y": fy(lb["y"] - (len(lb["lines"]) - 1) * 0.59 * fs),
                            "w": round((lb["w"] - 0.2 * s) / W, 5), "text": "\n".join(lb["lines"]), "size": u(fs), "weight": "semibold",
                            "align": "center", "color": readable(col, _PANEL), "opacity": 100, "anim": "appear",
                            "at": round(t_e + 0.28, 3), "dur": 0.3, **d_})
            if e["certainty"] == "disputed":            # „?” jelvény: a címke bal szélén, címke nélkül az él közepén
                r = 0.52 * s
                c = [lb["x"] - lb["w"] / 2 - r * 0.55, lb["y"] - lb["h"] / 2 + r * 0.3] if lb else _along(pts, 0.5)
                top.append({"shape": "circle", "x": fx(c[0]), "y": fy(c[1]), "r": round(r / m, 5), "fill": BADGE, "opacity": 100,
                            "anim": "pop", "at": round(t_e + 0.3, 3), "dur": 0.25, **d_})
                top.append({"shape": "text", "x": fx(c[0]), "y": fy(c[1]), "w": round(2 * r / W, 5), "text": "?", "size": u(r * 1.35),
                            "weight": "bold", "align": "center", "color": "#16181D", "opacity": 100, "anim": "appear",
                            "at": round(t_e + 0.35, 3), "dur": 0.2, **d_})
    # csúcsok: a dobozuk felpattan, a szöveg utána
    for n in spec["nodes"]:
        b = lay["nodes"][n["id"]]
        fill, stroke, tcol = _node_colors(b, actors, acc, _DARK_BG, False)
        t_n = at[n["id"]] + (0.0 if spec["type"] == "network" else 0.35)
        d_ = dim(at[n["id"]] + (0.25 if spec["type"] == "network" else 0.0))
        fs = s * b["font"]
        sw = 4 if b["kind"] in ("question", "decision") or b["bold"] else 3
        common = {"fill": fill, "stroke": stroke, "width": sw, "opacity": 100, **d_}
        if b["shape"] == "hex":                         # döntés: hatszög (két csúcsos vég), sokszögként
            cx, cy, w2, h2 = b["x"], b["y"], b["w"] / 2, b["h"] / 2
            cut = min(h2, 0.9 * s)
            poly = [[cx - w2, cy], [cx - w2 + cut, cy - h2], [cx + w2 - cut, cy - h2], [cx + w2, cy], [cx + w2 - cut, cy + h2], [cx - w2 + cut, cy + h2]]
            out.append({"shape": "polygon", "points": [pt(p) for p in poly], **common, "anim": "appear", "at": round(t_n, 3), "dur": 0.3})
        else:
            rad = b["h"] / 2 if b["shape"] == "pill" else 0.35 * s
            out.append({"shape": "rect", "x": fx(b["x"]), "y": fy(b["y"]), "w": round(b["w"] / W, 5), "h": round(b["h"] / H, 5),
                        "radius": u(rad), **common, "anim": "pop", "at": round(t_n, 3), "dur": 0.35})
        n_l = len(b["lines"])
        y_first = b["y"] - (n_l - 1) * 0.59 * fs
        chunks = [b["lines"][i:i + 2] for i in range(0, n_l, 2)] if n_l > 3 else [b["lines"]]
        for j, ch_ in enumerate(chunks):          # a graphic szövegdoboza legfeljebb 3 sort fogad: 4 sor = 2 + 2
            out.append({"shape": "text", "x": fx(b["x"]), "y": fy(y_first + 2 * j * 1.18 * fs), "w": round((b["tw"] + 0.5 * s) / W, 5),
                        "text": "\n".join(ch_), "size": u(fs), "weight": "bold" if b["bold"] else "semibold", "align": "center",
                        "color": tcol, "opacity": 100, "anim": "appear", "at": round(t_n + 0.12, 3), "dur": 0.3, **d_})
    out += top
    # jelmagyarázat: minden tétele akkor jelenik meg, amikor a fajtája először
    fs = max(s * LEG, 18 * k)
    for it in lay["legend"]:
        x, y = it["x"], it["y"]
        t_l = round(first_use.get(it["key"], t_end) + 0.3, 3) if spec.get("reveal") != "all" else 1.2
        if it["what"] == "group":
            col = it.get("color") or acc
            out.append({"shape": "rect", "x": fx(x + 0.5 * s), "y": fy(y), "w": round(0.9 * s / W, 5), "h": round(0.9 * s / H, 5),
                        "fill": mix(col, _DARK_BG, 0.5), "stroke": col, "width": 3, "radius": u(0.2 * s), "anim": "pop", "at": t_l, "dur": 0.3})
            tx = x + 1.2 * s + 0.5 * s
        else:
            kind = it["key"] if it["what"] == "kind" else "link"
            st = STYLE[kind]
            col = st["dark"] if it["what"] == "kind" else "#C9CDD6"
            dash = {"dash": [round(DASH[it["key"]][0] * st["width"], 2), round(DASH[it["key"]][1] * st["width"], 2)]} if it["what"] == "cert" else {}
            a_, b_ = [x, y], [x + 2.6 * s, y]
            shape = "arrow" if st["head"] == "triangle" and it["what"] == "kind" else "line"
            out.append({"shape": shape, "from": pt(a_), "to": pt(b_), "stroke": col, "width": st["width"], **dash,
                        "anim": "draw", "at": t_l, "dur": 0.4})
            if it["what"] == "kind" and st["head"] == "chevron":
                L_ = max(16.0 * k, st["width"] * k * 3.0)
                out.append({"shape": "path", "points": [pt([b_[0] - L_, y - L_ * 0.7]), pt(b_), pt([b_[0] - L_, y + L_ * 0.7])],
                            "stroke": col, "width": st["width"], "anim": "appear", "at": round(t_l + 0.3, 3), "dur": 0.2})
            if it["what"] == "kind" and st["head"] == "dots":
                for p in (a_, b_):
                    out.append({"shape": "circle", "x": fx(p[0]), "y": fy(p[1]), "r": round(st["width"] * k * 1.5 / m, 5), "fill": col,
                                "anim": "pop", "at": round(t_l + 0.3, 3), "dur": 0.2})
            if it["what"] == "cert" and it["key"] == "disputed":
                r = 0.42 * s
                out.append({"shape": "circle", "x": fx(x + 1.3 * s), "y": fy(y), "r": round(r / m, 5), "fill": BADGE, "anim": "pop",
                            "at": round(t_l + 0.2, 3), "dur": 0.2})
                out.append({"shape": "text", "x": fx(x + 1.3 * s), "y": fy(y), "w": round(2 * r / W, 5), "text": "?", "size": u(r * 1.35),
                            "weight": "bold", "align": "center", "color": "#16181D", "anim": "appear", "at": round(t_l + 0.25, 3), "dur": 0.2})
            tx = x + 2.6 * s + 0.5 * s
        tw = len(it["text"]) * 0.56 * fs + 0.4 * s
        out.append({"shape": "text", "x": fx(tx + tw / 2), "y": fy(y), "w": round(tw / W, 5), "text": it["text"], "size": u(fs),
                    "weight": "semibold", "align": "left", "color": "#C9CDD6", "anim": "appear",
                    "at": round(t_l + 0.1, 3), "dur": 0.3})
    return out


# ---------- pakli: natív PowerPoint-alakzatok ----------

def to_pptx(slide, spec: dict[str, Any], t: dict[str, str], box: tuple[float, float, float, float],
            actors: Optional[dict[str, Any]] = None) -> list[Any]:
    """Ugyanaz az elrendezés natív, szerkeszthető alakzatokkal (box: x, y, w, h hüvelykben). A csúcs egy lekerekített
    téglalap a saját szövegével; az egyenes él a két dobozhoz KÖTÖTT összekötő (a doboz mozgatásakor követi), a
    törött él szerkeszthető vonal; a stílusok a videóéval azonosak (világos változat), a jelmagyarázat is natív."""
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    X0, Y0, BW, BH = (v * 72 for v in box)                # hüvelyk → pont
    bg = "#" + t.get("bg", "#FFFFFF").lstrip("#")
    acc = "#" + t.get("accent", "#FCA311").lstrip("#")
    actors = actors or register(spec, {}, acc)
    lang = lang_of(spec)
    lay = layout(spec, (BW, BH), cw=0.56, s_max=16.0, s_min=9.5, lang=lang, actors=actors)
    s = lay["s"]
    E = lambda v: Emu(int(round(v * 12700)))          # noqa: E731
    rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())   # noqa: E731
    shapes: list[Any] = []
    font = t.get("font_body", "Calibri")            # sablonnál None: a téma betűje öröklődik
    text_col = "#" + t.get("text", "#1E2430").lstrip("#")

    def textframe(shp, lines: list[str], size: float, color: str, bold: bool, align: str = "center", wrap: bool = True) -> None:
        """Soronként egy bekezdés (a mi tördelésünk: minden motor ugyanígy rajzolja); a rövid címke nem tördel."""
        tf = shp.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = E(0.12 * size)
        tf.margin_top = tf.margin_bottom = E(0.04 * size)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for k, line in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT}[align]
            r = p.add_run()
            r.text = line
            f = r.font
            f.size, f.bold = Pt(round(size, 1)), bold
            if font:
                f.name = font
            f.color.rgb = rgb(color)

    def plain(shp) -> None:
        """Saját stílus: a témából semmi (árnyék, alapvonal) — a kitöltés és a vonal kifejezetten megadva."""
        st = shp._element.find(qn("p:style"))
        if st is not None:
            shp._element.remove(st)

    def rect(name: str, cx: float, cy: float, w: float, h: float, fill: Optional[str], line: Optional[str], lw: float = 1.25,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius: Optional[float] = None):
        shp = slide.shapes.add_shape(shape, E(X0 + cx - w / 2), E(Y0 + cy - h / 2), E(w), E(h))
        shp.name = name
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(lw)
        else:
            shp.line.fill.background()
        plain(shp)
        if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = max(0.0, min(0.5, radius / max(min(w, h), 1e-6)))
        shapes.append(shp)
        return shp

    def ends(ln_el, head: Optional[str], tail_too: bool = False) -> None:
        for tag, kind in (("a:headEnd", "oval" if head == "dots" else None), ("a:tailEnd", {"triangle": "triangle", "chevron": "arrow", "dots": "oval"}.get(head or ""))):
            if kind:
                el = ln_el.find(qn(tag))
                if el is None:
                    el = etree.SubElement(ln_el, qn(tag))
                el.set("type", kind)
                el.set("w", "med")
                el.set("len", "med")

    def style_line(line_fmt, kind: str, cert: str) -> None:
        st = STYLE[kind]
        line_fmt.color.rgb = rgb(st["light"])
        line_fmt.width = Pt(round(st["width"] * 0.5, 2))
        if cert == "uncertain":
            line_fmt.dash_style = MSO_LINE_DASH_STYLE.DASH
        elif cert == "disputed":
            line_fmt.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT

    # sávok
    for i, ln in enumerate(lay["lanes"]):
        col = (actors.get(ln["name"]) or {}).get("color") or acc
        bx, by, bw, bh = ln["band"]
        rect(f"Lane – {ln['name']}", bx, by, bw, bh, mix(col, bg, 0.9 if i % 2 == 0 else 0.94), None, shape=MSO_SHAPE.RECTANGLE)
        hx, hy, hw, hh = ln["head"]
        hd = rect(f"Lane header – {ln['name']}", hx, hy, hw, hh, None, None, shape=MSO_SHAPE.RECTANGLE)
        textframe(hd, ln["lines"], s * HEAD, readable(col, mix(col, bg, 0.9)), True)
    # csúcsok előbb (az összekötő ezekhez kötődik), az élek utánuk kerülnek alájuk a z-sorrendben
    node_shp = {}
    for n in spec["nodes"]:
        b = lay["nodes"][n["id"]]
        fill, stroke, tcol = _node_colors(b, actors, acc, bg, True)
        shape = MSO_SHAPE.HEXAGON if b["shape"] == "hex" else MSO_SHAPE.ROUNDED_RECTANGLE
        shp = rect(f"Node – {n['id']}", b["x"], b["y"], b["w"], b["h"], fill, stroke, 2.0 if b["bold"] else 1.25, shape=shape,
                   radius=b["h"] / 2 if b["shape"] == "pill" else 0.35 * s)
        textframe(shp, b["lines"], s * b["font"], tcol, b["bold"])
        node_shp[n["id"]] = shp
    sites = {"top": 0, "left": 1, "bottom": 2, "right": 3}
    edge_shapes = []
    for e in lay["edges"]:
        st = STYLE[e["kind"]]
        pts = e["pts"]
        name = f"Edge – {e['from']}→{e['to']}"
        if len(pts) == 2:
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(X0 + pts[0][0]), E(Y0 + pts[0][1]), E(X0 + pts[1][0]), E(Y0 + pts[1][1]))
            for which, p, nid in (("begin", pts[0], e["from"]), ("end", pts[1], e["to"])):
                b = lay["nodes"][nid]
                side = min(sites, key=lambda sd: math.dist(_side(b, sd), p))
                if b["shape"] != "hex":
                    (c.begin_connect if which == "begin" else c.end_connect)(node_shp[nid], sites[side])
        else:
            fb = slide.shapes.build_freeform(E(X0 + pts[0][0]), E(Y0 + pts[0][1]), scale=1.0)
            fb.add_line_segments([(E(X0 + p[0]), E(Y0 + p[1])) for p in pts[1:]], close=False)
            c = fb.convert_to_shape()
            c.fill.background()
            for path in c._element.iter(qn("a:path")):    # nyitott vonal: kitöltés nélkül (a DeckCraft is így rajzolja)
                path.set("fill", "none")
        c.name = name
        plain(c)
        style_line(c.line, e["kind"], e["certainty"])
        ends(c.line._get_or_add_ln(), st["head"])
        edge_shapes.append(c)
        shapes.append(c)
        lb = e.get("label")
        fs = s * LAB
        if lb:
            sh = rect(f"Label – {e['from']}→{e['to']}", lb["x"], lb["y"], lb["w"], lb["h"], bg, st["light"], 0.75, radius=lb["h"] / 2)
            textframe(sh, lb["lines"], fs, readable(st["light"], bg), False, wrap=False)
        if e["certainty"] == "disputed":
            r = 0.55 * s
            c_ = [lb["x"] - lb["w"] / 2 - r * 0.55, lb["y"] - lb["h"] / 2 + r * 0.3] if lb else _along(pts, 0.5)
            bd = rect(f"Disputed – {e['from']}→{e['to']}", c_[0], c_[1], 2 * r, 2 * r, BADGE, None, shape=MSO_SHAPE.OVAL)
            textframe(bd, ["?"], r * 1.2, "#16181D", True, wrap=False)
    # az élek a csúcsok alá (z-sorrend): a spTree-ben a sávok után, a csúcsok elé
    tree = slide.shapes._spTree
    first_node = node_shp[spec["nodes"][0]["id"]]._element if spec["nodes"] else None
    for c in edge_shapes:
        if first_node is not None:
            tree.remove(c._element)
            first_node.addprevious(c._element)
    # jelmagyarázat
    fs = max(s * LEG, 9.0)
    for it in lay["legend"]:
        x, y = it["x"], it["y"]
        if it["what"] == "group":
            col = it.get("color") or acc
            rect(f"Legend – {it['key']}", x + 0.5 * s, y, 0.9 * s, 0.9 * s, mix(col, bg, 0.75), col, 1.0, radius=0.2 * s)
            tx = x + 1.7 * s
        else:
            kind = it["key"] if it["what"] == "kind" else "link"
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(X0 + x), E(Y0 + y), E(X0 + x + 2.6 * s), E(Y0 + y))
            c.name = f"Legend – {it['key']}"
            plain(c)
            style_line(c.line, kind, it["key"] if it["what"] == "cert" else "certain")
            if it["what"] == "cert":
                c.line.color.rgb = rgb(text_col)
            ends(c.line._get_or_add_ln(), STYLE[kind]["head"] if it["what"] == "kind" else None)
            shapes.append(c)
            if it["what"] == "cert" and it["key"] == "disputed":
                r = 0.42 * s
                bd = rect("Legend – disputed mark", x + 1.3 * s, y, 2 * r, 2 * r, BADGE, None, shape=MSO_SHAPE.OVAL)
                textframe(bd, ["?"], r * 1.2, "#16181D", True, wrap=False)
            tx = x + 3.1 * s
        tw = len(it["text"]) * 0.58 * fs + 0.6 * s
        tb = rect(f"Legend text – {it['key']}", tx + tw / 2, y, tw, fs * 1.5, None, None, shape=MSO_SHAPE.RECTANGLE)
        textframe(tb, [it["text"]], fs, "#" + t.get("muted", "#6B7280").lstrip("#"), False, "left", wrap=False)
    return shapes
