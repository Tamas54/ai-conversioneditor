"""Word-dokumentum iteratív szerkesztése szavakkal: „a 22. oldalra kellene egy fotó”, „a harmadik bekezdés után
tabbal kezdődjenek a bekezdések”, „a második fejezet címe legyen …”.

A modell (AGENT_MODEL, DeepSeek V4.1 Flash) csak TERVET ír egy szűk lépéskészletből a dokumentum vázlata alapján
(bekezdés-sorszám, oldal, stílus, szövegeleje); a végrehajtás python-docx-szel determinisztikus. Az oldalszámot a
WordCraft PDF-renderjéből számoljuk (48 oldal ~70 ms), így „a 22. oldal” valódi bekezdésre fordítható.
Ami nem fér a lépéskészletbe → a régi szerkesztő-ügynök (run_agent) kapja meg.
"""
from __future__ import annotations

import asyncio
import json
import logging
import re
import subprocess
from pathlib import Path
from typing import Any, Optional

from app.config import settings

log = logging.getLogger("aice.doc_revise")

_PROMPT = """You edit a Word document by choosing steps. You do not rewrite the document yourself.
Outline of the document — one line per paragraph: #index p<page> [style] beginning of the text
{outline}

The user asks: {request}

Reply with ONE JSON object {{"steps": [...], "summary": "one short sentence in the language of the request — English if the request is only a few words or its language is unclear"}}.
Steps (use the #index numbers above; "to" is inclusive; omit "to" for a single paragraph):
  {{"op":"indent","from":i,"to":j,"first_line_cm":1.0}}             first-line indent (like a tab) on body paragraphs
  {{"op":"insert_image","after":i,"stock":"2-4 English words"}}      a real photo (free) after paragraph i (main noun last)
  {{"op":"insert_image","after":i,"ai":"prompt"}}                    a generated picture (costs a credit) — only if asked
      optional for images: "width_cm": 6-16, "caption": "text under it"
  {{"op":"insert_text","after":i,"text":"...","style":"Normal|Heading 1|Heading 2|Heading 3|Quote"}}
  {{"op":"replace_text","find":"exact text","replace":"new text","from":i,"to":j}}   (no from/to = whole document)
  {{"op":"delete","from":i,"to":j}}
  {{"op":"style","from":i,"to":j,"font":"font family name","color":"#RRGGBB","size_pt":n,"bold":true|false,"italic":true|false,"align":"left|center|right|justify","line_spacing":1.15,"space_after_pt":n}}
      (a colour goes in "color" as a hex code, never in "font")
  {{"op":"page_break_before","at":i}}
  {{"op":"heading","at":i,"level":1|2|3}}
  {{"op":"unsupported","reason":"…"}}   when no step can do it
"Page N" means the page shown as p<N>. Body paragraphs carry {{Section ¶n}}: "after the third paragraph of X" means
from the paragraph tagged {{X ¶4}}. Worked example: lines "#10 … {{Intro ¶3}} …" and "#11 … {{Intro ¶4}} …" →
"after the third paragraph of Intro" gives "from": 11 (NOT 10). A range that has a start but no end ("from … on", "after …", "every paragraph after …")
runs to the last index. For a picture "on page N", put it after the first full body paragraph on page N."""


# ---------- vázlat + oldal-megfeleltetés ----------

def _norm(s: str) -> str:
    return re.sub(r"\W+", " ", s.lower()).strip()


def page_map(docx: Path, paragraphs: list[str]) -> list[int]:
    """Bekezdésenként az oldal (1-től), a WordCraft PDF-jéből: az első 6 szó első előfordulása, monoton."""
    from app.tools import wordcraft
    pdf = docx.with_suffix(".pagemap.pdf")
    try:
        subprocess.run([wordcraft.cli_path() or "wordcraft-cli", "convert", str(docx), str(pdf)], check=True,
                       capture_output=True, timeout=120, env={**__import__("os").environ, "WORDCRAFT_NO_PREFS": "1"})
        n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout).group(1))
        pages = [_norm(subprocess.run(["pdftotext", "-f", str(k), "-l", str(k), str(pdf), "-"], capture_output=True, text=True).stdout)
                 for k in range(1, n + 1)]
    except Exception as e:
        log.warning("oldaltérkép nem készült: %s", e)
        return [0] * len(paragraphs)
    finally:
        pdf.unlink(missing_ok=True)
    out, cur = [], 1
    for text in paragraphs:
        words = _norm(text).split()[:6]
        if len(words) >= 3:
            key = " ".join(words)
            for k in range(cur, min(cur + 6, len(pages) + 1)):     # előre keresünk, néhány oldalon belül
                if key in pages[k - 1]:
                    cur = k
                    break
        out.append(cur)
    return out


def outline(doc, pages: list[int], limit: int = 2500) -> str:
    """#index p<oldal> [stílus] {Fejezet ¶n} szöveg — a szövegbekezdések fejezeten belüli sorszámával, hogy a modellnek
    ne kelljen számolnia („a harmadik bekezdés után” elszámolta)."""
    lines, section, k = [], "start", 0
    for i, p in enumerate(doc.paragraphs[:limit]):
        t = p.text.strip()
        if not t and not p._p.xpath(".//w:drawing"):
            continue
        style = p.style.name if p.style is not None else "Normal"
        tag = ""
        if style.lower().startswith(("heading", "title")):
            section, k = t[:40], 0
        elif _is_body(p):
            k += 1
            tag = f"{{{section} ¶{k}}} "
        lines.append(f"#{i} p{pages[i] if i < len(pages) else 0} [{style}] {tag}{t[:90] if t else '(picture)'}")
    return "\n".join(lines) + f"\n(last index: #{min(len(doc.paragraphs), limit) - 1})"


# ---------- terv ----------

async def plan(doc, pages: list[int], request: str) -> dict[str, Any]:
    from app.llm.siliconflow import _call_model, _strip_json_fences
    prompt = _PROMPT.format(outline=outline(doc, pages), request=request.strip()[:2000])
    last: Exception = RuntimeError("no plan")
    for model in (settings.AGENT_MODEL, settings.AGENT_MODEL_FALLBACK):
        try:
            raw = await _call_model(model, [{"role": "user", "content": prompt}], thinking=False)
            data = json.loads(_strip_json_fences(raw)[_strip_json_fences(raw).find("{"):])
            if isinstance(data.get("steps"), list):
                return data
        except Exception as e:
            last = e
            log.warning("doc-terv (%s): %s", model, e)
    raise RuntimeError(f"Could not plan the change: {last}")


# ---------- végrehajtás ----------

# színnév → hex (ha a modell mégis névvel vagy a "font" mezőben adja a színt)
_COLORS = {"black": "000000", "white": "FFFFFF", "gray": "7F7F7F", "grey": "7F7F7F", "dark gray": "404040",
           "dark grey": "404040", "red": "C00000", "dark red": "8B0000", "orange": "ED7D31", "yellow": "FFC000",
           "green": "00B050", "dark green": "1E5631", "blue": "2E74B5", "dark blue": "1F3864", "navy": "1F3864",
           "light blue": "9DC3E6", "purple": "7030A0", "pink": "E83E8C", "brown": "843C0C", "teal": "008080"}


def _color(v: Any) -> Optional[str]:
    """„#1F3864”, „1f3864”, „dark blue” → "1F3864"; ami nem szín: None."""
    t = str(v or "").strip().lower()
    if re.fullmatch(r"#?[0-9a-f]{6}", t):
        return t.lstrip("#").upper()
    return _COLORS.get(t.replace("-", " "))


def _replace_in_runs(p, find: str, repl: str, start: int = 0) -> int:
    """Csere a bekezdésben úgy, hogy a többi futás formázása megmaradjon: csak a `find`-ot lefedő futások
    érintett része változik, a csere az első érintett futás formázását kapja (a „Helyszín:” félkövér címkéje így
    nem nyeli el a mögötte álló szöveget — UI-próba 10-18). → a csere utáni pozíció, vagy -1 (nincs találat)."""
    runs = p.runs
    text = "".join(r.text for r in runs)
    at = text.find(find, start) if find else -1
    if at < 0:
        return -1
    end, pos, first = at + len(find), 0, True
    for r in runs:
        a, b = pos, pos + len(r.text)
        pos = b
        if b <= at or a >= end:
            continue
        lo, hi = max(at, a) - a, min(end, b) - a           # az érintett szelet a futáson belül
        r.text = r.text[:lo] + (repl if first else "") + r.text[hi:]
        first = False
    return at + len(repl)                               # innen keressük a következőt (a csere belsejét nem)


def _is_body(p) -> bool:
    name = (p.style.name if p.style is not None else "").lower()
    return bool(p.text.strip()) and not name.startswith(("heading", "title", "subtitle", "toc", "caption", "image caption", "centered")) \
        and (p.alignment is None or int(p.alignment) in (0, 3))       # bal / sorkizárt (a középre zárt címlap nem)


def _insert_after(par, style: Optional[str] = None):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    new = OxmlElement("w:p")
    par._p.addnext(new)
    p = Paragraph(new, par._parent)
    if style:
        try:
            p.style = style
        except KeyError:
            pass
    return p


async def apply(doc, steps: list[dict[str, Any]]) -> tuple[int, list[str]]:
    """Lépések a docx-en. Vissza: (végrehajtott lépések, nem támogatott kérések)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    pars = list(doc.paragraphs)                  # az eredeti bekezdés-objektumok (a beszúrások ezt nem tolják el)
    n = len(pars)
    rng = lambda s, a="from", b="to": range(max(0, int(s.get(a, 0))), min(n, int(s.get(b, s.get(a, 0))) + 1))
    done, unsupported = 0, []
    ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
             "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
    structural = []
    for s in steps:
        op = s.get("op")
        try:
            if op == "indent":
                for i in rng(s):
                    if _is_body(pars[i]):
                        pars[i].paragraph_format.first_line_indent = Cm(float(s.get("first_line_cm") or 1.0))
                done += 1
            elif op == "style":
                for i in rng(s):
                    pf = pars[i].paragraph_format
                    if s.get("align") in ALIGN:
                        pars[i].alignment = ALIGN[s["align"]]
                    if s.get("line_spacing"):
                        pf.line_spacing = float(s["line_spacing"])
                    if s.get("space_after_pt") is not None:
                        pf.space_after = Pt(float(s["space_after_pt"]))
                    color = _color(s.get("color")) or _color(s.get("font"))
                    font = s.get("font") if s.get("font") and not _color(s.get("font")) else None   # szín ≠ betűtípus
                    for r in pars[i].runs:
                        if font:
                            r.font.name = font
                        if color:
                            from docx.shared import RGBColor
                            r.font.color.rgb = RGBColor.from_string(color)
                        if s.get("size_pt"):
                            r.font.size = Pt(float(s["size_pt"]))
                        if s.get("bold") is not None:
                            r.font.bold = bool(s["bold"])
                        if s.get("italic") is not None:
                            r.font.italic = bool(s["italic"])
                done += 1
            elif op == "replace_text" and s.get("find"):
                targets = rng(s) if "from" in s else range(n)
                for i in targets:
                    p = pars[i]
                    if s["find"] not in p.text:
                        continue
                    pos = _replace_in_runs(p, s["find"], s.get("replace", ""))
                    if pos < 0:
                        if p.runs:                  # a szöveg nem a futásokban (pl. hivatkozásban): a régi út
                            p.runs[0].text = p.text.replace(s["find"], s.get("replace", ""))
                            for r in p.runs[1:]:
                                r.text = ""
                        continue
                    while pos >= 0:                 # a többi előfordulás is, a már cserélt szöveg után folytatva
                        pos = _replace_in_runs(p, s["find"], s.get("replace", ""), pos)
                done += 1
            elif op == "page_break_before" and 0 <= int(s.get("at", -1)) < n:
                pars[int(s["at"])].paragraph_format.page_break_before = True
                done += 1
            elif op == "heading" and 0 <= int(s.get("at", -1)) < n:
                pars[int(s["at"])].style = f"Heading {int(s.get('level') or 1)}"
                done += 1
            elif op in ("insert_image", "insert_text", "delete"):
                structural.append(s)
            elif op == "unsupported":
                unsupported.append(s.get("reason") or "")
        except Exception as e:
            log.warning("doc-lépés %s: %s", op, e)
    # szerkezeti lépések hátulról előre (a „after” indexek az eredeti bekezdésekre mutatnak)
    for s in sorted(structural, key=lambda s: -int(s.get("after", s.get("from", 0)))):
        op = s["op"]
        try:
            if op == "delete":
                for i in reversed(list(rng(s))):
                    el = pars[i]._p
                    el.getparent().remove(el)
                done += 1
            elif op == "insert_text" and 0 <= int(s.get("after", -1)) < n:
                p = _insert_after(pars[int(s["after"])], s.get("style"))
                p.add_run(str(s.get("text", "")))
                done += 1
            elif op == "insert_image" and 0 <= int(s.get("after", -1)) < n:
                from app.routes.make import _materialize
                from app.storage import read_meta
                from app.tools.crafts import existing
                spec = {"stock": s["stock"]} if s.get("stock") else {"generate": s["ai"]} if s.get("ai") else None
                if not spec:
                    continue
                fid = await _materialize(spec)
                anchor = pars[int(s["after"])]
                pic = _insert_after(anchor)
                pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.add_run().add_picture(str(existing(fid)), width=Cm(min(16.0, max(4.0, float(s.get("width_cm") or 12)))))
                credit = ((read_meta(fid) or {}).get("extra") or {}).get("credit") or (read_meta(fid) or {}).get("credit")
                cap_text = " — ".join(x for x in (s.get("caption"), credit) if x)
                if cap_text:
                    cap = _insert_after(pic)
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = cap.add_run(cap_text)
                    run.italic, run.font.size = True, Pt(8.5)
                done += 1
        except Exception as e:
            log.warning("doc-lépés %s: %s", op, e)
    return done, unsupported


async def revise(file_id: str, request: str, dst: Path) -> dict[str, Any]:
    from docx import Document
    from app.tools.crafts import existing
    src = existing(file_id)
    work = dst.with_suffix(".src.docx")
    work.write_bytes(src.read_bytes())
    try:
        doc = Document(str(work))
        texts = [p.text for p in doc.paragraphs]
        pages = await asyncio.to_thread(page_map, work, texts)
        p = await plan(doc, pages, request)
        done, unsupported = await apply(doc, p.get("steps") or [])
        if not done:
            return {"handled": False, "reason": "; ".join(unsupported) or "no step fits", "plan": p.get("steps")}
        doc.save(str(dst))
    finally:
        work.unlink(missing_ok=True)
    return {"handled": True, "summary": p.get("summary") or "", "plan": p.get("steps"), "steps_done": done,
            "not_done": unsupported}
