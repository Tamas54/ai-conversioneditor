"""Dokumentum-stílus átvétele mintából: „olyan legyen, mint ez a dokumentumom”.

Minta → referencia-docx (stílusok + oldalbeállítás + téma):
  docx            pontosan: maga a minta (a pandoc --reference-doc csak a stílusait, oldalbeállítását,
                  élőfejét/lábát használja, a tartalmát nem)
  odt/rtf/doc     WordCraft-tal docx-szé, aztán mint fent
  pdf             KÖZELÍTŐ: PyMuPDF-fel mérve — oldalméret, margók, a leggyakoribb (szöveg-) betű és méret, a nagyobb
                  méretek címsornak, sorköz, első soros behúzás → pandoc-alapú referencia-docx ezekkel
  kép / pptx      a style_ref témájából (szín, címsor- és szövegbetű)
Alkalmazás:
  új dokumentum (pandoc)   → --reference-doc = a referencia (Scan to text, Document „from a brief”)
  meglévő docx             → apply_style: a minta stílusai + témája + oldalbeállítása; a cél által használt, de a mintában
                             nem létező stílusok (pl. „Body Text”, „Centered”) megmaradnak, betű/méret/szín nélkül,
                             így a minta „Normal”-jától örökölnek
"""
from __future__ import annotations

import copy
import logging
import re
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any

log = logging.getLogger("aice.doc_style")


# ---------- referencia a mintából ----------

async def reference_for(style_file_id: str) -> Path:
    """A mintából pandoc-referencia (docx) egy ideiglenes helyre."""
    import asyncio
    from app.tools.crafts import existing
    src = existing(style_file_id)
    ext = src.suffix.lower()
    out = Path(tempfile.mkdtemp()) / "style_reference.docx"
    if ext == ".docx":
        shutil.copy(src, out)
    elif ext in (".odt", ".rtf", ".doc"):
        from app.tools import wordcraft
        await wordcraft.convert_file(src, out)
    elif ext == ".pdf":
        await asyncio.to_thread(reference_from_pdf, src, out)
    else:
        from app.tools import style_ref
        theme = await style_ref.theme_from(style_file_id)
        await asyncio.to_thread(reference_from_theme, theme, out)
    return out


def _pandoc_default(dst: Path) -> None:
    subprocess.run(["pandoc", "-o", str(dst), "--print-default-data-file", "reference.docx"], capture_output=True, check=True)


def _family(name: str) -> str:
    """PDF-betűnév → családnév: „ABCDEF+TimesNewRomanPS-BoldMT” → „Times New Roman”."""
    n = name.split("+")[-1]
    n = re.split(r"[-,]", n)[0]
    n = re.sub(r"(PSMT|PS|MT|Std|Pro|LT|Regular|Roman)$", "", n) if n not in ("TimesNewRoman",) else n
    n = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", n).strip()
    return {"Times New Roman PS": "Times New Roman", "Arial MT": "Arial"}.get(n, n) or "Times New Roman"


def measure_pdf(src: Path, pages: int = 6) -> dict[str, Any]:
    """Oldalméret, margók, szövegbetű/méret, címsor-méretek, sorköz, első soros behúzás — az első oldalakból."""
    import pymupdf as fitz
    doc = fitz.open(str(src))
    sizes: Counter = Counter()
    fonts: Counter = Counter()
    head: dict[float, Counter] = {}
    xs, x1s, ys, y1s, gaps, indents = [], [], [], [], [], []
    W = H = 0.0
    for pno in range(min(pages, len(doc))):
        page = doc[pno]
        W, H = page.rect.width, page.rect.height
        for b in page.get_text("dict")["blocks"]:
            if b.get("type") != 0:
                continue
            lines = b["lines"]
            for li, l in enumerate(lines):
                spans = [s for s in l["spans"] if s["text"].strip()]
                if not spans:
                    continue
                s0 = spans[0]
                size = round(s0["size"], 1)
                n = sum(len(s["text"]) for s in spans)
                sizes[size] += n
                fonts[(_family(s0["font"]), size)] += n
                if s0["flags"] & 16 or size >= 0:          # minden méretnél tároljuk a betűt és a színt
                    head.setdefault(size, Counter())[(_family(s0["font"]), bool(s0["flags"] & 16), s0.get("color", 0))] += n
                xs.append(l["bbox"][0]); x1s.append(l["bbox"][2]); ys.append(l["bbox"][1]); y1s.append(l["bbox"][3])
                if li:
                    gaps.append(l["bbox"][1] - lines[li - 1]["bbox"][1])
            if len(lines) >= 2:
                indents.append(lines[0]["bbox"][0] - lines[1]["bbox"][0])
    body = sizes.most_common(1)[0][0] if sizes else 11.0
    body_font = next((f for (f, s), _ in fonts.most_common() if s == body), "Times New Roman")
    bigger = sorted((s for s in sizes if s >= body * 1.15 and sizes[s] > 20), reverse=True)[:3]
    headings = []
    for s in bigger:
        (f, bold, col), _ = head[s].most_common(1)[0]
        headings.append({"size": s, "font": f, "bold": bold, "color": f"#{col:06X}" if isinstance(col, int) else "#000000"})
    q = lambda v, p: sorted(v)[int(len(v) * p)] if v else 0
    line_ratio = (q([g for g in gaps if 0 < g < body * 3], 0.5) / body) if gaps else 1.2
    indent = q([i for i in indents if 2 < i < body * 5], 0.5) if len([i for i in indents if 2 < i < body * 5]) > 3 else 0
    return {"page_w": W, "page_h": H, "left": q(xs, 0.05), "right": W - q(x1s, 0.95), "top": q(ys, 0.03), "bottom": H - q(y1s, 0.97),
            "body_size": body, "body_font": body_font, "headings": headings, "line_spacing": round(max(1.0, min(2.2, line_ratio / 1.17)), 2),
            "first_indent_pt": round(indent, 1)}


def reference_from_pdf(src: Path, dst: Path) -> dict[str, Any]:
    from docx import Document
    from docx.shared import Pt, RGBColor
    m = measure_pdf(src)
    _pandoc_default(dst)
    doc = Document(str(dst))
    for sec in doc.sections:
        if m["page_w"]:
            sec.page_width, sec.page_height = Pt(m["page_w"]), Pt(m["page_h"])
            sec.left_margin, sec.right_margin = Pt(max(36, m["left"])), Pt(max(36, m["right"]))
            sec.top_margin, sec.bottom_margin = Pt(max(36, m["top"] - 6)), Pt(max(36, m["bottom"] - 6))
    st = {s.name: s for s in doc.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in st:
            st[name].font.name, st[name].font.size = m["body_font"], Pt(m["body_size"])
            st[name].font.color.rgb = RGBColor(0x11, 0x11, 0x11)
            pf = st[name].paragraph_format
            pf.line_spacing = m["line_spacing"]
            if name == "Body Text" and m["first_indent_pt"]:
                pf.first_line_indent = Pt(m["first_indent_pt"])
                pf.space_after = Pt(0)
    for i, h in enumerate(m["headings"], 1):
        name = f"Heading {i}"
        if name in st:
            s = st[name]
            s.font.name, s.font.size, s.font.bold = h["font"], Pt(h["size"]), h["bold"]
            s.font.color.rgb = RGBColor.from_string(h["color"].lstrip("#")) if re.fullmatch(r"#[0-9A-F]{6}", h["color"]) else RGBColor(0, 0, 0)
    _drop_theme_fonts(doc)
    doc.save(str(dst))
    return m


def reference_from_theme(theme: dict[str, Any], dst: Path) -> None:
    from docx import Document
    from docx.shared import RGBColor
    _pandoc_default(dst)
    doc = Document(str(dst))
    st = {s.name: s for s in doc.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in st and theme.get("font_body"):
            st[name].font.name = theme["font_body"]
            if theme.get("text"):
                st[name].font.color.rgb = RGBColor.from_string(theme["text"].lstrip("#"))
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if name in st:
            if theme.get("font_heading"):
                st[name].font.name = theme["font_heading"]
            if theme.get("primary"):
                st[name].font.color.rgb = RGBColor.from_string(theme["primary"].lstrip("#"))
    _drop_theme_fonts(doc)
    doc.save(str(dst))


def _drop_theme_fonts(doc) -> None:
    """A pandoc-alap témabetűt (asciiTheme…) használ, ami felülírná a megadott betűt."""
    from docx.oxml.ns import qn
    for el in doc.styles.element.iter(qn("w:rFonts")):
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            el.attrib.pop(qn(a), None)


# ---------- meglévő docx átstílusozása ----------

def apply_style(target: Path, reference: Path, dst: Path) -> dict[str, Any]:
    """A minta stílusai (docDefaults, stílusok, téma) és oldalbeállítása a célra. A cél által használt, a mintából
    hiányzó stílusok megmaradnak, de betű/méret/szín nélkül (a minta Normal-jától örökölnek)."""
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    tgt, ref = Document(str(target)), Document(str(reference))
    used = {p.style.style_id for p in tgt.paragraphs if p.style is not None}
    used |= {r.style.style_id for p in tgt.paragraphs for r in p.runs if r.style is not None}
    used |= {t.style.style_id for t in tgt.tables if t.style is not None}
    new = copy.deepcopy(ref.styles.element)
    ref_ids = {s.get(qn("w:styleId")) for s in new.findall(qn("w:style"))}
    kept = []
    for s in tgt.styles.element.findall(qn("w:style")):
        sid = s.get(qn("w:styleId"))
        if sid in used and sid not in ref_ids:
            s2 = copy.deepcopy(s)
            rpr = s2.find(qn("w:rPr"))
            if rpr is not None:
                for tag in ("w:rFonts", "w:sz", "w:szCs", "w:color"):
                    for el in rpr.findall(qn(tag)):
                        rpr.remove(el)
            new.append(s2)
            kept.append(sid)
    tgt_styles_part = tgt.part.part_related_by(RT.STYLES)
    tgt_styles_part._element = new
    try:                                            # téma (a témabetűk ebből oldódnak fel)
        tgt.part.part_related_by(RT.THEME)._blob = ref.part.part_related_by(RT.THEME).blob
    except (KeyError, AttributeError):
        pass
    rs = ref.sections[0]
    for sec in tgt.sections:          # csak amit a minta ténylegesen megad (a témából épített mintában nincs oldalméret)
        for attr in ("page_width", "page_height", "orientation", "left_margin", "right_margin", "top_margin", "bottom_margin"):
            v = getattr(rs, attr)
            if v is not None:
                setattr(sec, attr, v)
    tgt.save(str(dst))
    return {"styles_from_example": len(ref_ids), "kept_own": kept}
