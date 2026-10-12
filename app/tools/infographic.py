"""Adat → infografika (VectorCraft): CSV / XLSX / táblázat-fotó → 2–4 kiemelt szám, 1–2 diagram, rövid magyarázat,
forrásmegjelölés — szerkeszthető grafika (.vectorcraft) + PNG + PDF, álló és fekvő változatban.

A SZÁMOKAT A KÓD SZÁMOLJA, nem a nyelvi modell: a táblázatból kiszámolt „tények” (összeg, átlag, legnagyobb,
legkisebb, első→utolsó változás, részarány) mind cellahivatkozással; a modell (DeepSeek-V4.1-Flash, tartalék GLM)
csak VÁLASZT közülük (fact-id), címet, rövid feliratot és számjegy nélküli magyarázatot ír. Ha mégis számot ír,
amit a táblázat nem tartalmaz, azt a mondatot kivesszük, és a check-ben jelezzük. Minden kiírt szám a check
„verified” listájában szerepel a forráscellákkal.

Diagram: matplotlib → SVG (svg.fonttype=none: a feliratok SZÖVEGKÉNT maradnak), a VectorCraft vektorként helyezi
el (mérve 10-12: egy oszlopdiagram 177 szerkeszthető objektum) — a grafikában minden átírható, átszínezhető.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

from app.config import settings

log = logging.getLogger("aice.infographic")

SIZES = {"portrait": (1080.0, 1350.0), "landscape": (1600.0, 900.0)}
# a VectorCraft BEÉPÍTETT betűje (Regular/Semibold/Bold/Italic) — a böngészős szerkesztő is pontosan ezt mutatja
# (a szerver Noto Serif-je ott hiányzó betűként, rózsaszínnel jelenik meg; mérve 10-12)
F_HEAD, F_BODY = "Source Sans 3", "Source Sans 3"
MAX_ROWS = 500
_MONTHS = ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec", "már", "ápr", "máj",
           "jún", "júl", "aug", "szep", "okt", "nov", "dec", "január", "február", "március", "április", "május",
           "június", "július", "augusztus", "szeptember", "október", "november", "december")
_HU = re.compile(r"[őűŐŰáéíóúöüÁÉÍÓÚÖÜ]")


# ---------- táblázat ----------

@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]
    refs: list[list[str]]                     # minden cella A1-hivatkozása (a táblázat SAJÁT koordinátáiban)
    sheet: str = ""
    name: str = ""

    def col(self, j: int) -> list[str]:
        return [r[j] if j < len(r) else "" for r in self.rows]


def _letter(j: int) -> str:
    s = ""
    j += 1
    while j:
        j, r = divmod(j - 1, 26)
        s = chr(65 + r) + s
    return s


def _from_grid(grid: list[list[Any]], top: int = 1, left: int = 0, sheet: str = "") -> Table:
    """Nyers cellarács → Table: az első nem üres sor a fejléc; az üres sorok/oszlopok kimaradnak (a hivatkozás az
    eredeti helyre mutat)."""
    def s(v: Any) -> str:
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v).strip()
    grid = [[s(v) for v in row] for row in grid]
    rows_idx = [i for i, r in enumerate(grid) if any(r)]
    if len(rows_idx) < 2:
        raise ValueError("The table needs a header row and at least one row of data.")
    width = max(len(r) for r in grid)
    cols_idx = [j for j in range(width) if any(j < len(grid[i]) and grid[i][j] for i in rows_idx)]
    hi = rows_idx[0]
    headers = [grid[hi][j] if j < len(grid[hi]) else "" for j in cols_idx]
    headers = [h or f"Column {_letter(left + j)}" for h, j in zip(headers, cols_idx)]
    body = rows_idx[1:MAX_ROWS + 1]
    rows = [[grid[i][j] if j < len(grid[i]) else "" for j in cols_idx] for i in body]
    refs = [[f"{_letter(left + j)}{top + i}" for j in cols_idx] for i in body]
    return Table(headers, rows, refs, sheet)


def load_table(path: Path) -> Table:
    ext = path.suffix.lower().lstrip(".")
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(str(path), data_only=True, read_only=True)
        ws = max(wb.worksheets, key=lambda w: (w.max_row or 0) * (w.max_column or 0))
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return _from_grid(grid, sheet=ws.title)
    if ext in ("csv", "tsv", "txt"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1250", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t|")
            delim = dialect.delimiter
        except csv.Error:
            delim = "\t" if ext == "tsv" else ","
        return _from_grid([r for r in csv.reader(io.StringIO(text), delimiter=delim)])
    raise ValueError("The data must be a .csv, .xlsx or a photo of a table.")


def num(s: str) -> Optional[float]:
    """Cellaszöveg → szám (magyar és angol írásmód: 1 234,5 · 1.234,5 · 1,234.5 · 12% · €1,200 · 3 500 Ft)."""
    t = str(s or "").strip().replace(" ", " ").replace(" ", " ").replace(" ", " ")
    t = re.sub(r"^[^\d\-+−.,]+|[^\d%.,]+$", "", t).replace("−", "-").replace(" ", "").rstrip("%")
    if not t or not re.fullmatch(r"[+-]?[\d.,]*\d[\d.,]*", t):
        return None
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", "") if re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+", t) else t.replace(",", ".")
    elif t.count(".") > 1:
        t = t.replace(".", "")
    try:
        return float(t)
    except ValueError:
        return None


@dataclass
class Column:
    j: int
    name: str
    values: list[Optional[float]]
    percent: bool = False


def numeric_columns(tb: Table) -> list[Column]:
    out = []
    for j, h in enumerate(tb.headers):
        cells = [c for c in tb.col(j) if c.strip()]
        vals = [num(c) for c in tb.col(j)]
        good = sum(1 for c in cells if num(c) is not None)
        if cells and good >= 0.7 * len(cells) and good >= 2:
            pct = "%" in h or sum(1 for c in cells if c.strip().endswith("%")) >= 0.5 * len(cells)
            out.append(Column(j, h, vals, pct))
    return out


def label_column(tb: Table, numeric: list[Column]) -> Optional[int]:
    nums = {c.j for c in numeric}
    for j in range(len(tb.headers)):
        if j not in nums:
            return j
    # csak számok: az első oszlop címke, ha évszám-szerű (1900–2100, egész)
    first = tb.col(0)
    if numeric and all(re.fullmatch(r"(19|20)\d\d", c.strip() or "x") for c in first if c.strip()):
        return 0
    return None


def time_like(labels: list[str]) -> bool:
    ls = [x.strip().lower() for x in labels if x.strip()]
    if len(ls) < 3:
        return False
    pat = re.compile(r"((19|20)\d\d)|(q[1-4])|(\d{4}[-./]\d{1,2})|(\d{1,2}[-./]\d{4})")
    return sum(1 for x in ls if pat.search(x) or x.rstrip(".").startswith(_MONTHS)) >= 0.8 * len(ls)


_UNIT = re.compile(r"%|\(|\bper\b|/|\brate\b|ráta|kamat|infl|index|átlag|average|price|\bár\b|arány|ratio|yield|hozam", re.I)


def indicator_table(tb: Table) -> bool:
    """„Mutató × egység" tábla: minden sor MÁS mutató (mértékegységgel: „Infláció (%)", „Kamat (%)"), az oszlopok
    országok/cégek — vagy az oszlopfejlécek évek. Ilyenkor oszloponként (egy ország „legmagasabb" értéke a
    különböző mutatói közül) értelmetlen számolni, ezért megfordítjuk: a sor = egység, az oszlop = mutató.
    (10-12, gazdasági szemle: „Legmagasabb magyar jegybanki kamat", „Amerikai kamat aránya 41%" — ráták összeadva.)"""
    if len(tb.headers) < 3 or len(tb.rows) < 2:
        return False
    first = tb.col(0)
    if sum(1 for c in first if num(c) is not None) > 0.3 * len(first):
        return False                                   # az első oszlop maga is szám: nem címkeoszlop
    rest = tb.headers[1:]
    if time_like(rest):
        return True                                    # évek az oszlopokban: idősor sorokban → fordítva
    if time_like(first):
        return False
    return sum(1 for c in first if _UNIT.search(c)) >= max(2, 0.6 * len([c for c in first if c.strip()]))


def transpose(tb: Table) -> Table:
    """Sor ↔ oszlop; a cellahivatkozások az EREDETI cellákra mutatnak (a check ugyanazokat a cellákat nevezi meg)."""
    headers = [tb.headers[0]] + [r[0] for r in tb.rows]
    rows, refs = [], []
    for j in range(1, len(tb.headers)):
        rows.append([tb.headers[j]] + [r[j] if j < len(r) else "" for r in tb.rows])
        refs.append([f"{_letter(j)}1"] + [rf[j] if j < len(rf) else "" for rf in tb.refs])
    return Table(headers, rows, refs, tb.sheet, tb.name)


def additive(c: "Column") -> bool:
    """Összeadható-e (darab, összeg) — ráta, százalék, ár, index, átlag NEM: annak „részaránya az összesből" értelmetlen."""
    return not c.percent and not _UNIT.search(c.name)


# ---------- tények (a kód számol) ----------

@dataclass
class Fact:
    id: str
    text: str                      # angol leírás a modellnek
    value: float
    display: str
    cells: str
    column: str
    kind: str
    extra: dict = field(default_factory=dict)


def fmt(v: float, lang: str, percent: bool = False, signed: bool = False) -> str:
    """Szám kiírása a nyelv szerint (hu: 1 234,5 — en: 1,234.5); százalék: 10 alatt egy tizedes."""
    a = abs(v)
    dec = 1 if (percent and a < 10) or (not percent and a < 100 and not float(v).is_integer()) else 0
    s = f"{a:,.{dec}f}"
    if lang == "hu":
        s = s.replace(",", " ").replace(".", ",")
    sign = "-" if v < 0 else ("+" if signed and v > 0 else "")
    return f"{sign}{s}{' %' if percent and lang == 'hu' else '%' if percent else ''}"


def facts_for(tb: Table, cols: list[Column], lab: Optional[int], lang: str) -> list[Fact]:
    labels = tb.col(lab) if lab is not None else [f"row {i + 1}" for i in range(len(tb.rows))]
    timey = time_like(labels)
    out: list[Fact] = []
    for c in cols[:4]:
        pts = [(i, v) for i, v in enumerate(c.values) if v is not None]
        if len(pts) < 2:
            continue
        refs = [tb.refs[i][c.j] for i, _ in pts]
        rng = f"{refs[0]}:{refs[-1]}"
        F = lambda kind, text, v, cells, pct=c.percent, signed=False, **ex: out.append(
            Fact(f"{kind}:{c.j}", text, v, fmt(v, lang, pct, signed), cells, c.name, kind, ex))
        if additive(c) and all(v >= 0 for _, v in pts):
            tot = sum(v for _, v in pts)
            F("total", f"total of {c.name} over all {len(pts)} rows", tot, f"sum of {rng}")
            if tot > 0 and not timey:
                i, v = max(pts, key=lambda t: t[1])
                F("share", f"{labels[i]} is this share of the total {c.name}", 100 * v / tot,
                  f"{tb.refs[i][c.j]} / sum of {rng}", pct=True, label=labels[i])
        F("mean", f"average {c.name} per row", sum(v for _, v in pts) / len(pts), f"average of {rng}")
        i, v = max(pts, key=lambda t: t[1])
        F("max", f"highest {c.name}: {labels[i]}", v, tb.refs[i][c.j], label=labels[i])
        i, v = min(pts, key=lambda t: t[1])
        F("min", f"lowest {c.name}: {labels[i]}", v, tb.refs[i][c.j], label=labels[i])
        if timey:
            (i0, v0), (i1, v1) = pts[0], pts[-1]
            F("last", f"latest {c.name} ({labels[i1]})", v1, tb.refs[i1][c.j], label=labels[i1])
            if v0:
                F("change", f"change of {c.name} from {labels[i0]} to {labels[i1]} (percent)",
                  100 * (v1 - v0) / abs(v0), f"({tb.refs[i1][c.j]} − {tb.refs[i0][c.j]}) / {tb.refs[i0][c.j]}",
                  pct=True, signed=True, frm=labels[i0], to=labels[i1])
            F("diff", f"change of {c.name} from {labels[i0]} to {labels[i1]} (absolute)", v1 - v0,
              f"{tb.refs[i1][c.j]} − {tb.refs[i0][c.j]}", signed=True)
    return out


# ---------- a modell: válogatás + szöveg ----------

_PROMPT = """You design a one-page infographic from a data table. The numbers are already calculated by code — you
NEVER calculate or invent numbers. Pick the most telling facts by their id and write short words around them.
{brief}
Table "{name}" ({n} rows). Columns: {headers}
First rows:
{preview}

Facts (id — what it is — value):
{facts}

Reply with ONE json object, all text in {lang}:
{{"title": "headline that states the main finding, at most 60 characters",
  "subtitle": "one line of context (what, where, when), at most 110 characters",
  "highlights": [{{"fact": "<fact id>", "label": "what this number is, at most 6 words"}}],   // 2 to 4, the most telling
  "charts": [{{"type": "column|bar|line|area|doughnut", "column": "<numeric column name>", "title": "at most 7 words"}}],  // 1 or 2
  "explanation": "two short sentences on what the data shows — WITHOUT any digits (the numbers are shown next to it)"}}
Chart choice: line/area for values over time, column for a few categories, bar for many or long category names,
doughnut only for parts of a whole (at most 6 slices). Keep the language of the table and the brief."""


def _lang(tb: Table, brief: str) -> str:
    sample = " ".join(tb.headers) + " " + " ".join(" ".join(r) for r in tb.rows[:20]) + " " + brief
    return "hu" if len(_HU.findall(sample)) >= 3 else "en"


async def _ask(tb: Table, facts: list[Fact], brief: str, lang: str, numeric: list[Column]) -> dict[str, Any]:
    from app.llm.siliconflow import _call_model, _strip_json_fences
    preview = "\n".join(" | ".join(r) for r in [tb.headers] + tb.rows[:25])
    prompt = _PROMPT.format(brief=f"The person asked: {brief.strip()[:600]}" if brief.strip() else "",
                            name=tb.name or "data", n=len(tb.rows), headers=", ".join(tb.headers),
                            preview=preview[:4000], lang={"hu": "Hungarian"}.get(lang, "English"),
                            facts="\n".join(f"{f.id} — {f.text} — {f.display}" for f in facts))
    last: Exception = RuntimeError("no answer")
    for model in (settings.AGENT_MODEL, settings.AGENT_MODEL_FALLBACK):
        try:
            raw = _strip_json_fences(await _call_model(model, [{"role": "user", "content": prompt}], thinking=False))
            d = json.loads(raw[raw.find("{"):raw.rfind("}") + 1])
            if isinstance(d, dict) and d.get("highlights"):
                return d
        except Exception as e:
            last = e
            log.warning("infografika-terv (%s): %s", model, e)
    log.warning("infografika: a modell nem válaszolt (%s) — alapértelmezett válogatás", last)
    return {}


def _allowed_numbers(tb: Table, facts: list[Fact], context: str = "") -> set[str]:
    """A szövegben megengedett számok: a táblázat cellái/fejlécei, a kiírt tények, és amit a felhasználó maga adott
    (brief, cím, forrás — pl. a „2026-10-12" dátum évszáma; 10-12: emiatt az alcímből csak „augusztus" maradt)."""
    toks: set[str] = set()
    for s in [*tb.headers, *(c for r in tb.rows for c in r), *(f.display for f in facts), context]:
        toks.update(re.findall(r"\d+", str(s)))
    return toks


def _scrub(text: str, allowed: set[str]) -> tuple[str, list[str]]:
    """Kiveszi azokat a mondatokat, amelyekben a táblázatban nem szereplő szám van."""
    kept, dropped = [], []
    for sent in re.split(r"(?<=[.!?])\s+", (text or "").strip()):
        bad = [d for d in re.findall(r"\d+", sent) if d not in allowed]
        (dropped if bad else kept).append(sent)
    return " ".join(kept).strip(), dropped


def plan(tb: Table, facts: list[Fact], numeric: list[Column], lab: Optional[int], answer: dict[str, Any],
         lang: str, context: str = "") -> dict[str, Any]:
    """A modell válasza → ellenőrzött terv (csak létező tények és oszlopok); hiányzó részre ésszerű alapérték."""
    by_id = {f.id: f for f in facts}
    allowed = _allowed_numbers(tb, facts, context)
    findings: list[str] = []
    hl = []
    for h in answer.get("highlights") or []:
        f = by_id.get(str((h or {}).get("fact")))
        if f and f not in [x[0] for x in hl]:
            hl.append((f, str(h.get("label") or f.text)[:48]))
    if len(hl) < 2:                                    # alapérték: az első oszlop legbeszédesebb tényei
        pref = ["change", "last", "total", "max", "share", "mean"]
        for kind in pref:
            for f in facts:
                if f.kind == kind and f not in [x[0] for x in hl] and len(hl) < 3:
                    hl.append((f, f.text[:48]))
    hl = hl[:4]
    # a fejléc mértékegysége a feliratba („Vendégéjszakák (ezer)” → „… (ezer)”), ha a modell kihagyta
    for k, (f, lb) in enumerate(hl):
        u = re.search(r"\(([^()]{1,24})\)\s*$", f.column)
        if u and f.kind not in ("share", "change") and u.group(1).lower() not in lb.lower():
            hl[k] = (f, f"{lb} ({u.group(1)})")
    names = {c.name: c for c in numeric}
    charts = []
    for ch in answer.get("charts") or []:
        c = names.get(str((ch or {}).get("column")))
        typ = str(ch.get("type") or "column")
        if c and typ in ("column", "bar", "line", "area", "doughnut") and c.name not in [x[1].name for x in charts]:
            charts.append((typ, c, str(ch.get("title") or c.name)[:60]))
    if not charts and numeric:
        labels = tb.col(lab) if lab is not None else []
        charts.append(("line" if time_like(labels) else "column", numeric[0], numeric[0].name))
    charts = charts[:2]
    title, d1 = _scrub(str(answer.get("title") or (numeric[0].name if numeric else tb.name)), allowed)
    subtitle, d2 = _scrub(str(answer.get("subtitle") or ""), allowed)
    expl = str(answer.get("explanation") or "")
    expl, d3 = _scrub(expl, allowed)
    for d in d1 + d2 + d3:
        findings.append(f"Removed a sentence with a number that isn't in the data: “{d[:120]}”")
    return {"title": (title or (numeric[0].name if numeric else "Infographic"))[:80], "subtitle": subtitle[:140],
            "highlights": hl, "charts": charts, "explanation": expl[:420], "findings": findings}


# ---------- diagram (SVG, a feliratok szövegként) ----------

def chart_svg(typ: str, labels: list[str], col: Column, palette: list[str], text: str, muted: str, lang: str,
              w_px: int, h_px: int, out: Path) -> list[str]:
    """Egy diagram SVG-be; visszaadja a kiírt értékfeliratokat (a check-hez)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"svg.fonttype": "none", "font.family": "DejaVu Sans", "font.size": 13})
    pts = [(lb, v) for lb, v in zip(labels, col.values) if v is not None]
    xs, ys = [str(p[0]) for p in pts], [p[1] for p in pts]
    shown: list[str] = []
    fig = plt.figure(figsize=(w_px / 100, h_px / 100), dpi=100)
    fig.patch.set_alpha(0)
    ax = fig.add_subplot(111)
    ax.patch.set_alpha(0)
    lab = lambda v: fmt(v, lang, col.percent)
    if typ == "doughnut":
        tot = sum(y for y in ys if y > 0) or 1
        wedges, _ = ax.pie([max(0, y) for y in ys], colors=(palette * 10)[:len(ys)], startangle=90, counterclock=False,
                           wedgeprops={"width": 0.38, "edgecolor": "white", "linewidth": 2})
        ax.set_aspect("equal")
        legend = [f"{x} — {fmt(100 * y / tot, lang, True)}" for x, y in zip(xs, ys)]
        ax.legend(wedges, legend, loc="center left", bbox_to_anchor=(1.0, 0.5), frameon=False, labelcolor=text)
        shown += legend
    else:
        many = len(xs) > 12
        if typ == "bar":
            pos = list(range(len(xs)))[::-1]
            bars = ax.barh(pos, ys, color=palette[0], height=0.62)
            ax.set_yticks(pos, xs)
            for b, y in zip(bars, ys):
                ax.text(b.get_width(), b.get_y() + b.get_height() / 2, "  " + lab(y), va="center", ha="left", color=text, fontsize=12)
                shown.append(lab(y))
            ax.set_xticks([])
            ax.margins(x=0.18)
        elif typ in ("line", "area"):
            ax.plot(range(len(xs)), ys, color=palette[0], linewidth=3.2, marker="o" if not many else None, markersize=6)
            if typ == "area":
                ax.fill_between(range(len(xs)), ys, color=palette[0], alpha=0.16)
            lo = ys.index(min(ys))
            for k in sorted({0, len(ys) - 1, ys.index(max(ys)), lo}):     # eleje, vége, csúcs, mélypont (az a pont ALÁ)
                below = k == lo and k not in (0, len(ys) - 1)
                ax.annotate(lab(ys[k]), (k, ys[k]), textcoords="offset points", xytext=(0, -18 if below else 10),
                            ha="center", va="top" if below else "bottom", color=text, fontsize=12, fontweight="bold")
                shown.append(lab(ys[k]))
            step = max(1, math.ceil(len(xs) / 10))
            ax.set_xticks(range(0, len(xs), step), xs[::step])
            rng = (max(ys) - min(ys)) or abs(max(ys)) or 1
            ax.set_ylim(min(ys) - 0.32 * rng, max(ys) + 0.22 * rng)     # hely a mélypont alatti feliratnak
        else:
            bars = ax.bar(range(len(xs)), ys, color=palette[0], width=0.64)
            if not many:
                for b, y in zip(bars, ys):
                    ax.text(b.get_x() + b.get_width() / 2, b.get_height(), lab(y), ha="center", va="bottom", color=text, fontsize=12)
                    shown.append(lab(y))
            step = max(1, math.ceil(len(xs) / 12))
            ax.set_xticks(range(0, len(xs), step), xs[::step], rotation=0 if max(len(x) for x in xs) < 8 else 30, ha="right" if max(len(x) for x in xs) >= 8 else "center")
            ax.margins(y=0.15)
        for side in ("top", "right", "left"):
            ax.spines[side].set_visible(False)
        ax.spines["bottom"].set_color(muted)
        ax.tick_params(colors=text, length=0)
        if typ != "bar":
            ax.set_yticks([])
    fig.tight_layout()
    fig.savefig(out, format="svg", transparent=True)
    plt.close(fig)
    # a VectorCraft beépített betűje: a feliratok a grafika betűjével jelenjenek meg
    svg = out.read_text("utf-8").replace("DejaVu Sans", F_BODY)
    out.write_text(svg, "utf-8")
    return shown


# ---------- elrendezés ----------

def _layout(d, W: float, H: float, t: dict[str, str], p: dict[str, Any], charts: list[tuple[Path, str]],
            source: str, data_name: str, landscape: bool) -> None:
    from app.tools.design import Doc  # noqa: F401  (a Doc-példány jön paraméterként)
    m = W * 0.06
    d.rect(0, 0, W, H, t["bg"])
    d.rect(0, 0, W, H * (0.2 if not landscape else 0.2), t["primary"])
    d.text(m, H * 0.035, p["title"], size=W * (0.052 if not landscape else 0.036), font=t["font_heading"], style="Bold",
           color="#FFFFFF", w=W - 2 * m, h=H * (0.11 if not landscape else 0.1), fit="shrinkText", leading=W * (0.058 if not landscape else 0.04))
    if p["subtitle"]:
        d.text(m, H * (0.148 if not landscape else 0.14), p["subtitle"], size=W * (0.022 if not landscape else 0.015),
               font=t["font_body"], color="#DDE3EC", w=W - 2 * m, h=H * 0.04, fit="shrinkText")
    d.rect(0, H * 0.2, W, max(4, H * 0.005), t["accent"])
    hl = p["highlights"]
    n = len(hl)
    if not landscape:
        top, tile_h = H * 0.235, H * 0.135
        gap = W * 0.025
        tw = (W - 2 * m - gap * (n - 1)) / max(1, n)
        for k, (f, label) in enumerate(hl):
            x = m + k * (tw + gap)
            d.rect(x, top, tw, tile_h, t["panel"])
            d.rect(x, top, tw, max(4, H * 0.004), t["accent"])
            d.text(x + tw * 0.08, top + tile_h * 0.14, f.display, size=min(W * 0.06, tw * 0.3), font=t["font_heading"],
                   style="Bold", color=t["primary"], w=tw * 0.84, h=tile_h * 0.46, fit="shrinkText")
            d.text(x + tw * 0.08, top + tile_h * 0.64, label, size=W * 0.017, font=t["font_body"], color=t["muted"],
                   w=tw * 0.84, h=tile_h * 0.32, fit="shrinkText", leading=W * 0.021)
        y = top + tile_h + H * 0.03
        ch_h = (H * 0.78 - y) / max(1, len(charts)) if charts else 0
        for path, ctitle in charts:
            d.text(m, y, ctitle, size=W * 0.022, font=t["font_body"], style="Bold", color=t["text"], w=W - 2 * m, h=H * 0.03, fit="shrinkText")
            d.image(path, m, y + H * 0.035, W - 2 * m, ch_h - H * 0.045)
            y += ch_h
        if p["explanation"]:
            d.text(m, H * 0.8, p["explanation"], size=W * 0.02, font=t["font_body"], color=t["text"], w=W - 2 * m,
                   h=H * 0.11, fit="shrinkText", leading=W * 0.028)
    else:
        lx, lw = m, W * 0.28
        top = H * 0.25
        tile_h = (H * 0.66 - top) / max(1, n) - H * 0.015 if n else 0
        tile_h = min(tile_h, H * 0.2)
        for k, (f, label) in enumerate(hl):
            y = top + k * (tile_h + H * 0.015)
            d.rect(lx, y, lw, tile_h, t["panel"])
            d.rect(lx, y, max(4, W * 0.003), tile_h, t["accent"])
            d.text(lx + lw * 0.08, y + tile_h * 0.1, f.display, size=min(W * 0.034, tile_h * 0.5), font=t["font_heading"],
                   style="Bold", color=t["primary"], w=lw * 0.86, h=tile_h * 0.52, fit="shrinkText")
            d.text(lx + lw * 0.08, y + tile_h * 0.62, label, size=W * 0.011, font=t["font_body"], color=t["muted"],
                   w=lw * 0.86, h=tile_h * 0.34, fit="shrinkText")
        if p["explanation"]:
            d.text(lx, H * 0.69, p["explanation"], size=W * 0.0115, font=t["font_body"], color=t["text"], w=lw,
                   h=H * 0.2, fit="shrinkText", leading=W * 0.016)
        cx, cw = lx + lw + W * 0.04, W - (lx + lw + W * 0.04) - m
        y = top
        ch_h = (H * 0.9 - top) / max(1, len(charts)) if charts else 0
        for path, ctitle in charts:
            d.text(cx, y, ctitle, size=W * 0.014, font=t["font_body"], style="Bold", color=t["text"], w=cw, h=H * 0.04, fit="shrinkText")
            d.image(path, cx, y + H * 0.045, cw, ch_h - H * 0.06)
            y += ch_h
    foot = " · ".join(x for x in ((f"Source: {source}" if source else ""), f"Data: {data_name}") if x)
    d.rect(0, H * 0.94, W, H * 0.06, t["panel"])
    d.text(m, H * 0.953, foot, size=W * (0.014 if not landscape else 0.0095), font=t["font_body"], color=t["muted"],
           w=W - 2 * m, h=H * 0.035, fit="shrinkText")


async def make(*, data_file_id: Optional[str] = None, table_image_file_id: Optional[str] = None,
               title: Optional[str] = None, brief: Optional[str] = None, source: Optional[str] = None,
               orientation: str = "portrait", theme: Optional[dict] = None, brand: Optional[str] = None) -> dict[str, Any]:
    """CSV/XLSX (vagy táblázatfotó) → infografika. orientation: portrait | landscape | both."""
    import tempfile
    from app.storage import display_label, output_path, public_url, write_meta
    from app.tools import crafts, design, vision_ops
    from app.tools.crafts import existing

    if not data_file_id and not table_image_file_id:
        raise ValueError("Add the data: a .csv or .xlsx file, or a photo of a table.")
    src_id = data_file_id or table_image_file_id
    ext = src_id.rsplit(".", 1)[-1].lower()
    via_photo = bool(table_image_file_id) or ext in ("jpg", "jpeg", "png", "webp", "gif", "bmp")
    if via_photo:                                  # táblázatfotó → xlsx (látómodell), utána ugyanaz az út
        x = await vision_ops.image_to_xlsx(src_id)
        data_id = x.get("xlsx_file_id") or x.get("file_id")
        write_meta(data_id, source_file_id=src_id, operation="table from photo",
                   label=f"{Path(display_label(src_id)).stem} (table).xlsx")
    else:
        data_id = src_id
    tb = load_table(existing(data_id))
    tb.name = Path(display_label(data_id)).stem
    if indicator_table(tb):
        tb = transpose(tb)
    numeric = numeric_columns(tb)
    if not numeric:
        raise ValueError("No numbers found in the table — the infographic needs at least one column of numbers.")
    lab = label_column(tb, numeric)
    lang = _lang(tb, (brief or "") + " " + (title or ""))
    facts = facts_for(tb, numeric, lab, lang)
    answer = await _ask(tb, facts, (brief or "") + (f" Title: {title}" if title else ""), lang, numeric)
    p = plan(tb, facts, numeric, lab, answer, lang, " ".join(x for x in (brief, title, source) if x))
    if title and title.strip():
        p["title"] = title.strip()[:80]
    t = {**design._theme({}, theme, brand), "font_heading": F_HEAD, "font_body": F_BODY}
    palette = [t["primary"], t["accent"], t["muted"], "#5D7A1A", "#8B1F1F", "#16A085"]
    labels = tb.col(lab) if lab is not None else [str(i + 1) for i in range(len(tb.rows))]
    orients = ["portrait", "landscape"] if orientation == "both" else [orientation if orientation in SIZES else "portrait"]

    from app.routes.make import _new_group
    group = _new_group()
    results, verified = [], []
    for f, label in p["highlights"]:
        verified.append(f"“{f.display}” ({label}) = {f.text}, from {f.cells}")
    with tempfile.TemporaryDirectory() as td:
        for orient in orients:
            W, H = SIZES[orient]
            land = orient == "landscape"
            svgs: list[tuple[Path, str]] = []
            for k, (typ, col, ctitle) in enumerate(p["charts"]):
                cw = int((W - 2 * W * 0.06) if not land else W * 0.56)
                chh = int((H * 0.78 - H * 0.4) / len(p["charts"])) if not land else int(H * 0.62 / len(p["charts"]))
                out = Path(td) / f"{orient}-{k}.svg"
                shown = chart_svg(typ, labels, col, palette[k:] + palette[:k], t["text"], t["muted"], lang,   # 2. diagram: akcentszín
                                  max(400, cw), max(260, chh), out)
                svgs.append((out, ctitle))
                if orient == orients[0]:
                    refs = [tb.refs[i][col.j] for i, v in enumerate(col.values) if v is not None]
                    verified.append(f"Chart “{ctitle}”: {len(refs)} values of {col.name} from {refs[0]}:{refs[-1]}"
                                    + (f" (labels: {', '.join(shown[:6])}{'…' if len(shown) > 6 else ''})" if shown else ""))
            d = design.Doc(W, H, p["title"])
            _layout(d, W, H, t, p, svgs, (source or "").strip(), display_label(src_id), land)
            res = await crafts.vector_run(commands=d.cmds, export=["vectorcraft", "png", "pdf"])
            stem = f"{p['title'][:60]} (infographic{', landscape' if land else ''})"
            for fobj in res["files"]:
                e = fobj["file_id"].rsplit(".", 1)[-1]
                fobj["label"] = f"{stem}.{e}"
                write_meta(fobj["file_id"], source_file_id=data_id, operation="infographic", label=fobj["label"],
                           extra={"group": group, "variant": orient})
            results.append({"orientation": orient, **res})
    files = [f for r in results for f in r["files"]]
    check = {"status": "ok" if not p["findings"] else "needs_check", "findings": p["findings"], "verified": verified}
    if via_photo:
        check["findings"].append("The numbers were read from a photo — compare them with the photo once (the table is in your files).")
        check["status"] = "needs_check"
    return {"files": files, "file_id": results[0]["file_id"], "template": "infographic", "title": p["title"],
            "highlights": [{"value": f.display, "label": lb, "cells": f.cells} for f, lb in p["highlights"]],
            "charts": [{"type": typ, "column": c.name, "title": ct} for typ, c, ct in p["charts"]],
            "data_file_id": data_id, "check": check, "variants": [r["orientation"] for r in results]}
