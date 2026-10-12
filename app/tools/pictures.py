"""A fájl képei — a KÉPVÁLASZTÓ háttere (Kommandant, 10-12: „ne szűrj ki semmit — a user választ, vagy AI generál”).

Slotok (egy fájlon belül stabil azonosítók):
  s7 / s7.2   saját pakli (a meta deck_spec-je): a 7. dia képe / galéria 2. képe — csere = a leírásban + újraépítés
  p7.1        bármilyen pptx (kézi mentés után a deck_spec elvész): a 7. dia 1. képalakzata (p:pic) — csere a fájlban
  d3          docx: a dokumentum 3. képe (a:blip, törzsben, sorrendben) — csere a fájlban, a képaláírás-kredit frissül

Csere a fájlban: a kép a keret arányára vágva ÚJ képrészként kerül be (r:embed átírva), a keret mérete/helye marad,
a régi kivágás (a:srcRect) törlődik — így nem torzul. Minden csere ÚJ verzió (operation „picture”, a revise-láncban).
"""
from __future__ import annotations

import asyncio
import copy
import hashlib
import io
import json
import logging
import re
from pathlib import Path
from typing import Any, Optional

log = logging.getLogger("aice.pictures")

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_EMBED = f"{{{_R}}}embed"
_PAGES: dict[tuple[str, float], list[int]] = {}       # docx oldaltérkép-gyorsítótár (file_id, mtime)


def _meta(fid: str) -> dict[str, Any]:
    from app.storage import read_meta
    return read_meta(fid) or {}


def _deck_refs(spec: dict[str, Any]) -> list[tuple[str, int, int, dict[str, Any]]]:
    out = []
    for i, s in enumerate(spec.get("slides") or []):
        if isinstance(s.get("image"), dict) and s["image"].get("file_id"):
            out.append((f"s{i + 1}", i, -1, s["image"]))
        if isinstance(s.get("images"), list):
            for j, im in enumerate(s["images"]):
                if isinstance(im, dict) and im.get("file_id"):
                    out.append((f"s{i + 1}.{j + 1}", i, j, im))
    return out


# ---------- a fájl képalakzatai (pptx p:pic, docx a:blip) ----------

def _pptx_pics(prs) -> list[tuple[str, int, Any, Any]]:
    """[(slot, dia-index, a:blip elem, a dia part-ja)] — diánként dokumentum-sorrendben (csoportokon belül is)."""
    out = []
    for i, slide in enumerate(prs.slides):
        k = 0
        for pic in slide.shapes._spTree.iter(f"{{{_P}}}pic"):
            blip = pic.find(f".//{{{_A}}}blip")
            if blip is None or not blip.get(_EMBED):
                continue
            k += 1
            out.append((f"p{i + 1}.{k}", i, blip, slide.part))
    return out


def _docx_pics(doc) -> list[tuple[str, int, Any]]:
    """[(slot, bekezdés-index, a:blip)] — a törzs bekezdéseiben (táblázatokon kívül), sorrendben."""
    out, k = [], 0
    for pi, p in enumerate(doc.paragraphs):
        for blip in p._p.iter(f"{{{_A}}}blip"):
            if not blip.get(_EMBED):
                continue
            k += 1
            out.append((f"d{k}", pi, blip))
    return out


def _frame_aspect(blip) -> Optional[float]:
    """A kép keretének aránya: pptx p:pic/p:spPr/a:xfrm/a:ext, docx wp:inline|wp:anchor/wp:extent."""
    el = blip
    while el is not None:
        tag = el.tag.split("}")[-1]
        if tag == "pic" and el.tag.startswith(f"{{{_P}}}"):
            ext = el.find(f"{{{_P}}}spPr/{{{_A}}}xfrm/{{{_A}}}ext")
            if ext is not None and int(ext.get("cy") or 0):
                return int(ext.get("cx")) / int(ext.get("cy"))
        if tag in ("inline", "anchor") and el.tag.startswith(f"{{{_WP}}}"):
            ext = el.find(f"{{{_WP}}}extent")
            if ext is not None and int(ext.get("cy") or 0):
                return int(ext.get("cx")) / int(ext.get("cy"))
        el = el.getparent()
    return None


def _blob(part, blip) -> bytes:
    rid = blip.get(_EMBED)          # python-pptx: related_part(rId); python-docx: related_parts[rId]
    return (part.related_part(rid) if hasattr(part, "related_part") else part.related_parts[rid]).blob


def _thumb_jpeg(data: bytes, px: int = 480) -> bytes:
    from PIL import Image, ImageOps
    im = ImageOps.exif_transpose(Image.open(io.BytesIO(data)))
    im = im.convert("RGBA")
    bg = Image.new("RGB", im.size, (255, 255, 255)); bg.paste(im, mask=im.split()[3])
    bg.thumbnail((px, px))
    out = io.BytesIO(); bg.save(out, "JPEG", quality=82)
    return out.getvalue()


# ---------- lista ----------

def _docx_pages(fid: str, src: Path, doc) -> list[int]:
    key = (fid, src.stat().st_mtime)
    if key not in _PAGES:
        from app.tools.doc_revise import page_map
        _PAGES[key] = page_map(src, [p.text for p in doc.paragraphs]) if len(doc.paragraphs) <= 4000 else []
    return _PAGES[key]


def _list_sync(file_id: str) -> list[dict[str, Any]]:
    from app.tools.crafts import existing
    src = existing(file_id)
    ext = src.suffix.lstrip(".").lower()
    m = _meta(file_id)
    if ext == "pptx" and m.get("deck_spec"):
        out = []
        for slot, i, j, ref in _deck_refs(m["deck_spec"]):
            out.append({"slot": slot, "where": f"Slide {i + 1}" + (f" · picture {j + 1}" if j >= 0 else ""),
                        "query": ref.get("query"), "prompt": ref.get("prompt"), "credit": ref.get("credit")})
        return out
    if ext == "pptx":
        from pptx import Presentation
        pics = _pptx_pics(Presentation(str(src)))
        per = {}
        for _, i, _, _ in pics:
            per[i] = per.get(i, 0) + 1
        return [{"slot": slot, "where": f"Slide {i + 1}" + (f" · picture {slot.split('.')[1]}" if per[i] > 1 else "")}
                for slot, i, _, _ in pics]
    if ext == "docx":
        from docx import Document
        doc = Document(str(src))
        pics = _docx_pics(doc)
        pages = _docx_pages(file_id, src, doc) if pics else []
        return [{"slot": slot, "where": (f"Page {pages[pi]} · picture {slot[1:]}" if pages and pages[pi] else f"Picture {slot[1:]}")}
                for slot, pi, _ in pics]
    return []


async def list_pictures(file_id: str) -> list[dict[str, Any]]:
    return await asyncio.to_thread(_list_sync, file_id)


def _image_sync(file_id: str, slot: str) -> bytes:
    """A slot képe kicsiben (a választó listájához)."""
    from app.tools.crafts import existing
    src = existing(file_id)
    ext = src.suffix.lstrip(".").lower()
    if slot.startswith("s"):
        ref = next((r for s, _, _, r in _deck_refs(_meta(file_id).get("deck_spec") or {}) if s == slot), None)
        if not ref:
            raise KeyError(slot)
        return _thumb_jpeg(existing(ref["file_id"]).read_bytes())
    if ext == "pptx":
        from pptx import Presentation
        hit = next((x for x in _pptx_pics(Presentation(str(src))) if x[0] == slot), None)
        if not hit:
            raise KeyError(slot)
        return _thumb_jpeg(_blob(hit[3], hit[2]))
    from docx import Document
    doc = Document(str(src))
    hit = next((x for x in _docx_pics(doc) if x[0] == slot), None)
    if not hit:
        raise KeyError(slot)
    return _thumb_jpeg(_blob(doc.part, hit[2]))


async def picture_image(file_id: str, slot: str) -> bytes:
    return await asyncio.to_thread(_image_sync, file_id, slot)


# ---------- mi változott (a revise után a választó erre nyílik) ----------

def _hashes_sync(file_id: str) -> dict[str, str]:
    """slot → a kép tartalmának hash-e (pakli-leírásnál: a kép file_id-je)."""
    from app.tools.crafts import existing
    src = existing(file_id)
    ext = src.suffix.lstrip(".").lower()
    m = _meta(file_id)
    if ext == "pptx" and m.get("deck_spec"):
        return {slot: ref["file_id"] for slot, _, _, ref in _deck_refs(m["deck_spec"])}
    if ext == "pptx":
        from pptx import Presentation
        return {slot: hashlib.sha1(_blob(part, blip)).hexdigest() for slot, _, blip, part in _pptx_pics(Presentation(str(src)))}
    if ext == "docx":
        from docx import Document
        doc = Document(str(src))
        return {slot: hashlib.sha1(_blob(doc.part, blip)).hexdigest() for slot, _, blip in _docx_pics(doc)}
    return {}


async def changed(old_id: str, new_id: str) -> list[dict[str, Any]]:
    """Az új verzió azon képei, amelyek a régiben nem voltak meg (új vagy kicserélt kép) — a hol/kereső adataival."""
    try:
        old, new = await asyncio.gather(asyncio.to_thread(_hashes_sync, old_id), asyncio.to_thread(_hashes_sync, new_id))
        if not new:
            return []
        seen = set(old.values())
        fresh = {s for s, h in new.items() if h not in seen}
        if not fresh:
            return []
        return [p for p in await list_pictures(new_id) if p["slot"] in fresh]
    except Exception as e:
        log.warning("kép-változás nem mérhető (%s → %s): %s", old_id, new_id, e)
        return []


# ---------- csere ----------

def _pick_spec(pick: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    """A választó választása → (images.resolve spec, a pakli-leírásba eltett adatok: query/prompt)."""
    keep = {k: pick[k] for k in ("query", "prompt") if pick.get(k)}
    if pick.get("ai"):
        return {"generate": str(pick["ai"])[:1500]}, {"prompt": str(pick["ai"])[:1500]}
    if pick.get("url"):
        return {"url": pick["url"], **({"credit": pick["credit"]} if pick.get("credit") else {})}, keep
    if pick.get("file_id"):
        return {"file_id": pick["file_id"]}, keep
    if pick.get("stock"):
        return {"stock": str(pick["stock"])}, {"query": str(pick["stock"])}
    raise ValueError("Pick a picture — a stock photo, an AI prompt or one of your images.")


def _set_picture(part, blip, img_path: Path, add_image) -> None:
    """A képet új képrészként köti a blip-hez; a keret marad, a régi kivágás (srcRect) el (a kép már a keret arányú)."""
    rid = add_image(part, img_path)
    blip.set(_EMBED, rid)
    fill = blip.getparent()
    src_rect = fill.find(f"{{{_A}}}srcRect")
    if src_rect is not None:
        fill.remove(src_rect)


def _find(src: Path, slot: str):
    """(dokumentum, part, blip, dia/bekezdés-index, a dián lévő képek száma) — a slot képalakzata a fájlban."""
    ext = src.suffix.lstrip(".").lower()
    if ext == "pptx":
        from pptx import Presentation
        prs = Presentation(str(src))
        pics = _pptx_pics(prs)
        hit = next((x for x in pics if x[0] == slot), None)
        if hit:
            return prs, hit[3], hit[2], hit[1], sum(1 for x in pics if x[1] == hit[1])
    else:
        from docx import Document
        doc = Document(str(src))
        hit = next((x for x in _docx_pics(doc) if x[0] == slot), None)
        if hit:
            return doc, doc.part, hit[2], hit[1], 1
    raise ValueError("That picture isn't in this file any more — open the newest version and try again.")


def _slot_aspect_sync(src: Path, slot: str) -> Optional[float]:
    return _frame_aspect(_find(src, slot)[2])


def _replace_in_file_sync(src: Path, dst: Path, slot: str, img: Path, credit: Optional[str]) -> None:
    """pptx/docx fájlban a slot képe → img (már a keret arányára vágva); a forrásmegjelölés frissül."""
    ext = src.suffix.lstrip(".").lower()
    doc, part, blip, idx, on_slide = _find(src, slot)
    if ext == "pptx":
        _set_picture(part, blip, img, lambda p, path: p.get_or_add_image_part(str(path))[1])
        if credit or on_slide == 1:                  # forrásmegjelölés a dia jegyzetébe (mint a saját paklinál)
            notes = doc.slides[idx].notes_slide.notes_text_frame
            lines = [ln for ln in (notes.text or "").split("\n") if not (on_slide == 1 and ln.startswith("Photo:"))]
            notes.text = "\n".join([*filter(None, lines), *([credit] if credit else [])])
    else:
        _set_picture(part, blip, img, lambda p, path: p.get_or_add_image(str(path))[0])
        _docx_credit(doc, idx, credit)
    doc.save(str(dst))


def _docx_credit(doc, pi: int, credit: Optional[str]) -> None:
    """A kép utáni képaláírásban a „Photo: …” forrásmegjelölés cseréje (ha nincs aláírás és van kredit: új, kicsi, dőlt)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    from app.tools.doc_revise import _insert_after
    pars = doc.paragraphs
    cap = pars[pi + 1] if pi + 1 < len(pars) else None
    if cap is not None and "Photo:" in cap.text:
        head = cap.text.split("Photo:")[0].rstrip(" —-")
        text = " — ".join(x for x in (head, credit) if x)
        for r in cap.runs[1:]:
            r._r.getparent().remove(r._r)
        if cap.runs:
            cap.runs[0].text = text
        else:
            cap.add_run(text)
        if not text:                                  # se aláírás, se kredit: az üres bekezdés el
            cap._p.getparent().remove(cap._p)
        return
    if credit:
        p = _insert_after(pars[pi])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(credit)
        run.italic, run.font.size = True, Pt(8.5)


async def replace(file_id: str, slot: str, pick: dict[str, Any]) -> dict[str, Any]:
    """A slot képe a választott képre → ÚJ verzió {file_id, url, label, where}."""
    from app.storage import display_label, output_path, public_url, write_meta
    from app.storage.files import new_file_id
    from app.tools import images
    from app.tools.crafts import existing
    from app.tools.revise import history
    src = existing(file_id)
    ext = src.suffix.lstrip(".").lower()
    spec, keep = _pick_spec(pick)
    where = next((p["where"] for p in await list_pictures(file_id) if p["slot"] == slot), None)
    if where is None:
        raise ValueError("That picture isn't in this file any more — open the newest version and try again.")
    m = _meta(file_id)
    if slot.startswith("s") and ext == "pptx" and m.get("deck_spec"):
        from app.tools import pptx_deck
        deck = copy.deepcopy(m["deck_spec"])
        hit = next((x for x in _deck_refs(deck) if x[0] == slot), None)
        _, i, j, _ = hit
        new_img = {**spec, **keep}
        if j < 0:
            deck["slides"][i]["image"] = new_img
        else:
            deck["slides"][i]["images"][j] = new_img
        out_id = (await pptx_deck.build(deck))["file_id"]
    elif ext in ("pptx", "docx"):
        aspect = await asyncio.to_thread(_slot_aspect_sync, src, slot)
        r = await images.resolve(spec, aspect=aspect, max_px=2000)      # AI-nál itt vonódik le a kredit
        credit = r.attribution if (spec.get("url") or spec.get("stock")) else (
            _meta(spec["file_id"]).get("credit") if spec.get("file_id") else None)
        out_id = new_file_id(ext)
        try:
            await asyncio.to_thread(_replace_in_file_sync, src, output_path(out_id), slot, r.path, credit)
        finally:
            r.path.unlink(missing_ok=True)
    else:
        raise ValueError("Pictures can be changed in decks and Word documents.")
    stem = Path(display_label(file_id)).stem.split(" (")[0]
    label = f"{stem} (v{len(history(file_id)) + 1}).{out_id.rsplit('.', 1)[1]}"
    how = "an AI picture" if spec.get("generate") else "your picture" if spec.get("file_id") else "a stock photo"
    write_meta(out_id, source_file_id=file_id, operation="picture", label=label,
               extra={"summary": f"New picture on {where.split(' · ')[0].lower()}: {how}.", "picture_slot": slot})
    return {"file_id": out_id, "url": public_url(out_id), "label": label, "where": where,
            "slot": slot}          # a slot az új verzióban is ugyanaz (a leírás / a képalakzatok sorrendje marad)


# ---------- „Kérj változtatást”: ha a kérés lényegében EGY kép cseréje, nem kell ügynök ----------

_SWAP_PROMPT = """A user asked for a change to a file. Decide if the request is ONLY about replacing one existing picture
with a different picture (nothing else changes). Pictures in the file:
{pictures}
Request: {request}
Reply with a json object: {{"swap": true|false, "slot": "<slot of the picture to replace>", "query": "2-4 English search
words for a stock photo, main noun last", "prompt": "a detailed English prompt for an AI picture of the same thing",
"ai": true|false}}. "ai" is true only if the user asks for an AI / generated / made-up picture. If the request changes
anything else too (text, layout, colours, adds a new picture) → "swap": false."""


async def detect_swap(file_id: str, request: str) -> Optional[dict[str, Any]]:
    """{slot, query, ai} ha a kérés egy kép cseréje (pl. „on page 7 change the picture to a shiny blonde girl”), különben None."""
    if not re.search(r"\b(pic|picture|photo|image|img|kép|fotó)", request, re.I):
        return None
    pics = await list_pictures(file_id)
    if not pics:
        return None
    from app.config import settings
    from app.llm.siliconflow import _call_model, _strip_json_fences
    listing = "\n".join(f"- {p['slot']}: {p['where']}" for p in pics[:80])
    prompt = _SWAP_PROMPT.format(pictures=listing, request=request.strip()[:600])
    for model in (settings.AGENT_MODEL, settings.AGENT_MODEL_FALLBACK):
        try:
            raw = _strip_json_fences(await _call_model(model, [{"role": "user", "content": prompt}], thinking=False))
            d = json.loads(raw[raw.find("{"):raw.rfind("}") + 1])
            break
        except Exception as e:
            log.warning("képcsere-felismerés (%s): %s", model, e)
    else:
        return None
    if not d.get("swap") or d.get("slot") not in {p["slot"] for p in pics}:
        return None
    return {"slot": d["slot"], "query": (d.get("query") or "").strip() or request[:60],
            "prompt": (d.get("prompt") or "").strip() or request, "ai": bool(d.get("ai"))}
