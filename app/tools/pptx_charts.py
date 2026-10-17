"""A videó mozgó grafikonjai a pakliban (10-17): ugyanaz a leírás, mint a videóé (app/tools/motion_charts.py) — egy kliens
AI ugyanazt az objektumot adhatja a videónak és a diának. Natív, szerkeszthető PowerPoint-tartalom lesz belőle:

  waterfall → halmozott oszlopdiagram: átlátszó alap + növekedés / csökkenés / kezdő- és végösszeg, előjeles feliratok
  lines     → a meglévő vonaldiagram (az időszakok a kategóriák, jelmagyarázattal)
  race      → az utolsó időszak rangsora vízszintes sávokon (a verseny mozgása csak a videóban)
  slope, band, scatter → natív alakzatok (vonal, pont, sáv, buborék, felirat), mint a magyarázó ábráknál

Miért alakzat a három (10-17-i mérés, mindkét motorban): a szerkesztő motorja (DeckCraft) vonaldiagramon semmilyen
adatfeliratot nem rajzol, a tengely min/max-ot nem veszi figyelembe, a terület + vonal kombót vonalaknak, a buborékot
vonalnak, a pontdiagramot kategóriáknak rajzolja (az x egyenközű, a pontnév nincs). A lejtőn a nevek és az értékek, a
sávon a sáv és a szaggatott előrejelzés, a pontdiagramon a hely és a név maga az üzenet — diagramként a szerkesztőben
névtelen vonalak, egy hamis „Low” vonal, rossz helyen álló névtelen pontok maradtak volna. Alakzatként a szerkesztő, a
PowerPoint, a LibreOffice és a Keynote ugyanazt mutatja (a buborékméret is megmarad). A vízesésnél a szerkesztő a
sorozatfeliratokat nem mutatja, de a sávok, a tengely és a rács ott is helyesen olvasható (az átlátszó alapot alfa 0-val
kapja: a „nincs kitöltés” sorozatot a DeckCraft színnel rajzolja). A régi pontdiagram (series/points) natív marad.
"""
from __future__ import annotations

import colorsys
import math
from typing import Any, Callable, Optional

from app.tools import motion_charts

HINTS = {
    "waterfall": '{"items": [{"label", "value"}…] — a start, then + / − changes (3+ items), "total_label"?, "total"?}',
    "slope": '{"labels": [before, after], "items": [{"label", "from", "to"}…]}',
    "scatter": '{"x_title", "y_title", "items": [{"label", "x", "y", "size"?}…] (3+ points)}',
    "band": '{"items": [{"label": period, "value", "low"?, "high"?}…] (3+ points, some with low–high), "forecast_from"?}',
    "lines": '{"labels": [periods], "series": [{"name", "values"}…] (2–4 series)}',
    "race": '{"labels": [periods] (3–12), "items": [{"label", "values"}…] (3–10 items)}',
}
GREEN = "#2E9E6B"


# ---------- a dia leírása → a videó ábrája ----------

def _obj(s: dict[str, Any]) -> tuple[dict[str, Any], str]:
    """A dia ábra-mezői: a dián közvetlenül (chart_type + a videó mezői) VAGY a videó objektuma a "chart" kulcson."""
    c = s.get("chart")
    if isinstance(c, dict):
        return c, str(c.get("type") or s.get("chart_type") or "")
    return s, str(s.get("chart_type") or "")


def kind_of(s: dict[str, Any]) -> Optional[str]:
    """A dia mozgó-grafikon típusa (a videó álneveivel együtt) — None: a meglévő natív diagram (column, line, scatter
    [x, y, név] pontokkal…)."""
    c, raw = _obj(s)
    typ = raw.strip().lower()
    typ = motion_charts.ALIASES.get(typ, typ)
    if typ not in motion_charts.TYPES:
        return None
    if typ == "scatter" and not c.get("items"):       # a pakli régi pontdiagramja (series/points) marad a régi úton
        return None
    return typ


def _source(title: str, c: dict[str, Any]) -> str:
    """A dia saját adatai mint „forrás”: a pakliban a számot a kliens adja, nincs mihez mérni — így a videó
    ellenőrzője (motion_charts.parse) csak az alakot nézi és ugyanúgy normalizál (a nyelvet a szavakból veszi)."""
    words, nums = [str(title or "")], []

    def walk(x: Any) -> None:
        if isinstance(x, dict):
            for v in x.values():
                walk(v)
        elif isinstance(x, (list, tuple)):
            for v in x:
                walk(v)
        else:
            if isinstance(x, str):
                words.append(x)
            v = motion_charts._num(x)
            if v is not None and math.isfinite(v):
                nums.append(v)
    walk(c)
    return " · ".join(words) + " · " + " ; ".join(f"{abs(v):.6f}".rstrip("0").rstrip(".") for v in nums)


def normalize(s: dict[str, Any], n: Optional[int] = None) -> tuple[dict[str, Any], list[str], str]:
    """→ (az ábra a videó alakjában, figyelmeztetések, nyelv). ValueError érthető üzenettel, ha nem rajzolható."""
    from app.tools.video import _lang_of
    typ = kind_of(s)
    c, _ = _obj(s)
    where = f"Slide {n} ({typ})" if n else f"The {typ} chart"
    src = _source(s.get("title") or "", c)
    out, notes = motion_charts.parse(typ, {**c, "title": c.get("title") or s.get("title")}, src)
    raw = [str(x) for x in (notes if isinstance(notes, list) else [notes] if notes else [])]
    if out is None:
        why = next((m for m in raw if "not all its numbers" not in m), None)
        why = why.replace("⚠ ", "").split(": ", 1)[-1].rstrip(".") if why else "a value is missing or not a number"
        raise ValueError(f"{where}: {why} — give {HINTS[typ]}.")
    # csak a figyelmeztetés kell (a „…in the source” megerősítés itt értelmetlen), pl. a vízesés megadott végösszege
    # nem egyezik, a pontdiagramnak nincs tengelycíme
    warn = [m.replace("⚠ ", "") for m in raw if m.startswith("⚠")]
    return out, warn, _lang_of(src)


def check_slides(spec: dict[str, Any]) -> tuple[list[str], list[str]]:
    """Építés ELŐTT: minden ilyen dia rajzolható-e (különben gyors, érthető hiba: melyik dia, mit adjon) →
    (figyelmeztetések a check-be, tájékoztatás: mi marad a videónak)."""
    warn, info = [], []
    for i, s in enumerate(spec.get("slides") or [], 1):
        if not (isinstance(s, dict) and s.get("layout") == "chart" and kind_of(s)):
            continue
        ch, w, _ = normalize(s, i)
        warn += [f"Slide {i}: {x}" for x in w]
        name = ch.get("title") or s.get("title") or ""
        if ch["type"] == "race":
            info.append(f"Slide {i}: the ranking race “{name}” shows its last period ({ch['labels'][-1]}) as a sorted bar "
                        "chart — the race itself moves only in the video.")
    return warn, info


# ---------- közös segédek ----------

def _fmt(ch: dict[str, Any], vals: list[float], lang: str) -> Callable[[float], str]:
    """A videó számformája: egy ábrán belül egyforma tizedes, az egységgel és a nyelv tizedesjelével."""
    from app.tools.video import _decimals, _fmt_fixed
    dec = max([_decimals(v) for v in vals] + [0])
    return lambda v: _fmt_fixed(v, dec, ch.get("unit") or "", lang)


def _numfmt(dec: int, unit: str, lang: str, sign: Optional[str] = None) -> str:
    """Excel-számformátum ugyanazzal a formával (előjel, tizedes, egység) — a felirat így az ÉRTÉKHEZ kötött: a
    PowerPointban átírt adat a feliratot is viszi. sign: up → mindig „+”, down → mindig „−” (a sáv magassága pozitív)."""
    body = "#,##0" + ("." + "0" * dec if dec else "")
    u = (unit or "").replace('"', "")
    if u == "%":
        num = body + ('" %"' if lang == "hu" else '"%"')
    elif u in ("$", "€", "£", "¥") and lang == "en":
        num = f'"{u}"' + body
    else:
        num = body + (f'" {u}"' if u else "")
    if sign == "up":
        return f'"+"{num};"+"{num};{num}'
    if sign == "down":
        return f'"−"{num};"−"{num};{num}'
    return f'{num};"−"{num};{num}'


def _hue_near_red(hexv: str) -> bool:
    h = hexv.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    hue, light, sat = colorsys.rgb_to_hls(r, g, b)
    return sat > 0.3 and (hue < 25 / 360 or hue > 335 / 360)


def colors(t: dict[str, str]) -> dict[str, str]:
    """Növekedés = akcentus (ha az akcentus maga piros, zöld — különben a fel és a le egyforma lenne), csökkenés =
    a háttéren is látszó piros, összeg = a fő szín (ha a háttéren nem látszik, a szövegszín)."""
    from app.tools.pptx_deck import contrast
    bg = "#" + (t.get("bg") or "#FFFFFF").lstrip("#")
    acc = "#" + (t.get("accent") or "#FCA311").lstrip("#")
    red = next((c for c in ("#D64541", "#F0706A", "#9E2A24") if contrast(c, bg) >= 3), "#D64541")
    prim = "#" + (t.get("primary") or t.get("text") or "#14213D").lstrip("#")
    text = "#" + (t.get("text") or "#1E2430").lstrip("#")
    return {"up": GREEN if _hue_near_red(acc) else acc, "down": red, "total": prim if contrast(prim, bg) >= 1.6 else text,
            "text": text, "muted": "#" + (t.get("muted") or "#6B7280").lstrip("#"), "bg": bg, "accent": acc}


def _plain(shp) -> None:
    """Saját stílus: a témából semmi (árnyék, alapvonal) — a kitöltés és a vonal kifejezetten megadva."""
    from pptx.oxml.ns import qn
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)


def _alpha(fill, a: float) -> None:
    from pptx.oxml.ns import qn
    clr = fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(a * 100000))}))


def _line(slide, name: str, pts: list[tuple[float, float]], color: str, width: float, dash: bool = False):
    """Egyenes (2 pont: összekötő) vagy törött vonal (szabadkézi, kitöltés nélkül), hüvelykben."""
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
    from app.tools.pptx_deck import _rgb
    if len(pts) == 2:
        shp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(pts[0][0]), Inches(pts[0][1]),
                                         Inches(pts[1][0]), Inches(pts[1][1]))
    else:
        E = lambda v: Emu(int(round(v * 914400)))     # noqa: E731
        fb = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
        fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=False)
        shp = fb.convert_to_shape()
        shp.fill.background()
        for p in shp._element.iter(qn("a:path")):      # nyitott vonal: kitöltés nélkül (a DeckCraft is így rajzolja)
            p.set("fill", "none")
    shp.name = name
    _plain(shp)
    shp.line.color.rgb = _rgb(color)
    shp.line.width = Pt(width)
    if dash:
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shp


def _polygon(slide, name: str, pts: list[tuple[float, float]], color: str, alpha: float):
    from pptx.util import Emu
    from app.tools.pptx_deck import _rgb
    E = lambda v: Emu(int(round(v * 914400)))         # noqa: E731
    fb = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=True)
    shp = fb.convert_to_shape()
    shp.name = name
    _plain(shp)
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    _alpha(shp.fill, alpha)
    shp.line.fill.background()
    return shp


def _dot(slide, name: str, x: float, y: float, d: float, color: str):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    from app.tools.pptx_deck import _rgb
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - d / 2), Inches(y - d / 2), Inches(d), Inches(d))
    shp.name = name
    _plain(shp)
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    shp.line.fill.background()
    return shp


def _label(slide, name: str, x: float, y: float, w: float, h: float, text: str, **kw):
    """Egysoros felirat (nem tördel: a doboz a becsült szélességnél bővebb) — a pptx_deck szövegdoboza, névvel."""
    from app.tools.pptx_deck import _text
    tb = _text(slide, x, y, w, h, text, anchor="middle", **kw)
    tb.text_frame.word_wrap = False
    tb.name = name
    return tb


def _em(text: str, size: float, k: float = 0.6) -> float:
    """A szöveg becsült szélessége hüvelykben (félkövérrel is bőven)."""
    return len(text) * k * size / 72


# ---------- vízesés: natív halmozott oszlopdiagram ----------

def _waterfall(slide, ch: dict[str, Any], s: dict[str, Any], t: dict[str, str], box, lang: str) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
    from pptx.util import Inches, Pt
    from app.tools.pptx_deck import _axis_titles, _rgb, _wrap_categories, contrast, dlbl_numfmt
    from app.tools.video import _decimals
    steps = motion_charts.waterfall_steps(ch)
    col = colors(t)
    n = len(steps)
    base, up, down, tot, up_neg, down_neg = ([None] * n for _ in range(6))
    split: dict[int, str] = {}                      # a nullát átlépő oszlopok: két részből, a felirat a nagyobbon
    for k, (_, a, b, kind) in enumerate(steps):
        top, bot = max(a, b), min(a, b)
        if kind == "total":
            tot[k] = b
            continue
        main, neg = (up, up_neg) if kind == "up" else (down, down_neg)
        if bot >= 0:                                # a nulla fölött: átlátszó alap + a változás
            base[k], main[k] = bot, top - bot
        elif top <= 0:                              # alatta: a negatív oszlop is lefelé halmozódik a nullától
            base[k], main[k] = top, bot - top
        else:                                       # átlépi: a pozitív rész a fő sorozatban, a negatív külön
            main[k], neg[k] = top, bot
            split[k] = kind
    cd = CategoryChartData()
    cd.categories = _wrap_categories([x[0] for x in steps], XL_CHART_TYPE.COLUMN_CLUSTERED)
    series = [("Base (hidden)", base, None, None), ("Increase", up, col["up"], "up"), ("Decrease", down, col["down"], "down"),
              ("Total", tot, col["total"], None)]
    series += [x for x in (("Increase below zero", up_neg, col["up"], ""), ("Decrease below zero", down_neg, col["down"], ""))
               if any(v is not None for v in x[1])]
    for name, vals, _, _ in series:
        cd.add_series(name, vals)
    x, y, w, h = box
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = f"Waterfall – {ch.get('title') or s.get('title') or ''}".rstrip(" –")
    c = gf.chart
    c.font.size, c.font.name = Pt(14), t.get("font_body")
    c.has_title, c.has_legend = False, False
    plot = c.plots[0]
    plot.gap_width, plot.overlap = 55, 100
    lo = min(0.0, *[min(a, b) for _, a, b, _ in steps])
    hi = max(0.0, *[max(a, b) for _, a, b, _ in steps])
    vlo, vhi, step = motion_charts._nice(lo, hi)
    va = c.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = vlo, vhi, step   # a LibreOffice különben 90-nél kezdené
    va.has_major_gridlines = True
    c.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW     # negatív oszlopnál se a nullára írja
    _axis_titles(c, {"x_title": s.get("x_title"), "y_title": s.get("y_title") or ch.get("unit")}, XL_CHART_TYPE.COLUMN_STACKED)
    nums = [st[2] for st in steps] + [it["value"] for it in ch["items"]]
    dec = max(_decimals(v) for v in nums)           # egy ábrán egyforma tizedes (mint a videóban)
    plot_h = h * 72 - 60                            # a rajzterület becsült magassága (pt): a tengelyfeliratok nélkül
    pt_per = plot_h / ((vhi - vlo) or 1)
    slot = (w * 72 - 70) / n                        # egy oszlop helye és a sáv szélessége (pt)
    bar_w = slot / (1 + plot.gap_width / 100)
    # sok oszlopnál kisebb betű; ha a felirat így sem fér az oszlop helyére, az egység csak a tengelycímben marad
    unit = ch.get("unit") or ""
    for unit in ([unit, ""] if unit else [""]):
        fv = _fmt({**ch, "unit": unit}, nums, lang)
        texts = [fv(b) if kind == "total" else ("+" if b >= a else "−") + fv(abs(b - a)) for _, a, b, kind in steps]
        size = next((z for z in (13, 12, 11) if max(_em(x, z) for x in texts) * 72 <= slot - 6), 0)
        if size:
            break
    size = size or 11
    lab_h = size * 1.25 + 6
    lab_w = [_em(x, size) * 72 for x in texts]

    def own_text(dl, txt: str, color: str) -> None:
        dl.text_frame.text = txt
        r = dl.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold = Pt(size), True
        r.font.color.rgb = _rgb(color)
        dl.position = XL_LABEL_POSITION.INSIDE_END

    for ser, (name, vals, color, sign) in zip(plot.series, series):
        if color is None:                           # az alap: átlátszó (alfa 0) — a „nincs kitöltés” sorozatot a
            ser.format.fill.solid()                 # DeckCraft a paletta színével rajzolja, az alfát mindenki érti
            ser.format.fill.fore_color.rgb = _rgb(t.get("bg") or "#FFFFFF")
            _alpha(ser.format.fill, 0)
            ser.format.line.fill.background()
            continue
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = _rgb(color)
        ser.format.line.fill.background()
        ser.invert_if_negative = False
        if sign == "":                              # a nullát átlépő oszlop alsó része: felirat nélkül
            continue
        inside = "#FFFFFF" if contrast("#FFFFFF", color) >= contrast("#1B1F24", color) else "#1B1F24"
        dl = ser.data_labels                        # sorozatszintű felirat: az értékhez kötve, a videó előjelével
        dl.show_value = True
        dl.number_format, dl.number_format_is_linked = _numfmt(dec, unit, lang, sign), False
        dl.position = XL_LABEL_POSITION.INSIDE_END  # halmozott oszlopnál a „kívül” érvénytelen (a PowerPoint javítaná)
        dl.font.size, dl.font.bold = Pt(size), True
        dl.font.color.rgb = _rgb(inside)
        _leader_lines_off(dl._element)
        for k, v in enumerate(vals):
            if v is None:
                continue
            a, b = steps[k][1], steps[k][2]
            bar = abs(b - a) * pt_per
            short = bar < lab_h or lab_w[k] > bar_w - 6
            p = ser.points[k].data_label
            if k in split:                          # a teljes változás (a két rész összege) egy feliratban — ritka
                own_text(p, texts[k], col["text"] if short else inside)
            elif short:
                # alacsony / keskeny sáv: a felirat nem fér bele — fölé (növekedés, pozitív összeg) vagy alá (csökkenés), a háttéren
                # olvasható színnel, az ÉRTÉKHEZ kötve (egyedi szöveg helyett: ami a szöveget nem ismeri, pl. a Keynote, ott
                # is a valódi szám látszik). A kiindulópont a sáv közepe — minden motorban ugyanaz.
                p.font.size, p.font.bold = Pt(size), True
                p.font.color.rgb = _rgb(col["text"])
                p.position = XL_LABEL_POSITION.CENTER
                down = sign == "down" if sign else b < 0
                _offset(p._dLbl, (bar / 2 + size * 0.75 + 3) * (1 if down else -1) / (h * 72))
                dlbl_numfmt(p._dLbl, _numfmt(dec, unit, lang, sign))


def _offset(dLbl, dy: float) -> None:
    """A pont feliratának függőleges eltolása az alaphelyéről (a diagram magasságának arányában, lefelé pozitív)."""
    from lxml import etree
    from pptx.oxml.ns import qn
    old = dLbl.find(qn("c:layout"))
    if old is not None:
        dLbl.remove(old)
    lay = etree.Element(qn("c:layout"))
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    etree.SubElement(ml, qn("c:x")).set("val", "0")
    etree.SubElement(ml, qn("c:y")).set("val", f"{dy:.4f}")
    dLbl.find(qn("c:idx")).addnext(lay)                  # a séma sorrendje: idx, layout?, tx?, numFmt?…


def _leader_lines_off(dLbls) -> None:
    """Az eltolt felirathoz a LibreOffice vezetővonalat húzna a sáv közepéből — nem kell."""
    from pptx.oxml.ns import qn
    el = dLbls.find(qn("c:showLeaderLines"))
    if el is not None:
        el.set("val", "0")


# ---------- előtte → utána: natív alakzatok ----------

def _slope(slide, ch: dict[str, Any], t: dict[str, str], box, lang: str) -> None:
    from app.tools import diagrams
    from app.tools.motion_charts import _spread
    items, labs = ch["items"], ch["labels"]
    col = colors(t)
    font = t.get("font_body")
    vals = [v for it in items for v in (it["from"], it["to"])]
    fv = _fmt(ch, vals, lang)
    lo, hi = min(vals), max(vals)
    pad = (hi - lo) * 0.08 or abs(hi) * 0.1 or 1
    lo, hi = lo - pad, hi + pad
    size = 16 if len(items) <= 6 else 14
    lh = size * 1.35 / 72
    x0, y0, w, h = box
    top, bottom = y0 + 0.85, y0 + h - 0.2
    val_w = max(_em(fv(v), size) for v in vals) + 0.15
    name_w = min(max(_em(it["label"], size, 0.56) for it in items) + 0.15, w * 0.3)
    gap = 0.18
    left_w, right_w = name_w + val_w + 2 * gap, val_w + gap
    span = min(w - left_w - right_w, 7.0)
    xa = x0 + (w - (left_w + span + right_w)) / 2 + left_w
    xb = xa + span
    Y = lambda v: bottom - (v - lo) / (hi - lo) * (bottom - top)   # noqa: E731
    axis = diagrams.mix(col["muted"], col["bg"], 0.55)
    for nm, x, lab in (("Axis A", xa, labs[0]), ("Axis B", xb, labs[1])):
        _line(slide, nm, [(x, top - 0.15), (x, bottom + 0.05)], axis, 1.0)
        _label(slide, f"Heading – {lab}", x - 1.4, y0 + 0.1, 2.8, 0.5, lab, size=18, font=font, color=col["text"], bold=True,
               align="center")
    la = _spread([Y(it["from"]) for it in items], lh, top, bottom)
    lb = _spread([Y(it["to"]) for it in items], lh, top, bottom)
    for k, it in enumerate(items):
        c = col["up"] if it["to"] >= it["from"] else col["down"]
        _line(slide, f"Change – {it['label']}", [(xa, Y(it["from"])), (xb, Y(it["to"]))], c, 3.0)
    for k, it in enumerate(items):
        c = col["up"] if it["to"] >= it["from"] else col["down"]
        _dot(slide, f"From – {it['label']}", xa, Y(it["from"]), 0.16, c)
        _dot(slide, f"To – {it['label']}", xb, Y(it["to"]), 0.16, c)
        _label(slide, f"Name – {it['label']}", xa - gap - val_w - name_w, la[k] - lh / 2, name_w, lh, it["label"], size=size,
               font=font, color=col["muted"], align="right")
        _label(slide, f"From value – {it['label']}", xa - gap - val_w, la[k] - lh / 2, val_w, lh, fv(it["from"]), size=size,
               font=font, color=col["text"], bold=True, align="right")
        _label(slide, f"To value – {it['label']}", xb + gap, lb[k] - lh / 2, val_w, lh, fv(it["to"]), size=size, font=font,
               color=diagrams.readable(c, col["bg"], 3.0), bold=True, align="left")


# ---------- sáv-előrejelzés: natív alakzatok ----------

def _band(slide, ch: dict[str, Any], t: dict[str, str], box, lang: str) -> None:
    from app.tools import diagrams
    from app.tools.video import _decimals, _fmt_fixed
    it = ch["items"]
    n = len(it)
    labs, vals = [i["label"] for i in it], [i["value"] for i in it]
    col = colors(t)
    font = t.get("font_body")
    acc = col["accent"]
    rngv = [v for i in it for v in (i.get("low"), i.get("high")) if v is not None]
    fv = _fmt(ch, vals + rngv, lang)
    dec = max([_decimals(v) for v in vals + rngv] + [0])
    ff = ch.get("forecast_from")
    kf = labs.index(ff) if ff in labs else n          # az első előrejelzett időszak
    last = max(0, min(kf, n) - 1)                     # az utolsó mért pont: innen szaggatott
    lo, hi = min(vals + rngv), max(vals + rngv)
    pad = (hi - lo) * 0.1 or abs(hi) * 0.1 or 1
    lo, hi = (0.0 if lo >= 0 and lo - pad < 0 else lo - pad), hi + pad
    size, small = 16, 14
    lh = size * 1.35 / 72
    ks = [k for k, i in enumerate(it) if i.get("low") is not None]
    k0, k1 = ks[0], ks[-1]
    if k0 > 0 and k0 >= kf:                           # a sáv az utolsó mért pontból indul (nem a semmiből)
        k0 -= 1
    rng_txt = f"{_fmt_fixed(it[k1]['low'], dec, '', lang)}–{fv(it[k1]['high'])}"
    end_txt = fv(vals[-1])
    x0, y0, w, h = box
    right_w = max(_em(end_txt, size), _em(rng_txt, small) if k1 == n - 1 else 0) + 0.3
    px0 = x0 + max(0.4, _em(fv(vals[0]), size) / 2 + 0.1)
    px1 = x0 + w - right_w
    top, bottom = y0 + 0.8, y0 + h - 0.5
    X = lambda k: px0 + (px1 - px0) * k / max(n - 1, 1)                 # noqa: E731
    Y = lambda v: bottom - (v - lo) / ((hi - lo) or 1) * (bottom - top)  # noqa: E731
    grey = diagrams.mix(col["muted"], col["bg"], 0.45)
    _line(slide, "Baseline", [(px0 - 0.15, bottom), (px1 + 0.15, bottom)], grey, 1.0)
    lab_w = max(_em(x, small, 0.56) for x in labs) + 0.2
    every = max(1, math.ceil(lab_w / max((px1 - px0) / max(n - 1, 1), 1e-6)))
    shown = [k for k in range(n) if k % every == 0]
    if shown[-1] != n - 1:                            # az utolsó időszak mindig látszik (az előző kimarad, ha ütközne)
        if (n - 1 - shown[-1]) < every:
            shown.pop()
        shown.append(n - 1)
    for k in shown:
        _label(slide, f"Period – {labs[k]}", X(k) - lab_w / 2, bottom + 0.08, lab_w, 0.32, labs[k], size=small, font=font,
               color=col["muted"], align="center")
    # a sáv a vonal mögött: felső él balról jobbra, alsó vissza (mért pontnál a sáv a vonalra húzódik)
    if k1 > k0:
        upper = [(X(k), Y(it[k].get("high", it[k]["value"]))) for k in range(k0, k1 + 1)]
        lower = [(X(k), Y(it[k].get("low", it[k]["value"]))) for k in range(k1, k0 - 1, -1)]
        _polygon(slide, "Range band", upper + lower, acc, 0.26)
    else:                                             # egyetlen pont tartománnyal: függőleges sáv
        _line(slide, "Range band", [(X(k1), Y(it[k1]["high"])), (X(k1), Y(it[k1]["low"]))], diagrams.mix(acc, col["bg"], 0.5), 6.0)
    if kf < n:
        _line(slide, "Forecast line", [(X(last), top - 0.45), (X(last), bottom)], grey, 1.0, dash=True)
        word = "előrejelzés" if lang == "hu" else "forecast"
        _label(slide, "Forecast", X(last) + 0.1, top - 0.6, _em(word, small, 0.56) + 0.3, 0.32, word, size=small, font=font,
               color=col["muted"], italic=True)
    if last >= 1:
        _line(slide, "Line", [(X(k), Y(vals[k])) for k in range(0, last + 1)], acc, 3.0)
    if last < n - 1:
        _line(slide, "Forecast", [(X(k), Y(vals[k])) for k in range(last, n)], acc, 3.0, dash=True).name = "Forecast path"
    for k in sorted({0, last, n - 1}):
        px, py = X(k), Y(vals[k])
        _dot(slide, f"Point – {labs[k]}", px, py, 0.15, acc)
        txt = fv(vals[k])
        tw = _em(txt, size) + 0.2
        if k == n - 1:                                # a vég jobbra a ponttól — a sáv vége alatta/fölötte marad
            _label(slide, f"Value – {labs[k]}", px + 0.14, py - lh / 2, tw, lh, txt, size=size, font=font, color=col["text"],
                   bold=True)
        elif k == last and k > 0:                     # az utolsó mért pont: balra fent (jobbra a választóvonal és felirata)
            _label(slide, f"Value – {labs[k]}", px - 0.06 - tw, py - 0.08 - lh, tw, lh, txt, size=size, font=font,
                   color=col["text"], bold=True, align="right")
        else:
            _label(slide, f"Value – {labs[k]}", px - tw / 2, py - 0.1 - lh, tw, lh, txt, size=size, font=font, color=col["text"],
                   bold=True, align="center")
    rw = _em(rng_txt, small) + 0.25
    if k1 == n - 1:                                   # a tartomány a végérték alatt
        _label(slide, "Range", X(k1) + 0.14, Y(vals[-1]) + lh / 2, rw, small * 1.35 / 72, rng_txt, size=small, font=font,
               color=col["muted"])
    else:
        _label(slide, "Range", X(k1) - rw / 2, Y(it[k1]["high"]) - 0.08 - small * 1.35 / 72, rw, small * 1.35 / 72, rng_txt,
               size=small, font=font, color=col["muted"], align="center")


# ---------- pont- / buborékdiagram: natív alakzatok ----------

def _scatter(slide, ch: dict[str, Any], t: dict[str, str], box, lang: str) -> None:
    from app.tools import diagrams
    from app.tools.video import _fmt_val
    items = ch["items"]
    col = colors(t)
    font = t.get("font_body")
    tick, size = 12, 13
    x0, y0, w, h = box
    xl, xh, xs = motion_charts._nice(min(i["x"] for i in items), max(i["x"] for i in items))
    yl, yh, ys = motion_charts._nice(min(i["y"] for i in items), max(i["y"] for i in items))
    ticks = lambda lo, hi, st: [lo + k * st for k in range(int(round((hi - lo) / st)) + 1)]   # noqa: E731
    ytw = max(_em(_fmt_val(v, "", lang), tick) for v in ticks(yl, yh, ys)) + 0.15
    px0, px1 = x0 + ytw + 0.1, x0 + w - 0.3
    py0, py1 = y0 + 0.55, y0 + h - 0.75
    X = lambda v: px0 + (v - xl) / ((xh - xl) or 1) * (px1 - px0)   # noqa: E731
    Y = lambda v: py1 - (v - yl) / ((yh - yl) or 1) * (py1 - py0)   # noqa: E731
    grid, axis = diagrams.mix(col["muted"], col["bg"], 0.78), diagrams.mix(col["muted"], col["bg"], 0.35)
    th = tick * 1.35 / 72
    for v in ticks(yl, yh, ys):
        if v > yl:
            _line(slide, f"Grid y {v:g}", [(px0, Y(v)), (px1, Y(v))], grid, 0.75)
        txt = _fmt_val(v, "", lang)
        _label(slide, f"Tick y {v:g}", px0 - 0.08 - ytw, Y(v) - th / 2, ytw, th, txt, size=tick, font=font, color=col["muted"],
               align="right")
    for v in ticks(xl, xh, xs):
        txt = _fmt_val(v, "", lang)
        tw = _em(txt, tick) + 0.2
        _label(slide, f"Tick x {v:g}", X(v) - tw / 2, py1 + 0.06, tw, th, txt, size=tick, font=font, color=col["muted"],
               align="center")
    _line(slide, "X axis", [(px0, py1), (px1, py1)], axis, 1.0)
    _line(slide, "Y axis", [(px0, py0), (px0, py1)], axis, 1.0)
    if ch.get("x_title"):
        _label(slide, "X title", px0, py1 + 0.08 + th, px1 - px0, 0.34, ch["x_title"], size=14, font=font, color=col["muted"],
               align="center")
    if ch.get("y_title"):
        yt = "↑ " + ch["y_title"]
        _label(slide, "Y title", px0 - 0.05, y0 + 0.05, min(_em(yt, 14) + 0.3, px1 - px0), 0.34, yt, size=14, font=font,
               color=col["muted"])
    # buborék: a terület arányos a mérettel (mint a videóban); a nagy alul, a kicsi fölötte látszik
    smax = max((i.get("size") or 0) for i in items)
    dmin, dmax = 0.16, min(0.95, (py1 - py0) * 0.24)
    D = [dmin + (dmax - dmin) * math.sqrt(i["size"] / smax) if smax and i.get("size") else dmin for i in items]
    for i in sorted(range(len(items)), key=lambda i: -D[i]):
        it = items[i]
        shp = _dot(slide, f"Point – {it['label'] or i + 1}", X(it["x"]), Y(it["y"]), D[i], col["accent"])
        if smax:
            _alpha(shp.fill, 0.8)
    # nevek: ne takarjanak pontot és egymást — a videó (és a régi pakli-pontdiagram) közös elhelyezője, pontban
    order = sorted(range(len(items)), key=lambda i: items[i]["x"])
    P = [(X(items[i]["x"]) * 72, Y(items[i]["y"]) * 72) for i in order]
    lh = size * 1.35 / 72
    sides = motion_charts.label_sides(P, [items[i]["label"] for i in order], (px0 * 72, py0 * 72 - 10, px1 * 72, py1 * 72),
                                      0.56 * size, size * 1.3, dot=max(D) * 36 if smax else dmin * 36)
    for i, side in zip(order, sides):
        it = items[i]
        if not it["label"]:
            continue
        px, py, g = X(it["x"]), Y(it["y"]), D[i] / 2 + 0.06
        lw = _em(it["label"], size, 0.56) + 0.2
        bx, by = {"right": (px + g, py - lh / 2), "left": (px - g - lw, py - lh / 2), "above": (px - lw / 2, py - g - lh),
                  "below": (px - lw / 2, py + g)}[side]
        _label(slide, f"Name – {it['label']}", bx, by, lw, lh, it["label"], size=size, font=font, color=col["text"],
               align={"right": "left", "left": "right"}.get(side, "center"))


# ---------- belépési pont ----------

def draw(slide, s: dict[str, Any], t: dict[str, str], chart_box, shape_box) -> Optional[dict[str, Any]]:
    """A dia ábrája: a vízesést, a lejtőt, a sávot és a pontdiagramot itt rajzolja (→ None); a vonalakat és a versenyt a
    meglévő natív diagram mezőire fordítja (→ a dia új leírása, amit a hívó natív diagramja rajzol tovább)."""
    ch, _, lang = normalize(s)
    typ = ch["type"]
    if typ == "waterfall":
        _waterfall(slide, ch, s, t, chart_box, lang)
        return None
    drawn = {"slope": _slope, "band": _band, "scatter": _scatter}.get(typ)
    if drawn:
        drawn(slide, ch, t, shape_box, lang)
        return None
    unit = ch.get("unit") or ""
    if typ == "lines":
        return {**s, "chart_type": "line", "categories": ch["labels"],
                "series": [{"name": x["label"], "values": x["values"]} for x in ch["items"]],
                "y_title": s.get("y_title") or unit}
    # race: az utolsó időszak rangsora (a legnagyobb felül), az értékek a sávokon
    labs, items = ch["labels"], ch["items"]
    order = sorted(items, key=lambda i: -i["values"][-1])
    return {**s, "chart_type": "bar", "categories": [i["label"] for i in order],
            "series": [{"name": labs[-1], "values": [i["values"][-1] for i in order]}], "data_labels": True,
            "x_title": s.get("x_title") or (f"{labs[-1]}, {unit}" if unit else labs[-1])}
