"""Egész PPTX-pakli EGY hívásból: 16:9, egységes dizájn, képekkel (file_id / URL / FLUX / stock).

    spec = {
      "title": "…", "footer": "…"?, "brand": "<brand név>"? | "theme": {primary, accent, bg, text, muted,
               font_heading, font_body}?,
      "slides": [ {"layout": "cover", "title", "subtitle"?, "image"?}, … ]
    }

Elrendezések: cover · section · bullets (+image) · image_left · image_right · full_image · quote (+image)
· stats · chart · table · gallery · closing. Minden diához "notes" (előadói jegyzet) adható; a stockképek
forrásmegjelölése automatikusan a jegyzetbe kerül.

A betűk alapból Cambria / Calibri: a PowerPointban megvannak, a szerveren a metrikailag egyező
Caladea / Carlito rajzolja az előnézetet, így a tördelés mindkét helyen ugyanaz.
"""
from __future__ import annotations

import time
from pathlib import Path
from typing import Any, Optional

from app.storage import output_path, public_url, write_meta
from app.storage.files import new_file_id
from app.tools import images

LAYOUTS = {"cover", "section", "bullets", "image_left", "image_right", "full_image", "quote",
           "stats", "chart", "table", "gallery", "closing"}
DEFAULT_THEME = {
    "primary": "#14213D", "accent": "#FCA311", "bg": "#FFFFFF", "panel": "#F4F5F7",
    "text": "#1E2430", "muted": "#6B7280", "font_heading": "Cambria", "font_body": "Calibri",
}
W_IN, H_IN, M_IN = 13.333, 7.5, 0.6
_FULL = W_IN / H_IN


async def _brand_kit(spec: dict[str, Any]) -> tuple[dict[str, Any], Optional[dict[str, Any]]]:
    """→ (spec, a pakli arculata). Ugyanaz a feloldás, mint a dizájnnál és a videónál (app.brands.user.resolve): a
    felhasználó mentett arculata név/„mine” szerint, az alapértelmezett automatikusan (ha be van kapcsolva),
    "none" = ki; a beépített katalógus a tartalék."""
    theme = spec.get("theme")
    if isinstance(theme, str):                       # „theme": "executive" — név, nem színek: arculatként próbáljuk
        if spec.get("brand"):
            raise ValueError(f'"theme" is an object of colors, e.g. {{"primary": "#1F3A5F", "accent": "#D9822B"}} '
                             f'(keys: {", ".join(DEFAULT_THEME)}); for a saved brand use "brand": "<name>"')
        spec = {**spec, "brand": theme, "theme": None}
    elif theme is not None and not isinstance(theme, dict):
        raise ValueError('"theme" must be an object of colors, e.g. {"primary": "#1F3A5F", "accent": "#D9822B"}')
    from app.brands import user as user_brands
    try:
        kit = await user_brands.resolve(spec.get("brand"))
    except ValueError:
        if isinstance(theme, str):
            raise ValueError(f'"theme": {theme!r} is not a brand — give colors, e.g. {{"primary": "#1F3A5F", '
                             f'"accent": "#D9822B"}} (keys: {", ".join(DEFAULT_THEME)})') from None
        raise
    return spec, kit


def _theme(spec: dict[str, Any], kit: Optional[dict[str, Any]] = None) -> dict[str, str]:
    t = dict(DEFAULT_THEME)
    if kit and kit.get("_registry"):                 # beépített katalógus-arculat: a pontos színei és betűi
        from app.brands import get_brand
        b = get_brand(kit["_registry"])
        t.update(primary=b.colors.primary, accent=b.colors.accent, bg=b.colors.body_bg, panel=b.colors.panel_bg,
                 text=b.colors.text_primary, muted=b.colors.text_muted,
                 font_heading=b.fonts.heading, font_body=b.fonts.body)
    elif kit:                                        # a felhasználó mentett arculata — ugyanúgy, mint a dizájnon
        from app.tools.design import apply_brand
        t = apply_brand(t, kit)
    t.update({k: v for k, v in (spec.get("theme") or {}).items() if k in DEFAULT_THEME and v})
    return t


def _image_aspect(s: dict[str, Any]) -> Optional[float]:
    lay = s.get("layout")
    if lay in ("cover", "full_image", "section", "closing"):
        return _FULL
    if lay == "bullets":
        return 5.2 / H_IN
    if lay in ("image_left", "image_right"):
        return (W_IN / 2) / H_IN
    if lay == "quote":
        return 4.6 / H_IN
    return None


def _requests(slides: list[dict[str, Any]]) -> list[tuple[tuple[int, int], Any, Optional[float]]]:
    reqs = []
    for i, s in enumerate(slides):
        if s.get("layout") == "gallery":
            for j, img in enumerate((s.get("images") or [])[:4]):
                reqs.append(((i, j), img, 1.5))
        elif s.get("image"):
            reqs.append(((i, -1), s["image"], _image_aspect(s)))
    return reqs


async def _pick_stock(reqs: list[tuple[tuple[int, int], Any, Optional[float]]],
                      slides: list[dict[str, Any]]) -> list[tuple[tuple[int, int], Any, Optional[float]]]:
    """Stockkép a látómodell rangsorával, a dia címe a kontextus (10-13: a kereső első találata egy Deutsche Post-bélyeg
    és egy „Lorem Ipsum” mockup volt). Sorrend, nem szűrés; két dia nem kapja ugyanazt; ami nem sikerül, a régi úton megy."""
    jobs = [(slot, spec["stock"], str(slides[slot[0]].get("title") or ""), aspect) for slot, spec, aspect in reqs
            if isinstance(spec, dict) and isinstance(spec.get("stock"), str) and spec["stock"].strip()]
    if not jobs:
        return reqs
    chosen = await images.choose_stock(jobs, medium="slide")
    return [(slot, {**spec, "url": chosen[slot]["url"], "credit": chosen[slot]["credit"]} if slot in chosen else spec, aspect)
            for slot, spec, aspect in reqs]


# ---------- rajzoló segédek (python-pptx) ----------

def _rgb(hexstr: str):
    from pptx.dml.color import RGBColor
    h = hexstr.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ---------- a pakli arculata a fájlban: ami később kerül bele, az örökli ----------
# A ChatGPT-teszt (10-17): zöld–krém paklihoz adott videós dia FEHÉR lett — a háttér diánként volt festve, az új dia
# az elrendezés (fehér) hátterét és a Calibrit kapta. Két általános javítás: (1) a pptx_build a mesterdiára és a
# témába is beírja a hátteret, a betűket és a színeket — így a PowerPointban/DeckCraftban/nálunk később hozzáadott
# dia is ilyen; (2) egy MEGLÉVŐ (bárki által készített) paklihoz adott dia a szomszédja hátterét és címstílusát kapja.

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def bake_theme(prs, t: dict[str, str]) -> None:
    """A téma (háttér, címsor- és szövegbetű, szöveg-/akcentusszín) a mesterdián és a theme1.xml-ben."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    m = prs.slide_master
    m.background.fill.solid()
    m.background.fill.fore_color.rgb = _rgb(t["bg"])
    try:
        tp = m.part.part_related_by(RT.THEME)
        root = etree.fromstring(tp.blob)
        for tag, font in (("majorFont", t["font_heading"]), ("minorFont", t["font_body"])):
            el = root.find(f".//{{{_A}}}fontScheme/{{{_A}}}{tag}/{{{_A}}}latin")
            if el is not None and font:
                el.set("typeface", font)
        for slot, hexv in (("dk1", t["text"]), ("dk2", t["primary"]), ("accent1", t["accent"]), ("accent2", t["primary"])):
            el = root.find(f".//{{{_A}}}clrScheme/{{{_A}}}{slot}")
            if el is not None and hexv:
                for ch in list(el):
                    el.remove(ch)
                etree.SubElement(el, f"{{{_A}}}srgbClr", val=hexv.lstrip("#").upper())
        tp._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    except Exception:          # a téma kényelmi — a háttér a mesterdián így is megvan
        pass


def _bg(slide):
    return slide._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}bg")


def inherit_look(slide, ref) -> bool:
    """Az új dia (saját háttér nélkül) a ref dia hátterét kapja — a képes háttér képével együtt. → átvette-e."""
    from copy import deepcopy
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    if ref is None or _bg(slide) is not None or _bg(ref) is None:
        return False
    bg = deepcopy(_bg(ref))
    r_embed = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    for blip in bg.iter(f"{{{_A}}}blip"):
        rid = blip.get(r_embed)
        if rid:
            blip.set(r_embed, slide.part.relate_to(ref.part.related_part(rid), RT.IMAGE))
    slide._element.cSld.insert(0, bg)
    return True


def title_style(ref) -> dict[str, Any]:
    """A ref dia címének stílusa (a legnagyobb betűs szöveg): betű, méret, szín, félkövér — az új dia címéhez."""
    best: dict[str, Any] = {}
    for shp in getattr(ref, "shapes", []):
        if not shp.has_text_frame:
            continue
        for para in shp.text_frame.paragraphs:
            for r in para.runs:
                size = r.font.size.pt if r.font.size else None
                if r.text.strip() and size and size > best.get("size", 0):
                    color = None
                    try:
                        color = str(r.font.color.rgb) if r.font.color is not None and r.font.color.type is not None else None
                    except Exception:
                        pass
                    best = {"size": size, "font": r.font.name, "color": color, "bold": r.font.bold}
    return best


def _lum(hexv: str) -> float:
    h = hexv.lstrip("#")
    c = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def contrast(a: str, b: str) -> float:
    la, lb = sorted((_lum(a), _lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def bg_color(slide) -> Optional[str]:
    """A dia tényleges egyszínű háttere (saját → elrendezés → mesterdia), ha egyszínű."""
    for el in (slide, getattr(slide, "slide_layout", None), getattr(getattr(slide, "slide_layout", None), "slide_master", None)):
        bg = _bg(el) if el is not None else None
        if bg is not None:
            clr = bg.find(f".//{{{_A}}}solidFill/{{{_A}}}srgbClr")
            return clr.get("val") if clr is not None else None
    return None


def theme_text_color(prs) -> Optional[str]:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    try:
        import re as _re
        xml = prs.slide_master.part.part_related_by(RT.THEME).blob.decode("utf-8", "replace")
        m = _re.search(r"<a:dk1>.*?(?:val|lastClr)=\"([0-9A-Fa-f]{6})\"", xml, _re.S)
        return m.group(1) if m else None
    except Exception:
        return None


def readable(color: Optional[str], slide, prs) -> Optional[str]:
    """A cím színe csak akkor öröklődik, ha az új dia hátterén olvasható (≥ 3:1); különben a téma szövegszíne,
    végső esetben fekete/fehér — fehér cím krém háttéren nem maradhat."""
    bg = bg_color(slide) or "FFFFFF"
    if color and contrast(color, bg) >= 3:
        return color
    dk = theme_text_color(prs)
    if dk and contrast(dk, bg) >= 3:
        return dk
    return "000000" if _lum(bg) > 0.4 else "FFFFFF"


def neighbour(prs, pos: int):
    """A pakli „tartalmi” mintadiája az új dia mellett: előbb a pos előtti, különben az utáni."""
    slides = list(prs.slides)
    for i in [pos - 1, pos + 1, *range(len(slides) - 1, -1, -1)]:
        if 0 <= i < len(slides) and i != pos:
            return slides[i]
    return None


def inherit_new_slides(before: Optional[Path], after: Path) -> int:
    """Szerkesztés után (DeckCraft slide.new…): a BESZÚRT diák, amelyeknek nincs saját hátterük, a szomszédjukét
    kapják. Beszúrt = a két dia-sorrend összevetésében új (a módosított meglévő dia nem; annak háttere marad).
    → hány dia kapott hátteret."""
    import difflib
    import hashlib
    from pptx import Presentation
    from app.tools import content_digest
    if before is None or before.suffix.lower() != ".pptx":
        return 0
    try:
        a = [hashlib.sha1(repr(x).encode()).hexdigest() for x in content_digest._pptx(before)[1]]
        b = [hashlib.sha1(repr(x).encode()).hexdigest() for x in content_digest._pptx(after)[1]]
    except Exception:
        return 0
    added = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op == "insert":
            added += range(j1, j2)
        elif op == "replace" and j2 - j1 > i2 - i1:
            added += range(j1 + (i2 - i1), j2)
    if not added:
        return 0
    prs = Presentation(str(after))
    slides = list(prs.slides)
    if len(slides) != len(b):
        return 0
    n = 0
    for i in added:
        if inherit_look(slides[i], neighbour(prs, i)):
            n += 1
    if n:
        prs.save(str(after))
    return n


def _rect(slide, x, y, w, h, color: str, alpha: Optional[float] = None):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if alpha is not None:  # átlátszóság: <a:alpha> a srgbClr alá (a python-pptx nem ad rá API-t)
        clr = shp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
        clr.append(a)
    return shp


def _text(slide, x, y, w, h, text, *, size, font, color, bold=False, italic=False, align="left",
          anchor="top", spacing=1.1, tracking=None):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    for n, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        f = r.font
        f.size, f.name, f.bold, f.italic, f.color.rgb = Pt(size), font, bold, italic, _rgb(color)
    return tb


def _bullets(slide, x, y, w, h, items, t, size):
    from pptx.util import Inches, Pt
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for n, item in enumerate(items):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.line_spacing = 1.15
        p.space_after = Pt(size * 0.7)
        dot = p.add_run(); dot.text = "■  "
        dot.font.size, dot.font.color.rgb, dot.font.name = Pt(size * 0.55), _rgb(t["accent"]), t["font_body"]
        r = p.add_run(); r.text = str(item)
        r.font.size, r.font.name, r.font.color.rgb = Pt(size), t["font_body"], _rgb(t["text"])


def _picture(slide, path, x, y, w, h):
    from pptx.util import Inches
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def _title_size(text: str, big: int = 36) -> int:
    n = len(text or "")
    return big if n <= 32 else big - 6 if n <= 55 else big - 10


def _bullet_size(items: list) -> int:
    longest = max((len(str(i)) for i in items), default=0)
    if len(items) <= 4 and longest <= 70:
        return 22
    return 19 if len(items) <= 6 else 16


def _heading(slide, s, t, x=M_IN, w=W_IN - 2 * M_IN):
    title = s.get("title") or ""
    _text(slide, x, 0.55, w, 1.0, title, size=_title_size(title), font=t["font_heading"], color=t["primary"],
          bold=True, anchor="bottom")
    _rect(slide, x, 1.68, 0.9, 0.07, t["accent"])


# ---------- elrendezések ----------

def _cover(slide, s, t, img):
    if img:
        _picture(slide, img.path, 0, 0, W_IN, H_IN)
        # a kép látsszon: enyhe teljes sötétítés + a cím mögött sötétebb sáv (olvashatóság)
        _rect(slide, 0, 0, W_IN, H_IN, "#000000", alpha=0.18)
        _rect(slide, 0, 1.35, W_IN * 0.62, 4.6, t["primary"], alpha=0.62)
    else:
        _rect(slide, 0, 0, W_IN, H_IN, t["primary"])
    _rect(slide, M_IN + 0.05, 4.05, 1.2, 0.09, t["accent"])
    title = s.get("title") or ""
    text_w = (W_IN * 0.62 - M_IN - 0.35) if img else (W_IN - 2 * M_IN)      # képes címdián a sötét panel szélessége
    _text(slide, M_IN, 1.6, text_w, 2.35, title, size=_title_size(title, 54), font=t["font_heading"],
          color="#FFFFFF", bold=True, anchor="bottom", spacing=1.0)
    if s.get("subtitle"):
        _text(slide, M_IN, 4.35, text_w if img else W_IN - 2.5 * M_IN, 1.4, s["subtitle"], size=22, font=t["font_body"],
              color="#F1F1F1")
    if s.get("kicker"):
        _text(slide, M_IN, 1.6, 7, 0.5, s["kicker"].upper(), size=14, font=t["font_body"], color=t["accent"], bold=True)


def _section(slide, s, t, img, closing=False):
    if img:
        _picture(slide, img.path, 0, 0, W_IN, H_IN)
        _rect(slide, 0, 0, W_IN, H_IN, t["primary"], alpha=0.78)
    else:
        _rect(slide, 0, 0, W_IN, H_IN, t["primary"])
    if s.get("number") is not None and not closing:
        _text(slide, M_IN, 1.5, 4, 1.4, f"{int(s['number']):02d}", size=72, font=t["font_heading"],
              color=t["accent"], bold=True)
    title = s.get("title") or ("Thank you!" if closing else "")
    _text(slide, M_IN, 2.9, W_IN - 2 * M_IN, 1.6, title, size=_title_size(title, 46), font=t["font_heading"],
          color="#FFFFFF", bold=True, anchor="bottom", align="center" if closing else "left")
    if s.get("subtitle"):
        _text(slide, M_IN, 4.7, W_IN - 2 * M_IN, 1.2, s["subtitle"], size=20, font=t["font_body"], color="#E8E8E8",
              align="center" if closing else "left")


def _bullets_layout(slide, s, t, img):
    items = s.get("bullets") or []
    text_w = (W_IN - 5.2 - 2 * M_IN + 0.2) if img else (W_IN - 2 * M_IN)
    _heading(slide, s, t, w=text_w)
    if s.get("text"):
        _text(slide, M_IN, 2.05, text_w, 1.0, s["text"], size=18, font=t["font_body"], color=t["muted"])
    _bullets(slide, M_IN, 2.15 + (0.95 if s.get("text") else 0), text_w, 4.6, items, t, _bullet_size(items))
    if img:
        _picture(slide, img.path, W_IN - 5.2, 0, 5.2, H_IN)


def _image_side(slide, s, t, img, left: bool):
    half = W_IN / 2
    if img:
        _picture(slide, img.path, 0 if left else half, 0, half, H_IN)
    x = half + M_IN if left else M_IN
    w = half - 2 * M_IN
    _heading(slide, s, t, x=x, w=w)
    if s.get("bullets"):
        _bullets(slide, x, 2.15, w, 4.6, s["bullets"], t, min(20, _bullet_size(s["bullets"])))
    elif s.get("text"):
        _text(slide, x, 2.15, w, 4.6, s["text"], size=19, font=t["font_body"], color=t["text"], spacing=1.25)


def _full_image(slide, s, t, img):
    if img:
        _picture(slide, img.path, 0, 0, W_IN, H_IN)
    # A sötét sáv a cím sorainak számával nő (10-12: egy háromsoros idézet első sora kilógott a világos égre)
    title = s.get("title") or ""
    size = 30 if len(title) <= 130 else 24
    cpl = max(20, int((W_IN - 2 * M_IN) * 72 / (size * 0.55)))          # becsült betű/sor (talpas betű, félkövér)
    lines = max(1, -(-int(len(title) * 1.08) // cpl)) if title else 0   # +8%: a szótördelés vesztesége
    line_h = size * 1.22 / 72
    band = max(1.9, 0.95 + lines * line_h + 0.25)
    _rect(slide, 0, H_IN - band, W_IN, band, "#000000", alpha=0.55)
    if title:
        _text(slide, M_IN, H_IN - 0.95 - lines * line_h - 0.1, W_IN - 2 * M_IN, lines * line_h + 0.1, title, size=size,
              font=t["font_heading"], color="#FFFFFF", bold=True, anchor="bottom")
    if s.get("caption"):
        _text(slide, M_IN, H_IN - 0.9, W_IN - 2 * M_IN, 0.6, s["caption"], size=15, font=t["font_body"], color="#E5E5E5")


def _quote(slide, s, t, img):
    x0 = 4.6 + M_IN if img else 1.4
    if img:
        _picture(slide, img.path, 0, 0, 4.6, H_IN)
    _text(slide, x0 - 0.05, 0.55, 1.5, 1.3, "„", size=110, font=t["font_heading"], color=t["accent"], bold=True)
    q = s.get("quote") or ""
    _text(slide, x0, 2.45, W_IN - x0 - M_IN - (0 if img else 0.8), 3.1, q,
          size=34 if len(q) < 120 else 26, font=t["font_heading"], color=t["primary"], italic=True, spacing=1.15)
    if s.get("author"):
        _text(slide, x0, 5.7, W_IN - x0 - M_IN, 0.6, "— " + s["author"], size=18, font=t["font_body"], color=t["muted"])


def _stats(slide, s, t):
    _heading(slide, s, t)
    stats = (s.get("stats") or [])[:4]
    if not stats:
        return
    col = (W_IN - 2 * M_IN) / len(stats)
    for n, st in enumerate(stats):
        x = M_IN + n * col
        _rect(slide, x + 0.05, 2.55, col - 0.3, 3.6, t["panel"])
        _rect(slide, x + 0.05, 2.55, 0.08, 3.6, t["accent"])
        val = str(st.get("value", ""))
        _text(slide, x + 0.35, 2.9, col - 0.8, 1.5, val, size=48 if len(val) <= 7 else 36, font=t["font_heading"],
              color=t["primary"], bold=True, anchor="bottom")
        _text(slide, x + 0.35, 4.55, col - 0.8, 1.4, st.get("label", ""), size=17, font=t["font_body"], color=t["muted"])


def _wrap_categories(cats: list[Any], kind) -> list[str]:
    """Hosszú kategórianév sortöréssel (10-13: a DeckCraft a sávdiagram feliratainak fix, ~94 pt-os sávot hagy, a kézi
    rajzterületet nem veszi figyelembe → a hosszú név a dia bal szélén levágódott; sortöréssel mindkét motor jól rajzol)."""
    import textwrap
    from pptx.enum.chart import XL_CHART_TYPE
    out = [str(c) for c in cats]
    if kind == XL_CHART_TYPE.BAR_CLUSTERED:
        width = 16
    elif kind in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.AREA) and out:
        width = max(10, int((W_IN - 2 * M_IN) * 72 / len(out) / (14 * 0.56)))      # egy oszlopnyi hely, 14 pt-os betűvel
    else:
        return out
    return ["\n".join(textwrap.wrap(c, width)) if len(c) > width else c for c in out]


def _chart(slide, s, t):
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
    _heading(slide, s, t)
    kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA,
            "scatter": XL_CHART_TYPE.XY_SCATTER}.get(s.get("chart_type", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
    series = s.get("series") or []
    if kind == XL_CHART_TYPE.XY_SCATTER:
        data = XyChartData()
        for se in series:
            ser = data.add_series(se.get("name", "Sorozat"))
            for x, y in se.get("points") or []:
                ser.add_data_point(float(x), float(y))
    else:
        data = CategoryChartData()
        data.categories = _wrap_categories(s.get("categories") or [], kind)
        for se in series:
            data.add_series(se.get("name", "Sorozat"), se.get("values", []))
    gf = slide.shapes.add_chart(kind, Inches(M_IN), Inches(2.0), Inches(W_IN - 2 * M_IN), Inches(5.0), data)
    ch = gf.chart
    ch.font.size, ch.font.name = Pt(14), t["font_body"]
    if kind == XL_CHART_TYPE.BAR_CLUSTERED:
        # PowerPoint alulról felfelé rajzolja a sávokat, a DeckCraft (szerkesztő, PDF) felülről lefelé (10-13) → a
        # kategóriatengely fordítva, az értéktengely a maximumnál keresztez (lent marad): mindkettőben az első van felül
        from pptx.oxml.ns import qn
        ch.category_axis.reverse_order = True
        vcr = ch._chartSpace.chart.plotArea.find(qn("c:valAx")).find(qn("c:crosses"))
        if vcr is not None:
            vcr.set("val", "max")
    ch.has_title = False  # a dia címe a cím; a sorozatnévből gyártott diagramcím zavaró
    ch.has_legend = len(series) > 1 or kind in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT)
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    pal = [t["primary"], t["accent"], "#2E7D6B", "#8E44AD", "#C0392B", "#7F8C8D"]
    plot = ch.plots[0]
    if kind in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT):
        for k, pt in enumerate(plot.series[0].points if plot.series else []):
            pt.format.fill.solid(); pt.format.fill.fore_color.rgb = _rgb(pal[k % len(pal)])
        return
    for k, ser in enumerate(plot.series):
        c = _rgb(pal[k % len(pal)])
        if kind in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.XY_SCATTER):
            from pptx.enum.chart import XL_MARKER_STYLE
            ser.marker.style, ser.marker.size = XL_MARKER_STYLE.CIRCLE, 9
            ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = c
            ser.marker.format.line.color.rgb = c
            if kind == XL_CHART_TYPE.XY_SCATTER:
                ser.format.line.fill.background()
            else:
                ser.format.line.color.rgb, ser.format.line.width = c, Pt(3)
        else:
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = c


def _table(slide, s, t):
    from pptx.util import Inches, Pt
    _heading(slide, s, t)
    headers, rows = s.get("headers") or [], s.get("rows") or []
    ncols = max([len(headers)] + [len(r) for r in rows] + [1])
    shape = slide.shapes.add_table(len(rows) + (1 if headers else 0), ncols, Inches(M_IN), Inches(2.1),
                                   Inches(W_IN - 2 * M_IN), Inches(0.5 * (len(rows) + 1)))
    tbl = shape.table
    size = 16 if len(rows) <= 8 else 13
    all_rows = ([headers] if headers else []) + rows
    for r, row in enumerate(all_rows):
        for c in range(ncols):
            cell = tbl.cell(r, c)
            cell.text = str(row[c]) if c < len(row) else ""
            head = bool(headers) and r == 0
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(t["primary"] if head else (t["panel"] if r % 2 == 0 else t["bg"]))
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size, run.font.name, run.font.bold = Pt(size), t["font_body"], head
                    run.font.color.rgb = _rgb("#FFFFFF" if head else t["text"])


def _gallery(slide, s, t, imgs):
    _heading(slide, s, t)
    caps = s.get("captions") or []
    n = max(1, len(imgs))
    gap = 0.3
    w = (W_IN - 2 * M_IN - gap * (n - 1)) / n
    h = min(w / 1.5, 4.3)
    for k, img in enumerate(imgs):
        x = M_IN + k * (w + gap)
        _picture(slide, img.path, x, 2.2, w, h)
        if k < len(caps):
            _text(slide, x, 2.3 + h, w, 0.7, caps[k], size=14, font=t["font_body"], color=t["muted"])


def _footer(slide, n, total, t, footer: Optional[str], dark: bool, left: float = M_IN, right: float = W_IN - M_IN):
    """left–right: a lábléc sávja a dia szabad részén (10-13: image_left-nél a 8" széles szöveg lelógott a diáról,
    jobb oldali képnél az oldalszám a fotóra került)."""
    color = "#D0D4DA" if dark else t["muted"]
    _text(slide, right - 1.5, H_IN - 0.45, 1.5, 0.3, f"{n} / {total}", size=11, font=t["font_body"],
          color=color, align="right")
    if footer:
        _text(slide, left, H_IN - 0.45, max(1.5, min(8, right - left - 1.7)), 0.3, footer, size=11, font=t["font_body"],
              color=color)


_PAGE_NO = None


def renumber_footers(prs) -> int:
    """Diát beszúró/törlő műveletek után (pl. video to_pptx) a pakli „n / N” oldalszámai a valós sorrendet kövessék
    (10-13: az elejére tett videós dia után „2 / 10” maradt egy 11 diás paklin). → hány számot írt át."""
    import re
    global _PAGE_NO
    _PAGE_NO = _PAGE_NO or re.compile(r"^\s*(\d+)\s*/\s*(\d+)\s*$")
    total, changed = len(prs.slides), 0
    for i, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if not shp.has_text_frame or not _PAGE_NO.match(shp.text_frame.text or ""):
                continue
            want = f"{i} / {total}"
            runs = [r for para in shp.text_frame.paragraphs for r in para.runs]
            if runs and shp.text_frame.text.strip() != want:
                runs[0].text = want
                for r in runs[1:]:
                    r.text = ""
                changed += 1
    return changed


# ---------- belépési pont ----------

def _persist_spec(spec: dict[str, Any], by_slot: dict) -> dict[str, Any]:
    """A spec másolata, amelyben minden kép {file_id, credit}: a stockképet rejtett (ephemeral) fájlként mentjük,
    így az újraépítés ugyanazt a képet kapja (nem keres újra, a generáltat nem fizeti újra)."""
    import copy
    import shutil
    saved = copy.deepcopy(spec)
    for (i, j), x in by_slot.items():
        fid = x.source_file_id
        if not fid:
            fid = new_file_id("jpg")
            shutil.copy(x.path, output_path(fid))
            write_meta(fid, operation="deck image", label=f"slide {i + 1} image.jpg", extra={"ephemeral": True, "credit": x.attribution})
        ref = {"file_id": fid, **({"credit": x.attribution} if x.attribution else {})}
        orig = spec["slides"][i].get("image") if j < 0 else (spec["slides"][i].get("images") or [None] * (j + 1))[j]
        if isinstance(orig, dict):     # a képválasztó ebből tölti elő a keresést / az AI-promptot
            q, pr = orig.get("query") or orig.get("stock"), orig.get("prompt") or orig.get("generate")
            ref.update({k: v for k, v in (("query", q), ("prompt", pr)) if isinstance(v, str) and v})
        if j < 0:
            saved["slides"][i]["image"] = ref
        else:
            saved["slides"][i]["images"][j] = ref
    return saved


async def build(spec: dict[str, Any]) -> dict[str, Any]:
    if spec.get("template_file_id"):          # a minta-pptx TELJES sablonként (mesterdiák, elrendezés, háttér, logó)
        from app.tools import pptx_template
        from app.tools.crafts import existing
        out = await pptx_template.build(spec, existing(spec["template_file_id"]))
        title = spec.get("title") or ((spec.get("slides") or [{}])[0].get("title")) or "deck"
        write_meta(out["file_id"], operation="pptx_build", label=f"{title[:50]}.pptx", extra={"deck_spec": out.pop("saved_spec")})
        return out
    from pptx import Presentation
    from pptx.util import Inches
    t0 = time.time()
    slides = spec.get("slides") or []
    if not slides:
        raise ValueError("slides: at least one slide is needed")
    bad = [s.get("layout") for s in slides if s.get("layout") not in LAYOUTS]
    if bad:
        raise ValueError(f"Unknown layout: {bad} (choose from {', '.join(sorted(LAYOUTS))})")
    spec, kit = await _brand_kit(spec)
    t = _theme(spec, kit)

    reqs = await _pick_stock(_requests(slides), slides)
    resolved = await images.resolve_many([(r[1], r[2]) for r in reqs])
    by_slot = {r[0]: img for r, img in zip(reqs, resolved)}
    try:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
        blank = prs.slide_layouts[6]
        total = len(slides)
        for i, s in enumerate(slides):
            slide = prs.slides.add_slide(blank)
            lay = s["layout"]
            img = by_slot.get((i, -1))
            bg = slide.background.fill
            bg.solid(); bg.fore_color.rgb = _rgb(t["bg"])
            if lay == "cover":
                _cover(slide, s, t, img)
            elif lay in ("section", "closing"):
                _section(slide, s, t, img, closing=lay == "closing")
            elif lay == "bullets":
                _bullets_layout(slide, s, t, img)
            elif lay in ("image_left", "image_right"):
                _image_side(slide, s, t, img, left=lay == "image_left")
            elif lay == "full_image":
                _full_image(slide, s, t, img)
            elif lay == "quote":
                _quote(slide, s, t, img)
            elif lay == "stats":
                _stats(slide, s, t)
            elif lay == "chart":
                _chart(slide, s, t)
            elif lay == "table":
                _table(slide, s, t)
            elif lay == "gallery":
                _gallery(slide, s, t, [by_slot[(i, j)] for j in range(len((s.get("images") or [])[:4]))])
            if lay not in ("cover",):
                # bal oldali képnél a lábléc a kép mellé kerül (ne a fotóra)
                left = (W_IN / 2 if lay == "image_left" else 4.6 if lay == "quote" and img else 0) + M_IN
                right = (W_IN / 2 if lay == "image_right" and img else W_IN - 5.2 if lay == "bullets" and img else W_IN) - M_IN
                _footer(slide, i + 1, total, t, spec.get("footer"), dark=lay in ("section", "closing", "full_image"),
                        left=left, right=right)
            credits = [x.attribution for (si, _), x in by_slot.items() if si == i and x.attribution]
            notes = "\n".join(filter(None, [s.get("notes"), *credits]))
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
        bake_theme(prs, t)
        fid = new_file_id("pptx")
        prs.save(str(output_path(fid)))
        saved = _persist_spec(spec, by_slot)
    finally:
        for x in resolved:
            x.path.unlink(missing_ok=True)
    title = spec.get("title") or (slides[0].get("title") if slides else "") or "pakli"
    # a pakli leírása (a képek már fájlként) — ebből iterál a „kérj változtatást” (újraépítés új képkeresés nélkül)
    write_meta(fid, operation="pptx_build", label=f"{title[:50]}.pptx", extra={"deck_spec": saved})
    return {
        "file_id": fid, "url": public_url(fid), "slides": len(slides),
        "size_kb": output_path(fid).stat().st_size // 1024,
        "images": [{"slot": f"s{i + 1}" + (f".{j + 1}" if j >= 0 else ""), "file_id": x.source_file_id,
                    "credit": x.attribution} for ((i, j), x) in by_slot.items()],
        "ms_elapsed": int((time.time() - t0) * 1000),
        **({"brand": kit.get("name")} if kit and kit.get("name") else {}),
    }
