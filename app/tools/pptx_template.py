"""Pakli egy minta-pptx (vagy .potx) TELJES sablonjával: mesterdiák, elrendezések, háttér, logó, betűk, színek.

Mint a PowerPoint „Új bemutató sablonból”: a minta az alap, a diái törlődnek, a pakli-leírás (deck spec) diái a
minta SAJÁT elrendezéseibe kerülnek (címdia, tartalom, két hasáb, szakaszcím, kép, csak cím), a helyőrzőikbe.
Diagram / táblázat / kép a helyőrző helyére kerül; színt és betűt nem írunk felül → minden a sablon témájából
öröklődik (a diagramsorozatok a téma accent-színeit kapják). Kevés elrendezésű sablonnál (pl. csak
„cím + tartalom”) a legközelebbi elrendezés és a helyőrző területe szolgál.

Ha a mesterdián NINCS dizájn, mert az a diákon ül (Canva/Gamma-export, saját paklik), a minta DIÁIT klónozzuk:
pptx_clone. Ha azok sem klónozhatók (nincs dizájn a diákon, nincs szöveghely), usable() hamis → a hívó a téma-utat
választja (színek, betűk).
"""
from __future__ import annotations

import shutil
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Any, Optional

from app.storage import output_path, public_url
from app.storage.files import new_file_id
from app.tools import images


# ---------- sablon megnyitása ----------

def _as_pptx(src: Path) -> Path:
    """A .potx-et a python-pptx nem nyitja (más a fő tartalomtípus) → ideiglenes pptx-másolat."""
    if src.suffix.lower() != ".potx":
        return src
    dst = Path(tempfile.mkdtemp()) / "template.pptx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            zout.writestr(item, data)
    return dst


def _layouts(prs) -> list:
    """Minden mesterdia minden elrendezése (egy pakliban több mesterdia is lehet, pl. külön borító-mester)."""
    return [l for m in prs.slide_masters for l in m.slide_layouts]


def _fix_idx(prs) -> None:
    """Egyes (pl. LibreOffice-ból mentett) sablonokban a cím és a tartalom helyőrzője ugyanazt az idx-et kapja; a dia
    a helyőrzőt idx alapján örökli → a tartalom a cím helyére esne. Elrendezésenként egyedi idx (a cím marad 0)."""
    from pptx.enum.shapes import PP_PLACEHOLDER as T
    for layout in _layouts(prs):
        phs = sorted(layout.placeholders, key=lambda p: p.placeholder_format.type not in (T.TITLE, T.CENTER_TITLE))
        seen, nxt = set(), 100
        for sp in phs:
            ph = sp._element.ph
            if ph.idx in seen:
                ph.idx, nxt = nxt, nxt + 1
            seen.add(ph.idx)


def _plain(el) -> bool:
    """Háttér-elem „sima” (fehér/semleges): nincs kép, átmenet, minta, se nem-alap kitöltőszín."""
    from lxml import etree
    xml = etree.tostring(el).decode()
    if any(k in xml for k in ("blipFill", "gradFill", "pattFill")):
        return False
    import re
    for v in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml):
        if v.upper() not in ("FFFFFF", "FEFEFE", "FAFAFA"):
            return False
    return all(v in ("bg1", "lt1") for v in re.findall(r'schemeClr val="(\w+)"', xml))


def has_design(prs) -> bool:
    """A sablon hordoz-e látványt a mesterdián / elrendezéseken (háttér, kép, alakzat) — ha nem, a dizájn a diákon
    van (pl. Canva-export, a saját paklijaink), és a teljes-sablon út fehér paklit adna → téma-út kell."""
    from pptx.oxml.ns import qn
    for m in prs.slide_masters:                   # sötét sablon: a „bg1” háttér a sötét témaszínre képeződik
        cm = m._element.find(qn("p:clrMap"))
        if cm is not None and cm.get("bg1") in ("dk1", "dk2"):
            return True
    for part in [*prs.slide_masters, *_layouts(prs)]:
        bg = part._element.find(qn("p:cSld")).find(qn("p:bg"))
        if bg is not None and not _plain(bg):
            return True
        if any(not sh.is_placeholder for sh in part.shapes):
            return True
    return False


def usable(path: Path) -> bool:
    """Teljes sablonként használható-e: van cím+tartalom elrendezése ÉS a mesterdián van dizájn — vagy a dizájn a
    diákon ül, és a diák klónozhatók (pptx_clone). Ha egyik sem, a hívó a téma-utat választja (színek, betűk)."""
    from pptx import Presentation
    try:
        tpl = _as_pptx(path)
        prs = Presentation(str(tpl))
        if has_design(prs):
            ok = pick_layouts(prs)["content"] is not None
        else:
            from app.tools import pptx_clone
            ok = pptx_clone.clonable(prs)
        if tpl != path:
            shutil.rmtree(tpl.parent, ignore_errors=True)
        return ok
    except Exception:
        return False


def _master_rects(master) -> dict[str, tuple[int, int, int, int]]:
    from pptx.enum.shapes import PP_PLACEHOLDER as T
    out = {}
    for sh in master.placeholders:
        t = sh.placeholder_format.type
        key = "title" if t in (T.TITLE, T.CENTER_TITLE) else "body" if t in (T.BODY, T.OBJECT) else None
        if key and sh.width:
            out[key] = (sh.left, sh.top, sh.width, sh.height)
    return out


def _place(shape, rect) -> None:
    shape.left, shape.top, shape.width, shape.height = rect


def _clear_slides(prs) -> None:
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


# ---------- elrendezések osztályozása ----------

def _real(ph, SW: int, SH: int) -> bool:
    """Valódi tartalomhely-e egy szöveg-helyőrző — a lábléc-/felirat-csík (pl. 11 × 0,3 hüvelyk) nem az."""
    return (ph.width or 0) >= SW * 0.3 and (ph.height or 0) >= SH * 0.15


def _kinds(layout, SW: int, SH: int) -> dict[str, Any]:
    from pptx.enum.shapes import PP_PLACEHOLDER as T
    phs = list(layout.placeholders)
    types = [p.placeholder_format.type for p in phs]
    bodies = [p for p in phs if p.placeholder_format.type in (T.BODY, T.OBJECT) and _real(p, SW, SH)]
    return {"name": (layout.name or "").lower(), "title": any(t in (T.TITLE, T.CENTER_TITLE) for t in types),
            "center": T.CENTER_TITLE in types, "subtitle": T.SUBTITLE in types, "bodies": len(bodies),
            "picture": any(t == T.PICTURE for t in types)}


# elrendezés-nevek a gyakori nyelveken (a PowerPoint a felület nyelvén nevezi el őket)
_NAMES = {
    "title": ("title slide", "címdia", "titelfolie", "diapositive de titre", "diapositiva de título", "diapositiva titolo"),
    "content": ("title and content", "cím és tartalom", "titel und inhalt", "titre et contenu", "título y objetos", "titolo e contenuto"),
    "two": ("two content", "comparison", "két tartalom", "összehasonlítás", "zwei inhalte", "vergleich", "deux contenus", "comparaison", "dos objetos"),
    "section": ("section", "szakaszfejléc", "szakasz", "abschnitts", "en-tête de section", "encabezado de sección"),
    "picture": ("picture", "kép", "bild", "image", "imagen", "immagine"),
    "title_only": ("title only", "csak cím", "nur titel", "titre seul", "solo el título", "solo titolo"),
    "blank": ("blank", "üres", "leer", "vide", "en blanco", "vuota"),
}


def _named(k: dict[str, Any], role: str) -> bool:
    return any(n in k["name"] for n in _NAMES[role])


def pick_layouts(prs, usage: Optional[dict[str, int]] = None) -> dict[str, Any]:
    """Szerepenként a legjobb elrendezés: title, content, two, section, picture, title_only, blank. Ha több illik,
    azt, amelyiket a minta szerzője a diáin a legtöbbször használta (usage: elrendezés-partname → darab)."""
    L = [(l, _kinds(l, prs.slide_width, prs.slide_height)) for l in _layouts(prs)]
    usage = usage or {}

    def first(*preds):
        for pred in preds:
            cands = [l for l, k in L if pred(k)]
            if cands:
                return max(cands, key=lambda l: usage.get(str(l.part.partname), 0))
        return None
    content = first(lambda k: _named(k, "content") and k["title"] and k["bodies"] >= 1,
                    lambda k: k["title"] and k["bodies"] == 1 and not k["subtitle"], lambda k: k["title"] and k["bodies"] >= 1,
                    # a tartalom a diára rajzolt szövegdobozokban él (nincs tartalom-helyőrző): a cím alá írunk
                    lambda k: _named(k, "content") and k["title"], lambda k: k["title"] and not k["center"], lambda k: k["title"])
    title = first(lambda k: _named(k, "title") and k["title"], lambda k: k["center"], lambda k: k["title"] and k["subtitle"]) or content
    return {
        "title": title, "content": content,
        "two": first(lambda k: _named(k, "two") and k["bodies"] >= 2, lambda k: k["title"] and k["bodies"] >= 2) or content,
        "section": first(lambda k: _named(k, "section") and k["title"]) or title,
        "picture": first(lambda k: _named(k, "picture") and k["picture"], lambda k: k["picture"] and k["title"]) or content,
        "title_only": first(lambda k: _named(k, "title_only") and k["title"],
                            lambda k: k["title"] and not k["bodies"] and not k["subtitle"]) or content,
        "blank": first(lambda k: _named(k, "blank"), lambda k: not k["title"]),
    }


# ---------- kitöltés ----------

def _phs(slide, SW: int, SH: int):
    from pptx.enum.shapes import PP_PLACEHOLDER as T
    title = sub = None
    bodies, pics = [], []
    for p in list(slide.placeholders):
        t = p.placeholder_format.type
        if t in (T.TITLE, T.CENTER_TITLE):
            title = p
        elif t == T.SUBTITLE:
            sub = p
        elif t in (T.BODY, T.OBJECT):
            if _real(p, SW, SH):
                bodies.append(p)
            else:
                _drop(p)
        elif t == T.PICTURE:
            pics.append(p)
    bodies.sort(key=lambda p: (p.left, p.top))
    return title, sub, bodies, pics


def _drop(shape) -> tuple[int, int, int, int]:
    """Helyőrző eltávolítása, a területét visszaadja (oda kerül a kép/diagram/táblázat)."""
    rect = (shape.left, shape.top, shape.width, shape.height)
    el = shape._element
    el.getparent().remove(el)
    return rect


def _fill_text(ph, lines: list[str], level0: bool = True) -> None:
    from pptx.enum.text import MSO_AUTO_SIZE
    tf = ph.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(line)
        p.level = 0
        if len(lines) == 1:                   # egyetlen sor (alcím, szerző): ne kapjon felsorolásjelet
            from pptx.oxml.ns import qn
            ppr = p._p.get_or_add_pPr()
            ppr.set("indent", "0"); ppr.set("marL", "0")
            if ppr.find(qn("a:buNone")) is None:
                ppr.insert(0, ppr.makeelement(qn("a:buNone"), {}))
    try:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE      # hosszú szövegnél a PowerPoint/LibreOffice kicsinyít
    except Exception:
        pass


def _area_below(title_ph, SW: int, SH: int) -> tuple[int, int, int, int]:
    """A cím alatti szabad terület (ha az elrendezésnek nincs tartalom-helyőrzője)."""
    if title_ph is not None and title_ph.width:
        x, w, top = title_ph.left, title_ph.width, title_ph.top + title_ph.height
    else:
        x, w, top = int(SW * 0.06), int(SW * 0.88), int(SH * 0.18)
    y = top + int(SH * 0.03)
    return x, y, w, max(int(SH * 0.2), int(SH * 0.9) - y)


def _textbox(slide, rect, lines: list[str], size: Optional[int] = None) -> None:
    """Saját szövegdoboz a sablon betűivel (nem írunk felül betűt/színt), több sornál felsorolás-jellel."""
    from pptx.util import Pt
    tb = slide.shapes.add_textbox(*rect)
    tf = tb.text_frame
    tf.word_wrap = True
    n = len(lines)
    pt = size or (24 if n <= 3 else 20 if n <= 5 else 16)
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = par.add_run()
        r.text = (f"•  {line}" if n > 1 else str(line))
        r.font.size = Pt(pt)
        par.space_after = Pt(pt * 0.5)


def _set_title(slide, title_ph, text: Optional[str]) -> None:
    if title_ph is not None:
        if text:
            title_ph.text = str(text)
        else:
            _drop(title_ph)


def _chart(slide, s: dict[str, Any], rect) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA}.get(
        s.get("chart_type", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
    data = CategoryChartData()
    data.categories = s.get("categories") or []
    for se in s.get("series") or []:
        data.add_series(se.get("name", "Series"), se.get("values", []))
    ch = slide.shapes.add_chart(kind, *rect, data).chart
    ch.has_title = False
    ch.has_legend = len(s.get("series") or []) > 1 or kind in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT)
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False


def _table(slide, s: dict[str, Any], rect) -> None:
    from pptx.util import Pt
    headers, rows = s.get("headers") or [], s.get("rows") or []
    if not headers and not rows:
        return
    ncol = max(len(headers), max((len(r) for r in rows), default=0))
    tbl = slide.shapes.add_table(len(rows) + (1 if headers else 0), ncol, *rect).table     # a sablon táblastílusa
    r0 = 0
    if headers:
        for j, h in enumerate(headers):
            tbl.cell(0, j).text = str(h)
        r0 = 1
    for i, row in enumerate(rows):
        for j, v in enumerate(row[:ncol]):
            tbl.cell(r0 + i, j).text = str(v)
    size = Pt(14 if len(rows) <= 6 else 11)
    for row in tbl.rows:
        for c in row.cells:
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = size


def _picture(slide, img: Optional[images.ResolvedImage], rect) -> None:
    if img is not None:
        slide.shapes.add_picture(str(img.path), *rect)


def _stats(slide, s: dict[str, Any], rect) -> None:
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    stats = (s.get("stats") or [])[:4]
    if not stats:
        return
    x, y, w, h = rect
    cw = w // len(stats)
    for i, st in enumerate(stats):
        tb = slide.shapes.add_textbox(x + i * cw, y + h // 5, cw, h // 2)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(st.get("value", ""))
        r.font.size, r.font.bold = Pt(44), True
        r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = str(st.get("label", ""))
        r2.font.size = Pt(16)


# ---------- építés ----------

def _plan(prs, slides: list[dict[str, Any]], L: dict[str, Any]) -> list[tuple[Any, list[tuple[Any, float, str]]]]:
    """Diánként: (elrendezés, [(kép-spec, képarány, hová)]) — a képarány a választott helyőrzőből jön."""
    plan = []
    for s in slides:
        lay = s.get("layout")
        layout = {"cover": L["title"], "section": L["section"], "closing": L["section"], "quote": L["content"],
                  "bullets": L["two"] if s.get("image") else L["content"], "image_left": L["two"], "image_right": L["two"],
                  "full_image": L["picture"] or L["title_only"], "stats": L["title_only"], "chart": L["title_only"],
                  "table": L["title_only"], "gallery": L["title_only"], "diagram": L["title_only"]}.get(lay, L["content"])
        need = []
        if lay == "gallery":
            for img in (s.get("images") or [])[:4]:
                need.append((img, 1.5, "gallery"))
        elif s.get("image") and lay not in ("cover", "section", "closing", "quote", "stats", "chart", "table"):
            need.append((s["image"], _aspect_of(layout, lay, prs.slide_width, prs.slide_height), "main"))
        plan.append((layout, need))
    return plan


def _aspect_of(layout, lay: str, SW: int, SH: int) -> float:
    from pptx.enum.shapes import PP_PLACEHOLDER as T
    cands = [p for p in layout.placeholders if p.placeholder_format.type == T.PICTURE
             or (p.placeholder_format.type in (T.BODY, T.OBJECT) and _real(p, SW, SH))]
    if lay != "full_image" and len(cands) >= 2:
        p = sorted(cands, key=lambda p: p.left)[0 if lay == "image_left" else -1]
    elif cands:
        p = max(cands, key=lambda p: p.width * p.height)
    else:
        return 1.5
    return max(0.5, min(3.0, p.width / max(1, p.height)))


async def build(spec: dict[str, Any], template: Path) -> dict[str, Any]:
    from pptx import Presentation
    t0 = time.time()
    tpl = _as_pptx(template)
    prs = Presentation(str(tpl))
    if not has_design(prs):                       # a dizájn a diákon ül: a minta diáit klónozzuk
        from app.tools import pptx_clone
        if pptx_clone.clonable(prs):
            try:
                return await pptx_clone.build(spec, prs, t0)
            finally:
                if tpl != template:
                    shutil.rmtree(tpl.parent, ignore_errors=True)
    usage: dict[str, int] = {}
    for sl in prs.slides:                         # melyik elrendezést használta a minta szerzője
        key = str(sl.slide_layout.part.partname)
        usage[key] = usage.get(key, 0) + 1
    _clear_slides(prs)
    _fix_idx(prs)
    L = pick_layouts(prs, usage)
    if L["content"] is None:
        raise ValueError("This pptx has no layout with a title — it can't be used as a template.")
    master = _master_rects(L["content"].slide_master)
    # Címdia-szerű (lent/középen ülő címmel) tartalom-elrendezés, miközben a mesterdián a cím fent van: a tartalmas
    # diák a mesterdia geometriáját kapják
    lt = next((p for p in L["content"].placeholders if p.placeholder_format.idx == 0), None)
    content_like_master = (lt is not None and "title" in master and "body" in master
                           and lt.top > prs.slide_height * 0.2 and master["title"][1] < prs.slide_height * 0.2)
    slides = spec.get("slides") or []
    plan = _plan(prs, slides, L)
    reqs = [(img, asp) for _, need in plan for img, asp, _ in need]
    resolved = await images.resolve_many(reqs) if reqs else []
    it = iter(resolved)
    SW, SH = prs.slide_width, prs.slide_height
    try:
        for s, (layout, need) in zip(slides, plan):
            slide = prs.slides.add_slide(layout)
            title_ph, sub_ph, bodies, pics = _phs(slide, SW, SH)
            lay = s.get("layout")
            if content_like_master and lay not in ("cover", "section", "closing"):
                if title_ph is not None:
                    _place(title_ph, master["title"])
                if bodies:
                    _place(bodies[0], master["body"])
            imgs = [next(it) for _ in need]
            if lay in ("cover", "section", "closing"):
                _set_title(slide, title_ph, s.get("title"))
                text = " · ".join(x for x in (s.get("kicker") if lay == "cover" else None, s.get("subtitle"),
                                              f"{s['number']}" if s.get("number") else None) if x)
                target = sub_ph or (bodies[0] if bodies else None)
                if target is not None:
                    _fill_text(target, [text]) if text else _drop(target)
                elif text:
                    x, y, w, _ = _area_below(title_ph, SW, SH)
                    _textbox(slide, (x, y, w, int(SH * 0.15)), [text], size=20)
                for b in bodies[1 if target in bodies else 0:]:
                    _drop(b)
            elif lay == "quote":
                _set_title(slide, title_ph, f"“{s.get('quote', '')}”")
                if bodies:
                    _fill_text(bodies[0], [f"— {s['author']}"]) if s.get("author") else _drop(bodies[0])
                    for b in bodies[1:]:
                        _drop(b)
                elif s.get("author"):
                    x, y, w, _ = _area_below(title_ph, SW, SH)
                    _textbox(slide, (x, y, w, int(SH * 0.12)), [f"— {s['author']}"], size=20)
            elif lay in ("bullets", "image_left", "image_right"):
                _set_title(slide, title_ph, s.get("title"))
                lines = s.get("bullets") or ([s["text"]] if s.get("text") else [])
                text_ph = img_rect = None
                if imgs and len(bodies) >= 2:
                    text_ph, img_ph = (bodies[1], bodies[0]) if lay == "image_left" else (bodies[0], bodies[-1])
                    img_rect = _drop(img_ph)
                elif imgs and pics:
                    text_ph, img_rect = (bodies[0] if bodies else None), _drop(pics[0])
                elif bodies:
                    text_ph = bodies[0]
                    if imgs:          # egyetlen tartalomhely: a szöveg bal fele, a kép jobb fele
                        x, y, w, h = text_ph.left, text_ph.top, text_ph.width, text_ph.height
                        _place(text_ph, (x, y, int(w * 0.55), h))
                        img_rect = (x + int(w * 0.58), y, int(w * 0.42), h)
                if text_ph is not None:
                    _fill_text(text_ph, lines)
                elif not bodies and not (imgs and pics):     # nincs tartalom-helyőrző: saját szövegdoboz a cím alatt
                    x, y, w, h = _area_below(title_ph, SW, SH)
                    if imgs:
                        tx = x + int(w * 0.45) if lay == "image_left" else x
                        _textbox(slide, (tx, y, int(w * 0.55), h), lines)
                        img_rect = (x if lay == "image_left" else x + int(w * 0.6), y, int(w * 0.4), h)
                    else:
                        _textbox(slide, (x, y, w, h), lines)
                for b in [b for b in bodies if b is not text_ph and b._element.getparent() is not None]:
                    _drop(b)
                if imgs and img_rect:
                    _picture(slide, imgs[0], img_rect)
            elif lay == "full_image":
                _set_title(slide, title_ph, s.get("title") or s.get("caption"))
                area = pics[0] if pics else (bodies[0] if bodies else None)
                rect = _drop(area) if area is not None else _area_below(title_ph, SW, SH)
                for b in bodies:
                    if b._element.getparent() is not None:
                        _drop(b)
                if imgs and imgs[0] is not None:
                    _picture(slide, imgs[0], rect)
                elif s.get("caption") and s.get("title"):      # kép nélkül (pl. „no pictures”): a felirat szövegként
                    _textbox(slide, _area_below(title_ph, SW, SH), [s["caption"]])
            else:                      # stats · chart · table · gallery: a cím alatti terület
                _set_title(slide, title_ph, s.get("title"))
                area = bodies[0] if bodies else None
                if area is not None:
                    rect = _drop(area)
                else:
                    rect = _area_below(title_ph, SW, SH)
                for b in bodies[1:]:
                    _drop(b)
                if lay == "stats":
                    _stats(slide, s, rect)
                elif lay == "chart":
                    _chart(slide, s, rect)
                elif lay == "table":
                    _table(slide, s, rect)
                elif lay == "diagram":            # magyarázó ábra natív alakzatokkal, a sablon hátterén és betűivel
                    from app.tools import diagrams, pptx_deck
                    acc = pptx_deck.theme_accent(prs)
                    t_ = {"bg": "#" + (pptx_deck.bg_color(slide) or "FFFFFF"), "text": "#" + (pptx_deck.theme_text_color(prs) or "1E2430"),
                          "accent": "#" + (acc or "2F6FDE"), "muted": "#6B7280", "font_body": None}
                    reg = pptx_deck.diagram_actors(spec, t_["accent"])
                    diagrams.to_pptx(slide, s["diagram"], t_, tuple(v / 914400 for v in rect),
                                     {k: reg[k] for k in diagrams.actor_names(s["diagram"])})
                elif lay == "gallery" and imgs:
                    x, y, w, h = rect
                    n = len(imgs)
                    cols = 2 if n > 1 else 1
                    rows = -(-n // cols)
                    gw, gh = w // cols, h // rows
                    for k, img in enumerate(imgs):
                        _picture(slide, img, (x + (k % cols) * gw + gw // 40, y + (k // cols) * gh + gh // 40, gw - gw // 20, gh - gh // 20))
            for p in pics:
                if p._element.getparent() is not None:
                    _drop(p)
            credits = [x.attribution for x in imgs if x and x.attribution]
            notes = "\n".join(filter(None, [s.get("notes"), *credits]))
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
        fid = new_file_id("pptx")
        prs.save(str(output_path(fid)))
        by_slot, k = {}, 0
        for i, (_, need) in enumerate(plan):
            for j, (_, _, where) in enumerate(need):
                by_slot[(i, -1 if where == "main" else j)] = resolved[k]
                k += 1
        from app.tools.pptx_deck import _persist_spec
        saved = _persist_spec(spec, by_slot)          # a képek fájlként (törlés ELŐTT), az újraépítéshez
    finally:
        for x in resolved:
            x.path.unlink(missing_ok=True)
        if tpl != template:
            shutil.rmtree(tpl.parent, ignore_errors=True)
    return {"file_id": fid, "url": public_url(fid), "slides": len(slides), "saved_spec": saved,
            "size_kb": output_path(fid).stat().st_size // 1024, "ms_elapsed": int((time.time() - t0) * 1000)}
