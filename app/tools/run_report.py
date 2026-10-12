"""Részletes műveleti eredmény a Craft-motorok parancsfuttatásához (photocraft / vectorcraft / office / wordcraft).

A ChatGPT-teszt (10-12): „a »commands_run: 2« kevés ahhoz, hogy megítéljem, a felhasználói cél teljesült-e".
Ezért parancsonként: mi futott, mit adott vissza a motor (új réteg, határok, jelölések száma…), és a fájl
ELŐTTE–UTÁNA (méret, rétegek, diák, oldalak). Hibánál: melyik lépés bukott, mi futott le előtte, és hasonló
parancsazonosítók (elírásnál ez a leggyorsabb javítás).
"""
from __future__ import annotations

import difflib
import json
import re
from pathlib import Path
from typing import Any, Optional


def json_stream(text: str) -> list[Any]:
    """Egymás után írt (akár többsoros, behúzott) JSON-értékek egy szövegben → lista."""
    dec, i, out = json.JSONDecoder(), 0, []
    while True:
        starts = [x for x in (text.find("{", i), text.find("[", i)) if x >= 0]
        if not starts:
            return out
        try:
            val, i = dec.raw_decode(text, min(starts))
            out.append(val)
        except ValueError:
            i = min(starts) + 1


def _short(v: Any, n: int = 160) -> Any:
    """Paraméter/eredmény röviden a jelentésbe (nagy base64, hosszú szöveg nélkül)."""
    s = json.dumps(v, ensure_ascii=False) if not isinstance(v, str) else v
    if len(s) <= n:
        return v
    return s[:n] + "…"


def steps(commands: list[tuple[str, Optional[dict]]], results: list[Any]) -> list[dict[str, Any]]:
    """Parancsonként egy sor: n, command, params (röviden), result (amit a motor visszaadott)."""
    out = []
    for i, (cid, params) in enumerate(commands):
        row: dict[str, Any] = {"n": i + 1, "command": cid, "ok": True}
        if params:
            row["params"] = {k: _short(v, 80) for k, v in params.items() if k != "path"} | (
                {"file": Path(str(params["path"])).name} if params.get("path") else {})
        if i < len(results) and results[i] is not None:
            row["result"] = _short(results[i], 300)
        out.append(row)
    return out


def engine_results(log_text: str) -> list[Any]:
    """A photocraft/vectorcraft `run` kimenetéből a parancsonkénti eredmények (export/mentés sorai nélkül)."""
    out = []
    for v in json_stream(log_text):
        if isinstance(v, dict) and "command" in v and v.get("step", "cmd") == "cmd":
            out.append(v.get("result"))
    return out


_UNKNOWN = re.compile(r"unknown command `([^`]+)`")
_STEP = re.compile(r"step (\d+) \(([\w.]+)\)")


def failure(engine: str, err: str, stdout: str, commands: list[tuple[str, Optional[dict]]],
            known_ids: Optional[list[str]] = None, offset: int = 0) -> str:
    """Érthető hibaüzenet: melyik lépés bukott, mi futott le előtte, hasonló azonosítók; „semmi nem mentődött"."""
    done = engine_results(stdout)
    ran = len(done)
    bad_id = None
    m = _UNKNOWN.search(err)
    if m:
        bad_id = m.group(1)
    s = _STEP.search(err)
    idx = None
    if s:
        idx = max(0, int(s.group(1)) - 1 - offset)
    elif bad_id:
        idx = next((i for i, (c, _) in enumerate(commands) if c == bad_id), None)
    elif ran < len(commands):
        idx = ran
    parts = []
    if idx is not None and idx < len(commands):
        msg = err.strip().splitlines()[-1][:300].rstrip()
        parts.append(f"Step {idx + 1} of {len(commands)} ({commands[idx][0]}) failed: {msg}{'' if msg.endswith(('.', '!', '?')) else '.'}")
    else:
        parts.append(f"{engine} failed: {err.strip()[-300:]}")
    if idx:
        parts.append(f"Steps 1–{idx} ran ({', '.join(c for c, _ in commands[:idx][:8])}), but nothing was saved — fix the step and send all commands again.")
    else:
        parts.append("Nothing was saved.")
    if bad_id and known_ids:
        near = similar_ids(bad_id, known_ids)
        if near:
            parts.append(f"Similar commands: {', '.join(near)}. Look one up with read='commands', query='<id>'.")
    return " ".join(parts)


def similar_ids(bad: str, ids: list[str], n: int = 8) -> list[str]:
    """Elírt / kitalált azonosítóhoz a valódiak: előbb amelyek a szó MINDEN részét tartalmazzák
    („comment.new" → review.newComment), aztán a betű-hasonlók, végül amelyek a legtöbb részt."""
    parts = [x for x in re.split(r"[._\-\s]+", bad.lower()) if len(x) > 1]
    low = {i: i.lower() for i in ids}
    every = [i for i in ids if parts and all(x in low[i] for x in parts)]
    close = difflib.get_close_matches(bad, ids, n=n, cutoff=0.55)
    some = sorted((i for i in ids if parts and any(x in low[i] for x in parts if len(x) > 3)),
                  key=lambda i: -sum(x in low[i] for x in parts))
    out: list[str] = []
    for i in every + close + some:
        if i not in out:
            out.append(i)
    return out[:n]


# ---------- előtte–utána összegzők ----------

def photo_summary(info: dict[str, Any]) -> dict[str, Any]:
    layers = info.get("layers") or []
    return {"size": [info.get("width"), info.get("height")], "layers": len(layers),
            "layer_names": [f"{l.get('name')} ({l.get('kind', '').lower()}{'' if l.get('visible', True) else ', hidden'})"
                            for l in layers][:20],
            "selection": bool(info.get("hasSelection"))}


def vector_summary(info: dict[str, Any]) -> dict[str, Any]:
    fonts = info.get("fonts") or []
    meta = info.get("info") or {}
    return {"artboards": len(info.get("artboards") or []) or meta.get("artboards"),
            "fonts": [f.get("family") for f in fonts][:12],
            "missing_fonts": [f.get("family") for f in fonts if f.get("missing")]}


def deck_summary(p: Path) -> dict[str, Any]:
    from pptx import Presentation
    prs = Presentation(str(p))
    titles = []
    for s in prs.slides:
        t = s.shapes.title.text.strip() if s.shapes.title is not None and s.shapes.title.has_text_frame else ""
        titles.append(t[:60] or f"({s.slide_layout.name})")
    return {"slides": len(prs.slides), "titles": titles[:40]}


def pdf_summary(p: Path) -> dict[str, Any]:
    import fitz
    with fitz.open(str(p)) as d:
        fields = sum(len(list(pg.widgets() or [])) for pg in d)
        return {"pages": d.page_count, "form_fields": fields, "chars": sum(len(pg.get_text()) for pg in d)}


def pdf_step_report(script_out: str, steps_in: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """PdfCraft lépéssor: a kimenet első értéke a doc_open, utána lépésenként egy (a doc_save is) — a fontos számok."""
    vals = json_stream(script_out)[1:]
    out = []
    for i, st in enumerate(steps_in):
        row: dict[str, Any] = {"n": i + 1, "tool": st.get("tool"), "args": {k: _short(v, 80) for k, v in (st.get("args") or {}).items()}}
        if i < len(vals) and isinstance(vals[i], dict):
            v = vals[i]
            keep = {k: v[k] for k in ("marked", "pages_marked", "applied", "removed", "matched_words", "filled", "pages",
                                      "changed", "count", "added", "deleted", "warnings") if k in v}
            row["result"] = keep or _short(v, 200)
        out.append(row)
    return out
