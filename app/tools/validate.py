"""Dokumentum-ellenőrzés (validate_document): oldalképek + GÉPPEL ellenőrzött hibajegyzék, együtt.

10-13 ChatGPT-teszt (Integritás Hatóság-tanulmány): a docx_render_pages → describe_image lánc egy belső látómodell
SZÖVEGES leírását adta vissza, és az hibamentesnek mondta a hibás oldalakat. Ezért itt:
  - a képek MAGUK mennek vissza (áttekintő lap, hibánként kivágás, megjelölt oldalak) — a kliens-AI maga lát;
  - a hibát nem modell „látja”: a tördelt PDF geometriája (szavak, képek helye) + a docx szerkezete mondja meg;
  - KÉT tördelés: „editor” = WordCraft-motor (ugyanaz rajzol a böngészős szerkesztőben) és „word” = LibreOffice
    (≈ ahogy a letöltött fájlt a Word mutatja); a kettő összevetve (oldalszám, tördelés, számozás, melyikben hibás).

error:   üres listapont, árva cím (a lap utolsó sora), kettészakadó táblázatsor, lapról lelógó tartalom, ismétlődő szám
warning: szokatlan számozás (ugrás, újrakezdés, folytatás új fejezetben, kézi szám), margóba lógás, fejléc nélkül
         folytatódó táblázat, képaláírás külön lapon, üres cím, üres oldal, a két tördelés eltérése
Pakli: dobozából kifutó szöveg (a TÉNYLEGES tördelésből), dián kívüli elem, üres helyőrző. PDF: üres/lelógó oldal.

A blokk-sorszám (`block`) a WordCraft `inspect` indexe; a kurzor oda: edit.goto {"paragraph": block + 1}.
"""
from __future__ import annotations

import asyncio
import difflib
import io
import json
import logging
import os
import re
import subprocess
import unicodedata
from bisect import bisect_left
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

log = logging.getLogger("aice.validate")

WORDISH = {"docx", "docm", "dotx", "odt", "rtf", "md", "txt", "html", "htm", "doc"}
VIEW_NAMES = {"editor": "the built-in editor and Studio's PDF (WordCraft engine)",
              "word": "the downloaded file in Word (LibreOffice layout, close to Word)",
              "powerpoint": "the downloaded file in PowerPoint (LibreOffice layout)",
              "pdf": "the PDF itself"}
ERROR, WARNING = "error", "warning"


def view_name(L: "Layout") -> str:
    if L.view == "editor" and L.unit == "slide":
        return "the built-in editor and Studio's PDF (DeckCraft engine)"
    return VIEW_NAMES.get(L.view, L.view)
_RED, _ORANGE = (220, 38, 38), (234, 140, 0)

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"


# ---------- szöveg → szavak (ékezet, kis-nagybetű, írásjel nem számít) ----------

_TOK = re.compile(r"[^\W_]+")


def _fold(s: str) -> str:
    s = unicodedata.normalize("NFKD", (s or "").replace("­", ""))
    return "".join(c for c in s if not unicodedata.combining(c)).lower()


def _toks(s: str) -> list[str]:
    return _TOK.findall(_fold(s))


def _short(s: str, n: int = 60) -> str:
    s = re.sub(r"\s+", " ", s or "").strip()
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


# ---------- a tördelés: PDF → szavak, képek, oldalak ----------

@dataclass
class Tok:
    t: str
    page: int                                    # 1-től
    box: tuple


@dataclass
class Layout:
    view: str
    pdf: Path
    unit: str = "page"                                 # page | slide
    sizes: list = field(default_factory=list)          # oldalanként (szélesség, magasság), pt
    toks: list = field(default_factory=list)
    words: list = field(default_factory=list)          # oldalanként [(x0, y0, x1, y1, szó)]
    images: list = field(default_factory=list)         # oldalanként [doboz]
    inked: list = field(default_factory=list)          # van-e bármi az oldalon
    spans: dict = field(default_factory=dict)          # egység → (első, utolsó tok-index)
    pics: dict = field(default_factory=dict)           # blokk-pozíció → [(oldal, doboz)]
    page_sec: list = field(default_factory=list)       # oldalanként a szakasz sorszáma
    _doc: Any = None

    @property
    def pages(self) -> int:
        return len(self.sizes)

    def fitz(self):
        if self._doc is None:
            import fitz
            self._doc = fitz.open(str(self.pdf))
        return self._doc

    def close(self) -> None:
        if self._doc is not None:
            self._doc.close()
            self._doc = None


def _page_words(pg) -> list[tuple]:
    """Az oldal szavai dobozzal. A DeckCraft-PDF-en a PyMuPDF „words” bontása szavakat VESZÍT (10-13: a 3. dián
    „Az éves beszámoló kötelező …” sorból csak „Az” maradt, a „dict” span-szövege viszont teljes) → ha a szavak
    kevesebb betűt adnak, mint az oldal szövege, a span-okból bontunk (a sor y-ja pontos, az x arányos)."""
    ws = [(w[0], w[1], w[2], w[3], w[4]) for w in pg.get_text("words")]
    full = len(re.sub(r"\s+", "", pg.get_text("text")))
    if not full or sum(len(w[4]) for w in ws) >= 0.95 * full:
        return ws
    out = []
    for b in pg.get_text("dict").get("blocks", []):
        for line in b.get("lines", []):
            for sp in line.get("spans", []):
                txt, (x0, y0, x1, y1) = sp.get("text") or "", sp["bbox"]
                n = len(txt) or 1
                for m in re.finditer(r"\S+", txt):
                    out.append((x0 + (x1 - x0) * m.start() / n, y0, x0 + (x1 - x0) * m.end() / n, y1, m.group(0)))
    return out if sum(len(w[4]) for w in out) > sum(len(w[4]) for w in ws) else ws


def read_pdf(pdf: Path, view: str, unit: str = "page") -> Layout:
    import fitz
    L = Layout(view, pdf, unit)
    with fitz.open(str(pdf)) as d:
        for n, pg in enumerate(d, 1):
            L.sizes.append((pg.rect.width, pg.rect.height))
            ws = _page_words(pg)
            L.words.append(ws)
            for w in ws:
                for t in _toks(w[4]):
                    L.toks.append(Tok(t, n, (w[0], w[1], w[2], w[3])))
            ims = [tuple(im["bbox"]) for im in pg.get_image_info()]
            L.images.append(ims)
            L.inked.append(bool(ws) or bool(ims) or bool(pg.get_drawings()))
    return L


def _boxes(L: Layout, span: Optional[tuple], allowed: Optional[set] = None) -> dict[int, list[float]]:
    """Egy egység (bekezdés/cella) doboza oldalanként. `allowed`: csak az egység saját szavai — a laptörésnél közbeeső
    élőfej / megismételt táblázatfejléc így nem nyújtja meg a dobozt."""
    out: dict[int, list[float]] = {}
    if not span:
        return out
    for t in L.toks[span[0]: span[1] + 1]:
        if allowed is not None and t.t not in allowed:
            continue
        b = out.get(t.page)
        out[t.page] = ([min(b[0], t.box[0]), min(b[1], t.box[1]), max(b[2], t.box[2]), max(b[3], t.box[3])]
                       if b else list(t.box))
    return out


_G = 3


def align(units: list[tuple[Any, list[str]]], toks: list[Tok]) -> dict[Any, tuple[int, int]]:
    """A dokumentum egységeit (sorrendben) a PDF szófolyamára illeszti: egység → (első, utolsó tok-index).
    Szó-hármasokkal horgonyoz, előre haladó kurzorral (az ismétlődő bekezdések így sem csúsznak egymásra)."""
    seq = [t.t for t in toks]
    grams: dict[tuple, list[int]] = defaultdict(list)
    for i in range(len(seq) - _G + 1):
        grams[(seq[i], seq[i + 1], seq[i + 2])].append(i)
    single: dict[str, list[int]] = defaultdict(list)
    for i, t in enumerate(seq):
        single[t].append(i)

    def hits(g: tuple, lo: int, hi: int):
        cand = grams.get(g) if len(g) == _G else single.get(g[0])
        if not cand:
            return
        k = bisect_left(cand, max(0, lo))
        while k < len(cand) and cand[k] <= hi:
            i = cand[k]
            if len(g) == _G or tuple(seq[i:i + len(g)]) == g:
                yield i
            k += 1

    out, cur = {}, 0
    for key, T in units:
        n = len(T)
        if not n:
            continue
        g = min(_G, n)
        hi = cur + (1500 if n >= _G else 120) + n
        start = None
        for o in range(0, min(n - g, 8) + 1):
            p = next(hits(tuple(T[o:o + g]), cur, hi), None)
            if p is not None:
                start = max(0, p - o)
                break
        if start is None:
            continue
        end = min(len(seq) - 1, start + n - 1)
        for o in range(0, min(n - g, 8) + 1):
            exp = start + n - g - o                 # ahol a záró hármasnak lennie kellene
            near = [p for p in hits(tuple(T[n - g - o:n - o]), max(start, exp - 30 - n // 10), exp + 40 + n // 5)]
            if near:
                end = min(near, key=lambda p: abs(p - exp)) + g - 1 + o
                break
        out[key] = (start, max(start, min(end, len(seq) - 1)))
        cur = out[key][1] + 1
    return out


# ---------- a dokumentum szerkezete (docx) ----------

@dataclass
class Block:
    pos: int                                     # saját sorszám a törzsben
    index: int                                   # a WordCraft inspect indexe (edit.goto paragraph = index + 1)
    kind: str                                    # paragraph | table
    text: str = ""
    style: str = ""
    heading: Optional[int] = None                # 0 = Title, n = Heading n
    caption: bool = False
    num: Optional[tuple] = None                  # (numId, ilvl, numFmt, kezdőérték)
    pictures: list = field(default_factory=list)  # beágyazott képek szélessége (pt)
    floating: int = 0
    keep_next: bool = False                      # „együtt a következővel” (közvetlenül vagy a stílusból)
    page_break: bool = False                     # kézi oldaltörés / szakaszhatár a bekezdésben
    break_before: bool = False
    rows: list = field(default_factory=list)     # táblázat: [[cellaszöveg]]
    cant_split: list = field(default_factory=list)
    header_row: bool = False
    section: int = 0

    @property
    def empty(self) -> bool:
        return self.kind == "paragraph" and not self.text.strip() and not self.pictures and not self.floating

    @property
    def numbered(self) -> bool:
        return bool(self.num) and self.num[2] not in ("bullet", "none")

    def obj(self) -> str:
        if self.kind == "table":
            return f"table (block {self.index})"
        what = ("heading" if self.heading is not None else "caption" if self.caption
                else "list item" if self.num else "picture" if self.pictures and not self.text.strip() else "paragraph")
        return f"{what} (block {self.index})" + (f" “{_short(self.text, 40)}”" if self.text.strip() else "")


def _val(el, tag: str) -> Optional[str]:
    c = el.find(_W + tag) if el is not None else None
    return None if c is None else c.get(_W + "val", "")


def _on(el, tag: str) -> bool:
    c = el.find(_W + tag) if el is not None else None
    return c is not None and c.get(_W + "val", "true") not in ("0", "false", "off")


def _body(el):
    for ch in el.iterchildren():
        if ch.tag in (_W + "p", _W + "tbl"):
            yield ch
        elif ch.tag in (_W + "sdt", _W + "customXml"):
            c = ch.find(_W + "sdtContent") if ch.tag == _W + "sdt" else ch
            if c is not None:
                yield from _body(c)


def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(_W + "t"))


def _section(sp) -> dict[str, float]:
    sz, mar = sp.find(_W + "pgSz"), sp.find(_W + "pgMar")

    def tw(el, k, d):
        try:
            return float(el.get(_W + k)) / 20 if el is not None and el.get(_W + k) else d
        except ValueError:
            return d
    return {"pageW": tw(sz, "w", 612), "pageH": tw(sz, "h", 792), "mL": tw(mar, "left", 72), "mR": tw(mar, "right", 72),
            "mT": tw(mar, "top", 72), "mB": tw(mar, "bottom", 72)}


def read_docx(path: Path) -> tuple[list[Block], list[dict]]:
    from docx import Document
    d = Document(str(path))
    styles = {st.get(_W + "styleId"): st for st in d.styles.element.findall(_W + "style")}
    default_p = next((sid for sid, st in styles.items() if st.get(_W + "type") == "paragraph"
                      and st.get(_W + "default") in ("1", "true")), None)
    numfmt: dict[str, dict[int, tuple]] = {}             # numId → {szint: (formátum, kezdőérték)}
    try:
        nel = d.part.numbering_part.element
        abstract = {a.get(_W + "abstractNumId"): {int(l.get(_W + "ilvl") or 0): (_val(l, "numFmt") or "decimal",
                                                                               int(_val(l, "start") or 1))
                                                  for l in a.findall(_W + "lvl")} for a in nel.findall(_W + "abstractNum")}
        for n in nel.findall(_W + "num"):
            lv = dict(abstract.get(_val(n, "abstractNumId"), {}))
            for o in n.findall(_W + "lvlOverride"):
                il, so = int(o.get(_W + "ilvl") or 0), _val(o, "startOverride")
                if so:
                    lv[il] = (lv.get(il, ("decimal", 1))[0], int(so))
            numfmt[n.get(_W + "numId")] = lv
    except Exception:
        pass

    def chain(p):
        ppr = p.find(_W + "pPr")
        if ppr is not None:
            yield ppr
        sid, seen = _val(ppr, "pStyle") or default_p, set()
        while sid and sid not in seen:
            seen.add(sid)
            st = styles.get(sid)
            if st is None:
                return
            sp = st.find(_W + "pPr")
            if sp is not None:
                yield sp
            sid = _val(st, "basedOn")

    blocks: list[Block] = []
    sections: list[dict] = []
    for pos, el in enumerate(_body(d.element.body)):
        if el.tag == _W + "tbl":
            rows, cant = [], []
            for tr in el.findall(_W + "tr"):
                rows.append([" ".join(_text(p) for p in tc.iter(_W + "p")) for tc in tr.findall(_W + "tc")])
                cant.append(_on(tr.find(_W + "trPr"), "cantSplit"))
            first = el.find(_W + "tr")
            blocks.append(Block(pos, pos, "table", rows=rows, cant_split=cant, section=len(sections),
                                header_row=first is not None and _on(first.find(_W + "trPr"), "tblHeader")))
            continue
        ppr = el.find(_W + "pPr")
        sid = _val(ppr, "pStyle") or default_p
        name = ((styles.get(sid) is not None and _val(styles[sid], "name")) or sid or "").lower()
        num_id, ilvl, outline, brk_before, keep = None, None, None, False, None
        for pp in chain(el):
            npr = pp.find(_W + "numPr")
            if npr is not None:
                if num_id is None:
                    num_id = _val(npr, "numId")
                if ilvl is None and npr.find(_W + "ilvl") is not None:
                    ilvl = int(_val(npr, "ilvl") or 0)
            if outline is None and _val(pp, "outlineLvl") not in (None, ""):
                outline = int(_val(pp, "outlineLvl"))
            brk_before = brk_before or _on(pp, "pageBreakBefore")
            if keep is None and pp.find(_W + "keepNext") is not None:
                keep = _on(pp, "keepNext")
        m = re.match(r"heading\s*(\d)", name)
        heading = 0 if name == "title" else int(m.group(1)) if m else (outline + 1 if outline is not None and outline < 9 else None)
        num = None
        if num_id not in (None, "", "0"):
            num = (num_id, ilvl or 0) + numfmt.get(num_id, {}).get(ilvl or 0, ("decimal", 1))
        sect = ppr.find(_W + "sectPr") if ppr is not None else None
        blocks.append(Block(
            pos, pos, "paragraph", text=_text(el), style=name, heading=heading, caption=name == "caption", num=num,
            pictures=[float((inl.find(_WP + "extent").get("cx") or 0)) / 12700 for inl in el.iter(_WP + "inline")
                      if inl.find(_WP + "extent") is not None],
            floating=sum(1 for _ in el.iter(_WP + "anchor")),
            page_break=any(br.get(_W + "type") == "page" for br in el.iter(_W + "br"))
            or (sect is not None and _val(sect, "type") != "continuous"),
            break_before=brk_before, keep_next=bool(keep), section=len(sections)))
        if sect is not None:
            sections.append({**_section(sect), "end": pos})
    final = d.element.body.find(_W + "sectPr")
    sections.append({**(_section(final) if final is not None else _section(d.element.body)), "end": len(blocks)})
    return blocks, sections


def _wordcraft_index(path: Path, blocks: list[Block]) -> None:
    """A WordCraft-inspect sorszámai a blokkokra (edit.goto ezt számolja); eltérő szerkezetnél szövegillesztéssel."""
    from app.config import settings
    from app.tools import wordcraft
    exe = wordcraft.cli_path()
    if not exe:
        return
    try:
        r = subprocess.run([exe, "inspect", str(path)], capture_output=True, timeout=settings.WORDCRAFT_TIMEOUT_S,
                           env={**os.environ, "WORDCRAFT_NO_PREFS": "1"})
        theirs = json.loads(r.stdout)["blocks"] if r.returncode == 0 else None
    except Exception as e:
        log.info("wordcraft inspect kimaradt: %s", e)
        return
    if not theirs:
        return

    def sig(kind, text):
        return kind + ":" + "".join(_toks(text))[:24]
    mine = [sig(b.kind, b.text) for b in blocks]
    other = [sig(t.get("type", ""), t.get("text") or "") for t in theirs]
    if mine == other:
        for b, t in zip(blocks, theirs):
            b.index = int(t.get("index", b.pos))
        return
    for a, bb, n in difflib.SequenceMatcher(None, mine, other, autojunk=False).get_matching_blocks():
        for k in range(n):
            blocks[a + k].index = int(theirs[bb + k].get("index", bb + k))


# ---------- hibák ----------

def _issue(kind: str, sev: str, message: str, view: Optional[str] = None, page: Optional[int] = None,
           block: Optional[Block] = None, row: Optional[int] = None, fix: Optional[str] = None,
           regions: Optional[list] = None, marks: Optional[list] = None) -> dict[str, Any]:
    return {"kind": kind, "severity": sev, "message": message, "view": view, "page": page,
            "block": block.index if block else None, "_pos": block.pos if block else None, "row": row,
            "object": (block.obj() + (f", row {row}" if row else "")) if block else None, "fix": fix,
            "_regions": regions or [], "_marks": marks or []}


def _goto(b: Block, then: str) -> str:
    return f'wordcraft: edit.goto {{"paragraph": {b.index + 1}}}, then {then}'


class DocCheck:
    """Egy Word-dokumentum: szerkezet + egy vagy két tördelés → hibajegyzék."""

    def __init__(self, blocks: list[Block], sections: list[dict], layouts: list[Layout]):
        self.blocks, self.sections, self.layouts = blocks, sections, layouts
        self.issues: list[dict] = []
        self._label_x: dict[tuple, float] = {}
        for L in layouts:
            self._place(L)

    # --- illesztés: egységek, képek, szakaszok ---
    def _place(self, L: Layout) -> None:
        units = []
        for b in self.blocks:
            if b.kind == "table":
                units += [(("c", b.pos, r, c), _toks(cell)) for r, row in enumerate(b.rows) for c, cell in enumerate(row)]
            else:
                units.append((("p", b.pos), _toks(b.text)))
        L.spans = align(units, L.toks)
        sec_of = lambda b: self.sections[min(b.section, len(self.sections) - 1)]
        # képek: a beágyazott kép az előző elhelyezett blokk vége UTÁNI első (fej-/láblécen kívüli) kép
        flat = sorted((pg, box[1], i, box) for pg, ims in enumerate(L.images, 1) for i, box in enumerate(ims))
        used, after = set(), (1, 0.0)
        for b in self.blocks:
            if b.pictures:
                s = sec_of(b)
                got = []
                for pg, y0, i, box in flat:
                    if len(got) >= len(b.pictures):
                        break
                    if (pg, i) in used or (pg, y0 + 5) < after:
                        continue
                    h = L.sizes[pg - 1][1]
                    if box[3] <= s["mT"] * 0.9 or box[1] >= h - s["mB"] * 0.9:      # fej-/lábléc logója
                        continue
                    used.add((pg, i))
                    got.append((pg, box))
                if got:
                    L.pics[b.pos] = got
                    after = max(after, (got[-1][0], got[-1][1][3]))
            end = self.end(L, b)
            if end:
                after = max(after, end)
        # oldal → szakasz (az oldalon kezdődő első blokk szakasza)
        first = {}
        for b in self.blocks:
            p = self.first_page(L, b)
            if p and p not in first:
                first[p] = b.section
        cur = 0
        L.page_sec = []
        for p in range(1, L.pages + 1):
            cur = first.get(p, cur)
            L.page_sec.append(min(cur, len(self.sections) - 1))

    def span(self, L: Layout, b: Block) -> Optional[tuple]:
        if b.kind == "table":
            sp = [L.spans[k] for k in (("c", b.pos, r, c) for r, row in enumerate(b.rows) for c in range(len(row)))
                  if k in L.spans]
            return (min(s[0] for s in sp), max(s[1] for s in sp)) if sp else None
        return L.spans.get(("p", b.pos))

    def boxes(self, L: Layout, b: Block) -> dict[int, list[float]]:
        words = set(_toks(" ".join(c for row in b.rows for c in row) if b.kind == "table" else b.text))
        out = _boxes(L, self.span(L, b), words)
        for pg, box in L.pics.get(b.pos, []):
            o = out.get(pg)
            out[pg] = [min(o[0], box[0]), min(o[1], box[1]), max(o[2], box[2]), max(o[3], box[3])] if o else list(box)
        return out

    def first_page(self, L: Layout, b: Block) -> Optional[int]:
        bx = self.boxes(L, b)
        return min(bx) if bx else None

    def end(self, L: Layout, b: Block) -> Optional[tuple]:
        bx = self.boxes(L, b)
        if not bx:
            return None
        p = max(bx)
        return (p, bx[p][3])

    def sec(self, L: Layout, page: int) -> dict[str, float]:
        return self.sections[L.page_sec[page - 1]] if L.page_sec else self.sections[0]

    def text_area(self, L: Layout, page: int) -> list[float]:
        s, (w, h) = self.sec(L, page), L.sizes[page - 1]
        return [s["mL"], s["mT"], w - s["mR"], h - s["mB"]]

    def strip(self, L: Layout, page: int, y0: float, y1: float, pad: float = 24) -> tuple:
        """Kivágási sáv: a szövegterület szélessége (kicsit túl), y0–y1 között."""
        a, (w, h) = self.text_area(L, page), L.sizes[page - 1]
        return (page, (max(0, a[0] - pad), max(0, y0), min(w, a[2] + pad), min(h, y1)))

    def gap(self, L: Layout, i: int) -> Optional[tuple[int, list[float]]]:
        """Egy szöveg nélküli blokk helye: az előző elhelyezett blokk vége és a következő eleje között."""
        prev = next((self.end(L, b) for b in reversed(self.blocks[:i]) if self.end(L, b)), None)
        nxt = None
        for b in self.blocks[i + 1:]:
            bx = self.boxes(L, b)
            if bx:
                p = min(bx)
                nxt = (p, bx[p][1])
                break
        if prev is None and nxt is None:
            return None
        page = prev[0] if prev else nxt[0]
        a = self.text_area(L, page)
        y0 = prev[1] if prev else a[1]
        y1 = nxt[1] if nxt and nxt[0] == page else a[3]
        return page, [a[0] - 20, y0, a[2], max(y1, y0 + 14)]

    def next_content(self, i: int) -> Optional[int]:
        return next((j for j in range(i + 1, len(self.blocks)) if not self.blocks[j].empty), None)

    # --- szerkezeti hibák (tördeléstől független) ---
    def structural(self) -> None:
        n0 = len(self.issues)
        self._structural()
        self._typed_numbers()
        for it in self.issues[n0:]:
            it["_structural"] = True

    def _structural(self) -> None:
        L = self.layouts[-1] if self.layouts else None          # a hely a „word” nézetből (ha van), különben a szerkesztőé
        for i, b in enumerate(self.blocks):
            if b.kind != "paragraph":
                continue
            if b.num and b.empty:
                g = self.gap(L, i) if L else None
                label = self._label_in(L, g) if g else None
                shows = (f" — only “{label}” shows" if label and _TOK.search(label) else " — only the bullet shows" if not b.numbered else "")
                self.issues.append(_issue(
                    "empty_list_item", ERROR,
                    f"Empty {'numbered' if b.numbered else 'bulleted'} list item" + shows
                    + (f" (after “{_short(self.blocks[i - 1].text, 40)}”)" if i and self.blocks[i - 1].text.strip() else "") + ".",
                    L.view if g else None, g[0] if g else None, b,
                    fix=_goto(b, "text.delete (removes the empty item) — or give it its text"),
                    regions=[self.strip(L, g[0], g[1][1] - 40, g[1][3] + 40)] if g else [],
                    marks=[(g[0], g[1])] if g else []))
            elif b.heading is not None and b.empty:
                g = self.gap(L, i) if L else None
                self.issues.append(_issue(
                    "empty_heading", WARNING, "Empty heading — it shows as a blank line and a blank entry in a table of contents.",
                    L.view if g else None, g[0] if g else None, b, fix=_goto(b, "text.delete"),
                    regions=[self.strip(L, g[0], g[1][1] - 40, g[1][3] + 40)] if g else [], marks=[(g[0], g[1])] if g else []))
            if b.numbered and not b.empty and _absurd_start(b.num[2], b.num[3]):
                lab = self.labels(L).get(b.pos) if L else None
                p, box = self.item_box(L, b) if L else (None, None)
                self.issues.append(_issue(
                    "numbering", ERROR,
                    f"This item's number starts at {b.num[3]} in a {_FMT.get(b.num[2], b.num[2])} list"
                    + (f" — it shows as “{_short(lab, 14)}”" if lab else "") + ". A number from the start of the text "
                    f"(a year?) has probably turned into the list number: put “{b.num[3]}.” back into the text and switch "
                    "the numbering off for this item.",
                    L.view if p else None, p, b, fix=_goto(b, f'para.numbering {{"off": true}}, then text.insert {{"text": "{b.num[3]}. "}}'),
                    regions=[self.strip(L, p, box[1] - 60, box[3] + 40)] if p else [], marks=[(p, box)] if p else []))
            for w in b.pictures:
                tw = self._text_width(b)
                if w > tw + 1:
                    self.issues.append(_issue(
                        "picture_too_wide", ERROR,
                        f"A picture is {w:.0f} pt wide, the text area only {tw:.0f} pt — in Word it sticks out past the "
                        f"margin{' and off the page' if w > tw + self.sections[b.section]['mR'] else ''}.",
                        block=b, fix=_goto(b, f'select the picture and picture.size {{"width": {int(tw)}, "lockAspect": true}}')))

    def _text_width(self, b: Block) -> float:
        s = self.sections[min(b.section, len(self.sections) - 1)]
        return s["pageW"] - s["mL"] - s["mR"]

    def _typed_numbers(self) -> None:
        """Kézzel gépelt számozás: szomszédos „1.” „2.” „4.” bekezdések, és a fejezetcímek számai."""
        L = self.layouts[-1] if self.layouts else None
        run: list[tuple[int, Block]] = []
        heads: dict[int, tuple] = {}
        prev_auto = None

        def where(b: Block):
            bx = self.boxes(L, b) if L else {}
            if not bx:
                return None, [], []
            p = min(bx)
            return p, [self.strip(L, p, bx[p][1] - 50, bx[p][3] + 50)], [(p, bx[p])]

        for b in self.blocks:
            if b.kind != "paragraph" or b.empty:
                continue
            m = re.match(r"\s*(\d{1,3})[.)]\s+\S", b.text)
            if b.heading is not None:
                hm = re.match(r"\s*((?:\d+\.)*\d+)\.?\s+\S", b.text)
                run = []
                if hm and not b.num:
                    t = tuple(int(x) for x in hm.group(1).split("."))
                    prev = heads.get(b.heading)
                    bad = None
                    if prev and len(prev) == len(t):
                        if t[:-1] == prev[:-1] and t[-1] != prev[-1] + 1:
                            bad = f"comes after “{'.'.join(map(str, prev))}”"
                        elif t[:-1] != prev[:-1] and t[-1] != 1:
                            bad = "should start a new numbering at 1"
                    if bad:
                        p, reg, mk = where(b)
                        self.issues.append(_issue(
                            "heading_numbering", WARNING if t[-1] != (prev[-1] if prev else -1) else ERROR,
                            f"Heading number “{'.'.join(map(str, t))}” {bad} — a number is skipped or repeated.",
                            L.view if p else None, p, b, fix=_goto(b, "fix the typed number (edit.replace …)"), regions=reg, marks=mk))
                    heads[b.heading] = t
                    for lv in [k for k in heads if k > b.heading]:
                        heads.pop(lv)
                continue
            if m and not b.num:
                v = int(m.group(1))
                if run and v != run[-1][0] + 1 and v != 1:
                    p, reg, mk = where(b)
                    self.issues.append(_issue(
                        "typed_numbering", ERROR if v == run[-1][0] else WARNING,
                        f"Typed numbering goes “{run[-1][0]}.” → “{v}.”" + (" (repeated)" if v == run[-1][0] else " (skips)") + ".",
                        L.view if p else None, p, b, fix=_goto(b, "correct the number, or make the items a real numbered list (para.numbering)"),
                        regions=reg, marks=mk))
                if prev_auto is not None and not run:
                    p, reg, mk = where(b)
                    self.issues.append(_issue(
                        "numbering_mixed", WARNING, "Typed numbers right after an automatic numbered list — they won't follow on.",
                        L.view if p else None, p, b, fix=_goto(b, "turn it into the same list (para.numbering) or remove the typed number"),
                        regions=reg, marks=mk))
                run.append((v, b))
                prev_auto = None
                continue
            run = []
            prev_auto = b if b.numbered else None
            if b.num and not b.numbered:
                prev_auto = None

    # --- tördelési hibák (nézetenként) ---
    def layout(self, L: Layout) -> None:
        self._headings(L)
        self._tables(L)
        self._overflow(L)
        self._numbering(L)
        for p in range(1, L.pages + 1):
            if not L.inked[p - 1]:
                self.issues.append(_issue("blank_page", WARNING, f"Page {p} is empty.", L.view, p,
                                          fix="remove the empty paragraphs or the page break before it",
                                          regions=[(p, (0, 0) + tuple(L.sizes[p - 1]))]))

    def _headings(self, L: Layout) -> None:
        for i, b in enumerate(self.blocks):
            if b.kind != "paragraph" or b.empty or not (b.heading is not None or b.caption or self._is_caption(b)):
                continue
            j = self.next_content(i)
            hb = self.boxes(L, b)
            if j is None or not hb:
                continue
            nb = self.blocks[j]
            nbx = self.boxes(L, nb)
            if not nbx:
                continue
            hp, np_ = max(hb), min(nbx)
            between = self.blocks[i:j + 1]
            if np_ <= hp or b.page_break or nb.break_before or any(x.page_break for x in between[:-1]):
                continue
            regions = [self.strip(L, hp, hb[hp][1] - 70, self.text_area(L, hp)[3] + 12),
                       self.strip(L, np_, self.text_area(L, np_)[1] - 12, nbx[np_][3] + 50)]
            marks = [(hp, hb[hp]), (np_, nbx[np_])]
            if b.heading is not None and L.view == "editor" and b.keep_next:
                self.issues.append(_issue(          # a fájl rendben: Word együtt tartja, csak a WordCraft-motor nem
                    "orphan_heading", WARNING,
                    f"Heading “{_short(b.text, 50)}” is the last line on page {hp}, its text starts on page {np_}. The "
                    "heading is already set to keep with its text, so Word keeps them together — only the built-in "
                    "editor's layout leaves it behind.",
                    L.view, hp, b, fix="nothing to fix in the file; if a PDF made here is needed, a page break before "
                                       "the heading (para.set {\"props\": {\"pageBreakBefore\": true}}) keeps them together",
                    regions=regions, marks=marks))
            elif b.heading is not None:
                self.issues.append(_issue(
                    "orphan_heading", ERROR,
                    f"Heading “{_short(b.text, 50)}” is the last line on page {hp}; its text starts on page {np_}.",
                    L.view, hp, b, fix=_goto(b, 'para.set {"props": {"keepNext": true}} (keep with next)'),
                    regions=regions, marks=marks))
            elif nb.pictures or nb.kind == "table":
                self.issues.append(_issue(
                    "caption_split", WARNING,
                    f"Caption “{_short(b.text, 50)}” is on page {hp}, its {'table' if nb.kind == 'table' else 'picture'} on page {np_}.",
                    L.view, hp, b, fix=_goto(b, 'para.set {"props": {"keepNext": true}}'), regions=regions, marks=marks))
        # képaláírás a kép UTÁN: a kép és az aláírás külön lapon
        for i, b in enumerate(self.blocks):
            if not b.pictures or b.text.strip():
                continue
            j = self.next_content(i)
            if j is None:
                continue
            c = self.blocks[j]
            if not (c.caption or self._is_caption(c)):
                continue
            pb, cb = self.boxes(L, b), self.boxes(L, c)
            if pb and cb and min(cb) > max(pb) and not b.page_break:
                pp, cp = max(pb), min(cb)
                self.issues.append(_issue(
                    "caption_split", WARNING, f"The picture is on page {pp}, its caption “{_short(c.text, 50)}” on page {cp}.",
                    L.view, pp, c, fix=_goto(b, 'para.set {"props": {"keepNext": true}} on the picture'),
                    regions=[self.strip(L, pp, pb[pp][1] - 30, self.text_area(L, pp)[3] + 12),
                             self.strip(L, cp, self.text_area(L, cp)[1] - 12, cb[cp][3] + 40)],
                    marks=[(pp, pb[pp]), (cp, cb[cp])]))

    @staticmethod
    def _is_caption(b: Block) -> bool:
        return bool(re.match(r"\s*((figure|fig\.|table|chart|exhibit)\s*\d+|\d+\.\s*(ábra|táblázat|diagram))", b.text, re.I))

    def _tables(self, L: Layout) -> None:
        for b in self.blocks:
            if b.kind != "table":
                continue
            all_pages: set[int] = set()
            for r, row in enumerate(b.rows):
                bx: dict[int, list[float]] = {}
                for c in range(len(row)):
                    for pg, box in _boxes(L, L.spans.get(("c", b.pos, r, c)), set(_toks(row[c]))).items():
                        o = bx.get(pg)
                        bx[pg] = [min(o[0], box[0]), min(o[1], box[1]), max(o[2], box[2]), max(o[3], box[3])] if o else box
                all_pages |= set(bx)
                if len(bx) > 1:
                    p0, p1 = min(bx), max(bx)
                    tall = r < len(b.cant_split) and b.cant_split[r]
                    self.issues.append(_issue(
                        "table_row_split", WARNING if (tall and L.view == "editor") else ERROR,
                        f"Row {r + 1} of the table (“{_short(' '.join(row), 40)}”) splits across pages {p0}–{p1}"
                        + (" although it is set not to break — the row is taller than the space left; shorten it." if tall else "."),
                        L.view, p0, b, row=r + 1,
                        fix=_goto(b, f'caret.down to row {r + 1}, then table.properties {{"allowBreak": false}} '
                                     "(Word: Table Properties → Row → untick “Allow row to break across pages”)"),
                        regions=[self.strip(L, p0, bx[p0][1] - 40, self.text_area(L, p0)[3] + 12),
                                 self.strip(L, p1, self.text_area(L, p1)[1] - 12, bx[p1][3] + 40)],
                        marks=[(p0, bx[p0]), (p1, bx[p1])]))
            if len(all_pages) > 1 and len(b.rows) >= 3 and not b.header_row:
                p1 = sorted(all_pages)[1]
                top = min((t.box[1] for k, sp in L.spans.items() if k[0] == "c" and k[1] == b.pos
                           for t in L.toks[sp[0]:sp[1] + 1] if t.page == p1), default=0)
                self.issues.append(_issue(
                    "table_header_not_repeated", WARNING,
                    f"The table continues on page {p1} without its header row (“{_short(' | '.join(b.rows[0]), 50)}”).",
                    L.view, p1, b, row=1, fix=_goto(b, 'table.properties {"headerRow": true} (caret in the first row)'),
                    regions=[self.strip(L, p1, top - 60, top + 90)], marks=[(p1, [self.text_area(L, p1)[0], top - 2, self.text_area(L, p1)[2], top + 14])]))

    def _overflow(self, L: Layout) -> None:
        owner = {}                                  # tok-index → blokk
        for b in self.blocks:
            sp = self.span(L, b)
            if b.kind == "table":
                for (k, sp2) in ((k, v) for k, v in L.spans.items() if k[0] == "c" and k[1] == b.pos):
                    for t in range(sp2[0], sp2[1] + 1):
                        owner[t] = b
            elif sp:
                for t in range(sp[0], sp[1] + 1):
                    owner[t] = b
        seen = set()
        for i, t in enumerate(L.toks):
            w, h = L.sizes[t.page - 1]
            b = owner.get(i)
            off = t.box[2] > w + 1 or t.box[0] < -1 or t.box[3] > h + 1
            a = self.text_area(L, t.page)
            tol = 12 if b is not None and b.kind == "table" else 6
            margin = b is not None and (t.box[2] > a[2] + tol or t.box[0] < a[0] - tol)
            if not (off or margin):
                continue
            key = ("off" if off else "margin", b.pos if b else None, t.page)
            if key in seen:
                continue
            seen.add(key)
            what = "table" if b is not None and b.kind == "table" else "text"
            box = list(t.box)
            if off:
                self.issues.append(_issue(
                    "off_page", ERROR, f"{what.capitalize()} runs off the edge of page {t.page} (“{_short(t.t, 30)}…”) — it is cut off.",
                    L.view, t.page, b, fix=(_goto(b, "make it narrower (table.properties {\"widthPct\": 100})" if what == "table"
                                                  else "check its indents and long unbroken text") if b else None),
                    regions=[(t.page, (0, max(0, box[1] - 60), w, min(h, box[3] + 60)))], marks=[(t.page, box)]))
            else:
                self.issues.append(_issue(
                    "into_margin", WARNING, f"{what.capitalize()} reaches into the {'right' if t.box[2] > a[2] + tol else 'left'} "
                                            f"margin on page {t.page}.",
                    L.view, t.page, b, fix=(_goto(b, 'table.properties {"widthPct": 100}' if what == "table"
                                                  else "check the paragraph's indents") if b else None),
                    regions=[(t.page, (0, max(0, box[1] - 50), w, min(h, box[3] + 50)))], marks=[(t.page, box)]))
        mapped = {(pg, tuple(box)): b for b in self.blocks for pg, box in L.pics.get(b.pos, [])}
        for pg, ims in enumerate(L.images, 1):
            w, h = L.sizes[pg - 1]
            a = self.text_area(L, pg)
            for box in ims:
                b = mapped.get((pg, tuple(box)))
                off = box[2] > w + 2 or box[0] < -2 or box[3] > h + 2
                if not off and not (b is not None and (box[2] > a[2] + 3 or box[0] < a[0] - 3)):
                    continue
                sev = ERROR if (off and b is not None) else WARNING
                over = max(box[2] - w, -box[0], 0)
                self.issues.append(_issue(
                    "off_page" if off else "into_margin", sev,
                    (f"A picture runs {over:.0f} pt off the edge of page {pg} — part of it is cut off." if off
                     else f"A picture reaches into the margin on page {pg}."),
                    L.view, pg, b, fix=(_goto(b, f'select the picture and picture.size {{"width": {int(a[2] - a[0])}, "lockAspect": true}}')
                                        if b else None),
                    regions=[(pg, (0, max(0, box[1] - 30), w, min(h, box[3] + 30)))], marks=[(pg, list(box))]))

    # --- számozás: a tördelt címkék („1.” „2.” …) sorrendje ---
    def labels(self, L: Layout) -> dict[int, str]:
        out = {}
        for b in self.blocks:
            if b.kind == "paragraph" and b.num and not b.empty:
                sp = self.span(L, b)
                if sp:
                    w = _label_left_of(L, L.toks[sp[0]])
                    if w:
                        out[b.pos] = w[4]
                        self._label_x[(L.view, b.pos)] = w[0]
        return out

    def item_box(self, L: Layout, b: Block) -> tuple[Optional[int], Optional[list[float]]]:
        """Egy listapont doboza a tördelt számmal („3.”) együtt."""
        bx = self.boxes(L, b)
        if not bx:
            return None, None
        p = min(bx)
        box = list(bx[p])
        x = self._label_x.get((L.view, b.pos))
        if x is not None:
            box[0] = min(box[0], x)
        return p, box

    def _label_in(self, L: Layout, g: tuple) -> Optional[str]:
        page, (x0, y0, x1, y1) = g
        ws = [w for w in L.words[page - 1] if w[1] >= y0 - 1 and w[3] <= y1 + 1 and w[0] >= x0 - 4]
        return " ".join(w[4] for w in ws[:2]) if ws and len(" ".join(w[4] for w in ws)) <= 6 else None

    def _numbering(self, L: Layout) -> None:
        labels = self.labels(L)
        last: dict[tuple, tuple] = {}             # (numId, ilvl) → (érték, blokk, fejezet-számláló)
        empties: dict[tuple, int] = defaultdict(int)
        chapter, prev = 0, None                   # prev: a KÖZVETLENÜL előző (számozott) listapont
        run = None                                # az épp nyitott hiba (egymás utáni „1. 1. 1.” = EGY tétel)
        for b in self.blocks:
            if b.kind == "paragraph" and b.heading is not None and not b.num:
                chapter += 1
            if b.kind == "paragraph" and b.empty and not b.num:
                continue
            if not (b.kind == "paragraph" and b.numbered):
                if not (b.kind == "paragraph" and b.num):          # felsorolásjel nem töri meg a listát
                    prev = run = None
                continue
            k = (b.num[0], b.num[1])
            for kk in [x for x in last if x[0] == k[0] and x[1] > k[1]]:
                last.pop(kk)
            if b.empty:
                empties[k] += 1
                continue
            v = _label_value(labels.get(b.pos, ""), b.num[2])
            if v is None and len(labels.get(b.pos, "")) > 8 and not _absurd_start(b.num[2], b.num[3]):
                p, box = self.item_box(L, b)
                self.issues.append(_issue(
                    "numbering", ERROR, f"This item is numbered “{_short(labels[b.pos], 14)}” — the numbering has gone wrong.",
                    L.view, p, b, fix=_goto(b, "para.numbering / para.restartNumbering, or switch numbering off for it"),
                    regions=[self.strip(L, p, box[1] - 60, box[3] + 40)] if p else [], marks=[(p, box)] if p else []))
            if v is None:
                continue
            lab = labels[b.pos]
            bad = None
            adjacent_other = prev is not None and prev[1] == k[1] and prev[0] != k[0]
            if k in last:
                lv, lb, ch = last[k]
                exp = lv + 1 + empties[k]
                if v == lv and not empties[k]:
                    bad = (ERROR, "repeat", f"Numbered items repeat “{lab}” (as the item before it).")
                elif v > exp:
                    bad = (WARNING, "skip", f"Numbering skips: “{lab}” comes after {lv}.")
                elif v < exp:
                    bad = (WARNING, "restart", f"Numbering starts over at “{lab}” in the middle of the list (the item before is {lv}).")
                elif ch < chapter and v > 1:
                    bad = (WARNING, "continues", f"The list under a new heading starts at “{lab}” — it continues the earlier list.")
            elif adjacent_other and v == prev[2]:
                bad = (ERROR, "separate", f"Numbered items repeat “{lab}” — each item is a separate list.")
            elif adjacent_other and v == 1 and prev[2] >= 1:
                bad = (WARNING, "restart", f"A new list starts right after the previous one: numbering starts over at “{lab}”.")
            elif v != b.num[3]:
                bad = (WARNING, "continues", f"This list starts at “{lab}”, not {b.num[3]} — it continues the numbering of an "
                                             "earlier list.")
            p, box = self.item_box(L, b)
            if bad and run is not None and run["_tag"] == bad[1] and prev is not None and run["_last"] == prev[3]:
                run["_n"] += 1                                         # ugyanaz a hiba a következő ponton is: bővül
                run["_last"] = b.pos
                run["message"] = re.sub(r"^Numbered items repeat “[^”]*”", f"{run['_n']} numbered items in a row show "
                                        f"“{lab}”", run["message"]) if bad[1] == "separate" else run["message"]
                if p and p == run["page"]:
                    run["_marks"].append((p, box))
                    pg, r = run["_regions"][0]
                    run["_regions"][0] = (pg, (r[0], r[1], r[2], max(r[3], box[3] + 40)))
            elif bad:
                marks = [(p, box)] if p else []
                if p and bad[1] in ("repeat", "separate") and prev is not None:       # az előző (azonos számú) pont is
                    pp, pbox = self.item_box(L, self.blocks[prev[3]])
                    if pp == p:
                        marks.insert(0, (pp, pbox))
                top = min(m[1][1] for m in marks) if marks else 0
                run = _issue(
                    "numbering", bad[0], bad[2], L.view, p, b,
                    fix=_goto(b, "para.restartNumbering (or put the items in one list: para.numbering on all of them)"
                              if bad[1] in ("repeat", "separate", "continues") else "check the list (para.numbering, para.restartNumbering)"),
                    regions=[self.strip(L, p, top - 60, box[3] + 40)] if p else [], marks=marks)
                run.update({"_tag": bad[1], "_last": b.pos, "_n": 2 if bad[1] == "separate" else 1})
                self.issues.append(run)
            else:
                run = None
            last[k] = (v, b, chapter)
            empties[k] = 0
            prev = (k[0], k[1], v, b.pos)

    # --- összegzés ---
    def run(self) -> list[dict]:
        self.structural()
        for L in self.layouts:
            self.layout(L)
        if len(self.layouts) == 2:
            self._compare(*self.layouts)
        return merge(self.issues, [L.view for L in self.layouts])

    def _compare(self, E: Layout, W: Layout) -> None:
        if E.pages != W.pages:
            self.issues.append(_issue(
                "views_differ", WARNING,
                f"The built-in editor lays the document out on {E.pages} pages, Word (LibreOffice) on {W.pages} — "
                "page breaks, and so headings at a page bottom or split rows, can differ between them.", None, None))
        le, lw = self.labels(E), self.labels(W)
        for b in self.blocks:
            a, c = le.get(b.pos), lw.get(b.pos)
            if a and c and b.numbered and a.strip(".)") != c.strip(".)"):
                bx = self.boxes(W, b)
                p = min(bx) if bx else None
                self.issues.append(_issue(
                    "numbering_differs", WARNING, f"This item is numbered “{a}” in the editor but “{c}” in Word.",
                    W.view, p, b, fix=_goto(b, "para.restartNumbering / para.numbering so both agree"),
                    regions=[self.strip(W, p, bx[p][1] - 50, bx[p][3] + 40)] if p else [], marks=[(p, bx[p])] if p else []))

    def page_starts(self, L: Layout) -> list[Optional[tuple]]:
        """Oldalanként: melyik egység (bekezdés/cella) hányadik szavával kezdődik — a két tördelés így vethető össze."""
        starts: dict[int, tuple] = {}
        for key, (s, e) in sorted(L.spans.items(), key=lambda kv: kv[1][0]):
            for i in range(s, e + 1):
                pg = L.toks[i].page
                if pg not in starts:
                    starts[pg] = (key, i - s)
        return [starts.get(p) for p in range(1, L.pages + 1)]


def _label_left_of(L: Layout, t: Tok) -> Optional[tuple]:
    """A listapont tördelt címkéje („3.”, „b)”, „•”): ugyanazon a soron, közvetlenül a szöveg előtt."""
    x0, y0, x1, y1 = t.box
    cy = (y0 + y1) / 2
    best = None
    for w in L.words[t.page - 1]:
        if w[2] <= x0 + 0.5 and abs((w[1] + w[3]) / 2 - cy) < 3 and x0 - w[2] < 90:
            if best is None or w[2] > best[2]:
                best = w
    return best


_ROMAN = {"i": 1, "v": 5, "x": 10, "l": 50, "c": 100}
_FMT = {"decimal": "numbered", "lowerLetter": "lettered (a, b, c)", "upperLetter": "lettered (A, B, C)",
        "lowerRoman": "roman (i, ii, iii)", "upperRoman": "roman (I, II, III)"}


def _absurd_start(fmt: str, start: int) -> bool:
    """Egy lista nem indul 2012-től: betűsnél 26, rómainál 39, számosnál 100 fölött gyanús."""
    f = (fmt or "").lower()
    return start > (26 if "letter" in f else 39 if "roman" in f else 99)


def _label_value(raw: str, fmt: str) -> Optional[int]:
    s = (raw or "").strip().strip("()[]").rstrip(".):")
    if not s:
        return None
    if re.fullmatch(r"\d+(?:\.\d+)*", s):
        return int(s.split(".")[-1])
    if "roman" in fmt.lower() and re.fullmatch(r"[ivxlc]+", s, re.I):
        vals = [_ROMAN[ch] for ch in s.lower()]
        return sum(-v if i + 1 < len(vals) and v < vals[i + 1] else v for i, v in enumerate(vals))
    if re.fullmatch(r"[a-z]", s, re.I):
        return ord(s.lower()) - 96
    return None


def merge(issues: list[dict], views: list[str]) -> list[dict]:
    """Ugyanaz a hiba két nézetben → egy tétel (`seen_in`, oldalanként); a szerkezeti képszélesség-hiba elmarad,
    ha a tördelés már megmutatta. Sorrend: hibák, aztán figyelmeztetések, oldalszám szerint; azonosító E1…, W1…"""
    shown = {i["_pos"] for i in issues if i["kind"] in ("off_page", "into_margin") and i["_pos"] is not None}
    broken = {i["_pos"] for i in issues if i["kind"] == "numbering" and i["severity"] == ERROR}
    out: dict[tuple, dict] = {}
    for it in issues:
        if (it["kind"] == "picture_too_wide" and it["_pos"] in shown) or (it["kind"] == "numbering_differs" and it["_pos"] in broken):
            continue
        key = (it["kind"], it["_pos"], it["row"], it["message"] if it["_pos"] is None else None)
        if it["kind"] in ("off_page", "into_margin") and it["_pos"] is not None:
            key = (it["kind"], it["_pos"], None, None)
        cur = out.get(key)
        if cur is None:
            out[key] = {**it, "seen_in": [it["view"]] if it["view"] else [], "pages": {it["view"]: it["page"]} if it["view"] else {}}
            continue
        if it["view"] and it["view"] not in cur["seen_in"]:
            cur["seen_in"].append(it["view"])
            cur["pages"][it["view"]] = it["page"]
            if it["view"] != "editor" and not cur.get("_structural"):     # a képek, oldalszám a letöltött fájl nézetéből
                for k in ("message", "page", "view", "_regions", "_marks", "fix"):
                    cur[k] = it[k]
            if it["severity"] == ERROR:
                cur["severity"] = ERROR
    order = {ERROR: 0, WARNING: 1}
    items = sorted(out.values(), key=lambda i: (order[i["severity"]], i["page"] or 10 ** 6, i["_pos"] or 0))
    ne = nw = 0
    for it in items:
        if it["severity"] == ERROR:
            ne += 1
            it["id"] = f"E{ne}"
        else:
            nw += 1
            it["id"] = f"W{nw}"
        if it.get("_structural"):                      # a fájl szerkezetéből: minden nézetre igaz
            it["seen_in"] = list(views)
        elif len(views) == 2 and len(it["seen_in"]) == 1:
            it["message"] += {"word": " (only in Word)", "powerpoint": " (only in PowerPoint)"}.get(
                it["seen_in"][0], " (only in the built-in editor and Studio's PDF)")
    return items


def public(it: dict) -> dict:
    return {k: v for k, v in it.items() if not k.startswith("_") and k != "view" and v not in (None, [], {})}


# ---------- pakli ----------

def _pptx_shapes(path: Path) -> tuple[list[list[dict]], float, float]:
    from pptx import Presentation
    prs = Presentation(str(path))
    sw, sh = prs.slide_width / 12700, prs.slide_height / 12700
    slides = []
    for s in prs.slides:
        items = []
        for shp in s.shapes:
            if shp.left is None or shp.width is None:
                continue
            try:
                pic = getattr(shp, "image", None) is not None
            except Exception:
                pic = True
            size, grows = None, False
            if shp.has_text_frame:
                from pptx.enum.text import MSO_AUTO_SIZE
                size = max((r.font.size.pt for p in shp.text_frame.paragraphs for r in p.runs if r.font.size), default=None)
                grows = shp.text_frame.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
            # id = a pptx saját alakzat-azonosítója — a DeckCraft is ezt használja (office text.set / edit.select)
            items.append({"id": shp.shape_id, "name": shp.name or "shape",
                          "box": [shp.left / 12700, shp.top / 12700, (shp.left + shp.width) / 12700, (shp.top + shp.height) / 12700],
                          "text": shp.text_frame.text if shp.has_text_frame else "", "placeholder": shp.is_placeholder,
                          "picture": pic, "table": getattr(shp, "has_table", False), "size": size, "grows": grows})
        slides.append(items)
    return slides, sw, sh


def _overlap(a: list[float], b: list[float], pad: float = 2) -> bool:
    return a[0] < b[2] - pad and b[0] < a[2] - pad and a[1] < b[3] - pad and b[1] < a[3] - pad


def _cmd(c: str, **params) -> str:
    return json.dumps({"command": c, "params": params}, ensure_ascii=False)


def _deck_obj(n: int, it: dict) -> dict:
    """A hibás alakzat azonosítása úgy, hogy a következő hívás közvetlenül javíthassa (dia + shape id)."""
    return {"object": f"slide {n}, shape {it['id']} “{it['name']}”", "slide": n, "shape_id": it["id"], "_key": (n, it["id"])}


def _deck_fix(kind: str, n: int, it: dict, sw: float, sh: float, over: float = 0) -> str:
    """Kész javító parancs(ok) az office eszközhöz — csak ez az alakzat változik, a pakli többi része nem."""
    sel = _cmd("edit.select", on_slide=n, ids=[it["id"]])
    if kind == "text_overflow":
        h = it["box"][3] - it["box"][1]
        size = it.get("size")
        # a tördelt szöveg magassága ~ betűméret² → a beférő méret ~ gyök(arány); 0,95 tartalék, egész pontra
        smaller = max(8, int((size or 18) * ((h / (h + over)) ** 0.5 if h > 0 else 0.8) * 0.95))
        return (f'office(app="deck", file_id, commands=[{_cmd("text.set", on_slide=n, id=it["id"], text="<shorter text>")}]) '
                f'— or a smaller font: commands=[{sel}, {_cmd("format.size", size=smaller)}]'
                + (f" (now {size:g} pt)" if size else "") + " — or a bigger box: shape.setBounds after the select.")
    if kind == "off_slide":
        b = it["box"]
        w, h = min(b[2] - b[0], sw * 0.96), min(b[3] - b[1], sh * 0.96)
        x, y = min(max(b[0], 0), sw - w), min(max(b[1], 0), sh - h)
        return f'office(app="deck", file_id, commands=[{sel}, {_cmd("shape.setBounds", x=round(x, 1), y=round(y, 1), w=round(w, 1), h=round(h, 1))}])'
    if kind == "empty_placeholder":
        return (f'office(app="deck", file_id, commands=[{_cmd("text.set", on_slide=n, id=it["id"], text="…")}]) '
                f'— or remove it: commands=[{sel}, {_cmd("edit.delete")}]')
    return ""


def deck_issues(path: Path, layouts: list[Layout]) -> list[dict]:
    slides, sw, sh = _pptx_shapes(path)
    issues = []
    for n, items in enumerate(slides, 1):
        for it in items:
            b = it["box"]
            if b[0] < -0.02 * sw or b[1] < -0.02 * sh or b[2] > 1.02 * sw or b[3] > 1.02 * sh:
                issues.append({**_issue("off_slide", WARNING, f"Slide {n}: “{it['name']}” reaches outside the slide.", None, n,
                                        fix=_deck_fix("off_slide", n, it, sw, sh)), **_deck_obj(n, it)})
            if it["placeholder"] and not it["text"].strip() and not it["picture"] and not it["table"]:
                issues.append({**_issue("empty_placeholder", WARNING,
                                        f"Slide {n}: empty placeholder “{it['name']}” (shows “Click to add text” in PowerPoint).", None, n,
                                        fix=_deck_fix("empty_placeholder", n, it, sw, sh)), **_deck_obj(n, it)})
    for L in layouts:
        if not L.pages:
            continue
        sc = L.sizes[0][0] / sw if sw else 1.0
        by_page: dict[int, list[int]] = defaultdict(list)
        for i, t in enumerate(L.toks):
            by_page[t.page].append(i)
        for n, items in enumerate(slides, 1):
            if n > L.pages:
                break
            idx = by_page.get(n, [])
            sub = [L.toks[i] for i in idx]
            for it in items:
                T = _toks(it["text"])
                if not T or it["table"]:
                    continue
                sp = align([("s", T)], sub).get("s")
                if not sp:
                    continue
                own, span = set(T), sub[sp[0]:sp[1] + 1]
                hit = sum(1 for t in span if t.t in own)
                if hit < 0.7 * len(T):          # kevés saját szó a tartományban: nincs bizonyíték, nincs ítélet
                    continue
                tb = None
                for t in span:
                    if t.t not in own:          # a közbeeső lábléc/cím szavai ne nyújtsák meg a dobozt
                        continue
                    tb = list(t.box) if tb is None else [min(tb[0], t.box[0]), min(tb[1], t.box[1]), max(tb[2], t.box[2]), max(tb[3], t.box[3])]
                if tb is None:
                    continue
                box = [v * sc for v in it["box"]]
                w, h = L.sizes[n - 1]
                reg = [(n, (max(0, min(box[0], tb[0]) - 30), max(0, min(box[1], tb[1]) - 30),
                            min(w, max(box[2], tb[2]) + 30), min(h, max(box[3], tb[3]) + 30)))]
                if tb[3] > h + 1 or tb[2] > w + 1:
                    over = max(tb[3] - h, tb[2] - w) / sc if sc else 0
                    issues.append({**_issue("off_page", ERROR, f"Slide {n}: text of “{it['name']}” runs off the slide.",
                                            L.view, n, regions=reg, marks=[(n, box), (n, tb)],
                                            fix=_deck_fix("text_overflow", n, it, sw, sh, over)), **_deck_obj(n, it)})
                elif tb[3] > box[3] + 6 or tb[2] > box[2] + 6:
                    by = max(tb[3] - box[3], tb[2] - box[2])
                    # más szövegre rálóg → hiba; csak a saját dobozán nyúlik túl (pl. a kártyán belül) → figyelmeztetés
                    hits = [o["name"] for o in items if o is not it and o["text"].strip() and _overlap(tb, [v * sc for v in o["box"]])]
                    big = by > max(24, 0.25 * (box[3] - box[1]))
                    if it["grows"] and not hits:
                        # „méret a szöveghez” doboz: PowerPointban a doboz nő a szöveggel — a két motor betűmérése
                        # tér el, nem hiba; csak ha másra rácsúszik (vagy lelóg a diáról: fent, off_page)
                        continue
                    issues.append({**_issue("text_overflow", ERROR if (hits or big) else WARNING,
                                            f"Slide {n}: the text of “{it['name']}” runs {by:.0f} pt past its box"
                                            + (f" and over “{hits[0]}”." if hits else "." if big else
                                               " (it stays on the slide and covers nothing)."),
                                            L.view, n, regions=reg, marks=[(n, box), (n, tb)],
                                            fix=_deck_fix("text_overflow", n, it, sw, sh, by / sc if sc else by)), **_deck_obj(n, it)})
        for n in range(1, L.pages + 1):
            if not L.inked[n - 1]:
                issues.append({**_issue("blank_page", WARNING, f"Slide {n} is empty.", L.view, n,
                                        regions=[(n, (0, 0) + tuple(L.sizes[n - 1]))]), "_key": (n, "")})
    if len(layouts) == 2 and layouts[0].pages != layouts[1].pages:
        issues.append({**_issue("views_differ", WARNING,
                                f"The editor shows {layouts[0].pages} slides, PowerPoint (LibreOffice) {layouts[1].pages}.", None, None),
                       "_key": None})
    for it in issues:                               # a merge kulcsa: dia + alakzat
        it["_pos"] = hash(it.pop("_key")) if it.get("_key") else None
    return merge(issues, [L.view for L in layouts])


# ---------- PDF ----------

def pdf_issues(L: Layout) -> list[dict]:
    issues = []
    for n in range(1, L.pages + 1):
        w, h = L.sizes[n - 1]
        if not L.inked[n - 1]:
            issues.append(_issue("blank_page", WARNING, f"Page {n} is empty.", L.view, n, regions=[(n, (0, 0, w, h))]))
        off = [wd for wd in L.words[n - 1] if wd[2] > w + 1 or wd[0] < -1 or wd[3] > h + 1]
        if off:
            box = [min(x[0] for x in off), min(x[1] for x in off), max(x[2] for x in off), max(x[3] for x in off)]
            issues.append(_issue("off_page", ERROR, f"Text runs off the edge of page {n} (“{_short(off[0][4], 30)}”).", L.view, n,
                                 regions=[(n, (0, max(0, box[1] - 60), w, min(h, box[3] + 60)))], marks=[(n, box)]))
        for box in L.images[n - 1]:
            if box[2] > w + 2 or box[0] < -2 or box[3] > h + 2:
                issues.append(_issue("off_page", WARNING, f"A picture runs off the edge of page {n}.", L.view, n,
                                     regions=[(n, (0, max(0, box[1] - 30), w, min(h, box[3] + 30)))], marks=[(n, list(box))]))
    return merge(issues, [L.view])


# ---------- képek ----------

def _font(size: int):
    from PIL import ImageFont
    for name in ("DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except Exception:
            pass
    return ImageFont.load_default()


def _jpeg(im, q: int = 78) -> bytes:
    buf = io.BytesIO()
    im.convert("RGB").save(buf, "JPEG", quality=q, optimize=True)
    return buf.getvalue()


def _render(L: Layout, page: int, rect: tuple, zoom: float):
    import fitz
    from PIL import Image
    pg = L.fitz()[page - 1]
    clip = fitz.Rect(*rect) & pg.rect
    pix = pg.get_pixmap(matrix=fitz.Matrix(zoom, zoom), clip=clip, alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples), clip


def _draw_marks(im, clip, zoom: float, marks: list[tuple]) -> None:
    """marks: [(doboz, szín, címke)] — a dobozt a kivágásba vágva rajzolja (a lapról lelógó rész a szélén)."""
    from PIL import ImageDraw
    dr = ImageDraw.Draw(im)
    f = _font(max(12, int(7 * zoom)))
    for box, color, tag in marks:
        x0 = (max(box[0], clip.x0) - clip.x0) * zoom - 3
        y0 = (max(box[1], clip.y0) - clip.y0) * zoom - 3
        x1 = (min(box[2], clip.x1) - clip.x0) * zoom + 3
        y1 = (min(box[3], clip.y1) - clip.y0) * zoom + 3
        if x1 <= x0 or y1 <= y0:
            continue
        dr.rectangle([x0, y0, x1, y1], outline=color, width=3)
        for cut, edge in ((box[2] > clip.x1 + 1, [x1 - 7, y0, x1, y1]), (box[0] < clip.x0 - 1, [x0, y0, x0 + 7, y1]),
                          (box[3] > clip.y1 + 1, [x0, y1 - 7, x1, y1])):
            if cut:                                          # ide már nem fér: a lap széle vágja le
                dr.rectangle(edge, fill=color)
        if tag:
            tw = dr.textlength(tag, font=f) if hasattr(dr, "textlength") else 8 * len(tag)
            if x0 - tw - 14 >= 0:                            # a bal margóba, a doboz mellé
                tx, ty = x0 - tw - 14, y0
            else:                                            # különben a doboz fölé
                tx, ty = x0, max(0, y0 - f.size - 6)
            dr.rectangle([tx, ty, tx + tw + 8, ty + f.size + 5], fill=color)
            dr.text((tx + 4, ty + 1), tag, fill=(255, 255, 255), font=f)


def _header(im, text: str, color=(70, 70, 70)):
    from PIL import Image, ImageDraw
    f = _font(15)
    out = Image.new("RGB", (im.width, im.height + 26), (245, 245, 245))
    ImageDraw.Draw(out).text((8, 5), text, fill=color, font=f)
    out.paste(im, (0, 26))
    return out


def crop(L: Layout, it: dict) -> Optional[bytes]:
    """Egy hiba kivágása: a hely(ek) a lapról, a hiba bekeretezve; két oldalon átnyúló hibánál a két sáv egymás alatt."""
    from PIL import Image
    regions = [r for r in it.get("_regions") or [] if r and r[0] and 0 < r[0] <= L.pages]
    if not regions:
        return None
    width = max(r[1][2] - r[1][0] for r in regions) or 1
    zoom = max(0.8, min(2.2, 860 / width))
    color = _RED if it["severity"] == ERROR else _ORANGE
    parts = []
    for page, rect in regions:
        if rect[3] - rect[1] < 30:
            rect = (rect[0], max(0, rect[1] - 20), rect[2], rect[3] + 20)
        im, clip = _render(L, page, rect, zoom)
        mk = [m for m in it.get("_marks") or [] if m and m[0] == page]
        _draw_marks(im, clip, zoom, [(m[1], color, it["id"] if k == 0 else "") for k, m in enumerate(mk)])
        parts.append(_header(im, f"{it['id']} · {L.unit} {page}"))
    if len(parts) == 1:
        return _jpeg(parts[0])
    w = max(p.width for p in parts)
    out = Image.new("RGB", (w, sum(p.height for p in parts) + 8 * (len(parts) - 1)), (190, 190, 190))
    y = 0
    for p in parts:
        out.paste(p, (0, y))
        y += p.height + 8
    return _jpeg(out)


def page_image(L: Layout, page: int, issues: list[dict], width: int = 1000) -> bytes:
    w, h = L.sizes[page - 1]
    zoom = width / w
    im, clip = _render(L, page, (0, 0, w, h), zoom)
    marks = [(m[1], _RED if it["severity"] == ERROR else _ORANGE, it["id"] if k == 0 else "")
             for it in issues if it.get("view") == L.view
             for k, m in enumerate([m for m in it.get("_marks") or [] if m and m[0] == page])]
    _draw_marks(im, clip, zoom, marks)
    return _jpeg(_header(im, f"{L.unit.capitalize()} {page} of {L.pages} — {view_name(L)}"), q=72)


def side_by_side(A: Layout, B: Layout, page: int, issues: list[dict]) -> bytes:
    from PIL import Image
    ims = [Image.open(io.BytesIO(page_image(L, page, issues, 760))) for L in (A, B) if page <= L.pages]
    if len(ims) == 1:
        return _jpeg(ims[0])
    out = Image.new("RGB", (ims[0].width + ims[1].width + 10, max(i.height for i in ims)), (190, 190, 190))
    out.paste(ims[0], (0, 0))
    out.paste(ims[1], (ims[0].width + 10, 0))
    return _jpeg(out)


def contact_sheet(L: Layout, issues: list[dict], limit: int = 30) -> tuple[bytes, list[int]]:
    """Minden oldal kicsiben, oldalszámmal; a hibás oldal piros, a figyelmeztetéses narancs keretet kap."""
    from PIL import Image, ImageDraw
    sev: dict[int, str] = {}
    tags: dict[int, list[str]] = defaultdict(list)
    for it in issues:
        p = (it.get("pages") or {}).get(L.view) or (it["page"] if it.get("view") == L.view else None)
        if p:
            tags[p].append(it["id"])
            if sev.get(p) != ERROR:
                sev[p] = it["severity"]
    pages = list(range(1, L.pages + 1))
    if len(pages) > limit:
        pages = sorted(set(sorted(sev)[:limit] + pages[: max(0, limit - len(sev))]))[:limit]
    cols = 4 if len(pages) <= 8 else 5 if len(pages) <= 20 else 6
    tw = 200 if cols <= 5 else 170
    f = _font(14)
    thumbs = []
    for p in pages:
        w, h = L.sizes[p - 1]
        im, _ = _render(L, p, (0, 0, w, h), tw / w)
        color = _RED if sev.get(p) == ERROR else _ORANGE if p in sev else (200, 200, 200)
        frame = Image.new("RGB", (im.width + 8, im.height + 30), (255, 255, 255))
        dr = ImageDraw.Draw(frame)
        dr.rectangle([0, 0, frame.width - 1, im.height + 7], outline=color, width=4 if p in sev else 1)
        frame.paste(im, (4, 4))
        label = f"{p}" + (f"  {' '.join(tags[p][:4])}" + ("…" if len(tags[p]) > 4 else "") if tags[p] else "")
        dr.text((6, im.height + 9), label, fill=color if p in sev else (90, 90, 90), font=f)
        thumbs.append(frame)
    cw, ch = max(t.width for t in thumbs), max(t.height for t in thumbs)
    rows = -(-len(thumbs) // cols)
    sheet = Image.new("RGB", (cols * (cw + 10) + 10, rows * (ch + 10) + 10), (235, 235, 235))
    for k, t in enumerate(thumbs):
        sheet.paste(t, (10 + (k % cols) * (cw + 10), 10 + (k // cols) * (ch + 10)))
    return _jpeg(_header(sheet, f"All {L.unit}s — {view_name(L)}" if len(pages) == L.pages
                         else f"{len(pages)} of {L.pages} {L.unit}s (those with problems first)"), q=74), pages


def _source(credit: Optional[str]) -> str:
    """„Photo: X (Wikimedia Commons), CC BY-SA — …/File:Név.jpg” → „Wikimedia Commons: Név.jpg”."""
    if not credit:
        return ""
    m = re.search(r"\(([^)]+)\).*?File:([^\s]+)", credit)
    if m:
        from urllib.parse import unquote
        return f"{m.group(1)}: {unquote(m.group(2)).replace('_', ' ')}"
    return _short(credit, 70)


def _notes_credits(file_id: str) -> dict[int, list[str]]:
    """Diánként a jegyzetekbe írt képforrások („Photo: …”, a pptx_build és az insert_image teszi oda)."""
    from app.tools.crafts import existing
    src = existing(file_id)
    if src.suffix.lower() != ".pptx":
        return {}
    from pptx import Presentation
    out = {}
    for i, sl in enumerate(Presentation(str(src)).slides, 1):
        if sl.has_notes_slide:
            got = [ln.strip() for ln in sl.notes_slide.notes_text_frame.text.splitlines() if ln.strip().startswith("Photo:")]
            if got:
                out[i] = got
    return out


def picture_sheet(file_id: str, limit: int = 24) -> tuple[Optional[bytes], list[dict[str, Any]]]:
    """A fájl MINDEN képe egy lapon, a helyével (dia/oldal, slot), a keresőszóval és a forrással. Az ügynök maga nézi meg
    és ítél (10-13, Kommandant: „tudja utólag is ellenőrizni az agent minden diát — képet, stb.”: a Deutsche Post-bélyeg
    és a „Lorem Ipsum” mockup átment, mert a tördelés-ellenőrző a képek tartalmát nem nézi). Gépi ítélet NINCS benne."""
    from PIL import Image, ImageDraw, ImageOps
    from app.tools import pictures
    try:
        items = pictures._list_sync(file_id)
    except Exception as e:
        log.info("képlista kimaradt (%s): %s", file_id, e)
        return None, []
    rows, tiles = [], []
    f, fb = _font(15), _font(17)
    notes = _notes_credits(file_id) if any(not it.get("credit") for it in items) else {}
    for it in items[:limit]:
        if not it.get("credit"):
            m = re.match(r"Slide (\d+)", it.get("where") or "")
            got = notes.get(int(m.group(1))) if m else None
            if got:
                it = {**it, "credit": got[0] if len(got) == 1 else " · ".join(got)}
        try:
            im = Image.open(io.BytesIO(pictures._image_sync(file_id, it["slot"]))).convert("RGB")
        except Exception:
            continue
        info = {"slot": it["slot"], "where": it.get("where"), "query": it.get("query") or None,
                "source": _source(it.get("credit")) or None, "slide_title": it.get("hint") or None}
        rows.append({k: v for k, v in info.items() if v})
        tile = Image.new("RGB", (430, 400), "white")
        tile.paste(ImageOps.contain(im, (420, 290)), (5, 5))
        dr = ImageDraw.Draw(tile)
        dr.text((8, 300), f"{it.get('where') or ''} · slot {it['slot']}", fill=(20, 20, 20), font=fb)
        y = 324
        for label, val in (("slide", it.get("hint")), ("searched", it.get("query")), ("source", _source(it.get("credit")))):
            if val:
                dr.text((8, y), _short(f"{label}: {val}", 54), fill=(80, 80, 80), font=f)
                y += 22
        tiles.append(tile)
    if not tiles:
        return None, rows
    cols = 3 if len(tiles) > 4 else 2 if len(tiles) > 1 else 1
    sheet = Image.new("RGB", (cols * 440 + 10, ((len(tiles) + cols - 1) // cols) * 410 + 10), (230, 230, 230))
    for k, t in enumerate(tiles):
        sheet.paste(t, (10 + (k % cols) * 440, 10 + (k // cols) * 410))
    more = f" (first {limit} of {len(items)})" if len(items) > limit else ""
    return _jpeg(_header(sheet, f"Every picture in the file{more} — where it is, what was searched, where it comes from"), q=76), rows


# ---------- belépési pontok ----------

async def _cached(src: Path, tag: str, make) -> Path:
    from app.storage import temp_path
    from app.tools.previews import _fresh
    dst = temp_path(f"vd-{src.stem}-{src.suffix.lstrip('.')}.{tag}")
    dst.parent.mkdir(parents=True, exist_ok=True)
    if not _fresh(dst, src):
        await make(src, dst)
    return dst


async def _as_docx(src: Path) -> Path:
    from app.tools import anyconvert, wordcraft
    ext = src.suffix.lstrip(".").lower()
    if ext == "docx":
        return src
    return await _cached(src, "docx", anyconvert.lo_convert if ext == "doc" else wordcraft.convert_file)


async def _layouts(src: Path, kind: str, want_file: bool = True) -> tuple[list[Layout], list[str]]:
    """→ ([szerkesztő, letöltött fájl] tördelés, megjegyzések). A LibreOffice-út hibája nem állítja meg az ellenőrzést."""
    from app.tools import anyconvert, previews
    notes, out = [], []
    jobs = [previews.as_pdf(src)] + ([_cached(src, "lo.pdf", anyconvert.lo_convert)] if want_file else [])
    res = await asyncio.gather(*jobs, return_exceptions=True)
    names = ["editor", "word" if kind == "document" else "powerpoint"]
    unit = "slide" if kind == "deck" else "page"
    for name, r in zip(names, res):
        if isinstance(r, Exception):
            label = view_name(Layout(name, Path("."), unit))
            notes.append(f"{label[0].upper()}{label[1:]} could not be laid out ({str(r)[:120]}) — checked without it.")
            continue
        out.append(await asyncio.to_thread(read_pdf, r, name, unit))
    return out, notes


def _analyze_doc(docx: Path, layouts: list[Layout]) -> tuple[list[dict], DocCheck]:
    blocks, sections = read_docx(docx)
    _wordcraft_index(docx, blocks)
    dc = DocCheck(blocks, sections, layouts)
    return dc.run(), dc


async def quick_layout(path: Path) -> Optional[dict[str, Any]]:
    """Mentés utáni gyors tördelés-ellenőrzés (csak a szerkesztő motorja, kép nélkül) — a `check` mezőhöz."""
    layouts, _ = await _layouts(path, "document", want_file=False)
    if not layouts:
        return None
    issues, _ = await asyncio.to_thread(_analyze_doc, path, layouts)
    return {"errors": [i for i in issues if i["severity"] == ERROR], "warnings": [i for i in issues if i["severity"] == WARNING],
            "pages": layouts[0].pages}


def counts(issues: list[dict]) -> str:
    ne = sum(1 for i in issues if i["severity"] == ERROR)
    nw = len(issues) - ne
    return f"{ne} error{'' if ne == 1 else 's'}, {nw} warning{'' if nw == 1 else 's'}"


CHECKED = {
    "document": ["empty list items and headings", "numbering (repeats, skips, restarts, a list continuing under a new heading, "
                 "typed numbers, heading numbers)", "table rows split across pages, table headers not repeated",
                 "headings left as the last line of a page, captions apart from their picture/table",
                 "text, tables and pictures off the page or into the margins", "empty pages",
                 "the built-in editor's layout against Word's (LibreOffice) — page count, numbering, where problems appear"],
    "deck": ["text overflowing its box (in the actual rendered slide)", "things off the slide", "empty placeholders",
             "empty slides", "the editor's slides against PowerPoint's (LibreOffice)"],
    "pdf": ["empty pages", "text and pictures off the page"],
}


async def validate(file_id: str, pages: Optional[list[int]] = None, view: str = "word",
                   max_crops: int = 6) -> tuple[dict[str, Any], list[tuple[str, bytes]]]:
    """→ (összegzés, [(képaláírás, JPEG)]). A képek: áttekintő lap, hibánként kivágás, a kért (vagy hibás) oldalak."""
    from app.storage import display_label
    from app.tools.crafts import existing
    src = existing(file_id)
    ext = src.suffix.lstrip(".").lower()
    notes: list[str] = []
    if ext in WORDISH:
        kind = "document"
        docx = await _as_docx(src)
        layouts, notes = await _layouts(docx, kind)
        if not layouts:
            raise RuntimeError("The document could not be laid out in either engine. " + " ".join(notes))
        issues, dc = await asyncio.to_thread(_analyze_doc, docx, layouts)
    elif ext in ("pptx", "deckcraft", "ppt", "odp"):
        kind = "deck"
        if ext != "pptx":
            raise ValueError("validate_document checks .pptx decks — convert this one to pptx first (convert).")
        layouts, notes = await _layouts(src, kind)
        if not layouts:
            raise RuntimeError("The deck could not be laid out. " + " ".join(notes))
        issues, dc = await asyncio.to_thread(deck_issues, src, layouts), None
    elif ext == "pdf":
        kind = "pdf"
        layouts = [await asyncio.to_thread(read_pdf, src, "pdf", "page")]
        issues, dc = await asyncio.to_thread(pdf_issues, layouts[0]), None
    else:
        raise ValueError(f"validate_document checks Word documents, PowerPoint decks and PDFs — for a .{ext} file look at "
                         "files(action=\"preview\").")
    try:
        return await asyncio.to_thread(_assemble, file_id, display_label(file_id), kind, layouts, issues, dc, notes,
                                       pages, view, max_crops)
    finally:
        for L in layouts:
            L.close()


def _assemble(file_id: str, name: str, kind: str, layouts: list[Layout], issues: list[dict], dc: Optional[DocCheck],
              notes: list[str], pages: Optional[list[int]], view: str, max_crops: int) -> tuple[dict, list[tuple[str, bytes]]]:
    by_view = {L.view: L for L in layouts}
    v = (view or "word").lower()
    v = {"file": "word", "export": "word", "downloaded": "word", "office": "word"}.get(v, v)
    if kind == "deck" and v == "word":
        v = "powerpoint"
    main = by_view.get(v) or layouts[-1]                     # alapból a letöltött fájl nézete
    images: list[tuple[str, bytes]] = []
    sheet, on_sheet = contact_sheet(main, issues)
    what = (f"the {main.unit}" if main.pages == 1 else f"all {main.pages} {main.unit}s" if len(on_sheet) == main.pages
            else f"{len(on_sheet)} of {main.pages} {main.unit}s (those with problems first)")
    images.append((f"Overview: {what} as {view_name(main)} — red frame = error, orange = warning.", sheet))
    pic_sheet, pic_rows = picture_sheet(file_id) if kind in ("deck", "document") else (None, [])
    if pic_sheet:
        images.append(("Pictures: every picture in the file, with where it is, what was searched and its source. LOOK AT "
                       "EACH ONE — the layout check above does not judge pictures: does it fit its slide or page? Watch for "
                       "placeholder text (“lorem ipsum”), watermarks (“designed by …”), stamps, logos or institutions from "
                       "elsewhere, and unrelated subjects. Swap one with insert_image(file_id, slot=…, image_url | "
                       "image_file_id | ai_prompt), or show the person candidates with generate_image(stock_search=…).", pic_sheet))
    shown = 0
    for it in issues:
        if shown >= max(0, int(max_crops)):
            break
        L = by_view.get(it.get("view")) or main
        jpg = crop(L, it) if it.get("view") else None
        if jpg:
            shown += 1
            it["picture"] = len(images)
            images.append((f"{it['id']} ({it['severity']}) · {L.unit} {it['page']} in {('the editor' if L.view == 'editor' else 'the file')}: "
                           f"{it['message']}", jpg))
    want = [p for p in (pages or []) if isinstance(p, int)][:4]
    if not want:
        want = sorted({it["page"] for it in issues if it["severity"] == ERROR and it.get("page") and it.get("view") == main.view})[:2]
    for p in want:
        if not 0 < p <= main.pages:
            continue
        if v == "both" and len(layouts) == 2:
            images.append((f"{main.unit.capitalize()} {p}: the built-in editor (left) and the downloaded file (right), problems marked.",
                           side_by_side(layouts[0], layouts[1], p, issues)))
        else:
            L = by_view.get(v) or main
            images.append((f"{L.unit.capitalize()} {p} of {L.pages} as {view_name(L)}, problems marked.", page_image(L, p, issues)))
    ne = sum(1 for i in issues if i["severity"] == ERROR)
    nw = len(issues) - ne
    status = "failed" if ne else "needs_check" if nw else "ok"
    summary = counts(issues)
    agree = None
    if dc is not None and len(layouts) == 2:
        a, b = dc.page_starts(layouts[0]), dc.page_starts(layouts[1])
        k = next((i for i, (x, y) in enumerate(zip(a, b))
                  if (x is None) != (y is None) or (x and y and (x[0] != y[0] or abs(x[1] - y[1]) > 3))), None)
        agree = ("The editor and Word break every page in the same place." if k is None and len(a) == len(b) else
                 f"The editor and Word break the pages the same way up to page {k if k is not None else min(len(a), len(b))}; "
                 "after that the page breaks differ.")
    verified = [f"Checked in {len(layouts)} layout{'s' if len(layouts) > 1 else ''}: "
                + "; ".join(f"{view_name(L)} — {L.pages} {L.unit}{'s' if L.pages != 1 else ''}" for L in layouts) + "."]
    if not issues:
        verified.append("No problems found: " + "; ".join(CHECKED[kind]) + ".")
    out = {
        "file_id": file_id, "name": name, "status": status, "summary": summary, "errors": ne, "warnings": nw,
        "pages": {L.view: L.pages for L in layouts},
        "views": {L.view: view_name(L) for L in layouts},
        **({"page_breaks": agree} if agree else {}),
        "issues": [public(i) for i in issues],
        **({"pictures": pic_rows} if pic_rows else {}),
        "checked": CHECKED[kind] + (["pictures: all shown on one sheet for YOU to judge — not judged by the machine"]
                                    if pic_rows else []),
        "images": [c for c, _ in images],
        "check": {"status": status, "findings": [f"{i['id']} page {i['page']}: {i['message']}" if i.get("page")
                                                 else f"{i['id']}: {i['message']}" for i in issues][:30],
                  "verified": verified},
        "note": ("Look at the picture sheet as well — pictures are yours to judge. " if pic_rows else "")
                + ("Fix every error, then call validate_document again. A file that saved is not proof the layout is right — "
                   "these results are. Word itself was not run: LibreOffice lays the file out like Word in almost every case, "
                   "but fonts and spacing can shift a line or two." if kind == "document" else
                   "Fix every error, then call validate_document again."),
    }
    if notes:
        out["warnings_about_the_check"] = notes
    if pages and not want:
        out["note"] += f" There is no page {pages[0]}."
    return out, images
