"""Videó a Studióban — storytold EffectCraft (After Effects-szerű, .ecproj) és FilmCraft (Premiere-szerű).

Kész upstream release-ek (Apache-2.0), mint a többi Craft-motor: a CLI a konténerben fut, a szerkesztő
a böngészőben (/editor/effectcraft/, /editor/filmcraft/). A videó SZERKESZTHETŐ idővonallal készül:
az eredmény egy .ecproj (jelenetenként beágyazott kompozíció, feliratsávok, hang) + a renderelt MP4.

    make()    cikk / URL / dokumentum / saját szöveg → forgatókönyv (DeepSeek, JSON) → képek (a felhasználóé,
              stock, MCP-n a kliens AI-képe) → video_builder.jsx → `effectcraft-cli script` → .ecproj
              → `effectcraft-cli render` → MP4
    overlay() a felhasználó felvétele (mp4/mov/webm/m4v) az alapsáv, fölötte időkódos rétegek (névsáv, fejezetcím,
              forrásfelirat, szám, idézet, ábra, teljes képes bevágás, osztott kép) → .ecproj + MP4
    render()  egy (akár kézzel szerkesztett) projekt újrarenderelése MP4-re
    blank()   üres videóprojekt (álló / fekvő / négyzetes) a szerkesztőbe
    info()    kompozíciók, hossz, jelenetek, rétegek, médiafájlok

Rétegek (layers): EGY modell a make- és az overlay-videón — abszolút időben a fő idővonalon, rétegenként egy
kompozíció, a meglévő ábra-építőkkel (chart_elements); alpha = átlátszó, csak grafikás változat (ProRes 4444 / WebM).

A projekt saját médiafájljai (képek, zene, hang, felvétel) rejtett fájlok a munkaterületen; a projekt metája tartja a
név → file_id párokat („assets”). Az .ecproj relatív névvel hivatkozik rájuk, így renderkor egy ideiglenes
mappába rakjuk össze (stage), a böngészős szerkesztőnek pedig a bridge adja át őket (addFile) a projekt előtt.

Mérve 10-12 (16 mag): 8 s 1080×1920 render 16 s; a CLI a termék Docker-képében (jammy) gond nélkül fut.
AI-narráció: offline Piper-hangok (narration.py), alapból BE; saját hangfájl is mehet alá, feliratok mindig vannak.
Időtartam: duration_mode fixed (a kért hossz) | narration (a hang szabja); az eredmény őszintén jelzi az eltérést.
Remotion NEM (céges licenc).
"""
from __future__ import annotations

import asyncio
import contextvars
import copy
import functools
import json
import logging
import math
import os
import re
import shutil
import tempfile
import time
from pathlib import Path
from typing import Any, Callable, Optional

from app.config import settings
from app.storage import display_label, output_path, public_url, read_meta, temp_path, write_meta
from app.storage.files import new_file_id
from app.tools import crafts

log = logging.getLogger("aice.video")

FORMATS = {"vertical": (1080, 1920), "landscape": (1920, 1080), "square": (1080, 1080)}
_ALIASES = {"9:16": "vertical", "story": "vertical", "stories": "vertical", "reel": "vertical", "reels": "vertical",
            "shorts": "vertical", "tiktok": "vertical", "portrait": "vertical", "16:9": "landscape", "youtube": "landscape",
            "wide": "landscape", "horizontal": "landscape", "1:1": "square", "post": "square", "instagram": "square"}
FPS = 30
XFADE = 0.5                            # jelenetek közti áttűnés (s)
MIN_TOTAL = 4.0                        # a legrövidebb kérhető videó (cím + egy jelenet + zárókép) — korábban 10 s volt a
                                       # padló, és a 8 s-os kérésből 10–13 s lett
MIN_SCENE = 1.2                        # szoros fix hossznál ennél rövidebb jelenet nem olvasható
TICKS_PER_S = 254_016_000_000          # az .ecproj időegysége (mint a Premiere „tick”-je)
MOVIE_EXTS = {"mp4", "mov", "webm", "m4v", "mkv"}
AUDIO_EXTS = {"mp3", "wav", "m4a", "aac", "ogg", "oga", "flac", "aiff", "aif", "opus"}
EFFECT_OPEN = {"ecproj"}
FILM_OPEN = {"fcproj"} | MOVIE_EXTS
# amit a két szerkesztő menthet / renderelhet (a mentés-végpont ezt engedi be)
VIDEO_SAVE = {"ecproj", "fcproj", "mp4", "mov", "webm", "gif", "wav", "aiff", "otio", "lottie"}
FONT = {"bold": "Inter-Bold", "semibold": "Inter-SemiBold", "regular": "Inter-Regular"}   # az EffectCraftba csomagolt Inter
DEFAULT_ACCENT = "#FF5A36"
BUILDER_JS = Path(__file__).with_name("video_builder.jsx")
MAX_FETCH = 6 * 1024 * 1024

_gate: Optional[asyncio.Semaphore] = None     # egyszerre egy render/építés: CPU-zabáló, a párhuzamosság nem gyorsít


def _render_gate() -> asyncio.Semaphore:
    global _gate
    if _gate is None:
        _gate = asyncio.Semaphore(1)
    return _gate


def effect_cli() -> Optional[str]:
    return crafts.cli_path(settings.EFFECTCRAFT_CLI, "effectcraft-cli")


def film_cli() -> Optional[str]:
    return crafts.cli_path(settings.FILMCRAFT_CLI, "filmcraft-cli")


def available() -> bool:
    return effect_cli() is not None


def fmt_name(fmt: Optional[str]) -> str:
    f = (fmt or "vertical").strip().lower()
    f = _ALIASES.get(f, f)
    if f not in FORMATS:
        raise ValueError(f"format: vertical (9:16) | landscape (16:9) | square (1:1) — not {fmt!r}")
    return f


_JAIL: list[bool] = []


async def _jail_ok(exe: str) -> bool:
    """Egyszeri önteszt: a motor elindul-e a zárkában (1+1). Ha nem — a megbízható munka zárka nélkül fut tovább
    (naplózva), idegen szkript viszont nem."""
    from app.tools import jail
    if not _JAIL:
        ok = False
        if jail.available():
            work = Path(tempfile.mkdtemp(prefix="jailtest-", dir=str(temp_path("").parent)))
            try:
                argv, env = jail.command(exe, ["script", "--eval", "1+1", "--empty"], work=str(work))
                proc = await asyncio.create_subprocess_exec(*argv, cwd=str(work), env=env, stdout=asyncio.subprocess.PIPE,
                                                            stderr=asyncio.subprocess.PIPE)
                out, err = await asyncio.wait_for(proc.communicate(), timeout=30)
                ok = proc.returncode == 0 and out.decode("utf-8", "replace").strip().endswith("2")
                if not ok:
                    log.error("a videómotor zárkája nem indul (rc %s): %s", proc.returncode, err.decode("utf-8", "replace")[-300:])
            except Exception as e:
                log.error("a videómotor zárkája nem indul: %s", e)
            finally:
                shutil.rmtree(work, ignore_errors=True)
        _JAIL.append(ok)
        log.info("videómotor zárkában: %s (Landlock ABI %s)", ok, jail.abi())
    return _JAIL[0]


async def _run(*args: str, cwd: Path, timeout: Optional[int] = None, untrusted: bool = False) -> str:
    """effectcraft-cli a megadott mappában, ZÁRKÁBAN (jail.py): csak a munkamappát írja, rajta kívül csak a
    rendszerkönyvtárakat és a mappába linkelt médiát olvassa, hálózat és alfolyamat nélkül, a környezeti változók
    (kulcsok) nélkül. A motor saját kapuja is él (a szkriptek csak a projekt mappájából olvasnak, nem írnak, nem
    hálóznak) — a zárka a második fal. untrusted (az ügynök saját szkriptje): zárka nélkül NEM fut."""
    from app.tools import jail
    exe = effect_cli()
    if exe is None:
        raise RuntimeError("The video engine (EffectCraft) isn't installed on this server")
    argv, env = [exe, *args], {**os.environ, "EFFECTCRAFT_NO_PREFS": "1"}
    if await _jail_ok(exe):
        argv, env = jail.command(exe, list(args), work=str(cwd), env={"EFFECTCRAFT_NO_PREFS": "1"},
                                 mem_gb=8.0 if untrusted else None)
    elif untrusted:
        raise RuntimeError("Custom scripts need the sandbox, and it isn't available on this server — use a graphic (type \"graphic\") instead.")
    else:
        log.warning("zárka nélkül fut a videómotor (a kernel nem tud Landlockot)")
    proc = await asyncio.create_subprocess_exec(*argv, cwd=str(cwd), env=env,
                                                stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
    try:
        out, err = await asyncio.wait_for(proc.communicate(), timeout=timeout or settings.VIDEO_TIMEOUT_S)
    except asyncio.TimeoutError:
        proc.kill()
        await proc.wait()
        raise RuntimeError(f"effectcraft-cli {args[0]}: timed out")
    if proc.returncode != 0:
        msg = err.decode("utf-8", "replace").strip() or out.decode("utf-8", "replace").strip()
        raise RuntimeError(f"effectcraft-cli {args[0]}: {msg[-600:] or f'exit code {proc.returncode}'}")
    return out.decode("utf-8", "replace")


# ---------- a projekt médiafájljai ----------

def assets_of(file_id: str) -> dict[str, str]:
    return dict((read_meta(file_id) or {}).get("assets") or {})


def add_asset(project_id: str, name: str, src: Path) -> str:
    """Egy médiafájl a projekthez: rejtett fájl a munkaterületen + név → file_id a projekt metájában."""
    name = _asset_name(name)
    fid = new_file_id(name.rsplit(".", 1)[-1].lower() if "." in name else "bin")
    shutil.copyfile(src, output_path(fid))
    write_meta(fid, operation="video asset", label=name, source_file_id=project_id, extra={"ephemeral": True})
    write_meta(project_id, extra={"assets": {**assets_of(project_id), name: fid}})
    return fid


def carry_assets(source_id: Optional[str], new_id: str) -> None:
    """Szerkesztőbeli mentés (.ecproj) → az új változat ugyanazokat a médiafájlokat látja."""
    if not source_id:
        return
    meta = read_meta(source_id) or {}
    extra = {k: meta[k] for k in ("assets", "video", "spec") if meta.get(k)}
    if extra.get("spec"):
        extra["hand_edited"] = True            # az ügynöki szerkesztés a forgatókönyvből épít újra — ezt jelezni kell
    if extra:
        write_meta(new_id, extra=extra)


def _asset_name(name: str) -> str:
    name = Path(str(name).replace("\\", "/")).name.strip() or "media"
    return re.sub(r"[^\w.\- ]+", "_", name)[:120]


def stage(project_id: str, workdir: Path) -> Path:
    """Projekt + médiafájlok egy mappába; a médiahivatkozások a fájlnévre igazítva (a böngészőben importált
    fájl útvonala lehet /név is). → a projekt útvonala."""
    src = crafts.existing(project_id)
    assets = assets_of(project_id)
    for name, fid in assets.items():
        try:
            p = crafts.existing(fid)
        except FileNotFoundError:
            log.warning("hiányzó videó-médiafájl: %s (%s)", name, fid)
            continue
        dst = workdir / name
        if not dst.exists():
            try:
                os.symlink(p, dst)
            except OSError:
                shutil.copyfile(p, dst)
    data = json.loads(src.read_text("utf-8"))
    for item in (data.get("items") or {}).values():
        kind = item.get("kind") or {}
        if kind.get("type") == "Footage" and kind.get("path"):
            base = Path(str(kind["path"]).replace("\\", "/")).name
            if base in assets:
                kind["path"] = base
    out = workdir / "project.ecproj"
    out.write_text(json.dumps(data, ensure_ascii=False), "utf-8")
    return out


def _comps(project: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for item in (project.get("items") or {}).values():
        k = item.get("kind") or {}
        if k.get("type") == "Comp":
            out.append({"name": item.get("name"), "width": k.get("width"), "height": k.get("height"),
                        "duration": round((k.get("duration") or 0) / TICKS_PER_S, 2), "layers": len(k.get("layers") or [])})
    return out


def _main_comp(project_path: Path) -> str:
    comps = _comps(json.loads(project_path.read_text("utf-8")))
    names = [c["name"] for c in comps]
    if "Video" in names:
        return "Video"
    return max(comps, key=lambda c: (c["layers"], c["duration"]))["name"] if comps else "-"


async def probe_movie(path: Path) -> dict[str, Any]:
    """Egy felvétel adatai (ffprobe): méret (az elforgatás szerint), képkocka/s, hossz, van-e hang, pixelformátum."""
    from app.convert.engines import run
    out, _ = await run("ffprobe", "-v", "error", "-print_format", "json", "-show_format", "-show_streams", str(path), timeout=60)
    d = json.loads(out.decode("utf-8", "replace") or "{}")
    vs = [s for s in d.get("streams") or [] if s.get("codec_type") == "video" and not (s.get("disposition") or {}).get("attached_pic")]
    if not vs:
        raise ValueError("That file has no video picture — give an MP4, MOV, WebM or M4V video.")
    v = vs[0]
    w, h = int(v.get("width") or 0), int(v.get("height") or 0)
    rot = 0
    for sd in v.get("side_data_list") or []:
        if sd.get("rotation") is not None:
            rot = int(float(sd["rotation"]))
    rot = int(float((v.get("tags") or {}).get("rotate") or rot))
    if abs(rot) % 180 == 90:
        w, h = h, w
    fr = lambda x: (lambda a, b: float(a) / float(b) if float(b) else 0.0)(*(str(x or "0/1").split("/") + ["1"])[:2])
    fps = fr(v.get("avg_frame_rate")) or fr(v.get("r_frame_rate")) or FPS
    dur = float((d.get("format") or {}).get("duration") or v.get("duration") or 0)
    tags = v.get("tags") or {}
    alpha_tag = str(tags.get("alpha_mode") or tags.get("ALPHA_MODE") or "") == "1"
    # az első DEKÓDOLHATÓ hangsáv (a hangsávok közti sorszám): az iPhone a sima AAC mellé térhangot (APAC) is ír,
    # amit az ffmpeg nem ismer — „-map 0:a” vele együtt elhasal (10-18, IMG_9271: „Decoder (codec none) not found”)
    aud = [s for s in d.get("streams") or [] if s.get("codec_type") == "audio"]
    known = [k for k, s in enumerate(aud) if s.get("codec_name") not in (None, "", "unknown", "none")]
    return {"audio_stream": known[0] if known else None,
            "width": w, "height": h, "fps": round(fps, 3), "duration": round(dur, 3), "rotation": rot,
            "audio": any(s.get("codec_type") == "audio" for s in d.get("streams") or []),
            "codec": v.get("codec_name"), "pix_fmt": v.get("pix_fmt"), "alpha_mode": alpha_tag}


def _tc(t: float) -> str:
    """Másodperc → időkód a felhasználónak („00:06”, „1:02:03”, tört másodperccel, ha van)."""
    t = max(0.0, float(t))
    h, rest = divmod(t, 3600)
    mnt, s = divmod(rest, 60)
    ss = f"{s:04.1f}" if abs(s - round(s)) > 0.05 else f"{int(round(s)):02d}"
    return (f"{int(h)}:{int(mnt):02d}:" if h else f"{int(mnt):02d}:") + ss


def timecode(v: Any, what: str = "at") -> float:
    """Időpont: másodperc (84 · 84.5 · "84.5s") vagy időkód ("01:24" · "1:02:03.5" · "00:01:24,5") → másodperc."""
    bad = ValueError(f"{what}: {v!r} is not a time — give seconds (84.5) or a timecode (\"01:24\", \"1:02:03.5\").")
    if isinstance(v, bool) or v is None:
        raise bad
    if isinstance(v, (int, float)):
        if v < 0 or v != v:
            raise bad
        return float(v)
    s = str(v).strip().lower().rstrip("s").strip().replace(",", ".")
    parts = s.split(":")
    if not s or len(parts) > 3 or not all(re.fullmatch(r"\d+(?:\.\d+)?", p) for p in parts):
        raise bad
    if any("." in p for p in parts[:-1]) or (len(parts) > 1 and float(parts[-1]) >= 60) or (len(parts) == 3 and int(parts[1]) >= 60):
        raise bad
    out = 0.0
    for p in parts:
        out = out * 60 + float(p)
    return out


# ---------- forrás: cikk, URL, dokumentum ----------

async def fetch_article(url: str) -> dict[str, Any]:
    """Egy weboldal cikk-szövege (címsorok + bekezdések), címe, fő képe (og:image) és a webhely neve."""
    import httpx
    from bs4 import BeautifulSoup
    from app import net
    browser = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/130 Safari/537.36"
    host = (httpx.URL(url).host or "").lower()
    # 10-12 élesben: a Wikipedia 403-mal tiltja az adatközpontból jövő, böngészőnek álcázott kérést — a Wikimedia
    # a kapcsolatot megadó, őszinte User-Agentet kéri; máshol a böngésző-UA az első, 403/429 után az őszinte.
    # A letöltés a közös kimenő klienssel megy (SSRF-védelem: belső címre nem kér).
    uas = ([net.user_agent(), browser] if host.endswith(("wikipedia.org", "wikimedia.org", "wikinews.org", "wikivoyage.org"))
           else [browser, net.user_agent()])
    got, code = None, "no answer"
    for ua in uas:
        try:
            got = await net.fetch(url, max_bytes=max(MAX_FETCH, 5 * 1024 * 1024), timeout=25, retries=1,
                                  headers={"User-Agent": ua, "Accept-Language": "en,hu;q=0.8,*;q=0.5"})
            break
        except httpx.HTTPStatusError as e:
            code = e.response.status_code
            if code not in (401, 403, 429):
                break
    if got is None:
        raise ValueError(f"That page wouldn't let us read it ({code}) — paste the article's text instead.")
    m = re.search(r"charset=([\w-]+)", got.content_type or "")
    html = got.content[:MAX_FETCH].decode(m.group(1) if m else "utf-8", "replace")
    soup = BeautifulSoup(html, "lxml")
    meta = lambda *keys: next((m.get("content") for k in keys for m in soup.find_all("meta", attrs={"property": k}) + soup.find_all("meta", attrs={"name": k}) if m.get("content")), None)
    title = meta("og:title", "twitter:title") or (soup.title.string.strip() if soup.title and soup.title.string else "")
    for t in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "figure", "iframe", "svg"]):
        t.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    parts = []
    for el in root.find_all(["h1", "h2", "h3", "p", "li", "blockquote"]):
        txt = re.sub(r"\s+", " ", el.get_text(" ", strip=True))
        if len(txt) >= (12 if el.name.startswith("h") else 40):
            parts.append(("# " if el.name.startswith("h") else "") + txt)
    text = "\n\n".join(dict.fromkeys(parts))
    if len(text) < 200:
        raise ValueError("Couldn't read an article on that page — paste its text instead.")
    return {"title": title or "", "text": text[:30_000], "image": meta("og:image", "twitter:image"),
            "site": meta("og:site_name") or re.sub(r"^www\.", "", httpx.URL(url).host or "")}


# ---------- forgatókönyv ----------

_PLAN = """You plan short social videos as JSON for a motion-graphics builder: full-screen pictures, an on-screen
headline per scene and timed captions. Reply with ONE json object:
{"title": str, "kicker": str, "image_query": str, "scenes": [{"headline": str, "caption": str, "image_query": str, "stat": {"value": str, "label": str} | null,
  "chart": {"type": "column|bar|line|area|donut|ring|icons|timeline|steps|counter|compare|progress|gauge|funnel|proscons|quote|table|waterfall|slope|lines|race|sankey",
            "title": str, "unit": str, "total": number | null, "min": number | null, "max": number | null,
            "items": [{"label": str, "value": number | null, "text": str}]} | null}],
 "end": {"title": str, "subtitle": str}}
Rules:
- exactly {n} scenes; write in the language of the source (unless told otherwise); no invented facts, names or numbers.
- title ≤ 60 characters, catchy; kicker ≤ 24 characters (the topic or the source's name); the top-level image_query is
  the title card's picture — a different one from every scene's.
- headline ≤ 55 characters — the scene's point, not a label. caption = what a narrator would say: 1–2 short sentences,
  ≤ 170 characters, and the captions together tell the story in order.
- image_query: 2–4 plain English words for a stock-photo search — concrete and visual ("wind turbines at sunset",
  "crowded city street"), never abstract.
- stat only when the source gives a striking number (at most 2 scenes); value ≤ 8 characters ("72%", "3.4 bn"), label ≤ 40.
- chart: an animated graphic for ANY topic (science, health, history, sport, culture, business, how-to…), only when
  the source supports it — pick the type that fits:
    column / bar — 2–8 comparable numbers of the SAME kind (bar for long names or many items);
    line — a value over time, in time order;
    donut — parts of one whole (2–6 shares the source gives; leave out "the rest" — it is added automatically);
    ring — ONE percentage worth showing (unit "%", one item: label = what it is);
    icons — "x in N" or a percentage of people/things (one item: value = x; total = N, or unit "%" for x out of 100);
    timeline — 2–6 dated events in order (label = the date exactly as written in the source, text = what happened);
    steps — 3–6 steps of a process or how-to in order (label = short step, text = one short line);
    area — a value over time where the amount matters (filled under the line);
    counter — ONE striking number that counts up (one item: label = what it is);
    compare — 2–3 numbers of the same kind side by side (before/after, A vs B);
    progress — 2–6 separate percentages that do NOT add up to a whole (e.g. "64% agree", "41% use daily");
    gauge — ONE value on a scale (min/max given in the source, or 0–100%), e.g. a risk level or a score;
    funnel — 3–6 shrinking stages of a process (applicants → interviewed → hired);
    proscons — arguments for (label "+") and against (label "-"), each a short text; labels? = the two headings;
    quote — ONE striking sentence quoted WORD FOR WORD from the source (text = the quote, label = who said it);
    waterfall — a starting amount and what added to / took from it (items: the start first, then the changes with
      their sign, e.g. -7); the end is computed;
    slope — the same 2–6 things at two moments: "labels": [the two dates], items [{"label", "from", "to"}];
    lines — 2–4 series over the same periods (e.g. scenarios): "labels": [periods], "series": [{"name", "values"}];
    race — 3–10 things ranked over 3–12 periods: "labels": [periods], "items": [{"label", "values"}];
    sankey — amounts flowing between stages (money from sources → programmes → projects, energy, people): "flows":
      [{"from", "to", "value"}] (2–24 flows, at most 14 nodes, one way only); the node totals are computed.
  Copy every number and date EXACTLY from the source (value = a plain number, no unit); unit = "%" or a short unit
  ("$", "km", "bn") or ""; item labels ≤ 18 characters (ring/icons: ≤ 60); at most 3 chart scenes (4 for 80 s or longer) and never two of the
  same type in a row; a scene has a stat OR a chart, never both; headline = what the graphic shows.
- end.title ≤ 50 characters (the takeaway or a call to action); end.subtitle ≤ 80 (e.g. where to read more)."""

_KEEP = """
The user wrote this text as the video's SCRIPT: keep their wording for the captions (fix only typos), one scene per
paragraph or idea, in their order; write the headlines and image queries yourself."""


_NUM_IN_TEXT = re.compile(r"(?<![\d.,])-?\d{1,3}(?:[ \u00a0\u202f]\d{3})+(?:[.,]\d+)?(?![\d])|-?\d+(?:[.,]\d+)?")


_NUM_WORDS = {w: i for i, ws in enumerate((
    (), ("one", "egy"), ("two", "kettő", "két"), ("three", "három"), ("four", "négy"), ("five", "öt"), ("six", "hat"),
    ("seven", "hét"), ("eight", "nyolc"), ("nine", "kilenc"), ("ten", "tíz"), ("eleven", "tizenegy"), ("twelve", "tizenkettő"),
    ("thirteen", "tizenhárom"), ("fourteen", "tizennégy"), ("fifteen", "tizenöt"), ("sixteen", "tizenhat"),
    ("seventeen", "tizenhét"), ("eighteen", "tizennyolc"), ("nineteen", "tizenkilenc"), ("twenty", "húsz"))) for w in ws}
# a „hat” / „hét” többértelmű (ige, illetve a hét nap) — mégis számít: egy jó ábra ne essen ki, mert a forrás betűvel
# írja („hat hónapon belül” → 6), a laza egyezés legfeljebb egy 6-ost enged át (10-17)
_NUM_WORDS.update({**{f"huszon{w}": 20 + v for w, v in (("egy", 1), ("kettő", 2), ("két", 2), ("három", 3), ("négy", 4),
                                                         ("öt", 5), ("hat", 6), ("hét", 7), ("nyolc", 8), ("kilenc", 9))},
                   "thirty": 30, "harminc": 30, "forty": 40, "negyven": 40, "fifty": 50, "ötven": 50, "sixty": 60,
                   "hatvan": 60, "seventy": 70, "hetven": 70, "eighty": 80, "nyolcvan": 80, "ninety": 90, "kilencven": 90,
                   "hundred": 100, "száz": 100, "thousand": 1000, "ezer": 1000})
_NUM_WORD_RE = re.compile(r"\b(" + "|".join(sorted(_NUM_WORDS, key=len, reverse=True)) + r")\b", re.I)


def _text_numbers(text: str) -> set[float]:
    """A szöveg összes száma (1,8 · 1.8 · 4 194 · 41 670 · 3,35, és a kis számnevek: „one in four", „négyből egy")
    — a grafikon minden értékének itt kell lennie."""
    out = {float(_NUM_WORDS[w.lower()]) for w in _NUM_WORD_RE.findall(text or "")}
    for t in _NUM_IN_TEXT.findall(text or ""):
        t = re.sub(r"[ \u00a0\u202f]", "", t)
        for cand in (t.replace(",", "."), t.replace(",", "")):
            try:
                out.add(round(float(cand), 6))
            except ValueError:
                pass
    return out


_GFX_SHAPES = ("rect", "circle", "line", "arrow", "path", "polygon", "text", "image", "pin", "lottie", "map")
_GFX_ANIMS = ("appear", "pop", "grow", "draw", "slide-left", "slide-right", "slide-up", "slide-down", "pulse", "none")


def _graphic_from_plan(c: dict[str, Any], source: str) -> tuple[Optional[dict[str, Any]], Optional[str]]:
    """Saját grafika elemekből (az ügynök rajzolja le): alakzatok, nyilak, útvonalak, szövegek, képek, tűk, Lottie —
    a képkocka arányában (0–1), mozgással. Ellenőrzés: ismert elemek, ésszerű koordináták, legfeljebb 80 elem, a
    szövegekben csak forrásbeli szám (különben AZ az elem kimarad — a többi marad, a check jelzi)."""
    nums = _text_numbers(source)
    num = lambda x, d=None: float(x) if isinstance(x, (int, float)) else d
    pt = lambda q: [max(-0.2, min(1.2, num(q[0], 0.5))), max(-0.2, min(1.2, num(q[1], 0.5)))] if isinstance(q, (list, tuple)) and len(q) >= 2 else None

    def geo(q: Any) -> Optional[dict[str, Any]]:      # földrajzi hely: {"place": "Dóm tér, Szeged"} | {"lat", "lon"}
        if isinstance(q, dict) and isinstance(q.get("place"), str) and q["place"].strip():
            return {"place": q["place"].strip()[:160]}
        if isinstance(q, dict) and num(q.get("lat")) is not None and num(q.get("lon")) is not None:
            return {"lat": max(-85.0, min(85.0, float(q["lat"]))), "lon": max(-180.0, min(180.0, float(q["lon"])))}
        return None
    out, dropped, warn, expanded = [], [], [], 0
    alias = {"marker": "pin", "point": "pin", "place": "pin", "label": "text", "title": "text", "route": "path",
             "polyline": "path", "dot": "circle", "ellipse": "circle", "square": "rect", "rectangle": "rect", "box": "rect",
             "picture": "image", "photo": "image", "animation": "lottie", "basemap": "map"}
    flat: list[Any] = []
    for e in (c.get("elements") or [])[:80]:
        if isinstance(e, dict) and isinstance(e.get("shape"), str):
            e = {**e, "shape": alias.get(e["shape"].lower(), e["shape"].lower())}
        flat.append(e)
        if isinstance(e, dict) and e.get("shape") == "map":       # a térképbe írt tűk és útvonalak külön elemként
            for key in ("pins", "markers", "places", "points_of_interest"):
                for q in e.get(key) or []:
                    q = {"place": q} if isinstance(q, str) else q
                    if isinstance(q, dict):
                        flat.append({"shape": "pin", "anim": "pop", **q, "text": q.get("text") or q.get("label") or q.get("name") or ""})
                        expanded += 1
            for key in ("route", "routes", "paths", "lines"):
                rs = e.get(key) or []
                rs = [rs] if rs and not isinstance(rs[0], list) and not (isinstance(rs[0], dict) and "points" in rs[0]) else rs
                for r_ in rs:
                    pts_ = r_.get("points") if isinstance(r_, dict) else r_
                    if isinstance(pts_, list) and len(pts_) >= 2:
                        pts_ = [{"place": q} if isinstance(q, str) else q for q in pts_]
                        flat.append({"shape": "path", "stroke": "accent", "width": 8, "anim": "draw",
                                     **({k: v for k, v in r_.items() if k != "points"} if isinstance(r_, dict) else {}), "points": pts_})
                        expanded += 1
    for n_, e in enumerate(flat[:120], 1):
        if not isinstance(e, dict) or e.get("shape") not in _GFX_SHAPES:
            what = e.get("shape") if isinstance(e, dict) else type(e).__name__
            warn.append(f"element {n_}: unknown shape “{what}” — left out (use {', '.join(_GFX_SHAPES)})")
            continue
        for k_ in ("location", "position"):                          # rokon kulcsok a helyre
            if e.get(k_) is not None and not e.get("place") and "lat" not in e:
                q = e[k_]
                if isinstance(q, str):
                    e = {**e, "place": q}
                elif isinstance(q, dict) and (q.get("place") or "lat" in q):
                    e = {**e, **{k: q[k] for k in ("place", "lat", "lon") if k in q}}
                elif isinstance(q, (list, tuple)) and len(q) == 2 and "x" not in e:
                    e = {**e, "x": q[0], "y": q[1]}
        el: dict[str, Any] = {"shape": e["shape"]}
        for k in ("x", "y", "w", "h", "r"):
            if num(e.get(k)) is not None:
                el[k] = max(-0.2, min(1.2, float(e[k])))
        for k in ("size", "width", "radius", "at", "dur", "opacity", "move_dur"):
            if num(e.get(k)) is not None:
                el[k] = float(e[k])
        if "opacity" in el:                      # 0–1 (1 = tömör), a mentett specben is így; a 0–100 százalékot is érti
            # (10-18: 0.55 → 5 % = láthatatlan; 10-19: a 0.94-ből a specben 94.0 lett) — a builder kapja százalékban
            el["opacity"] = round(max(0.0, min(1.0, el["opacity"] / 100 if el["opacity"] > 1 else el["opacity"])), 4)
        dash = {"dashed": [16, 11], "dotted": [0.5, 11]}.get(e["dash"]) if isinstance(e.get("dash"), str) else e.get("dash")
        if isinstance(dash, (list, tuple)) and len(dash) == 2 and all(num(v) is not None for v in dash):   # szaggatott vonal
            el["dash"] = [max(0.1, min(200.0, float(dash[0]))), max(1.0, min(200.0, float(dash[1])))]
        if isinstance(e.get("dim"), list):        # irányított figyelem: ezekben a szakaszokban [tól, ig] mp halványabb
            el["dim"] = [[float(w[0]), float(w[1])] for w in e["dim"][:8] if isinstance(w, (list, tuple)) and len(w) == 2
                         and all(num(v) is not None for v in w) and w[1] > w[0]]
        for k in ("from", "to", "move_to"):
            if pt(e.get(k)):
                el[k] = pt(e[k])
            elif geo(e.get(k)) or geo(e.get("geo_" + k)):
                el["geo_" + k] = geo(e.get(k)) or geo(e.get("geo_" + k))
        g_ = geo(e.get("geo")) or (geo(e) if e["shape"] != "map" else None)   # a tű / szöveg helye a jelenet térképén
        if g_:
            el["geo"] = g_
        if isinstance(e.get("points"), list) or isinstance(e.get("geo_points"), list):
            raw = e.get("points") if isinstance(e.get("points"), list) else e["geo_points"]
            gp = [geo(q) for q in raw[:200]]
            if any(gp):
                el["geo_points"] = [q for q in gp if q]
            else:
                el["points"] = [q for q in (pt(q) for q in raw[:200]) if q]
        if e["shape"] == "map":                          # valódi térkép: hely vagy bbox [s, w, n, e]; a tűk is rajta legyenek
            mp = e.get("map") if isinstance(e.get("map"), dict) else e
            el["map"] = {k: v for k, v in {
                "place": mp.get("place").strip()[:160] if isinstance(mp.get("place"), str) and mp["place"].strip() else None,
                "bbox": [float(v) for v in mp["bbox"]] if isinstance(mp.get("bbox"), list) and len(mp["bbox"]) == 4
                and all(isinstance(v, (int, float)) for v in mp["bbox"]) else None,
                "detail": mp.get("detail") if mp.get("detail") in ("streets", "outline") else None}.items() if v is not None}
        for k in ("fill", "stroke", "color", "align", "weight"):
            if isinstance(e.get(k), str):
                el[k] = e[k][:24]
        el["anim"] = e.get("anim") if e.get("anim") in _GFX_ANIMS else "appear"
        if e["shape"] in ("text", "pin"):
            txt = _clean(e.get("text") or e.get("label") or "", 160, "Graphic text")
            if [x for x in _text_numbers(txt) if x not in nums]:
                dropped.append(txt)
                continue
            el["text"] = txt
        if e["shape"] in ("image", "lottie"):
            src = e.get("image") or e.get("lottie") or e.get("file_id") or e.get("url") or e.get("src")   # src: mentett alak
            if not src:
                warn.append(f"element {n_} ({e['shape']}): no {e['shape']} given — left out")
                continue
            el["fit"] = "cover" if e.get("fit") == "cover" else "contain"
            el["loop"] = e.get("loop") is not False
            cl = e.get("clip")
            if isinstance(cl, (list, tuple)) and cl and num(cl[0]) is not None:     # videó: a rész [tól, ig?] mp-ben
                a_ = max(0.0, float(cl[0]))
                b_ = float(cl[1]) if len(cl) > 1 and num(cl[1]) is not None and float(cl[1]) > a_ else None
                el["clip"] = [a_, b_]
            if e.get("sound") in ("replace", "mix", "off"):
                el["sound"] = e["sound"]
            for k_ in ("sound_file", "clip_len"):
                if e.get(k_) is not None:
                    el[k_] = e[k_]
            el["src"] = src if isinstance(src, (str, dict)) else None
            if e.get("asset"):
                el["asset"] = str(e["asset"])
            if isinstance(e.get("credit"), str) and e["credit"].strip():
                el["credit"] = e["credit"].strip()[:200]
        if e["shape"] in ("path", "polygon") and len(el.get("points") or el.get("geo_points") or []) < 2:
            warn.append(f"element {n_} ({e['shape']}): fewer than 2 usable points — left out")
            continue
        if e["shape"] in ("line", "arrow") and not ((el.get("from") or el.get("geo_from")) and (el.get("to") or el.get("geo_to"))):
            warn.append(f"element {n_} ({e['shape']}): needs from and to ([x, y], {{\"place\"}} or {{\"lat\", \"lon\"}}) — left out")
            continue
        if e["shape"] == "map" and not (el.get("map") or {}).get("place") and not (el.get("map") or {}).get("bbox"):
            pass                                          # hely nélkül a tűk köré rajzol (geo.build) — ha nincs tű, ott jelez
        out.append(el)
    title = _clean(c.get("title"), 60, "Chart title")
    warns = [f"⚠ Graphic “{title or 'custom'}”: {w}." for w in warn[:8]]
    if dropped:
        warns.append(f"⚠ Graphic “{title or 'custom'}”: left out {len(dropped)} text(s) with a number not in the source, "
                     f"e.g. “{dropped[0][:60]}”.")
    if not out:
        return None, [f"⚠ A suggested graphic “{title}” had nothing to draw."] + warns if c.get("elements") else None
    g = {"type": "graphic", "title": title, "unit": "", "items": [], "elements": out}
    if num(c.get("backdrop")) is not None:
        g["backdrop"] = max(0.0, min(1.0, float(c["backdrop"])))
    note = f"Graphic “{title or 'custom'}”: {len(out)} elements" + (f" ({expanded} pins/routes taken from inside the map)" if expanded else "") + "."
    return g, [note] + warns


def _script_from_plan(c: dict[str, Any]) -> tuple[Optional[dict[str, Any]], Optional[str]]:
    """Az ügynök saját EffectCraft-szkriptje egy jelenetre (After Effects-szerű JS: comp.layers.addShape…). ZÁRKÁBAN fut
    (jail.py) — a motor saját kapuja mellett. A számait a build után a jelenet szövegrétegeiből ellenőrizzük."""
    code = c.get("code") if isinstance(c.get("code"), str) else ""
    title = _clean(c.get("title"), 60, "Chart title")
    if not code.strip():
        return None, f"A suggested script scene “{title}” had no code."
    if len(code) > 20000:
        return None, f"A suggested script scene “{title}” is too long ({len(code)} characters; 20 000 at most)."
    g = {"type": "script", "title": title, "unit": "", "items": [], "code": code}
    if isinstance(c.get("backdrop"), (int, float)):
        g["backdrop"] = max(0.0, min(1.0, float(c["backdrop"])))
    return g, f"Script scene “{title or 'custom'}”: {code.count(chr(10)) + 1} lines, run in the sandbox."


_CHART_TYPES = ("column", "bar", "line", "area", "donut", "ring", "icons", "timeline", "steps",
                "counter", "compare", "progress", "gauge", "funnel", "proscons", "quote", "table", "graphic", "script",
                "decision_tree", "swimlane", "network",
                "waterfall", "slope", "scatter", "band", "lines", "race", "sankey",
                "threejs")                                                     # a kliens three.js-jelenete: app/tools/threejs.py     # az utóbbi hét: app/tools/motion_charts.py


def _chart_from_plan(c: Any, source: str) -> tuple[Optional[dict[str, Any]], Optional[str]]:
    """A modell javasolta ábra → ellenőrzött ábra | None + megjegyzés. Szabály minden típusra: szám és dátum CSAK a
    forrásból (számos ábra: minden érték; ikonrács: az „N-ből" is; idővonal: a dátum szó szerint; lépések: a szövegük
    számai) — különben az ábra kimarad. A modell nem talál ki számot."""
    if not isinstance(c, dict):
        return None, None
    from app.tools import diagrams, motion_charts
    typ = str(c.get("type") or "column").lower()
    typ = diagrams.ALIASES.get(typ, motion_charts.ALIASES.get(typ, typ))
    if typ not in _CHART_TYPES:
        return None, None
    if typ in motion_charts.TYPES:     # vízesés, előtte→utána, pont, sáv, több vonal, verseny, sankey: saját alakjuk van
        return motion_charts.parse(typ, c, source)
    if typ in diagrams.TYPES:          # magyarázó ábra: a leírása marad a jelenetben, a rajzolás a graphic elemeivel
        nums = _text_numbers(source)
        return diagrams.from_plan({**c, "type": typ}, lambda t: [x for x in _text_numbers(t) if x not in nums], _clean)
    if typ == "graphic":
        return _graphic_from_plan(c, source)
    if typ == "script":
        return _script_from_plan(c)
    if typ == "threejs":
        from app.tools import threejs
        return threejs.from_plan(c, _clean)
    title = _clean(c.get("title"), 60, "Chart title")
    nums = _text_numbers(source)
    src = re.sub(r"\s+", " ", source or "").lower()
    bad = lambda t: [x for x in _text_numbers(t) if x not in nums]
    left_out = f"A suggested chart “{title}” was left out: not all its numbers are in the source."
    raw = [it for it in c.get("items") or [] if isinstance(it, dict)]
    if typ == "quote":                 # idézet: CSAK szó szerint a forrásból (kitalált idézet nincs)
        it = raw[0] if raw else {}
        q = _clean(it.get("text") or it.get("label"), 220, "Quote")
        norm = lambda x: re.sub(r"[^\w ]+", "", re.sub(r"\s+", " ", x or "")).lower().strip()
        if not q or norm(q.rstrip("…")) not in norm(source):
            return None, f"A suggested quote “{q[:60]}” was left out: it is not word for word in the source."
        who = _clean(it.get("label") if it.get("text") else "", 40, "Quote author")
        return {"type": typ, "title": title, "unit": "", "items": [{"text": q, "label": who}]}, \
            f"Quote “{q[:50]}…”: word for word in the source."
    if typ == "table":                 # táblázat: fejléc + legfeljebb 7 sor; a szám csak a forrásból (mint minden ábrán)
        cols = [_clean(h, 24, "Column") for h in (c.get("columns") or c.get("labels") or []) if isinstance(h, (str, int, float))][:5]
        rows = []
        for r_ in (c.get("rows") or [])[:7]:
            if not isinstance(r_, list):
                continue
            cells = [_clean(x, 32, "Cell") for x in r_[:max(len(cols), 2)]]
            if any(bad(x) for x in cells):
                return None, left_out
            if any(cells):
                rows.append(cells + [""] * (len(cols) - len(cells)))
        if len(cols) < 2 or not rows:
            return None, None
        return ({"type": typ, "title": title, "unit": "", "columns": cols, "rows": rows,
                 "items": [{"label": r_[0], "value": None, "text": " | ".join(r_[1:])} for r_ in rows]},
                f"Table “{title}”: {len(cols)} columns, {len(rows)} rows, no number that isn't in the source.")
    if typ == "proscons":
        items = []
        for it in raw[:8]:
            txt = _clean(it.get("text"), 70, "Point")
            if bad(txt):
                return None, left_out
            side = "+" if str(it.get("label") or "").strip().lower() in ("+", "pro", "for", "mellette", "előny", "yes") else "-"
            if txt:
                items.append({"label": side, "text": txt})
        if len(items) < 2:
            return None, None
        heads = [h for h in (c.get("labels") or []) if isinstance(h, str)][:2]
        return ({"type": typ, "title": title, "unit": "", "items": items, **({"labels": [_clean(h, 18) for h in heads]} if heads else {})},
                f"For/against “{title}”: {len(items)} points, no number that isn't in the source.")
    if typ in ("timeline", "steps"):
        items = []
        for it in raw[:6]:
            lab, txt = _clean(it.get("label"), 26 if typ == "timeline" else 30, "Chart item"), _clean(it.get("text"), 70, "Chart text")
            # a dátum szó szerint a forrásban — vagy legalább MINDEN száma ott van („15. mp” ← „15 és 18 másodperc”)
            if typ == "timeline" and re.sub(r"\s+", " ", lab).lower() not in src and (not _text_numbers(lab) or bad(lab)):
                return None, f"A suggested timeline “{title}” was left out: the date “{lab}” is not in the source."
            if bad(txt) or (typ == "steps" and bad(lab)):
                return None, left_out
            if lab:
                items.append({"label": lab, "text": txt})
        if len(items) < (2 if typ == "timeline" else 3):
            return None, None
        if typ == "timeline" and sum(bool(re.search(r"\d", it["label"])) for it in items) < 2:
            # „Kr. e. 100 körül” → „ettől kezdve”: a szövegben szerepel, de nem időpont (UI-próba 10-18) — egy dátumból
            # nem lesz idővonal
            return None, f"A suggested timeline “{title}” was left out: it has fewer than two real dates."
        return ({"type": typ, "title": title, "unit": "", "items": items},
                f"{'Timeline' if typ == 'timeline' else 'Steps'} “{title}”: {len(items)} items"
                + (", every date is in the source text." if typ == "timeline" else ", no number that isn't in the source."))
    items, missing = [], []
    for it in raw:
        try:
            v = float(str(it.get("value")).replace(",", ".").replace(" ", ""))
        except ValueError:
            continue
        (items if round(v, 6) in nums else missing).append({"label": _clean(it.get("label"), 22 if typ not in ("ring", "icons") else 60, "Chart label"), "value": v})
    unit = _clean(c.get("unit"), 14, "Chart unit")   # „millió" is beférjen (10-12: „1,1 milli…”)
    out: dict[str, Any] = {"type": typ, "title": title, "unit": unit}
    if typ in ("ring", "icons", "counter", "gauge"):
        if missing or not items:
            return None, left_out
        out["items"] = items[:1]
        if typ == "ring" and not (0 <= items[0]["value"] <= 100):
            return None, None
        if typ == "counter":
            return out, f"Counter “{title}”: the number is in the source text."
        if typ == "gauge":            # a skála vége is a forrásból (a 0–100 % alapértelmezés kivétel)
            for key in ("min", "max"):
                if c.get(key) not in (None, ""):
                    try:
                        val = float(str(c[key]).replace(",", "."))
                    except ValueError:
                        return None, None
                    if round(val, 6) not in nums and not (val in (0, 100) and (unit == "%" or key == "min")):
                        return None, left_out
                    out[key] = val
            lo_, hi_ = out.get("min", 0), out.get("max", 100)
            if not lo_ <= items[0]["value"] <= hi_:
                return None, None
            return out, f"Gauge “{title}”: the value (and its scale) are in the source text."
        if typ == "icons":
            tot = c.get("total")
            if tot not in (None, "", 0):
                try:
                    tot = float(str(tot).replace(",", "."))
                except ValueError:
                    return None, None
                if round(tot, 6) not in nums and not (unit == "%" and tot == 100):
                    return None, left_out
                out["total"] = tot
            elif unit != "%":
                return None, None
        return out, f"{'Ring' if typ == 'ring' else 'Icons'} “{title}”: the number is in the source text."
    checked, rest = len(items), ""
    # fánk: a „többiek" sora (100% − a többi) KISZÁMOLT érték, nem kitalált — a modell címkéjével marad
    # (10-12: „48% telefon, 22% olvas, 18% tévé, a többiek…" — a 12% nincs a szövegben, a fánk emiatt kiesett)
    if typ == "donut" and unit == "%" and len(missing) == 1 and abs(missing[0]["value"] - (100 - sum(i["value"] for i in items))) < 0.01:
        items.append(missing.pop())
        rest = f" “{items[-1]['label']}” is the rest of 100%."
    if missing or len(items) < 2:
        return None, left_out
    if typ == "compare":
        items = items[:3]
    if typ == "progress" and not all(0 <= i["value"] <= 100 for i in items):
        return None, None
    if typ == "funnel" and len(items) < 3:
        return None, None
    if typ == "area" and len(items) < 3:
        typ = out["type"] = "column"
    if typ == "donut":
        tot = sum(i["value"] for i in items)
        if any(i["value"] < 0 for i in items) or len(items) > 6:
            return None, None
        if unit == "%" and 60 <= tot < 99.5 and not rest:   # a hiányzó rész „Egyéb" — kiszámolt, nem kitalált
            items.append({"label": "Egyéb" if _lang_of(source) == "hu" else "Other", "value": round(100 - tot, 6)})
            rest = f" “{items[-1]['label']}” is the rest of 100%."
    if typ == "column" and (len(items) > 7 or max(len(i["label"]) for i in items) > 14):
        typ = out["type"] = "bar"
    out["items"] = items[:12]
    return out, f"Chart “{title}”: all {checked} values are in the source text.{rest}"


def charts_from_table(file_id: str, limit: int = 2) -> list[dict[str, Any]]:
    """Táblázatból (CSV/XLSX) grafikon-jelenetek: az értékek a cellákból (az infografika táblázat-kezelésével —
    a „mutató × ország/év" tábla megfordítva), oszloponként egy grafikon."""
    from app.tools import infographic as ig
    tb = ig.load_table(crafts.existing(file_id))
    if ig.indicator_table(tb):
        tb = ig.transpose(tb)
    cols = ig.numeric_columns(tb)
    lab = ig.label_column(tb, cols)
    cols = [c for c in cols if c.j != lab]          # az évszám-oszlop címke, nem adat (10-12: „Year” grafikon lett)
    labels = tb.col(lab) if lab is not None else [str(i + 1) for i in range(len(tb.rows))]
    timey = ig.time_like(labels)
    out = []
    for c in cols[:limit]:
        items = [{"label": _clean(labels[i], 22), "value": v} for i, v in enumerate(c.values) if v is not None][:12]
        if len(items) < 2:
            continue
        tot = sum(x["value"] for x in items)
        share = c.percent and 95 <= tot <= 105 and 2 <= len(items) <= 6 and all(x["value"] >= 0 for x in items)
        typ = "donut" if share and not timey else "line" if timey and len(items) >= 3 else (
            "bar" if len(items) > 7 or max(len(x["label"]) for x in items) > 14 else "column")
        unit = "%" if c.percent else ""
        out.append({"type": typ, "title": _clean(re.sub(r"\s*\(%\)|\s*%", "", c.name), 60), "unit": unit, "items": items,
                    "cells": f"{tb.refs[0][c.j]}:{tb.refs[-1][c.j]}"})
    return out


# egy gyártás / szerkesztés alatt: mely szövegek rövidültek (a check jelzi — 10-12: „millió” → „milli…” észrevétlenül)
_TRUNC: "contextvars.ContextVar[Optional[list[str]]]" = contextvars.ContextVar("video_trunc", default=None)


def _clean(s: Any, n: int, what: Optional[str] = None) -> str:
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    if len(s) <= n:
        return s
    out = s[:n - 1].rsplit(" ", 1)[0].rstrip(",;:—-") + "…"
    notes = _TRUNC.get()
    if what and notes is not None:
        notes.append(f"{what} “{s[:70]}{'…' if len(s) > 70 else ''}” was shortened to “{out}” (limit {n} characters)")
    return out


# ábraszövegek, amelyek a dobozukba csak kisebb betűvel fértek be (a check megemlíti)
_FIT: "contextvars.ContextVar[Optional[list[str]]]" = contextvars.ContextVar("video_fit", default=None)


def _wrap(st: str, size: float, box_w: float) -> tuple[int, bool, float]:
    """A szöveg tördelve a doboz szélességére (~0,56 em átlagos betűszélesség) → (sorok, van-e szó, ami magában
    sem fér ki, a legszélesebb sor px-ben)."""
    cw, lines, cur, wide, widest = 0.56 * size, 1, 0.0, False, 0.0
    for w in str(st or "").split() or [""]:
        ww = len(w) * cw
        wide = wide or ww > box_w
        nxt = ww if cur == 0 else cur + cw + ww
        if nxt > box_w and cur > 0:
            widest, lines, cur = max(widest, cur), lines + 1, ww
        else:
            cur = nxt
    return lines, wide, max(widest, cur)


def _fit_size(st: str, size: float, box_w: float, box_h: float, name: str = "") -> int:
    """Elfér-e a szöveg a dobozában (szótördeléssel, ~0,56 em átlagos betűszélesség, 1,22 sormagasság)? Ha nem,
    kisebb betű — legfeljebb az eredeti 55%-áig; az illesztést feljegyzi."""
    s0, size = int(size), float(size)
    while size > s0 * 0.55:
        lines, wide, _ = _wrap(st, size, box_w)
        if not wide and lines * size * 1.22 <= box_h * 1.05:
            break
        size *= 0.93
    out = int(round(size))
    notes = _FIT.get()
    if out < s0 and notes is not None:
        notes.append(f"{name or 'text'} {s0}→{out} px")
    return out


# a renderer betűje (Inter) méréshez: a rendszer fonts-inter csomagja — a motorba csomagolt Interrel mérve (10-19) a
# sorszélesség ±2%-on belül egyezik, az ő/ű is (a 0,56 em-es becslés a nagybetűs címet ~15%-kal alábecsülte)
_INTER = {"bold": "Inter-Bold", "semibold": "Inter-SemiBold", "regular": "Inter-Regular"}
_FONT_DIRS = ("/usr/share/fonts/opentype/inter", "/usr/share/fonts/truetype/inter", "/usr/local/share/fonts")


@functools.lru_cache(maxsize=64)
def _inter_font(weight: str, size: int) -> Any:
    from PIL import ImageFont
    for d in _FONT_DIRS:
        for ext in (".otf", ".ttf"):
            p = Path(d) / (_INTER.get(weight, "Inter-SemiBold") + ext)
            if p.exists():
                return ImageFont.truetype(str(p), size)
    return None


def _text_px(st: str, size: float, weight: str = "semibold") -> float:
    """A szöveg legszélesebb (\\n-nel tagolt) sora EGY sorban, px — a renderer betűjével; betű nélkül becslés."""
    lines = str(st or "").split("\n")
    f = _inter_font(weight if weight in _INTER else "semibold", max(1, int(round(size))))
    if f is None:
        return max(sum(0.68 if c.isupper() else 0.55 for c in x) * size for x in lines)
    return max(f.getlength(x) for x in lines) * size / max(1, int(round(size)))


def _auto_text_w(st: str, size: float, weight: str, x: float, align: str, W: float) -> float:
    """A saját grafika szövegdoboza, ha nincs w: 0,6 W (mint eddig) — ha a szöveg egy sorban ennél szélesebb, annyi,
    amennyi kell, legfeljebb a képkocka 0,92-e, és a képkockán belül (balra / jobbra igazítva a 0,6-os doboz széle a
    horgony). Ami így sem fér el, az tördel — a check megmondja (10-19: a rejtett 0,6 két sorba törte a címet)."""
    base = 0.6 * W
    need = _text_px(st, size, weight) * 1.03 + 0.01 * W
    if need <= base:
        return base
    lo, hi = 0.04 * W, 0.96 * W
    room = {"left": hi - (x - 0.3 * W), "right": (x + 0.3 * W) - lo}.get(align, 2 * min(x - lo, hi - x))
    return max(base, min(need, 0.92 * W, room))


def scrim_of(v: Any, what: str = "scrim") -> Optional[float]:
    """A képek fölötti sötétítés: 0–1 (0 = semmi, 1 = fekete); 0–100 százalékként is; None = a beépített."""
    if v in (None, "", "auto"):
        return None
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        raise ValueError(f"{what}: 0–1 (0 = no darkening over the pictures, 1 = black) — not {v!r}.")
    v = float(v) / 100 if float(v) > 1 else float(v)
    return round(max(0.0, min(1.0, v)), 3)


def _with_ids(scenes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Stabil jelenetazonosítók (s1, s2…): beszúrás/törlés után is ugyanarra a jelenetre mutatnak."""
    used = {str(sc.get("id")) for sc in scenes if sc.get("id")}
    k = max([int(m.group(1)) for u in used if (m := re.fullmatch(r"s(\d+)", u))] or [0])
    for sc in scenes:
        if not sc.get("id") or str(sc["id"]) in ("title", "end"):
            k += 1
            while f"s{k}" in used:
                k += 1
            sc["id"] = f"s{k}"
            used.add(sc["id"])
    return scenes


_ABBR = {"dr", "kr", "e", "i", "sz", "st", "szt", "mr", "mrs", "ms", "jr", "sr", "vs", "pl", "ld", "kb", "prof", "ifj",
         "id", "özv", "no", "nr", "ca", "approx", "etc", "gen", "u", "ú", "hr", "ny", "kft", "zrt", "bt"}


def _sentences(text: str) -> list[str]:
    """Mondatok: írásjel + szóköz a határ, de a kisbetűvel vagy számmal folytatódó rész az előzőhöz tartozik
    (a „Kr. e. 100”, „dr. Kiss”, „a 3. fejezet” nem vág ketté). Bekezdéshatáron nem olvaszt össze."""
    out: list[str] = []
    for para in re.split(r"\n+", text or ""):
        cur: list[str] = []
        for part in re.split(r"(?<=[.!?…])\s+", para.strip()):
            if not part:
                continue
            last = cur[-1].rsplit(None, 1)[-1] if cur else ""
            short_abbr = cur and (last.rstrip(".").lower() in _ABBR or re.fullmatch(r"[^\W\d_]\.", last))   # „Dr.”, „J.”
            if cur and (short_abbr or not (part[:1].isupper() or part[:1] in "\"„«(")):
                cur[-1] += " " + part
            else:
                cur.append(part)
        out += cur
    return [x for x in out if re.search(r"[^\W\d_]", x)]


def _restore_wording(text: str, scenes: list[dict[str, Any]], cap: int = 200) -> int:
    """keep_wording: a forrás kimaradt mondatai vissza, a sorrendjük szerinti helyre — a legutóbbi lefedett mondat
    jelenetének feliratához fűzve, vagy (ha nem fér) új jelenetként utána. → hány mondatot pótolt."""
    flat = lambda t: " ".join(re.findall(r"\w+", (t or "").lower()))
    def covered_in(snt: str) -> int:
        ws = {_STEM(w) for w in _words(snt)}
        for i, sc in enumerate(scenes):
            vt = f"{sc.get('headline', '')} {sc.get('caption', '')}"
            if (len(ws & {_STEM(w) for w in _words(vt)}) / len(ws) >= 0.6) if ws else flat(snt) in flat(vt):
                return i                       # rövid mondat („Ez jó.”): szó szerint
        return -1
    added, cur = 0, 0
    for snt in _sentences(text):
        at = covered_in(snt)
        if at >= 0:
            cur = at
            continue
        sc = scenes[cur]
        if len(sc.get("caption") or "") + 1 + len(snt) <= cap:
            sc["caption"] = f"{sc.get('caption') or ''} {snt}".strip()
        else:
            cur += 1
            scenes.insert(cur, {"headline": "", "caption": _clean(snt, cap, "Caption"),
                                "image_query": sc.get("image_query", ""), "stat": None, "chart": None})
        added += 1
    return added


def _paragraphs(text: str) -> list[str]:
    paras = [re.sub(r"\s+", " ", p).strip(" #") for p in re.split(r"\n\s*\n|\n(?=#)", text) if p.strip()]
    return [p for p in paras if len(p) > 3]


def _fallback_plan(text: str, n: int, title: str) -> dict[str, Any]:
    """Ha a modell nem elérhető: bekezdésenként egy jelenet (zsákutca helyett egyszerűbb, de kész videó)."""
    paras = _paragraphs(text) or [text]
    sentences = [s for p in paras for s in re.split(r"(?<=[.!?…])\s+", p) if len(s) > 3]
    chunk = max(1, math.ceil(len(sentences) / n))
    scenes = []
    for i in range(0, len(sentences), chunk):
        cap = " ".join(sentences[i:i + chunk])
        scenes.append({"headline": _clean(cap.split(".")[0], 55), "caption": _clean(cap, 170), "image_query": "", "stat": None})
    return {"title": title or _clean(paras[0], 60), "kicker": "", "scenes": scenes[:max(n, 1)], "end": {"title": title or "", "subtitle": ""}}


async def plan(text: str, *, n: int, keep_wording: bool, title_hint: str = "", language: Optional[str] = None) -> dict[str, Any]:
    from app.llm.siliconflow import _call_model, _strip_json_fences
    system = _PLAN.replace("{n}", str(n)) + (_KEEP if keep_wording else "")
    user = (f"Language: {language}\n" if language else "") + (f"Working title: {title_hint}\n" if title_hint else "") + "SOURCE:\n" + text[:24_000]
    last: Exception = RuntimeError("no model")
    for model in (settings.LLM_MODEL_PRIMARY, settings.LLM_MODEL_FALLBACK):
        try:
            raw = await _call_model(model, [{"role": "system", "content": system}, {"role": "user", "content": user}], thinking=False)
            data = json.loads(_strip_json_fences(raw))
            if isinstance(data, dict) and data.get("scenes"):
                return data
            raise ValueError("no scenes in the plan")
        except Exception as e:   # következő modell
            last = e
            log.warning("videó-forgatókönyv %s modellel nem sikerült: %s", model, e)
    log.warning("videó-forgatókönyv: egyszerű bontás (%s)", last)
    return _fallback_plan(text, n, title_hint)


# ---------- elrendezés és időzítés ----------

def _layout(W: int, H: int) -> dict[str, Any]:
    m = min(W, H)
    vertical, landscape = H > W * 1.2, W > H * 1.2
    L: dict[str, Any] = {"bar": [round(m * 0.11), max(6, round(m * 0.009))], "kickerSize": round(m * 0.032),
                         "capSize": round(m * (0.042 if not landscape else 0.038)), "creditSize": round(m * 0.022),
                         "statSize": round(m * 0.24), "statLabelSize": round(m * 0.052), "titleAlign": "left"}
    if vertical:
        L.update(titleW=W * 0.84, titleX=W / 2, kickerTop=H * 0.575, barY=H * 0.618, titleTop=H * 0.635,
                 headW=W * 0.84, headX=W / 2, headBarY=H * 0.655, headTop=H * 0.67, capW=W * 0.84, capY=H * 0.865)
    elif landscape:
        L.update(titleW=W * 0.62, titleX=W * 0.07 + W * 0.31, kickerTop=H * 0.50, barY=H * 0.565, titleTop=H * 0.59,
                 headW=W * 0.56, headX=W * 0.07 + W * 0.28, headBarY=H * 0.60, headTop=H * 0.62, capW=W * 0.74, capY=H * 0.895)
    else:
        L.update(titleW=W * 0.84, titleX=W / 2, kickerTop=H * 0.53, barY=H * 0.58, titleTop=H * 0.60,
                 headW=W * 0.84, headX=W / 2, headBarY=H * 0.585, headTop=H * 0.60, capW=W * 0.84, capY=H * 0.885)
    L.update(statY=H * 0.42, endY=H * 0.42, endBarY=H * 0.52, endSubY=H * 0.585, creditY=H * 0.935, capPad=round(m * 0.022))
    # a nagy szám alatti címke a szám MÉRETÉHEZ igazodik (az a rövidebb oldalból jön) — a magasság fix aránya (0,53 H)
    # fekvő 16:9-ben a számba csúsztatta a címkét (10-12, gazdasági videó); állóban ez ~ugyanott marad
    L["statLabelY"] = L["statY"] + L["statSize"] * 0.5 + L["statLabelSize"] * 1.2
    return L


def _size_for(text: str, base: float, ref: int, floor: float) -> int:
    return round(max(floor, base * min(1.0, math.sqrt(ref / max(len(text), 1)))))


def _chart_seconds(ch: Any) -> Optional[float]:
    """Az időzített ábra (rangsor-verseny, megadott tempójú kibontás) kényelmes hossza — None, ha bármilyen hosszba belefér."""
    from app.tools import diagrams, motion_charts
    return diagrams.seconds_needed(ch) or motion_charts.seconds_needed(ch or {}) if isinstance(ch, dict) else None


def _scene_seconds(s: dict[str, Any]) -> Optional[float]:
    """A jelenet kért hossza: a megadott, vagy — ha nincs — egy kifejezett tempójú (pace) ábra-kibontás igénye."""
    if isinstance(s.get("duration"), (int, float)) and s["duration"] > 0:
        return float(s["duration"])
    return _chart_seconds(s.get("chart"))


def _chunks(text: str, max_chars: int) -> list[str]:
    """Felirat-darabok: mondathatáron törünk (a rövid mondatok egy darabba kerülnek); a túl hosszú mondat
    tagmondathatáron (, ; : —), végső esetben szóhatáron. (10-12: a puszta karakterszámos bontás mondat
    közepén tört: „…keep millions from going” / „solar. Community…”.)"""
    def words(s: str) -> list[str]:
        out, cur = [], ""
        for w in s.split():
            if cur and len(cur) + 1 + len(w) > max_chars:
                out.append(cur); cur = w
            else:
                cur = f"{cur} {w}".strip()
        return out + ([cur] if cur else [])
    pieces: list[str] = []
    for sent in re.split(r"(?<=[.!?…])\s+", text.strip()):
        if len(sent) <= max_chars:
            pieces.append(sent)
            continue
        cur = ""
        for clause in re.split(r"(?<=[,;:—–])\s+", sent):
            if len(clause) > max_chars:
                if cur:
                    pieces.append(cur); cur = ""
                pieces += words(clause)
            elif cur and len(cur) + 1 + len(clause) > max_chars:
                pieces.append(cur); cur = clause
            else:
                cur = f"{cur} {clause}".strip()
        if cur:
            pieces.append(cur)
    out: list[str] = []
    for p in pieces:              # rövid mondatok összevonva, ha együtt is elférnek
        if out and len(out[-1]) + 1 + len(p) <= max_chars and (len(p) < max_chars * 0.45 or len(out[-1]) < max_chars * 0.45):
            out[-1] = f"{out[-1]} {p}"
        else:
            out.append(p)
    return [o for o in out if o]


def _shade_png(path: Path, W: int, H: int, top: float, full: float, opacity: float) -> None:
    """Alul sötétedő átmenet előre megrajzolva (fekete, alfa top·H-tól full·H-ig 0 → opacity, lágy görbével).
    A Linear Wipe-os változat képkockánként ~20 ms volt; a kész PNG csak összeillesztés."""
    from PIL import Image
    import numpy as np
    y = np.linspace(0, 1, H, dtype=np.float32)
    t = np.clip((y - top) / max(full - top, 1e-3), 0, 1)
    a = (t * t * (3 - 2 * t) * opacity * 255).astype(np.uint8)
    alpha = np.repeat(a[:, None], W, axis=1)
    img = Image.fromarray(np.dstack([np.zeros((H, W, 3), np.uint8), alpha]), "RGBA")
    img.save(path, "PNG", optimize=True)


def _blurred(src: Path, dst: Path, radius: float) -> None:
    from PIL import Image, ImageFilter
    with Image.open(src) as im:
        im.convert("RGB").filter(ImageFilter.GaussianBlur(radius)).save(dst, "JPEG", quality=85)


def _rgb(hex_color: Optional[str]) -> list[float]:
    h = (hex_color or DEFAULT_ACCENT).lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", h):
        h = DEFAULT_ACCENT.lstrip("#")
    return [round(int(h[i:i + 2], 16) / 255, 4) for i in (0, 2, 4)]


def _decimals(v: float) -> int:
    return 0 if float(v).is_integer() else 1 if round(v, 1) == round(v, 6) else 2


def _fmt_fixed(x: float, dec: int, unit: str, lang: str) -> str:
    """Számláló köztes értéke ugyanannyi tizedessel, mint a végső (ne ugráljon a szélessége)."""
    num = f"{abs(x):,.{dec}f}"
    num = num.replace(",", " ").replace(".", ",") if lang == "hu" else num
    t = ("-" if x < 0 else "") + num
    if unit == "%":
        return t + (" %" if lang == "hu" else "%")
    if not unit:
        return t
    return f"{unit}{t}" if unit in ("$", "€", "£", "¥") and lang == "en" else f"{t} {unit}"


def _fmt_val(v: float, unit: str, lang: str) -> str:
    """Értékcímke a forrás pontosságával (3,35 % marad 3,35 — az infografika fmt-je egy tizedesre kerekítene)."""
    from app.tools.infographic import fmt
    pct = unit == "%"
    t = fmt(v, lang, percent=pct)
    if round(v, 1) != round(v, 2):
        num = f"{abs(v):,.2f}"
        num = num.replace(",", " ").replace(".", ",") if lang == "hu" else num
        t = ("-" if v < 0 else "") + num + ((" %" if lang == "hu" else "%") if pct else "")
    if pct or not unit:
        return t
    return f"{unit}{t}" if unit in ("$", "€", "£", "¥") and lang == "en" else f"{t} {unit}"


def _palette(acc: list[float]) -> list[list[float]]:
    """Szeletszínek az akcentusból: akcentus, világos, sötétebb akcentus, szürke, halvány akcentus, középszürke."""
    mix = lambda a, b, t: [round(a[i] * (1 - t) + b[i] * t, 4) for i in range(3)]
    return [acc, [0.93, 0.94, 0.96], mix(acc, [0, 0, 0], 0.45), [0.55, 0.58, 0.64], mix(acc, [1, 1, 1], 0.5), [0.35, 0.37, 0.42]]


def _arc(out: list, name: str, cx: float, cy: float, r: float, th: float, a0: float, a1: float, color: Any,
         t0: float, dur: float, opacity: int = 100) -> None:
    """Körív apró, elforgatott, lekerekített szakaszokból, amelyek egymás után tűnnek fel (rajzolás-hatás).
    Szögek fokban, 0 = fent, az óramutató járásával."""
    span = a1 - a0
    k = max(3, math.ceil(abs(span) / 5))
    for i in range(k):
        a = math.radians(a0 + span * (i + 0.5) / k - 90)
        chord = 2 * r * math.sin(math.radians(abs(span) / k) / 2)
        out.append({"k": "rect", "name": f"{name} {i + 1}", "size": [round(chord * 1.18 + th * 0.2, 1), round(th, 1)],
                    "pos": [round(cx + r * math.cos(a), 1), round(cy + r * math.sin(a), 1)], "rot": round(math.degrees(a) + 90, 2),
                    "color": color, "opacity": opacity, "round": round(th * 0.25, 1), "grow": "fade",
                    "t0": round(t0 + dur * i / k, 3), "dur": 0.08})


def _gfx_color(v: Any, acc: list[float], default: Any = "accent") -> Any:
    if isinstance(v, (list, tuple)) and len(v) >= 3 and all(isinstance(c, (int, float)) for c in v[:3]):
        return [max(0.0, min(1.0, float(c))) for c in v[:3]]          # [r, g, b] 0–1 (a térkép saját színei)
    v = v or default
    named = {"accent": "accent", "white": [1, 1, 1], "muted": [0.62, 0.65, 0.72], "dark": [0.06, 0.06, 0.08],
             "black": [0, 0, 0], "red": [0.9, 0.2, 0.2], "green": [0.2, 0.7, 0.35], "blue": [0.2, 0.45, 0.9],
             "yellow": [0.95, 0.78, 0.2], "none": None}
    if isinstance(v, str) and v.lower() in named:
        return named[v.lower()]
    if isinstance(v, str) and re.fullmatch(r"#?[0-9a-fA-F]{6}", v):
        return _rgb(v if v.startswith("#") else "#" + v)
    return named.get(str(default), "accent")


def _pin_sides(els: list[tuple[int, dict[str, Any]]], W: int, H: int, m: float, k: float) -> dict[int, str]:
    """Tűfeliratok helye: jobbra | balra | fölötte | alatta — ahol a legkevésbé fed más feliratot, tűt, és kint sem lóg
    (10-12: álló országtérképen a „Budapest” a „Debrecen” tűbe futott)."""
    pins = [(i, e) for i, e in els if e.get("shape") == "pin" and e.get("text") and not e.get("_skip")]
    taken = []                                         # (x0, y0, x1, y1): a tűk és a már lerakott feliratok
    for _, e in pins:
        r = e.get("r", 0.03) * m
        x, y = e.get("x", 0.5) * W, e.get("y", 0.5) * H
        taken.append((x - r, y - r, x + r, y + r))
    def overlap(a: tuple, b: tuple) -> float:
        return max(0.0, min(a[2], b[2]) - max(a[0], b[0])) * max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    out: dict[int, str] = {}
    for i, e in pins:
        ts = e.get("size", 34) * k
        w, h = min(W * 0.3, 0.56 * ts * len(e["text"]) + 4), ts * 1.25
        r = e.get("r", 0.03) * m
        x, y = e.get("x", 0.5) * W, e.get("y", 0.5) * H
        own = (x - r, y - r, x + r, y + r)
        cands = {"right": (x + r + 6, y - h / 2, x + r + 6 + w, y + h / 2), "left": (x - r - 6 - w, y - h / 2, x - r - 6, y + h / 2),
                 "above": (x - w / 2, y - r - 8 - h, x + w / 2, y - r - 8), "below": (x - w / 2, y + r + 8, x + w / 2, y + r + 8 + h)}
        def cost(bx: tuple) -> float:
            out_ = max(0.0, -bx[0]) + max(0.0, bx[2] - W) + max(0.0, -bx[1]) + max(0.0, bx[3] - H)
            return sum(overlap(bx, t) for t in taken if t != own) + out_ * h * 4
        best = min(("right", "left", "above", "below"), key=lambda sd: (cost(cands[sd]), ("right", "left", "above", "below").index(sd)))
        out[i] = best
        taken.append(cands[best])
    return out


def _graphic_elements(ch: dict[str, Any], W: int, H: int, acc: list[float]) -> list[dict[str, Any]]:
    """A „graphic” elemei builder-rétegekké: arányos koordináták → pixel; méretek 1080-as rövid oldalra értve."""
    m = min(W, H)
    k = m / 1080
    px = lambda q: [round(q[0] * W, 1), round(q[1] * H, 1)]
    out: list[dict[str, Any]] = []
    els = list(enumerate(ch.get("elements") or []))
    sides = _pin_sides(els, W, H, m, k)
    for i, e in sorted(els, key=lambda ie: ie[1].get("shape") != "map"):     # a térkép legalul: ne takarja a tűket
        if e.get("_skip"):
            continue
        sh, name = e["shape"], f"{e['shape'].capitalize()} {i + 1}"
        t0, dur = _el_t0(e, i), max(0.1, e.get("dur", 0.6))
        anim = {"anim": e.get("anim", "appear"), "t0": round(t0, 3), "dur": round(dur, 3)}
        if e.get("move_to"):              # odasiklik: a megjelenés után, move_dur alatt (alap: dur) — a builder animate()
            anim["move_to"] = px(e["move_to"])
            if e.get("move_dur"):
                anim["move_dur"] = round(max(0.1, e["move_dur"]), 3)
        if e.get("dim"):                  # irányított figyelem: a rétegre (és a nyílhegyére) is
            anim["dim"] = [[round(a, 3), round(b, 3)] for a, b in e["dim"]]
        ov = float(e.get("opacity", 1.0))
        op = int(max(5, min(100, round(ov * 100 if ov <= 1 else ov))))     # 0–1 (a régi mentett spec: százalék)
        if sh in ("rect", "circle") and anim["anim"] == "draw" and e.get("stroke") and _gfx_color(e.get("fill"), acc, "none") is None:
            # körvonal, ami kirajzolódik: zárt útvonal (a builder a „draw”-t csak útvonalon ismeri — téglalapon az időpont
            # elveszett, a keret rögtön ott volt; 10-18)
            cx, cy = e.get("x", 0.5) * W, e.get("y", 0.5) * H
            if sh == "rect":
                hw, hh = e.get("w", 0.2) * W / 2, e.get("h", 0.1) * H / 2
                pts = [[cx - hw, cy - hh], [cx + hw, cy - hh], [cx + hw, cy + hh], [cx - hw, cy + hh], [cx - hw, cy - hh]]
            else:
                rr = e.get("r", 0.06) * m
                pts = [[cx + rr * math.sin(2 * math.pi * q / 48), cy - rr * math.cos(2 * math.pi * q / 48)] for q in range(49)]
            out.append({"k": "path", "name": name, "points": [[round(a, 1), round(b, 1)] for a, b in pts], "closed": False,
                        "fill": None, "stroke": _gfx_color(e["stroke"], acc), "stroke_w": round(e.get("width", 4) * k, 1),
                        "opacity": op, **anim})
            continue
        if sh in ("rect", "circle", "pin"):
            if sh == "rect":
                size = [e.get("w", 0.2) * W, e.get("h", 0.1) * H]
                rnd = e.get("radius", 0) * k
            else:
                d = 2 * e.get("r", 0.03 if sh == "pin" else 0.06) * m
                size, rnd = [d, d], d / 2
            el = {"k": "rect", "name": name, "size": [round(max(size[0], 2), 1), round(max(size[1], 2), 1)],
                  "pos": px([e.get("x", 0.5), e.get("y", 0.5)]), "color": _gfx_color(e.get("fill"), acc, "accent" if sh != "rect" or not e.get("stroke") else "none"),
                  "opacity": op, "round": round(rnd, 1), **anim}
            if e.get("stroke"):
                el["stroke"], el["stroke_w"] = _gfx_color(e["stroke"], acc), round(e.get("width", 4) * k, 1)
            if sh == "rect" and anim["anim"] == "grow":       # oszlop a talpáról felfelé, sáv balról — nem középről
                el["grow"] = "up" if size[1] >= size[0] else "right"
            if sh == "rect":
                el["gfx"] = i                                 # a szövegek kilógását ehhez méri a check (_text_report)
            out.append(el)
            if sh == "pin" and e.get("text"):
                ts = round(e.get("size", 34) * k)
                side = sides.get(i, "right")
                px_, py_, r_ = e.get("x", 0.5) * W, e.get("y", 0.5) * H, size[0] / 2
                cx_ = {"right": px_ + r_ + 6 + W * 0.15, "left": px_ - r_ - 6 - W * 0.15}.get(side, px_)
                cy_ = {"above": py_ - r_ - 8 - ts * 0.6, "below": py_ + r_ + 8 + ts * 0.6}.get(side, py_)
                out.append({"k": "text", "name": f"Pin label {i + 1}", "str": e["text"], "size": ts, "box": [round(W * 0.3, 1), round(ts * 2.8, 1)],
                            "pos": [round(cx_, 1), round(cy_ + ts * 0.65, 1)],
                            "align": {"right": "left", "left": "right"}.get(side, "center"), "font": "semibold", "color": _gfx_color(e.get("color"), acc, "white"),
                            "anim": "appear", "t0": round(t0 + 0.2, 3), "dur": 0.4})
        elif sh in ("line", "arrow", "path", "polygon"):
            pts = [e["from"], e["to"]] if sh in ("line", "arrow") else e["points"]
            el = {"k": "path", "name": name, "points": [px(q) for q in pts], "closed": sh == "polygon",
                  "fill": _gfx_color(e.get("fill"), acc, "accent" if sh == "polygon" else "none"),
                  "stroke": _gfx_color(e.get("stroke"), acc, "none" if sh == "polygon" and e.get("fill") else "accent"),
                  "stroke_w": round(e.get("width", 8) * k, 1), "opacity": op, **anim}
            if e.get("dash"):
                el["dash"] = [round(e["dash"][0] * k, 2), round(e["dash"][1] * k, 2)]
            if sh == "arrow":           # nyílhegy: kis háromszög a végpontnál, a vonal irányába fordítva
                (ax, ay), (bx, by) = px(pts[0]), px(pts[1])
                ang = math.atan2(by - ay, bx - ax)
                L = max(18, e.get("width", 8) * k * 3.2)
                tip = [bx, by]
                l = [bx - L * math.cos(ang) + L * 0.6 * math.sin(ang), by - L * math.sin(ang) - L * 0.6 * math.cos(ang)]
                r_ = [bx - L * math.cos(ang) - L * 0.6 * math.sin(ang), by - L * math.sin(ang) + L * 0.6 * math.cos(ang)]
                out.append(el)
                out.append({"k": "path", "name": f"Arrowhead {i + 1}", "points": [[round(v, 1) for v in tip], [round(v, 1) for v in l], [round(v, 1) for v in r_]],
                            "closed": True, "fill": el["stroke"] if el["stroke"] else "accent", "stroke": None, "stroke_w": 0, "opacity": op,
                            "anim": "appear", "t0": round(t0 + dur * 0.85, 3), "dur": 0.2, **({"dim": anim["dim"]} if anim.get("dim") else {})})
                continue
            out.append(el)
        elif sh == "text":
            ts = round(e.get("size", 44) * k)
            st = e.get("text") or ""
            font = {"bold": "bold", "regular": "regular"}.get(e.get("weight") or "", "semibold")
            align = e.get("align") if e.get("align") in ("left", "center", "right") else "center"
            x = e.get("x", 0.5) * W
            auto = e.get("w") is None               # w nélkül: egy sor, amíg a képkocka 0,92-ébe elfér (_auto_text_w)
            bw = _auto_text_w(st, ts, font, x, align, W) if auto else e["w"] * W
            if auto and align != "center":          # balra / jobbra igazítva a régi 0,6-os doboz széle marad a helyén
                x += (bw / 2 - 0.3 * W) * (1 if align == "left" else -1)
            fs = _fit_size(st, ts, bw, ts * 3.6, name)
            out.append({"k": "text", "name": name, "str": st, "size": fs, "box": [round(bw, 1), round(fs * 3.6, 1)],
                        "pos": [round(x, 1), round(e.get("y", 0.5) * H + fs * 1.15, 1)], "align": align, "font": font,
                        "color": _gfx_color(e.get("color"), acc, "white"), "gfx": i, **({"auto_w": True} if auto else {}), **anim})
        elif sh == "map":                # valódi térkép: rétegcsoportok (utak, folyók, országok), a forrás feltüntetve
            if not e.get("_groups"):
                continue
            t0 = max(0.0, e.get("at", 0.1))                  # a térkép a talaj: azonnal indul, nem a sorszáma szerint
            anim = {**anim, "t0": round(t0, 3)}
            pxs = lambda g: [[[round(x * W, 1), round(y * H, 1)] for x, y in path] for path in g["paths"]]
            how = e.get("anim") if e.get("anim") in ("draw", "appear", "none") else "draw"
            # két réteg: a területek (országok, parkok, víz) beúsznak; a vonalak (utak, folyók, határok) kirajzolódnak —
            # a kitöltött sokszögre tett Trim Paths szilánkos foltokat rajzol
            areas = [{"name": g["name"], "closed": True, "fill": _gfx_color(g["fill"], acc), "stroke": None, "stroke_w": 0,
                      "paths": pxs(g)} for g in e["_groups"] if g["closed"] and g.get("fill")]
            lines = [{"name": g["name"] + (" (borders)" if g["closed"] else ""), "closed": g["closed"], "fill": None,
                      "stroke": _gfx_color(g["stroke"], acc), "stroke_w": round(g["width"], 1), "paths": pxs(g)}
                     for g in e["_groups"] if g.get("stroke")]
            if areas:
                out.append({"k": "multipath", "name": "Map areas", "groups": areas, "opacity": op,
                            **{**anim, "anim": "appear" if how == "draw" else how, "dur": 0.8}})
            if lines:
                out.append({"k": "multipath", "name": "Map lines", "groups": lines, "opacity": op,
                            **{**anim, "anim": how, "dur": round(max(0.6, e.get("dur", 1.6)), 3)}})
            bx, by = (e.get("x", 0.5) + e.get("w", 1.0) / 2) * W, (e.get("y", 0.5) + e.get("h", 1.0) / 2) * H
            fs = round(max(14, 17 * k))
            out.append({"k": "text", "name": "Map credit", "str": "OpenFreeMap © OpenMapTiles © OpenStreetMap contributors" if "OpenMapTiles" in (e.get("_credit") or "")
                        else "© OpenStreetMap contributors" if "OpenStreetMap" in (e.get("_credit") or "") else "Natural Earth", "size": fs, "box": [round(W * 0.5, 1), round(fs * 1.6, 1)],
                        "pos": [round(min(bx, W) - W * 0.25 - 12 * k, 1), round(min(by, H) - fs * 0.9 - 8 * k, 1)], "align": "right",
                        "font": "regular", "color": [0.62, 0.65, 0.72], "anim": "appear", "t0": round(t0, 3), "dur": 0.4})
        elif sh in ("image", "lottie"):
            if not e.get("file"):
                continue
            out.append({"k": sh, "name": name, "file": e["file"], "comp": e.get("comp"), "fit": e.get("fit", "contain"),
                        "loop": e.get("loop", True), "pos": px([e.get("x", 0.5), e.get("y", 0.5)]), "gfx": i,
                        **({"credit": e["credit"]} if sh == "lottie" and e.get("credit") else {}),
                        **({"pick": e["_pick"]["pick"]} if sh == "lottie" and e.get("_pick") else {}),
                        **({"movie": True} if Path(e["file"]).suffix.lower().lstrip(".") in MOVIE_EXTS else {}),
                        "size": [round(e.get("w", 0.3) * W, 1), round(e.get("h", 0.3) * H, 1)], "opacity": op, **anim})
    return out


def chart_elements(ch: dict[str, Any], W: int, H: int, lang: str, accent: Optional[list[float]] = None,
                   dur: Optional[float] = None) -> list[dict[str, Any]]:
    """Egy mozgó grafikon rétegei (a builder csak lerakja): oszlop / sáv / vonal, értékcímkék, kategóriacímkék,
    alapvonal. Idő: az elemek egymás után nőnek ki (t0, dur). A felirat-sáv (alul) és a cím (fent) helye szabad."""
    if ch["type"] == "script":           # a jelenetet az ügynök szkriptje rajzolja (builder: runScript)
        return []
    if ch["type"] == "threejs":          # 3D: a klip a render után kerül a helyére (_render_3d), addig jelölő elem
        return [{"k": "threejs", "name": "3D scene", "code": ch["code"], "transparent": ch.get("transparent", True),
                 "size": [W, H], "pos": [round(W / 2, 1), round(H / 2, 1)]}]
    from app.tools import diagrams, motion_charts
    if ch["type"] in diagrams.TYPES:     # magyarázó ábra → a saját grafika elemei, a képkocka arányára rendezve
        return _graphic_elements({"elements": diagrams.to_graphic(ch, W, H, accent, dur, lang)}, W, H, accent or _rgb(None))
    if ch["type"] in motion_charts.TYPES:
        return motion_charts.elements(ch, W, H, lang, accent or _rgb(None), dur)
    m, vertical = min(W, H), H > W * 1.2
    items = ch["items"]
    n, vals = len(items), [float(i["value"]) if i.get("value") is not None else 0.0 for i in items]
    unit, typ = ch.get("unit") or "", ch["type"]
    x0, x1 = W * 0.09, W * 0.91
    yT, yB = H * (0.2 if vertical else 0.22), H * (0.66 if vertical else 0.7)
    vl, cl = round(m * 0.034), round(m * (0.031 if not vertical else 0.033))
    soft = [0.86, 0.88, 0.92]
    out: list[dict[str, Any]] = []
    rect = lambda name, size, pos, **o: out.append({"k": "rect", "name": name, "size": [round(size[0], 1), round(max(size[1], 2), 1)],
                                                     "pos": [round(pos[0], 1), round(pos[1], 1)], **o})
    text = lambda name, st, size, box, pos, **o: out.append({"k": "text", "name": name, "str": st,
                                                             "size": _fit_size(st, size, box[0], box[1], name),
                                                             "box": [round(box[0], 1), round(box[1], 1)], "pos": [round(pos[0], 1), round(pos[1], 1)], **o})
    stagger, grow = min(0.18, 1.6 / max(n, 1)), 0.7
    if typ == "column":
        vmin, vmax = min(0.0, *vals), max(0.0, *vals)
        top = yT + vl * 1.8
        bottom = yB - (vl * 1.8 if vmin < 0 else 0)
        sc = (bottom - top) / ((vmax - vmin) or 1)
        y0 = top + vmax * sc
        slot = (x1 - x0) / n
        bw = min(slot * 0.6, m * 0.17)
        cat_y = yB + cl * 1.4
        rect("Baseline", (x1 - x0, max(2, m * 0.003)), ((x0 + x1) / 2, y0), color=[1, 1, 1], opacity=45, t0=0.1, dur=0.4, grow="right")
        for k, (it, v) in enumerate(zip(items, vals)):
            cx, h, t0 = x0 + slot * (k + 0.5), abs(v) * sc, 0.35 + k * stagger
            rect(f"Bar – {it['label']}", (bw, h), (cx, y0 - h / 2 if v >= 0 else y0 + h / 2), color="accent", opacity=100,
                 round=min(bw * 0.08, 8), grow="up" if v >= 0 else "down", t0=t0, dur=grow)
            text(f"Value – {it['label']}", _fmt_val(v, unit, lang), vl, (slot, vl * 1.4),
                 (cx, (y0 - h - vl * 0.75) if v >= 0 else (y0 + h + vl * 0.95)), align="center", font="bold", t0=t0 + grow * 0.8)
            text(f"Label – {it['label']}", it["label"], cl, (slot * 0.96, cl * 2.7), (cx, cat_y + cl * 0.6), align="center",
                 font="semibold", color=soft, t0=0.2)
    elif typ == "bar":
        vmin, vmax = min(0.0, *vals), max(0.0, *vals)
        lab_w = (x1 - x0) * (0.34 if not vertical else 0.36)
        px0, px1 = x0 + lab_w + m * 0.02, x1 - vl * 4.2
        sc = (px1 - px0) / ((vmax - vmin) or 1)
        xz = px0 - vmin * sc
        row = (yB - yT) / n
        bh = min(row * 0.62, m * 0.075)
        rect("Baseline", (max(2, m * 0.003), yB - yT), (xz, (yT + yB) / 2), color=[1, 1, 1], opacity=45, t0=0.1, dur=0.4, grow="down")
        for k, (it, v) in enumerate(zip(items, vals)):
            cy, w, t0 = yT + row * (k + 0.5), abs(v) * sc, 0.35 + k * stagger
            rect(f"Bar – {it['label']}", (w, bh), (xz + w / 2 if v >= 0 else xz - w / 2, cy), color="accent", opacity=100,
                 round=min(bh * 0.12, 8), grow="right" if v >= 0 else "left", t0=t0, dur=grow)
            vx = xz + w + vl * 2.1 if v >= 0 else xz - w - vl * 2.1
            text(f"Value – {it['label']}", _fmt_val(v, unit, lang), vl, (vl * 4, vl * 1.4), (vx, cy + vl * 0.1),
                 align="center", font="bold", t0=t0 + grow * 0.8)
            text(f"Label – {it['label']}", it["label"], cl, (lab_w, cl * 1.4), (x0 + lab_w / 2, cy + cl * 0.1), align="right",
                 font="semibold", color=soft, t0=0.2)
    elif typ in ("donut", "ring"):
        pal = _palette(accent or _rgb(None))
        r = m * (0.2 if not vertical else 0.24)
        th = r * (0.32 if typ == "donut" else 0.22)
        cx = W * (0.32 if not vertical and typ == "donut" else 0.5)
        cy = (yT + yB) / 2 if not vertical or typ == "ring" else yT + r + th
        _arc(out, "Track", cx, cy, r, th, 0, 360, [1, 1, 1], 0.1, 0.3, opacity=14)
        if typ == "ring":
            v = max(0.0, min(100.0, vals[0]))
            _arc(out, "Ring", cx, cy, r, th, 0, 360 * v / 100, "accent", 0.4, 1.6)
            big = round(r * 0.62)
            text("Ring value", _fmt_val(vals[0], unit or "%", lang), big, (r * 1.9, big * 1.3), (cx, cy + big * 0.05),
                 align="center", font="bold", t0=0.5)
            text("Ring label", items[0]["label"], round(m * 0.036), (W * 0.8, m * 0.12), (cx, cy + r + th + m * 0.07),
                 align="center", font="semibold", color=soft, t0=0.9)
        else:
            tot = sum(max(v, 0) for v in vals) or 1
            a, leg_x = 0.0, W * 0.58
            leg_y0 = cy - (n - 1) * m * 0.045
            for k, (it, v) in enumerate(zip(items, vals)):
                sweep = 360 * max(v, 0) / tot
                col = pal[k % len(pal)]
                _arc(out, f"Slice {k + 1} – {it['label']}", cx, cy, r, th, a, a + sweep, col, 0.35 + 1.4 * a / 360, 1.4 * sweep / 360)
                a += sweep
                ly = (leg_y0 + k * m * 0.09) if not vertical else (cy + r + th + m * (0.06 + 0.075 * k))
                lx = leg_x if not vertical else W * 0.16
                t1 = 0.35 + 1.4 * (a - sweep / 2) / 360
                rect(f"Key {k + 1}", (m * 0.03, m * 0.03), (lx, ly), color=col, opacity=100, round=m * 0.006, grow="pop", t0=t1, dur=0.25)
                text(f"Key label {k + 1}", f"{it['label']}  ·  {_fmt_val(v, unit, lang)}", round(m * 0.034),
                     (W * (0.36 if not vertical else 0.72), m * 0.06), (lx + m * 0.03 + W * (0.18 if not vertical else 0.36), ly + m * 0.004),
                     align="left", font="semibold", t0=t1)
    elif typ == "icons":
        part, total = vals[0], float(ch.get("total") or (100 if unit == "%" else 0))
        if total <= 0 or part < 0 or part > total:
            return out
        cells = int(total) if total <= 20 and float(total).is_integer() else 100
        on = round(part / total * cells)
        cols = 10 if cells == 100 else (cells if cells <= 10 else math.ceil(cells / 2))
        rows = math.ceil(cells / cols)
        area_w = (x1 - x0) * (0.62 if not vertical else 0.9)
        cell = min(area_w / cols, (yB - yT - m * 0.16) / rows)
        dot = cell * 0.72
        gx0 = W / 2 - cell * cols / 2 + cell / 2
        gy0 = yT + m * 0.02 + cell / 2
        for k in range(cells):
            gx, gy = gx0 + (k % cols) * cell, gy0 + (k // cols) * cell
            hot = k < on
            rect(f"Dot {k + 1}", (dot, dot), (gx, gy), color="accent" if hot else [1, 1, 1], opacity=100 if hot else 22,
                 round=dot / 2, grow="pop", t0=round(0.3 + (1.6 * k / cells if hot else 0.05 * (k % cols)), 3), dur=0.22)
        say = (f"{_fmt_val(part, '', lang)} / {_fmt_val(total, '', lang)}" if cells != 100 or unit != "%" else _fmt_val(part, "%", lang))
        ty = gy0 + rows * cell + m * 0.03
        text("Icons value", say, round(m * 0.07), (W * 0.8, m * 0.1), (W / 2, ty + m * 0.03), align="center", font="bold",
             color=accent or _rgb(None), t0=1.6)
        text("Icons label", items[0]["label"], round(m * 0.034), (W * 0.84, m * 0.11), (W / 2, ty + m * 0.12), align="center",
             font="semibold", color=soft, t0=1.9)
    elif typ in ("timeline", "steps"):
        k_n = len(items)
        hor = not vertical and k_n <= 6
        th = max(4, m * 0.006)
        d = m * (0.03 if typ == "timeline" else 0.07)
        if hor:
            ly, xa, xb = (yT + yB) / 2 - m * 0.03, x0 + m * 0.03, x1 - m * 0.03
            if typ == "timeline":
                out.append({"k": "rect", "name": "Timeline", "size": [round(xb - xa, 1), round(th, 1)], "pos": [round(xa, 1), round(ly, 1)],
                            "rot": 0, "color": [1, 1, 1], "opacity": 55, "grow": "segment", "t0": 0.2, "dur": 0.9})
            slot = (xb - xa) / k_n
            for k, it in enumerate(items):
                px, t0 = xa + slot * (k + 0.5), 0.5 + k * min(0.55, 2.6 / k_n)
                rect(f"Point {k + 1}", (d, d), (px, ly), color="accent", opacity=100, round=d / 2, grow="pop", t0=t0, dur=0.3)
                if typ == "steps":
                    text(f"Step number {k + 1}", str(k + 1), round(d * 0.55), (d, d * 0.8), (px, ly + d * 0.04), align="center",
                         font="bold", t0=t0 + 0.1)
                up = typ == "timeline" and k % 2 == 1
                head_y = ly - m * 0.09 if up else ly + d / 2 + m * 0.05
                text(f"Label {k + 1}", it["label"], round(m * 0.042), (slot * 0.94, m * 0.06), (px, head_y), align="center",
                     font="bold", color=(accent or _rgb(None)) if typ == "timeline" else None, t0=t0 + 0.15)
                if it.get("text"):
                    body_y = head_y - m * 0.07 if up else head_y + m * 0.075
                    text(f"Text {k + 1}", it["text"], round(m * 0.03), (slot * 0.94, m * 0.11), (px, body_y + (0 if up else m * 0.02)),
                         align="center", font="regular", color=soft, t0=t0 + 0.3)
        else:
            lx, ya, yb = x0 + m * 0.06, yT + m * 0.02, yB + m * 0.02
            if typ == "timeline":
                out.append({"k": "rect", "name": "Timeline", "size": [round(yb - ya, 1), round(th, 1)], "pos": [round(lx, 1), round(ya, 1)],
                            "rot": 90, "color": [1, 1, 1], "opacity": 55, "grow": "segment", "t0": 0.2, "dur": 0.9})
            row = (yb - ya) / k_n
            for k, it in enumerate(items):
                py, t0 = ya + row * (k + 0.5), 0.5 + k * min(0.55, 2.6 / k_n)
                rect(f"Point {k + 1}", (d, d), (lx, py), color="accent", opacity=100, round=d / 2, grow="pop", t0=t0, dur=0.3)
                if typ == "steps":
                    text(f"Step number {k + 1}", str(k + 1), round(d * 0.55), (d, d * 0.8), (lx, py + d * 0.04), align="center",
                         font="bold", t0=t0 + 0.1)
                tx = lx + d / 2 + m * 0.04
                tw = x1 - tx
                text(f"Label {k + 1}", it["label"], round(m * 0.044), (tw, m * 0.065), (tx + tw / 2, py - (m * 0.02 if it.get("text") else 0)),
                     align="left", font="bold", color=(accent or _rgb(None)) if typ == "timeline" else None, t0=t0 + 0.15)
                if it.get("text"):
                    text(f"Text {k + 1}", it["text"], round(m * 0.032), (tw, m * 0.1), (tx + tw / 2, py + m * 0.055), align="left",
                         font="regular", color=soft, t0=t0 + 0.3)
    elif typ == "graphic":            # saját grafika: az elemek a képkocka arányában (0–1) → pixel
        out += _graphic_elements(ch, W, H, accent or _rgb(None))
    elif typ == "counter":            # felpörgő szám: 0 → érték, egymást váltó rétegekben, alatta a címke
        v, dec = vals[0], _decimals(vals[0])
        final = _fmt_val(v, unit, lang)
        big = round(min(m * 0.21, W * 0.9 / (0.6 * max(len(final), 1))))
        cy = (yT + yB) / 2 - m * 0.04
        steps, T = 26, 1.7
        for k in range(1, steps + 1):
            x = v * (1 - (1 - k / steps) ** 3)
            st = final if k == steps else _fmt_fixed(x, dec, unit, lang)
            el = {"k": "text", "name": f"Counter {k}", "str": st, "size": big, "box": [round(W * 0.94, 1), round(big * 1.35, 1)],
                  "pos": [round(W / 2, 1), round(cy, 1)], "align": "center", "font": "bold", "color": accent or _rgb(None),
                  "t0": round(0.3 + T * (k - 1) / steps, 3), "cut": True}
            if k < steps:
                el["t1"] = round(0.3 + T * k / steps, 3)
            out.append(el)
        text("Counter label", items[0]["label"], round(m * 0.042), (W * 0.84, m * 0.12), (W / 2, cy + big * 0.85),
             align="center", font="semibold", color=soft, t0=0.5)
    elif typ == "compare":            # A vs B (vs C): nagy számok, alattuk arányos oszlopok
        k_n = len(items)
        xs = [W * x for x in ((0.3, 0.7) if k_n == 2 else (0.2, 0.5, 0.8))]
        mx = max(max(vals), 1e-9)
        bw = min(m * 0.17, W * 0.22)
        big = round(min(m * 0.085, W * 0.3 / (0.6 * max(len(_fmt_val(v, unit, lang)) for v in vals))))
        top = yT + m * 0.02 - big * 0.2            # a legmagasabb oszlop fölött a szám — ennyi a hely
        rect("Baseline", (x1 - x0, max(2, m * 0.003)), ((x0 + x1) / 2, yB), color=[1, 1, 1], opacity=45, t0=0.1, dur=0.4, grow="right")
        for k, (it, v, x) in enumerate(zip(items, vals, xs)):
            h, t0 = max(2.0, max(v, 0) / mx * (yB - top - big * 1.4)), 0.35 + k * 0.35
            hot = v == mx
            rect(f"Bar – {it['label']}", (bw, h), (x, yB - h / 2), color="accent" if hot else [0.86, 0.88, 0.92],
                 opacity=100 if hot else 70, round=min(bw * 0.06, 8), grow="up", t0=t0, dur=0.8)
            text(f"Value – {it['label']}", _fmt_val(v, unit, lang), big, (W * 0.3, big * 1.3), (x, yB - h - big * 0.75),
                 align="center", font="bold", color=(accent or _rgb(None)) if hot else None, t0=t0 + 0.6)
            text(f"Label – {it['label']}", it["label"], round(m * 0.034), (W * 0.3, m * 0.1), (x, yB + m * 0.055),
                 align="center", font="semibold", color=soft, t0=0.2)
        if k_n == 2:
            text("Versus", "vs", round(m * 0.05), (m * 0.2, m * 0.08), (W / 2, (top + yB) / 2), align="center", font="bold",
                 color=soft, t0=0.3)
    elif typ == "progress":           # egymástól független százalékok: sín + kitöltés
        row = (yB - yT) / n
        bh = min(row * 0.22, m * 0.032)
        for k, (it, v) in enumerate(zip(items, vals)):
            cy, t0 = yT + row * (k + 0.55), 0.35 + k * min(0.3, 1.5 / n)
            text(f"Label – {it['label']}", it["label"], round(m * 0.036), ((x1 - x0) * 0.75, m * 0.06),
                 (x0 + (x1 - x0) * 0.375, cy - bh - m * 0.028), align="left", font="semibold", t0=0.2)
            text(f"Value – {it['label']}", _fmt_val(v, unit or "%", lang), round(m * 0.04), ((x1 - x0) * 0.25, m * 0.06),
                 (x1 - (x1 - x0) * 0.125, cy - bh - m * 0.028), align="right", font="bold", color=accent or _rgb(None), t0=t0 + 0.5)
            rect(f"Track – {it['label']}", (x1 - x0, bh), ((x0 + x1) / 2, cy), color=[1, 1, 1], opacity=16, round=bh / 2, t0=0.1, dur=0.3, grow="fade")
            w = max(bh, (x1 - x0) * max(0.0, min(100.0, v)) / 100)
            rect(f"Fill – {it['label']}", (w, bh), (x0 + w / 2, cy), color="accent", opacity=100, round=bh / 2, grow="right", t0=t0, dur=0.9)
    elif typ == "gauge":              # mérőóra: félkör-skála, kitöltés és mutató a helyére fordul
        lo_, hi_ = float(ch.get("min") or 0), float(ch.get("max") or ch.get("total") or 100)
        frac = max(0.0, min(1.0, (vals[0] - lo_) / ((hi_ - lo_) or 1)))
        r = min((x1 - x0) * 0.32, (yB - yT) * 0.78)
        th = r * 0.16
        cx, cy = W / 2, yB - m * 0.05
        _arc(out, "Gauge track", cx, cy, r, th, -90, 90, [1, 1, 1], 0.1, 0.3, opacity=16)
        _arc(out, "Gauge fill", cx, cy, r, th, -90, -90 + 180 * frac, "accent", 0.4, 1.3)
        out.append({"k": "rect", "name": "Needle", "size": [round(r * 0.86, 1), round(max(4, th * 0.28), 1)], "pos": [round(cx, 1), round(cy, 1)],
                    "rot0": -180, "rot": round(-180 + 180 * frac, 2), "color": [0.95, 0.95, 0.97], "opacity": 100,
                    "round": round(th * 0.14, 1), "grow": "rotate", "t0": 0.4, "dur": 1.3})
        rect("Hub", (th * 0.9, th * 0.9), (cx, cy), color=[0.95, 0.95, 0.97], opacity=100, round=th * 0.45, grow="pop", t0=0.3, dur=0.25)
        big = round(m * 0.09)
        text("Gauge value", _fmt_val(vals[0], unit, lang), big, (r * 2, big * 1.3), (cx, cy - r * 0.42), align="center", font="bold",
             color=accent or _rgb(None), t0=1.2)
        text("Gauge label", items[0]["label"], round(m * 0.036), (W * 0.8, m * 0.1), (cx, cy + m * 0.07), align="center",
             font="semibold", color=soft, t0=1.4)
        for nm, val, px in (("Gauge min", lo_, cx - r), ("Gauge max", hi_, cx + r)):
            text(nm, _fmt_val(val, unit, lang), round(m * 0.03), (m * 0.25, m * 0.06), (px, cy + th * 0.9 + m * 0.02), align="center",
                 font="regular", color=soft, t0=0.3)
    elif typ == "funnel":             # tölcsér: szűkülő, középről nyíló sávok
        row = (yB - yT) / n
        mx = max(max(vals), 1e-9)
        bh = row * 0.74
        for k, (it, v) in enumerate(zip(items, vals)):
            w, cy, t0 = max(bh, (x1 - x0) * 0.7 * max(v, 0) / mx), yT + row * (k + 0.5), 0.35 + k * 0.3
            col = _palette(accent or _rgb(None))[0] if k == 0 else [round(c * (1 - 0.13 * k) + 0.05 * k, 4) for c in (accent or _rgb(None))]
            rect(f"Stage – {it['label']}", (w, bh), (W / 2, cy), color=col, opacity=100, round=min(bh * 0.12, 10), grow="center", t0=t0, dur=0.7)
            text(f"Label – {it['label']}", it["label"], round(min(m * 0.036, bh * 0.42)), (w * 0.92, bh * 0.9), (W / 2, cy + m * 0.004),
                 align="center", font="bold", color=[0.06, 0.06, 0.08], t0=t0 + 0.45)
            text(f"Value – {it['label']}", _fmt_val(v, unit, lang), round(m * 0.036), ((x1 - x0) * 0.14, bh), (W / 2 + w / 2 + (x1 - x0) * 0.08, cy),
                 align="center", font="bold", color=accent or _rgb(None), t0=t0 + 0.55)
    elif typ == "proscons":           # mellette / ellene: két oszlop (állóban egymás alatt), tételenként megjelenve
        heads = ch.get("labels") or (["Mellette", "Ellene"] if lang == "hu" else ["For", "Against"])
        pros = [it for it in items if str(it.get("label")).strip() in ("+", "pro", "for", "mellette", "előny")]
        cons = [it for it in items if it not in pros]
        cols_ = [(pros, "+", "accent", heads[0]), (cons, "–", [0.6, 0.62, 0.68], heads[1] if len(heads) > 1 else "")]
        ts = 0.35
        # közös sormagasság (a két fejléc egy vonalban); a szövegdoboz a tetejétől ír — a dobozt úgy tesszük le, hogy az
        # első sor a jel közepére essen (10-12: a pontok szövege a jel fölé csúszott)
        rows_ = max(len(pros), len(cons)) + 1
        row = (yB - yT) / (rows_ * (2 if vertical else 1))
        fs, hs = round(m * 0.034), round(m * 0.045)
        line = lambda name, st, size, w, xl, yc, **o: text(name, st, size, (w, size * 2.8), (xl + w / 2, yc + size * 0.65), **o)
        for ci, (lst, mark, col, head) in enumerate(cols_):
            if vertical:
                hx, hy, colw = x0 + m * 0.02, yT + ci * (yB - yT) / 2, x1 - x0
            else:
                hx, hy, colw = x0 + ci * (x1 - x0) / 2 + m * 0.02, yT, (x1 - x0) / 2 - m * 0.04
            line(f"Heading {ci + 1}", head, hs, colw, hx, hy + row * 0.5, align="left", font="bold",
                 color=(accent or _rgb(None)) if ci == 0 else soft, t0=0.2)
            for k, it in enumerate(lst):
                cy, d = hy + row * (k + 1.5), m * 0.045
                rect(f"Mark {ci + 1}.{k + 1}", (d, d), (hx + d / 2, cy), color=col, opacity=100, round=d / 2, grow="pop", t0=ts, dur=0.25)
                text(f"Mark sign {ci + 1}.{k + 1}", mark, round(d * 0.8), (d, d * 1.1), (hx + d / 2, cy + d * 0.08), align="center",
                     font="bold", color=[0.06, 0.06, 0.08], t0=ts + 0.05)
                line(f"Point {ci + 1}.{k + 1}", it.get("text") or "", fs, colw - d * 1.6, hx + d * 1.4, cy, align="left",
                     font="semibold", t0=ts + 0.1)
                ts += 0.4
    elif typ == "table":              # táblázat: kiemelt fejléc, a sorok egymás után úsznak be (zebracsíkkal)
        cols, rows_ = ch.get("columns") or [], ch.get("rows") or []
        nc, nr = max(len(cols), 1), len(rows_)
        rh = min((yB - yT) / (nr + 1), m * 0.11)
        top = yT + ((yB - yT) - rh * (nr + 1)) / 2
        cw = (x1 - x0) / nc
        fs = round(min(rh * 0.42, m * (0.036 if nc <= 3 else 0.03)))
        rect("Header band", (x1 - x0, rh), ((x0 + x1) / 2, top + rh / 2), color="accent", opacity=100, round=min(rh * 0.18, 10),
             grow="right", t0=0.15, dur=0.4)
        for j, h in enumerate(cols):
            text(f"Head {j + 1}", h, fs, (cw * 0.92, rh * 0.9), (x0 + cw * (j + 0.5), top + rh / 2 + fs * 0.35), align="center",
                 font="bold", color=[0.06, 0.06, 0.08], t0=0.35)
        for k, r_ in enumerate(rows_):
            cy, t0 = top + rh * (k + 1.5), 0.6 + k * min(0.35, 2.4 / max(nr, 1))
            if k % 2 == 0:
                rect(f"Row band {k + 1}", (x1 - x0, rh * 0.94), ((x0 + x1) / 2, cy), color=[1, 1, 1], opacity=10,
                     round=min(rh * 0.15, 8), grow="right", t0=t0, dur=0.3)
            for j, cell in enumerate(r_[:nc]):
                text(f"Cell {k + 1}.{j + 1}", str(cell), fs, (cw * 0.92, rh * 0.9), (x0 + cw * (j + 0.5), cy + fs * 0.35),
                     align="center", font="semibold" if j == 0 else "regular", color=[1, 1, 1] if j == 0 else soft, t0=t0 + 0.1)
    elif typ == "quote":              # idézet: nagy idézőjel, a szöveg, alatta a szerző
        q = items[0].get("text") or ""
        size = round(m * (0.06 if len(q) < 90 else 0.05 if len(q) < 160 else 0.042))
        text("Quote mark", "“", round(m * 0.24), (m * 0.3, m * 0.26), (W / 2, yT + m * 0.06), align="center", font="bold",
             color=accent or _rgb(None), t0=0.2)
        text("Quote", q, size, (W * 0.82, (yB - yT) * 0.7), (W / 2, (yT + yB) / 2 + m * 0.04), align="center", font="semibold", t0=0.5)
        if items[0].get("label"):
            text("Quote author", "— " + str(items[0]["label"]), round(m * 0.036), (W * 0.8, m * 0.08), (W / 2, yB + m * 0.03),
                 align="center", font="regular", color=soft, t0=1.1)
    else:   # line / area
        lo, hi = min(vals), max(vals)
        pad = (hi - lo) * 0.12 or abs(hi) * 0.1 or 1
        lo, hi = lo - pad, hi + pad
        top, bottom = yT + vl * 1.8, yB
        xs = [x0 + (x1 - x0) * (0.06 + 0.88 * k / max(n - 1, 1)) for k in range(n)]
        ys = [bottom - (v - lo) / (hi - lo) * (bottom - top) for v in vals]
        rect("Baseline", (x1 - x0, max(2, m * 0.003)), ((x0 + x1) / 2, bottom), color=[1, 1, 1], opacity=45, t0=0.1, dur=0.4, grow="right")
        seg = min(0.45, 2.4 / max(n - 1, 1))
        if typ == "area":             # kitöltés a vonal alatt: keskeny, halvány oszlopok balról jobbra nőnek
            K = 48
            span = xs[-1] - xs[0]
            for j in range(K):
                xj = xs[0] + span * (j + 0.5) / K
                kk = min(n - 2, int((xj - xs[0]) / (span / max(n - 1, 1))))
                f_ = (xj - xs[kk]) / ((xs[kk + 1] - xs[kk]) or 1)
                yj = ys[kk] + (ys[kk + 1] - ys[kk]) * f_
                rect(f"Area {j + 1}", (span / K * 1.04, bottom - yj), (xj, (yj + bottom) / 2), color="accent", opacity=30, grow="up",
                     t0=round(0.35 + (n - 1) * seg * j / K, 3), dur=0.35)
        th, d = max(4, m * 0.006), max(10, m * 0.016)
        show = set(range(n)) if n <= 8 else {0, n - 1, vals.index(max(vals)), vals.index(min(vals))}
        every = 1 if n <= 8 else math.ceil(n / 8)
        for k in range(n):
            t0 = 0.35 + k * seg
            if k < n - 1:
                dx, dy = xs[k + 1] - xs[k], ys[k + 1] - ys[k]
                length = math.hypot(dx, dy)
                out.append({"k": "rect", "name": f"Line {k + 1}", "size": [round(length, 1), round(th, 1)],
                            "pos": [round(xs[k], 1), round(ys[k], 1)], "rot": round(math.degrees(math.atan2(dy, dx)), 2),
                            "color": "accent", "opacity": 100, "grow": "segment", "t0": t0, "dur": seg})
            rect(f"Point – {items[k]['label']}", (d, d), (xs[k], ys[k]), color="accent", opacity=100, round=d / 2, grow="pop", t0=t0, dur=0.25)
            if k in show:
                text(f"Value – {items[k]['label']}", _fmt_val(vals[k], unit, lang), vl, (vl * 5, vl * 1.4), (xs[k], ys[k] - vl * 1.0),
                     align="center", font="bold", t0=t0 + 0.15)
            if k % every == 0 or k == n - 1:
                text(f"Label – {items[k]['label']}", items[k]["label"], cl, ((x1 - x0) / min(n, 8), cl * 2.7),
                     (xs[k], bottom + cl * 2.0), align="center", font="semibold", color=soft, t0=0.2)
    return out


def _focus(path: Path) -> list[float]:
    """A kép „érdekes pontja" (0–1 koordináták) a Ken Burns-mozgás céljának: spectral residual szaliencia
    (Hou–Zhang; csak numpy + OpenCV-alap — az OpenCV 5-ből kikerült az arcfelismerő), enyhe középre húzással."""
    import cv2
    import numpy as np
    from PIL import Image
    try:
        im = Image.open(path).convert("L")
        im.thumbnail((128, 128))
        g = np.asarray(im, dtype=np.float32)
        F = np.fft.fft2(g)
        la = np.log(np.abs(F) + 1e-6)
        sal = np.abs(np.fft.ifft2(np.exp((la - cv2.blur(la, (3, 3))) + 1j * np.angle(F)))) ** 2
        sal = cv2.GaussianBlur(sal.astype(np.float32), (0, 0), 3)
        h, w = sal.shape
        yy, xx = np.mgrid[0:h, 0:w]
        sal *= np.exp(-(((xx - w / 2) / w) ** 2 + ((yy - h / 2) / h) ** 2))      # enyhe középre húzás
        top = (sal >= np.quantile(sal, 0.95)).astype(np.uint8)
        n, lab, stats, cent = cv2.connectedComponentsWithStats(top)
        if n > 1:      # a legerősebb összefüggő folt közepe (a szétszórt csúcsok átlaga középre húzna)
            k = 1 + int(np.argmax([sal[lab == j].sum() for j in range(1, n)]))
            cx, cy = cent[k]
        else:
            cx, cy = w / 2, h / 2
        return [float(round(min(0.8, max(0.2, cx / w)), 3)), float(round(min(0.8, max(0.2, cy / h)), 3))]
    except Exception:
        return [0.5, 0.45]


def _motions(n: int, ken_burns: str) -> list[str]:
    """Jelenetenkénti képmozgás. ken_burns: some (alap: kb. minden második jelenet határozott Ken Burns —
    közelítés + úsztatás a kép érdekes pontja felé, minden második ilyen fordítva, kifelé —, a többi finom
    mozgás) | all | off (csak a finom mozgás)."""
    gentle = ["in", "left", "out", "right"]
    out = []
    for i in range(n):
        kb = ken_burns == "all" or (ken_burns == "some" and i % 2 == 0)
        out.append(("kb-out" if (i // 2 if ken_burns == "some" else i) % 2 else "kb-in") if kb else gentle[i % 4])
    return out


def _cards(total: Optional[float]) -> tuple[float, float]:
    """A címkártya és a zárókép hossza — ezek is a kért hosszba számítanak, ezért rövid videón arányosan rövidebbek
    (8 s-nál 1,3 + 1,4 s)."""
    t = total or 60
    if t >= 22:
        return 3.6, 4.2
    return round(min(2.6, max(1.0, 0.16 * t)), 2), round(min(3.0, max(1.2, 0.18 * t)), 2)


def _scene_floor(total: float, n: int, Tt: float, Te: float) -> float:
    """A jelenet legrövidebb hossza: rendesen 3 s; ha a kért hossz ennél szorosabb, ami jut (legalább MIN_SCENE)."""
    return max(MIN_SCENE, min(3.0, (total - Tt - Te + XFADE * (n + 1)) / max(n, 1)))


def compose(story: dict[str, Any], *, fmt: str, total: Optional[float], images: list[Optional[str]],
            accent: Optional[str], credits: str, voice: Optional[str], voice_s: Optional[float], music: Optional[str],
            focus: Optional[dict[str, list[float]]] = None, ken_burns: str = "some",
            fits: Optional[dict[str, str]] = None, fill: bool = False, scrim: Optional[float] = None) -> dict[str, Any]:
    """A forgatókönyvből a builder adatcsomagja (D): méretek, jelenet-hosszak, feliratok időzítve. fill (fix hossz): ha
    minden jelenetnek saját hossza van, és együtt rövidebbek a kértnél, arányosan nyúlnak — az áttűnések átfedését is
    kiegyenlítve. scrim: a képek fölötti sötétítés (0–1; None = a beépített), a jelenet saját scrim-je felülírja."""
    W, H = FORMATS[fmt]
    L = _layout(W, H)
    m = min(W, H)
    scenes = story["scenes"]
    n = len(scenes)
    X = XFADE
    if voice_s:
        total = max(8.0, voice_s + 1.0)
    Tt, Te = _cards(total)
    weights = [max(len(s.get("caption") or s.get("headline") or ""), 45) for s in scenes]
    # a jelenetterv (scenes) adhat fix hosszt jelenetenként; a többi a maradékon osztozik a felirat hossza szerint
    fixed = [_scene_seconds(s) for s in scenes]
    if total is None:                        # saját szöveg, hossz nélkül: olvasási tempó (~14 karakter/s)
        durs = [f or max(3.5, w / 14 + 1.0) for f, w in zip(fixed, weights)]
        total = Tt + sum(durs) + Te - X * (n + 1)
    else:
        lo = _scene_floor(total, n, Tt, Te)
        room = max(total - Tt - Te + X * (n + 1), lo * n)
        free = room - sum(f for f in fixed if f)
        fw = sum(w for f, w in zip(fixed, weights) if not f) or 1
        durs = [f or max(lo, free * w / fw) for f, w in zip(fixed, weights)]
        loose = [i for i, f in enumerate(fixed) if not f]
        if loose:                            # a szabad jelenetek igazodnak, hogy a teljes hossz kijöjjön
            k = max(0.2, (room - sum(f for f in fixed if f)) / max(sum(durs[i] for i in loose), 1e-6))
            for i in loose:
                durs[i] = max(lo, durs[i] * k)
        elif fill and sum(durs) < room - 0.01:
            # minden jelenet hossza megadva, de a cím + zárókép + (n + 1) × 0,5 s áttűnés átfedése után rövidebb a kértnél
            # (10-19: 10 s kérve, 3 × 2,5 s → 8,9 s): a jelenetek arányosan nyúlnak (a megadott hossz minimum, mint a
            # narrációs illesztésben)
            k = room / sum(durs)
            durs = [d * k for d in durs]
        total = Tt + sum(durs) + Te - X * (n + 1)
    total = round(total, 2)
    cpl = L["capW"] / (0.53 * L["capSize"])
    max_chunk = int(cpl * 2 * 0.9)
    captions, t = [], Tt - X
    d_scenes = []
    motions = _motions(len(scenes), ken_burns)
    lang = _lang_of(" ".join([story.get("title") or ""] + [x.get("caption") or "" for x in scenes]))
    for i, (s, d) in enumerate(zip(scenes, durs)):
        cap = s.get("caption") or ""
        parts = _chunks(cap, max_chunk) if cap else []
        a, b = t + X + 0.15, t + d - X - 0.05
        tot = sum(len(p) for p in parts) or 1
        cur = a
        for p in parts:
            span = (b - a) * len(p) / tot
            captions.append({"t0": round(cur, 3), "t1": round(cur + span, 3), "text": p, "lines": 1 if len(p) <= cpl * 0.95 else 2})
            cur += span
        head = s.get("headline") or ""
        stat = s.get("stat") or None
        if stat:      # a nagy szám EGY sorban férjen el (10-12: „7,800 MW” két sorba tört és a címkére csúszott)
            stat = {**stat, "size": round(min(L["statSize"], W * 0.88 / (0.62 * max(len(stat["value"]), 1))))}
        chart = s.get("chart")
        if chart:
            stat = None
            # cím nélküli graphic/script is a jelenet címét kapja (10-13: a headline nyomtalanul elveszett — az ábrás
            # jelenet a headline-t nem rajzolja, csak az ábra címét)
            chart = {"title": chart.get("title") or head,
                     "elements": chart_elements(chart, W, H, lang, _rgb(accent), dur=d),
                     **({"backdrop": chart["backdrop"]} if chart.get("backdrop") is not None else {}),
                     **({"code": chart["code"], "id": s.get("id") or f"s{i + 1}"} if chart["type"] == "script" else {})}
            if s["chart"]["type"] == "graphic":      # a saját grafika szövegeit a renderer megméri (_text_report)
                _probe_texts(chart["elements"], f"scene:{s.get('id') or f's{i + 1}'}")
        slot = _slots(images[i + 1] if i + 1 < len(images) else None)
        mv = MOTIONS.get(s.get("motion") or "auto") or motions[i]
        # képsor: a megadott mozgás mindegyikre; különben váltakozva (közelít, úszik, távolodik…), hogy legyen ritmusa
        pics = [{"file": f, "motion": mv if k == 0 or s.get("motion") else ("kb-in", "left", "kb-out", "right")[k % 4],
                 "focus": s.get("focus") or (focus or {}).get(f) or [0.5, 0.45],
                 "fit": s.get("fit") or (fits or {}).get(f) or "cover"} for k, f in enumerate(slot) if f]
        d_scenes.append({"headline": head, "stat": stat, "chart": chart, "dur": round(d, 3), "image": pics[0]["file"] if pics else None,
                         "motion": pics[0]["motion"] if pics else mv, "focus": pics[0]["focus"] if pics else [0.5, 0.45],
                         "fit": pics[0]["fit"] if pics else "cover",
                         **({"pics": pics} if len(pics) > 1 else {}), **({"zoom": s["zoom"]} if s.get("zoom") else {}),
                         **({"camera": _camera_keys(s["camera"], d, W, H)} if s.get("camera") else {}),
                         **({"bg": s["background"]} if s.get("background") else {}),
                         **({"scrim": s["scrim"]} if s.get("scrim") is not None else {"scrim": scrim} if scrim is not None else {}),
                         "size": _size_for(head, m * (0.066 if fmt != "landscape" else 0.058), 34, m * 0.046)})
        t += d - X
    title = story.get("title") or ""
    end = story.get("end") or {}
    end_title = end.get("title") or title
    # zárókép: a csík a cím ALATT, az alcím a csík alatt — a cím becsült sorszámából (10-12: fix helyen az alcímre csúszott)
    end_size = _size_for(end_title, m * 0.072, 30, m * 0.05)
    end_lines = max(1, math.ceil(len(end_title) / max(1.0, (W * 0.84) / (0.55 * end_size))))
    end_top = H * (0.30 if H >= W else 0.26)
    L["endY"] = end_top + m * 0.22
    L["endBarY"] = end_top + end_lines * end_size * 1.1 + H * 0.03
    L["endSubY"] = L["endBarY"] + H * 0.035 + L["capSize"] * 1.7
    tl, t0_ = [{"id": "title", "start": 0.0, "end": round(Tt, 2)}], Tt - X
    for s_, d_ in zip(scenes, durs):
        tl.append({"id": s_.get("id") or "", "start": round(t0_, 2), "end": round(t0_ + d_, 2)})
        t0_ += d_ - X
    tl.append({"id": "end", "start": round(t0_, 2), "end": round(total, 2)})
    D = {
        "timeline": tl,
        "W": W, "H": H, "fps": FPS, "total": total, "xfade": X, "accent": _rgb(accent), "dark": [0.055, 0.06, 0.08],
        "font": FONT,
        "L": {**L, "titleH": m * 0.44, "headH": m * 0.36,
              "kickerY": L["kickerTop"] + L["kickerSize"] * 0.8, "titleY": L["titleTop"] + m * 0.22, "headY": L["headTop"] + m * 0.18},
        "title": {"text": title, "kicker": story.get("kicker") or "", "image": _slots(images[0])[0] if images else None, "dur": Tt,
                  **({"scrim": scrim} if scrim is not None else {}),
                  "motion": "in", "size": _size_for(title, m * (0.088 if fmt != "landscape" else 0.075), 26, m * 0.056),
                  "fit": (fits or {}).get(_slots(images[0])[0] if images else "") or "cover",
                  **({"pics": [{"file": f, "motion": ("in", "left", "out", "right")[k % 4], "focus": (focus or {}).get(f) or [0.5, 0.45]}
                               for k, f in enumerate(_slots(images[0])) if f]} if images and isinstance(images[0], list) else {})},
        "scenes": d_scenes,
        "end": {"title": end_title, "subtitle": end.get("subtitle") or "", "credits": credits, "dur": Te,
                **({"scrim": scrim} if scrim is not None else {}),
                "image": _slots(images[0])[0] if images else None, "size": end_size},
        "captions": captions,
        "audio": {"voice": voice, "music": music, "musicDb": -20 if voice else -10},
    }
    return D


# ---------- időkódos rétegek (layers): EGY modell a make- és az overlay-videón ----------
# Réteg = {id (L1…), type, at, dur, …}: abszolút idő a fő idővonalon, saját kompozíció (Layer L1 – …), a geometriát
# itt számoljuk (mint a chart_elements), a builder csak lerakja (overlayLayers). Az ábra, a szám és az idézet a forrással
# ellenőrizve — forrás nélkül a réteg a saját szövegével megy át (NEM ellenőrizve; a check ezt ki is mondja).

LAYER_TYPES = ("lower_third", "title", "chapter", "source", "stat", "quote", "chart", "fullscreen", "split")
_LAYER_DUR = {"lower_third": 5.0, "title": 4.0, "chapter": 3.5, "source": 5.0, "stat": 5.0, "quote": 6.0, "chart": 6.0,
              "fullscreen": 6.0, "split": 8.0}
_LAYER_ANIM = {"lower_third": ("slide", "fade"), "title": ("fade", "fade"), "chapter": ("slide", "fade"), "source": ("fade", "fade"),
               "stat": ("pop", "fade"), "quote": ("fade", "fade"), "chart": ("pop", "fade"), "fullscreen": ("fade", "fade"),
               "split": ("fade", "fade")}
_LAYER_ANIMS = ("fade", "slide", "pop", "none")
_LAYER_POS = ("auto", "top", "bottom", "left", "right", "center")
# z-sorrend: a teljes képes és az osztott alul, a névsáv és a forrásfelirat legfelül (egy fajtán belül az idő szerint)
_LAYER_Z = {"fullscreen": 0, "split": 1, "chart": 2, "quote": 2, "stat": 2, "title": 3, "chapter": 3, "lower_third": 4, "source": 5}
_LAYER_ALIAS = {"lowerthird": "lower_third", "lower-third": "lower_third", "full_screen": "fullscreen", "split_screen": "split",
                "splitscreen": "split", "credit": "source"}
# a rövidített réteg ({"type": "race", …}) ezeket viszi át az ábrába — a saját alakú ábrák mezői is (10-17: a „series”,
# az „x_title”, a magyarázó ábrák „nodes”/„edges” mezője eddig elveszett a rövidítésben)
_CHART_KEYS = ("title", "unit", "items", "elements", "code", "total", "min", "max", "backdrop", "labels", "series", "x_title",
               "y_title", "forecast_from", "total_label", "flows", "nodes", "edges", "root", "branches", "steps", "lanes",
               "actors", "reveal", "pace", "direction", "center", "focus", "columns", "rows")
_PANEL = [0.035, 0.04, 0.055]
# a hibaüzenetek minimális példája (10-19: a csonkolt leírás után a „type None” hiba nem mondta meg, mi a jó alak)
_LAYER_EXAMPLE = '{"at": "00:02", "dur": 5, "type": "lower_third", "name": "Anna Kovács", "role": "Mayor"}'


def _own_text(x: Any) -> str:
    """A réteg saját szövegei és számai egy szövegben (forrás nélkül ezzel megy át az ellenőrzésen)."""
    if isinstance(x, dict):
        return "\n".join(_own_text(v) for v in x.values())
    if isinstance(x, (list, tuple)):
        return "\n".join(_own_text(v) for v in x)
    return "" if x is None or isinstance(x, bool) else str(x)


def _layer_from(raw: dict[str, Any], source: str, notes: list[str], base: Optional[dict[str, Any]] = None) -> Optional[dict[str, Any]]:
    """Egy réteg (a kliensé, vagy szerkesztésnél base + a változó mezők) tisztítva: időkód → s, szöveghatárok, az ábra /
    szám / idézet a forrás számaival (ami nincs benne, kimarad — None, a ⚠ megjegyzés a check-be kerül)."""
    if not isinstance(raw, dict):
        raise ValueError("layers[]: each layer is an object {at, dur | until, type, …}")
    if not raw.get("image"):                    # rokon kulcsok a réteg képére (10-18: a tervező a jelenetek „image_query”-jét adta)
        q = next((raw[k] for k in ("image_prompt", "generate", "ai_image", "ai_picture") if isinstance(raw.get(k), str) and raw[k].strip()), None)
        if q:
            raw = {**raw, "image": {"generate": q}}
        elif isinstance(raw.get("image_query"), str) and raw["image_query"].strip():
            raw = {**raw, "image": {"stock": raw["image_query"]}}
    r = {**(base or {}), **raw}
    if r.get("type") in (None, "") and isinstance(r.get("chart"), dict):
        r["type"] = "chart"                     # ábra-objektum típus nélkül: ábra-réteg (10-19: „type None” hiba lett)
    elif r.get("type") in (None, "") and isinstance(r.get("elements"), list):
        r["type"] = "graphic"                   # elemek típus nélkül: saját grafika
    lid = re.sub(r"[^\w-]", "", str(r.get("id") or ""))[:20] or "L?"
    what = f"Layer {lid}"
    typ = _LAYER_ALIAS.get(str(r.get("type") or "").strip().lower().replace(" ", "_"),
                           str(r.get("type") or "").strip().lower().replace(" ", "_"))
    if typ in _CHART_TYPES and typ != "quote":          # rövidítés: {"type": "column", "items": …} = ábra-réteg
        # (szerkesztésnél a meglévő ábrára: {"id": "L3", "type": "bar"} csak a típust cseréli, a tételek maradnak)
        r = {**r, "chart": {**((base or {}).get("chart") or {}), "type": typ, **{k: raw[k] for k in _CHART_KEYS if k in raw}}}
        raw = {**raw, "chart": r["chart"]}
        typ = "chart"
    if typ not in LAYER_TYPES:
        raise ValueError(f"{what}: type {r.get('type')!r} — use {' | '.join(LAYER_TYPES)} (or a chart type such as \"column\"), "
                         f"e.g. {_LAYER_EXAMPLE} — the full list: video(action=\"schema\", topic=\"layers\").")
    if r.get("at") is None:
        raise ValueError(f"{what}: give at — when it starts (seconds, or a timecode such as \"01:24\"), e.g. {_LAYER_EXAMPLE}.")
    at = timecode(r["at"], f"{what} at")
    if raw.get("until") is not None:
        dur = timecode(raw["until"], f"{what} until") - at
    elif r.get("dur") is not None:
        dur = timecode(r["dur"], f"{what} dur")
    else:
        dur = _LAYER_DUR[typ]
    if dur < 0.5:
        raise ValueError(f"{what}: it must last at least 0.5 s (dur, or an until after its at).")
    ly: dict[str, Any] = {"id": lid, "type": typ, "at": round(at, 3), "dur": round(min(dur, 6 * 3600.0), 3)}
    pos = r.get("position")
    if isinstance(pos, (list, tuple)):
        if len(pos) != 4 or not all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in pos):
            raise ValueError(f"{what}: position [x, y, w, h] = centre and size, 0–1 of the frame — or auto | top | bottom | left | right | center.")
        x, y, w, h = (float(v) for v in pos)
        ly["position"] = [max(0.0, min(1.0, x)), max(0.0, min(1.0, y)), max(0.05, min(1.0, w)), max(0.03, min(1.0, h))]
    elif pos not in (None, "", "auto"):
        if str(pos).lower() not in _LAYER_POS:
            raise ValueError(f"{what}: position {pos!r} — auto | top | bottom | left | right | center | [x, y, w, h].")
        ly["position"] = str(pos).lower()
    anim = r.get("anim")
    if anim not in (None, ""):
        a_in, a_out = (anim.get("in"), anim.get("out")) if isinstance(anim, dict) else (anim, anim)
        for a_ in (a_in, a_out):
            if a_ not in (None, "") and str(a_).lower() not in _LAYER_ANIMS:
                raise ValueError(f"{what}: anim {a_!r} — {' | '.join(_LAYER_ANIMS)} (or {{\"in\": …, \"out\": …}}).")
        ly["anim"] = {k: str(v).lower() for k, v in (("in", a_in), ("out", a_out)) if v not in (None, "")}
    checked = bool((source or "").strip())
    src = source if checked else _own_text(r)
    nums = _text_numbers(src)
    cnotes: list[str] = []

    def chart_of(cc: Any, quiet: bool = False) -> Optional[dict[str, Any]]:
        if isinstance(cc, dict) and not cc.get("type") and (base or {}).get("chart"):
            cc = {**base["chart"], **cc}          # részleges javítás: csak a megadott mezők változnak (pl. a cím)
        if not isinstance(cc, dict):
            raise ValueError(f"{what}: chart = {{type, title, unit, items}} — any scene chart type (graphic, map, lottie and script "
                             "too), e.g. {\"at\": 9, \"chart\": {\"type\": \"column\", \"title\": \"Budget\", \"unit\": \"bn Ft\", "
                             "\"items\": [{\"label\": \"2024\", \"value\": 44.8}, {\"label\": \"2025\", \"value\": 47.5}]}} — "
                             "video(action=\"schema\", topic=\"chart\").")
        chart, note = _chart_from_plan(cc, src)
        for n_ in (note if isinstance(note, list) else [note] if note else []):
            if n_.startswith("⚠ ") or not chart:
                cnotes.append("⚠ " + n_.lstrip("⚠ "))
            elif not quiet:
                cnotes.append(n_ if checked else f"“{chart.get('title') or chart['type']}” ({chart['type']}): not checked against a source.")
        if cc and chart is None and not note:
            cnotes.append(f"⚠ The {cc.get('type') or 'chart'} “{cc.get('title') or ''}” was left out: "
                          f"{'unknown type' if cc.get('type') not in _CHART_TYPES else 'not enough usable items (check items / elements)'}.")
        return chart

    if typ == "lower_third":
        ly["name"] = _clean(r.get("name") or r.get("text"), 60, "Lower-third name")
        if not ly["name"]:
            raise ValueError(f"{what} (lower_third): give name (and role).")
        ly["role"] = _clean(r.get("role") or r.get("subtitle"), 90, "Lower-third role")
        if str(r.get("side") or "").lower() in ("left", "right"):
            ly["side"] = str(r["side"]).lower()
    elif typ in ("title", "chapter"):
        ly["text"] = _clean(r.get("text") or r.get("title"), 80, "Title")
        if not ly["text"]:
            raise ValueError(f"{what} ({typ}): give text (and subtitle?).")
        ly["subtitle"] = _clean(r.get("subtitle"), 110, "Subtitle")
        ly["kicker"] = _clean(r.get("kicker"), 30, "Kicker")
    elif typ == "source":
        ly["text"] = _clean(r.get("text"), 140, "Source caption")
        if not ly["text"]:
            raise ValueError(f"{what} (source): give text, e.g. \"Source: Eurostat, 2025\".")
    elif typ == "stat":
        st = r.get("stat") if isinstance(r.get("stat"), dict) else r
        ly["value"], ly["label"] = _clean(st.get("value"), 12, "Stat value"), _clean(st.get("label"), 70, "Stat label")
        if not ly["value"]:
            raise ValueError(f"{what} (stat): give value (and label), e.g. {{\"value\": \"72%\", \"label\": \"turnout\"}}.")
        if [x for x in _text_numbers(f"{ly['value']} {ly['label']}") if x not in nums]:
            notes.append(f"⚠ {what}: the number “{ly['value']}” is not in the source — the layer was left out.")
            return None
        cnotes.append(f"Stat “{ly['value']}”: the number is in the source text." if checked else f"Stat “{ly['value']}”: not checked against a source.")
    elif typ == "quote":
        chart = chart_of({"type": "quote", "items": [{"text": r.get("text"), "label": r.get("label") or r.get("author") or r.get("who")}]})
        if not chart:
            notes += [f"⚠ {what}: {n_[2:]}" for n_ in cnotes if n_.startswith("⚠ ")] or [f"⚠ {what}: the quote was left out."]
            return None
        ly["chart"] = chart
    elif typ == "chart":
        chart = chart_of(raw.get("chart") if "chart" in raw else r.get("chart"))
        if not chart:
            notes += [f"⚠ {what}: {n_[2:]}" for n_ in cnotes if n_.startswith("⚠ ")] or [f"⚠ {what}: the chart was left out."]
            return None
        ly["chart"] = chart
        if r.get("panel") is not None:
            ly["panel"] = bool(r["panel"])
    else:                                                   # fullscreen | split: ábra, kép vagy szöveg
        kind = next((k for k in ("chart", "image", "text") if raw.get(k)), None) or (base or {}).get("kind") \
            or next((k for k in ("chart", "image", "text") if r.get(k)), None)
        if not kind:
            raise ValueError(f"{what} ({typ}): give chart, image or text.")
        ly["kind"] = kind
        if kind == "chart":
            chart = chart_of(raw.get("chart") if "chart" in raw else r.get("chart"))
            if not chart:
                notes += [f"⚠ {what}: {n_[2:]}" for n_ in cnotes if n_.startswith("⚠ ")] or [f"⚠ {what}: the chart was left out."]
                return None
            ly["chart"] = chart
        elif kind == "image":
            if raw.get("image") or not (base or {}).get("chart"):
                ly["image"] = r["image"]
                cc = {"type": "graphic", "backdrop": 0, "elements": [{"shape": "image", "image": _img_spec(r["image"]), "fit": "cover",
                                                                      "x": 0.5, "y": 0.5, "w": 1, "h": 1, "anim": "none",
                                                                      **({"clip": r["clip"]} if r.get("clip") else {}),
                                                                      "sound": r.get("sound") if r.get("sound") in ("replace", "mix", "off")
                                                                      else ("replace" if typ == "fullscreen" else "off")}]}
            else:                                     # szerkesztés: a már letöltött kép marad (asset), nem jön le újra
                ly["image"], cc = base.get("image"), base["chart"]
            ly["chart"] = chart_of(cc, quiet=True)
            if not ly["chart"]:
                raise ValueError(f"{what} ({typ}): image = a file_id, a URL or {{\"stock\": \"words\"}}.")
        else:
            ly["text"] = _clean(r.get("text"), 80, "Title")
            ly["subtitle"] = _clean(r.get("subtitle"), 110, "Subtitle")
            ly["kicker"] = _clean(r.get("kicker"), 30, "Kicker")
        if typ == "split" and str(r.get("side") or "").lower() in ("left", "right", "top", "bottom"):
            ly["side"] = str(r["side"]).lower()
    notes += [f"⚠ {what}: {n_[2:]}" if n_.startswith("⚠ ") else f"{what}: {n_}" for n_ in cnotes]
    return ly


def layers_from(raw: Any, source: str, notes: list[str], existing: Optional[list[dict[str, Any]]] = None) -> list[dict[str, Any]]:
    """A kliens rétegei → tisztított rétegek stabil azonosítóval (L1, L2… — szerkesztés után is ugyanarra mutatnak)."""
    if raw in (None, "", []):
        return []
    if not isinstance(raw, list):
        raise ValueError("layers: a list of {at, dur | until, type, …} objects.")
    if len(raw) + len(existing or []) > 60:
        raise ValueError("layers: 60 at most in one video.")
    given = [re.sub(r"[^\w-]", "", str(x.get("id")))[:20] for x in raw if isinstance(x, dict) and x.get("id")]
    used = {ly["id"] for ly in existing or []}
    dup = sorted({i for i in given if given.count(i) > 1 or i in used})
    if dup:
        raise ValueError(f"Layer id {dup[0]} is used twice — ids must be unique (leave id out and one is given).")
    used |= set(given)
    k = max([int(m_.group(1)) for u in used if (m_ := re.fullmatch(r"L(\d+)", u))] or [0])
    out = []
    for x in raw:
        if isinstance(x, dict) and not x.get("id"):
            k += 1
            while f"L{k}" in used:
                k += 1
            x = {**x, "id": f"L{k}"}
            used.add(x["id"])
        ly = _layer_from(x, source, notes)
        if ly:
            out.append(ly)
    return _drop_echo_titles(out, existing or [], notes)


def _drop_echo_titles(layers: list[dict[str, Any]], existing: list[dict[str, Any]], notes: list[str]) -> list[dict[str, Any]]:
    """Az a címréteg, amely egy vele egy időben látszó ábra CÍMÉT ismétli, kimarad: az ábra a címét maga is kiírja
    (UI-próba 10-18: „Három fontos gondolat” kétszer, a dupla panel a beszélő homlokába lógott)."""
    norm = lambda t: re.sub(r"\W+", " ", str(t or "")).strip().lower()
    charts = [(c["at"], c["at"] + c["dur"], norm((c.get("chart") or {}).get("title")))
              for c in existing + layers if c.get("type") == "chart" and (c.get("chart") or {}).get("title")]
    keep = []
    for ly in layers:
        t = norm(ly.get("text")) if ly.get("type") == "title" and not (ly.get("subtitle") or ly.get("kicker")) else ""
        if t and any(a < ly["at"] + ly["dur"] and ly["at"] < b and t == ct for a, b, ct in charts):
            notes.append(f"Layer {ly['id']}: left out — it only repeated the title of the chart shown at the same time.")
            continue
        keep.append(ly)
    return keep


def apply_layer_changes(layers: list[dict[str, Any]], changes: dict[str, Any], source: str, notes: list[str]) -> list[dict[str, Any]]:
    """Célzott szerkesztés id alapján: {"layers": [{"id": "L2", …a változó mezők} | {"id": "L2", "remove": true}],
    "add_layers": [ … ]} → az új rétegsor (a többi réteg érintetlen)."""
    out = [dict(ly) for ly in layers]
    ids = [ly["id"] for ly in out]
    for ch in changes.get("layers") or []:
        if not isinstance(ch, dict) or not ch.get("id"):
            raise ValueError("changes.layers[]: give the layer's id (\"L2\" — video action=info lists them), plus what changes "
                             "or \"remove\": true. New layers go in add_layers.")
        if ch["id"] not in ids:
            raise ValueError(f"There is no layer {ch['id']} (layers: {', '.join(ids) or 'none'}). New layers go in add_layers.")
        k = ids.index(ch["id"])
        if ch.get("remove"):
            out[k] = None                                     # type: ignore[call-overload]
            continue
        out[k] = _layer_from({k_: v for k_, v in ch.items() if k_ != "remove"}, source, notes, base=out[k])
    kept = [ly for ly in out if ly]
    # az új réteg nem kapja meg egy most törölt id-ját (a válaszban az „L4” ne mást jelentsen, mint a kérésben)
    return kept + layers_from(changes.get("add_layers") or [], source, notes, existing=layers)


def _layer_text(ly: dict[str, Any]) -> str:
    """A réteg fő szövege a listázáshoz és a kompozíció nevéhez."""
    t, ch = ly["type"], ly.get("chart") or {}
    if t == "lower_third":
        return ly["name"] + (f" — {ly['role']}" if ly.get("role") else "")
    if t == "stat":
        return f"{ly['value']} {ly.get('label') or ''}".strip()
    if ch.get("type") == "quote":
        return (ch.get("items") or [{}])[0].get("text") or ""
    if ly.get("kind") == "image":
        return "picture"
    if ch:                                       # a saját grafika a címével, vagy az első szövegével (10-19: mind „graphic “””)
        els = ch.get("elements") or []
        lab = ch.get("title") or next((str(e["text"]) for e in els if e.get("shape") in ("text", "pin") and e.get("text")), "")
        if not lab and ch.get("type") == "graphic":
            return f"graphic ({len(els)} element{'s' if len(els) != 1 else ''})"
        return f"{ch['type']} “{lab}”"
    return ly.get("text") or ""


def layer_listing(layers: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [{"id": ly["id"], "type": ly["type"], "at": ly["at"], "dur": ly["dur"], "until": round(ly["at"] + ly["dur"], 3),
             "timecode": f"{_tc(ly['at'])}–{_tc(ly['at'] + ly['dur'])}", "text": _layer_text(ly)}
            for ly in sorted(layers, key=lambda x: (x["at"], x["id"]))]


def _tx(name: str, st: str, size: float, x0: float, y_cap: float, w: float, *, lines: int = 1, align: str = "left",
        font: str = "semibold", color: Any = None, t0: float = 0.3, **o: Any) -> dict[str, Any]:
    """Szövegréteg a builder-nek: a doboz bal széle x0, szélessége w; az ELSŐ sor nagybetűinek közepe y_cap (mérve: a
    dobozszöveg első sorának nagybetű-közepe a doboz tetejétől 0,61 em-re van)."""
    size = int(round(size))
    bh = size * (1.22 * max(1, lines) + 0.6)
    top = y_cap - 0.61 * size
    return {"k": "text", "name": name, "str": st, "size": size, "box": [round(w, 1), round(bh, 1)],
            "pos": [round(x0 + w / 2, 1), round(top + bh / 2, 1)], "align": align, "font": font, "color": color,
            "leading": round(size * 1.22, 1), "t0": t0, **o}


def _stack(ly: dict[str, Any], x0: float, w: float, m: float, acc: list[float], *, size: float, align: str = "center") -> tuple[list[dict[str, Any]], float]:
    """Felcím · cím · akcentus-csík · alcím egymás alatt (0-tól lefelé), a szélességbe illesztve → (rétegek, magasság)."""
    out: list[dict[str, Any]] = []
    y = 0.0
    if ly.get("kicker"):
        ks = _fit_size(ly["kicker"].upper(), m * 0.03, w, m * 0.04, "Kicker")
        out.append(_tx("Kicker", ly["kicker"].upper(), ks, x0, y + 0.365 * ks, w, align=align, color=acc, t0=0.12, tracking=60))
        y += ks * 1.35
    s = _fit_size(ly["text"], size, w, size * 1.22 * 3, "Title")
    lines = min(3, _wrap(ly["text"], s, w)[0])
    out.append(_tx("Title", ly["text"], s, x0, y + 0.365 * s, w, lines=lines, align=align, font="bold", t0=0.2))
    y += s * (0.73 + 1.22 * (lines - 1)) + s * 0.42
    bw, bh = m * 0.1, max(5.0, m * 0.008)
    bx = x0 + w / 2 if align == "center" else x0 + bw / 2
    out.append({"k": "rect", "name": "Accent bar", "size": [round(bw, 1), round(bh, 1)], "pos": [round(bx, 1), round(y + bh / 2, 1)],
                "color": "accent", "opacity": 100, "round": round(bh / 2, 1), "grow": "center" if align == "center" else "right",
                "t0": 0.3, "dur": 0.4})
    y += bh
    if ly.get("subtitle"):
        s2 = _fit_size(ly["subtitle"], m * 0.038, w, m * 0.038 * 1.22 * 3, "Subtitle")
        sl = min(3, _wrap(ly["subtitle"], s2, w)[0])
        y += s2 * 0.62
        out.append(_tx("Subtitle", ly["subtitle"], s2, x0, y + 0.365 * s2, w, lines=sl, align=align, font="regular",
                       color=[0.86, 0.88, 0.92], t0=0.42))
        y += s2 * (0.73 + 1.22 * (sl - 1)) + s2 * 0.22
    return out, y


def _shift(els: list[dict[str, Any]], dx: float, dy: float) -> list[dict[str, Any]]:
    for e in els:
        e["pos"] = [round(e["pos"][0] + dx, 1), round(e["pos"][1] + dy, 1)]
    return els


def _even(v: float) -> int:
    return max(4, int(round(v / 2)) * 2)


def _layer_box(ly: dict[str, Any], W: int, H: int, safe: float, acc: list[float]) -> dict[str, Any]:
    """A réteg kompozíciója a képkockán: méret (w, h), közép (cx, cy), honnan csúszik be (slide), osztott képnél az
    alapsáv helye (split); a szöveges rétegeknél a szövegrétegek is (els) — a szövegek mérete a dobozt is méretezi."""
    t, m = ly["type"], min(W, H)
    mg = round(m * 0.05)
    land, vert = W > H * 1.2, H > W * 1.2
    pos = ly.get("position") or "auto"
    custom = pos if isinstance(pos, list) else None
    els: list[dict[str, Any]] = []
    split = None
    side = ly.get("side")
    if t == "lower_third":
        s1, s2, pad = m * 0.052, m * 0.032, m * 0.026
        bar = max(6.0, m * 0.009)
        inner = (custom[2] * W if custom else W * (0.86 if vert else 0.6)) - bar - 2 * pad
        s1 = _fit_size(ly["name"], s1, inner / 1.04, s1 * 1.25, "Lower-third name")
        tw, rl = _wrap(ly["name"], s1, inner)[2] * 1.04, 0
        if ly.get("role"):
            s2 = _fit_size(ly["role"], s2, inner, s2 * 1.22 * 2, "Lower-third role")
            rl, _, rw = _wrap(ly["role"], s2, inner)
            rl, tw = min(2, rl), max(tw, rw)
        w = _even(bar + 2 * pad + max(tw + s1 * 0.3, W * 0.14))
        top = pad * 0.95
        nb = top + 0.73 * s1                                       # a név alapvonala
        rt = nb + 0.22 * s1 + 0.42 * s2                            # a tisztség nagybetűinek teteje
        bottom = (rt + 0.73 * s2 + 1.22 * s2 * (rl - 1) + 0.22 * s2) if rl else nb + 0.22 * s1
        h = _even(bottom + pad * 0.95)
        side = side or ("right" if pos == "right" else "left")
        els = [{"k": "rect", "name": "Panel", "size": [w - bar, h], "pos": [round(bar + (w - bar) / 2, 1), h / 2], "color": _PANEL,
                "opacity": 86, "round": 0, "grow": "right", "t0": 0.05, "dur": 0.35},
               {"k": "rect", "name": "Accent bar", "size": [round(bar, 1), h], "pos": [round(bar / 2, 1), h / 2], "color": "accent",
                "opacity": 100, "round": 0, "grow": "up", "t0": 0.0, "dur": 0.3},
               _tx("Name", ly["name"], s1, bar + pad, top + 0.365 * s1, w - bar - 2 * pad, font="bold", t0=0.2)]
        if rl:
            els.append(_tx("Role", ly["role"], s2, bar + pad, rt + 0.365 * s2, w - bar - 2 * pad, lines=rl, font="regular",
                           color=[0.84, 0.86, 0.9], t0=0.32))
        cx = W / 2 if pos == "center" else (mg + w / 2 if side == "left" else W - mg - w / 2)
        cy = mg + h / 2 if pos == "top" else safe - h / 2
        slide = [-m * 0.05 if side == "left" else m * 0.05, 0]
    elif t == "source":
        s = m * 0.025
        maxw = (custom[2] * W if custom else W * 0.5)
        pad = s * 0.62
        s = _fit_size(ly["text"], s, maxw - 2 * pad, s * 1.22 * 2, "Source caption")
        lines, _, tw = _wrap(ly["text"], s, maxw - 2 * pad)
        lines = min(2, lines)
        w = _even(tw + 2 * pad + s * 0.3)
        h = _even(pad * 1.5 + s * (0.95 + 1.22 * (lines - 1)))
        els = [{"k": "rect", "name": "Panel", "size": [w, h], "pos": [w / 2, h / 2], "color": _PANEL, "opacity": 70,
                "round": round(min(h * 0.3, m * 0.012), 1)},
               _tx("Source", ly["text"], s, pad, pad * 0.75 + 0.365 * s, w - 2 * pad, lines=lines, font="regular",
                   color=[0.9, 0.91, 0.94], t0=0.0)]
        cx = mg + w / 2 if pos == "left" else W / 2 if pos == "center" else W - mg - w / 2
        cy = mg + h / 2 if pos == "top" else safe - h / 2
        slide = [0, 0]
    elif t in ("title", "chapter"):
        if t == "title":
            w = _even(custom[2] * W if custom else W)
            inner = min(w * 0.84, W * 0.84)
            stack, sh = _stack(ly, (w - inner) / 2, inner, m, acc, size=_size_for(ly["text"], m * 0.085, 24, m * 0.055))
            pad = m * 0.05
            h = _even(sh + 2 * pad)
            els = [{"k": "rect", "name": "Band", "size": [w, h], "pos": [w / 2, h / 2], "color": _PANEL, "opacity": 62, "round": 0}] \
                + _shift(stack, 0, pad)
            cx = W / 2
            cy = mg + h / 2 if pos == "top" else safe - h / 2 if pos == "bottom" else H / 2
            slide = [0, 0]
        else:
            pad = m * 0.04
            maxw = custom[2] * W if custom else W * (0.56 if land else 0.84)
            size = _size_for(ly["text"], m * 0.07, 22, m * 0.05)
            widest = max(_wrap(ly["text"], size, maxw - 2 * pad)[2] * 1.05,
                         _wrap(ly.get("subtitle") or "", m * 0.038, maxw - 2 * pad)[2], m * 0.2)
            w = _even(min(maxw, widest + 2 * pad))
            stack, sh = _stack(ly, pad, w - 2 * pad, m, acc, size=size, align="left")
            h = _even(sh + 2 * pad)
            els = [{"k": "rect", "name": "Panel", "size": [w, h], "pos": [w / 2, h / 2], "color": _PANEL, "opacity": 80,
                    "round": round(m * 0.01, 1)}] + _shift(stack, 0, pad)
            right = pos == "right"
            cx = W / 2 if pos == "center" else (W - mg - w / 2 if right else mg + w / 2)
            cy = mg + h / 2 if pos == "top" else safe - h / 2 if pos == "bottom" else H * (0.5 if land else 0.42)
            slide = [m * 0.05 if right else -m * 0.05, 0]
    elif t == "stat":
        w = _even(custom[2] * W if custom else W * (0.34 if land else 0.62 if not vert else 0.8))
        pad = m * 0.035
        vs = min(m * 0.16, (w - 2 * pad) / (0.62 * max(len(ly["value"]), 1)))
        s2 = _fit_size(ly.get("label") or "", m * 0.036, w - 2 * pad, m * 0.036 * 1.22 * 3, "Stat label")
        ll = min(3, _wrap(ly.get("label") or "", s2, w - 2 * pad)[0]) if ly.get("label") else 0
        vb = pad + 0.73 * vs                                         # a szám alapvonala
        h = _even(vb + (0.3 * vs + 0.42 * s2 + s2 * (0.73 + 1.22 * (ll - 1)) + 0.22 * s2 if ll else 0.15 * vs) + pad)
        els = [{"k": "rect", "name": "Panel", "size": [w, h], "pos": [w / 2, h / 2], "color": _PANEL, "opacity": 84,
                "round": round(m * 0.012, 1)},
               _tx("Number", ly["value"], vs, pad, pad + 0.365 * vs, w - 2 * pad, align="center", font="bold", color=acc, t0=0.15)]
        if ll:
            els.append(_tx("Number label", ly["label"], s2, pad, vb + 0.3 * vs + 0.42 * s2 + 0.365 * s2, w - 2 * pad, lines=ll,
                           align="center", t0=0.35))
        cx, cy = (W - mg - w / 2, H * 0.42) if land else (W / 2, mg + h / 2 + H * 0.04)
        if pos == "left":
            cx = mg + w / 2
        elif pos == "center":
            cx, cy = W / 2, H / 2
        elif pos == "top":
            cy = mg + h / 2
        elif pos == "bottom":
            cy = safe - h / 2
        slide = [0, 0]
    elif t in ("chart", "quote"):
        transparent = ly.get("panel") is False or (ly.get("panel") is None and (ly.get("chart") or {}).get("type") in ("graphic", "script"))
        if transparent and pos == "auto" and not custom:          # saját grafika a felvétel fölött: az egész képkocka
            w, h, cx, cy = W, H, W / 2, H / 2
        elif land and pos in ("auto", "right", "left"):
            w, h = _even(W * 0.46), _even(min(H * 0.84, safe - mg * 0.6))
            cx = mg + w / 2 if pos == "left" else W - mg - w / 2
            cy = (mg * 0.6 + safe) / 2
        elif pos == "center":
            w, h, cx, cy = _even(W * (0.62 if land else 0.9)), _even(H * (0.74 if land else 0.5)), W / 2, H / 2
        else:                                                    # álló / négyzetes, vagy top | bottom: széles sáv
            w, h = _even(W - 2 * mg), _even(H * (0.42 if (vert or land) else 0.5))
            cx = W / 2
            cy = safe - h / 2 if pos == "bottom" else mg + h / 2 + (H * 0.04 if vert else 0)
        if not transparent:
            els = [{"k": "rect", "name": "Panel", "size": [w, h], "pos": [w / 2, h / 2], "color": _PANEL,
                    "opacity": int(round(100 * (ly["chart"].get("backdrop") if ly["chart"].get("backdrop") is not None else 0.86))),
                    "round": round(m * 0.014, 1)}]
        slide = [0, 0]
    elif t == "fullscreen":
        w, h, cx, cy, slide = W, H, W / 2, H / 2, [0, 0]
        if ly.get("kind") != "image":                 # zárókártya: a kimerevített utolsó kocka halványan átlátszik
            els = [{"k": "rect", "name": "Background", "size": [W, H], "pos": [W / 2, H / 2], "color": [0.055, 0.06, 0.08],
                    "opacity": 80 if ly.get("end") else 100, "round": 0}]
        if ly.get("credits"):
            fs = max(12, round(m * 0.022))
            els.append(_tx("Picture credits", ly["credits"], fs, W * 0.05, H - mg - fs * 2.6, W * 0.9, lines=2, align="center",
                           font="regular", color=[0.75, 0.76, 0.8], t0=0.3))
    else:                                                          # split: az alapsáv a fél képre, a grafika a másikra
        side = side or ("bottom" if vert else "right")
        if side in ("left", "right"):
            w, h = _even(W / 2 - mg * 1.5), _even(H - 2 * mg)
            cx, cy = (W * 0.75 - mg * 0.25, H / 2) if side == "right" else (W * 0.25 + mg * 0.25, H / 2)
            k = min((W / 2 - mg * 1.5) / W, (H - 2 * mg) / H)
            to = [W * 0.25 + mg * 0.25, H / 2] if side == "right" else [W * 0.75 - mg * 0.25, H / 2]
            slide = [m * 0.06 if side == "right" else -m * 0.06, 0]
        else:
            w, h = _even(W - 2 * mg), _even(H / 2 - mg * 1.5)
            cx, cy = (W / 2, H * 0.75 - mg * 0.25) if side == "bottom" else (W / 2, H * 0.25 + mg * 0.25)
            k = min((W - 2 * mg) / W, (H / 2 - mg * 1.5) / H)
            to = [W / 2, H * 0.25 + mg * 0.25] if side == "bottom" else [W / 2, H * 0.75 - mg * 0.25]
            slide = [0, m * 0.06 if side == "bottom" else -m * 0.06]
        split = {"scale": round(k, 4), "to": [round(to[0], 1), round(to[1], 1)]}
        els = [{"k": "rect", "name": "Panel", "size": [w, h], "pos": [w / 2, h / 2], "color": _PANEL, "opacity": 94,
                "round": round(m * 0.014, 1)}]
    if custom and t not in ("lower_third", "source", "fullscreen", "split"):
        cx, cy = custom[0] * W, custom[1] * H
        if t in ("chart", "quote"):
            w, h = _even(custom[2] * W), _even(custom[3] * H)
            if els:
                els[0].update(size=[w, h], pos=[w / 2, h / 2])
    elif custom:
        cx, cy = custom[0] * W, custom[1] * H
    if t in ("fullscreen", "split") and ly.get("kind") == "text":
        size = _size_for(ly["text"], min(w, h) * 0.1, 24, min(w, h) * 0.06)
        stack, sh = _stack(ly, w * 0.08, w * 0.84, min(w, h) * 1.1, acc, size=size)
        els += _shift(stack, 0, (h - sh) / 2)
    return {"w": int(w), "h": int(h), "cx": round(cx), "cy": round(cy), "slide": [round(slide[0], 1), round(slide[1], 1)],
            "split": split, "els": els}


def _layer_elements(ly: dict[str, Any], b: dict[str, Any], lang: str, acc: list[float],
                    dur: Optional[float] = None) -> list[dict[str, Any]]:
    """A réteg összes builder-eleme: a doboz (panel, szövegek) + az ábra a meglévő építővel (chart_elements), a réteg
    kompozíciójának méretében; az ábra címe fölül, mint a jelenetekben."""
    els = list(b["els"])
    ch = ly.get("chart")
    if not ch:
        return els
    w, h = b["w"], b["h"]
    m = min(w, h)
    over = not any(e.get("name") in ("Panel", "Background") for e in els)       # átlátszó: árnyék a felvétel fölött
    # panelen nincs felirat-sáv alul: az ábra a panel aljáig ér (a jelenet-elrendezés alja a feliratoké) — ha ettől nem
    # vált álló elrendezésre; a saját grafika koordinátái a panel arányai maradnak
    hv = h / 0.9 if any(e.get("name") == "Panel" for e in els) and ch["type"] not in ("graphic", "script") \
        and (h / 0.9 > w * 1.2) == (h > w * 1.2) else h
    ce = chart_elements(ch, w, hv, lang, acc, dur=dur)      # a réteg hossza: az időzített ábra (verseny, kibontás) ebbe fér
    if over:
        for e in ce:
            if e["k"] == "text":
                e["shadow"] = 18
    els += ce
    if ch.get("title") and ch["type"] != "quote":
        cs = round(m * 0.056)
        els.append({"k": "text", "name": "Chart title", "str": ch["title"], "size": _fit_size(ch["title"], cs, w * 0.84, cs * 2.6, "Chart title"),
                    "box": [round(w * 0.84, 1), round(cs * 2.6, 1)], "pos": [round(w / 2, 1), round(h * 0.1, 1)], "align": "center",
                    "font": "bold", "leading": round(cs * 1.1, 1), "t0": 0.1, **({"shadow": 18} if over else {})})
    credit = next((e.get("credit") for e in ch.get("elements") or [] if e.get("shape") == "image" and e.get("credit")), None)
    if credit:
        fs = max(12, round(m * 0.022))
        short = re.sub(r"\s+—\s+https?://\S+", "", re.sub(r"^Photo:\s*", "", credit)).strip()[:120]
        els.append(_tx("Picture credit", short, fs, w * 0.04, h - fs * 1.4, w * 0.92, align="right", font="regular",
                       color=[0.86, 0.88, 0.92], t0=0.2, shadow=12))
    return els


def _loose(ly: dict[str, Any], b: dict[str, Any]) -> bool:
    """Szándékosan mindent takaró (teljes képes) vagy panel nélküli saját grafika (nyíl, jelölés a felvételen): ezek nem
    „ütköznek” a többi réteggel — a teljes képes alul van, a jelölésnek csak a rajzolt elemei látszanak."""
    return ly["type"] == "fullscreen" or (ly["type"] == "chart" and not any(e.get("name") == "Panel" for e in b["els"]))


def _overlaps(placed: list[tuple[dict[str, Any], float, float, dict[str, Any]]]) -> list[str]:
    """Két réteg ugyanott, ugyanakkor (a teljes képes kivételével — az szándékosan takar) → figyelmeztetés."""
    out = []
    for i, (a, a0, a1, ba) in enumerate(placed):
        for b, b0, b1, bb in placed[i + 1:]:
            if _loose(a, ba) or _loose(b, bb) or min(a1, b1) - max(a0, b0) < 0.3:
                continue
            ra = (ba["cx"] - ba["w"] / 2, ba["cy"] - ba["h"] / 2, ba["cx"] + ba["w"] / 2, ba["cy"] + ba["h"] / 2)
            rb = (bb["cx"] - bb["w"] / 2, bb["cy"] - bb["h"] / 2, bb["cx"] + bb["w"] / 2, bb["cy"] + bb["h"] / 2)
            ix = max(0.0, min(ra[2], rb[2]) - max(ra[0], rb[0])) * max(0.0, min(ra[3], rb[3]) - max(ra[1], rb[1]))
            if ix > 0.15 * min(ba["w"] * ba["h"], bb["w"] * bb["h"]):
                out.append(f"Layers {a['id']} ({a['type']}) and {b['id']} ({b['type']}) cover the same place from "
                           f"{_tc(max(a0, b0))} to {_tc(min(a1, b1))} — move one (position) or change its time.")
    return out


def _rect(b: dict[str, Any]) -> tuple[float, float, float, float]:
    return b["cx"] - b["w"] / 2, b["cy"] - b["h"] / 2, b["cx"] + b["w"] / 2, b["cy"] + b["h"] / 2


def _inter(r1: tuple[float, ...], r2: tuple[float, ...]) -> float:
    return max(0.0, min(r1[2], r2[2]) - max(r1[0], r2[0])) * max(0.0, min(r1[3], r2[3]) - max(r1[1], r2[1]))


def _cover(b: dict[str, Any], t0: float, t1: float, placed: list[tuple[dict[str, Any], float, float, dict[str, Any]]],
           zones: list[tuple[float, float, tuple[float, float, float, float], str]]) -> float:
    """Mennyire takarná a réteg (a doboza) a vele egy időben látszó rétegeket és a videó saját szövegeit (terület × idő)."""
    r = _rect(b)
    tot = sum(_inter(r, _rect(bb)) * (min(t1, a1) - max(t0, a0)) for ly_, a0, a1, bb in placed
              if min(t1, a1) > max(t0, a0) and not _loose(ly_, bb))
    return tot + sum(_inter(r, zr) * (min(t1, z1) - max(t0, z0)) for z0, z1, zr, _ in zones if min(t1, z1) > max(t0, z0))


def _auto_place(ly: dict[str, Any], b: dict[str, Any], t0: float, t1: float, placed: list, zones: list,
                W: int, H: int, safe: float, acc: list[float]) -> dict[str, Any]:
    """Hely nélkül megadott névsáv és forrásfelirat: az alaphely, ha szabad; különben a sarkok közül az, amelyiket a vele egy
    időben látszó rétegek és a videó saját szövegei (jelenetcím, ábra) a legkevésbé takarnak — a szövegek körül egy kis
    levegővel (ami épp csak nem ér hozzá, az is zsúfolt)."""
    pad = min(W, H) * 0.03
    zones = [(z0, z1, (r[0] - pad, r[1] - pad, r[2] + pad, r[3] + pad), w_) for z0, z1, r, w_ in zones]
    if ly["type"] == "source":
        mg = round(min(W, H) * 0.05)
        w, h = b["w"], b["h"]
        cands = [{**b, "cx": round(x), "cy": round(y)} for x, y in ((W - mg - w / 2, safe - h / 2), (mg + w / 2, safe - h / 2),
                                                                     (W - mg - w / 2, mg + h / 2), (mg + w / 2, mg + h / 2))]
    else:
        cands = [b] + [_layer_box({**ly, "side": sd, **({"position": ps} if ps else {})}, W, H, safe, acc)
                       for sd, ps in (("right", None), ("left", "top"), ("right", "top"))]
    return min(enumerate(cands), key=lambda ic: (_cover(ic[1], t0, t1, placed, zones), ic[0]))[1]


def _scene_zones(D: dict[str, Any]) -> list[tuple[float, float, tuple[float, float, float, float], str]]:
    """A make-videó saját szövegei és ábrái időben (címkártya, jelenetcím, ábra, nagy szám, zárókép): a hely nélkül megadott
    rétegek ezeket kerülik, és a check szól, ha egy réteg rájuk kerül."""
    W, H, L = D["W"], D["H"], D["L"]
    m = min(W, H)
    out = []
    for k, row in enumerate(D.get("timeline") or []):
        t0, t1 = row["start"] + (D["xfade"] if k else 0), row["end"]
        if row["id"] == "title":
            ti = D["title"]
            lines = min(3, _wrap(ti["text"], ti["size"], L["titleW"])[0])
            r = (L["titleX"] - L["titleW"] / 2, L["kickerTop"] - m * 0.02, L["titleX"] + L["titleW"] / 2,
                 L["titleTop"] + ti["size"] * (0.4 + 1.08 * lines))
            what = "the title card's title"
        elif row["id"] == "end":
            r, what = (W * 0.08, L["endY"] - m * 0.25, W * 0.92, L["endSubY"] + L["capSize"] * 1.5), "the end card's text"
        else:
            sc = D["scenes"][k - 1] if 0 < k <= len(D["scenes"]) else {}
            if sc.get("chart"):
                r, what = (W * 0.05, H * 0.04, W * 0.95, H * 0.78), f"scene {row['id']}'s chart"
            elif sc.get("stat"):
                r = (W * 0.05, L["statY"] - L["statSize"] * 0.6, W * 0.95, L["statLabelY"] + L["statLabelSize"] * 1.5)
                what = f"scene {row['id']}'s number"
            else:
                lines = min(3, _wrap(sc.get("headline") or "", sc.get("size") or 40, L["headW"])[0])
                r = (L["headX"] - L["headW"] / 2, L["headBarY"] - m * 0.02, L["headX"] + L["headW"] / 2,
                     L["headTop"] + (sc.get("size") or 40) * (0.4 + 1.12 * lines))
                what = f"scene {row['id']}'s headline"
        out.append((t0, t1, r, what))
    return out


def layer_comps(layers: list[dict[str, Any]], W: int, H: int, total: float, lang: str, acc: list[float],
                safe: float, zones: Optional[list] = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[str]]:
    """A rétegek a builder-nek (D.layers), az idővonal-soraik (előnézet, render scene=L2), és ami nem stimmel
    (a videó vége utánra esik, levágódik a végén, két réteg egymást, vagy egy réteg a videó saját szövegét takarja).
    zones: a make-videó saját szövegei időben (_scene_zones) — az overlay-videón nincs ilyen."""
    zones = zones or []
    out, rows, problems, placed = [], [], [], []
    for ly in sorted(layers, key=lambda x: (_LAYER_Z[x["type"]], x["at"], x["id"])):
        at, dur = float(ly["at"]), float(ly["dur"])
        if at >= total - 0.3:
            problems.append(f"Layer {ly['id']} starts at {_tc(at)}, but the video ends at {_tc(total)} — it was left out. "
                            "Give it an earlier at.")
            continue
        if at + dur > total + 0.05:
            problems.append(f"Layer {ly['id']} was asked to last until {_tc(at + dur)}; the video ends at {_tc(total)}, so it ends there.")
            dur = round(total - at, 3)
        b = _layer_box(ly, W, H, safe, acc)
        if ly["type"] in ("source", "lower_third") and not ly.get("position") and not ly.get("side"):
            b = _auto_place(ly, b, at, at + dur, placed, zones, W, H, safe, acc)
        a_in, a_out = _LAYER_ANIM[ly["type"]]
        if ly["type"] == "chart" and not any(e.get("name") == "Panel" for e in b["els"]):
            a_in = "fade"
        anim = {"in": (ly.get("anim") or {}).get("in") or a_in, "out": (ly.get("anim") or {}).get("out") or a_out}
        item: dict[str, Any] = {"id": ly["id"], "name": f"Layer {ly['id']} – {_layer_text(ly)[:32]}", "at": round(at, 3), "dur": round(dur, 3),
                                "w": b["w"], "h": b["h"], "pos": [b["cx"], b["cy"]], "elements": _layer_elements(ly, b, lang, acc, dur),
                                "anim": anim, "slide": b["slide"]}
        need = _chart_seconds(ly.get("chart"))
        if need and dur < need - 0.3:
            problems.append(f"Layer {ly['id']}: its {ly['chart']['type'].replace('_', ' ')} needs about {need:.0f} s to play "
                            f"through; the layer is {dur:.1f} s, so it runs fast — give it \"dur\": {math.ceil(need)}.")
        if b.get("split"):
            item["split"] = b["split"]
        if (ly.get("chart") or {}).get("type") == "script":
            item["code"] = ly["chart"]["code"]
        if (ly.get("chart") or {}).get("type") == "graphic":
            _probe_texts(item["elements"], ly["id"])
        out.append(item)
        rows.append({"id": ly["id"], "start": round(at, 2), "end": round(at + dur, 2), "type": ly["type"], "layer": True})
        placed.append((ly, at, at + dur, b))
    rows.sort(key=lambda r_: (r_["start"], r_["id"]))
    for ly, a0, a1, b in placed:                      # a videó saját szövegét / ábráját takaró réteg
        if _loose(ly, b):
            continue
        hit: dict[str, list[float]] = {}
        for z0, z1, zr, what in zones:
            if min(a1, z1) - max(a0, z0) >= 0.3 and _inter(_rect(b), zr) > 0.15 * b["w"] * b["h"]:
                span = hit.setdefault(what, [max(a0, z0), min(a1, z1)])
                span[0], span[1] = min(span[0], max(a0, z0)), max(span[1], min(a1, z1))
        problems += [f"Layer {ly['id']} ({ly['type']}) covers {what} from {_tc(sp[0])} to {_tc(sp[1])} — move it (position) or "
                     "change its time." for what, sp in hit.items()]
    return out, rows, problems + _overlaps(placed)


def layer_size(ly: dict[str, Any], W: int, H: int, safe: float) -> tuple[int, int]:
    """A réteg ábrájának mérete (a térképek és a saját grafikák ebben a méretben kerülnek a helyükre)."""
    b = _layer_box(ly, W, H, safe, _rgb(None))
    return b["w"], b["h"]


def _caption_top(W: int, H: int) -> float:
    """A make-videó felirat-sávjának teteje (két sorral): fölötte maradnak a rétegek."""
    L = _layout(W, H)
    return L["capY"] - (2 * L["capSize"] * 1.28 + L["capPad"] * 1.4) / 2 - L["capSize"] * 0.35


# ---------- építés ----------

async def _probe_duration(workdir: Path, name: str) -> Optional[float]:
    code = f'app.project.importFile(new ImportOptions(new File({json.dumps(name)}))).duration'
    try:
        out = await _run("script", "--eval", code, "--empty", cwd=workdir, timeout=120)
        return float(re.findall(r"[\d.]+", out.strip().splitlines()[-1])[0])
    except Exception as e:
        log.warning("hanghossz nem olvasható (%s): %s", name, e)
        return None


def _lost_picture(i: int, story: dict[str, Any]) -> str:
    where = "the title card" if i == 0 else f"scene {story['scenes'][i - 1].get('id') or i}"
    fix = ('changes={"title_image": …}' if i == 0 else
           f'changes={{"scenes": [{{"id": "{story["scenes"][i - 1].get("id")}", "image": …}}]}}')
    return (f"The picture for {where} couldn't be downloaded (the picture site refused or timed out), so it shows without "
            f"one — give it again with action=edit, {fix} (a file_id, a URL or {{\"stock\": \"words\"}}).")


async def _materialize_images(specs: list[Any], fmt: str, workdir: Path,
                              parent: Optional[str] = None) -> tuple[list[Optional[str]], list[str], list[Optional[str]]]:
    """Képleírások (file_id / {url} / {stock} / {generate}, vagy {"asset": név} = a szülőprojekt meglévő képe, letöltés
    nélkül) → fájlnevek a munkamappában + forrásmegjelölések (egyedi) + jelenetenkénti kredit."""
    from app.tools import images as imgs
    W, H = FORMATS[fmt]
    names: list[Optional[str]] = []
    credits: list[str] = []
    per: list[Optional[str]] = []
    parent_assets = assets_of(parent) if parent else {}

    async def one(i: int, spec: Any) -> tuple[Optional[str], Optional[str]]:
        if not spec:
            return None, None
        if isinstance(spec, dict) and spec.get("asset"):
            fid = parent_assets.get(spec["asset"])
            if fid:
                name = f"picture-{i + 1:02d}{Path(spec['asset']).suffix or '.jpg'}"
                shutil.copyfile(crafts.existing(fid), workdir / name)
                return name, spec.get("credit")
            return None, None
        try:
            # a saját fotó vágatlanul jön (a builder dönt: kitölti a keretet, vagy egészében mutatja — fit); a stock- és
            # AI-kép a keret arányára vágva, mint eddig
            own = _own_picture(spec)
            r = await imgs.resolve(spec, aspect=None if own else W / H, max_px=int(max(W, H) * 1.6) if own else max(W, H))
        except Exception as e:
            log.warning("videókép %s nem jött le (%s): %s", i, spec, e)
            return None, None
        name = f"picture-{i + 1:02d}.jpg"
        shutil.move(str(r.path), workdir / name)
        return name, r.attribution
    sem = asyncio.Semaphore(4)

    async def guarded(i, s):
        async with sem:
            return await one(i, s)
    # a képsor (lista) elemei külön fájlok: picture-02.jpg, picture-02b.jpg… — a hely maga lista marad
    flat = [(i, j, sp) for i, slot in enumerate(specs) for j, sp in enumerate(_slots(slot))]
    res = await asyncio.gather(*(guarded(i if j == 0 else 1000 * (j + 1) + i, sp) for i, j, sp in flat))
    by_slot: dict[int, list[tuple[Optional[str], Optional[str]]]] = {}
    for (i, j, _), (name, credit) in zip(flat, res):
        if name and j:
            new = f"picture-{i + 1:02d}{chr(ord('a') + j)}{Path(name).suffix}"
            shutil.move(str(workdir / name), workdir / new)
            name = new
        by_slot.setdefault(i, []).append((name, credit))
        if credit and credit not in credits:
            credits.append(credit)
    for i in range(len(specs)):
        got = [(nm, cr) for nm, cr in by_slot.get(i, []) if nm]
        if isinstance(specs[i], list) and len(got) > 1:
            names.append([nm for nm, _ in got])
            per.append([cr for _, cr in got])
        else:
            names.append(got[0][0] if got else None)
            per.append(got[0][1] if got else None)
    return names, credits, per


# a stock-jelöltek látómodell-rangsora közös (videó, pakli, Lottie): app/tools/images.py
from app.tools.images import PICK_N as _PICK_N, contact_sheet as _contact_sheet, rank as _rank  # noqa: E402


async def _choose_stock(jobs: list[tuple[int, str, str]]) -> dict[int, dict[str, Any]]:
    from app.tools import images as imgs
    return await imgs.choose_stock(jobs, medium="video")


async def _pick_stock(specs: list[Any], ctx: Callable[[int], str]) -> None:
    """A stock-keresések jelöltjeiből a látómodell választ (egy kép nem kerül két helyre) — a képsorok elemeire is."""
    jobs = [((i, j), sp["stock"], ctx(i)) for i, slot in enumerate(specs) for j, sp in enumerate(_slots(slot))
            if isinstance(sp, dict) and sp.get("stock")]
    if not jobs:
        return
    for (i, j), chosen in (await _choose_stock(jobs)).items():
        if isinstance(specs[i], list):
            specs[i][j] = chosen
        else:
            specs[i] = chosen


def _title_stem(title: str) -> str:
    t = re.sub(r"[\\/:*?\"<>|\s]+", " ", title or "").strip(" .")
    return (t[:60].rsplit(" ", 1)[0] if len(t) > 60 else t) or "Video"


async def make(body: dict[str, Any], progress: Optional[Callable[..., None]] = None) -> dict[str, Any]:
    """Automata videó. body: {mode: article|text|new, text?, url?, source_file_id?, format?, duration?,
    images?: [file_id | {url} | {stock} | {generate}], pictures?: stock|mine|none, music_file_id?, voice_file_id?,
    accent?, language?, title?, render?: bool (alap true)}."""
    step = progress or (lambda *a, **k: None)
    mode = (body.get("mode") or "article").lower()
    fmt = fmt_name(body.get("format"))
    if mode == "new":
        return await blank(fmt, body.get("duration"))
    if not available():
        raise RuntimeError("The video engine isn't installed on this server")
    text = (body.get("text") or "").strip()
    url = (body.get("url") or "").strip()
    if not url and re.fullmatch(r"https?://\S+", text):
        url, text = text, ""
    title_hint, og_image, site, doc_pics = (body.get("title") or "").strip(), None, "", []
    step(0, 5, "reading the source")
    if url:
        art = await fetch_article(url)
        text, title_hint, og_image, site = art["text"], title_hint or art["title"], art["image"], art["site"]
    elif body.get("source_file_id"):
        from app.tools import doc_source
        try:
            src = await doc_source.read(body["source_file_id"])
        except ValueError as e:
            raise ValueError(str(e).replace("Slides", "A video").replace("make slides from", "make a video from"))
        text, title_hint = src["text"], title_hint or src["title"]
        doc_pics = [p["file_id"] for p in src.get("pictures") or []]
    trunc_tok, fit_tok = _TRUNC.set([]), _FIT.set([])
    try:
        return await _make_body(body, step, fmt, mode, text, title_hint, og_image, site, doc_pics)
    finally:
        _TRUNC.reset(trunc_tok)
        _FIT.reset(fit_tok)


async def _queries_for(scenes: list[dict[str, Any]]) -> list[str]:
    """Képkereső szavak (angolul, konkrétan) jelenetekhez, amelyekhez a kliens nem adott képet — a magyar cím a
    képtárban mellényúl. Egy modellhívás; hiba esetén a cím."""
    from app.llm.siliconflow import _call_model, _strip_json_fences
    lines = "\n".join(f"{i + 1}. {sc.get('headline')} — {sc.get('caption')}" for i, sc in enumerate(scenes))
    prompt = ("For each scene give a stock-photo search of 2–4 plain, concrete English words (\"wind turbines at sunset\"). "
              f"Reply with json: {{\"queries\": [str, …]}} — exactly {len(scenes)} items, in order.\n{lines}")
    for model in (settings.LLM_MODEL_PRIMARY, settings.LLM_MODEL_FALLBACK):
        try:
            q = json.loads(_strip_json_fences(await _call_model(model, [{"role": "user", "content": prompt}], thinking=False)))["queries"]
            if isinstance(q, list) and len(q) == len(scenes):
                return [_clean(x, 60) for x in q]
        except Exception as e:
            log.info("képkereső szavak (%s): %s", model, e)
    return [sc.get("headline") or "" for sc in scenes]


async def _brand_accent(body: dict[str, Any]) -> Optional[str]:
    """A videó akcentusszíne: a megadott, különben a felhasználó (alapértelmezett / megnevezett) arculatáé; brand="none" = ki."""
    if body.get("accent"):
        return body["accent"]
    b = body.get("brand")
    if str(b or "").strip().lower() in ("none", "off", "no"):
        return None
    try:
        from app.brands import user as user_brands
        kit = await user_brands.resolve(b)
    except Exception as e:
        log.info("videó-arculat: %s", e)
        return None
    cols = (kit or {}).get("colors") or {}
    return cols.get("accent") or cols.get("primary")


async def _make_body(body: dict[str, Any], step: Callable[..., None], fmt: str, mode: str, text: str, title_hint: str,
                     og_image: Optional[str], site: str, doc_pics: list[str]) -> dict[str, Any]:
    pictures = (body.get("pictures") or "stock").lower()
    if pictures == "ai":
        from app import tenancy
        if tenancy.client_makes_images.get():
            raise ValueError(tenancy.CLIENT_IMAGE_MSG)
    body = {**body, "accent": await _brand_accent(body)}
    pre_v: list[str] = []
    pre_f: list[str] = []
    alpha_kinds(body.get("alpha"))           # a hibás érték még a munka előtt derüljön ki
    duration_mode(body.get("duration_mode"), body.get("duration"))
    lnotes: list[str] = []
    layers = layers_from(body.get("layers"), text, lnotes)      # időkódos rétegek a fő idővonalon (a forrás számaival)
    if body.get("scenes"):                  # a kliens jelenetterve: nem tervezünk újra, pontosan ez lesz
        return await _make_from_scenes(body, step, fmt, text, title_hint, site, pictures, layers, lnotes)
    if len(text) < 20:
        raise ValueError("Give the video something to say — an article link, a document, or a few sentences of your own.")
    keep = mode == "text"
    dur = body.get("duration")
    total = None if (keep and not dur) else float(min(max(float(dur or 60), MIN_TOTAL), 180))
    # saját szövegnél is a hossz szerint bontunk (10-12: egy bekezdésként jött 60 s → 1 jelenet, a szöveg nagy része
    # kimaradt); a bekezdések száma csak alsó határ — a modell a mondatokat gondolatonként csoportosítja
    by_len = max(2, min(12, round(((total or 60) - 7) / 6.5)))
    # saját szöveg időtartam nélkül: legalább két mondatonként egy jelenet (10-18: 4 mondat → 1 jelenet → 2 kimaradt)
    n = max(len(_paragraphs(text)), -(-len(_sentences(text)) // 2), by_len if (total or len(text) > 400) else 1) if keep else by_len
    n = max(1, min(n, 14))

    step(1, 5, "writing the script")
    story = await plan(text, n=n, keep_wording=keep, title_hint=title_hint, language=body.get("language"))
    story["title"] = _clean(story.get("title") or title_hint or "Video", 70, "Title")
    story["kicker"] = _clean(story.get("kicker") or site, 28, "Kicker")
    chart_notes: list[str] = []
    # minden jelenet egy úton: szöveghatárok, ábra a forrás számaival (különben kimarad — a modell nem talál ki számot)
    story["scenes"] = _with_ids([_scene_from(s, text, chart_notes) for s in story.get("scenes") or [] if isinstance(s, dict)][:14])
    if keep and story["scenes"]:              # a saját szöveg MINDEN mondata a feliratokban (a modell néha elhagy)
        added = _restore_wording(text, story["scenes"])
        if added:
            story["scenes"] = _with_ids(story["scenes"])
            pre_v.append(f"Your text in full: {added} sentence(s) the script writer left out were put back in their place.")
    if body.get("data_file_id"):           # táblázatból: a cellák értékei, a 2. jelenettől beszúrva
        for k, ch in enumerate(charts_from_table(body["data_file_id"])):
            story["scenes"].insert(min(len(story["scenes"]), 1 + 2 * k),
                                   {"headline": ch["title"], "caption": "", "image_query": story["scenes"][0]["image_query"],
                                    "stat": None, "chart": ch})
            chart_notes.append(f"Chart “{ch['title']}”: {len(ch['items'])} values from the table, cells {ch['cells']}.")
    if not story["scenes"]:
        raise ValueError("Couldn't plan any scenes from that text.")
    story["end"] = {"title": _clean((story.get("end") or {}).get("title"), 60, "End title"),
                    "subtitle": _clean((story.get("end") or {}).get("subtitle"), 90, "End subtitle")}

    step(2, 5, "finding pictures")
    mine = [s for s in (body.get("images") or []) if s] + doc_pics
    specs: list[Any] = []
    for i in range(len(story["scenes"]) + 1):       # 0 = címkép, 1… = jelenetek
        if pictures == "none":
            specs.append(None)
        elif mine and (i < len(mine) or pictures == "mine"):
            specs.append(mine[i % len(mine)])
        elif i == 0 and og_image:
            specs.append({"url": og_image, "credit": site or None})
        elif pictures == "ai":                     # AI-fotó a jelenet keresőszavaiból (webes felület, kreditért)
            specs.append(_ai_photo((_clean(story.get("image_query"), 60) or story["scenes"][-1]["image_query"]) if i == 0
                                   else story["scenes"][i - 1]["image_query"] or story["scenes"][i - 1]["headline"]))
        elif pictures != "mine":
            # a címkártya saját képet kap (ne az 1. jelenetét); ha a terv nem adott, az utolsó jelenet keresőszava
            q = (_clean(story.get("image_query"), 60) or story["scenes"][-1]["image_query"]) if i == 0 \
                else story["scenes"][i - 1]["image_query"]
            specs.append({"stock": q or story["title"]})
        else:
            specs.append(None)

    # a kliens képei hova kerültek (10-12, ChatGPT: „a képek sorrendje nem volt egyértelmű”)
    if body.get("images") and pictures != "none":
        places = ["title card"] + [f"scene {sc['id']}" for sc in story["scenes"]]
        n_img = len([x for x in body.get("images") or [] if x])
        pre_v.append("Your pictures: " + ", ".join(f"#{k + 1} → {places[k]}" for k in range(min(n_img, len(places))))
                     + (" (then repeated)" if pictures == "mine" and n_img < len(places) else ""))
        if n_img > len(places):
            pre_f.append(f"Pictures #{len(places) + 1}–#{n_img} were not used: the video has {len(places)} places "
                         "(title card + scenes). Use scenes=[…] with an image per scene, or fewer pictures.")

    # a stock-keresések jelöltjeiből a látómodell választ (és egy kép nem kerül két jelenetbe)
    await _pick_stock(specs, lambda i: story["title"] if i == 0 else story["scenes"][i - 1]["headline"])

    story["layers"] = layers
    opts = {"fmt": fmt, "total": total, "accent": body.get("accent"), "ken_burns": body.get("ken_burns"),
            "voice_file_id": body.get("voice_file_id"), "music_file_id": body.get("music_file_id"), "music": body.get("music"),
            "keep_wording": keep,
            "narration": body.get("narration"), "narration_voice": body.get("narration_voice"), "language": body.get("language"),
            "requested": float(dur) if dur else None, "duration_mode": body.get("duration_mode"), "alpha": body.get("alpha"),
            "layer_source": text, "scrim": scrim_of(body.get("scrim"))}
    return await _assemble(story, specs, opts, chart_notes=chart_notes, source_text=text, draft=bool(body.get("draft")),
                           source_file_id=body.get("source_file_id"), step=step, render=body.get("render", True) is not False,
                           pre_verified=pre_v, pre_findings=pre_f, layer_notes=lnotes)


async def _make_from_scenes(body: dict[str, Any], step: Callable[..., None], fmt: str, text: str, title_hint: str,
                            site: str, pictures: str, layers: Optional[list[dict[str, Any]]] = None,
                            lnotes: Optional[list[str]] = None) -> dict[str, Any]:
    """A kliens kész jelenetterve (scenes): szöveg, kép, ábra, hossz jelenetenként — pontosan ez készül el.
    Az ábrák számai a forrásban (text), vagy ha az nincs, a jelenetek saját szövegében legyenek."""
    raw = [s for s in body.get("scenes") or [] if isinstance(s, dict)][:20]
    if not raw:
        raise ValueError("scenes: a list of {headline, caption, image?, image_query?, stat?, chart?, duration?}")
    source = text or "\n".join(f"{s.get('headline') or ''}. {s.get('caption') or ''}" for s in raw)
    chart_notes: list[str] = []
    scenes = _with_ids([_scene_from(s, source, chart_notes) for s in raw])
    story = {"title": _clean(body.get("title") or title_hint or scenes[0].get("headline") or "Video", 70, "Title"),
             "kicker": _clean(body.get("kicker") or site, 28, "Kicker"), "image_query": "", "scenes": scenes,
             "end": {"title": _clean(body.get("end_title") or "", 60, "End title"),
                     "subtitle": _clean(body.get("end_subtitle") or "", 90, "End subtitle")}, "layers": layers or []}
    pre_v = [f"Your scene plan: {len(scenes)} scenes ({', '.join(sc['id'] for sc in scenes)}), used as given."]
    step(2, 5, "finding pictures")
    specs: list[Any] = [None] * (len(scenes) + 1)
    # a saját hátteret („background”) kérő jelenet kép nélkül marad — annak nem kell keresőszó sem
    need = [k for k, s in enumerate(raw) if not (s.get("image") or s.get("images") or s.get("background"))
            and pictures not in ("none", "ai")]
    pre_f: list[str] = []
    if need and any(not scenes[k].get("image_query") for k in need):
        from app import tenancy
        if tenancy.client_makes_images.get():
            # MCP: a kereső szavakat a kliens is meg tudja adni — a szerver modellje nem fut rá (nem AI-lépés)
            miss = [scenes[k]["id"] for k in need if not scenes[k].get("image_query")]
            pre_f.append(f"Scene{'s' if len(miss) > 1 else ''} {', '.join(miss)} had no image and no image_query — the stock "
                         "search used the headline. Give each such scene \"image_query\" (2–4 plain English words) or its own "
                         "image for better pictures.")
        else:
            qs = await _queries_for([scenes[k] for k in need])
            for k, q in zip(need, qs):
                scenes[k]["image_query"] = scenes[k].get("image_query") or q
    for k, s in enumerate(raw):
        if (s.get("images") or s.get("image")) and pictures != "none":
            specs[k + 1] = _img_spec(s.get("images") or s["image"])        # a lista: a jelenet hátterei egymás után
        elif s.get("background"):
            specs[k + 1] = None
        elif pictures == "ai":
            specs[k + 1] = _ai_photo(scenes[k].get("image_query") or scenes[k]["headline"])
        elif pictures not in ("none", "mine"):
            specs[k + 1] = {"stock": scenes[k].get("image_query") or scenes[k]["headline"]}
    if body.get("title_image") and pictures != "none":
        specs[0] = _img_spec(body["title_image"])
    elif pictures == "ai":
        specs[0] = _ai_photo(scenes[0].get("image_query") or story["title"])
    elif pictures != "none":
        # a címkártya stockképe csak értelmes keresőszóból: a puszta címre (10-17: „Éles próba”) egy idegen ember
        # fotója jött — ha egyik jelenetnek sincs képe és keresőszava, a címkártya is kép nélküli háttér
        query = next((sc.get("image_query") for sc in scenes if sc.get("image_query")), "")
        has_pic = any(specs[1:])
        specs[0] = _slots(specs[1])[0] if specs[1] and not (isinstance(specs[1], dict) and specs[1].get("stock")) else \
            ({"stock": query or story["title"]} if (query or has_pic) else None)
    kind_of = lambda k: ("yours" if (raw[k].get("image") or raw[k].get("images")) else "AI photo" if pictures == "ai"   # noqa: E731
                         else "stock" if specs[k + 1] else "none") + (f" ×{len(specs[k + 1])}" if isinstance(specs[k + 1], list) else "")
    pre_v.append("Pictures: title card ← " + ("your title_image" if body.get("title_image") else "AI photo" if pictures == "ai"
                                              else "none" if specs[0] is None else "scene s1's picture" if specs[0] is _slots(specs[1])[0]
                                              else "stock")
                 + "; " + ", ".join(f"{sc['id']} ← {kind_of(k)}" for k, sc in enumerate(scenes)))
    await _pick_stock(specs, lambda i: story["title"] if i == 0 else scenes[i - 1]["headline"])
    dur = body.get("duration")
    total = float(min(max(float(dur), MIN_TOTAL), 180)) if dur else None
    opts = {"fmt": fmt, "total": total, "accent": body.get("accent"), "ken_burns": body.get("ken_burns"),
            "voice_file_id": body.get("voice_file_id"), "music_file_id": body.get("music_file_id"), "music": body.get("music"),
            "keep_wording": False,
            "narration": body.get("narration"), "narration_voice": body.get("narration_voice"), "language": body.get("language"),
            "requested": float(dur) if dur else None, "duration_mode": body.get("duration_mode"), "alpha": body.get("alpha"),
            "layer_source": text, "scrim": scrim_of(body.get("scrim"))}
    return await _assemble(story, specs, opts, chart_notes=chart_notes, source_text=text or source, draft=bool(body.get("draft")),
                           source_file_id=body.get("source_file_id"), step=step, render=body.get("render", True) is not False,
                           pre_verified=pre_v, pre_findings=pre_f, scene_plan=True, layer_notes=lnotes)


def _script_report(out: str, source: str, layer_ids: Optional[set[str]] = None) -> list[str]:
    """A build kimenetéből: a szkriptes jelenetek és rétegek hibái, és a szövegeikben a forrásban nem szereplő számok."""
    nums, found = _text_numbers(source), []
    who = lambda sid: f"Layer {sid}" if sid in (layer_ids or set()) else f"Scene {sid}"
    for ln in out.splitlines():
        if ln.startswith("SCRIPT_ERROR "):
            sid, _, msg = ln[13:].partition(": ")
            found.append(f"{who(sid)}: the script failed — {msg[:300]}. "
                         + ("The layer shows nothing; fix its code (chart.code) with action=edit, changes.layers."
                            if sid in (layer_ids or set()) else
                            "The scene shows only its picture and title; fix the code (chart.code) with action=edit."))
        elif ln.startswith("SCRIPT_TEXT "):
            sid, _, raw = ln[12:].partition(" ")
            try:
                texts = [str(t) for t in json.loads(raw)]
            except ValueError:
                continue
            bad = sorted({t for t in texts for x in _text_numbers(t) if x not in nums})
            if bad and (source.strip() or sid not in (layer_ids or set())):       # forrás nélküli réteg: nem ellenőrzött
                found.append(f"{who(sid)}: the script shows numbers that aren't in the source: "
                             + "; ".join(f"“{b[:60]}”" for b in bad[:4]) + ".")
    return found


def _probe_texts(els: list[dict[str, Any]], holder: str) -> None:
    """A saját grafika szövegei mérésre jelölve: a builder TEXT_BOX sort ír róluk (a kirajzolt tinta doboza)."""
    for e in els:
        if e.get("k") == "text" and e.get("gfx") is not None:
            e["probe"] = f"{holder}#{e['gfx']}"


def _holders(D: dict[str, Any]) -> list[tuple[str, tuple[float, float], tuple[float, float], float, float, list[dict[str, Any]]]]:
    """A builder-csomag (D) grafikái: [(név, eltolás a képkockán, időszak, szélesség, magasság, elemek)] — a rétegek a
    saját kompozíciójukkal, a jelenetek a teljes képkockán (a jelenetek egymással nem „egyszerre látszók”)."""
    out = [(f"Layer {ly['id']}", (ly["pos"][0] - ly["w"] / 2, ly["pos"][1] - ly["h"] / 2),
            (float(ly["at"]), float(ly["at"]) + float(ly["dur"])), ly["w"], ly["h"], ly.get("elements") or [])
           for ly in D.get("layers") or []]
    return out + [("", (0.0, 0.0), (-9.0 - k, -9.0 - k), D["W"], D["H"], (sc.get("chart") or {}).get("elements") or [])
                  for k, sc in enumerate(D.get("scenes") or [])]


def _measured_texts(out: str, D: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[int, list]]:
    """A build kimenetének TEXT_BOX sorai → a mért szövegek (tinta a képkockán) és a saját téglalapok rétegenként."""
    boxes: dict[str, tuple[float, float, float, float]] = {}
    for ln in out.splitlines():
        if ln.startswith("TEXT_BOX "):
            try:
                key, l_, t_, w_, h_ = json.loads(ln[9:])
                boxes[str(key)] = (float(l_), float(t_), float(w_), float(h_))
            except (ValueError, TypeError):
                continue
    texts: list[dict[str, Any]] = []
    rects: dict[int, list] = {}
    if not boxes:
        return texts, rects
    for hi, (name, (ox, oy), span, Wh, Hh, els) in enumerate(_holders(D)):
        for e in els:
            if e.get("k") == "rect" and e.get("gfx") is not None:
                cx, cy, sw, sh = ox + e["pos"][0], oy + e["pos"][1], e["size"][0] / 2, e["size"][1] / 2
                rects.setdefault(hi, []).append((e["gfx"] + 1, (cx - sw, cy - sh, cx + sw, cy + sh), e["size"][1] / Hh))
            elif e.get("k") == "text" and e.get("probe") in boxes:
                l_, t_, w_, h_ = boxes[e["probe"]]
                x0, y0, fs = ox + e["pos"][0] + l_, oy + e["pos"][1] + t_, float(e["size"])
                lead = float(e.get("leading") or fs * 1.2)
                texts.append({"h": hi, "who": name or "Scene " + e["probe"].split(":", 1)[-1].rsplit("#", 1)[0], "n": e["gfx"] + 1,
                              "e": e, "ink": (x0, y0, x0 + w_, y0 + h_), "fs": fs, "k": min(Wh, Hh) / 1080, "W": Wh, "H": Hh,
                              "lines": 1 + max(0, round((h_ - 0.95 * fs) / lead)), "span": span, "o": (ox, oy)})
    return texts, rects


def _text_report(out: str, D: dict[str, Any]) -> list[str]:
    """A renderer MÉRT szövegdobozaiból (TEXT_BOX: a kirajzolt tinta a szövegréteg horgonyához képest) a saját grafikák
    szöveghibái: két szöveg egymáson (egy rétegben, vagy két egyszerre látszó rétegben), kilóg a befoglaló téglalapjából
    (rect) vagy a képkockából, w nélkül váratlanul több sorra tört. A finding megnevezi a réteget (jelenetet) és az elemet,
    és javasol w / size (/ y) értéket. (10-19: a cím két sorba tört és az alcímre csúszott — a check mégis „ok” volt.)"""
    texts, rects = _measured_texts(out, D)
    q = lambda s: "“" + re.sub(r"\s+", " ", s)[:40] + ("…" if len(s) > 40 else "") + "”"       # noqa: E731
    over = lambda a, b: (min(a[2], b[2]) - max(a[0], b[0]), min(a[3], b[3]) - max(a[1], b[1]))  # noqa: E731
    found = []
    for i, t in enumerate(texts):
        e, fs, (x0, y0, x1, y1) = t["e"], t["fs"], t["ink"]
        st = str(e.get("str") or "")
        wrapped = t["lines"] > st.count("\n") + 1
        tol = max(2.0, 0.06 * fs)
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        box = min((r for r in rects.get(t["h"], []) if r[1][0] <= cx <= r[1][2] and r[1][1] <= cy <= r[1][3]),
                  key=lambda r: (r[1][2] - r[1][0]) * (r[1][3] - r[1][1]), default=None)
        issues, move = [], None
        if box and (x0 < box[1][0] - tol or x1 > box[1][2] + tol or y0 < box[1][1] - tol or y1 > box[1][3] + tol):
            issues.append(f"runs out of element {box[0]} (a rect {box[2]:.2g} of the frame high)")
        elif not box and (x0 < -tol or y0 < -tol or x1 > D["W"] + tol or y1 > D["H"] + tol):
            issues.append("runs off the frame")
        for j, u in enumerate(texts):
            if j == i or (u["h"] != t["h"] and min(t["span"][1], u["span"][1]) - max(t["span"][0], u["span"][0]) < 0.3):
                continue
            iw, ih = over(t["ink"], u["ink"])
            u_wrapped = u["lines"] > str(u["e"].get("str") or "").count("\n") + 1
            if iw > 2 and ih > max(2.0, 0.1 * min(fs, u["fs"])) and ((wrapped and not u_wrapped) or (wrapped == u_wrapped and i < j)):
                where = "" if u["h"] == t["h"] else f" of {u['who']}"
                issues.append(f"runs into element {u['n']}{where} (text {q(str(u['e'].get('str') or ''))})")
                if u["ink"][1] >= y0 and u["h"] == t["h"] and not move:            # a lejjebb kezdődőt még lejjebb
                    dy = y1 + 0.25 * u["fs"] - u["ink"][1]
                    move = (u["n"], (u["e"]["pos"][1] - u["fs"] * 1.15 + dy) / u["H"])
        if not issues and not (wrapped and e.get("auto_w")):
            continue
        # javaslat a réteg / jelenet saját egységeiben: w = a szélesség 0–1-e, size = px a rövidebb oldal 1080-ára vetítve
        one = _text_px(st, fs, e.get("font") or "semibold") * 1.04
        avail = (box[1][2] - box[1][0]) - 0.4 * fs if box else 0.92 * t["W"]
        size_u = fs / t["k"]
        if wrapped and one + 0.01 * t["W"] <= avail:
            fix = f"give it \"w\": {math.ceil((one + 0.01 * t['W']) / t['W'] * 100) / 100:.2f} (one line) or a smaller \"size\": " \
                  f"{max(8, math.floor(size_u * e['box'][0] / one))}"
        elif move and not (wrapped and e.get("auto_w")):        # a kért (w-s) tördelés marad: a másik szöveg kerüljön lejjebb
            fix = f"move element {move[0]} lower (\"y\": {math.ceil(move[1] * 1000) / 1000:.3f}) or give this one a smaller \"size\""
        elif wrapped or one > avail:
            fix = f"give it a smaller \"size\": {max(8, math.floor(size_u * avail / one))} (one line)"
        elif box and y1 - y0 > (box[1][3] - box[1][1]) * 0.9:
            fix = f"give it a smaller \"size\": {max(8, math.floor(size_u * (box[1][3] - box[1][1]) * 0.8 / (y1 - y0)))} or a taller rect"
        else:
            fix = "move it (x, y) or give a smaller \"size\""
        what = ([f"broke into {t['lines']} lines"] if wrapped else []) + issues
        found.append(f"{t['who']}, element {t['n']} (text {q(st)}): {' and '.join(what)} — {fix}.")
    return found


async def _lottie_report(D: dict[str, Any], work: Path, out: str, project: str = "project.ecproj") -> tuple[list[str], list[str]]:
    """A check MINTÁT vesz a Lottie-elemekből: a kész projektben a Lottie saját kompozíciója három pillanatban, átlátszó
    háttérrel → mennyit fed a dobozából, és rákerül-e egy (mért) szövegre. (10-19: a „confetti” lila háromszögfoltként
    ült a szövegen, a check nem látta.) → (verified, findings)"""
    import numpy as np
    from PIL import Image
    texts, _ = _measured_texts(out, D)
    verified, findings = [], []
    for hi, (name, (ox, oy), span, Wh, Hh, els) in enumerate(_holders(D)):
        for e in els:
            if e.get("k") != "lottie" or not e.get("comp") or not e.get("file"):
                continue
            try:
                lj = json.loads((work / e["file"]).read_text("utf-8"))
                lw, lh, fr = float(lj.get("w") or 1), float(lj.get("h") or 1), float(lj.get("fr") or 30)
                ip, op = float(lj.get("ip") or 0), float(lj.get("op") or fr)
                s_ = (max if e.get("fit") == "cover" else min)(e["size"][0] / lw, e["size"][1] / lh)
                rx0, ry0 = ox + e["pos"][0] - lw * s_ / 2, oy + e["pos"][1] - lh * s_ / 2
                covs, union = [], None
                for f in (0.25, 0.5, 0.75):
                    t = (ip + (op - ip) * f) / fr
                    png = work / f"lottie-sample-{hi}-{e.get('gfx', 0)}-{int(f * 100)}.png"
                    await _run("render-frame", "--project", project, "--comp", e["comp"], "--time", f"{t:.3f}", "--max-side", "240",
                               "--transparent", "--out", png.name, cwd=work, timeout=120)
                    a_ = np.asarray(Image.open(png).convert("RGBA"))[:, :, 3] > 16
                    png.unlink(missing_ok=True)
                    covs.append(float(a_.mean()))
                    if a_.any():
                        ys, xs = np.nonzero(a_)
                        kx, ky = lw * s_ / a_.shape[1], lh * s_ / a_.shape[0]
                        bb = (rx0 + xs.min() * kx, ry0 + ys.min() * ky, rx0 + (xs.max() + 1) * kx, ry0 + (ys.max() + 1) * ky)
                        union = bb if union is None else (min(union[0], bb[0]), min(union[1], bb[1]), max(union[2], bb[2]), max(union[3], bb[3]))
            except Exception as ex:
                log.info("lottie-minta kimaradt (%s): %s", e.get("comp"), ex)
                continue
            who = (name or "Scene") + f", element {int(e.get('gfx', 0)) + 1}"
            title = re.sub(r"\s*\(LottieFiles\)$", "", str(e.get("credit") or e["comp"]))
            hits = [t for t in texts if union and (t["h"] == hi or min(t["span"][1], span[1]) - max(t["span"][0], span[0]) >= 0.3)
                    and (lambda w_, h_: w_ > 0 and h_ > 0 and w_ * h_ > 0.15 * (t["ink"][2] - t["ink"][0]) * (t["ink"][3] - t["ink"][1]))(
                        min(union[2], t["ink"][2]) - max(union[0], t["ink"][0]), min(union[3], t["ink"][3]) - max(union[1], t["ink"][1]))]
            lo, hi_ = min(covs) * 100, max(covs) * 100
            pct = f"{lo:.0f}%" if round(lo) == round(hi_) else f"{lo:.0f}–{hi_:.0f}%"
            other = f'(lottie {{"search": …, "pick": {2 if e.get("pick", 1) == 1 else 1}}}, or video action="lottie" to look)'
            if min(covs) > 0.85:
                findings.append(f"{who} ({title}): the animation fills {pct} of its box at every sampled moment — it has an "
                                f"opaque background and hides what is under it. Choose another one {other}, or make its box smaller.")
            if hits:
                findings.append(f"{who} ({title}): it sits on " + ", ".join(f"element {t['n']}" + ("" if t["h"] == hi else f" of {t['who']}")
                                + f" (text “{str(t['e'].get('str') or '')[:30]}”)" for t in hits[:3])
                                + f" — move it (x, y), make it smaller (w, h), or choose another one {other}.")
            if min(covs) <= 0.85 and not hits:
                verified.append(f"{who} ({title}): sampled at 3 moments — it draws on {pct} of its canvas"
                                + (", clear of the texts." if texts else "."))
    return verified, findings


def _gfx_failures(story: dict[str, Any]) -> list[str]:
    """A be nem töltött képek / Lottie-k / bevágások (a grafika többi része elkészült) — findings, nem csak napló."""
    out = []
    for kind, holders in (("Scene", story.get("scenes") or []), ("Layer", story.get("layers") or [])):
        for k, h in enumerate(holders):
            for i, e in enumerate((h.get("chart") or {}).get("elements") or []):
                if e.get("_failed"):
                    out.append(f"{kind} {h.get('id') or k + 1}, element {i + 1} ({e.get('shape')}): it couldn't be loaded "
                               f"({e['_failed']}) — it was left out. Give another file_id / URL, or {{\"search\": …}} for a Lottie.")
    return out


def _lottie_picks(story: dict[str, Any]) -> list[dict[str, Any]]:
    """A kereséssel választott Lottie-k (név, előnézet, hányadik a sorrendben, a többi jelölt) a válaszba."""
    out = []
    for kind, holders in (("Scene", story.get("scenes") or []), ("Layer", story.get("layers") or [])):
        for k, h in enumerate(holders):
            for i, e in enumerate((h.get("chart") or {}).get("elements") or []):
                if e.get("_pick"):
                    out.append({"where": f"{kind} {h.get('id') or k + 1}, element {i + 1}", **e["_pick"]})
    return out


def _without_runtime(story: dict[str, Any]) -> dict[str, Any]:
    """A forgatókönyv mentéshez: a futásidejű kulcsok (_groups, _skip, _credit…) nélkül."""
    def clean(x: Any) -> Any:
        if isinstance(x, dict):
            return {k: clean(v) for k, v in x.items() if not str(k).startswith("_")}
        if isinstance(x, list):
            return [clean(v) for v in x]
        return x
    return clean(story)


def _gfx_holders(story: dict[str, Any], frame: tuple[int, int, float]) -> list[tuple[str, dict[str, Any], int, int]]:
    """Ahol ábra lehet: a jelenetek (a képkocka méretében) és a rétegek (a saját kompozíciójuk méretében) →
    [(név a hibaüzenethez, a jelenet / réteg, W, H)]."""
    W, H, safe = frame
    out = [(f"Scene {sc.get('id') or f's{i + 1}'}", sc, W, H) for i, sc in enumerate(story.get("scenes") or [])]
    for ly in story.get("layers") or []:
        if ly.get("chart"):
            w, h = layer_size(ly, W, H, safe)
            out.append((f"Layer {ly['id']}", ly, w, h))
    return out


async def _place_on_maps(story: dict[str, Any], frame: tuple[int, int, float]) -> tuple[list[str], list[str]]:
    """A saját grafikák térképei és földrajzi elemei (jelenetek és rétegek): helynév → lat/lon → a képkocka (rétegnél a
    réteg kompozíciója) arányában (x, y, from, to, points); a térkép rétegei az elem "_groups" kulcsába (a mentett
    forgatókönyvbe nem kerül). frame = (W, H, a rétegek alsó határa). → (kreditek, hibák)"""
    from app.tools import geo as geomod
    credits, problems = [], []
    for sid, sc, W, H in _gfx_holders(story, frame):
        ch = sc.get("chart") or {}
        if ch.get("type") != "graphic":
            continue
        els = ch.get("elements") or []
        refs = [e for e in els if e.get("geo") or e.get("geo_from") or e.get("geo_to") or e.get("geo_points")]
        maps = [e for e in els if e.get("shape") == "map"]
        if not refs and not maps:
            continue
        if not maps:
            problems.append(f"{sid}: elements placed by place name need a map element (shape \"map\") — left out.")
            for e in refs:
                e["_skip"] = True
            continue
        coords: dict[str, tuple[float, float]] = {}

        async def where(q: dict[str, Any]) -> Optional[tuple[float, float]]:
            if "lat" in q:
                return q["lat"], q["lon"]
            if q["place"] not in coords:
                try:
                    g = await geomod.geocode(q["place"])
                    coords[q["place"]] = (g["lat"], g["lon"])
                except Exception as ex:
                    problems.append(f"{sid}: “{q['place']}” — {ex}.")
                    coords[q["place"]] = None          # type: ignore[assignment]
            return coords[q["place"]]
        for e in refs:                                   # előbb minden hely: a térkép mindet mutassa
            e["_ll"] = {k: await where(e[k]) for k in ("geo", "geo_from", "geo_to") if e.get(k)}
            if e.get("geo_points"):
                e["_ll"]["points"] = [await where(q) for q in e["geo_points"]]
        pts = [v for e in refs for k, v in e["_ll"].items() if k != "points" and v] + \
              [v for e in refs for v in (e["_ll"].get("points") or []) if v]
        m = maps[0]
        box = (W * (m.get("x", 0.5) - m.get("w", 1.0) / 2), H * (m.get("y", 0.5) - m.get("h", 1.0) / 2), W * m.get("w", 1.0), H * m.get("h", 1.0))
        # a tűk és az útvonal a cím és a felirat közé (10-12: a vasútállomás a felirat alá csúszott); a jobb szélen
        # álló tű felirata balra kerül (_graphic_elements)
        sx0, sy0, sx1, sy1 = max(box[0], W * 0.07), max(box[1], H * 0.17), min(box[0] + box[2], W * 0.93), min(box[1] + box[3], H * 0.76)
        fit = (sx0, sy0, sx1 - sx0, sy1 - sy0) if sx1 - sx0 > W * 0.2 and sy1 - sy0 > H * 0.2 else None
        try:
            built = await geomod.build(m.get("map") or {}, pts, box, W, H, fit=fit)
        except Exception as ex:
            problems.append(f"{sid}: the map could not be made — {ex}.")
            for e in refs + maps:
                e["_skip"] = True
            continue
        m["_groups"], m["_credit"] = built["groups"], built["credit"]
        if built["credit"] not in credits:
            credits.append(built["credit"])
        pr = built["proj"]
        for e in refs:
            ll = e.pop("_ll")
            if e.get("geo") and ll.get("geo"):
                e["x"], e["y"] = pr.frac(*ll["geo"])
            for k in ("from", "to"):
                if e.get("geo_" + k) and ll.get("geo_" + k):
                    e[k] = pr.frac(*ll["geo_" + k])
            if e.get("geo_points"):
                e["points"] = [pr.frac(*v) for v in ll.get("points") or [] if v]
            if (e.get("geo") and not ll.get("geo")) or (e.get("geo_points") and len(e.get("points") or []) < 2) or \
               any(e.get("geo_" + k) and not ll.get("geo_" + k) for k in ("from", "to")):
                e["_skip"] = True
    return credits, problems


async def _materialize_gfx(story: dict[str, Any], work: Path, parent: Optional[str],
                           frame: Optional[tuple[int, int, float]] = None) -> tuple[list[str], list[tuple[str, str]], list[str]]:
    """A saját grafikák képei és Lottie-animációi → a munkamappába. → (asset-nevek, [(lottie-fájl, kompozíció-név)], kreditek).
    A szülőprojekt meglévő elemei ({"asset"}) letöltés nélkül jönnek; a Lottie JSON-t ellenőrizzük és egyedi nevet kap."""
    from app.tools import images as imgs
    parent_assets = assets_of(parent) if parent else {}
    names, lotties, credits = [], [], []
    n = 0
    holders = list(story.get("scenes") or []) + [ly for ly in story.get("layers") or [] if ly.get("chart")]   # jelenetek és rétegek
    sizes = {id(h): (w_, h_) for _, h, w_, h_ in _gfx_holders(story, frame)} if frame else {}      # a doboz aránya (Lottie)
    for si, sc in enumerate(holders):
        ch = sc.get("chart") or {}
        if ch.get("type") != "graphic":
            continue
        for ei, e in enumerate(ch.get("elements") or []):
            if e.get("shape") not in ("image", "lottie"):
                continue
            n += 1
            ext = ".json" if e["shape"] == "lottie" else ".jpg"
            name = e.get("asset") if e.get("asset") and parent_assets.get(e["asset"]) else f"gfx-{n:02d}-{sc.get('id') or si}{ext}"
            try:
                if e.get("asset") and parent_assets.get(e["asset"]):
                    shutil.copyfile(crafts.existing(parent_assets[e["asset"]]), work / name)
                    if e.get("sound_file") and parent_assets.get(e["sound_file"]):     # a bevágás hangja is marad
                        shutil.copyfile(crafts.existing(parent_assets[e["sound_file"]]), work / e["sound_file"])
                        names.append(e["sound_file"])
                    if e.get("credit"):                                    # újraépítés: a kredit marad
                        credits.append(str(e["credit"])[:200])
                elif e["shape"] == "image" and _movie_src(e.get("src")):
                    name = f"gfx-{n:02d}-{sc.get('id') or si}.mp4"
                    src_p = crafts.existing(_movie_src(e["src"]))
                    span = _clip_span(e.get("clip"), sc.get("dur") if sc in (story.get("layers") or []) else None)
                    await _cut_clip(src_p, span, work / name)
                    e["clip_len"] = span[1]
                    e.pop("sound_file", None)
                    if (e.get("sound") or "off") != "off" and sc in (story.get("layers") or []):    # a hangja: a réteg idejére
                        sname = name[:-4] + ".wav"
                        if await _cut_sound(src_p, span, work / sname):
                            e["sound_file"] = sname
                            names.append(sname)
                elif e["shape"] == "image":
                    r = await imgs.resolve(_img_spec(e.get("src")), max_px=1600)
                    name = f"gfx-{n:02d}-{sc.get('id') or si}{r.path.suffix or '.jpg'}"
                    shutil.move(str(r.path), work / name)
                    if r.attribution:
                        credits.append(r.attribution)
                        e["credit"] = r.attribution[:200]
                else:
                    from app.tools import lottie as lot
                    src = e.get("src")
                    sid = src.get("file_id") if isinstance(src, dict) else src
                    if isinstance(src, dict) and src.get("url"):
                        sid = src["url"]
                    if isinstance(src, dict) and src.get("search"):      # a LottieFiles könyvtárából: átlátszó, illő alakú előre
                        hw, hh = sizes.get(id(sc), (16.0, 9.0))
                        box = (e.get("w", 0.3) * hw) / max(1e-6, e.get("h", 0.3) * hh)
                        data, cand = await lot.pick(str(src["search"]), sc.get("headline") or sc.get("text") or ch.get("title") or "",
                                                    aspect=box, n=int(src["pick"]) if str(src.get("pick") or "").isdigit() else None)
                        e["src"] = {"url": cand["json_url"], "name": cand["name"], "thumb": cand.get("thumb_url")}
                        e["credit"] = cand["credit"]
                        e["_pick"] = {"search": str(src["search"]), **{k: cand.get(k) for k in (
                            "name", "author", "thumb_url", "gif_url", "page", "pick", "of", "transparent", "aspect", "alternatives")}}
                        credits.append(cand["credit"])
                    elif isinstance(sid, str) and sid.startswith("http"):
                        data = await lot.fetch(sid)
                        if isinstance(src, dict) and src.get("credit"):
                            e["credit"] = str(src["credit"])[:200]
                            credits.append(e["credit"])
                    else:
                        data = lot.check(json.loads(crafts.existing(str(sid)).read_bytes()[:lot.MAX_BYTES].decode("utf-8")))
                    data["nm"] = f"Lottie {n}"
                    (work / name).write_text(json.dumps(data), "utf-8")
                    lotties.append((name, data["nm"]))
                    e["comp"] = data["nm"]
                e["file"] = e["asset"] = name
                names.append(name)
                if e["shape"] == "lottie" and not e.get("comp"):
                    e["comp"] = json.loads((work / name).read_text("utf-8")).get("nm")
                    lotties.append((name, e["comp"]))
            except Exception as ex:
                log.warning("grafika-elem (%s) nem jött le: %s", e.get("shape"), ex)
                e["file"] = None
                e["_failed"] = re.sub(r"\s+", " ", str(ex))[:160] or type(ex).__name__     # a check megnevezi (_gfx_failures)
    return names, lotties, credits


def _movie_src(src: Any) -> Optional[str]:
    """Egy kép-elem forrása videó-e (a felhasználó egy másik felvétele) → a file_id-ja, különben None."""
    sid = src.get("file_id") if isinstance(src, dict) else src
    return sid if isinstance(sid, str) and not sid.startswith("http") and Path(sid).suffix.lower().lstrip(".") in MOVIE_EXTS else None


def _clip_span(clip: Any, limit: Optional[float] = None) -> tuple[float, float]:
    """A bevágás része → (kezdet, hossz) mp: [tól, ig] (ig nélkül legfeljebb 60 mp), legfeljebb a réteg hossza."""
    a = float(clip[0]) if isinstance(clip, (list, tuple)) and clip else 0.0
    b = float(clip[1]) if isinstance(clip, (list, tuple)) and len(clip) > 1 and clip[1] is not None else a + 60.0
    n = max(0.2, b - a)
    return a, round(min(n, float(limit)) if limit else n, 3)


async def _cut_sound(src: Path, span: tuple[float, float], dst: Path) -> bool:
    """A bevágás hangja külön sávnak (az első dekódolható hangsáv — az iPhone térhangja nélkül); False, ha nincs hangja."""
    from app.convert.engines import run
    k = (await probe_movie(src)).get("audio_stream")
    if k is None:
        return False
    await run("ffmpeg", "-v", "error", "-y", "-ss", f"{span[0]:.3f}", "-i", str(src), "-t", f"{span[1]:.3f}", "-map", f"0:a:{k}",
              "-vn", "-ac", "2", "-ar", "48000", str(dst), timeout=settings.VIDEO_TIMEOUT_S)
    return dst.exists() and dst.stat().st_size > 1000


def _clip_sounds(layers: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """A rétegek bevágásainak hangja a fő idővonalra: [{file, t0, dur, mode}] — replace: közben a felvétel hangja elhalkul."""
    out = []
    for ly in layers:
        for i, e in enumerate((ly.get("chart") or {}).get("elements") or []):
            if e.get("sound_file") and e.get("sound") in ("replace", "mix"):
                t0 = float(ly["at"]) + _el_t0(e, i)
                dur = min(float(e.get("clip_len") or ly["dur"]), float(ly["at"]) + float(ly["dur"]) - t0)
                if dur > 0.1:
                    out.append({"file": e["sound_file"], "t0": round(t0, 3), "dur": round(dur, 3), "mode": e["sound"]})
    return out


def _el_t0(e: dict[str, Any], i: int) -> float:
    """Egy grafika-elem megjelenése (mp a jelenet / réteg elejétől): a megadott, különben sorban; a bevágott videó azonnal."""
    movie = e.get("clip") is not None or Path(str(e.get("file") or "")).suffix.lower().lstrip(".") in MOVIE_EXTS
    return max(0.0, float(e["at"])) if e.get("at") is not None else (0.0 if movie else 0.3 + 0.25 * i)


async def _cut_clip(src: Path, span: tuple[float, float], dst: Path) -> Path:
    """Egy másik felvétel része a grafikába (bevágás): span = (kezdet, hossz) mp; a kép NÉMÁN (a hangja külön sáv:
    _cut_sound), az elforgatás szerint egyenesen, a hosszabb oldal legfeljebb 1920 px."""
    from app.convert.engines import run
    a, n = span
    await run("ffmpeg", "-v", "error", "-y", "-ss", f"{a:.3f}", "-i", str(src), "-t", f"{n:.3f}", "-map", "0:v:0", "-an",
              "-vf", "scale='if(gt(iw,ih),min(1920,iw),-2)':'if(gt(iw,ih),-2,min(1920,ih))'", "-c:v", "libx264",
              "-preset", "veryfast", "-crf", "18", "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(dst),
              timeout=settings.VIDEO_TIMEOUT_S)
    return dst


NARRATION_ASSET = "narration.wav"
_NO_NARRATION = {"none", "off", "false", "no", "0"}


def fit_to_narration(story: dict[str, Any], clips: list) -> dict[str, Any]:
    """A forgatókönyv másolata, amelyben minden jelenet legalább a narrációja + 1,4 s (átmenet és lélegzet)."""
    out = json.loads(json.dumps(story))
    for sc, (_, secs) in zip(out.get("scenes") or [], clips[1:]):
        if secs:
            sc["duration"] = round(max(float(sc.get("duration") or 0), secs + 1.4), 2)
    return out


def fit_to_length(story: dict[str, Any], clips: list, total: float) -> tuple[dict[str, Any], float]:
    """Fix hossz (duration_mode=fixed): minden jelenet legalább a narrációja + 1,4 s (a kliens saját jelenethossza is
    minimum), a kért hosszig maradó idő a többi jelenet közt oszlik meg arányosan (a jelenet tovább áll, a hang az
    elején szól) — így a videó a kért hosszú. → (a forgatókönyv másolata, ennyi másodperccel hosszabb a narráció a
    kértnél; 0, ha belefér)."""
    out = json.loads(json.dumps(story))
    scenes = out.get("scenes") or []
    n = len(scenes)
    if not n:
        return out, 0.0
    Tt, Te = _cards(total)
    room = total - Tt - Te + XFADE * (n + 1)
    lo = _scene_floor(total, n, Tt, Te)

    def needs(pad: float) -> tuple[list[float], list[bool]]:
        need, flex = [], []
        for i, sc in enumerate(scenes):
            secs = clips[i + 1][1] if i + 1 < len(clips) else 0.0
            given = float(sc["duration"]) if isinstance(sc.get("duration"), (int, float)) and sc["duration"] > 0 else None
            need.append(max(given or 0.0, secs + pad if secs else (given or lo)))
            flex.append(given is None)
        return need, flex
    need, flex = needs(1.4)
    if sum(need) > room:                    # szoros fix hossz: a narráció után rövidebb szünet (a hang a jelenet 0,45 s-ánál indul)
        need, flex = needs(0.6)
    spare = room - sum(need)
    if spare > 0:
        grow = [i for i in range(n) if flex[i]] or list(range(n))
        wsum = sum(need[i] for i in grow) or 1.0
        for i in grow:
            need[i] += spare * need[i] / wsum
    for sc, d in zip(scenes, need):
        sc["duration"] = round(d, 3)
    return out, round(max(0.0, -spare), 2)


def narration_segments(D: dict[str, Any], clips: list) -> list[tuple[float, Path]]:
    """A narráció darabjai az idővonalon: a cím a címkártyán (ha belefér), a jelenetek a saját elejükön."""
    segs: list[tuple[float, Path]] = []
    title = clips[0] if clips else (None, 0.0)
    if title[0] and title[1] <= D["title"]["dur"] - 0.8:
        segs.append((0.35, title[0]))
    segs += [(tl["start"] + 0.45, clip) for tl, (clip, _) in zip(D["timeline"][1:-1], clips[1:]) if clip]
    return segs


async def _narrate(story: dict[str, Any], opts: dict[str, Any], work: Path) -> tuple[Optional[str], list, Optional[str]]:
    """A cím és a jelenetek felolvasandó szövege → AI-narráció darabjai (offline Piper-hang, app/tools/narration.py).
    → (hang, [(wav, hossz)…] a címmel kezdve, megjegyzés ha kimaradt). Hiba esetén a videó narráció nélkül készül."""
    from app.tools import narration
    if not narration.available():
        return None, [], "AI narration isn't installed on this server — the video has no voice (give voice_file_id)."
    texts = [story.get("title") or ""] + [sc.get("caption") or sc.get("headline") or "" for sc in story.get("scenes") or []]
    mode = str(opts.get("narration") or "ai")
    want = opts.get("narration_voice") or (mode if "-" in mode and "_" in mode else None)     # szerkesztés: ugyanaz a hang
    lang = narration.lang_code(opts.get("language")) or narration.detect(" ".join(texts))
    try:
        key = await asyncio.to_thread(narration.pick, lang, want)
        clips = await narration.speak_all(texts, key, work, float(opts.get("narration_speed") or 1.0))
    except Exception as e:
        log.warning("AI-narráció kimaradt: %s", e)
        return None, [], f"AI narration was skipped ({str(e)[:160]}) — the video has no voice."
    return key, clips, None


def duration_mode(mode: Any, requested: Optional[float]) -> str:
    """fixed = a kért hossz (a jelenetek kitartanak, a narráció gyorsabban olvas — ésszerű határig) | narration = a hang
    szabja a hosszt. Alapból fixed, ha kértek hosszt, különben narration."""
    m_ = str(mode or "").strip().lower()
    if m_ in ("", "auto", "none"):
        return "fixed" if requested else "narration"
    if m_ not in ("fixed", "narration"):
        raise ValueError(f"duration_mode: fixed (the asked duration) | narration (the voice sets the length) — not {mode!r}.")
    return m_


def _lang_of(text: str) -> str:
    """A számformázás és a kiegészítő címkék nyelve (hu | en): a gyakori szavakból; ha a szöveg ahhoz túl rövid, az ékezet
    dönt. (10-17: egy „Kovács” név miatt az angol ábra „41,2”-t és „Egyéb”-et írt — az ékezet magában nem magyar.)"""
    from app.tools import narration
    from app.tools.infographic import _HU
    t = text or ""
    words = re.findall(r"[^\W\d_]+", t.lower())
    hu = sum(w in narration._STOPSETS["hu"] for w in words) + 2 * sum(t.count(c) for c in "őűŐŰ")
    en = sum(w in narration._STOPSETS["en"] for w in words)
    if hu or en:
        return "hu" if hu > en else "en"
    return "hu" if _HU.search(t) else "en"


def _layer_text_all(story: dict[str, Any]) -> str:
    return " ".join(_own_text({k: v for k, v in ly.items() if k not in ("id", "type", "at", "dur", "position", "anim")})
                    for ly in story.get("layers") or [])


def _kept_layer_notes(spec: dict[str, Any], layers: list[dict[str, Any]]) -> list[str]:
    """A mentett réteg-jegyzetek közül csak a VÁLTOZATLAN rétegeké öröklődik (a módosított réteg újra ellenőrződik)."""
    old = {ly.get("id"): ly for ly in (spec.get("story") or {}).get("layers") or []}
    same = {ly["id"] for ly in layers if old.get(ly["id"]) == ly}
    return [n for n in spec.get("layer_notes") or [] if (m_ := re.match(r"Layer (\S+?):", n)) and m_.group(1) in same]


def _unrendered_check(findings: list[str], verified: Optional[list[str]] = None) -> dict[str, Any]:
    """render=false: a projekt elmentve, render nélkül — a szerkezeti / elrendezési check (szövegek mérete, ütközés,
    kimaradt elem) így is jön; a kép maga nincs ellenőrizve."""
    return {"status": "needs_check" if findings else "ok", "findings": list(dict.fromkeys(findings)),
            "verified": ["Saved the project without rendering (render=false): the layout was checked (text boxes measured by "
                         "the video engine, overlaps, left-out elements) — render it with video(action=\"render\", file_id=…)."]
            + list(verified or [])}


def _split_notes(notes: list[str]) -> tuple[list[str], list[str]]:
    """⚠ = figyelmeztetés (findings), a többi az igazolt (verified)."""
    return [n for n in notes if not n.startswith("⚠ ")], [n[2:] for n in notes if n.startswith("⚠ ")]


async def _build(D: dict[str, Any], work: Path, lotties: list[tuple[str, str]], out_name: str = "project.ecproj") -> str:
    """D + video_builder.jsx → .ecproj (effectcraft-cli script). A Lottie-k előbb beimportálva (külön kompozíció, a
    szkript erre a projektre épít); saját szkriptes ábra (jelenet vagy réteg) → csak zárkában."""
    (work / "build.jsx").write_text("var D = " + json.dumps(D, ensure_ascii=False) + ";\n" + BUILDER_JS.read_text("utf-8"), "utf-8")
    base = ["base.ecproj"] if lotties and (work / "base.ecproj").exists() else []
    if not base:
        for lf, _nm in lotties:
            args = ["exec", "file.importLottie", "--params", json.dumps({"path": lf})]
            args += ["--project", "base.ecproj", "--save"] if base else ["--empty", "--save-as", "base.ecproj"]
            await _run(*args, cwd=work, timeout=120)
            base = ["base.ecproj"]
    scripted = any((sc.get("chart") or {}).get("code") for sc in D.get("scenes") or []) or any(ly.get("code") for ly in D.get("layers") or [])
    async with _render_gate():
        return await _run("script", "build.jsx", *base, "--save-as", out_name, cwd=work, untrusted=scripted,
                          timeout=300 if scripted else None)


def alpha_kinds(alpha: Any) -> list[str]:
    """alpha: true | "mov" → ProRes 4444 .mov · "webm" → WebM VP9 · "both" · false/None → nincs."""
    if alpha in (None, False, "", "false", "no", "none", "off", 0):
        return []
    a = "mov" if alpha is True or str(alpha).lower() in ("true", "yes", "1", "prores", "mov") else str(alpha).lower()
    kinds = {"mov": ["mov"], "webm": ["webm"], "both": ["mov", "webm"]}.get(a)
    if kinds is None:
        raise ValueError(f"alpha: true (ProRes 4444 .mov) | \"webm\" (WebM VP9) | \"both\" — not {alpha!r}.")
    return kinds


async def _alpha_proof(path: Path, at: float) -> dict[str, Any]:
    """Valóban átlátszó-e a kimenet: a pixelformátum (ffprobe; WebM-nél az alpha_mode jelölés) és egy dekódolt kocka
    alfája — hol teljesen átlátszó, hol látszik grafika."""
    from app.convert.engines import run
    info = await probe_movie(path)
    pf = info.get("pix_fmt") or ""
    tagged = bool(re.search(r"yuva|rgba|argb|bgra|abgr|gbrap|^ya", pf)) or info.get("alpha_mode")
    png = path.with_name(path.stem + "-alpha.png")
    clear = shown = None
    try:
        dec = ["-c:v", "libvpx-vp9"] if path.suffix.lower() == ".webm" else []        # a beépített VP9-dekóder eldobja az alfát
        await run("ffmpeg", "-v", "error", *dec, "-ss", f"{max(0.0, at):.2f}", "-i", str(path), "-frames:v", "1",
                  "-vf", "format=rgba", "-y", str(png), timeout=120)
        import numpy as np
        from PIL import Image
        a = np.asarray(Image.open(png).convert("RGBA"))[:, :, 3]
        clear, shown = float((a < 8).mean()), float((a >= 8).mean())
    except Exception as e:
        log.info("alfa-kocka nem olvasható (%s): %s", path.name, e)
    finally:
        png.unlink(missing_ok=True)
    return {"pix_fmt": pf, "codec": info.get("codec"), "alpha_mode": bool(info.get("alpha_mode")), "at": round(at, 2),
            "transparent": None if clear is None else round(clear, 3), "graphics": None if shown is None else round(shown, 3),
            "ok": bool(tagged) and (clear is None or (clear > 0.01 and (shown or 0) > 0.0005))}


async def _render_alpha(D: dict[str, Any], work: Path, lotties: list[tuple[str, str]], pid: str, stem: str,
                        group: Optional[str], kinds: list[str], probe_at: float) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    """Átlátszó, csak grafikás változat: ugyanaz a builder D.alpha-val (a felvétel, a képek, a sötétítés és a hang nélkül)
    → ProRes 4444 .mov és/vagy WebM VP9 alfával. → (fájlok, igazolt, gondok)"""
    await _build({**D, "alpha": True}, work, lotties, "alpha.ecproj")
    files, verified, findings = [], [], []
    for ext in kinds:
        out = work / f"alpha.{ext}"
        args = ["render", "--project", "alpha.ecproj", "--comp", "Video", "--out", out.name, "--channels", "rgba", "--audio", "off", "--json"]
        args += ["--format", "prores", "--prores", "4444"] if ext == "mov" else ["--format", "webm", "--webm-codec", "vp9"]
        async with _render_gate():
            await _run(*args, cwd=work)
        if not out.exists():
            findings.append(f"The transparent .{ext} render produced no file — try again, or use alpha=\"{'webm' if ext == 'mov' else 'mov'}\".")
            continue
        proof = await _alpha_proof(out, probe_at)
        fid = new_file_id(ext)
        shutil.move(str(out), output_path(fid))
        what = "ProRes 4444 .mov" if ext == "mov" else "WebM VP9"
        write_meta(fid, label=f"{stem} – transparent graphics.{ext}", operation="video render", source_file_id=pid,
                   extra={"group": group, "variant": "alpha", "alpha": proof})
        files.append({**_file_entry(fid), "format": what, "pix_fmt": proof["pix_fmt"], "alpha": proof["ok"]})
        seen = (f"; at {proof['at']:.1f} s {proof['transparent'] * 100:.0f}% of the frame is fully transparent, "
                f"{proof['graphics'] * 100:.1f}% shows graphics") if proof["transparent"] is not None else ""
        tag = proof["pix_fmt"] + (", alpha_mode=1" if proof["alpha_mode"] else "")
        if proof["ok"]:
            verified.append(f"Transparent graphics: {what} ({tag}){seen}.")
        else:
            findings.append(f"The transparent {what} may not be transparent ({tag}{seen}). Open it in your editor to check, "
                            f"or try alpha=\"{'webm' if ext == 'mov' else 'mov'}\".")
    return files, verified, findings


def _probe_time(D: dict[str, Any], rows: list[dict[str, Any]]) -> float:
    """Egy pillanat, amikor grafika látszik (az alfa-ellenőrzéshez): az első réteg közepe, különben az első jelenet vége."""
    if rows:
        r_ = rows[0]
        return r_["start"] + min(2.0, (r_["end"] - r_["start"]) / 2)
    tl = D.get("timeline") or []
    return scene_still(tl, 1 if len(tl) > 2 else 0) if tl else D.get("total", 2) / 2


async def _assemble(story: dict[str, Any], specs: list[Any], opts: dict[str, Any], *, chart_notes: list[str],
                    source_text: str, source_file_id: Optional[str], step: Callable[..., None], render: bool = True,
                    parent: Optional[str] = None, label_stem: Optional[str] = None, operation: str = "video",
                    pre_verified: Optional[list[str]] = None, pre_findings: Optional[list[str]] = None,
                    scene_plan: bool = False, layer_notes: Optional[list[str]] = None, draft: bool = False) -> dict[str, Any]:
    """Forgatókönyv + képleírások → .ecproj (+ MP4). A make és az edit közös útja; parent = a szerkesztett projekt
    (a képei és a hangja {"asset": név} hivatkozással, letöltés nélkül jönnek át)."""
    from app.tools import diagrams
    diagrams.prepare(story, opts.get("accent"))       # szereplő → szín: a projekt regisztere (story["actors"])
    fmt, total = opts["fmt"], opts.get("total")
    W, H = FORMATS[fmt]
    requested = opts.get("requested")
    mode = duration_mode(opts.get("duration_mode"), requested)
    kinds = alpha_kinds(opts.get("alpha"))
    work = Path(tempfile.mkdtemp(prefix="video-", dir=str(temp_path("").parent)))
    try:
        names, credit_list, per_credit = await _materialize_images(specs, fmt, work, parent)
        map_credits, gfx_findings = await _place_on_maps(story, (W, H, _caption_top(W, H)))
        # a kért, de le nem töltött kép (pl. a forrás 429-cel elutasította) nem lehet „ok”: megnevezzük a helyét
        gfx_findings = [_lost_picture(i, story) for i, (sp, nm) in enumerate(zip(specs, names)) if sp and not nm] + gfx_findings
        gfx_findings += [f"{len(_slots(sp)) - len(nm)} of the {len(_slots(sp))} pictures for "
                         f"{'the title card' if i == 0 else 'scene ' + str(story['scenes'][i - 1].get('id'))} couldn't be downloaded — "
                         "the others are shown." for i, (sp, nm) in enumerate(zip(specs, names))
                         if isinstance(sp, list) and isinstance(nm, list) and len(nm) < len(sp)]
        gfx_names, lotties, gfx_credits = await _materialize_gfx(story, work, parent, (W, H, _caption_top(W, H)))
        gfx_findings += _gfx_failures(story)
        credit_list = credit_list + [c for c in gfx_credits + map_credits if c not in credit_list]
        voice = music = None
        voice_s = None
        parent_assets = assets_of(parent) if parent else {}
        for key, base in (("voice_file_id", "voice"), ("music_file_id", "music")):
            fid = opts.get(key)
            kept = opts.get(f"{base}_asset")          # szerkesztés: a szülőprojekt hangja marad
            if base == "voice" and kept == NARRATION_ASSET:
                kept = None                           # …kivéve az AI-narrációt: az a (módosult) szövegből újra készül
            if base == "music" and _builtin_music(opts):
                kept = None                           # …és a beépített zenét: az a tárból, az új hosszra készül újra
            if not fid and kept and parent_assets.get(kept):
                shutil.copyfile(crafts.existing(parent_assets[kept]), work / kept)
                if base == "voice":
                    voice = kept
                else:
                    music = kept
                continue
            if fid:
                p = crafts.existing(fid)
                ext = p.suffix.lstrip(".").lower()
                if ext not in AUDIO_EXTS:
                    raise ValueError(f"The {base} file should be audio (mp3, wav, m4a…), not .{ext}")
                shutil.copyfile(p, work / f"{base}.{ext}")
                if base == "voice":
                    voice = f"voice.{ext}"
                else:
                    music = f"music.{ext}"
        if voice:
            voice_s = await _probe_duration(work, voice)
        narr_key, narr_clips, narr_note, speed = None, [], None, 1.0
        if not voice and str(opts.get("narration") or "ai").lower() not in _NO_NARRATION:
            step(3, 5, "recording the AI narration")
            narr_key, narr_clips, narr_note = await _narrate(story, opts, work)
        why = None                                    # miért tér el a hossz a kérttől (a check mondja ki)
        if narr_clips and mode == "fixed" and total:
            story_c, over = fit_to_length(story, narr_clips, total)
            if over > 0.5:                            # a narráció nem fér bele: gyorsabban olvas (legfeljebb 1,2×), újra illeszt
                said = sum(s_ for _, s_ in narr_clips[1:]) or 1.0
                speed = round(min(1.2, said / max(said - over, 1e-6)), 2)
                k2, c2, _n2 = await _narrate(story, {**opts, "narration_speed": speed}, work)
                if c2:
                    narr_key, narr_clips = k2, c2
                    story_c, over = fit_to_length(story, narr_clips, total)
                else:
                    speed = 1.0
                if over > 0.5:
                    why = (f"fixed mode: the narration needs about {over:.0f} s more even read {speed:g}× faster — shorten the "
                           "captions, or ask for a longer duration")
        else:
            story_c = fit_to_narration(story, narr_clips) if narr_clips else story
        c_total, c_voice = total, voice_s
        if voice_s and mode == "fixed" and total:    # saját hang + fix hossz: a kért hossz (ha a hang belefér)
            c_total, c_voice = max(float(total), voice_s + 1.0), None
            if voice_s + 1.0 > total + 0.5:
                why = f"fixed mode: your voice file is {voice_s:.0f} s long"
        # teljes (linkekkel) a projekt metájába; a képernyőn rövid: név + licenc, linkek nélkül
        credits, on_screen = _credit_line(credit_list, story)
        kb = str(opts.get("ken_burns") or "some").lower()
        kb = kb if kb in ("some", "all", "off") else "some"
        focus = {nm: await asyncio.to_thread(_focus, work / nm) for slot in names for nm in _slots(slot) if nm}
        fits = _auto_fits(names, specs, work, W / H)
        D = compose(story_c, fmt=fmt, total=c_total, images=names, accent=opts.get("accent"), credits=on_screen, focus=focus, ken_burns=kb,
                    fits=fits, fill=mode == "fixed" and bool(c_total), scrim=scrim_of(opts.get("scrim")),
                    voice=voice, voice_s=c_voice, music=music)
        if narr_clips:                   # a darabok a jelenetek elejére, egy sávban, a videó teljes hosszára
            from app.tools import narration
            segs = narration_segments(D, narr_clips)
            if segs:
                await narration.track(segs, D["total"], work / NARRATION_ASSET)
                voice = NARRATION_ASSET
                D["audio"].update(voice=voice, musicDb=-20)
                how = ("each scene lasts as long as its narration" if mode != "fixed" else
                       f"read {speed:g}× faster to come close to the asked {total:.0f} s (it doesn't fit — see the findings)" if why else
                       f"fixed length: the scenes hold to fill {total:.0f} s" + (f", read {speed:g}× faster to fit" if speed != 1.0 else ""))
                pre_verified = list(pre_verified or []) + [f"AI narration ({narr_key}): {len(segs)} part(s) read aloud — {how}."]
        if narr_note:
            pre_findings = list(pre_findings or []) + [narr_note]
        music, track, m_v, m_f = await _music_under(opts, work, D["total"], work / voice if voice else None, music,
                                                    story.get("title") or "")
        if track:
            D["audio"].update(music=track, musicDb=0 if track != music else (-20 if voice else -10))
        pre_verified, pre_findings = list(pre_verified or []) + m_v, list(pre_findings or []) + m_f
        # időkódos rétegek a fő idővonalon (abszolút időben, a narráció után kialakult hosszon)
        lang = _lang_of(" ".join([story.get("title") or ""] + [x.get("caption") or "" for x in story.get("scenes") or []]))
        D["layers"], layer_rows, layer_problems = layer_comps(story.get("layers") or [], W, H, D["total"], lang, D["accent"],
                                                              _caption_top(W, H), _scene_zones(D))
        # előre megrajzolt rétegek (képkockánként effekt helyett): sötét átmenet a cím/jelenet alá, elmosott zárókép
        _shade_png(work / "shade-title.png", W, H, 0.28, 0.80, 0.92)
        _shade_png(work / "shade-scene.png", W, H, 0.34, 0.84, 0.90)
        extra = ["shade-title.png", "shade-scene.png"]
        if names and names[0]:
            _blurred(work / _slots(names[0])[0], work / "end-blur.jpg", max(W, H) * 0.018)
            D["end"]["image"] = "end-blur.jpg"
            extra.append("end-blur.jpg")
        D["shade"] = {"title": "shade-title.png", "scene": "shade-scene.png"}
        if _has_3d(D):                                        # three.js-jelenetek: a végleges hosszukban, a build előtt
            step(3, 5, "rendering the 3D scenes")
            three_names, three_notes = await _render_3d(D, work, D["accent"])
            extra += three_names
            pre_verified += [n for n in three_notes if not n.startswith("⚠ ")]
            pre_findings += [n[2:] for n in three_notes if n.startswith("⚠ ")]
        step(3, 5, "building the timeline")
        out = await _build(D, work, lotties)
        layer_problems = layer_problems + _text_report(out, D)   # a saját grafikák mért szövegei (jelenetek és rétegek)
        lot_v, lot_f = await _lottie_report(D, work, out) if lotties else ([], [])   # a Lottie-k mintája
        layer_problems, pre_verified = layer_problems + lot_f, list(pre_verified or []) + lot_v
        scripted = any(((x.get("chart") or {}).get("type") == "script") for x in (story.get("scenes") or []) + (story.get("layers") or []))
        script_findings = _script_report(out, source_text or "", {ly["id"] for ly in story.get("layers") or []}) if scripted else []
        log.info("videó felépítve: %s", out.strip().splitlines()[-2:] if out.strip() else "")
        stem = label_stem or _title_stem(story["title"])
        pid = new_file_id("ecproj")
        shutil.copyfile(work / "project.ecproj", output_path(pid))
        from app.routes.make import _new_group
        group = _new_group()
        info_meta = {"format": fmt, "duration": D["total"], "scenes": len(D["scenes"]), "title": story["title"],
                     "requested_s": requested, "duration_mode": mode, "layers": len(D["layers"])}
        # a forgatókönyv a projekt mellett: ebből szerkeszt az ügynök (edit) — képek a projekt saját médiájára hivatkozva
        def _src(sp: Any) -> Any:
            sp = sp.get("src") if isinstance(sp, dict) and sp.get("asset") else sp
            return sp if isinstance(sp, (str, dict)) else None
        def _pic(nm: Any, cr: Any, sp: Any) -> Any:
            if isinstance(nm, list):              # képsor: elemenként (a forrás a listából, ha ugyanannyi)
                sps = _slots(sp) if isinstance(sp, list) and len(sp) == len(nm) else [None] * len(nm)
                return [{"asset": a, "credit": c, "src": _src(x)} for a, c, x in zip(nm, cr, sps)]
            return {"asset": nm, "credit": cr, "src": _src(sp)} if nm else None
        spec = {"story": _without_runtime(story), "pictures": [_pic(nm, cr, sp) for nm, cr, sp in zip(names, per_credit, specs)],
                "opts": {"fmt": fmt, "total": total, "accent": opts.get("accent"), "ken_burns": kb, "voice_asset": voice,
                         "music_asset": music, "music": opts.get("music") if _builtin_music(opts) else None,
                         "keep_wording": bool(opts.get("keep_wording")),
                         "narration": narr_key or ("none" if str(opts.get("narration") or "").lower() in _NO_NARRATION else None),
                         "narration_voice": opts.get("narration_voice"), "language": opts.get("language"),
                         "requested": requested, "duration_mode": opts.get("duration_mode"), "alpha": opts.get("alpha"),
                         "scrim": opts.get("scrim")},
                "source": (source_text or "")[:30000], "chart_notes": [n for n in chart_notes if not n.startswith("⚠ ")],
                "layer_notes": [n for n in layer_notes or [] if not n.startswith("⚠ ")],
                "layer_source": (opts.get("layer_source") or "")[:30000],
                "timeline": D.get("timeline"), "layer_timeline": layer_rows}
        write_meta(pid, label=f"{stem}.ecproj", operation=operation, source_file_id=source_file_id,
                   extra={"video": info_meta, "group": group, "credit": credits or None, "spec": spec})
        for name in [x for slot in names for x in _slots(slot) if x] + list(dict.fromkeys(x for x in (voice, music, track) if x)) + extra + gfx_names:
            add_asset(pid, name, work / name)
        result: dict[str, Any] = {"file_id": pid, "project": _file_entry(pid), "title": story["title"], "format": fmt,
                                  "duration": {"requested": requested, "actual": D["total"], "mode": mode},
                                  "scenes": len(D["scenes"]), "captions": len(D["captions"]),
                                  "credits": credits or None, "story": _without_runtime(story)}
        if _lottie_picks(story):
            result["lottie_picks"] = _lottie_picks(story)
        if D["layers"] or story.get("layers"):
            result["layers"] = layer_listing([ly for ly in story.get("layers") or [] if ly["id"] in {r_["id"] for r_ in layer_rows}])
        if render:
            step(4, 5, "rendering the video")
            vid = await _render_staged(pid, work / "project.ecproj", stem, group=group, want_s=D["total"], requested=requested,
                                       mode=mode, why=why, draft=draft)
            result["duration"] = vid.pop("duration", None) or result["duration"]
            result["check"] = vid.pop("check", None)          # a közös check-szerződés: a feladat (a videó) ellenőrzése
            chk = result["check"]
            if chk is not None:
                chk["verified"] = list(pre_verified or []) + chk.get("verified", [])
                trunc = _TRUNC.get() or []
                cov_v, cov_f, cov = _coverage(source_text, story, bool(opts.get("keep_wording")), scene_plan)
                result["coverage"] = cov
                chk["verified"] += cov_v
                fitted = _FIT.get() or []
                if fitted:
                    chk["verified"].append(f"Fitted to their boxes: {len(fitted)} graphic text(s) set smaller, e.g. "
                                           + "; ".join(fitted[:4]) + ".")
                lay_v, lay_f = _split_notes(layer_notes or [])
                chk["verified"] += lay_v
                chk["findings"] = chk.get("findings", []) + list(pre_findings or []) + script_findings + gfx_findings + cov_f + lay_f \
                    + layer_problems + ([f"Shortened text — {t}. Shorten it yourself, or change it with action=edit." for t in trunc[:6]])
                if chk["findings"] and chk.get("status") == "ok":
                    chk["status"] = "needs_check"
            # technikailag jó, tartalmilag gyanús: kevés jelenet a hosszhoz (10-12: 60 s, 1 jelenet — a szöveg nagy része kimaradt).
            # Az ÁTLAG számít: egy 19 s-os idővonal három eseménnyel és három felirattal rendben van.
            ns = max(1, len(D["scenes"]))
            if result["check"] and (D["total"] / ns > 15 or (ns == 1 and D["total"] > 20)):
                result["check"]["status"] = "needs_check"
                result["check"]["findings"] = result["check"].get("findings", []) + [
                    f"Only {ns} scene(s) for {D['total']:.0f} s, so much of the text may be missing. Add scenes "
                    "(action=edit, add) or make it again without keep_wording."]
            if result["check"] and chart_notes:
                ok_notes = [n for n in chart_notes if not n.startswith("⚠ ")]
                bad = [n[2:] for n in chart_notes if n.startswith("⚠ ")]
                result["check"]["verified"] = result["check"].get("verified", []) + ok_notes
                if bad:
                    result["check"]["findings"] = result["check"].get("findings", []) + bad
                    if result["check"].get("status") == "ok":
                        result["check"]["status"] = "needs_check"
            result["video"] = vid
            result["files"] = [result["project"], vid]
            if kinds:
                step(4, 5, "rendering the transparent graphics")
                a_files, a_v, a_f = await _render_alpha(D, work, lotties, pid, stem, group, kinds, _probe_time(D, layer_rows))
                result["alpha"] = a_files
                result["files"] += a_files
                if result["check"] is not None:
                    result["check"]["verified"] += a_v
                    result["check"]["findings"] += a_f
                    if a_f and result["check"].get("status") == "ok":
                        result["check"]["status"] = "needs_check"
        else:
            result["files"] = [result["project"]]
            result["check"] = _unrendered_check(list(pre_findings or []) + script_findings + gfx_findings
                                                + _split_notes(layer_notes or [])[1] + layer_problems
                                                + [n[2:] for n in chart_notes if n.startswith("⚠ ")], lot_v)
        step(5, 5, "done")
        return result
    finally:
        shutil.rmtree(work, ignore_errors=True)


# ---------- a felhasználó felvétele + időkódos rétegek (overlay) ----------

FOOTAGE_EXTS = {"mp4", "mov", "webm", "m4v"}


def overlay_input(body: dict[str, Any]) -> tuple[str, str]:
    """Az overlay-kérés gyors ellenőrzése a munka ELŐTT (a végpont is ezt hívja, hogy a hiba azonnal jöjjön, ne a jobból):
    a fájl (videó, nem projekt), az alpha és a rétegek. → (file_id, kiterjesztés)"""
    fid = str(body.get("file_id") or "").strip()
    ext = fid.rsplit(".", 1)[-1].lower() if "." in fid else ""
    if ext == "ecproj":
        raise ValueError("That is a video project: add layers to it with video(action=\"edit\", file_id=…, "
                         "changes={\"add_layers\": […]}) — overlay takes a video file (mp4, mov, webm, m4v).")
    if ext not in FOOTAGE_EXTS:
        raise ValueError("overlay puts graphics on your own video: give file_id = an uploaded mp4, mov, webm or m4v"
                         + (f" (not .{ext})" if ext else "") + ". To make a video from a text, use action=\"make\".")
    crafts.existing(fid)
    alpha_kinds(body.get("alpha"))
    card = end_card_of(body)                                                   # a zárókártya alakja is most derüljön ki
    if not body.get("layers") and not body.get("subtitles") and not card:
        raise ValueError("overlay needs layers=[…] and/or subtitles=[…] (or \"auto\"), e.g. [{\"at\": \"00:02\", \"type\": "
                         "\"lower_third\", \"name\": \"Anna Kovács\", \"role\": \"Mayor\"}, {\"at\": \"00:09\", \"dur\": 5, "
                         "\"type\": \"chart\", \"chart\": {…}}].")
    layers_from(_unsaid(body.get("layers")), body.get("text") or "", [])   # szerkezeti hiba (típus, idő, hely) most derüljön ki
    if body.get("subtitles") != "auto":
        subtitle_cues(body.get("subtitles"))                                   # a feliratok formája is
    return fid, ext


async def overlay(body: dict[str, Any], progress: Optional[Callable[..., None]] = None) -> dict[str, Any]:
    """A felhasználó felvétele az alapsáv (mérete, képkocka/s-e és hangja marad — nincs 9:16-ra kényszerítés), fölötte
    időkódos rétegek a meglévő ábra-építőkkel → .ecproj (szerkeszthető) + MP4, alpha esetén átlátszó grafika is.
    body: {file_id, layers, text? (a forrás — a számok, idézetek ellenőrzéséhez), title?, accent?, brand?, alpha?, render?}"""
    step = progress or (lambda *a, **k: None)
    if not available():
        raise RuntimeError("The video engine isn't installed on this server")
    fid, ext = overlay_input(body)
    step(0, 4, "reading the footage")
    info = await probe_movie(crafts.existing(fid))
    if info["duration"] < 1.0 or not info["width"]:
        raise ValueError("That video is too short (under 1 s) or has no picture — give another one.")
    tok, fit_tok = _TRUNC.set([]), _FIT.set([])
    try:
        notes: list[str] = []
        topts = transcript_opts(body.get("transcript"))       # a kliens súgása a felismerőnek (prompt, vocabulary, model)
        raw = await said_to_times(body.get("layers"), crafts.existing(fid), notes, step, body.get("language"), new_ids=True,
                                  topts=topts)
        layers = layers_from(raw, body.get("text") or "", notes)
        subs = await _auto_subtitles(fid, body.get("language"), step, topts) if body.get("subtitles") == "auto" \
            else subtitle_cues(body.get("subtitles"))
        card = end_card_of(body)
        if not layers and not subs and not card:
            raise ValueError("overlay needs layers=[…], e.g. [{\"at\": \"00:02\", \"type\": \"lower_third\", \"name\": \"Anna Kovács\", "
                             "\"role\": \"Mayor\"}, {\"at\": \"00:09\", \"dur\": 5, \"type\": \"chart\", \"chart\": {…}}]"
                             + (" — every layer given was left out: " + "; ".join(n[2:] for n in notes if n.startswith("⚠ ")) if notes else "") + ".")
        from app.storage.naming import result_stem, root_name
        base = {"asset": f"footage.{ext}", "file_id": fid, **{k: info[k] for k in ("width", "height", "fps", "duration", "audio")}}
        if info.get("rotation") and int(info["rotation"]) % 360:
            # telefonos felvétel elforgatás-jelöléssel: a motor a tárolt (oldalra fektetett) képet olvasná — egyszer
            # felállítjuk (ffmpeg autorotate), és ez lesz az alapsáv
            step(0, 4, "turning the footage upright")
            base["media_id"], base["asset"] = await _upright(crafts.existing(fid)), "footage.mp4"
            base["rotated"] = int(info["rotation"])
        if card:                                              # zárókártya: a felvétel vége kimerevítve, a hang csenddel
            step(0, 4, "holding the last frame for the end card")
        base = await _with_end_card(base, card)
        opts = {"accent": await _brand_accent(body), "alpha": body.get("alpha"), "music": body.get("music"),
                "title": _clean(body.get("title") or root_name(display_label(fid), "Video"), 70, "Title"),
                "camera": _footage_camera(body.get("camera")), "subtitles": subs, **({"transcript": topts} if topts else {})}
        return await _assemble_overlay(base, layers, opts, source=body.get("text") or "", layer_notes=notes, step=step,
                                       label_stem=result_stem(fid, "with graphics"), operation="video overlay",
                                       render=body.get("render", True) is not False, draft=bool(body.get("draft")))
    finally:
        _TRUNC.reset(tok)
        _FIT.reset(fit_tok)


def _say_of(v: Any) -> Optional[str]:
    if isinstance(v, dict) and isinstance(v.get("say"), str) and v["say"].strip():
        return v["say"].strip()
    if isinstance(v, str) and v.lower().startswith("say:"):
        return v[4:].strip() or None
    return None


def _unsaid(raw: Any) -> Any:
    """Az előzetes (szerkezeti) ellenőrzéshez: a {"say": …} idők helyén egy helyőrző — az igazi időt az átirat adja."""
    if not _says(raw):
        return raw
    out = []
    for ly in raw:
        if isinstance(ly, dict) and (_say_of(ly.get("at")) or _say_of(ly.get("until"))):
            ly = dict(ly)
            if _say_of(ly.get("at")):
                ly["at"] = 0.0
            if _say_of(ly.get("until")):
                try:
                    ly["until"] = timecode(ly.get("at", 0.0)) + 1.0
                except ValueError:
                    ly["until"] = 1.0
        out.append(ly)
    return out


def _says(raw: Any) -> bool:
    return isinstance(raw, list) and any(isinstance(ly, dict) and (_say_of(ly.get("at")) or _say_of(ly.get("until")))
                                         for ly in raw)


def transcript_opts(raw: Any) -> dict[str, Any]:
    """overlay transcript={prompt, vocabulary, model, force?, correct?} → a beszédfelismerő (transcript.transcribe)
    paraméterei — csak azok, amelyeket az ismer (a magyar átirat bővítése előtt a régi aláírás mellett nem hibázik)."""
    if not isinstance(raw, dict) or not raw:
        return {}
    import inspect
    from app.tools import transcript
    known = set(inspect.signature(transcript.transcribe).parameters)
    out = {k: raw[k] for k in ("prompt", "vocabulary", "model", "force", "correct") if raw.get(k) not in (None, "") and k in known}
    if len(out) < len([k for k in raw if raw.get(k) not in (None, "")]):
        log.info("overlay transcript: a felismerő nem ismeri: %s", sorted(set(raw) - set(out)))
    return out


async def said_to_times(raw: Any, footage: Path, notes: list[str], step: Callable[..., None] = lambda *a, **k: None,
                        language: Optional[str] = None, new_ids: bool = False, topts: Optional[dict[str, Any]] = None) -> Any:
    """Réteg mondathoz kötve: at / until = {"say": "az ügyfél fellebbez"} → másodperc a felvétel átiratából. A réteg
    a kimondás előtt egy hajszállal jelenik meg, és (ha nincs dur/until) a mondat végéig marad. Ha a mondat nem hangzik
    el, a réteg kimarad, és a jegyzet megmondja, mit hallott a gép helyette."""
    if not _says(raw):
        return raw
    from app.tools import transcript
    step(0, 4, "listening to the footage (transcript)")
    tr = await transcript.transcribe(footage, transcript_lang(language), **(topts or {}))
    out = []
    for n, ly in enumerate(raw, 1):
        if not isinstance(ly, dict):
            out.append(ly)
            continue
        ly = dict(ly)
        name = ly.get("id") or (f"L{n}" if new_ids else f"new layer {n}")     # új overlaynél az id-k sorban: L1, L2…
        ok = True
        for key in ("at", "until"):
            phrase = _say_of(ly.get(key))
            if not phrase:
                continue
            after = float(ly["at"]) if key == "until" and isinstance(ly.get("at"), (int, float)) else 0.0
            hit = transcript.find(tr, phrase, after=after)
            if hit is None:
                close = transcript.nearest(tr, phrase)
                notes.append(f"⚠ {name}: “{phrase}” is not said in the footage" + (f" (closest: “{close}”)" if close else
                             " (no speech was heard)") + " — left out. Give the time instead (at: \"01:24\"), or quote what is said.")
                ok = False
                break
            if key == "at":
                ly["at"] = round(max(0.0, hit["start"] - 0.15), 2)
                if ly.get("dur") is None and ly.get("until") is None:
                    ly["dur"] = round(max(3.0, hit["sentence_end"] - hit["start"] + 0.8), 2)
            else:
                ly["until"] = round(hit["end"] + 0.3, 2)
            notes.append(f"{name} {'starts' if key == 'at' else 'ends'} when “{phrase}” is said ({_tc(hit['start'])}"
                         + (f", heard as “{hit['said']}”" if _norm_said(hit['said']) != _norm_said(phrase) else "") + ").")
        if ok:
            out.append(ly)
    return out


def _norm_said(s: str) -> str:
    from app.tools.transcript import _norm
    return "".join(_norm(w) for w in s.split())


def transcript_lang(language: Optional[str]) -> Optional[str]:
    from app.tools import narration
    return narration.lang_code(language)


def check_music(choice: Any) -> None:
    """A zene-választás ellenőrzése a munka indítása ELŐTT (ne egy perc renderelés után bukjon el)."""
    c = str(choice or "").strip().lower()
    if not c or c in _NO_NARRATION:
        return
    from app.tools import music as music_lib
    if music_lib.find(c) is None and c not in music_lib.moods():
        raise ValueError(f"music: {choice!r} — choose a mood ({', '.join(music_lib.moods())}), a track id, or \"none\"")


def _builtin_music(opts: dict[str, Any]) -> bool:
    """A beépített zenetárból kér-e zenét (hangulat vagy darab), saját zenefájl nélkül."""
    c = str(opts.get("music") or "").strip().lower()
    return bool(c) and c not in _NO_NARRATION and not opts.get("music_file_id")


async def _ffmpeg_audio(src: Path, dst: Path) -> None:
    from app.convert.engines import run
    await run("ffmpeg", "-v", "error", "-y", "-i", str(src), "-vn", "-ac", "2", "-ar", "48000", str(dst),
              timeout=settings.VIDEO_TIMEOUT_S)


async def _music_under(opts: dict[str, Any], work: Path, total: float, speech: Optional[Path], user_music: Optional[str],
                       seed: str) -> tuple[Optional[str], Optional[str], list[str], list[str]]:
    """Háttérzene: a beépített tárból (opts.music = hangulat vagy darab, a videó hosszára hurkolva) vagy a saját zene;
    ha van beszéd (narráció, saját hang, a felvétel hangja), a zene alatta halkul (sidechain), a szünetekben visszajön.
    → (az eredeti zene neve — a spec ebből halkít újra szerkesztéskor —, a projektbe kerülő sáv, verified, findings)."""
    from app.tools import music as music_lib
    verified: list[str] = []
    orig, label = user_music, "Your music"
    if not user_music and _builtin_music(opts):
        t = music_lib.pick(str(opts["music"]).strip().lower(), seed)
        if t is None:
            return None, None, [], ["The built-in music library isn't installed on this server — the video has no music "
                                    "(give music_file_id with your own)."]
        await music_lib.bed(t, total, work / "music.wav")
        orig, label = "music.wav", f"Music: “{t['title']}” ({t['mood']}, {t.get('genre')}; CC0 — no credit needed)"
    if not orig:
        return None, None, [], []
    if speech is not None and speech.exists():
        await music_lib.duck(speech, work / orig, work / "music-ducked.wav")       # a beszéd mért szintje alatt 14 dB
        verified.append(f"{label} — it drops under the speech and comes back in the pauses.")
        return orig, "music-ducked.wav", verified, []
    verified.append(f"{label} — there is no speech, so it plays on its own.")
    return orig, orig, verified, []


def end_card_of(body: dict[str, Any]) -> Optional[dict[str, Any]]:
    """Zárókártya az overlay végére: end_card = {"title", "subtitle"?, "kicker"?, "dur"?} | "Cím" — vagy end_title
    (+ end_subtitle). {} / "none" / false: nincs. → {title, subtitle, kicker, dur} | None"""
    raw = body.get("end_card")
    if raw in (None, "", False, {}, [], "none", "off") and not body.get("end_title"):
        return None
    if raw in ("none", "off", False) and "end_card" in body:
        return None
    raw = {"title": raw} if isinstance(raw, str) else dict(raw) if isinstance(raw, dict) else {}
    title = _clean(raw.get("title") or raw.get("text") or body.get("end_title"), 60, "End title")
    if not title:
        raise ValueError('end_card: give its title, e.g. {"title": "Thanks for watching", "subtitle": "Tamás TV", "dur": 4} '
                         '(or end_title="…").')
    dur = timecode(raw.get("dur") if raw.get("dur") is not None else raw.get("duration", 4.0), "end_card dur")
    if not 1.5 <= dur <= 20:
        raise ValueError(f"end_card dur: 1.5–20 seconds — not {dur:g}.")
    return {"title": title, "subtitle": _clean(raw.get("subtitle") or body.get("end_subtitle"), 90, "End subtitle"),
            "kicker": _clean(raw.get("kicker"), 30, "Kicker"), "dur": round(dur, 2)}


async def _extend(src: Path, secs: float) -> str:
    """A felvétel vége kimerevítve secs másodpercre (tpad: az utolsó kocka marad), a hang csenddel kitöltve (apad) —
    a zárókártya alá (az overlay hossza különben mindig a felvételé). → a meghosszabbított munkafájl file_id-ja."""
    from app.convert.engines import run
    fid = new_file_id("mp4")
    k = (await probe_movie(src)).get("audio_stream")
    aud = ["-map", f"0:a:{k}", "-af", f"apad=pad_dur={secs:.3f}", "-c:a", "aac", "-b:a", "192k"] if k is not None else []
    await run("ffmpeg", "-v", "error", "-y", "-i", str(src), "-map", "0:v:0", *aud,
              "-vf", f"tpad=stop_mode=clone:stop_duration={secs:.3f}", "-c:v", "libx264", "-preset", "veryfast", "-crf", "17",
              "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(output_path(fid)), timeout=settings.VIDEO_TIMEOUT_S)
    write_meta(fid, operation="footage held for the end card", source_file_id=None, extra={"ephemeral": True})
    return fid


async def _with_end_card(base: dict[str, Any], card: Optional[dict[str, Any]]) -> dict[str, Any]:
    """Az overlay alapsávja a zárókártyához: a kártya hosszával meghosszabbított felvétel (új asset-névvel, hogy a
    szerkesztés ne a régi hosszút vegye át), vagy — ha a kártya kikerült — újra az eredeti. base["plain"] = a
    meghosszabbítás előtti felvétel (asset, media_id, duration)."""
    b = dict(base)
    plain = b.pop("plain", None) or {k: b.get(k) for k in ("asset", "media_id", "duration")}
    old = b.pop("end_card", None)
    b.update(asset=plain["asset"], duration=plain["duration"])
    if plain.get("media_id"):
        b["media_id"] = plain["media_id"]
    else:
        b.pop("media_id", None)
    if not card:
        return b
    name = f"footage-end-{card['dur']:g}s.mp4"
    if old and old.get("dur") == card["dur"] and base.get("asset") == name and base.get("media_id"):
        b["media_id"] = base["media_id"]                                # ugyanolyan hosszú kártya: a meglévő marad
    else:
        b["media_id"] = await _extend(crafts.existing(plain.get("media_id") or b["file_id"]), card["dur"])
    b.update(asset=name, duration=round(float(plain["duration"]) + card["dur"], 3), end_card=card, plain=plain)
    return b


def _has_lottie(story: dict[str, Any]) -> bool:
    return any(e.get("shape") == "lottie" and (e.get("file") or e.get("asset") or e.get("credit"))
               for h in (story.get("scenes") or []) + (story.get("layers") or []) for e in (h.get("chart") or {}).get("elements") or [])


def _credit_line(credit_list: list[str], story: dict[str, Any]) -> tuple[str, str]:
    """A kreditsor → (teljes, linkekkel: a projekt metájába; rövid, a képernyőre: név + licenc, linkek nélkül).
    „Pictures and animations:”, ha Lottie is van a videóban (10-19: a saját kreditű Lottie mellett „Pictures:” lett)."""
    if not credit_list:
        return "", ""
    anim = _has_lottie(story) or any(c.startswith("Animation ") for c in credit_list)
    label = ("Pictures and animations: " if len(credit_list) > 1 else "") if anim else "Pictures: "
    short = [re.sub(r"\s+—\s+https?://\S+", "", re.sub(r"^Photo:\s*", "", c)).strip() for c in credit_list]
    return label + " · ".join(credit_list), (label + " · ".join(dict.fromkeys(x for x in short if x)))[:300]


async def _upright(src: Path) -> str:
    """Az elforgatás-jelöléssel tárolt felvétel álló/fekvő helyzetbe újrakódolva (rejtett munkafájl) → file_id."""
    from app.convert.engines import run
    fid = new_file_id("mp4")
    k = (await probe_movie(src)).get("audio_stream")
    aud = ["-map", f"0:a:{k}"] if k is not None else []      # csak az első dekódolható hangsáv (térhang nélkül)
    await run("ffmpeg", "-v", "error", "-y", "-i", str(src), "-map", "0:v:0", *aud, "-c:v", "libx264", "-preset", "veryfast",
              "-crf", "17", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "192k", "-movflags", "+faststart", str(output_path(fid)),
              timeout=settings.VIDEO_TIMEOUT_S)
    write_meta(fid, operation="footage set upright", source_file_id=None, extra={"ephemeral": True})
    return fid


# ---------- feliratok a saját felvételen (10-17, Kommandant: „webes feliratozó, szövegjavítással és a saját stílusunkkal”) ----------
# A felirat a make-videó feliratsávja (sötét, lekerekített sáv, a saját betűnk, a keret aljának közepén) — a builder ugyanazzal
# a kóddal rajzolja; a szöveg a személy (vagy a kliens) javított átirata. A .srt is elkészül (a lejátszók feliratfájlja).

def subtitle_cues(raw: Any) -> list[dict[str, Any]]:
    """Feliratok: [{start, end, text}] (az átirat „at”/„until” időkódja is jó) → időrendben, átfedés nélkül, üres nélkül."""
    if raw in (None, "", [], False):
        return []
    if not isinstance(raw, list):
        raise ValueError('subtitles: a list of {"start", "end", "text"} (seconds or "MM:SS"), or "auto"')
    out = []
    for c in raw[:3000]:
        if not isinstance(c, dict):
            continue
        txt = re.sub(r"\s+", " ", str(c.get("text") or "")).strip()[:300]
        if not txt:
            continue
        a = timecode(c.get("start", c.get("at", 0)), "subtitle start")
        b = c.get("end", c.get("until"))
        b = timecode(b, "subtitle end") if b not in (None, "") else a + max(1.5, len(txt) / 15)
        out.append({"start": round(a, 3), "end": round(max(b, a + 0.4), 3), "text": txt})
    out.sort(key=lambda x: x["start"])
    for x, y in zip(out, out[1:]):                       # egyszerre egy felirat látszik
        x["end"] = round(min(x["end"], y["start"]), 3)
    return [x for x in out if x["end"] - x["start"] >= 0.2]


def caption_track(cues: list[dict[str, Any]], W: int, H: int, total: float) -> list[dict[str, Any]]:
    """A feliratok a make-videó feliratsávjának méretében: legfeljebb két sor, a hosszú mondat darabokra (mondat- és
    tagmondathatáron), a darabok a mondat idején belül a hosszuk arányában."""
    L = _layout(W, H)
    cpl = L["capW"] / (0.53 * L["capSize"])
    out = []
    for c in cues:
        a, b = max(0.0, c["start"]), min(total, c["end"])
        if b - a < 0.2:
            continue
        parts = _chunks(c["text"], int(cpl * 2 * 0.9)) or [c["text"]]
        tot, cur = sum(len(x) for x in parts) or 1, a
        for x in parts:
            span = (b - a) * len(x) / tot
            out.append({"t0": round(cur, 3), "t1": round(cur + span, 3), "text": x, "lines": 1 if len(x) <= cpl * 0.95 else 2})
            cur += span
    return out


def subtitle_force_style(W: int, H: int) -> str:
    """A gyors ráégetés (ffmpeg/libass) stílusa a videók feliratsávjához igazítva: Inter, fehér betű sötét, áttetsző
    sávon, alul középen, a make-videó méretében. A libass a .srt-t 384×288-as vásznon méretezi — ehhez arányítunk."""
    L = _layout(W, H)
    k = 288 / H
    size = max(6, round(L["capSize"] * k))
    margin_v = max(4, round((H - L["capY"]) * k - size * 0.55))
    side = max(4, round((W - L["capW"]) / 2 / W * 384 * (W / H) / (4 / 3)))
    box = "&H470F0B0A"                       # ~72% fedő, sötét (#0A0B0F) — a feliratsáv színe; ASS: &HAABBGGRR
    return (f"FontName=Inter,Bold=1,FontSize={size},PrimaryColour=&H00FFFFFF,OutlineColour={box},BackColour={box},"
            f"BorderStyle=3,Outline={max(1, round(size * 0.35))},Shadow=0,Alignment=2,MarginV={margin_v},MarginL={side},MarginR={side}")


def _srt_text(cues: list[dict[str, Any]]) -> str:
    from app.tools import transcript
    return transcript.srt({"sentences": cues})


async def _auto_subtitles(fid: str, language: Optional[str], step: Callable[..., None],
                          topts: Optional[dict[str, Any]] = None) -> list[dict[str, Any]]:
    """subtitles="auto": a felvétel átirata, javítás nélkül (a webes feliratozó előbb megmutatja és javíttatja)."""
    from app.tools import transcript
    step(1, 4, "listening (transcript)")
    tr = await transcript.transcribe(crafts.existing(fid), transcript_lang(language), **(topts or {}))
    return subtitle_cues(transcript.cues(tr))


def _has_3d(D: dict[str, Any]) -> bool:
    return any(el.get("k") == "threejs" for sc in D.get("scenes") or [] for el in ((sc.get("chart") or {}).get("elements") or [])) \
        or any(el.get("k") == "threejs" for ly in D.get("layers") or [] for el in ly.get("elements") or [])


async def _render_3d(D: dict[str, Any], work: Path, accent: list[float]) -> tuple[list[str], list[str]]:
    """A three.js-jelölők (jelenetek, rétegek) helyére a lerenderelt klip — a jelenet / réteg hosszában és méretében.
    → (a klipek fájlnevei a projekt médiájához, megjegyzések)."""
    from app.tools import threejs
    jobs = [(el, float(sc["dur"])) for sc in D.get("scenes") or [] for el in ((sc.get("chart") or {}).get("elements") or [])
            if el.get("k") == "threejs"]
    jobs += [(el, float(ly["dur"])) for ly in D.get("layers") or [] for el in ly.get("elements") or [] if el.get("k") == "threejs"]
    names, notes = [], []
    for el, dur in jobs:
        w, h = el["size"]
        clip, n_ = await threejs.render(el["code"], int(w), int(h), dur, transparent=el.get("transparent", True),
                                        accent=accent, dest=work)
        notes += n_
        if clip is None:                     # a hiba a check-ben; a jelenet a háttérrel és a címmel megy tovább
            el.clear()
            el.update({"k": "rect", "name": "3D scene (failed)", "size": [2, 2], "pos": [0, 0], "opacity": 0})
            continue
        names.append(clip.name)
        pos = el["pos"]
        el.clear()
        el.update({"k": "image", "name": "3D scene", "file": clip.name, "size": [w, h], "pos": pos, "fit": "contain", "t0": 0})
    return names, notes


def _footage_camera(cam: Any) -> Any:
    """A saját felvétel kamerája (ráközelítés egy pillanatra — „punch-in”): ugyanaz a leírás, mint a jelenet kamerája,
    abszolút időben (másodperc vagy időkód)."""
    if cam in (None, "", "none", False):
        return None
    tmp: dict[str, Any] = {}
    _scene_motion({"camera": cam}, tmp)
    return tmp.get("camera")


async def _assemble_overlay(base: dict[str, Any], layers: list[dict[str, Any]], opts: dict[str, Any], *, source: str,
                            layer_notes: list[str], step: Callable[..., None], label_stem: str, operation: str,
                            parent: Optional[str] = None, render: bool = True, draft: bool = False) -> dict[str, Any]:
    """Felvétel + rétegek → .ecproj (+ MP4, + átlátszó grafika). Az overlay és a szerkesztése közös útja; parent = a
    szerkesztett projekt (a felvétel és a rétegek képei onnan jönnek, nem másolódnak újra)."""
    W, H = int(base["width"]), int(base["height"])
    fps = float(base.get("fps") or FPS)
    fps = fps if 1 <= fps <= 120 else FPS
    total = round(float(base["duration"]), 3)
    kinds = alpha_kinds(opts.get("alpha"))
    work = Path(tempfile.mkdtemp(prefix="video-", dir=str(temp_path("").parent)))
    try:
        shared = assets_of(parent).get(base["asset"]) if parent else None
        src = crafts.existing(shared or base.get("media_id") or base["file_id"])
        try:                                                  # a felvétel linkként a munkamappába (nem másolat)
            os.symlink(src, work / base["asset"])
        except OSError:
            shutil.copyfile(src, work / base["asset"])
        subs = opts.get("subtitles") or []
        story = {"title": opts.get("title") or "Video", "layers": layers, **({"subtitles": subs} if subs else {})}
        safe = H - round(min(W, H) * 0.05) * 1.2
        if subs:                                              # a rétegek a feliratsáv fölött maradnak
            safe = min(safe, _caption_top(W, H))
        map_credits, gfx_findings = await _place_on_maps(story, (W, H, safe))
        gfx_names, lotties, gfx_credits = await _materialize_gfx(story, work, parent, (W, H, safe))
        gfx_findings = gfx_findings + _gfx_failures(story)
        credit_list = list(dict.fromkeys(gfx_credits + map_credits))
        lang = _lang_of(_layer_text_all({"layers": [{k: v for k, v in ly.items() if k != "name"} for ly in layers]}) + " " + (source or ""))
        acc = _rgb(opts.get("accent"))
        D: dict[str, Any] = {"W": W, "H": H, "fps": fps, "total": total, "accent": acc, "dark": [0.055, 0.06, 0.08], "font": FONT,
                             "base": {"file": base["asset"], "label": "Footage"}, "scenes": [], "captions": [], "audio": {}}
        if opts.get("camera"):
            D["camera"] = _camera_keys(opts["camera"], total, W, H)
        if subs:                                              # a make-videó feliratsávja: ugyanaz a builder-kód rajzolja
            Lc = _layout(W, H)
            D["L"] = {k: Lc[k] for k in ("capW", "capSize", "capY", "capPad")}
            D["captions"] = caption_track(subs, W, H, total)
        card, all_layers = base.get("end_card"), layers
        if card:                                              # zárókártya a kimerevített vég fölött, a kreditekkel
            _, on_screen = _credit_line(credit_list, story)
            all_layers = layers + [{"id": "end_card", "type": "fullscreen", "kind": "text", "end": True,
                                    "at": float(base["plain"]["duration"]), "dur": card["dur"], "text": card["title"],
                                    "subtitle": card.get("subtitle") or "", "kicker": card.get("kicker") or "",
                                    **({"credits": on_screen} if on_screen else {})}]
        D["layers"], rows, problems = layer_comps(all_layers, W, H, total, lang, acc, safe)
        cuts = _clip_sounds(layers)                           # bevágott felvételek hangja (teljes képen a felvételé helyett)
        if cuts:
            D["audio"]["clips"] = cuts
        m_v, m_f, track = [], [], None
        if _builtin_music(opts):                              # zene a felvétel saját hangja alá (ahol beszélnek, halkul)
            speech = None
            if base.get("audio"):
                speech = work / "footage-audio.wav"
                await _ffmpeg_audio(work / base["asset"], speech)
            _, track, m_v, m_f = await _music_under(opts, work, total, speech, None, story["title"])
            if track:
                D["audio"] = {**D["audio"], "music": track, "musicDb": 0 if speech else -10}
        problems = problems + m_f
        if _has_3d(D):                                        # three.js-rétegek: a réteg hosszában és méretében
            step(2, 4, "rendering the 3D scenes")
            three_names, three_notes = await _render_3d(D, work, acc)
            gfx_names = gfx_names + three_names
            m_v = m_v + [n for n in three_notes if not n.startswith("⚠ ")]
            problems = problems + [n[2:] for n in three_notes if n.startswith("⚠ ")]
        step(2, 4, "building the timeline")
        out = await _build(D, work, lotties)
        problems = problems + _text_report(out, D)            # a renderer mért szövegdobozai: ütközés, kilógás, tördelés
        lot_v, lot_f = await _lottie_report(D, work, out) if lotties else ([], [])   # a Lottie-k mintája
        problems, m_v = problems + lot_f, m_v + lot_v
        scripted = any((ly.get("chart") or {}).get("type") == "script" for ly in layers)
        script_findings = _script_report(out, source or "", {ly["id"] for ly in layers}) if scripted else []
        log.info("overlay felépítve: %s", out.strip().splitlines()[-1:] if out.strip() else "")
        pid = new_file_id("ecproj")
        shutil.copyfile(work / "project.ecproj", output_path(pid))
        from app.routes.make import _new_group
        group = _new_group()
        credits = _credit_line(credit_list, story)[0]
        spec = {"kind": "overlay", "story": _without_runtime(story),
                "base": {k: base.get(k) for k in ("asset", "file_id", "media_id", "rotated", "width", "height", "fps", "duration", "audio",
                                                  "end_card", "plain")
                         if base.get(k) is not None},
                "opts": {"accent": opts.get("accent"), "alpha": opts.get("alpha"), "title": story["title"],
                         "music": opts.get("music") if _builtin_music(opts) else None, "camera": opts.get("camera"),
                         "subtitles": subs or None, "transcript": opts.get("transcript") or None},
                "source": (source or "")[:30000], "layer_source": (source or "")[:30000],
                "layer_notes": [n for n in layer_notes if not n.startswith("⚠ ")], "timeline": [], "layer_timeline": rows}
        write_meta(pid, label=f"{label_stem}.ecproj", operation=operation, source_file_id=parent or base["file_id"],
                   extra={"video": {"format": "custom", "width": W, "height": H, "fps": fps, "duration": total, "duration_mode": "footage",
                                    "layers": len(D["layers"]), "title": story["title"]},
                          "group": group, "credit": credits or None, "spec": spec})
        if shared or base.get("media_id"):                    # a felvétel egy példányban: a változatok ugyanazt használják
            write_meta(pid, extra={"assets": {**assets_of(pid), base["asset"]: shared or base["media_id"]}})
        else:
            add_asset(pid, base["asset"], work / base["asset"])
        for name in gfx_names + ([track] if track else []):
            add_asset(pid, name, work / name)
        placed = {r_["id"] for r_ in rows}
        result: dict[str, Any] = {"file_id": pid, "project": _file_entry(pid), "title": story["title"], "format": "custom",
                                  "width": W, "height": H, "fps": fps,
                                  "duration": {"requested": None, "actual": total, "mode": "footage"},
                                  "layers": layer_listing([ly for ly in layers if ly["id"] in placed]), "credits": credits or None}
        if _lottie_picks(story):
            result["lottie_picks"] = _lottie_picks(story)
        if card:
            result["end_card"] = {**card, "at": float(base["plain"]["duration"]), "until": total}
        srt_entry = None
        if subs:                                              # a javított feliratok fájlként is (lejátszók, YouTube)
            sid = new_file_id("srt")
            output_path(sid).write_text(_srt_text(subs), "utf-8")
            write_meta(sid, label=f"{label_stem}.srt", operation="subtitles", source_file_id=pid, extra={"group": group})
            srt_entry = _file_entry(sid)
            result["subtitles"] = {"count": len(subs), "srt": srt_entry}
        if not render:
            result["files"] = [result["project"]]
            result["check"] = _unrendered_check(_split_notes(layer_notes)[1] + problems + gfx_findings + script_findings, m_v)
            return result
        step(3, 4, "rendering the video")
        vid = await _render_staged(pid, work / "project.ecproj", label_stem, group=group, want_s=total, mode="footage", draft=draft)
        result["duration"] = vid.pop("duration", None) or result["duration"]
        chk = vid.pop("check", None) or {"status": "needs_check", "findings": [], "verified": []}
        lay_v, lay_f = _split_notes(layer_notes)
        fitted, trunc = _FIT.get() or [], _TRUNC.get() or []
        chk["verified"] = m_v + [f"Your footage is the base track: {W}×{H}, {fps:g} fps, {_tc(total)}"
                           + (", with its own sound." if base.get("audio") else " (it has no sound track).")
                           + (f" It was stored turned by {base['rotated']}°, so it was set upright first." if base.get("rotated") else "")] \
            + ([f"End card “{card['title']}”: the footage's last frame is held for {card['dur']:g} s under it"
                 + (", the sound padded with silence" if base.get("audio") else "")
                 + f" — {_tc(base['plain']['duration'])} → {_tc(total)}."] if card else []) \
            + chk.get("verified", []) + [f"{len(rows)} layer(s) on the timeline: " + ", ".join(f"{r_['id']} {_tc(r_['start'])}" for r_ in rows) + "."] \
            + lay_v + ([f"Fitted to their boxes: {len(fitted)} text(s) set smaller, e.g. " + "; ".join(fitted[:4]) + "."] if fitted else [])
        chk["findings"] = chk.get("findings", []) + lay_f + problems + gfx_findings + script_findings + \
            [f"Shortened text — {t}. Shorten it yourself, or change it with action=edit." for t in trunc[:6]]
        result["video"] = vid
        result["files"] = [result["project"], vid] + ([srt_entry] if srt_entry else [])
        if subs:
            chk["verified"].append(f"{len(subs)} subtitle(s) in the Studio caption style, {len(D['captions'])} on screen "
                                   f"(long ones in two-line parts); the .srt is in the files.")
        if kinds:
            step(3, 4, "rendering the transparent graphics")
            a_files, a_v, a_f = await _render_alpha(D, work, lotties, pid, label_stem, group, kinds, _probe_time(D, rows))
            result["alpha"] = a_files
            result["files"] += a_files
            chk["verified"] += a_v
            chk["findings"] += a_f
        if chk["findings"] and chk.get("status") == "ok":
            chk["status"] = "needs_check"
        result["check"] = chk
        step(4, 4, "done")
        return result
    finally:
        shutil.rmtree(work, ignore_errors=True)


_STEM = lambda w: w[:5]


def _words(t: str) -> list[str]:
    import unicodedata
    t = unicodedata.normalize("NFKD", (t or "").lower())
    t = "".join(c for c in t if not unicodedata.combining(c))
    return [w for w in re.findall(r"[a-z0-9]+", t) if len(w) >= 4]


def _video_text(story: dict[str, Any]) -> str:
    parts = [story.get("title") or "", (story.get("end") or {}).get("title") or "", (story.get("end") or {}).get("subtitle") or ""]
    for sc in story.get("scenes") or []:
        parts += [sc.get("headline") or "", sc.get("caption") or ""]
        if sc.get("stat"):
            parts += [sc["stat"].get("value") or "", sc["stat"].get("label") or ""]
        ch = sc.get("chart") or {}
        parts.append(ch.get("title") or "")
        if ch.get("nodes"):                        # magyarázó ábra: a dobozok, az élcímkék, a sávok neve
            from app.tools import diagrams
            parts += diagrams.texts(ch)
        from app.tools import motion_charts
        if ch.get("type") in motion_charts.TYPES:
            parts += motion_charts.numbers(ch)
        for it in ch.get("items") or []:
            parts += [str(it.get("label") or ""), str(it.get("text") or ""), str(it.get("value") if it.get("value") is not None else "")]
        if ch.get("total"):
            parts.append(str(ch["total"]))
    return " ".join(parts) + " " + _layer_text_all(story)          # a rétegek szövege is a képernyőn van


def _coverage(source: str, story: dict[str, Any], keep: bool, scene_plan: bool) -> tuple[list[str], list[str], dict[str, Any]]:
    """Tartalmi lefedettség (10-12, ChatGPT: „egy technikailag hibátlan, de tartalmilag hiányos videó is sikeres lett”):
    a forrás mondatai és számai közül mi jelenik meg a videóban (felirat, cím, szám, ábra). Saját szövegnél (keep_wording)
    minden mondatnak benne kell lennie; cikknél a nagy hiány (< 60%) „needs_check”, a kimaradtak listája mindig ott van."""
    vt = _video_text(story)
    vstems = {_STEM(w) for w in _words(vt)}
    sents = [x.strip() for x in re.split(r"(?<=[.!?])\s+|\n+", source or "") if len(_words(x)) >= 4]
    missing = []
    for snt in sents:
        ws = {_STEM(w) for w in _words(snt)}
        if ws and len(ws & vstems) / len(ws) < 0.4:
            missing.append(snt)
    src_nums = _text_numbers(source or "") - {float(i) for i in range(0, 4)}
    vid_nums = _text_numbers(vt)
    miss_nums = sorted(src_nums - vid_nums)
    cov = {"sentences": {"total": len(sents), "covered": len(sents) - len(missing), "missing": missing[:12]},
           "numbers": {"total": len(src_nums), "covered": len(src_nums) - len(miss_nums),
                       "missing": [(int(x) if float(x).is_integer() else x) for x in miss_nums[:20]]}}
    if scene_plan and not source:
        return [], [], cov
    verified = [f"Covers {cov['sentences']['covered']} of {len(sents)} sentences and {cov['numbers']['covered']} of "
                f"{len(src_nums)} numbers of the source (coverage lists what is left out)."] if sents or src_nums else []
    findings = []
    s_ratio = (len(sents) - len(missing)) / len(sents) if sents else 1
    n_ratio = cov["numbers"]["covered"] / len(src_nums) if src_nums else 1
    if keep and missing:
        findings.append(f"{len(missing)} of your sentences are not in the video, e.g. “{missing[0][:90]}”.")
    elif s_ratio < 0.6 or n_ratio < 0.6:
        findings.append(f"Much of the source is left out ({cov['sentences']['covered']}/{len(sents)} sentences, "
                        f"{cov['numbers']['covered']}/{len(src_nums)} numbers), e.g. “{(missing or ['…'])[0][:90]}”"
                        + (f"; numbers not shown: {', '.join(str(x) for x in cov['numbers']['missing'][:8])}" if miss_nums else "")
                        + ". Add scenes with action=edit, or pass scenes=[…] yourself.")
    return verified, findings, cov


def _file_entry(fid: str) -> dict[str, Any]:
    return {"file_id": fid, "url": public_url(fid), "label": display_label(fid), "size_kb": output_path(fid).stat().st_size // 1024}


async def _render_staged(project_id: str, project_path: Path, stem: str, *, group: Optional[str] = None,
                         fmt: str = "mp4", draft: bool = False, want_s: Optional[float] = None, requested: Optional[float] = None,
                         mode: Optional[str] = None, why: Optional[str] = None) -> dict[str, Any]:
    """A projekt fő kompozíciója → MP4 (WebM / GIF) + check. A kért hossz (requested) és a mód a projekt metájából is
    jön (újrarenderelésnél): az eredmény duration = {requested, actual (a fájl mért hossza), mode}."""
    ext = "webm" if fmt == "webm" else "gif" if fmt == "gif" else "mp4"
    out = project_path.parent / f"render.{ext}"
    args = ["render", "--project", project_path.name, "--comp", _main_comp(project_path), "--out", out.name,
            "--format", {"mp4": "h264", "webm": "webm", "gif": "gif"}[ext], "--json"]
    if ext == "mp4":
        args += ["--bitrate", "2500" if draft else "8000", "--audio", "auto"]
    if draft:
        args += ["--resolution", "half", "--quality", "draft"]
    t0 = time.time()
    async with _render_gate():
        await _run(*args, cwd=project_path.parent)
    if not out.exists():
        raise RuntimeError("The render produced no file")
    from app.tools import ios_safe
    await ios_safe.finalize(out)                        # iPhone/Safari: a moov elöl (a renderer a végére írja)
    vid = new_file_id(ext)
    shutil.move(str(out), output_path(vid))
    meta = read_meta(project_id) or {}
    # a projekt és a renderje egy kártya (group); a gyors vázlat külön „változat”, hogy a Download sose azt adja
    write_meta(vid, label=f"{stem} (draft).{ext}" if draft else f"{stem}.{ext}", operation="video render", source_file_id=project_id,
               extra={"group": group or meta.get("group"), "variant": "draft" if draft else None, "credit": meta.get("credit")})
    if not draft:
        write_meta(project_id, extra={"render_file_id": vid, "render_draft": False})
    elif not meta.get("render_file_id") or meta.get("render_draft"):
        # gyors vázlat (draft=true): amíg nincs teljes render, az előnézet (preview, frame) ezt mutatja
        write_meta(project_id, extra={"render_file_id": vid, "render_draft": True})
    log.info("videó renderelve: %s → %s (%.1f s)", project_id, vid, time.time() - t0)
    from app.tools import checks
    v = (read_meta(project_id) or {}).get("video") or {}
    wh = FORMATS.get(v.get("format")) or ((v["width"], v["height"]) if v.get("width") and v.get("height") else None)
    requested = requested if requested is not None else v.get("requested_s")
    mode = mode or v.get("duration_mode")
    chk = await asyncio.to_thread(checks.check_video, output_path(vid), want_s, wh if ext == "mp4" else None, draft,
                                  requested, mode, why)
    try:
        actual = round((await probe_movie(output_path(vid)))["duration"], 2)
    except Exception:
        actual = want_s
    return {**_file_entry(vid), "render_seconds": round(time.time() - t0, 1), "check": chk,
            "duration": {"requested": requested, "actual": actual, "mode": mode}}


async def render(file_id: str, fmt: str = "mp4", draft: bool = False) -> dict[str, Any]:
    """Egy videóprojekt (akár kézzel szerkesztett) újrarenderelése."""
    if not file_id.lower().endswith(".ecproj"):
        raise ValueError("render needs a video project (.ecproj)")
    work = Path(tempfile.mkdtemp(prefix="video-", dir=str(temp_path("").parent)))
    try:
        proj = stage(file_id, work)
        stem = re.sub(r"( \(edited [^)]*\))+$", "", Path(display_label(file_id)).stem)   # a render neve a projekté, „(edited …)” nélkül
        return await _render_staged(file_id, proj, stem, fmt=fmt, draft=draft)
    finally:
        shutil.rmtree(work, ignore_errors=True)


async def blank(fmt: Optional[str] = None, duration: Any = None) -> dict[str, Any]:
    """Üres videóprojekt: egy „Video” kompozíció a választott méretben, a szerkesztőbe."""
    fmt = fmt_name(fmt)
    W, H = FORMATS[fmt]
    dur = float(min(max(float(duration or 30), 1), 600))
    work = Path(tempfile.mkdtemp(prefix="video-", dir=str(temp_path("").parent)))
    try:
        params = {"name": "Video", "width": W, "height": H, "frameRate": FPS, "duration": dur, "background": "#0e0f14"}
        await _run("exec", "comp.new", "--params", json.dumps(params), "--empty", "--save-as", "project.ecproj", cwd=work, timeout=120)
        fid = new_file_id("ecproj")
        shutil.copyfile(work / "project.ecproj", output_path(fid))
    finally:
        shutil.rmtree(work, ignore_errors=True)
    label = "Untitled video.ecproj"
    write_meta(fid, operation="new video", label=label, extra={"video": {"format": fmt, "duration": dur},
                                                                "blank": {"kind": "video", "size": fmt}})
    return {"file_id": fid, "url": public_url(fid), "label": label, "size_kb": output_path(fid).stat().st_size // 1024}


def info(file_id: str) -> dict[str, Any]:
    p = crafts.existing(file_id)
    ext = p.suffix.lstrip(".").lower()
    meta = read_meta(file_id) or {}
    if ext != "ecproj":
        src = meta.get("source_file_id")
        out = {"file_id": file_id, "kind": "movie" if ext in MOVIE_EXTS else ext, "size_kb": p.stat().st_size // 1024,
               "made_from": src}
        if src and str(src).lower().endswith(".ecproj"):
            out["project"] = src
            out["note"] = ("To change this video, use video(action=\"edit\") with this file or its project; info on the project "
                           "lists the scenes and layers.")
        elif ext in FOOTAGE_EXTS:
            out["note"] = ("To put graphics on this video (name bars, chapter titles, charts, a split screen…), use "
                           "video(action=\"overlay\", file_id=…, layers=[…]).")
        return out
    data = json.loads(p.read_text("utf-8"))
    footage = [Path(str((i.get("kind") or {}).get("path"))).name for i in (data.get("items") or {}).values()
               if (i.get("kind") or {}).get("type") == "Footage"]
    assets = assets_of(file_id)
    rid = meta.get("render_file_id")
    spec = meta.get("spec") or {}
    pics = spec.get("pictures") or []
    tl = {x["id"]: x for x in spec.get("timeline") or [] if x.get("id")}
    scenes = [{"scene": i, "id": sc.get("id") or f"s{i}", **({"start": tl[sc["id"]]["start"], "end": tl[sc["id"]]["end"]}
                                                              if sc.get("id") in tl else {}),
               "headline": sc.get("headline"), "caption": sc.get("caption"),
               **({"chart": sc["chart"]} if sc.get("chart") else {}),
               **({"stat": sc["stat"]} if sc.get("stat") else {}),
               "picture": _pic_credit(pics[i]) if i < len(pics) and pics[i] else None}
              for i, sc in enumerate((spec.get("story") or {}).get("scenes") or [], 1)]
    rows = {x["id"]: x for x in spec.get("layer_timeline") or []}
    layers = [{**x, **({"on_timeline": f"{_tc(rows[x['id']]['start'])}–{_tc(rows[x['id']]['end'])}"} if x["id"] in rows
                       else {"on_timeline": "left out (it starts after the video ends)"})}
              for x in layer_listing((spec.get("story") or {}).get("layers") or [])]
    return {"file_id": file_id, "comps": _comps(data), "footage": footage, "missing": [f for f in footage if f not in assets],
            "video": meta.get("video"), "last_render": ({**_file_entry(rid), **({"draft": True} if meta.get("render_draft") else {})}
                                                        if rid and output_path(rid).exists() else None),
            "credits": meta.get("credit"), "editable_by_description": bool(spec),
            **({"title": spec["story"].get("title"), "scenes": scenes, "layers": layers, "options": spec.get("opts")} if spec else {}),
            **({"base": spec["base"], "how_to_edit": "changes={\"layers\": [{\"id\": \"L2\", …fields} | {\"id\": \"L2\", "
                "\"remove\": true}], \"add_layers\": […]}; see one layer with files(action=\"preview\", scene=\"L2\", frames=6)."}
               if spec.get("kind") == "overlay" else {}),
            **({"note": "It was also changed by hand in the editor; an edit rebuilds from the script and leaves those hand changes out."}
               if meta.get("hand_edited") else {})}


async def thumb(src: Path, raw: Path) -> None:
    """Bélyegkép: projektnél egy képkocka a címjelenetből, filmnél az 1. másodpercből."""
    ext = src.suffix.lstrip(".").lower()
    work = Path(tempfile.mkdtemp(prefix="vthumb-", dir=str(temp_path("").parent)))
    try:
        if ext == "ecproj":
            proj = stage(src.name, work)
            comp, t = _main_comp(proj), "1.6"
        else:
            name = f"movie.{ext}"
            os.symlink(src, work / name)
            code = (f'var f = app.project.importFile(new ImportOptions(new File({json.dumps(name)})));'
                    f'var c = app.project.items.addComp("Thumb", f.width, f.height, 1, Math.max(f.duration, 0.2), 30); c.layers.add(f);')
            (work / "t.jsx").write_text(code, "utf-8")
            await _run("script", "t.jsx", "--save-as", "project.ecproj", cwd=work, timeout=120)
            proj, comp, t = work / "project.ecproj", "Thumb", "1"
        await _run("render-frame", "--project", proj.name, "--comp", comp, "--time", t, "--max-side", "640",
                   "--out", "thumb.png", cwd=work, timeout=120)
        shutil.move(str(work / "thumb.png"), str(raw))
    finally:
        shutil.rmtree(work, ignore_errors=True)


# ---------- utólagos szerkesztés (az ügynök és az „Ask for a change”) ----------

async def transcribe_file(file_id: str, language: Optional[str] = None, step: Callable[..., None] = lambda *a, **k: None) -> dict[str, Any]:
    """Egy felvétel (videó vagy hang; overlay-projektnél az alapsáv) időbélyeges átirata + SRT-feliratfájl. Ebből köthető
    réteg egy mondathoz (at: {"say": …}), és ebből javasolhat az ügynök grafikát (melyik mondatnál mi segítene)."""
    from app.storage.naming import result_stem
    from app.tools import transcript
    p = crafts.existing(file_id)
    ext = p.suffix.lstrip(".").lower()
    if ext == "ecproj":
        spec = ((read_meta(file_id) or {}).get("spec") or {})
        base = spec.get("base")
        if not base:
            raise ValueError("This video was made from text — its words are its captions (video action=info). A transcript "
                             "is for your own footage or a sound file.")
        p = crafts.existing(assets_of(file_id).get(base["asset"]) or base.get("media_id") or base["file_id"])
    elif ext not in MOVIE_EXTS | AUDIO_EXTS:
        raise ValueError(f"A transcript needs a video or a sound file, not .{ext}.")
    step(0, 2, "listening (transcript)")
    tr = await transcript.transcribe(p, transcript_lang(language))
    step(1, 2, "writing the subtitles")
    subs = transcript.cues(tr)               # feliratméretű sorok a szóidőkből (a mondat a rétegkötéshez marad)
    sid = new_file_id("srt")
    output_path(sid).write_text(transcript.srt({"sentences": subs}), "utf-8")
    label = f"{result_stem(file_id, 'transcript')}.srt"
    write_meta(sid, source_file_id=file_id, operation="transcript", label=label)
    n = len(tr["sentences"])
    check = ({"status": "ok", "findings": [], "verified": [f"{n} sentence(s) heard ({tr['language']}), {len(tr['words'])} words with "
                                                           f"their times, from {_tc(tr['duration'])} of sound (model {tr['model']})."]}
             if n else {"status": "needs_check", "findings": ["No speech was heard in this file."], "verified": []})
    return {"file_id": sid, "name": label, "url": public_url(sid), "language": tr["language"], "duration": tr["duration"],
            "sentences": [{"at": _tc(x["start"]), "until": _tc(x["end"]), "start": x["start"], "end": x["end"], "text": x["text"]}
                          for x in tr["sentences"]],
            "subtitles": [{"at": _tc(x["start"]), "until": _tc(x["end"]), "start": x["start"], "end": x["end"], "text": x["text"]}
                          for x in subs],
            "how": 'Tie a layer to what is said: {"at": {"say": "words from a sentence"}, …} — it starts as they are said and '
                   'stays to the end of that sentence (or give dur / until). Subtitles in the Studio caption style: show the '
                   'person these sentences, correct them together, then video(action="overlay", file_id, subtitles=[{"start", '
                   '"end", "text"}…]) — or quickly, as heard: file_tools add-subtitles with this .srt.',
            "check": check}


def project_of(file_id: str) -> str:
    """A videó projektje: .ecproj maga, MP4/WebM/GIF esetén az, amiből renderelték."""
    if file_id.lower().endswith(".ecproj"):
        return file_id
    src = (read_meta(file_id) or {}).get("source_file_id")
    if src and str(src).lower().endswith(".ecproj"):
        return src
    raise ValueError("This video has no project behind it — put graphics on it with video(action=\"overlay\", file_id=…, "
                     "layers=[…]), or make a new one with the video maker (video action=make).")


def _img_spec(x: Any) -> Any:
    """Kép a szerkesztéshez: file_id | {"file_id"} | {"url"} (pl. a kliens saját generált képe) | {"stock": szavak}
    | {"generate": leírás} (AI-fotó: a webes felületen a mi FLUX-unk, kreditért; MCP-n a kliens a saját képeszközével
    készíti — 10-12 elv) | None / "none" (kép nélkül). Lista (legfeljebb 4): a jelenet hátterei egymás után, áttűnéssel."""
    if x in (None, "", "none", False):
        return None
    if isinstance(x, list):
        out = [y for y in (_img_spec(v) for v in x[:4]) if y]
        return out if len(out) > 1 else (out[0] if out else None)
    if isinstance(x, str):
        return {"url": x} if x.startswith("http") else x
    if isinstance(x, dict) and x.get("generate"):
        from app import tenancy
        if tenancy.client_makes_images.get():
            raise ValueError(tenancy.CLIENT_IMAGE_MSG)
        return {"generate": _clean(x["generate"], 400)}
    if isinstance(x, dict) and (x.get("file_id") or x.get("url") or x.get("stock")):
        return x.get("file_id") or ({"url": x["url"], **({"credit": x["credit"]} if x.get("credit") else {})} if x.get("url")
                                    else {"stock": _clean(x["stock"], 60)})
    raise ValueError(f"image: a file_id, a URL, {{\"stock\": \"words\"}}, {{\"generate\": \"a description\"}}, a list "
                     f"of up to 4 of these, or none — not {x!r}")


def _pic_credit(pic: Any) -> str:
    """Egy képhely forrása szövegben (képsornál: hány kép, és a forrásaik)."""
    if isinstance(pic, list):
        return f"{len(pic)} pictures: " + " · ".join((x or {}).get("credit") or "yes" for x in pic)
    return (pic or {}).get("credit") or "yes"


def _ai_photo(words: str) -> dict[str, str]:
    """AI-fotó leírása egy jelenet képkereső szavaiból (fotószerű, szöveg és logó nélkül — a cím úgyis rákerül)."""
    return {"generate": f"{words}, realistic photograph, natural light, no text, no logo, no watermark"}


def _slots(sp: Any) -> list[Any]:
    """Egy képhely: egy kép vagy egy képsor (a jelenet hátterei) — mindig listaként."""
    return list(sp) if isinstance(sp, list) else [sp]


# ---------- kép- és kameramozgás, háttér (10-17, Kommandant: „a kameramozgás (Ken Burns), zoom, a hátterek is cserélődhetnek”) ----------
# motion: a jelenet KÉPÉNEK mozgása (a régi ken_burns választás jelenetenként felülírható); camera: a jelenet képének
# ÉS ábrájának közös virtuális kamerája — egy pillanatra ráközelít egy részletre, aztán vissza; a jelenet címe, nagy
# száma és a feliratok a helyükön maradnak (10-17: a mezre közelítéskor a cím kicsúszott); background: kép nélküli
# jelenet háttere.
MOTIONS = {"auto": None, "zoom_in": "kb-in", "zoom_out": "kb-out", "pan_left": "left", "pan_right": "right", "pan_up": "up",
           "pan_down": "down", "still": "still", "ken_burns": "kb-in", "ken_burns_out": "kb-out", "gentle": "in"}
CAMERA_PRESETS = ("push_in", "pull_out", "drift")
BACKGROUNDS = ("gradient", "motion", "dark")
# fit: a kép kitölti a keretet (cover — a széle levágódhat) vagy egészében látszik, mögötte önmaga elmosva (contain);
# auto: a SAJÁT fotónál contain, ha az aránya nagyon eltér a keretétől (10-17: a fekvő csoportkép álló videóban
# csak 1–2 arcot mutatott — a közösségi videók bevett megoldása az elmosott kitöltés), a stockkép cover marad
FITS = ("auto", "cover", "contain")


def _xy01(v: Any) -> Optional[list[float]]:
    if isinstance(v, (list, tuple)) and len(v) == 2 and all(isinstance(c, (int, float)) for c in v):
        return [round(min(1.0, max(0.0, float(c))), 3) for c in v]
    return None


def _scene_motion(ch: dict[str, Any], sc: dict[str, Any]) -> None:
    if "motion" in ch:
        m = str(ch.get("motion") or "auto").lower().replace("-", "_").replace(" ", "_")
        if m not in MOTIONS:
            raise ValueError(f"motion: {' | '.join(MOTIONS)}")
        sc["motion"] = None if m == "auto" else m
    if "focus" in ch:
        sc["focus"] = _xy01(ch.get("focus"))
    if "zoom" in ch:
        try:
            sc["zoom"] = round(min(1.8, max(1.03, float(ch["zoom"]))), 3) if ch["zoom"] not in (None, "") else None
        except (TypeError, ValueError):
            raise ValueError("zoom: a number from 1.03 to 1.8 (how far the picture moves in)") from None
    if "camera" in ch:
        cam = ch.get("camera")
        if cam in (None, "", "none", False):
            sc["camera"] = None
        elif isinstance(cam, str) and cam.lower() in CAMERA_PRESETS:
            sc["camera"] = cam.lower()
        elif isinstance(cam, list):
            moves = []
            for mv in cam[:8]:
                if not isinstance(mv, dict):
                    continue
                at = timecode(mv.get("at", 0), "camera at")
                try:
                    z, d = float(mv.get("zoom", 1)), float(mv.get("dur", 0.8))
                except (TypeError, ValueError):
                    raise ValueError('camera: [{"at": seconds, "zoom": 1–3 (1 = the whole frame), "to"?: [x, y] 0–1, "dur"?: s}]') from None
                moves.append({"at": round(max(0.0, at), 2), "zoom": round(min(3.0, max(1.0, z)), 3), "dur": round(min(4.0, max(0.2, d)), 2),
                              **({"to": _xy01(mv.get("to"))} if _xy01(mv.get("to")) else {})})
            sc["camera"] = sorted(moves, key=lambda x: x["at"]) or None
        else:
            raise ValueError(f'camera: {" | ".join(CAMERA_PRESETS)} | [{{"at", "zoom", "to"?, "dur"?}}] | "none"')
    if "fit" in ch:
        ft = str(ch.get("fit") or "auto").lower()
        if ft not in FITS:
            raise ValueError(f"fit: {' | '.join(FITS)}")
        sc["fit"] = None if ft == "auto" else ft
    if "background" in ch:
        bg = str(ch.get("background") or "gradient").lower()
        if bg not in BACKGROUNDS and not re.fullmatch(r"#?[0-9a-f]{6}", bg):
            raise ValueError(f"background (a scene without a picture): {' | '.join(BACKGROUNDS)} | \"#RRGGBB\"")
        sc["background"] = bg


def _own_picture(sp: Any) -> bool:
    """A személy saját képe (file_id, URL, a szerkesztésben megtartott saját kép) — nem stock / AI-kép."""
    if isinstance(sp, dict) and sp.get("asset"):
        sp = sp.get("src")
    return isinstance(sp, str) or (isinstance(sp, dict) and bool(sp.get("url") or sp.get("file_id")) and not sp.get("stock"))


def _auto_fits(names: list[Any], specs: list[Any], work: Path, frame: float) -> dict[str, str]:
    """fit=auto: a saját fotó egészében látszik (elmosott kitöltéssel), ha az aránya 1,5-szeresnél jobban eltér a keretétől."""
    from PIL import Image
    out = {}
    for slot, sp in zip(names, specs):
        sps = _slots(sp) if isinstance(sp, list) else [sp] * len(_slots(slot))
        for nm, one in zip(_slots(slot), sps):
            if not nm or not _own_picture(one):
                continue
            try:
                with Image.open(work / nm) as im:
                    a = im.width / im.height
            except Exception:
                continue
            if max(a / frame, frame / a) > 1.5:
                out[nm] = "contain"
    return out


def _camera_keys(cam: Any, dur: float, W: int, H: int) -> Optional[list[list[float]]]:
    """A jelenet virtuális kamerája → kulcskockák [[t, nagyítás, x, y]] (a jelenet-réteg helye a fő kompozícióban).
    A nagyított kép mindig kitölti a keretet (a hely úgy korlátozva), így szél soha nem látszik."""
    if not cam:
        return None

    def state(z: float, to: Optional[list[float]]) -> list[float]:
        fx, fy = (to or [0.5, 0.5])
        px, py = W / 2 - (fx * W - W / 2) * z, H / 2 - (fy * H - H / 2) * z
        lx, ly = W * (z - 1) / 2, H * (z - 1) / 2
        return [round(z, 4), round(min(W / 2 + lx, max(W / 2 - lx, px)), 1), round(min(H / 2 + ly, max(H / 2 - ly, py)), 1)]
    if cam == "push_in":
        return [[0.0, *state(1.0, None)], [round(dur, 3), *state(1.1, None)]]
    if cam == "pull_out":
        return [[0.0, *state(1.1, None)], [round(dur, 3), *state(1.0, None)]]
    if cam == "drift":
        return [[0.0, *state(1.06, [0.47, 0.5])], [round(dur, 3), *state(1.06, [0.53, 0.5])]]
    keys, cur = [[0.0, *state(1.0, None)]], state(1.0, None)
    for mv in cam:
        a = min(mv["at"], max(0.0, dur - 0.3))
        b = min(dur, a + mv["dur"])
        keys.append([round(a, 3), *cur])
        cur = state(mv["zoom"], mv.get("to"))
        keys.append([round(b, 3), *cur])
    return keys


def _scene_from(ch: dict[str, Any], source: str, notes: list[str], base: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    """Egy jelenet (a modell terve, a kliens jelenetterve vagy egy szerkesztés) tisztítva: szöveghatárok (a rövidülés
    a check-be kerül), ábra a forrásszámokkal ellenőrizve, stabil id, opcionális fix hossz."""
    sc = dict(base or {"headline": "", "caption": "", "image_query": "", "stat": None, "chart": None})
    if ch.get("id"):
        sc["id"] = re.sub(r"[^\w-]", "", str(ch["id"]))[:20] or sc.get("id")
    if "headline" in ch:
        sc["headline"] = _clean(ch["headline"], 60, "Headline")
    if "caption" in ch:
        sc["caption"] = _clean(ch["caption"], 200, "Caption")
    if ch.get("image_query"):
        sc["image_query"] = _clean(ch["image_query"], 60)
    if "duration" in ch:
        try:
            sc["duration"] = min(30.0, max(2.0, float(ch["duration"]))) if ch["duration"] not in (None, "", "auto") else None
        except (TypeError, ValueError):
            pass
    _scene_motion(ch, sc)
    if "scrim" in ch:                             # a jelenet saját sötétítése (0–1; None = a videóé)
        sc["scrim"] = scrim_of(ch["scrim"], "scene scrim")
    if "stat" in ch:
        st = ch["stat"]
        sc["stat"] = ({"value": _clean(st.get("value"), 10, "Stat value"), "label": _clean(st.get("label"), 48, "Stat label")}
                      if isinstance(st, dict) and st.get("value") else None)
        if sc["stat"]:
            sc["chart"] = None
    if "chart" in ch:
        cc = ch["chart"]
        if isinstance(cc, dict) and not cc.get("type") and (base or {}).get("chart"):
            from app.tools import diagrams        # ábránál a csúcsok / élek id (from+to) szerint frissülnek
            cc = diagrams.merge(base["chart"], cc) if base["chart"].get("type") in diagrams.TYPES else {**base["chart"], **cc}
        chart, note = _chart_from_plan(cc, source) if cc else (None, None)
        for n in (note if isinstance(note, list) else [note] if note else []):
            # ⚠ = figyelmeztetés (a check findings-be): ami kimaradt, nem mehet el csendben (10-12: a térkép tűi kiestek,
            # a check mégis ok volt — a megjegyzés a „verified” közé került)
            notes.append(n if n.startswith("⚠ ") or chart else "⚠ " + n)
        if cc and chart is None and not note:
            t_ = cc.get("type") if isinstance(cc, dict) else None
            notes.append(f"⚠ The suggested {t_ or 'chart'} “{(cc.get('title') if isinstance(cc, dict) else '') or ''}” was left out: "
                         f"{'unknown type' if t_ not in _CHART_TYPES else 'not enough usable items (check items / elements)'}.")
        sc["chart"] = chart
        if chart:
            sc["stat"] = None
    if (sc.get("chart") or {}).get("nodes") and ("chart" in ch or "duration" in ch):
        from app.tools import diagrams
        n_ = diagrams.duration_note(sc["chart"], sc.get("duration"))
        if n_:
            notes.append(n_)
    return sc


async def edit(file_id: str, changes: dict[str, Any], *, extra_source: str = "",
               progress: Optional[Callable[..., None]] = None, alpha: Any = None, render: bool = True,
               draft: bool = False) -> dict[str, Any]:
    tok, fit_tok = _TRUNC.set([]), _FIT.set([])
    try:
        return await _edit(file_id, changes, extra_source=extra_source, progress=progress, alpha=alpha, render=render, draft=draft)
    finally:
        _TRUNC.reset(tok)
        _FIT.reset(fit_tok)


async def _edit(file_id: str, changes: dict[str, Any], *, extra_source: str = "",
                progress: Optional[Callable[..., None]] = None, alpha: Any = None, render: bool = True,
                draft: bool = False) -> dict[str, Any]:
    """Egy elkészült videó módosítása a forgatókönyvéből → ÚJ változat (projekt + MP4). changes:
    {"scenes": [{"scene": n (1-től), "headline"?, "caption"?, "image"?, "stat"?, "chart"?, "remove"?}],
     "add": [{"after": n (0 = az első elé), "headline", "caption"?, "image"? | "image_query"?, "stat"?, "chart"?}],
     "title"?, "kicker"?, "title_image"?, "end_title"?, "end_subtitle"?,
     "format"?, "duration"?, "duration_mode"?, "ken_burns"?, "accent"?, "music_file_id"? | "none", "voice_file_id"? | "none",
     "layers"?: [{"id": "L2", …} | {"id": "L2", "remove": true}], "add_layers"?: […]}.
    A grafikon számai a forrásból (vagy az extra_source-ból — pl. a kérés szövegéből) jöhetnek. Az overlay-projekt
    (a felhasználó felvétele + rétegek) ugyanígy, csak rétegekkel."""
    step = progress or (lambda *a, **k: None)
    pid = project_of(file_id)
    meta = read_meta(pid) or {}
    spec = meta.get("spec")
    if not spec:
        raise ValueError("This video wasn't made by the video maker, so it can't be changed by description — change it in "
                         "the video editor (edit_url), or make a new one with video action=make.")
    if not isinstance(changes, dict) or (not changes and alpha in (None, False)):
        raise ValueError("Say what should change (changes).")
    if spec.get("kind") == "overlay":
        res = await _edit_overlay(pid, spec, changes, extra_source=extra_source, step=step, alpha=alpha, render=render, draft=draft)
        res["notes"] = ["Rebuilt from the layer list: changes made by hand in the video editor are not in this version "
                        "(they stay in the previous one)."] if meta.get("hand_edited") else []
        return res
    alpha_kinds(alpha)
    duration_mode(changes.get("duration_mode"), changes.get("duration"))
    story = copy.deepcopy(spec["story"])
    story["scenes"] = _with_ids(story["scenes"])          # régi projekt: s1, s2… a mostani sorrendben
    pics = copy.deepcopy(spec.get("pictures") or [])
    pics += [None] * (len(story["scenes"]) + 1 - len(pics))
    opts = dict(spec.get("opts") or {})
    source = (spec.get("source") or "") + "\n" + (extra_source or "")
    notes: list[str] = []
    keep1 = lambda pic: ({"asset": pic["asset"], "credit": pic.get("credit"), "src": pic.get("src")} if pic else None)   # noqa: E731
    keep = lambda pic: [keep1(x) for x in pic if x] if isinstance(pic, list) else keep1(pic)   # noqa: E731

    for key, lim in (("title", 70), ("kicker", 28)):
        if key in changes:
            story[key] = _clean(changes[key], lim, key.capitalize())
    end = dict(story.get("end") or {})
    if "end_title" in changes:
        end["title"] = _clean(changes["end_title"], 60, "End title")
    if "end_subtitle" in changes:
        end["subtitle"] = _clean(changes["end_subtitle"], 90, "End subtitle")
    story["end"] = end

    entries = [{"scene": sc, "pic": keep(pics[i + 1]), "remove": False} for i, sc in enumerate(story["scenes"])]
    title_pic = _img_spec(changes["title_image"]) if "title_image" in changes else keep(pics[0])
    n0 = len(entries)
    ids = [e["scene"].get("id") for e in entries]
    for ch in changes.get("scenes") or []:
        if ch.get("id"):                        # stabil azonosító (s3) — beszúrás/törlés után is ugyanaz a jelenet
            if ch["id"] not in ids:
                raise ValueError(f"There is no scene {ch['id']} (scenes: {', '.join(ids)}).")
            k = ids.index(ch["id"])
        else:
            try:
                k = int(ch.get("scene")) - 1
            except (TypeError, ValueError):
                raise ValueError("scenes[]: give the scene's id (\"s3\") or number (1 = the first scene after the title card)")
            if not 0 <= k < n0:
                raise ValueError(f"There is no scene {k + 1} (the video has {n0}).")
        e = entries[k]
        if ch.get("remove"):
            e["remove"] = True
            continue
        e["scene"] = _scene_from(ch, source, notes, e["scene"])
        if "images" in ch:
            e["pic"] = _img_spec(ch["images"])
        elif "image" in ch:
            e["pic"] = _img_spec(ch["image"])
        elif ch.get("image_query"):
            e["pic"] = {"stock": _clean(ch["image_query"], 60)}
    adds: dict[int, list[dict[str, Any]]] = {}
    for a in changes.get("add") or []:
        aft = a.get("after", n0)
        after = (ids.index(aft) + 1) if isinstance(aft, str) and aft in ids else 0 if aft == "title" else max(0, min(n0, int(aft)))
        sc = _scene_from(a, source, notes)
        # keresőszó híján a videó (angol) keresőszava — a magyar cím a képtárban mellényúl („Mit tehetsz a méhekért?”)
        near = next((e["scene"].get("image_query") for e in entries[max(0, after - 1):] if e["scene"].get("image_query")), "")
        pic = _img_spec(a.get("images") or a["image"]) if (a.get("image") or a.get("images")) else {"stock": _clean(a.get("image_query") or near or story.get("image_query")
                                                                            or sc["headline"] or story["title"], 60)}
        adds.setdefault(after, []).append({"scene": sc, "pic": pic, "remove": False})
    final = list(adds.get(0, []))
    for i, e in enumerate(entries, 1):
        if not e["remove"]:
            final.append(e)
        final += adds.get(i, [])
    if not final:
        raise ValueError("A video needs at least one scene.")

    if "format" in changes:
        opts["fmt"] = fmt_name(changes["format"])
    if "duration" in changes:
        d = changes["duration"]
        opts["total"] = None if d in (None, "", "auto") else float(min(max(float(d), MIN_TOTAL), 180))
        opts["requested"] = None if d in (None, "", "auto") else float(d)
        opts["duration_mode"] = None                     # új hossz: alapból fix (hacsak a duration_mode mást nem mond)
    if changes.get("duration_mode"):
        opts["duration_mode"] = changes["duration_mode"]
    if alpha is not None:
        opts["alpha"] = alpha
    layer_src = ((spec.get("layer_source") or "") + "\n" + (extra_source or "")).strip()
    lnotes: list[str] = []
    story["layers"] = apply_layer_changes(story.get("layers") or [], changes, layer_src, lnotes)
    for key in ("ken_burns", "accent", "narration", "narration_voice", "music"):
        if key in changes:
            opts[key] = changes[key]
    if "scrim" in changes:
        opts["scrim"] = scrim_of(changes["scrim"])
    if "music" in changes:                               # a tár zenéje (vagy "none") a saját zene helyére
        opts["music_asset"], opts["music_file_id"] = None, None
    elif changes.get("music_file_id") not in (None, "", "none"):
        opts["music"] = None
    if changes.get("voice_file_id") not in (None, "", "none"):
        opts["narration"] = "none"                       # saját hang jött: az AI-narráció kimarad
    for key, base in (("music_file_id", "music"), ("voice_file_id", "voice")):
        if key in changes:
            v = changes[key]
            opts[f"{base}_asset"] = None
            opts[key] = None if v in (None, "", "none") else v

    story["scenes"] = _with_ids([e["scene"] for e in final])
    specs: list[Any] = [title_pic] + [e["pic"] for e in final]
    reshape = opts["fmt"] != (spec.get("opts") or {}).get("fmt")
    if reshape:          # más képarány: a képeket az eredeti forrásból újra (különben nagyítva, elmosódva)
        src_of = lambda sp: (sp.get("src") or sp) if isinstance(sp, dict) and sp.get("asset") else sp   # noqa: E731
        specs = [[src_of(x) for x in sp] if isinstance(sp, list) else src_of(sp) for sp in specs]
    if any(isinstance(sp, dict) and sp.get("stock") for slot in specs for sp in _slots(slot)):
        step(1, 5, "finding pictures")
    await _pick_stock(specs, lambda i: story["title"] if i == 0 else story["scenes"][i - 1]["headline"])
    from app.storage.naming import result_stem
    opts["layer_source"] = layer_src
    res = await _assemble(story, specs, opts, chart_notes=_kept_chart_notes(spec, story) + notes, source_text=source,
                          source_file_id=pid, step=step, parent=pid, label_stem=result_stem(pid, "edited"), render=render, draft=draft,
                          operation="video edit", layer_notes=_kept_layer_notes(spec, story["layers"]) + lnotes)
    res["notes"] = notes + (["Rebuilt from the video's script: changes made by hand in the video editor are not in this "
                             "version (they stay in the previous one)."] if meta.get("hand_edited") else [])
    return res


async def _edit_overlay(pid: str, spec: dict[str, Any], changes: dict[str, Any], *, extra_source: str,
                        step: Callable[..., None], alpha: Any = None, render: bool = True, draft: bool = False) -> dict[str, Any]:
    """Overlay-projekt célzott szerkesztése: rétegek id alapján (layers / add_layers), akcentus, alpha → ÚJ változat
    (projekt + MP4); a felvétel ugyanaz a fájl marad, a változatlan rétegek érintetlenek."""
    from app.storage.naming import result_stem
    allowed = {"layers", "add_layers", "accent", "title", "alpha", "music", "camera", "subtitles", "end_card", "end_title",
               "end_subtitle", "transcript"}
    other = sorted(k for k in set(changes) - allowed if changes[k] not in (None, "", [], {}))
    if other:
        raise ValueError(f"This video is your own footage with layers on it: change its layers (changes.layers with an id, "
                         f"changes.add_layers), accent or title — not {', '.join(other)}. To recut the footage itself, use "
                         "file_tools (trim-video…) and run overlay again.")
    opts = dict(spec.get("opts") or {})
    if "camera" in changes:
        opts["camera"] = _footage_camera(changes["camera"])
    opts["subtitles"] = subtitle_cues(changes["subtitles"]) if "subtitles" in changes else \
        ((spec.get("story") or {}).get("subtitles") or [])
    for k in ("accent", "title", "music"):
        if k in changes:
            opts[k] = changes[k]
    if alpha is not None or "alpha" in changes:
        opts["alpha"] = changes.get("alpha", alpha)
    alpha_kinds(opts.get("alpha"))
    source = ((spec.get("layer_source") or "") + "\n" + (extra_source or "")).strip()
    notes: list[str] = []
    if _says(changes.get("layers")) or _says(changes.get("add_layers")):        # „amikor elhangzik…” a szerkesztésben is
        footage = crafts.existing(assets_of(pid).get(spec["base"]["asset"]) or spec["base"].get("media_id") or spec["base"]["file_id"])
        topts = transcript_opts(changes.get("transcript") or opts.get("transcript"))
        changes = {**changes, "layers": await said_to_times(changes.get("layers"), footage, notes, step, topts=topts),
                   "add_layers": await said_to_times(changes.get("add_layers"), footage, notes, step, topts=topts)}
    layers = apply_layer_changes(copy.deepcopy((spec.get("story") or {}).get("layers") or []), changes, source, notes)
    base = dict(spec["base"])
    card = base.get("end_card")
    if any(k in changes for k in ("end_card", "end_title", "end_subtitle")):
        ec = changes.get("end_card")
        if "end_card" in changes and not (isinstance(ec, dict) and ec):          # {} / "none" / null: ki; "Cím": új
            merged = ec
        else:                                                                    # részleges: a régi kártya mezői maradnak
            merged = {**(card or {}), **(ec if isinstance(ec, dict) else {})}
            merged.update({k_: changes[f"end_{k_}"] for k_ in ("title", "subtitle") if changes.get(f"end_{k_}")})
        card = end_card_of({"end_card": merged})
    if not layers and not opts["subtitles"] and not card:
        raise ValueError("That would leave no layers and no subtitles — keep at least one, or use the original video as it is.")
    base = await _with_end_card(base, card)                    # új / más hosszú kártya: a felvétel újra meghosszabbítva
    return await _assemble_overlay(base, layers, opts, source=source,
                                   layer_notes=_kept_layer_notes(spec, layers) + notes, step=step,
                                   label_stem=result_stem(pid, "edited"), operation="video edit", parent=pid, render=render,
                                   draft=draft)


_EDIT_PLAN = """You change a short video made of scenes. Reply with ONE json object of changes — only what the request asks:
{"summary": "what you changed, one short sentence in the request's language",
 "scenes": [{"scene": n, "headline"?: str, "caption"?: str, "image"?: {"stock": "2–4 plain English words"} | "none",
             "stat"?: {"value": str, "label": str} | null, "chart"?: <chart> | null, "remove"?: true}],
 "add": [{"after": n, "headline": str, "caption": str, "image_query": "2–4 plain English words", "stat"?: ..., "chart"?: ...}],
 "title"?: str, "end_title"?: str, "end_subtitle"?: str, "title_image"?: {"stock": "words"},
 "format"?: "vertical|landscape|square", "duration"?: seconds, "ken_burns"?: "some|all|off", "accent"?: "#RRGGBB",
 "layers": [{"id": "L2", "at"?: seconds, "dur"?: seconds, "name"?, "role"?, "text"?, "subtitle"?, "value"?, "label"?,
             "chart"?: <chart>, "remove"?: true}],
 "add_layers": [{"at": seconds, "dur": seconds, "type": "lower_third|title|chapter|source|stat|quote|chart|fullscreen|split", …}]}
A fullscreen layer (covers the picture, the sound goes on) or a split layer (+ "side": left|right|top|bottom) holds ONE of:
"image": {"stock": "2–4 plain English words"} (a real photo) | {"generate": "an English description"} (an AI picture —
when the person asks for an AI / generated picture) | a file_id from the request · "chart": <chart> · "text" (+ "subtitle").
Layers are time-coded graphics over the video (listed below with their ids); change them by id. A video of the person's
own footage has only layers: change or add layers there, never scenes. When the request points to what is said in the
footage ("what I ask for in the video", "the graphics I mention"), read the speech listed below and add every graphic the
speaker asks for or describes, each at the time it is said (it may last until the next one starts): e.g. "a table" → a
table chart, "an infographic above me" → a chart layer, "flying / animated infographics" → an animated chart (race,
waterfall, lines…) — with the speaker's own words and numbers, never invented ones.
<chart> = {"type": "column|bar|line|area|donut|ring|icons|timeline|steps|counter|compare|progress|gauge|funnel|proscons|quote|table|waterfall|slope|scatter|band|lines|race|sankey", "title": str, "unit": str, "total": number|null, "min"?: number, "max"?: number,
           "items": [{"label": str, "value": number|null, "text": str}]}
  (table: {"type": "table", "title": str, "columns": [2–5 headings], "rows": [[cells], … up to 7 rows]} — words or numbers)
(waterfall, slope, scatter, band, lines, race and sankey have their own fields — labels, series, from/to, x/y/size, low/high, flows;
to change one, send a chart WITHOUT "type" with only the fields that change.)
A scene with a diagram (decision_tree | swimlane | network, listed with its node ids): to change it send a chart WITHOUT
"type" holding only what changes — {"nodes": [{"id", "text"?, "kind"?, "remove"?: true} | a new node],
"edges": [{"from", "to", "kind"?, "certainty"?: "certain|uncertain|disputed", "label"?, "remove"?: true}], "pace"?: seconds
per reveal step}; nodes match by id, edges by from+to; new ones are added.
Scene numbers are the ones listed below (1 = the first scene after the title card). Numbers in charts and stats must
come from the video's source or from the request itself — never invent them. Keep the language of the video.

The video now:
{video}

Request: {request}"""


def _speech(spec: dict[str, Any]) -> list[dict[str, Any]]:
    """A saját felvétel beszéde időkóddal: a feliratok (ha a videónak van), különben az alapfelvétel átirata (gyorsítótárból)."""
    if spec.get("kind") != "overlay":
        return []
    subs = (spec.get("story") or {}).get("subtitles") or []
    return [c for c in subs if isinstance(c, dict) and c.get("text")][:300]


def _describe(spec: dict[str, Any]) -> str:
    st, o = spec["story"], spec.get("opts") or {}
    pics = spec.get("pictures") or []
    if spec.get("kind") == "overlay":
        b = spec.get("base") or {}
        lines = [f"the person's own footage: {b.get('width')}×{b.get('height')}, {_tc(b.get('duration') or 0)} — layers only"]
    else:
        lines = [f"title: {st.get('title')} · format {o.get('fmt')} · {o.get('total') or 'auto'} s · Ken Burns {o.get('ken_burns')}"]
    lines += [f"layer {x['id']} ({x['type']}) {x['timecode']}: {x['text']}" for x in layer_listing(st.get("layers") or [])]
    said = _speech(spec)
    if said:                  # saját felvétel: ami elhangzik, időkóddal (a „tedd bele, amit a videóban kérek”-hez)
        lines.append("what is said in the footage (start–end: words):")
        lines += [f"  {_tc(c['start'])}–{_tc(c['end'])}: {c['text']}" for c in said]
    for i, sc in enumerate(st.get("scenes") or [], 1):
        pic = pics[i] if i < len(pics) else None
        what = f"chart {sc['chart']['type']} “{sc['chart']['title']}”" if sc.get("chart") else (f"stat {sc['stat']['value']}" if sc.get("stat") else "")
        from app.tools import diagrams
        if (sc.get("chart") or {}).get("type") in diagrams.TYPES:    # a sankey csomópontjai nem ábra-dobozok (nincs id-jük)
            what += " — " + diagrams.describe(sc["chart"])
        lines.append(f"{i}. {sc.get('headline')} — caption: {sc.get('caption')}" + (f" — {what}" if what else "")
                     + (f" — picture: {_pic_credit(pic)}" if pic else " — no picture"))
    if st.get("scenes"):
        lines.append(f"end: {(st.get('end') or {}).get('title')} / {(st.get('end') or {}).get('subtitle')}")
    return "\n".join(lines)


async def revise(file_id: str, request: str, progress: Optional[Callable[..., None]] = None) -> dict[str, Any]:
    """„Ask for a change" egy videóra: a kérésből változtatás-lista (DeepSeek, JSON) → edit."""
    from app.llm.siliconflow import _call_model, _strip_json_fences
    pid = project_of(file_id)
    spec = (read_meta(pid) or {}).get("spec")
    if not spec:
        raise ValueError("This video wasn't made by the video maker — change it in the video editor instead.")
    (progress or (lambda *a, **k: None))(0, 5, "planning the change")
    prompt = _EDIT_PLAN.replace("{video}", _describe(spec)).replace("{request}", request)
    plan_, last = None, None
    for model in (settings.AGENT_MODEL, settings.AGENT_MODEL_FALLBACK):
        try:
            plan_ = json.loads(_strip_json_fences(await _call_model(model, [{"role": "user", "content": prompt}], thinking=False)))
            break
        except Exception as e:
            last = e
            log.warning("videó-szerkesztési terv %s modellel nem sikerült: %s", model, e)
    if not isinstance(plan_, dict):
        raise ValueError(f"Couldn't work out the change ({last}) — try saying it differently.")
    summary = _clean((plan_ or {}).pop("summary", "") if isinstance(plan_, dict) else "", 200)
    said = " ".join(c["text"] for c in _speech(spec))       # a felvételben elhangzott számok is a forrás részei
    res = await edit(pid, plan_ if isinstance(plan_, dict) else {}, extra_source=(request + "\n" + said).strip(),
                     progress=progress)
    res["summary"] = summary or "Video changed."
    return res


# ---------- videó egy diára ----------

async def to_pptx(file_id: str, deck_file_id: Optional[str] = None, after_slide: Optional[int] = None,
                  title: Optional[str] = None) -> dict[str, Any]:
    """A videó egy diára: új egydiás pakli, vagy egy meglévő pakli új diája (after_slide után, alapból a végén).
    A dián a videó maga (beágyazott MP4, az első kockája a borítóképe), a dia méretéhez igazítva."""
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from app.storage.naming import result_stem, root_name
    from app.tools import previews
    vid = file_id
    if file_id.lower().endswith(".ecproj"):                 # projekt → a legutóbbi renderje (ha nincs, renderel)
        rid = (read_meta(file_id) or {}).get("render_file_id")
        if not rid or not output_path(rid).exists():
            rid = (await render(file_id))["file_id"]
        vid = rid
    vpath = crafts.existing(vid)
    if vpath.suffix.lower() not in (".mp4", ".m4v", ".mov"):
        raise ValueError("Slides take an MP4 video (or a video project, which is rendered first).")
    import cv2
    cap = cv2.VideoCapture(str(vpath))
    vw, vh = cap.get(cv2.CAP_PROP_FRAME_WIDTH) or 16, cap.get(cv2.CAP_PROP_FRAME_HEIGHT) or 9
    secs = (cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0) / (cap.get(cv2.CAP_PROP_FPS) or 30)
    cap.release()
    # borítókép: a címkártya, amikor a cím már teljesen látszik (~1,6 s) — fix arány átúszásra eshet
    poster = temp_path(new_file_id("png"))
    await asyncio.to_thread(previews._video_frame, vpath, min(0.5, 1.6 / secs) if secs > 0 else 0.1, poster)
    if deck_file_id:
        if not deck_file_id.lower().endswith(".pptx"):
            raise ValueError("deck_file_id: a PowerPoint deck (.pptx)")
        prs = Presentation(str(crafts.existing(deck_file_id)))
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)          # 16:9
    SW, SH = int(prs.slide_width), int(prs.slide_height)
    from app.tools import pptx_deck
    n_before = len(prs.slides)
    # a minta: az a dia, ami után az új jön (alapból az utolsó) — annak hátterét és címstílusát kapja az új dia
    ref_i = int(after_slide) - 1 if after_slide is not None else n_before - 1
    ref = list(prs.slides)[max(0, min(n_before - 1, ref_i))] if n_before else None
    layout = ref.slide_layout if ref is not None else \
        min(prs.slide_layouts, key=lambda lo: (len(lo.placeholders), "blank" not in lo.name.lower()))
    slide = prs.slides.add_slide(layout)
    for ph in list(slide.placeholders):                     # üres helyőrző ne maradjon a dián
        ph._element.getparent().remove(ph._element)
    took_bg = pptx_deck.inherit_look(slide, ref)
    top_room = 0
    if title and title.strip():
        top_room = int(SH * 0.16)
        st = pptx_deck.title_style(ref) if ref is not None else {}
        tb = slide.shapes.add_textbox(int(SW * 0.06), int(SH * 0.04), int(SW * 0.88), int(SH * 0.1))
        tb.text_frame.text = title.strip()[:120]
        run = tb.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(min(32, st.get("size") or 28))
        run.font.bold = True if st.get("bold") is None else st["bold"]
        if st.get("font"):
            run.font.name = st["font"]
        color = pptx_deck.readable(st.get("color"), slide, prs)
        if color:
            run.font.color.rgb = pptx_deck._rgb(color)
    pad = 0 if (not deck_file_id and not top_room and vw / vh >= SW / SH * 0.98) else int(min(SW, SH) * 0.04)
    box_w, box_h = SW - 2 * pad, SH - top_room - 2 * pad
    k = min(box_w / vw, box_h / vh)
    w, h = int(vw * k), int(vh * k)
    left, top = (SW - w) // 2, top_room + pad + (box_h - h) // 2
    # a dián lévő videó iPhone-on (Keynote) is játszódjon: ha kell, iPhone-biztos másolat kerül be
    from app.tools import ios_safe
    embed = temp_path(new_file_id(vpath.suffix.lstrip(".").lower()))
    shutil.copy(vpath, embed)
    await ios_safe.finalize(embed)
    movie = slide.shapes.add_movie(str(embed), left, top, w, h, poster_frame_image=str(poster),
                                   mime_type="video/quicktime" if embed.suffix.lower() == ".mov" else "video/mp4")
    movie.name = "Video"
    from app.tools import pptx_media              # a PowerPoint saját időzítése: magától indul, koppintásra szünet
    pptx_media.autoplay(slide._element, movie.shape_id, int(secs * 1000))
    n = len(prs.slides)
    pos = n - 1
    if deck_file_id and after_slide is not None:            # a dia helye: after_slide után
        pos = max(0, min(n - 1, int(after_slide)))
        lst = prs.slides._sldIdLst
        el = lst[-1]
        lst.remove(el)
        lst.insert(pos, el)
    if deck_file_id:                                         # a pakli „n / N” oldalszámai az új sorrend szerint
        from app.tools.pptx_deck import renumber_footers
        renumber_footers(prs)
    out = new_file_id("pptx")
    await asyncio.to_thread(prs.save, str(output_path(out)))
    poster.unlink(missing_ok=True)
    embed.unlink(missing_ok=True)
    label = (result_stem(deck_file_id, "with video") if deck_file_id else root_name(display_label(vid), "Video")) + ".pptx"
    write_meta(out, source_file_id=deck_file_id or vid, operation="video on a slide", label=label)
    # ellenőrzés (közös check-szerződés): a mentett fájlban tényleg ott a videó, a megadott dián
    chk_prs = Presentation(str(output_path(out)))
    media = [sh for sh in chk_prs.slides[pos].shapes if sh.shape_type == 16]       # MSO_SHAPE_TYPE.MEDIA
    look = ([f"The new slide has the deck's look: {'background' if took_bg else 'layout'}{' and title style' if title else ''} "
             "from the slide next to it."] if ref is not None else [])
    check = ({"status": "ok", "findings": [], "verified": [f"The MP4 ({secs:.0f} s) is embedded on slide {pos + 1} of {len(chk_prs.slides)}, "
                                                          "with the title card as its poster.", *look]}
             if media else {"status": "failed", "findings": ["The video didn't end up on the slide."], "verified": []})
    return {"file_id": out, "url": public_url(out), "label": label, "slide": pos + 1, "video_file_id": vid, "check": check,
            "note": "The video is embedded and starts by itself when the slide shows (tap or click pauses it) in "
                    "PowerPoint and Keynote, iPhone included; the browser's slide editor shows its first frame."}


# ---------- előnézet jelenetenként ----------

def _ids_hint(spec: dict[str, Any]) -> str:
    sc = [str(x["id"]) for x in spec.get("timeline") or []]
    ly = [str(x["id"]) for x in spec.get("layer_timeline") or []]
    return "; ".join(x for x in ((f"scenes: {', '.join(sc)}" if sc else ""), (f"layers: {', '.join(ly)}" if ly else "")) if x) or "none"


def layer_still(row: dict[str, Any]) -> float:
    """A réteg „kész” pillanata: a beúszás és az ábra animációi lefutottak, a kiúszás még nem kezdődött."""
    return round(max(row["start"] + 0.1, row["end"] - 0.6), 2)


def frame_time(file_id: str, scene: Any = None, at: Optional[float] = None) -> tuple[str, float, str]:
    """→ (a renderelt videó file_id-ja, az időpont másodpercben, címke). scene: id (s3), sorszám, "title", "end", vagy egy
    réteg id-ja (L2); a jelenet / réteg VÉGE előtt (az animációk már lefutottak); at: tetszőleges másodperc."""
    if file_id.lower().endswith(".ecproj"):              # projekt → a legutóbbi renderje
        pid, vid = file_id, (read_meta(file_id) or {}).get("render_file_id") or ""
        if not vid or not output_path(vid).exists():
            raise ValueError("This video project has no render yet — render it first (video action=render).")
    else:                                                 # videó → a projektje (ha a videókészítő csinálta)
        src = str((read_meta(file_id) or {}).get("source_file_id") or "")
        pid, vid = (src if src.lower().endswith(".ecproj") else ""), file_id
    spec = ((read_meta(pid) or {}).get("spec") or {}) if pid else {}
    if at is not None:
        return vid, max(0.0, float(at)), f"at {float(at):.1f} s"
    tl, lrows = spec.get("timeline") or [], spec.get("layer_timeline") or []
    if not tl and not lrows:
        raise ValueError("This video has no scene timeline (made before it was recorded) — use at=seconds instead.")
    key = str(scene).strip().lower() if scene is not None else ("title" if tl else str(lrows[0]["id"]).lower())
    if key.isdigit() and tl:
        key = (tl[int(key)]["id"] if 0 < int(key) < len(tl) - 1 else "")
    hit = next((x for x in tl if str(x.get("id")).lower() == key), None)
    if hit:
        return vid, scene_still(tl, tl.index(hit)), f"{hit['id']} ({hit['start']:.1f}–{hit['end']:.1f} s)"
    lh = next((x for x in lrows if str(x.get("id")).lower() == key), None)
    if lh:
        return vid, layer_still(lh), f"{lh['id']} ({lh['type']}, {lh['start']:.1f}–{lh['end']:.1f} s)"
    raise ValueError(f"No scene or layer {scene!r} — {_ids_hint(spec)}. Or use at=seconds.")


def _kept_chart_notes(spec: dict[str, Any], story: dict[str, Any]) -> list[str]:
    """A mentett ábra-jegyzetek közül csak a VÁLTOZATLAN ábráké öröklődik (10-13: a cím javítása után a check a régi
    „Counter “fős kockázati lista …”” és az új címet is felsorolta)."""
    old = {sc.get("id"): sc.get("chart") for sc in (spec.get("story") or {}).get("scenes") or [] if isinstance(sc, dict)}
    new = {sc.get("id"): sc.get("chart") for sc in story.get("scenes") or [] if isinstance(sc, dict)}
    gone = {(c or {}).get("title") for sid, c in old.items() if c and new.get(sid) != c} - {None, ""}
    return [n for n in dict.fromkeys(spec.get("chart_notes") or []) if not any(f"“{t}”" in n for t in gone)]


def scene_still(tl: list[dict[str, Any]], k: int) -> float:
    """A jelenet „kész” pillanata: az animációk lefutottak, de a KÖVETKEZŐ jelenet áttűnése még nem kezdődött
    (10-13: a jelenet vége előtt 0,4 s a 0,5 s-os áttűnésbe esett — a cím és a felirat a kontaktlapon már elhalványult)."""
    x = tl[k]
    t = x["end"] - 0.4 if x["id"] != "title" else min(x["end"] - 0.4, 2.0)
    nxt = tl[k + 1]["start"] if k + 1 < len(tl) else None
    if nxt is not None and nxt < x["end"]:
        t = min(t, nxt - 0.15)
    return max(x["start"] + 0.1, t, 0.0)


def contact_sheet(file_id: str, width: int = 1400) -> tuple[bytes, list[dict[str, Any]]]:
    """Az összes jelenet egy képen (a címkártya, minden jelenet vége, a zárókép), utánuk a rétegek (mindegyik a „kész”
    pillanatában), feliratozva — egy pillantás az egészre."""
    import cv2
    import numpy as np
    pid = project_of(file_id)
    spec = (read_meta(pid) or {}).get("spec") or {}
    tl0 = spec.get("timeline") or []
    tl = [(x, scene_still(tl0, k)) for k, x in enumerate(tl0)] + [(x, layer_still(x)) for x in spec.get("layer_timeline") or []]
    if not tl:
        raise ValueError("This video has no scene timeline (made before it was recorded) — use at=seconds instead.")
    vid = (read_meta(pid) or {}).get("render_file_id") if file_id.lower().endswith(".ecproj") else file_id
    if not vid or not output_path(vid).exists():
        raise ValueError("This video project has no render yet — render it first (video action=render).")
    cap = cv2.VideoCapture(str(crafts.existing(vid)))
    fps = cap.get(cv2.CAP_PROP_FPS) or 30
    tiles, info_ = [], []
    cols = 4 if len(tl) > 6 else 3
    tw = width // cols
    for x, t in tl:
        cap.set(cv2.CAP_PROP_POS_FRAMES, int(max(0, t) * fps))
        ok, f = cap.read()
        if not ok:
            continue
        th = int(tw * f.shape[0] / f.shape[1])
        f = cv2.resize(f, (tw, th))
        cv2.rectangle(f, (0, 0), (tw, 30), (0, 0, 0), -1)
        cv2.putText(f, f"{x['id']}  {x['start']:.1f}-{x['end']:.1f}s", (8, 21), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (255, 255, 255), 1, cv2.LINE_AA)
        tiles.append(f)
        info_.append({"id": x["id"], "start": x["start"], "end": x["end"], "frame_at": round(t, 2)})
    cap.release()
    if not tiles:
        raise ValueError("No frame could be read from the video.")
    while len(tiles) % cols:
        tiles.append(np.zeros_like(tiles[0]))
    grid = np.vstack([np.hstack(tiles[i:i + cols]) for i in range(0, len(tiles), cols)])
    ok, buf = cv2.imencode(".jpg", grid, [cv2.IMWRITE_JPEG_QUALITY, 82])
    return buf.tobytes(), info_


def _scene_index(pid: str, scene: Any) -> tuple[int, dict[str, Any], dict[str, Any]]:
    """Jelenet (id / sorszám) vagy réteg (L2) → (1-től számolt index — rétegnél 0 —, jelenet / réteg, idővonal-sor)."""
    spec = (read_meta(pid) or {}).get("spec") or {}
    scenes = _with_ids(copy.deepcopy((spec.get("story") or {}).get("scenes") or []))
    tl = {x["id"]: x for x in spec.get("timeline") or [] if x.get("id")}
    key = str(scene or "").strip()
    k = int(key) if key.isdigit() else next((i + 1 for i, sc in enumerate(scenes) if sc["id"] == key), 0)
    if not 1 <= k <= len(scenes):
        row = next((x for x in spec.get("layer_timeline") or [] if str(x["id"]).lower() == key.lower()), None)
        if row:
            ly = next((x for x in (spec.get("story") or {}).get("layers") or [] if x["id"] == row["id"]), {"id": row["id"]})
            return 0, ly, row
        raise ValueError(f"No scene or layer {scene!r} — {_ids_hint(spec)}.")
    sc = scenes[k - 1]
    return k, sc, tl.get(sc["id"]) or {}


def filmstrip(file_id: str, scene: Any, frames: int = 6, width: int = 1400) -> tuple[bytes, dict[str, Any]]:
    """Egy jelenet mozgása egy képen: `frames` kocka a jelenet elejétől a végéig, időponttal feliratozva — így egy
    külső AI is látja, hogyan nő az oszlop, rajzolódik a vonal, fordul a mutató (nem csak a végállapotot)."""
    import cv2
    import numpy as np
    pid = project_of(file_id)
    _, sc, row = _scene_index(pid, scene)
    if not row:
        raise ValueError("This video has no scene timeline (made before it was recorded) — use at=seconds instead.")
    vid = (read_meta(pid) or {}).get("render_file_id") if file_id.lower().endswith(".ecproj") else file_id
    if not vid or not output_path(vid).exists():
        raise ValueError("This video project has no render yet — render it first (video action=render).")
    frames = max(2, min(12, int(frames or 6)))
    a, b = row["start"] + 0.55, row["end"] - 0.25           # az áttűnés (0,5 s) után: ne az előző jelenet látsszon
    if row.get("layer"):                                    # réteg: a beúszástól a kiúszásig
        a, b = row["start"] + 0.08, row["end"] - 0.12
    times = [a + (b - a) * i / (frames - 1) for i in range(frames)]
    cap = cv2.VideoCapture(str(crafts.existing(vid)))
    fps = cap.get(cv2.CAP_PROP_FPS) or 30
    cols = frames if frames <= 4 else math.ceil(frames / 2)
    tw = width // cols
    tiles = []
    for t in times:
        cap.set(cv2.CAP_PROP_POS_FRAMES, int(t * fps))
        ok, f = cap.read()
        if not ok:
            continue
        f = cv2.resize(f, (tw, int(tw * f.shape[0] / f.shape[1])))
        cv2.rectangle(f, (0, 0), (tw, 26), (0, 0, 0), -1)
        cv2.putText(f, f"{sc['id']}  {t:.1f}s", (6, 19), cv2.FONT_HERSHEY_SIMPLEX, 0.55, (255, 255, 255), 1, cv2.LINE_AA)
        tiles.append(f)
    cap.release()
    if not tiles:
        raise ValueError("No frame could be read from the video.")
    while len(tiles) % cols:
        tiles.append(np.zeros_like(tiles[0]))
    grid = np.vstack([np.hstack(tiles[i:i + cols]) for i in range(0, len(tiles), cols)])
    ok, buf = cv2.imencode(".jpg", grid, [cv2.IMWRITE_JPEG_QUALITY, 82])
    return buf.tobytes(), {"scene": sc["id"], "start": row["start"], "end": row["end"], "frames_at": [round(t, 2) for t in times]}


async def render_scene(file_id: str, scene: Any, draft: bool = True) -> dict[str, Any]:
    """Csak egy jelenet (a saját kompozíciója) vagy egy réteg (a fő videó a réteg idősávjában, a felvétellel / jelenettel
    együtt) renderelése egy rövid MP4-be — ellenőrzéshez gyors, az egész videó újrarenderelése nélkül. A klip rejtett
    munkafájl; a teljes videó a render / edit útján készül."""
    pid = project_of(file_id)
    k, sc, row = _scene_index(pid, scene)
    work = Path(tempfile.mkdtemp(prefix="video-", dir=str(temp_path("").parent)))
    try:
        proj = stage(pid, work)
        names = [c["name"] for c in _comps(json.loads(proj.read_text("utf-8")))]
        span: list[str] = []
        if row.get("layer"):
            comp = _main_comp(proj)
            span = ["--start", f"{row['start']:.3f}", "--end", f"{row['end']:.3f}"]
        else:
            comp = next((n for n in names if n.startswith(f"Scene {k} – ")), None)
            if not comp:
                raise ValueError(f"Scene {sc['id']} isn't a separate part of this project (was it rearranged by hand?).")
        args = ["render", "--project", proj.name, "--comp", comp, "--out", "scene.mp4", "--format", "h264", "--json",
                "--bitrate", "2500" if draft else "8000", "--audio", "auto", *span]
        if draft:
            args += ["--resolution", "half", "--quality", "draft"]
        t0 = time.time()
        async with _render_gate():
            await _run(*args, cwd=work)
        out = work / "scene.mp4"
        if not out.exists():
            raise RuntimeError("The scene render produced no file")
        from app.tools import ios_safe
        await ios_safe.finalize(out)
        vid = new_file_id("mp4")
        shutil.move(str(out), output_path(vid))
        stem = re.sub(r"( \(edited [^)]*\))+$", "", Path(display_label(pid)).stem)
        write_meta(vid, label=f"{stem} – {sc['id']} clip.mp4", operation="video scene render", source_file_id=pid,
                   extra={"ephemeral": True})
        if row.get("layer"):
            return {**_file_entry(vid), "layer": sc["id"], "type": row.get("type"), "start": row["start"], "end": row["end"],
                    "render_seconds": round(time.time() - t0, 1),
                    "note": f"Just {sc['id']}'s time span ({_tc(row['start'])}–{_tc(row['end'])}) of the video, with what is under it "
                            "(draft quality unless draft=false) — to check it quickly."}
        return {**_file_entry(vid), "scene": sc["id"], "headline": sc.get("headline"), "render_seconds": round(time.time() - t0, 1),
                "note": "Just this scene, rendered on its own (draft quality unless draft=false) — to check it quickly."}
    finally:
        shutil.rmtree(work, ignore_errors=True)


def frames_grid(file_id: str, times: list[float], width: int = 1500) -> tuple[bytes, list[dict[str, Any]]]:
    """Több pillanat EGY képen (files preview, at=[…]): minden csempén az időpont és egy halvány tizedes rács (0–1, x balról,
    y felülről) — a kliens-AI ebből látja, hova mutat / néz a szereplő, és a réteg helyét ("position": [x, y, w, h])
    egyenesen leolvashatja. (Kommandant 10-18: „ahova mutatok, oda akarok ábrákat”.)"""
    import cv2
    times = [max(0.0, float(t)) for t in times][:12]
    vid = frame_time(file_id, None, times[0])[0]
    cap = cv2.VideoCapture(str(crafts.existing(vid)))
    fps = cap.get(cv2.CAP_PROP_FPS) or 30
    tiles, info_ = [], []
    first = None
    for t in times:
        cap.set(cv2.CAP_PROP_POS_FRAMES, int(t * fps))
        ok, f = cap.read()
        if not ok:
            continue
        first = first if first is not None else f.shape
        cols = 3 if f.shape[0] > f.shape[1] * 1.2 else 2
        tw = width // cols
        th = int(tw * f.shape[0] / f.shape[1])
        f = cv2.resize(f, (tw, th))
        over = f.copy()
        for k in range(1, 10):                       # tizedes rács, a 0,5 erősebb
            x, y = int(tw * k / 10), int(th * k / 10)
            w_ = 2 if k == 5 else 1
            cv2.line(over, (x, 0), (x, th), (255, 255, 255), w_)
            cv2.line(over, (0, y), (tw, y), (255, 255, 255), w_)
            if k % 2 == 1 or k == 5:
                cv2.putText(over, f".{k}", (x + 2, 14), cv2.FONT_HERSHEY_SIMPLEX, 0.4, (255, 255, 255), 1, cv2.LINE_AA)
                cv2.putText(over, f".{k}", (2, y - 3), cv2.FONT_HERSHEY_SIMPLEX, 0.4, (255, 255, 255), 1, cv2.LINE_AA)
        f = cv2.addWeighted(over, 0.35, f, 0.65, 0)
        cv2.rectangle(f, (0, th - 30), (110, th), (0, 0, 0), -1)
        cv2.putText(f, f"{t:.1f} s", (8, th - 9), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (255, 255, 255), 2, cv2.LINE_AA)
        tiles.append(f)
        info_.append({"index": len(tiles), "at": round(t, 2)})
    cap.release()
    if not tiles:
        raise ValueError("No frame could be read at those times — check them against the video's length.")
    import numpy as np
    cols = 3 if first[0] > first[1] * 1.2 else 2
    rows = [tiles[i:i + cols] for i in range(0, len(tiles), cols)]
    blank = np.zeros_like(tiles[0])
    sheet = np.vstack([np.hstack(r + [blank] * (cols - len(r))) for r in rows])
    ok, buf = cv2.imencode(".jpg", sheet, [cv2.IMWRITE_JPEG_QUALITY, 85])
    return buf.tobytes(), info_
