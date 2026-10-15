"""WordCraft — a Rust Word-motor (storytold/wordcraft, MIT OR Apache-2.0) CLI-n át.

A Dockerfile a rögzített upstream release-ből teszi a `wordcraft-cli`-t a képbe; itt csak
subprocessként hívjuk. Három dolgot ad:
    - run:     a motor parancsai (ugyanaz a ~460, ami a szalagot hajtja) egy dokumentumon,
               új verzióba mentve; változáskövetés a megadott szerző nevén
    - read:    text | info | inspect (szerkezet JSON) | commands (parancslista)
    - convert: gyors konverzió LibreOffice nélkül (docx/odt/rtf/html/md/tex/txt → … / pdf / png)
"""
from __future__ import annotations

import asyncio
import json
import logging
import os
import shutil
import time
from pathlib import Path
from typing import Any, Optional, Union

from app.config import settings
from app.storage import display_label, input_path, output_path, public_url, write_meta
from app.storage.files import new_file_id

log = logging.getLogger("aice.wordcraft")

# `wordcraft-cli convert` / `--save` célformátumai, és amit megnyitni tud.
SAVE_FORMATS = {"docx", "pdf", "odt", "rtf", "html", "md", "tex", "txt", "json", "png"}
OPEN_FORMATS = {"docx", "docm", "dotx", "dotm", "doc", "dot", "odt", "rtf", "html", "htm", "md", "tex", "txt", "json"}
TEMPLATES = {"blank", "sample", "letter", "resume", "report"}
READ_MODES = {"text", "info", "inspect", "commands", "review", "comments", "changes"}

Command = Union[str, dict[str, Any]]


def cli_path() -> Optional[str]:
    p = settings.WORDCRAFT_CLI
    if p and Path(p).is_file() and os.access(p, os.X_OK):
        return p
    return shutil.which("wordcraft-cli")


def available() -> bool:
    return cli_path() is not None


async def _cli(*args: str) -> str:
    exe = cli_path()
    if exe is None:
        raise RuntimeError("wordcraft-cli isn't installed (WORDCRAFT_CLI)")
    # WORDCRAFT_NO_PREFS: ne írjon ~/.config-ot, a napló stderr-re megy.
    env = {**os.environ, "WORDCRAFT_NO_PREFS": "1"}
    proc = await asyncio.create_subprocess_exec(
        exe, *args, env=env,
        stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
    )
    try:
        out, err = await asyncio.wait_for(proc.communicate(), timeout=settings.WORDCRAFT_TIMEOUT_S)
    except asyncio.TimeoutError:
        proc.kill()
        await proc.wait()
        raise RuntimeError(f"wordcraft-cli {args[0]}: timed out after {settings.WORDCRAFT_TIMEOUT_S}s")
    if proc.returncode != 0:
        msg = err.decode("utf-8", "replace").strip() or out.decode("utf-8", "replace").strip()
        raise RuntimeError(msg or f"wordcraft-cli {args[0]}: exit code {proc.returncode}")
    return out.decode("utf-8", "replace")


def _existing(file_id: str) -> Path:
    for p in (input_path(file_id), output_path(file_id)):
        if p.exists():
            return p
    raise FileNotFoundError(f"file_id not found: {file_id}")


def _cmd_arg(c: Command) -> str:
    """{"command": "select.text", "params": {...}} vagy a CLI saját 'id={json}' alakja."""
    if isinstance(c, str):
        return c.strip()
    cmd = c.get("command") or c.get("id")
    if not cmd:
        raise ValueError(f"Command id missing: {c!r}")
    params = c.get("params")
    if isinstance(params, dict) and params.get("file_id"):
        # Feltöltött/generált fájl file_id-vel (pl. insert.picture): a szerver oldja fel útvonalra —
        # az ügynök nem ismeri a szerver fájlrendszerét, base64-et pedig nem visz el a parancssor.
        params = {k: v for k, v in params.items() if k != "file_id"} | {"path": str(_existing(params["file_id"]))}
    return f"{cmd}={json.dumps(params, ensure_ascii=False)}" if params else str(cmd)


async def run(
    *,
    file_id: Optional[str] = None,
    commands: Optional[list[Command]] = None,
    author: Optional[str] = None,
    save_format: str = "docx",
    template: Optional[str] = None,
) -> dict[str, Any]:
    """Parancsok futtatása egy dokumentumon (vagy sablonból indulva), új output file_id-be."""
    t0 = time.time()
    fmt = save_format.lower().lstrip(".")
    if fmt not in SAVE_FORMATS:
        raise ValueError(f"save_format: {fmt} (choose from {', '.join(sorted(SAVE_FORMATS))})")
    args = ["run"]
    if file_id:
        args += ["--file", str(_existing(file_id))]
    else:
        tpl = template or "blank"
        if tpl not in TEMPLATES:
            raise ValueError(f"template: {tpl} (choose from {', '.join(sorted(TEMPLATES))})")
        args += ["--template", tpl]
    if author:
        args += ["--author", author]
    cmd_args = [_cmd_arg(c) for c in (commands or [])]
    for a in cmd_args:
        args += ["--cmd", a]

    out_id = new_file_id(fmt)
    args += ["--save", str(output_path(out_id))]
    from app.tools import run_report
    parsed = [(a.partition("=")[0], _json_or_none(a.partition("=")[2])) for a in cmd_args]
    try:
        await _cli(*args)
    except RuntimeError as e:
        if "unknown command" in str(e) or "step" in str(e):
            ids = [c["id"] for c in json.loads(await _cli("commands", "--json"))]
            raise ValueError(run_report.failure("wordcraft", str(e), "", parsed, ids)) from None
        raise
    from app.storage.naming import result_stem
    from app.tools import ios_safe
    await ios_safe.finalize(output_path(out_id))          # iPhone (Pages, Quick Look): sémahelyes docx
    ids_used = " ".join(c for c, _ in parsed)
    what = "reviewed" if ("review." in ids_used or "comment" in ids_used) else "edited" if file_id else ""
    stem = result_stem(file_id, what, template or "document")
    write_meta(out_id, source_file_id=file_id, operation="wordcraft", label=f"{stem}.{fmt}")
    res = {
        "file_id": out_id,
        "url": public_url(out_id),
        "size_kb": output_path(out_id).stat().st_size // 1024,
        "commands_run": len(cmd_args),
        "steps": run_report.steps(parsed, []),
        "author": author,
        "ms_elapsed": int((time.time() - t0) * 1000),
    }
    if fmt == "docx":
        try:
            after = _doc_summary(output_path(out_id))
            src = _existing(file_id) if file_id else None
            before = _doc_summary(src) if src is not None and src.suffix.lower() == ".docx" else None
            res["before_after"] = {"before": before, "after": after} if before else {"after": after}
        except Exception:
            pass
    return res


def _json_or_none(raw: str):
    try:
        return json.loads(raw) if raw else None
    except ValueError:
        return {"value": raw}


def _doc_summary(path: Path) -> dict[str, Any]:
    """Bekezdések, szavak, követett változások, megjegyzések — az előtte–utána jelentéshez."""
    from docx import Document
    from app.tools.checks import _docx_review_counts
    d = Document(str(path))
    words = sum(len(par.text.split()) for par in d.paragraphs)
    changes, comments, authors = _docx_review_counts(path)
    return {"paragraphs": len(d.paragraphs), "words": words, "tables": len(d.tables),
            "tracked_changes": changes, "comments": comments, "authors": authors}


# Egy MCP-válasz felső határa: egy 345 oldalas regény szövege 1 MB, az inspect-je 800 KB —
# egyben szétvetné a kliens kontextusát. Lapozás: offset + max_chars, next_offset a válaszban.
DEFAULT_MAX_CHARS = 20_000


async def read(
    mode: str,
    file_id: Optional[str] = None,
    query: Optional[str] = None,
    offset: int = 0,
    max_chars: int = DEFAULT_MAX_CHARS,
    base_file_id: Optional[str] = None,
) -> dict[str, Any]:
    if mode not in READ_MODES:
        raise ValueError(f"read: {mode} (choose from {', '.join(sorted(READ_MODES))})")
    offset, max_chars = max(0, offset), max(1_000, max_chars)
    if mode == "commands":
        # 460 parancs teljes leírással 81 KB; az azonosítók ~9 KB. Paraméterek: query-vel (több kifejezés is,
        # a pontos azonosító elöl, hasonló azonosítók, ha nincs találat — közös a többi Craft-motorral).
        from app.tools.crafts import _list_commands
        return _list_commands(json.loads(await _cli("commands", "--json")), query, offset, max_chars)
    if not file_id:
        raise ValueError(f"read={mode} needs a file_id")
    if mode == "review":
        from app.tools import review
        return review.summary(file_id)
    if mode == "comments":           # a bíráló nyitott megjegyzései munkalistaként (query="all": a lezártak is)
        from app.tools import review_loop
        return await asyncio.to_thread(review_loop.worklist, file_id, offset, max_chars, (query or "").lower() == "all")
    if mode == "changes":            # mi változott a korábbi verzióhoz (base_file_id) képest
        if not base_file_id:
            raise ValueError("read=changes needs base_file_id: the earlier version (e.g. the one that was sent out)")
        from app.tools import review_loop
        return await asyncio.to_thread(review_loop.changes, file_id, base_file_id, offset, max_chars)
    out = await _cli(mode, str(_existing(file_id)))
    if mode == "text":
        chunk = out[offset:offset + max_chars]
        end = offset + len(chunk)
        return {"text": chunk, "offset": offset, "total_chars": len(out),
                "next_offset": end if end < len(out) else None}
    data = json.loads(out)
    if mode == "inspect":
        # A blokkokat lapozzuk (offset = blokk-index); a kis metaadatok mindig jönnek.
        blocks = data.pop("blocks", [])
        page, nxt = _page_items(blocks, offset, max_chars)
        data.update({"blocks": page, "blocks_total": len(blocks), "offset": offset, "next_offset": nxt})
    return data


def _page_items(items: list[Any], offset: int, max_chars: int) -> tuple[list[Any], Optional[int]]:
    """items[offset:]-ből annyit, amennyi JSON-ban max_chars alá fér (legalább egyet)."""
    page, size, i = [], 0, offset
    while i < len(items):
        n = len(json.dumps(items[i], ensure_ascii=False))
        if page and size + n > max_chars:
            break
        page.append(items[i]); size += n; i += 1
    return page, (i if i < len(items) else None)


async def convert_file(src: Path, dst: Path) -> None:
    """Konverzió a kiterjesztések alapján (a `convert` tool engine='wordcraft' útja)."""
    await _cli("convert", str(src), str(dst))
