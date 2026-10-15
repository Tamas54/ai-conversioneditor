"""Formátumtudatos kimenet-ellenőrzés: nem elég, hogy a fájl létezik — az legyen, aminek mondja magát.

    problems = check(path, "docx", text=True)      # [] = rendben

Rétegek (ami az adott formátumra értelmes):
  1. mágikus bájtok      a fájl eleje a formátumé (PDF → %PDF-, OOXML/ODF/EPUB → PK, MP4 → ftyp …)
  2. szerkezet           a saját olvasójával megnyílik: Pillow, pypdf (+ qpdf --check), python-docx / openpyxl /
                         python-pptx, ODF/EPUB csomagszabály („mimetype” elöl, tömörítetlenül), ffprobe (stream,
                         hossz), fontTools, tar/zip/7z/gz/bz2/xz/zstd kibontás, JSON/YAML/XML/HTML elemzés
  3. iPhone              MP4/MOV/M4A: moov elöl, iPhone-on lejátszható kodek (ios_safe.plan üres)
  4. ékezetek            text=True esetén a kinyert szövegben ott az „ő” és az „ű”, és nincs „õ”/„û” (Latin-1 csere)
A sweep_conversions.py --verify ezt futtatja minden irányra; a hiányzó eszközt (qpdf, ffprobe, 7z) kihagyja.
"""
from __future__ import annotations

import bz2
import gzip
import io
import json
import lzma
import re
import shutil
import subprocess
import tarfile
import zipfile
from pathlib import Path
from typing import Callable, Optional

OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP = b"PK\x03\x04"


def _riff(kind: bytes) -> Callable[[bytes], bool]:
    return lambda b: b[:4] == b"RIFF" and b[8:12] == kind


def _ftyp(b: bytes) -> bool:
    return b[4:8] == b"ftyp"


def _ebml(b: bytes) -> bool:
    return b[:4] == b"\x1a\x45\xdf\xa3"


def _mpeg_audio(b: bytes) -> bool:
    return b[:3] == b"ID3" or (b[0] == 0xFF and b[1] & 0xE0 == 0xE0)


def _ts(b: bytes) -> bool:              # MPEG-TS: 188 (m2ts/mts: 192) bájtos csomagok, 0x47 szinkronbájt
    return (len(b) > 376 and b[0] == 0x47 and b[188] == 0x47) or (len(b) > 388 and b[4] == 0x47 and b[196] == 0x47)


def _text(b: bytes) -> bool:
    try:
        b[:4096].decode("utf-8")
        return True
    except UnicodeDecodeError as e:     # a minta végén elvágott többbájtos karakter nem hiba
        return e.start >= 4090


MAGIC: dict[str, Callable[[bytes], bool]] = {
    "pdf": lambda b: b[:5] == b"%PDF-", "pdfa": lambda b: b[:5] == b"%PDF-",
    "ps": lambda b: b[:2] == b"%!", "eps": lambda b: b[:2] == b"%!" or b[:4] == b"\xc5\xd0\xd3\xc6",
    "ai": lambda b: b[:5] == b"%PDF-" or b[:2] == b"%!",
    "png": lambda b: b[:8] == b"\x89PNG\r\n\x1a\n", "apng": lambda b: b[:8] == b"\x89PNG\r\n\x1a\n",
    "jpg": lambda b: b[:3] == b"\xff\xd8\xff", "jfif": lambda b: b[:3] == b"\xff\xd8\xff",
    "gif": lambda b: b[:4] == b"GIF8", "webp": _riff(b"WEBP"), "bmp": lambda b: b[:2] == b"BM",
    "tiff": lambda b: b[:4] in (b"II*\0", b"MM\0*"), "ico": lambda b: b[:4] == b"\0\0\1\0",
    "cur": lambda b: b[:4] == b"\0\0\2\0", "icns": lambda b: b[:4] == b"icns", "psd": lambda b: b[:4] == b"8BPS",
    "dds": lambda b: b[:4] == b"DDS ", "exr": lambda b: b[:4] == b"\x76\x2f\x31\x01",
    "hdr": lambda b: b[:2] == b"#?", "fits": lambda b: b[:6] == b"SIMPLE", "sgi": lambda b: b[:2] == b"\x01\xda",
    "jp2": lambda b: b[:12] == b"\0\0\0\x0cjP  \r\n\x87\n" or b[:4] == b"\xff\x4f\xff\x51",
    "jxl": lambda b: b[:2] == b"\xff\x0a" or b[:12] == b"\0\0\0\x0cJXL \r\n\x87\n",
    "avif": _ftyp, "heic": _ftyp, "heif": _ftyp, "pcx": lambda b: b[0] == 0x0A,
    "pbm": lambda b: b[:2] in (b"P1", b"P4"), "pgm": lambda b: b[:2] in (b"P2", b"P5"),
    "ppm": lambda b: b[:2] in (b"P3", b"P6"), "pnm": lambda b: b[:1] == b"P" and b[1:2] in b"123456",
    "xbm": lambda b: b"#define" in b[:200], "xpm": lambda b: b"XPM" in b[:200],
    "wbmp": lambda b: b[:2] == b"\0\0", "emf": lambda b: b[40:44] == b" EMF", "wmf": lambda b: b[:4] in (b"\xd7\xcd\xc6\x9a", b"\x01\0\x09\0", b"\x02\0\x09\0"),
    "emz": lambda b: b[:2] == b"\x1f\x8b", "wmz": lambda b: b[:2] == b"\x1f\x8b", "svgz": lambda b: b[:2] == b"\x1f\x8b",
    "swf": lambda b: b[:3] in (b"FWS", b"CWS", b"ZWS"),
    # dokumentumok
    "doc": lambda b: b[:8] == OLE, "dot": lambda b: b[:8] == OLE, "xls": lambda b: b[:8] == OLE,
    "xlt": lambda b: b[:8] == OLE, "ppt": lambda b: b[:8] == OLE, "pot": lambda b: b[:8] == OLE, "pps": lambda b: b[:8] == OLE,
    "rtf": lambda b: b[:5] == b"{\\rtf", "lit": lambda b: b[:8] == b"ITOLITLS",
    "mobi": lambda b: b[60:68] in (b"BOOKMOBI", b"TEXtREAd"), "azw": lambda b: b[60:68] in (b"BOOKMOBI", b"TEXtREAd"),
    "azw3": lambda b: b[60:68] == b"BOOKMOBI", "prc": lambda b: b[60:68] in (b"BOOKMOBI", b"TEXtREAd"),
    "pdb": lambda b: len(b) > 78, "rb": lambda b: b[:4] == b"\xb0\x0c\xb0\x0c", "lrf": lambda b: b[:6] == "LRF\0".encode("utf-16-le")[:6] or b[:3] == b"L\0R",
    "snb": lambda b: b[:8] == b"SNBP000B", "tcr": lambda b: b[:9] == b"!!8-Bit!!",
    # hang, videó
    "mp3": _mpeg_audio, "mp2": _mpeg_audio, "wav": _riff(b"WAVE"), "avi": _riff(b"AVI "), "divx": _riff(b"AVI "),
    "flac": lambda b: b[:4] == b"fLaC", "ogg": lambda b: b[:4] == b"OggS", "ogv": lambda b: b[:4] == b"OggS",
    "opus": lambda b: b[:4] == b"OggS", "weba": _ebml, "webm": _ebml, "mkv": _ebml, "mka": _ebml,
    "aiff": lambda b: b[:4] == b"FORM" and b[8:12] in (b"AIFF", b"AIFC"), "caf": lambda b: b[:4] == b"caff",
    "au": lambda b: b[:4] == b".snd", "voc": lambda b: b[:19] == b"Creative Voice File",
    "wv": lambda b: b[:4] == b"wvpk", "flv": lambda b: b[:3] == b"FLV", "f4v": _ftyp,
    "asf": lambda b: b[:4] == b"\x30\x26\xb2\x75", "wmv": lambda b: b[:4] == b"\x30\x26\xb2\x75",
    "wma": lambda b: b[:4] == b"\x30\x26\xb2\x75", "wtv": lambda b: b[:4] == b"\xb7\xd8\x00\x20",
    "mpeg": lambda b: b[:4] in (b"\0\0\1\xba", b"\0\0\1\xb3"), "mpg": lambda b: b[:4] in (b"\0\0\1\xba", b"\0\0\1\xb3"),
    "vob": lambda b: b[:4] == b"\0\0\1\xba", "ts": _ts, "mts": _ts, "m2ts": _ts,
    "mxf": lambda b: b[:4] == b"\x06\x0e\x2b\x34", "dv": lambda b: b[:1] == b"\x1f",
    "aac": lambda b: b[:2] in (b"\xff\xf1", b"\xff\xf9") or b[:4] == b"ADIF", "ac3": lambda b: b[:2] == b"\x0b\x77",
    "dts": lambda b: b[:4] in (b"\x7f\xfe\x80\x01", b"\xfe\x7f\x01\x80"), "gsm": lambda b: len(b) >= 33,
    "hevc": lambda b: b[:4] == b"\0\0\0\1" or b[:3] == b"\0\0\1", "mjpeg": lambda b: b[:3] == b"\xff\xd8\xff",
    # archívumok
    "7z": lambda b: b[:6] == b"7z\xbc\xaf\x27\x1c", "tar.7z": lambda b: b[:6] == b"7z\xbc\xaf\x27\x1c",
    "gz": lambda b: b[:2] == b"\x1f\x8b", "tgz": lambda b: b[:2] == b"\x1f\x8b", "tar.gz": lambda b: b[:2] == b"\x1f\x8b",
    "bz2": lambda b: b[:3] == b"BZh", "bz": lambda b: b[:3] == b"BZh", "tbz": lambda b: b[:3] == b"BZh",
    "tbz2": lambda b: b[:3] == b"BZh", "tar.bz2": lambda b: b[:3] == b"BZh", "tar.bz": lambda b: b[:3] == b"BZh",
    "xz": lambda b: b[:6] == b"\xfd7zXZ\0", "txz": lambda b: b[:6] == b"\xfd7zXZ\0", "tar.xz": lambda b: b[:6] == b"\xfd7zXZ\0",
    "lzma": lambda b: b[:1] == b"\x5d", "zst": lambda b: b[:4] == b"\x28\xb5\x2f\xfd", "tar.zst": lambda b: b[:4] == b"\x28\xb5\x2f\xfd",
    "lz": lambda b: b[:4] == b"LZIP", "lzo": lambda b: b[:9] == b"\x89LZO\0\r\n\x1a\n", "tar.lzo": lambda b: b[:9] == b"\x89LZO\0\r\n\x1a\n",
    "tzo": lambda b: b[:9] == b"\x89LZO\0\r\n\x1a\n", "z": lambda b: b[:2] == b"\x1f\x9d", "tz": lambda b: b[:2] == b"\x1f\x9d",
    "tar.z": lambda b: b[:2] == b"\x1f\x9d", "tar": lambda b: b[257:262] == b"ustar",
    "cpio": lambda b: b[:6] in (b"070701", b"070702", b"070707") or b[:2] in (b"\xc7\x71", b"\x71\xc7"),
    "iso": lambda b: True, "wim": lambda b: b[:8] == b"MSWIM\0\0\0",
    # betűtípusok
    "ttf": lambda b: b[:4] in (b"\0\1\0\0", b"true"), "otf": lambda b: b[:4] in (b"OTTO", b"\0\1\0\0"),
    "ttc": lambda b: b[:4] == b"ttcf", "woff": lambda b: b[:4] == b"wOFF", "woff2": lambda b: b[:4] == b"wOF2",
    "eot": lambda b: b[34:36] == b"LP", "pfb": lambda b: b[:2] == b"\x80\x01", "pfa": lambda b: b[:2] == b"%!",
    "dfont": lambda b: b[:4] == b"\0\0\1\0", "bdf": lambda b: b[:9] == b"STARTFONT", "sfd": lambda b: b[:13] == b"SplineFontDB:",
    "plt": lambda b: _text(b), "slk": lambda b: b[:2] == b"ID", "dif": lambda b: b[:5] == b"TABLE",
}
ZIPPED = {"docx", "docm", "dotx", "xlsx", "xlsm", "xltx", "pptx", "pptm", "potx", "ppsx", "odt", "ott", "ods", "odp", "odg",
          "epub", "zip", "jar", "war", "cbz", "htmlz", "txtz", "kmz", "3mf", "xps", "oxps"}
TEXTUAL = {"txt", "md", "html", "htm", "xml", "json", "yaml", "yml", "csv", "tsv", "svg", "rst", "adoc", "org", "tex",
           "textile", "wiki", "fb2", "ipynb", "fodt", "fods", "fodp", "fodg", "srt", "vtt", "ass", "ics", "vcf",
           "plt", "slk", "dif", "xhtml", "opml", "mediawiki", "dokuwiki", "jira", "typst", "bib", "ris"}
ODF_MIME = {"odt": "application/vnd.oasis.opendocument.text", "ott": "application/vnd.oasis.opendocument.text-template",
            "ods": "application/vnd.oasis.opendocument.spreadsheet", "odp": "application/vnd.oasis.opendocument.presentation",
            "odg": "application/vnd.oasis.opendocument.graphics", "epub": "application/epub+zip"}
PIL_EXTS = {"png", "apng", "jpg", "jfif", "gif", "webp", "bmp", "tiff", "ico", "cur", "icns", "psd", "dds", "tga", "pcx",
            "pbm", "pgm", "ppm", "pnm", "xbm", "sgi", "jp2", "heic", "heif", "avif", "jxl", "wbmp", "fits", "xpm"}
MEDIA = {"mp3", "mp2", "wav", "flac", "ogg", "opus", "m4a", "m4b", "m4r", "aac", "ac3", "dts", "aiff", "caf", "au", "voc",
         "wv", "wma", "weba", "mka", "gsm", "alac", "mp4", "m4v", "mov", "mkv", "webm", "avi", "divx", "flv", "f4v", "wmv",
         "asf", "mpeg", "mpg", "vob", "ts", "mts", "m2ts", "mxf", "dv", "ogv", "3gp", "3g2", "wtv", "hevc", "mjpeg"}
VIDEO = {"mp4", "m4v", "mov", "mkv", "webm", "avi", "divx", "flv", "f4v", "wmv", "asf", "mpeg", "mpg", "vob", "ts", "mts",
         "m2ts", "mxf", "dv", "ogv", "3gp", "3g2", "wtv", "hevc", "mjpeg"}
APPLE_MP4 = {"mp4", "m4v", "mov", "m4a", "m4b", "m4r", "3gp", "3g2"}
FONTS = {"ttf", "otf", "woff", "woff2", "ttc", "dfont", "eot"}
# ezekben a kinyerhető szöveg ellenőrizhető (képes/hangos/bináris e-könyv formátumoknál nem)
TEXT_BEARING = TEXTUAL | {"pdf", "pdfa", "docx", "docm", "dotx", "xlsx", "xlsm", "xltx", "pptx", "pptm", "potx", "ppsx",
                          "odt", "ott", "ods", "odp", "odg", "epub", "htmlz", "txtz", "oeb", "rtf", "doc", "dot", "xls",
                          "xlt", "ppt", "pot", "pps"}


def _which(cmd: str) -> Optional[str]:
    return shutil.which(cmd)


def _run(args: list[str], timeout: int = 120) -> subprocess.CompletedProcess:
    return subprocess.run(args, capture_output=True, timeout=timeout)


# ---------- szövegkinyerés az ékezet-ellenőrzéshez ----------

def _zip_text(z: zipfile.ZipFile) -> str:
    out = []
    for n in z.namelist():
        if re.search(r"\.(xml|xhtml|html?|opf|ncx|txt|md|fb2|json)$", n, re.I) and not n.startswith("[Content_Types]"):
            raw = z.read(n)
            try:
                out.append(raw.decode("utf-8"))
            except UnicodeDecodeError:
                out.append(raw.decode("utf-16", "ignore") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else raw.decode("cp1252", "ignore"))
    t = "\n".join(out)
    return re.sub(r"&#(\d+);", lambda m: chr(int(m.group(1))), re.sub(r"&#x([0-9a-f]+);", lambda m: chr(int(m.group(1), 16)), t, flags=re.I))


def extract_text(path: Path, fmt: str) -> Optional[str]:
    """→ a fájl szövege (nagyjából), vagy None, ha ebből a formátumból nem nyerünk ki megbízhatóan."""
    b = path.read_bytes()
    if fmt == "svg" and b"<text" not in b:
        return None                            # a szöveg körvonalként (path) van rajta — látványra ép, kiolvasni nem lehet
    if fmt in TEXTUAL:
        t = b.decode("utf-8", "replace")
        t = re.sub(r"&#x([0-9a-f]+);", lambda m: chr(int(m.group(1), 16)), t, flags=re.I)
        t = re.sub(r"&#(\d+);", lambda m: chr(int(m.group(1))), t)
        if fmt == "tex":                       # LaTeX-átírás: \H{o} = ő, \H{u} = ű
            t = re.sub(r"\\H\{?o\}?", "ő", re.sub(r"\\H\{?u\}?", "ű", t))
        if fmt == "ipynb":
            t = json.dumps(json.loads(t), ensure_ascii=False)
        return t
    if fmt in ("pdf", "pdfa"):
        if _which("pdftotext"):
            r = _run(["pdftotext", "-enc", "UTF-8", str(path), "-"])
            if r.returncode == 0:
                return r.stdout.decode("utf-8", "replace")
        import pypdf
        return "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(str(path)).pages)
    if fmt == "rtf":                          # \u337? = ő, \u369? = ű; \'f5 csak Latin-1-ben (õ) — az hiba
        t = b.decode("latin-1")
        t = re.sub(r"\\u(-?\d+)\s?(\\'[0-9a-f]{2}|\?|.)?", lambda m: chr(int(m.group(1)) % 65536), t)
        return re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1250"), t)
    if b[:4] == ZIP:
        with zipfile.ZipFile(path) as z:
            return _zip_text(z)
    if b[:8] == OLE:                          # Word/Excel/PowerPoint 97: a szöveg UTF-16LE darabokban — csak a meglét
        t = b.decode("utf-16-le", "ignore") + b[1:].decode("utf-16-le", "ignore")     # számít (a bináris
        return "".join(ch for ch in t if ch in "őŐűŰ")                                   # táblákban akad õ/û)
    return None


def accents(text: str) -> list[str]:
    bad = []
    if "ő" not in text and "Ő" not in text:
        bad.append("accent ő missing")
    if "ű" not in text and "Ű" not in text:
        bad.append("accent ű missing")
    if re.search(r"[õû]", text):
        bad.append("Latin-1 substitute õ/û found (ő/ű lost)")
    if "Å‘" in text or "Å±" in text:
        bad.append("mojibake (UTF-8 read as Latin-1)")
    return bad


# ---------- szerkezet ----------

def _ooxml(path: Path, fmt: str) -> list[str]:
    p = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        if "[Content_Types].xml" not in names:
            return ["OOXML without [Content_Types].xml"]
        if names[0] != "[Content_Types].xml":
            p.append("[Content_Types].xml is not the first zip entry (Apple apps are strict)")
    try:
        if fmt in ("docx", "docm", "dotx"):
            import docx
            docx.Document(str(path))
        elif fmt in ("xlsx", "xlsm", "xltx"):
            import openpyxl
            openpyxl.load_workbook(str(path)).close()
        else:
            import pptx
            if len(pptx.Presentation(str(path)).slides) < 1:
                p.append("presentation has no slides")
    except Exception as e:
        if not (fmt in ("dotx", "xltx", "potx", "ppsx", "docm", "xlsm", "pptm") and "content type" in str(e).lower()):
            p.append(f"does not open: {str(e)[:120]}")
    p += _apple_fixes(path)
    return p


def _apple_fixes(path: Path) -> list[str]:
    """A Keynote/Pages/Numbers-javítók egy másolaton: ha még találnak teendőt, a kimenet nem ment át a kapun."""
    import tempfile
    from app.tools import keynote_safe, ooxml_safe
    with tempfile.TemporaryDirectory() as d:
        cp = Path(d) / path.name
        shutil.copyfile(path, cp)
        fixer = keynote_safe if cp.suffix.lower().lstrip(".") in keynote_safe.EXTS else ooxml_safe
        left = fixer.fix(cp)
    return [f"Apple fixer still had work: {', '.join(map(str, left))[:160]}"] if left else []


def _odf(path: Path, fmt: str) -> list[str]:
    p = []
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        first = infos[0] if infos else None
        want = ODF_MIME.get(fmt)
        if not first or first.filename != "mimetype":
            p.append("'mimetype' is not the first zip entry")
        elif first.compress_type != zipfile.ZIP_STORED:
            p.append("'mimetype' is compressed")
        elif want and z.read("mimetype").decode().strip() != want:
            p.append(f"mimetype is {z.read('mimetype')!r}, want {want}")
        if fmt == "epub":
            from lxml import etree
            if "META-INF/container.xml" not in z.namelist():
                return p + ["EPUB without META-INF/container.xml"]
            c = etree.fromstring(z.read("META-INF/container.xml"))
            opf = c.xpath("//*[local-name()='rootfile']/@full-path")
            if not opf or opf[0] not in z.namelist():
                return p + ["EPUB rootfile (OPF) missing"]
            o = etree.fromstring(z.read(opf[0]))
            base = opf[0].rsplit("/", 1)[0] + "/" if "/" in opf[0] else ""
            for href in o.xpath("//*[local-name()='item']/@href"):
                from urllib.parse import unquote
                if base + unquote(href) not in z.namelist():
                    p.append(f"EPUB manifest item missing: {href}")
                    break
            if not o.xpath("//*[local-name()='itemref']"):
                p.append("EPUB spine is empty")
        else:
            from lxml import etree
            for n in ("content.xml", "styles.xml", "META-INF/manifest.xml"):
                if n in z.namelist():
                    etree.fromstring(z.read(n))
                elif n != "styles.xml":
                    p.append(f"ODF part {n} missing")
    return p


def _zip_generic(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            return [f"zip entry corrupt: {bad}"]
        if not z.namelist():
            return ["empty archive"]
    return []


def _pdf(path: Path, fmt: str) -> list[str]:
    import pypdf
    p = []
    r = pypdf.PdfReader(str(path))
    if len(r.pages) < 1:
        p.append("PDF has no pages")
    if _which("qpdf"):
        q = _run(["qpdf", "--check", str(path)])
        if q.returncode not in (0, 3):                 # 3 = csak figyelmeztetés
            p.append("qpdf --check: " + q.stdout.decode("utf-8", "replace").strip().splitlines()[-1][:120])
    if fmt == "pdfa":
        meta = path.read_bytes()
        if b"pdfaid:part" not in meta and b"pdfaid" not in meta:
            p.append("PDF/A identification (pdfaid) missing")
    return p


def _image(path: Path, fmt: str) -> list[str]:
    from PIL import Image
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except Exception:
        pass
    try:
        import pillow_jxl  # noqa: F401
    except Exception:
        pass
    try:
        with Image.open(path) as im:
            im.load()
            if im.width < 1 or im.height < 1:
                return ["image has no pixels"]
        return []
    except Exception as e:
        if _which("magick") or _which("identify"):
            r = _run([_which("magick") or "identify"] + (["identify"] if _which("magick") else []) + [str(path)])
            if r.returncode == 0:
                return []
            return [f"unreadable image: {r.stderr.decode('utf-8', 'replace')[:120]}"]
        return [f"unreadable image: {str(e)[:120]}"]


def _media(path: Path, fmt: str) -> list[str]:
    if not _which("ffprobe"):
        return []
    r = _run(["ffprobe", "-v", "error", "-print_format", "json", "-show_streams", "-show_format", str(path)])
    if r.returncode != 0:
        return ["ffprobe: " + r.stderr.decode("utf-8", "replace")[:160]]
    info = json.loads(r.stdout or b"{}")
    streams = info.get("streams") or []
    p = []
    if not streams:
        p.append("no audio/video stream")
    if fmt in VIDEO and not any(s.get("codec_type") == "video" for s in streams):
        p.append("video container without a video stream")
    dur = float((info.get("format") or {}).get("duration") or 0) or max(
        (float(s.get("duration") or 0) for s in streams), default=0)
    if dur <= 0 and fmt not in ("hevc", "mjpeg", "gsm", "dts", "ac3", "aac"):
        p.append("zero duration")
    if fmt in APPLE_MP4:
        from app.tools import ios_safe
        _, fixes = ios_safe.plan(info, fmt, ios_safe.moov_first(path))
        if fixes:
            p.append("not iPhone-ready: " + ", ".join(fixes))
    return p


def _font(path: Path, fmt: str) -> list[str]:
    from fontTools.ttLib import TTCollection, TTFont
    if fmt == "eot":
        return []
    if fmt == "ttc":
        TTCollection(str(path))
        return []
    f = TTFont(str(path), res_name_or_index=1) if fmt == "dfont" else TTFont(str(path))
    if "cmap" not in f and "CFF " not in f:
        return ["font has no cmap"]
    cmap = f.getBestCmap() or {}
    return [] if (0x151 in cmap and 0x171 in cmap) or not cmap else ["font lost the ő/ű glyphs"]


def _archive(path: Path, fmt: str) -> list[str]:
    raw = path.read_bytes()
    try:
        if fmt in ("gz", "tgz", "tar.gz"):
            data = gzip.decompress(raw)
        elif fmt in ("bz2", "bz", "tbz", "tbz2", "tar.bz2", "tar.bz"):
            data = bz2.decompress(raw)
        elif fmt in ("xz", "txz", "tar.xz"):
            data = lzma.decompress(raw)
        elif fmt == "lzma":
            data = lzma.decompress(raw, format=lzma.FORMAT_ALONE)
        elif fmt in ("zst", "tar.zst") and _which("zstd"):
            data = _run(["zstd", "-dc", str(path)]).stdout
        elif fmt in ("7z", "tar.7z", "wim", "iso", "cpio", "lz", "lzo", "tar.lzo", "tzo", "z", "tz", "tar.z"):
            if _which("7z") and fmt not in ("lz", "lzo", "tar.lzo", "tzo"):
                r = _run(["7z", "t", "-bd", str(path)])
                return [] if r.returncode == 0 else ["7z t: " + r.stdout.decode("utf-8", "replace").strip()[-160:]]
            if _which("bsdtar"):
                r = _run(["bsdtar", "-tf", str(path)])
                return [] if r.returncode == 0 and r.stdout.strip() else ["bsdtar cannot list it"]
            return []
        elif fmt == "tar":
            data = raw
        else:
            return []
    except Exception as e:
        return [f"does not decompress: {str(e)[:120]}"]
    if not data:
        return ["decompresses to nothing"]
    if fmt.startswith("tar") or fmt in ("tgz", "tbz", "tbz2", "txz"):
        with tarfile.open(fileobj=io.BytesIO(data)) as t:
            if not t.getmembers():
                return ["empty tar"]
    return []


def _textual(path: Path, fmt: str) -> list[str]:
    b = path.read_bytes()
    try:
        t = b.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        return [f"not UTF-8 text: {e}"]
    try:
        if fmt in ("json", "ipynb"):
            json.loads(t)
        elif fmt in ("yaml", "yml"):
            import yaml
            yaml.safe_load(t)
        elif fmt in ("xml", "svg", "fb2", "fodt", "fods", "fodp", "fodg", "xhtml", "opml"):
            from lxml import etree
            root = etree.fromstring(b)
            if fmt == "svg" and not root.tag.endswith("svg"):
                return ["SVG root is not <svg>"]
        elif fmt in ("html", "htm"):
            if "<" not in t:
                return ["HTML without markup"]
    except Exception as e:
        return [f"does not parse as {fmt}: {str(e)[:120]}"]
    return []


def check(path: Path, fmt: str, text: bool = False) -> list[str]:
    """→ a talált hibák (üres: rendben). text=True: a forrásban ékezetes szöveg volt, a kimenetben is legyen."""
    path = Path(path)
    if path.is_dir():
        files = [f for f in sorted(path.rglob("*")) if f.is_file()]
        return [f"{f.name}: {x}" for f in files for x in check(f, fmt, text)] if files else ["empty output folder"]
    if not path.exists() or path.stat().st_size == 0:
        return ["empty output"]
    if fmt == "oeb":                           # kicsomagolt könyv: .opf + .html + .css + .ncx fájlok
        own = path.suffix.lower().lstrip(".")
        if own in ("opf", "ncx"):
            from lxml import etree
            try:
                etree.fromstring(path.read_bytes())
            except Exception as e:
                return [f"{own} does not parse: {str(e)[:100]}"]
            return []
        return check(path, "html", text) if own in ("html", "xhtml", "htm") else []
    head = path.read_bytes()[:4096]
    p: list[str] = []
    want = MAGIC.get(fmt)
    if fmt in ZIPPED and head[:4] != ZIP:
        p.append(f"not a zip package (starts {head[:8]!r})")
    elif want and not want(head):
        p.append(f"wrong signature for .{fmt} (starts {head[:12]!r})")
    elif fmt in TEXTUAL and not _text(head):
        p.append("not UTF-8 text")
    if p:
        return p
    try:
        if fmt in ("docx", "docm", "dotx", "xlsx", "xlsm", "xltx", "pptx", "pptm", "potx", "ppsx"):
            p += _zip_generic(path) + _ooxml(path, fmt)
        elif fmt in ODF_MIME:
            p += _zip_generic(path) + _odf(path, fmt)
        elif fmt in ZIPPED:
            p += _zip_generic(path)
        elif fmt in ("pdf", "pdfa"):
            p += _pdf(path, fmt)
        elif fmt in PIL_EXTS:
            p += _image(path, fmt)
        elif fmt in MEDIA:
            p += _media(path, fmt)
        elif fmt in FONTS:
            p += _font(path, fmt)
        elif fmt in TEXTUAL:
            p += _textual(path, fmt)
        elif fmt in ("svgz", "emz", "wmz"):
            gzip.decompress(path.read_bytes())
        else:
            p += _archive(path, fmt)
    except Exception as e:
        p.append(f"{type(e).__name__}: {str(e)[:140]}")
    if text and not p and fmt in TEXT_BEARING:
        t = extract_text(path, fmt)
        if t is not None:
            p += accents(t)
    return p
