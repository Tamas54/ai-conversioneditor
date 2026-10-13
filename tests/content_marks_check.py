"""Önellenőrzés a tartalmi jelölésekhez (modell és adatbázis nélkül): python tests/content_marks_check.py

    docx_marks   — idézet keresése (pontos és tűrő), több futáson átnyúló követett csere, megjegyzés, könyvjelző,
                   követett bekezdés-beszúrás/-törlés; az „összes elfogadva” szöveg a vártnak felel meg
    live_data    — számformátum (tizedesjel, ezres tagolás, tizedesek) olvasása és visszaírása, a kötendő szám kiválasztása
    content_check — a csere a pontos eltérésre szűkül, a záró írásjel nem kerül a mondat közepére, a számellenőrzés
A WordCraft-elfogadás külön (ha a wordcraft-cli elérhető): ugyanaz a szöveg jön ki, mint a saját olvasatunkból.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx  # noqa: E402

from app.tools.content_check import _numbers_ok, _same, minimal_change, tidy_replacement  # noqa: E402
from app.tools.docx_marks import Docx, accept_all_text  # noqa: E402
from app.tools.live_data import _parse, _pick_number, fmt_number  # noqa: E402


def check_marks(tmp: Path) -> Path:
    d = docx.Document()
    d.add_heading("Teszt", 1)
    p = d.add_paragraph("Az infláció ")
    p.add_run("szeptemberben 4,").bold = True
    p.add_run("3% volt, a kamat 6,5%.")
    d.add_paragraph("Ezt töröljük.")
    d.add_paragraph("Utolsó „idézőjellel” és – gondolatjellel.")
    src = tmp / "src.docx"
    d.save(src)
    x = Docx(src, author="test")
    i, a, b = x.find("szeptemberben 4,3%")
    assert x.texts()[i][a:b] == "szeptemberben 4,3%"
    x.comment(x.paragraphs()[i], a, b, "megjegyzés")
    i, a, b = x.find("4,3%")
    x.replace(x.paragraphs()[i], a, b, "1,6%")                      # két futáson átnyúló csere
    i, a, b = x.find("6,5")
    x.bookmark(x.paragraphs()[i], a, b, "aice_t1")
    assert x.bookmarks("aice_") == {"aice_t1": (i, a, b)}
    k = x.find("Ezt töröljük.")[0]
    x.delete_paragraph(x.paragraphs()[k])
    x.insert_paragraph_after(x.paragraphs()[k], "Új bekezdés.")
    i, a, b = x.find('"idézőjellel" és - gondolatjellel')         # idézőjel/gondolatjel-tűrő keresés
    x.replace(x.paragraphs()[i], a, b, "jelek nélkül")
    out = tmp / "out.docx"
    x.save(out)
    want = ["Teszt", "Az infláció szeptemberben 1,6% volt, a kamat 6,5%.", "Új bekezdés.", "Utolsó jelek nélkül."]
    assert accept_all_text(out) == want, accept_all_text(out)
    k = want[1].find("6,5")                                         # a könyvjelző a csere után is a számon ül
    assert Docx(out).bookmarks("aice_")["aice_t1"] == (1, k, k + 3)
    return out


def check_wordcraft(out: Path, tmp: Path) -> str:
    cli = os.getenv("WORDCRAFT_CLI") or shutil.which("wordcraft-cli")
    if not cli or not Path(cli).exists():
        return "skipped (no wordcraft-cli)"
    acc = tmp / "acc.docx"
    subprocess.run([cli, "run", "--file", str(out), "--cmd", "review.acceptAll", "--save", str(acc)], check=True, capture_output=True)
    got = subprocess.run([cli, "text", str(acc)], check=True, capture_output=True, text=True).stdout.strip().splitlines()
    assert got == accept_all_text(out), got
    return "ok"


def check_numbers() -> None:
    assert _parse("4,3", 1.6) == (4.3, {"dec": ",", "group": "", "decimals": 1})
    assert _parse("1 234,5", 1200)[0] == 1234.5
    assert _parse("1,234", 1200)[0] == 1234 and _parse("1,234", 1.2)[0] == 1.234   # a forrásérték dönt
    assert fmt_number(364.22, {"dec": ",", "group": "", "decimals": 0}) == "364"
    assert fmt_number(12345.678, {"dec": ",", "group": " ", "decimals": 1}) == "12 345,7"
    assert fmt_number(-0.5, {"dec": ".", "group": ",", "decimals": 2}) == "-0.50"
    s, e, v, _ = _pick_number("2026-ban az infláció 4,3% volt", 1.6)        # az évszám nem a kötendő szám
    assert (s, e, v) == (21, 24, 4.3)


def check_replacements() -> None:
    q = "A magyar infláció jelenleg 4,3%"
    r = tidy_replacement(q, "A magyar infláció jelenleg 1,6%.")
    assert r == "A magyar infláció jelenleg 1,6%"
    assert minimal_change(q, r) == (27, 30, "1,6")
    assert minimal_change("a Kisvaros Kft.", "a Kisváros Kft.") == (2, 10, "Kisváros")
    assert _same("2026. okt. 5.", "2026. október 5-én") and not _same("4,3%", "1,6%")
    ev = [{"type": "data", "value": 364.22, "id": "S1"}]
    assert _numbers_ok("364 forint", "410 forint", ev) and not _numbers_ok("370 forint", "410 forint", ev)


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        out = check_marks(tmp)
        print("docx_marks ok · WordCraft accept-all:", check_wordcraft(out, tmp))
    check_numbers()
    print("number formats ok")
    check_replacements()
    print("replacements ok")
