"""Szintetikus havi jelenléti ív (a 14–15. pont tesztjeihez) — valódi sablon vagy személyes adat nincs a repóban.

Elrendezés (minden cella keretezett — a 15. pont a fotóból ezt a rácsot méri vissza): 1. sor cím (A1:F1), 2. sor
munkáltató (A2:F2), 3–4. sor fejléc (Név, Munkakör, Hónap, Kód; az értékcellák B:C és E:F egyesítve), 5. sor
munkarend (A5:F5), 6. sor kitöltési útmutató (A6:F6), a 7. sorban a táblázat fejléce, a d. nap a 7+d. sorban
(A: nap, B: érkezés, C: távozás, D: távollét-jelölő „+”, E: ledolgozott idő = C − B, ha mindkettő idő), a 39. sorban
az összesítők: B39 ledolgozott napok (COUNT), D39 távollétnapok (COUNTIF „+”), E39 összes óra (SUM, [h]:mm).
"""
from __future__ import annotations

from pathlib import Path

MERGED = ["A1:F1", "A2:F2", "B3:C3", "E3:F3", "B4:C4", "E4:F4", "A5:F5", "A6:F6"]

PROFILE = {
    "sheet": "Jelenléti ív",
    "header": {"name": "B3", "job_title": "B4", "code": "E4", "month": "E3"},
    "month_format": "{year}. {month}",
    "days": {"first_row": 8, "step": 1},
    "columns": {"arrive": "B", "leave": "C", "absence": "D"},
    "rules": {"workday": {"arrive": "{start}", "leave": "{end}"},
              "kinds": {"szabadság": {"leave": "SZABADSÁG", "absence": "+"},
                        "betegség": {"leave": "BETEGSÉG", "absence": "+"}}},
    "totals": {"hours": "E39", "absence_days": "D39", "worked_days": "B39"},
    "header_values": {"name": "Próba Ödön", "job_title": "Ügyintéző", "code": "T-001"},
    "defaults": {"start": "8:00", "end": "16:00"},
}


def make_template(path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Jelenléti ív"
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "JELENLÉTI ÍV"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "Munkáltató: Szintetikus Próba Kft."
    ws["A5"] = "Munkarend: általános (hétfő–péntek)"
    ws["A6"] = "Szabadság, betegség esetén a Távollét oszlopba „+” kerül."
    for coord, text in (("A3", "Név:"), ("A4", "Munkakör:"), ("D3", "Hónap:"), ("D4", "Kód:")):
        ws[coord] = text
        ws[coord].font = Font(bold=True)
    for rng in MERGED:
        ws.merge_cells(rng)
    for r in range(1, 7):
        for c in "ABCDEF":
            ws[f"{c}{r}"].border = box
    for c, h in zip("ABCDEF", ("Nap", "Érkezés", "Távozás", "Távollét", "Ledolgozott idő", "Aláírás")):
        cell = ws[f"{c}7"]
        cell.value = h
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = box
    for d in range(1, 32):
        r = 7 + d
        ws[f"A{r}"] = d
        for c in "BC":
            ws[f"{c}{r}"].number_format = "h:mm"
        ws[f"E{r}"] = f"=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),C{r}-B{r},0)"
        ws[f"E{r}"].number_format = "[h]:mm"
        for c in "ABCDEF":
            ws[f"{c}{r}"].border = box
            ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
    ws["A39"] = "Összesen:"
    ws["A39"].font = Font(bold=True)
    ws["B39"] = "=COUNT(B8:B38)"
    ws["D39"] = '=COUNTIF(D8:D38,"+")'
    ws["E39"] = "=SUM(E8:E38)"
    ws["E39"].number_format = "[h]:mm"
    for c in "ABCDEF":
        ws[f"{c}39"].border = box
    for c, w in zip("ABCDEF", (8, 12, 14, 11, 16, 18)):
        ws.column_dimensions[c].width = w
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_area = "A1:F39"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def visible_texts() -> list[list[str]]:
    """Ami az ÜRES sablon nyomtatott képén látszik, cellánként (az egyesített cella csak a bal felsőben; a képletek
    eredménye: az óra-oszlop 0:00, az összesítők 0 / 0 / 0:00)."""
    rows = [[""] * 6 for _ in range(39)]
    rows[0][0] = "JELENLÉTI ÍV"
    rows[1][0] = "Munkáltató: Szintetikus Próba Kft."
    rows[2][0], rows[2][3] = "Név:", "Hónap:"
    rows[3][0], rows[3][3] = "Munkakör:", "Kód:"
    rows[4][0] = "Munkarend: általános (hétfő–péntek)"
    rows[5][0] = "Szabadság, betegség esetén a Távollét oszlopba „+” kerül."
    rows[6] = ["Nap", "Érkezés", "Távozás", "Távollét", "Ledolgozott idő", "Aláírás"]
    for d in range(1, 32):
        rows[6 + d][0], rows[6 + d][4] = str(d), "0:00"
    rows[38][0], rows[38][1], rows[38][3], rows[38][4] = "Összesen:", "0", "0", "0:00"
    return rows


def render_png(xlsx: Path, out: Path, dpi: int = 150) -> Path:
    import subprocess
    out.mkdir(parents=True, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{out}/lo", "--headless", "--convert-to", "pdf", "--outdir",
                    str(out), str(xlsx)], check=True, capture_output=True, timeout=180)
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", "-singlefile", str(out / f"{xlsx.stem}.pdf"), str(out / xlsx.stem)],
                   check=True, timeout=60)
    return out / f"{xlsx.stem}.png"


def distort(png: Path, dst: Path, seed: int = 7, rot: float = 1.5, persp: float = 0.035, noise: float = 9) -> Path:
    """Telefonfotó-szerű kép: asztalon fekvő lap, perspektíva, enyhe elforgatás, elmosás, zaj, JPEG."""
    import cv2
    import numpy as np
    img = cv2.imread(str(png))
    rng = np.random.default_rng(seed)
    h, w = img.shape[:2]
    pad = int(0.08 * max(h, w))
    canvas = np.full((h + 2 * pad, w + 2 * pad, 3), 120, np.uint8)
    canvas[pad:pad + h, pad:pad + w] = img
    H, W = canvas.shape[:2]
    src = np.float32([[pad, pad], [pad + w, pad], [pad + w, pad + h], [pad, pad + h]])
    dst_pts = src + np.float32([[rng.uniform(-persp, persp) * w for _ in range(2)] for _ in range(4)])
    out = cv2.warpPerspective(canvas, cv2.getPerspectiveTransform(src, dst_pts), (W, H), borderValue=(120, 120, 120))
    out = cv2.warpAffine(out, cv2.getRotationMatrix2D((W / 2, H / 2), rot, 1.0), (W, H), borderValue=(120, 120, 120))
    out = cv2.GaussianBlur(out, (3, 3), 0.8)
    out = np.clip(out.astype(np.int16) + rng.normal(0, noise, out.shape), 0, 255).astype(np.uint8)
    cv2.imwrite(str(dst), out, [cv2.IMWRITE_JPEG_QUALITY, 85])
    return dst
