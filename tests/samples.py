"""Mintafájlok a tesztekhez — minden futáskor generálva (a repóban nincs bináris fixture).

Minden szöveges minta ékezetes (ő, ű), hogy a kimenetekben ellenőrizni lehessen az épségüket.
"""
from __future__ import annotations

import shutil
import subprocess
import wave
from pathlib import Path

HU = "Árvíztűrő tükörfúrógép — őszi űrlap"


class Samples:
    def __init__(self, root: Path):
        self.root = Path(root)

    def _p(self, name: str) -> Path:
        return self.root / name

    # ---------- dokumentumok ----------

    def docx(self, name: str = "level.docx", text: str = HU) -> Path:
        from docx import Document
        d = Document()
        d.add_heading("Címsor: " + text, 1)
        d.add_paragraph("Első bekezdés: " + text)
        d.add_paragraph("Második bekezdés, ő és ű betűkkel.")
        p = self._p(name)
        d.save(p)
        return p

    def txt(self, name: str = "jegyzet.txt", text: str = HU) -> Path:
        p = self._p(name)
        p.write_text(text + "\nMásodik sor ő ű.\n", "utf-8")
        return p

    def md(self, name: str = "jegyzet.md") -> Path:
        p = self._p(name)
        p.write_text(f"# {HU}\n\nBekezdés **félkövér** ő ű.\n\n- egy\n- kettő\n", "utf-8")
        return p

    def html(self, name: str = "oldal.html") -> Path:
        p = self._p(name)
        p.write_text(f"<!doctype html><html lang=hu><meta charset=utf-8><title>{HU}</title><body><h1>{HU}</h1>"
                     "<p>Bekezdés ő ű.</p></body></html>", "utf-8")
        return p

    def pdf(self, name: str = "dok.pdf", pages: int = 2, text: str = HU) -> Path:
        import pymupdf as fitz
        d = fitz.open()
        font = _font()
        for i in range(pages):
            pg = d.new_page(width=595, height=842)
            if font:
                pg.insert_font(fontname="F0", fontfile=font)
            pg.insert_text((72, 100), f"{i + 1}. oldal: {text}", fontsize=14, fontname="F0" if font else "helv")
            pg.draw_rect(fitz.Rect(72, 150, 300, 300), color=(0.2, 0.4, 0.8), fill=(0.8, 0.9, 1))
        p = self._p(name)
        d.save(p)
        d.close()
        return p

    def fillable_pdf(self, name: str = "urlap.pdf") -> Path:
        """Kitölthető PDF: szövegmező (név, cím), jelölőnégyzet (hozzájárul), aláírásmező helye."""
        import pymupdf as fitz
        d = fitz.open()
        pg = d.new_page(width=595, height=842)
        font = _font()
        if font:
            pg.insert_font(fontname="F0", fontfile=font)
        fn = "F0" if font else "helv"
        pg.insert_text((72, 90), "Jelentkezési lap — őszi tábor", fontsize=16, fontname=fn)
        for i, (label, fname) in enumerate((("Név:", "nev"), ("Lakcím:", "lakcim"), ("Születési dátum:", "szuletesi_datum"))):
            y = 140 + i * 40
            pg.insert_text((72, y + 14), label, fontsize=11, fontname=fn)
            w = fitz.Widget()
            w.field_name, w.field_type = fname, fitz.PDF_WIDGET_TYPE_TEXT
            w.rect = fitz.Rect(200, y, 500, y + 20)
            pg.add_widget(w)
        pg.insert_text((72, 274), "Hozzájárulok az adatkezeléshez", fontsize=11, fontname=fn)
        cb = fitz.Widget()
        cb.field_name, cb.field_type = "hozzajarul", fitz.PDF_WIDGET_TYPE_CHECKBOX
        cb.rect = fitz.Rect(300, 260, 316, 276)
        pg.add_widget(cb)
        pg.insert_text((72, 340), "Aláírás:", fontsize=11, fontname=fn)
        pg.draw_line((140, 345), (400, 345))
        p = self._p(name)
        d.save(p)
        d.close()
        return p

    def flat_form_pdf(self, name: str = "lapos_urlap.pdf") -> Path:
        """Lapos (mező nélküli) nyomtatvány: címkék + aláhúzás-vonalak + jelölőnégyzet-keret."""
        import pymupdf as fitz
        d = fitz.open()
        pg = d.new_page(width=595, height=842)
        font = _font()
        if font:
            pg.insert_font(fontname="F0", fontfile=font)
        fn = "F0" if font else "helv"
        pg.insert_text((72, 90), "Adatlap", fontsize=16, fontname=fn)
        for i, label in enumerate(("Név:", "Lakcím:", "Adóazonosító jel:", "Dátum:")):
            y = 150 + i * 40
            pg.insert_text((72, y), label, fontsize=11, fontname=fn)
            pg.draw_line((200, y + 2), (500, y + 2), width=0.8)
        pg.insert_text((72, 330), "Hozzájárulok", fontsize=11, fontname=fn)
        pg.draw_rect(fitz.Rect(200, 318, 212, 330), width=0.8)
        pg.insert_text((72, 380), "Aláírás:", fontsize=11, fontname=fn)
        pg.draw_line((200, 382), (420, 382), width=0.8)
        p = self._p(name)
        d.save(p)
        d.close()
        return p

    def form_docx(self, name: str = "urlap.docx") -> Path:
        from docx import Document
        d = Document()
        d.add_heading("Meghatalmazás", 1)
        d.add_paragraph("Név: ____________________")
        d.add_paragraph("Születési dátum: __________")
        d.add_paragraph("Adóazonosító jel: [          ]")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(1, 0).text = "Lakcím", "Telefonszám"
        d.add_paragraph("Hozzájárulok: ☐")
        d.add_paragraph("Aláírás: ____________________")
        p = self._p(name)
        d.save(p)
        return p

    def word_review(self, name: str = "fonok_kommentjei.docx") -> Path:
        """Wordben kommentelt dokumentum, ahogy a Word 2019+ menti: a főnök megjegyzései (szál Beata válaszával,
        egy lezárt, egy két bekezdésen átívelő, egy táblázatcellában), a főnök saját követett változásai, és a
        Word-kiegészítő részek (commentsExtended: szál + lezárás, commentsIds, commentsExtensible, people)."""
        import zipfile
        from docx import Document
        base = self._p("_base_" + name)
        Document().save(base)
        boss, beata = "Dr. Főnök Ferenc", "Kiss Beata"
        d0, d1 = "2026-10-14T09:12:00Z", "2026-10-14T10:30:00Z"
        NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
              'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
              'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
              'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
              'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
              'xmlns:w16cid="http://schemas.microsoft.com/office/word/2016/wordml/cid" '
              'xmlns:w16cex="http://schemas.microsoft.com/office/word/2018/wordml/cex" '
              'mc:Ignorable="w14 w15 w16cid w16cex"')

        def run(t):
            return f'<w:r><w:t xml:space="preserve">{t}</w:t></w:r>'

        def ref(i):
            return f'<w:r><w:rPr><w:rStyle w:val="CommentReference"/></w:rPr><w:commentReference w:id="{i}"/></w:r>'

        def para(pid, inner, style=None):
            ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
            return f'<w:p w14:paraId="{pid}" w14:textId="77777777">{ppr}{inner}</w:p>'
        body = "".join([
            para("1A000001", run("Integritás Hatóság – javaslat"), "Heading1"),
            para("1A000002", run("A hatóság ") + '<w:commentRangeStart w:id="0"/><w:commentRangeStart w:id="1"/>'
                 + run("2022-ben jött létre") + '<w:commentRangeEnd w:id="0"/>' + ref(0) + '<w:commentRangeEnd w:id="1"/>'
                 + ref(1) + run(" az uniós források védelmére.")),
            para("1A000003", run("Az ") + '<w:commentRangeStart w:id="4"/>' + run("elnököt az Állami Számvevőszék elnöke javasolja.")),
            para("1A000004", run("A testület munkája ") + run("független") + '<w:commentRangeEnd w:id="4"/>' + ref(4)
                 + run(", és évente beszámol.")),
            para("1A000005", run("A költségvetés évi ")
                 + f'<w:del w:id="101" w:author="{boss}" w:date="{d0}"><w:r><w:delText>4</w:delText></w:r></w:del>'
                 + f'<w:ins w:id="102" w:author="{boss}" w:date="{d0}">{run("5")}</w:ins>' + run(" milliárd forint.")),
            para("1A000006", run("Ezt a mondatot a főnök már ") + '<w:commentRangeStart w:id="2"/>' + run("lezárt")
                 + '<w:commentRangeEnd w:id="2"/>' + ref(2) + run(" megjegyzéssel látta el.")),
            '<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
            '<w:tblGrid><w:gridCol w:w="4500"/><w:gridCol w:w="4500"/></w:tblGrid><w:tr>'
            '<w:tc><w:tcPr><w:tcW w:w="4500" w:type="dxa"/></w:tcPr>' + para("1A000007", run("Határidő")) + '</w:tc>'
            '<w:tc><w:tcPr><w:tcW w:w="4500" w:type="dxa"/></w:tcPr>' + para("1A000008", '<w:commentRangeStart w:id="3"/>'
            + run("március 31.") + '<w:commentRangeEnd w:id="3"/>' + ref(3)) + '</w:tc></w:tr></w:tbl>',
            para("1A000009", run("Záró bekezdés, ő és ű betűkkel.")),
        ])
        document = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {NS}><w:body>{body}'
                    '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1417" w:right="1417" w:bottom="1417" '
                    'w:left="1417" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr></w:body></w:document>')
        cs = [("0", boss, d0, "B", "2B000001", "Pontosítsd: melyik törvény hozta létre?", None, "0"),
              ("1", beata, d1, "KB", "2B000002", "Utánanézek.", "2B000001", "0"),
              ("2", boss, d0, "B", "2B000003", "Rendben.", None, "1"),
              ("3", boss, d0, "B", "2B000004", "Ez a határidő nem jó, április 15. legyen.", None, "0"),
              ("4", boss, d0, "B", "2B000005", "Ez a két mondat ellentmond egymásnak.", None, "0")]
        comments = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:comments {NS}>' + "".join(
            f'<w:comment w:id="{i}" w:author="{a}" w:date="{dt}" w:initials="{ini}"><w:p w14:paraId="{pid}" w14:textId="77777777">'
            f'<w:pPr><w:pStyle w:val="CommentText"/></w:pPr><w:r><w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>'
            f'<w:annotationRef/></w:r>{run(t)}</w:p></w:comment>' for i, a, dt, ini, pid, t, _, _ in cs) + '</w:comments>')
        ext = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w15:commentsEx {NS}>' + "".join(
            f'<w15:commentEx w15:paraId="{pid}"' + (f' w15:paraIdParent="{par}"' if par else "") + f' w15:done="{done}"/>'
            for _, _, _, _, pid, _, par, done in cs) + '</w15:commentsEx>')
        ids = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w16cid:commentsIds {NS}>' + "".join(
            f'<w16cid:commentId w16cid:paraId="{pid}" w16cid:durableId="{0x3C000001 + k:08X}"/>'
            for k, (_, _, _, _, pid, _, _, _) in enumerate(cs)) + '</w16cid:commentsIds>')
        cex = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w16cex:commentsExtensible {NS}>' + "".join(
            f'<w16cex:commentExtensible w16cex:durableId="{0x3C000001 + k:08X}" w16cex:dateUtc="{dt}"/>'
            for k, (_, _, dt, _, _, _, _, _) in enumerate(cs)) + '</w16cex:commentsExtensible>')
        people = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w15:people {NS}>' + "".join(
            f'<w15:person w15:author="{a}"><w15:presenceInfo w15:providerId="None" w15:userId="{a}"/></w15:person>'
            for a in (boss, beata)) + '</w15:people>')
        parts = {"word/comments.xml": ("application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
                                       "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments", comments),
                 "word/commentsExtended.xml": ("application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml",
                                               "http://schemas.microsoft.com/office/2011/relationships/commentsExtended", ext),
                 "word/commentsIds.xml": ("application/vnd.openxmlformats-officedocument.wordprocessingml.commentsIds+xml",
                                          "http://schemas.microsoft.com/office/2016/09/relationships/commentsIds", ids),
                 "word/commentsExtensible.xml": ("application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtensible+xml",
                                                 "http://schemas.microsoft.com/office/2018/08/relationships/commentsExtensible", cex),
                 "word/people.xml": ("application/vnd.openxmlformats-officedocument.wordprocessingml.people+xml",
                                     "http://schemas.microsoft.com/office/2011/relationships/people", people)}
        out = self._p(name)
        with zipfile.ZipFile(base) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/document.xml":
                    data = document.encode()
                elif item.filename == "[Content_Types].xml":
                    data = data.replace(b"</Types>", "".join(
                        f'<Override PartName="/{n}" ContentType="{ct}"/>' for n, (ct, _, _) in parts.items()).encode() + b"</Types>")
                elif item.filename == "word/_rels/document.xml.rels":
                    data = data.replace(b"</Relationships>", "".join(
                        f'<Relationship Id="rIdC{k}" Type="{rt}" Target="{n.split("/", 1)[1]}"/>'
                        for k, (n, (_, rt, _)) in enumerate(parts.items())).encode() + b"</Relationships>")
                zo.writestr(zipfile.ZipInfo(item.filename, date_time=item.date_time), data, zipfile.ZIP_DEFLATED)
            for n, (_, _, xml) in parts.items():
                zo.writestr(n, xml.encode())
        base.unlink()
        return out

    def form_xlsx(self, name: str = "urlap.xlsx") -> Path:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Adatlap"
        rows = [("Név", None), ("Lakcím", None), ("Születési dátum", None), ("IBAN", None)]
        for i, (k, v) in enumerate(rows, 1):
            ws.cell(i, 1, k)
        p = self._p(name)
        wb.save(p)
        return p

    def csv_rows(self, name: str = "adatok.csv") -> Path:
        p = self._p(name)
        p.write_text("nev;lakcim;hozzajarul\nKőváry Ödön;Győr, Fő utca 1.;igen\nŰrmös Éva;Pécs, Ős tér 2.;nem\n", "utf-8")
        return p

    def xlsx(self, name: str = "tabla.xlsx") -> Path:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Név", "Összeg"])
        ws.append(["Őrs", 12])
        ws.append(["Űr", 30])
        p = self._p(name)
        wb.save(p)
        return p

    def pptx(self, name: str = "dia.pptx") -> Path:
        from pptx import Presentation
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = HU
        s.placeholders[1].text = "Pont ő ű"
        p = self._p(name)
        prs.save(p)
        return p

    # ---------- képek ----------

    def png(self, name: str = "kep.png", size=(320, 200), alpha: bool = False) -> Path:
        from PIL import Image, ImageDraw
        im = Image.new("RGBA" if alpha else "RGB", size, (0, 0, 0, 0) if alpha else (240, 240, 230))
        dr = ImageDraw.Draw(im)
        dr.rectangle((20, 20, size[0] - 20, size[1] - 20), outline=(200, 30, 30), width=6)
        dr.ellipse((60, 40, 180, 160), fill=(30, 120, 200))
        p = self._p(name)
        im.save(p)
        return p

    def jpg(self, name: str = "foto.jpg", size=(640, 480)) -> Path:
        from PIL import Image
        im = Image.open(self.png("tmp_src.png", size)).convert("RGB")
        p = self._p(name)
        exif = Image.Exif()
        exif[0x010F] = "TesztGyártó"   # Make
        exif[0x0110] = "Modell ő"       # Model
        im.save(p, "JPEG", quality=92, exif=exif.tobytes())
        return p

    def signature_photo(self, name: str = "alairas.jpg") -> Path:
        """Papírra írt aláírás fotója: halvány, egyenetlen háttér + sötét tollvonás."""
        from PIL import Image, ImageDraw, ImageFilter
        import math
        im = Image.new("RGB", (600, 260), (236, 232, 222))
        dr = ImageDraw.Draw(im)
        for x in range(600):                    # árnyékos papír
            dr.line([(x, 0), (x, 260)], fill=(236 - x // 30, 232 - x // 30, 222 - x // 40))
        pts = [(60 + t * 4.6, 140 + 40 * math.sin(t / 7.0) - 20 * math.cos(t / 3.0)) for t in range(100)]
        dr.line(pts, fill=(20, 25, 70), width=5)
        dr.line([(80, 200), (520, 190)], fill=(25, 30, 80), width=3)
        im = im.filter(ImageFilter.GaussianBlur(0.6))
        p = self._p(name)
        im.save(p, "JPEG", quality=90)
        return p

    def svg(self, name: str = "abra.svg") -> Path:
        p = self._p(name)
        p.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="120"><rect x="10" y="10" width="180" '
                     'height="100" fill="#3a7"/><text x="20" y="70" font-size="20">ő ű</text></svg>', "utf-8")
        return p

    # ---------- média ----------

    def wav(self, name: str = "hang.wav", seconds: float = 1.0, rate: int = 16000) -> Path:
        import math
        import struct
        p = self._p(name)
        with wave.open(str(p), "wb") as w:
            w.setnchannels(1)
            w.setsampwidth(2)
            w.setframerate(rate)
            w.writeframes(b"".join(struct.pack("<h", int(8000 * math.sin(2 * math.pi * 440 * i / rate)))
                                   for i in range(int(seconds * rate))))
        return p

    def mp4(self, name: str = "video.mp4", seconds: int = 2) -> Path:
        p = self._p(name)
        _ff("-f", "lavfi", "-i", f"testsrc=size=320x240:rate=15:duration={seconds}", "-f", "lavfi", "-i",
            f"sine=frequency=440:duration={seconds}", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac",
            "-shortest", str(p))
        return p

    # ---------- archívum, font, CAD ----------

    def zip(self, name: str = "csomag.zip") -> Path:
        import zipfile
        p = self._p(name)
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("ő/első.txt", "első ő")
            z.writestr("második.txt", "második ű")
        return p

    def ttf(self, name: str = "betu.ttf") -> Path:
        """Egy rendszer-TTF másolata (DejaVu / Liberation), ha van."""
        for cand in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                     "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"):
            if Path(cand).exists():
                p = self._p(name)
                shutil.copy(cand, p)
                return p
        import pytest
        pytest.skip("no TTF font on this machine")

    def dxf(self, name: str = "rajz.dxf") -> Path:
        import ezdxf
        doc = ezdxf.new()
        msp = doc.modelspace()
        msp.add_line((0, 0), (100, 0))
        msp.add_circle((50, 50), 25)
        msp.add_text("ő ű", dxfattribs={"height": 5}).set_placement((10, 80))
        p = self._p(name)
        doc.saveas(p)
        return p


def _font() -> str | None:
    for cand in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                 "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"):
        if Path(cand).exists():
            return cand
    return None


def _ff(*args: str) -> None:
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", *args], check=True)
