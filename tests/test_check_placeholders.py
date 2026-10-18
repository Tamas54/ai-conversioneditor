"""Check: a kitöltetlen helyőrző biztos lelet, modell nélkül.

UI-próba (10-18): a briefből írt meghívóban „[a könyvtár neve és címe]”, „[elérhetőségek]” maradt, a Check mégis
„nothing found”-ot mondott (csak ellentmondást, hangnemet, szóhasználatot keresett).
"""
from __future__ import annotations

import docx

from app.tools.content_check import placeholders
from app.tools.docx_marks import Docx


def test_placeholders_found_with_places(tmp_path):
    p = tmp_path / "meghivo.docx"
    d = docx.Document()
    d.add_paragraph("Helyszín: [a könyvtár neve és címe]")
    d.add_paragraph("Lásd [1] és a […] részt; Hello {{name}}, ár: TBD.")
    d.save(p)
    f = placeholders(Docx(p))
    assert f["kind"] == "placeholder"
    assert [(x["paragraph"], x["quote"]) for x in f["places"]] == [
        (0, "[a könyvtár neve és címe]"), (1, "{{name}}"), (1, "TBD")]


def test_clean_document_has_none(tmp_path):
    p = tmp_path / "ok.docx"
    d = docx.Document()
    d.add_paragraph("Minden ki van töltve [1], a hivatkozás nem helyőrző.")
    d.save(p)
    assert placeholders(Docx(p)) is None
