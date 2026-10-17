"""Összetett pakli kliensként (10-17, a szegedi energiaprogram próbája) — ami ott elcsúszott, általánosan javítva:

- a címlap felső sora (kicker) és a többsoros cím nem kerül egymásra (a cím betűje a dobozához igazodik);
- a táblázat oszlopai a tartalomhoz szélesek, a sorok magassága a szöveghez, és az egész a cím és a lábléc közé fér;
- a kör- és fánkdiagram szeletein százalék van (nem a nyers érték), a pontdiagramnak tengelycíme és pontnevei lehetnek;
- a stockkép-keresés rövidítéskor sem dobja el a tulajdonnevet („Szeged city aerial” ≠ „city aerial”)."""
from __future__ import annotations

from pathlib import Path

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _deck(tmp_path: Path, slides: list[dict]) -> Path:
    from pptx import Presentation
    from pptx.util import Inches
    from app.tools import pptx_deck
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(pptx_deck.W_IN), Inches(pptx_deck.H_IN)
    t = dict(pptx_deck.DEFAULT_THEME)
    t.setdefault("panel", "#F3F4F6")
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        {"cover": lambda: pptx_deck._cover(slide, s, t, None), "table": lambda: pptx_deck._table(slide, s, t),
         "chart": lambda: pptx_deck._chart(slide, s, t)}[s["layout"]]()
    p = tmp_path / "d.pptx"
    prs.save(str(p))
    return p


def test_cover_kicker_stays_above_a_long_title(tmp_path):
    from pptx import Presentation
    p = _deck(tmp_path, [{"layout": "cover", "kicker": "Közgyűlési előterjesztés · 2026. november",
                          "title": "Szeged energiaátállása és távhőfejlesztési programja 2027–2030"}])
    boxes = {sh.text_frame.text: sh for sh in Presentation(str(p)).slides[0].shapes if sh.has_text_frame and sh.text_frame.text}
    kicker = next(v for k, v in boxes.items() if k.startswith("KÖZGYŰLÉSI"))
    title = next(v for k, v in boxes.items() if k.startswith("Szeged"))
    assert kicker.top + kicker.height <= title.top                                   # a két doboz nem fedi egymást
    size = title.text_frame.paragraphs[0].runs[0].font.size.pt
    assert size < 54                                                                 # a hosszú cím kisebb betűvel fér be


def test_table_columns_follow_the_content_and_fit_the_slide(tmp_path):
    from pptx import Presentation
    from app.tools import pptx_deck
    rows = [["Geotermikus kút és hőközpont – Tarján", "2027–2028", "12,5", "Szegedi Hőszolgáltató"],
            ["Közintézmények napelemes rendszerei", "2027", "4,2", "Városüzemeltetés"]] * 4
    p = _deck(tmp_path, [{"layout": "table", "title": "Projektek", "headers": ["Projekt", "Ütem", "Mrd Ft", "Felelős"], "rows": rows}])
    tbl = next(sh for sh in Presentation(str(p)).slides[0].shapes if sh.has_table)
    w = [c.width for c in tbl.table.columns]
    assert w[0] > w[3] > w[2] and w[1] > w[2]                                        # a szöveges oszlop széles, a szám keskeny
    assert (tbl.top + sum(r.height for r in tbl.table.rows)) / 914400 <= pptx_deck.H_IN - 0.7   # a lábléc fölött véget ér


def test_pie_slices_show_percent_and_scatter_has_axis_titles_and_names(tmp_path):
    from pptx import Presentation
    p = _deck(tmp_path, [
        {"layout": "chart", "title": "Források", "chart_type": "doughnut", "categories": ["EU", "Állam", "Város"],
         "series": [{"name": "Mrd Ft", "values": [17.3, 6.9, 9.2]}]},
        {"layout": "chart", "title": "Projektek", "chart_type": "scatter", "x_title": "Beruházás (Mrd Ft)",
         "y_title": "Megtakarítás (Mrd Ft/év)", "series": [{"name": "Projektek", "points": [[12.5, 1.9, "Tarján"], [4.2, 0.7]]}]}])
    prs = Presentation(str(p))
    pie = next(sh for sh in prs.slides[0].shapes if sh.has_chart).chart
    lbls = list(pie._chartSpace.iter(f"{C}dLbl"))
    assert len(lbls) == 3
    for d in lbls:                                                                    # szeletenként százalék, nem érték
        assert d.find(f"{C}showPercent").get("val") == "1" and d.find(f"{C}showVal").get("val") == "0"
        assert d.find(f"{C}numFmt").get("formatCode") == "0%"
    sc = next(sh for sh in prs.slides[1].shapes if sh.has_chart).chart
    assert sc.category_axis.axis_title.text_frame.text == "Beruházás (Mrd Ft)"
    assert sc.value_axis.axis_title.text_frame.text == "Megtakarítás (Mrd Ft/év)"
    names = [t.text for t in sc._chartSpace.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t")]
    assert "Tarján" in names


def test_stock_search_keeps_the_place_name_when_it_shortens():
    from app.tools.images import _tries
    tries = _tries("Szeged city aerial view")
    assert tries[0] == "Szeged city aerial view" and all("Szeged" in q for q in tries)
    assert "Szeged" in tries                                                          # végül csak a név
    assert _tries("gas pipeline valve") == ["gas pipeline valve", "pipeline valve", "gas pipeline"]   # név nélkül a régi
