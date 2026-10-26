"""Deck layer (N2, app/tools/deck_layer.py): the schema, the slide pictures, "switch when this phrase is spoken", the
builder output, the live preview's matte part, and the per-segment content hashes (only the retimed segments change)."""
from __future__ import annotations

import asyncio
import json
import shutil
import subprocess
from pathlib import Path

import pytest

DECK = "0123456789abcdef0123456789abcdef.pptx"


def _deck(at=150, **kw):
    return {"at": at, "type": "deck", "file_id": DECK, **kw}


# ---------- the schema ----------

def test_deck_layer_is_validated_and_keeps_its_slides():
    from app.tools import video
    notes: list[str] = []
    [ly] = video.layers_from([_deck(slides=[{"at": "02:30", "slide": 1}, {"at": "02:52", "slide": 4},
                                            {"at": 200, "slide": 7}], until="04:00")], "", notes)
    assert ly["type"] == "deck" and ly["placement"] == "window" and ly["transition"] == "fade"
    assert ly["at"] == 150 and ly["dur"] == 90
    assert [(s["at"], s["slide"]) for s in ly["slides"]] == [(150, 1), (172, 4), (200, 7)]
    # without at: the first slide's time; without dur / until: the last slide stays 20 s
    [l2] = video.layers_from([{"type": "slides", "file_id": DECK, "placement": "behind",
                               "slides": [{"at": 1500, "slide": 2}, {"at": 1530, "slide": 3}]}], "", [])
    assert l2["type"] == "deck" and l2["placement"] == "background" and l2["at"] == 1500 and l2["dur"] == 50
    assert video.layer_listing([l2])[0]["text"] == "deck (2 slides, background)"


@pytest.mark.parametrize("bad,msg", [
    ({"file_id": None}, "file_id = the PowerPoint"),
    ({"file_id": "abc.png"}, "isn't a slide deck"),
    ({"placement": "under"}, "placement 'under'"),
    ({"side": "middle"}, "side 'middle'"),
    ({"transition": "wipe"}, "transition 'wipe'"),
    ({"slides": [{"at": 100, "slide": 1}]}, "before the deck starts"),
    ({"slides": [{"at": 160, "slide": 0}]}, "slide numbers start at 1"),
    ({"slides": [{"at": 160}]}, "slide numbers from 1"),
])
def test_deck_layer_errors_say_what_is_wrong(bad, msg):
    from app.tools import video
    with pytest.raises(ValueError, match=msg):
        video.layers_from([{**_deck(dur=60), **bad}], "", [])


def test_deck_slides_out_of_span_or_too_close_are_left_out_with_a_note():
    from app.tools import video
    notes: list[str] = []
    [ly] = video.layers_from([_deck(dur=30, slides=[{"at": 150, "slide": 1}, {"at": 150.4, "slide": 2},
                                                    {"at": 170, "slide": 3}, {"at": 185, "slide": 4}])], "", notes)
    assert [s["slide"] for s in ly["slides"]] == [1, 3]
    assert any("slide 2" in n and "less than 1 s" in n for n in notes)
    assert any("slide 4" in n and "after the deck ends" in n for n in notes)


def test_deck_edit_moves_slides_with_the_deck_and_replaces_them_when_given():
    from app.tools import video
    [ly] = video.layers_from([_deck(dur=60, slides=[{"at": 150, "slide": 1}, {"at": 170, "slide": 2}])], "", [])
    [moved] = video.apply_layer_changes([ly], {"layers": [{"id": ly["id"], "at": 160}]}, "", [])
    assert moved["at"] == 160 and [s["at"] for s in moved["slides"]] == [160, 180]
    [re_] = video.apply_layer_changes([ly], {"layers": [{"id": ly["id"], "slides": [{"at": 150, "slide": 1},
                                                                                      {"at": 175, "slide": 5}]}]}, "", [])
    assert [(s["at"], s["slide"]) for s in re_["slides"]] == [(150, 1), (175, 5)] and re_["at"] == 150
    [front] = video.apply_layer_changes([ly], {"layers": [{"id": ly["id"], "placement": "front"}]}, "", [])
    assert front["placement"] == "front" and front["slides"] == ly["slides"]


def test_deck_switches_follow_a_cut():
    from app.tools import video
    [ly] = video.layers_from([_deck(at=10, dur=40, slides=[{"at": 10, "slide": 1}, {"at": 25, "slide": 2},
                                                           {"at": 40, "slide": 3}])], "", [])
    keep_old, keep_new = None, [[0.0, 20.0], [30.0, 60.0]]          # 20–30 s cut out
    out, _, _ = video._recut([ly], [], None, keep_old, keep_new, [])
    s = out[0]["slides"]
    assert out[0]["at"] == 10 and [x["slide"] for x in s] == [1, 2, 3]
    assert s[1]["at"] == 20 and s[2]["at"] == 30                     # 25 was cut out → the next kept time; 40 → 30


# ---------- "switch when this phrase is spoken" ----------

CUES = [{"start": 100.0, "end": 106.0, "text": "Jó napot, ma a mesterséges intelligenciáról beszélünk."},
        {"start": 106.0, "end": 112.0, "text": "Nézzük először, mit hoz az MI a munkaerőpiacra."},
        {"start": 112.0, "end": 118.0, "text": "A fiatalok körében az MI használat drámaian nő."}]


def test_follow_resolves_from_subtitle_words_deterministically():
    from app.tools import deck_layer, video
    notes: list[str] = []
    raw = [{"type": "deck", "file_id": DECK, "follow": [{"slide": 1, "said": "Jó napot"}, {"slide": 4, "said": "munkaerőpiacra"},
                                                        {"slide": 7, "said": "a fiatalok körében"}, {"slide": 9, "said": "kvantumgép"}]}]
    a = asyncio.run(deck_layer.resolve(raw, notes, cues=CUES))
    b = asyncio.run(deck_layer.resolve(raw, [], cues=CUES))
    assert a == b                                                   # deterministic
    s = a[0]["slides"]
    assert [x["slide"] for x in s] == [1, 4, 7] and "follow" not in a[0]
    assert s[0]["at"] == pytest.approx(99.85, abs=0.01)             # a hair before the first word
    assert 106 < s[1]["at"] < 112 and 111.8 < s[2]["at"] < 112.2       # inside its line (letter-count estimate) / its start
    assert s[1]["said"] == "munkaerőpiacra"
    assert any("kvantumgép" in n and n.startswith("⚠") for n in notes)
    assert any("estimated" in n for n in notes)
    [ly] = video.layers_from(a, "", [])
    assert ly["at"] == pytest.approx(99.85, abs=0.01) and [x["slide"] for x in ly["slides"]] == [1, 4, 7]
    # exact word times on the cues (from a transcript) are used as they are
    cues = [{**CUES[1], "words": [106.0, 106.5, 107.1, 107.4, 107.9, 108.2, 108.6, 109.3]}]
    [x] = asyncio.run(deck_layer.resolve([{"type": "deck", "file_id": DECK, "follow": [{"slide": 2, "said": "munkaerőpiacra"}]}],
                                         [], cues=cues))
    assert x["slides"][0]["at"] == pytest.approx(109.15, abs=0.01)


def test_unresolved_follow_is_an_error_not_a_silent_slide():
    from app.tools import video
    with pytest.raises(ValueError, match="couldn't be matched"):
        video.layers_from([{"at": 5, "type": "deck", "file_id": DECK, "follow": [{"slide": 2, "said": "x"}]}], "", [])


# ---------- the builder output ----------

def _built(placement="window", total=1800.0, **kw):
    from app.tools import video
    [ly] = video.layers_from([_deck(dur=60, placement=placement, slides=[{"at": 150, "slide": 1}, {"at": 172, "slide": 4},
                                                                         {"at": 200, "slide": 7}], **kw)], "", [])
    ly["_files"] = {1: "deck-aa-001.png", 4: "deck-aa-004.png", 7: "deck-aa-007.png"}
    ly["_aspect"] = 16 / 9
    items, rows, problems = video.layer_comps([ly], 1920, 1080, total, "hu", [0.9, 0.3, 0.2], 1000)
    return ly, items, rows, problems


def test_window_deck_builds_a_carrier_with_the_split_and_one_layer_per_slide_segment():
    ly, items, rows, problems = _built("window")
    car, *sl = items
    assert car["name"].startswith("Layer L1 – deck") and car["at"] == 150 and car["dur"] == 60
    assert car["split"]["scale"] == pytest.approx(0.4578, abs=1e-3)              # the speaker on the other half
    assert car["pos"][0] > 1920 / 2 and car["w"] < 960 and abs(car["w"] / car["h"] - 16 / 9) < 0.02
    assert [x["id"] for x in sl] == ["L1.1", "L1.2", "L1.3"] and not any(x.get("split") for x in sl)
    assert [(x["at"], x["dur"]) for x in sl] == [(150, 22.45), (172, 28.45), (200, 10)]   # + the crossfade under the next
    assert [x["anim"] for x in sl] == [{"in": "fade", "out": "none"}, {"in": "fade", "out": "none"}, {"in": "fade", "out": "fade"}]
    img = sl[1]["elements"][0]
    assert img["k"] == "image" and img["file"] == "deck-aa-004.png" and img["anim"] == "none" and img["fit"] == "contain"
    assert img["size"][0] <= car["w"] and rows == [{"id": "L1", "start": 150, "end": 210, "type": "deck", "layer": True}]
    assert not problems


def test_front_and_background_decks_cover_the_frame_and_cut_has_no_crossfade():
    _, items, _, _ = _built("front", transition="cut")
    car, *sl = items
    assert (car["w"], car["h"], car["pos"]) == (1920, 1080, [960, 540]) and "split" not in car
    assert car["elements"][0]["name"] == "Backdrop" and car["elements"][0]["opacity"] == 100
    assert [(x["at"], x["dur"]) for x in sl] == [(150, 22), (172, 28), (200, 10)]
    assert [x["anim"]["in"] for x in sl] == ["fade", "none", "none"]
    ly, items, _, problems = _built("background")
    assert items[0]["deck"]["placement"] == "background" and items[1]["elements"][0]["size"] == [1920, 1080]
    from app.tools import deck_layer
    assert any("background" in p and "doesn't cut you out yet" in p for p in deck_layer.final_notes([ly]))


def test_deck_is_clipped_by_the_end_of_the_video():
    _, items, rows, problems = _built("front", total=190.0)
    assert [x["id"] for x in items] == ["L1", "L1.1", "L1.2"] and items[-1]["at"] + items[-1]["dur"] == pytest.approx(190)
    assert "ends there" in " ".join(problems)


def test_matte_part_flags_the_background_deck_clips():
    from app.tools import deck_layer
    spec = {"story": {"layers": [{"id": "L5", "type": "deck", "placement": "background", "at": 1500, "dur": 60},
                                 {"id": "L6", "type": "deck", "placement": "window", "at": 150, "dur": 60}]}}
    scene = [{"id": "L5"}, {"id": "L5.1"}, {"id": "L5.2"}, {"id": "L6"}, {"id": "L6.1"}, {"id": "L50"}]
    m = deck_layer.matte_part(spec, scene, 1797.0)
    assert m["spans"] == [[1500, 1560]] and m["model"] == "selfie_segmenter_landscape"
    assert [bool(x.get("behind")) for x in scene] == [True, True, True, False, False, False]
    assert deck_layer.matte_part({"story": {"layers": []}}, [], 10) is None


# ---------- slides → pictures (LibreOffice, cached by content) ----------

def _make_pptx(path: Path, n: int = 3) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt
    p = Presentation()
    p.slide_width, p.slide_height = Emu(12192000), Emu(6858000)          # 16:9
    for i in range(n):
        s = p.slides.add_slide(p.slide_layouts[6])
        box = s.shapes.add_shape(1, Emu(600000), Emu(600000), Emu(5000000), Emu(3000000))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*[(40, 160, 90), (200, 60, 50), (40, 90, 200)][i % 3])
        tb = s.shapes.add_textbox(Emu(600000), Emu(4200000), Emu(9000000), Emu(1200000))
        tb.text_frame.text = f"Slide {i + 1}"
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    p.save(str(path))
    return path


@pytest.mark.engines
@pytest.mark.skipif(not shutil.which("soffice") or not shutil.which("pdftoppm"), reason="LibreOffice + poppler needed")
def test_slides_are_drawn_at_the_project_size_and_cached_by_content(tmp_path, put_file):
    from PIL import Image
    from app.tools import deck_layer
    fid = put_file(_make_pptx(tmp_path / "d.pptx"))
    got = asyncio.run(deck_layer.slide_files(fid, 640, 360))
    assert not got["cached"] and got["engine"] == "libreoffice" and len(got["files"]) == 3
    for f, rgb in zip(got["files"], [(40, 160, 90), (200, 60, 50), (40, 90, 200)]):
        with Image.open(f) as im:
            assert im.size == (640, 360)
            px = im.convert("RGB").getpixel((100, 70))
            assert all(abs(a - b) < 12 for a, b in zip(px, rgb)), (f.name, px)       # the shape fill is there
    fid2 = put_file(tmp_path / "d.pptx")                                            # the same content, another file
    again = asyncio.run(deck_layer.slide_files(fid2, 640, 360))
    assert again["cached"] and again["files"] == got["files"]


# ---------- end to end: the engine project and the live layers' hashes ----------

def _engine_ok() -> bool:
    from app.tools import video
    return video.available() and bool(shutil.which("ffmpeg")) and bool(shutil.which("soffice"))


@pytest.mark.engines
@pytest.mark.skipif(not _engine_ok(), reason="EffectCraft CLI, ffmpeg and LibreOffice needed")
def test_deck_in_the_engine_project_and_a_retime_changes_only_two_segment_clips(tmp_path, put_file):
    from app.storage import output_path
    from app.tools import live_layers, video
    clip = tmp_path / "talk.mp4"
    subprocess.run(["ffmpeg", "-y", "-v", "error", "-f", "lavfi", "-i", "testsrc2=size=640x360:rate=25", "-f", "lavfi", "-i",
                    "sine=frequency=330:sample_rate=48000", "-t", "20", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac",
                    "-shortest", str(clip)], check=True)
    fid = put_file(clip)
    deck = put_file(_make_pptx(tmp_path / "d.pptx"))
    layer = {"at": 2, "until": 18, "type": "deck", "file_id": deck, "placement": "window",
             "slides": [{"at": 2, "slide": 1}, {"at": 7, "slide": 2}, {"at": 12, "slide": 3}]}
    res = asyncio.run(video.overlay({"file_id": fid, "layers": [layer], "render": False}))
    assert [x["id"] for x in res["layers"]] == ["L1"] and res["layers"][0]["text"] == "deck (3 slides, window)"
    names = sorted(n for n in video.assets_of(res["file_id"]) if n.startswith("deck-"))
    assert len(names) == 3

    def scene(pid):
        d = json.loads(output_path(pid).read_text("utf-8"))
        return {e["id"]: e for e in live_layers.scene_from(d, proxy=None, assets=video.assets_of(pid))["layers"]}

    a = scene(res["file_id"])
    assert set(a) == {"L1", "L1.1", "L1.2", "L1.3"} and all(not e["unsupported"] for e in a.values())
    assert a["L1"]["z"] < a["L1.1"]["z"] < a["L1.2"]["z"] < a["L1.3"]["z"]          # each slide above the previous
    proj = json.loads(output_path(res["file_id"]).read_text("utf-8"))
    foot = next(L for it in proj["items"].values() if it.get("name") == "Video" for L in it["kind"]["layers"]
                if L.get("name") == "Footage")
    assert "keys" in json.dumps(foot.get("props") or foot)                          # the footage moved aside (split)
    ed = asyncio.run(video.edit(res["file_id"], {"layers": [{"id": "L1", "slides": [{"at": 2, "slide": 1}, {"at": 8, "slide": 2},
                                                                                   {"at": 12, "slide": 3}]}]},
                                render=False))
    b = scene(ed["file_id"])
    changed = sorted(k for k in a if a[k]["hash"] != b[k]["hash"])
    assert changed == ["L1.1", "L1.2"], changed                                     # the carrier and slide 3 are reused
