"""A pakli arculata a fájlban él, és ami később kerül bele, az örökli (a ChatGPT-teszt 10-17: zöld–krém paklihoz adott
videós dia fehér lett). Plusz: a pakli is a felhasználó mentett arculatát kapja, ugyanazon a feloldón át, mint a
dizájn és a videó — és az eszközleírások azt mondják, amit a kód csinál."""
from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest

CREAM, GREEN = "F3EEDC", "2F5D3A"
SPEC = {"title": "Ajándék", "theme": {"bg": f"#{CREAM}", "primary": f"#{GREEN}", "accent": "#7FA65A", "text": "#22301F"},
        "slides": [{"layout": "section", "title": "Kedves Beáta"},
                   {"layout": "bullets", "title": "Emlékek", "bullets": ["Egy", "Kettő"]}]}


def _prs(fid: str):
    from pptx import Presentation
    from app.storage import output_path
    return Presentation(str(output_path(fid)))


def _theme_xml(prs) -> str:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    return prs.slide_master.part.part_related_by(RT.THEME).blob.decode("utf-8")


async def test_built_deck_carries_its_look_in_master_and_theme():
    from app.tools import pptx_deck
    prs = _prs((await pptx_deck.build(SPEC))["file_id"])
    assert str(prs.slide_master.background.fill.fore_color.rgb) == CREAM
    xml = _theme_xml(prs)
    assert 'typeface="Cambria"' in xml and GREEN in xml                   # címbetű + elsődleges szín a témában
    # egy „idegen” eszközzel hozzáadott dia (python-pptx, saját háttér nélkül) is krém marad: a mesterdiától örököl
    new = prs.slides.add_slide(prs.slide_layouts[6])
    assert pptx_deck._bg(new) is None and str(new.slide_layout.slide_master.background.fill.fore_color.rgb) == CREAM


async def test_video_slide_takes_the_decks_background_and_title_style(put_file, tmp_path):
    if shutil.which("ffmpeg") is None:
        pytest.skip("ffmpeg not installed")
    from app.tools import pptx_deck, video
    deck = (await pptx_deck.build(SPEC))["file_id"]
    mp4 = tmp_path / "v.mp4"
    subprocess.run(["ffmpeg", "-y", "-v", "error", "-f", "lavfi", "-i", "testsrc2=size=320x180:rate=25", "-t", "1",
                    "-pix_fmt", "yuv420p", str(mp4)], check=True)
    out = await video.to_pptx(put_file(mp4), deck_file_id=deck, title="A videó", after_slide=1)
    slide = _prs(out["file_id"]).slides[1]
    bg = pptx_deck._bg(slide)
    assert bg is not None and CREAM in __import__("lxml.etree", fromlist=["tostring"]).tostring(bg).decode()
    runs = [r for sh in slide.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    # a minta (1. dia) sötét fejezetdia fehér címmel — krém háttéren a cím a téma szövegszínét kapja, nem fehéret
    assert runs and runs[0].font.name == "Cambria"
    assert pptx_deck.contrast(str(runs[0].font.color.rgb), CREAM) >= 3
    assert any("deck's look" in v for v in out["check"]["verified"])


async def test_slide_added_by_deckcraft_inherits_the_background(put_file, tmp_path):
    from app.tools import crafts, pptx_deck
    if crafts.office_cli("deck") is None:
        pytest.skip("deckcraft-cli not installed")
    from pptx import Presentation
    from pptx.util import Emu
    p = tmp_path / "foreign.pptx"                                         # idegen pakli: diánként festett háttér
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for _ in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = pptx_deck._rgb(f"#{CREAM}")
        s.shapes.add_textbox(Emu(500000), Emu(500000), Emu(4000000), Emu(800000)).text_frame.text = "Szöveg"
    prs.save(str(p))
    res = await crafts.office_run("deck", file_id=put_file(p), commands=[{"command": "slide.new", "params": {"layout": "blank"}}])
    new = _prs(res["file_id"]).slides[-1]
    assert pptx_deck._bg(new) is not None and "background" in res.get("note", "")


async def test_deck_uses_the_saved_brand(monkeypatch):
    from app.brands import user as user_brands
    from app.tools import pptx_deck
    seen = []

    async def fake(param=None):
        seen.append(param)
        return {"name": "Beáta", "colors": {"primary": "#8A1C3B", "accent": "#E0A526", "bg": "#FFF8EE", "text": "#2A1A12"},
                "fonts": {}}
    monkeypatch.setattr(user_brands, "resolve", fake)
    res = await pptx_deck.build({"title": "x", "slides": [{"layout": "bullets", "title": "a", "bullets": ["b"]}]})
    assert seen == [None] and res["brand"] == "Beáta"                     # alapértelmezett arculat: automatikusan
    assert "8A1C3B" in _theme_xml(_prs(res["file_id"]))


async def test_descriptions_say_what_the_code_does():
    from app import mcp_server
    tools = {t.name: t.description for t in await mcp_server.mcp.list_tools()}
    assert 'read="inspect"' in tools["office"] and "on_slide" in tools["office"]
    assert "same file_id comes back" in tools["office"]
    assert "PPTX" in tools["find_replace"] and "PPTX" in tools["delete_pages"] and "PPTX" in tools["extract"]
    assert "saved brand" in tools["pptx_build"] and "scatter" in tools["pptx_build"]
    assert "shape_id" in tools["validate_document"]
    assert "does not fetch" in tools["facts"]
    assert "no text-to-speech" not in tools["video"].lower()
