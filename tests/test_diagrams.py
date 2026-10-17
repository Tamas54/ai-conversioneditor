"""Magyarázó ábrák (app/tools/diagrams.py): döntési fa, sávos folyamat, kapcsolati háló — elrendezési invariánsok,
szemantikus nyilak, a szereplő-regiszter, a célzott szerkesztés, a videó meglévő „graphic” elemei és a natív pakli."""
from __future__ import annotations

import copy
import itertools
import json

import pytest

SWIM = {"type": "swimlane", "title": "Hatósági ügy", "lanes": ["Ügyfél", "Hatóság", "Bíróság"],
        "steps": [{"id": "k", "actor": "Ügyfél", "text": "Kérelmet nyújt be", "kind": "start", "order": 1},
                  {"id": "h1", "actor": "Hatóság", "text": "Elsőfokú határozat", "order": 2},
                  {"id": "d", "actor": "Ügyfél", "text": "Elfogadja a döntést?", "kind": "decision", "order": 3},
                  {"id": "f", "actor": "Ügyfél", "text": "Az ügyfél fellebbez 15 napon belül", "order": 4},
                  {"id": "j", "actor": "Hatóság", "text": "A határozat jogerős", "kind": "end", "order": 4},
                  {"id": "b", "actor": "Bíróság", "text": "Ítélet", "kind": "end", "order": 5}],
        "edges": [{"from": "k", "to": "h1"}, {"from": "h1", "to": "d"}, {"from": "d", "to": "f", "label": "nem"},
                  {"from": "d", "to": "j", "label": "igen"}, {"from": "f", "to": "b"},
                  {"from": "b", "to": "h1", "label": "új eljárás"}]}
TREE = {"type": "decision_tree", "title": "Hová tegyük?", "root": "Kell-e a pénz 3 éven belül?",
        "branches": [{"label": "igen", "text": "Állampapír", "kind": "outcome"},
                     {"label": "nem", "text": "Mekkora kockázat?", "branches": [
                         {"label": "kicsi", "text": "Vegyes alap"}, {"label": "nagy", "text": "Részvényindex-alap"}]}]}
NET = {"type": "network", "title": "Tulajdonosok", "center": "h",
       "nodes": [{"id": "h", "text": "Alfa Holding", "group": "Cég"}, {"id": "p", "text": "Kovács Péter", "group": "Tulajdonos"},
                 {"id": "tv", "text": "Hír TV", "group": "Média"}, {"id": "a", "text": "Állami hirdetések", "group": "Állam"},
                 {"id": "n", "text": "Nézettség", "group": "Média"}],
       "edges": [{"from": "p", "to": "h", "kind": "hierarchy", "label": "60%"},
                 {"from": "h", "to": "tv", "kind": "hierarchy", "certainty": "uncertain"},
                 {"from": "a", "to": "tv", "kind": "money", "label": "2,1 milliárd Ft"},
                 {"from": "tv", "to": "n", "kind": "cause", "certainty": "disputed"},
                 {"from": "n", "to": "a", "kind": "correlation"},
                 {"from": "p", "to": "a", "kind": "link"}]}
SOURCE = "Kovács Péter 60%-ban birtokolja. 2,1 milliárd Ft állami hirdetés. 15 napon belül fellebbez. 3 éven belül."
FORMATS = {"vertical": (1080, 1920), "landscape": (1920, 1080), "square": (1080, 1080)}


def _norm(c):
    from app.tools import diagrams, video
    nums = video._text_numbers(SOURCE)
    spec, notes = diagrams.from_plan(c, lambda t: [x for x in video._text_numbers(t) if x not in nums], video._clean)
    assert spec is not None, notes
    return spec


def _boxes(lay):
    return [(b["x"], b["y"], b["w"], b["h"]) for b in lay["nodes"].values()]


@pytest.mark.parametrize("c", [SWIM, TREE, NET], ids=["swimlane", "tree", "network"])
@pytest.mark.parametrize("fmt", list(FORMATS))
def test_layout_has_no_overlapping_boxes_and_stays_in_its_region(c, fmt):
    from app.tools import diagrams
    spec = _norm(c)
    W, H = FORMATS[fmt]
    rw, rh = W * 0.9, H * 0.63
    lay = diagrams.layout(spec, (rw, rh), cw=0.56, s_max=30, s_min=22, lang="hu")
    bx = _boxes(lay)
    for a, b in itertools.combinations(bx, 2):
        assert diagrams._overlap(a, b) == 0, (fmt, a, b)
    labels = [(e["label"]["x"], e["label"]["y"], e["label"]["w"], e["label"]["h"]) for e in lay["edges"] if e.get("label")]
    for lb in labels:                               # a címke nem takar dobozt
        assert all(diagrams._overlap(lb, b) < 0.02 * lb[2] * lb[3] for b in bx), (fmt, lb)
    x0, y0, x1, y1 = lay["bbox"]
    assert x0 >= -1 and y0 >= -1 and x1 <= rw + 1 and y1 <= rh + 1
    assert all(it["y"] <= rh + 1 for it in lay["legend"])
    assert lay["s"] >= 17                           # 1080-as rövid oldalon is olvasható betű


def test_the_layout_differs_by_orientation_and_is_deterministic():
    from app.tools import diagrams
    spec = _norm(TREE)
    tall = diagrams.layout(spec, (972, 1210), cw=0.56, s_max=30, s_min=22)
    wide = diagrams.layout(spec, (1728, 680), cw=0.56, s_max=30, s_min=22)
    asp = lambda L: (L["bbox"][2] - L["bbox"][0]) / (L["bbox"][3] - L["bbox"][1])           # noqa: E731
    assert asp(wide) > asp(tall)                                                              # képarányonként más elrendezés
    wider = _norm({**TREE, "branches": TREE["branches"] + [{"label": "talán", "text": "Lekötött betét"},
                                                           {"label": "később", "text": "Nyugdíjpénztár"}]})
    assert diagrams.layout(wider, (972, 1210), cw=0.56, s_max=30, s_min=22)["dir"] == "outline"   # álló: vázlat-fa
    assert diagrams.layout(wider, (1728, 680), cw=0.56, s_max=30, s_min=22)["dir"] != "outline"
    a = diagrams.to_graphic(copy.deepcopy(spec), 1080, 1920, [1, 0.35, 0.2], 12.0)
    b = diagrams.to_graphic(copy.deepcopy(spec), 1080, 1920, [1, 0.35, 0.2], 12.0)
    assert json.dumps(a, sort_keys=True) == json.dumps(b, sort_keys=True)


def test_every_edge_kind_has_one_style_and_correlation_never_looks_causal():
    from app.tools import diagrams
    assert set(diagrams.STYLE) == set(diagrams.EDGE_KINDS)
    assert len({(v["dark"], v["head"], v["width"]) for v in diagrams.STYLE.values()}) == len(diagrams.STYLE)
    assert diagrams.STYLE["correlation"]["head"] not in ("triangle", "chevron") and diagrams.STYLE["correlation"]["sym"]
    assert diagrams.STYLE["money"]["width"] > max(v["width"] for k, v in diagrams.STYLE.items() if k != "money")
    spec = _norm(NET)
    els = diagrams.to_graphic(spec, 1920, 1080, None, 12.0)
    corr = diagrams.STYLE["correlation"]["dark"]
    assert not [e for e in els if e["shape"] == "arrow" and e.get("stroke") == corr]          # nincs nyílhegye
    assert not [e for e in els if e["shape"] == "path" and e.get("stroke") == corr and len(e["points"]) == 3
                and e["anim"] == "appear"]                                                     # nincs „V” hegye sem
    lines = [e for e in els if e.get("stroke") == corr and e["shape"] == "line"]
    halves = [(a, b) for a, b in itertools.combinations(lines, 2) if a["from"] == b["from"]]
    assert len(halves) == 1 and halves[0][0]["at"] == halves[0][1]["at"]                       # középről kifelé, egyszerre
    for kind in ("hierarchy", "money", "cause", "link"):
        assert [e for e in els if e.get("stroke") == diagrams.STYLE[kind]["dark"]], kind
    legend = " ".join(e.get("text") or "" for e in els)
    assert "az okozás nem igazolt" in legend                    # a jelmagyarázat kimondja (magyar videón magyarul)
    dashed = {tuple(e["dash"]) for e in els if e.get("dash")}
    want = lambda cert, kind: tuple(round(v * diagrams.STYLE[kind]["width"], 2) for v in diagrams.DASH[cert])   # noqa: E731
    assert want("uncertain", "hierarchy") in dashed and want("disputed", "cause") in dashed      # szaggatott / pontozott
    assert any(e["shape"] == "text" and e["text"] == "?" for e in els)


def test_reveal_builds_up_step_by_step_with_the_rest_faded():
    from app.tools import diagrams
    spec = _norm(TREE)
    at, pace, end = diagrams.timing(spec, 12.0)
    order = diagrams._order(spec)
    root = next(i for i, k in order.items() if k == 0)
    assert at[root] == min(at.values()) and len(set(at.values())) == len(at)                  # gyökér, majd ágról ágra
    els = diagrams.to_graphic(spec, 1080, 1920, None, 12.0)
    dims = [e["dim"][0] for e in els if e.get("dim")]
    assert dims and all(a < b <= end + 1e-6 for a, b in dims)
    assert not any(e.get("dim") for e in diagrams.to_graphic({**spec, "reveal": "all"}, 1080, 1920, None, 12.0))
    assert not any(e.get("dim") for e in diagrams.to_graphic({**spec, "focus": False}, 1080, 1920, None, 12.0))
    slow = {**spec, "pace": 2.0}
    need = diagrams.seconds_needed(slow)
    assert need == pytest.approx(diagrams.T_START + diagrams.steps_of(spec) * 2.0 + diagrams.HOLD, abs=0.1)
    assert "needs" in (diagrams.duration_note(slow, 6.0) or "") and diagrams.duration_note(slow, need + 1) is None
    late = {**spec, "nodes": [{**n, "at": 9.0} if n["id"] == root else n for n in spec["nodes"]]}
    assert diagrams.timing(late, 12.0)[0][root] == 9.0                                         # csúcsonkénti at felülír


def test_compiled_elements_are_the_existing_graphic_vocabulary():
    """A fordító kimenete a meglévő graphic elemkészlet: a _graphic_from_plan változatlanul átengedi (fixpont), a
    _graphic_elements a builder rétegeivé rajzolja (szaggatással és halványítással)."""
    from app.tools import diagrams, video
    spec = _norm(NET)
    els = diagrams.to_graphic(spec, 1920, 1080, None, 12.0)
    assert {e["shape"] for e in els} <= set(video._GFX_SHAPES)
    chunk = els[:80]
    back, _ = video._graphic_from_plan({"elements": chunk}, SOURCE + " ? " + " ".join(e.get("text") or "" for e in chunk))
    strip = lambda e: {k: v for k, v in e.items() if k != "text"}                              # noqa: E731  (a szöveg szóközei)
    assert [strip(e) for e in back["elements"]] == [strip({**e, **{k: float(v) for k, v in e.items() if isinstance(v, int) and not isinstance(v, bool)}}) for e in chunk]
    layers = video.chart_elements(spec, 1920, 1080, "hu", None, dur=12.0)
    assert any(x.get("dash") for x in layers if x["k"] == "path")
    assert any(x.get("dim") for x in layers)
    tree = video.chart_elements(_norm(TREE), 1080, 1920, "hu", None, dur=12.0)
    assert any(x.get("dim") for x in tree if x["name"].startswith("Arrowhead"))               # a nyílhegy is halványul


def test_actor_colours_are_stable_across_scenes_and_edits():
    from app.tools import diagrams
    a = _norm({**SWIM, "actors": {"Bíróság": {"color": "#7744CC"}}})
    b = _norm({"type": "network", "title": "Ki kivel", "nodes": [{"id": "x", "text": "Bírói tanács", "actor": "Bíróság"},
                                                                {"id": "y", "text": "Ügyfél képviselője", "actor": "Ügyfél"}],
               "edges": [{"from": "x", "to": "y"}]})
    story = {"scenes": [{"chart": copy.deepcopy(a)}, {"chart": copy.deepcopy(b)}]}
    diagrams.prepare(story, "#FF5A36")
    reg = story["actors"]
    assert reg["Bíróság"]["color"] == "#7744CC"                                                # a megadott szín nyer
    c1, c2 = story["scenes"][0]["chart"]["_actors"], story["scenes"][1]["chart"]["_actors"]
    assert c1["Ügyfél"] == c2["Ügyfél"] and c1["Bíróság"] == c2["Bíróság"]
    assert len({v["color"] for v in reg.values()}) == len(reg)                                 # különböző szereplő, különböző szín
    later = {"actors": reg, "scenes": [{"chart": copy.deepcopy(b)}]}                          # az első jelenet törölve
    diagrams.prepare(later, "#2255FF")
    assert later["actors"]["Ügyfél"] == reg["Ügyfél"]                                           # a regiszter tárolt: marad
    assert diagrams.palette("#FF5A36", 5) == diagrams.palette("#FF5A36", 5)


def test_numbers_go_through_the_same_source_check():
    from app.tools import video
    bad_label = {**NET, "edges": NET["edges"] + [{"from": "h", "to": "n", "label": "99 millió Ft"}]}
    spec, notes = video._chart_from_plan(bad_label, SOURCE)
    assert spec and not any(e.get("label") == "99 millió Ft" for e in spec["edges"])
    assert any(n.startswith("⚠") and "edge label" in n for n in notes)
    bad_box = {**TREE, "root": "Kell-e a pénz 7 éven belül?"}
    spec, notes = video._chart_from_plan(bad_box, SOURCE)
    assert spec is None and "isn't in the source" in notes[0]
    spec, notes = video._chart_from_plan(TREE, SOURCE)
    assert spec and "no number that isn't in the source" in notes[0]


def test_broken_diagrams_are_refused_with_what_to_fix():
    from app.tools import diagrams
    with pytest.raises(ValueError, match="not a node .*node ids"):
        diagrams.from_plan({**NET, "edges": [{"from": "h", "to": "nincs"}]})
    with pytest.raises(ValueError, match="two parents"):
        diagrams.from_plan({"type": "decision_tree", "nodes": [{"id": "a", "text": "A?"}, {"id": "b", "text": "B"},
                                                               {"id": "c", "text": "C"}],
                            "edges": [{"from": "a", "to": "b"}, {"from": "a", "to": "c"}, {"from": "b", "to": "c"}]})
    with pytest.raises(ValueError, match="no actor"):
        diagrams.from_plan({"type": "swimlane", "steps": [{"text": "Lépés"}, {"text": "Másik", "actor": "Hatóság"}]})
    with pytest.raises(ValueError, match="edge kind .*correlation"):
        diagrams.from_plan({**NET, "edges": [{"from": "h", "to": "p", "kind": "magic"}]})
    with pytest.raises(ValueError, match="Split it"):
        diagrams.from_plan({"type": "network", "nodes": [f"N{i}" for i in range(20)]})


def test_shorthands_expand_to_ids_and_edges():
    from app.tools import diagrams
    t, _ = diagrams.from_plan(TREE)
    assert [n["id"] for n in t["nodes"]] == ["n1", "n1.1", "n1.2", "n1.2.1", "n1.2.2"]
    assert {(e["from"], e["to"], e.get("label")) for e in t["edges"]} >= {("n1", "n1.1", "igen"), ("n1.2", "n1.2.2", "nagy")}
    s, _ = diagrams.from_plan({"type": "swimlane", "steps": [{"text": "A", "actor": "X"}, {"text": "B", "actor": "Y"},
                                                             {"text": "C", "actor": "X"}]})
    assert [(e["from"], e["to"]) for e in s["edges"]] == [("n1", "n2"), ("n2", "n3")] and s["lanes"] == ["X", "Y"]
    n, _ = diagrams.from_plan({"type": "network", "edges": [{"from": "MNB", "to": "Bankok", "kind": "hierarchy"}]})
    assert [x["text"] for x in n["nodes"]] == ["MNB", "Bankok"] and n["edges"][0]["from"] == "n1"


def test_targeted_edit_changes_only_what_is_sent():
    from app.tools import video
    base = video._scene_from({"headline": "Háló", "caption": SOURCE, "chart": NET}, SOURCE, [])
    notes: list[str] = []
    patch = {"chart": {"edges": [{"from": "p", "to": "h", "certainty": "uncertain"}],
                       "nodes": [{"id": "uj", "text": "Új portál", "group": "Média"}],
                       "pace": 1.5}}
    patch["chart"]["edges"].append({"from": "h", "to": "uj", "kind": "hierarchy"})
    out = video._scene_from(patch, SOURCE, notes, base)
    old_e = {(e["from"], e["to"]): e for e in base["chart"]["edges"]}
    new_e = {(e["from"], e["to"]): e for e in out["chart"]["edges"]}
    assert new_e[("p", "h")]["certainty"] == "uncertain" and new_e[("p", "h")]["label"] == "60%"
    assert {k: v for k, v in new_e.items() if k not in (("p", "h"), ("h", "uj"))} == {k: v for k, v in old_e.items() if k != ("p", "h")}
    assert ("h", "uj") in new_e and out["chart"]["pace"] == 1.5 and len(out["chart"]["nodes"]) == len(base["chart"]["nodes"]) + 1
    gone = video._scene_from({"chart": {"nodes": [{"id": "n", "remove": True}]}}, SOURCE, [], base)
    assert "n" not in [x["id"] for x in gone["chart"]["nodes"]] and all("n" not in (e["from"], e["to"]) for e in gone["chart"]["edges"])
    tree = video._scene_from({"headline": "Fa", "caption": SOURCE, "chart": TREE}, SOURCE, [])
    split = video._scene_from({"chart": {"nodes": [{"id": "n1.1", "kind": "condition", "branches": [
        {"label": "rövid", "text": "Rövid lejáratú"}, {"label": "hosszú", "text": "Hosszú lejáratú"}]}]}}, SOURCE, [], tree)
    kids = [e["to"] for e in split["chart"]["edges"] if e["from"] == "n1.1"]
    assert kids == ["n1.1.1", "n1.1.2"] and len(split["chart"]["nodes"]) == len(tree["chart"]["nodes"]) + 2


def test_deck_diagram_is_native_editable_shapes(tmp_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from app.tools import diagrams, pptx_deck
    spec = pptx_deck.diagram_slides({"slides": [{"layout": "diagram", "title": "Háló", "diagram": NET},
                                                {"layout": "diagram", "title": "Folyamat", "diagram": SWIM}]})
    reg = pptx_deck.diagram_actors(spec, "#FCA311")
    prs = Presentation()
    t = dict(pptx_deck.DEFAULT_THEME)
    for s in spec["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        pptx_deck._diagram(slide, s, t, {k: reg[k] for k in diagrams.actor_names(s["diagram"])})
    p = tmp_path / "d.pptx"
    prs.save(str(p))
    prs = Presentation(str(p))
    for slide, s in zip(prs.slides, spec["slides"]):
        kinds = {sh.shape_type for sh in slide.shapes}
        assert MSO_SHAPE_TYPE.PICTURE not in kinds
        nodes = [sh for sh in slide.shapes if sh.name.startswith("Node – ")]
        assert len(nodes) == len(s["diagram"]["nodes"]) and all(sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE for sh in nodes)
        texts = {sh.name[7:]: sh.text_frame.text.replace("\n", " ") for sh in nodes}
        assert texts == {n["id"]: n["text"] for n in s["diagram"]["nodes"]}               # a szöveg a dobozban, szerkeszthetően
        edges = [sh for sh in slide.shapes if sh.name.startswith("Edge – ")]
        assert len(edges) == len(s["diagram"]["edges"])
        bound = [sh for sh in edges if sh._element.tag.endswith("cxnSp")
                 and (sh._element.find(".//{*}stCxn") is not None or sh._element.find(".//{*}endCxn") is not None)]
        assert bound                                                                  # az egyenes él a dobozokhoz kötött
        for sh in edges:
            ln = sh._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
            tail = ln.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
            kind = next(e.get("kind") or diagrams.default_kind(s["diagram"]["type"]) for e in s["diagram"]["edges"]
                        if sh.name == f"Edge – {e['from']}→{e['to']}")
            if kind == "correlation":
                assert tail is None or tail.get("type") == "oval"                     # együttmozgás: nincs nyílhegy
            elif diagrams.STYLE[kind]["head"] == "triangle":
                assert tail is not None and tail.get("type") == "triangle"
    lanes = [sh.fill.fore_color.rgb for sh in prs.slides[1].shapes if sh.name.startswith("Lane – Ügyfél")]
    assert lanes and reg["Ügyfél"]["color"]


@pytest.mark.engines
async def test_video_with_a_diagram_and_a_targeted_edit_changes_only_that_scene():
    """Valódi építés (render nélkül): a jelenet leírása a projektben marad; a célzott szerkesztés csak azt a jelenetet
    változtatja, a szereplő-regiszter a szerkesztés után is ugyanaz."""
    from app.storage import read_meta
    from app.tools import video
    if not video.available():
        pytest.skip("EffectCraft not installed")
    body = {"mode": "article", "format": "vertical", "pictures": "none", "narration": "none", "render": False, "title": "Ábrák",
            "scenes": [{"headline": "Folyamat", "caption": SOURCE, "chart": SWIM, "duration": 8},
                       {"headline": "Háló", "caption": SOURCE, "chart": NET, "duration": 8}]}
    res = await video.make(body)
    spec = read_meta(res["file_id"])["spec"]
    assert spec["story"]["scenes"][1]["chart"]["type"] == "network" and spec["story"]["actors"]
    out = await video.edit(res["file_id"], {"scenes": [{"id": "s2", "chart": {"edges": [{"from": "a", "to": "tv", "certainty": "uncertain"}]}}]})
    new = read_meta(out["file_id"])["spec"]
    assert new["story"]["scenes"][0] == spec["story"]["scenes"][0]
    assert new["story"]["scenes"][1] != spec["story"]["scenes"][1]
    edge = next(e for e in new["story"]["scenes"][1]["chart"]["edges"] if (e["from"], e["to"]) == ("a", "tv"))
    assert edge["certainty"] == "uncertain" and edge["label"] == "2,1 milliárd Ft"
    assert new["story"]["actors"] == spec["story"]["actors"]
