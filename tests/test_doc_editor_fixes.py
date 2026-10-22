"""A dokumentumszerkesztő hibái egy MCP-ügynök tesztjelentéséből (10-22): a [[N7]]-féle maradvány (a lábjegyzetben is)
a check-be kerül; a több bekezdés beszúrása egyértelmű (text.paragraphs), a \\r\\n nem hagy kóbor sortörést."""
from __future__ import annotations

import asyncio
import zipfile
from pathlib import Path

import pytest


def _doc(path: Path, body: str, footnote: str | None = None) -> Path:
    from docx import Document
    d = Document()
    d.add_paragraph(body)
    d.save(path)
    if footnote is not None:
        W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        tmp = path.with_suffix(".tmp")
        with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w") as zo:
            for it in zi.infolist():
                zo.writestr(it, zi.read(it.filename))
            zo.writestr("word/footnotes.xml", f'<w:footnotes {W}><w:footnote w:id="1"><w:p><w:r><w:t>{footnote}</w:t></w:r>'
                                              f'</w:p></w:footnote></w:footnotes>')
        tmp.replace(path)
    return path


def test_double_bracket_markers_are_placeholders_anywhere():
    from app.tools.content_check import _PLACEHOLDER
    hits = lambda t: [m.group() for m in _PLACEHOLDER.finditer(t)]                     # noqa: E731
    assert hits("A tétel [[N7]] szerint, lásd [[N12]].") == ["[[N7]]", "[[N12]]"]
    assert hits("Lásd [1], [12–14] és [sic].") == []                                # hivatkozás, nem helyőrző
    assert hits("Kedves [a vendég neve]!") == ["[a vendég neve]"]


def test_the_check_names_new_leftover_markers_also_in_footnotes(tmp_path):
    from app.tools import checks
    from app.tools.content_check import placeholder_quotes
    src = _doc(tmp_path / "a.docx", "A munkaerőhiány nem zárja ki a munkanélküliséget.")
    out = _doc(tmp_path / "b.docx", "A munkaerőhiány nem zárja ki a munkanélküliséget [[N7]].", footnote="KSH [[N8]], 2025.")
    assert placeholder_quotes(out) == [("the text", "[[N7]]"), ("a footnote", "[[N8]]")]
    chk = checks.check_docx(out, src, "wordcraft", {})
    f = " ".join(chk["findings"])
    assert "Leftover placeholders" in f and "“[[N7]]” (the text)" in f and "“[[N8]]” (a footnote)" in f
    same = checks.check_docx(out, out, "wordcraft", {})                             # ami már ott volt, nem újra
    assert "Leftover" not in " ".join(same["findings"])


def test_paragraphs_and_line_breaks_are_told_apart():
    from app.tools import wordcraft
    assert wordcraft.expand([{"command": "text.paragraphs", "params": {"text": "Első.\nMásodik.\n\nHarmadik."}}]) == [
        {"command": "text.insert", "params": {"text": "Első.\rMásodik.\rHarmadik."}}]
    assert wordcraft.expand([{"command": "text.paragraphs", "params": {"paragraphs": ["A", "B"]}}])[0]["params"]["text"] == "A\rB"
    assert wordcraft.expand([{"command": "text.insert", "params": {"text": "a\r\nb\nc"}}])[0]["params"]["text"] == "a\rb\nc"
    assert wordcraft.expand(["caret.docEnd"]) == ["caret.docEnd"]
    with pytest.raises(ValueError, match="nothing to insert"):
        wordcraft.expand([{"command": "text.paragraphs", "params": {"text": "\n\n"}}])


def _wc_ok() -> bool:
    from app.tools import wordcraft
    return wordcraft.available()


@pytest.mark.engines
@pytest.mark.skipif(not _wc_ok(), reason="WordCraft CLI needed")
def test_wordcraft_makes_real_paragraphs(tmp_path, put_file):
    from docx import Document
    from app.tools import checks, wordcraft
    fid = put_file(_doc(tmp_path / "c.docx", "Bevezetés."))
    r = asyncio.run(wordcraft.run(file_id=fid, commands=[
        {"command": "caret.docEnd", "params": {}}, {"command": "text.newParagraph", "params": {}},
        {"command": "text.paragraphs", "params": {"text": "Első bekezdés.\nMásodik bekezdés.\r\nHarmadik bekezdés."}}]))
    paras = [p.text for p in Document(str(checks._path(r["file_id"]))).paragraphs if p.text.strip()]
    assert paras == ["Bevezetés.", "Első bekezdés.", "Második bekezdés.", "Harmadik bekezdés."]
