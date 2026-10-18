"""„Kérj változtatást” Word-dokumentumon: a szín színként megy be, a csere nem rontja a formázást.

UI-próba (10-18): „A cím legyen sötétkék, és a hely legyen: Városi Könyvtár…” → a modell a "font" mezőbe írta a
„dark blue”-t (a style-lépésnek nem volt színe) → a cím betűtípusa „dark blue” lett, a színe nem változott;
a csere az egész bekezdést a félkövér „Helyszín:” futásba írta.
"""
from __future__ import annotations

import asyncio

import docx
from docx.shared import RGBColor

from app.tools import doc_revise as R


def _doc():
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Helyszín:").bold = True
    p.add_run(" ")
    p.add_run("[a könyvtár neve]")
    d.add_paragraph("Meghívó az olvasóestre", style="Heading 1")
    return d


def test_replace_keeps_run_formatting():
    d = _doc()
    asyncio.run(R.apply(d, [{"op": "replace_text", "find": "[a könyvtár neve]", "replace": "Városi Könyvtár"}]))
    assert [(r.text, r.bold) for r in d.paragraphs[0].runs] == [("Helyszín:", True), (" ", None), ("Városi Könyvtár", None)]


def test_replace_across_runs_and_every_occurrence():
    d = _doc()
    asyncio.run(R.apply(d, [{"op": "replace_text", "find": ": [a", "replace": ": a"}]))
    assert d.paragraphs[0].text == "Helyszín: a könyvtár neve]" and d.paragraphs[0].runs[0].bold
    d = docx.Document()
    d.add_paragraph("a b a")
    asyncio.run(R.apply(d, [{"op": "replace_text", "find": "a", "replace": "aa"}]))
    assert d.paragraphs[0].text == "aa b aa"


def test_colour_is_never_a_font_name():
    d = _doc()
    asyncio.run(R.apply(d, [{"op": "style", "from": 1, "font": "dark blue"},
                            {"op": "style", "from": 0, "color": "#C00000", "font": "Georgia"}]))
    title, label = d.paragraphs[1].runs[0], d.paragraphs[0].runs[0]
    assert title.font.name is None and title.font.color.rgb == RGBColor(0x1F, 0x38, 0x64)
    assert label.font.name == "Georgia" and label.font.color.rgb == RGBColor(0xC0, 0x00, 0x00)
