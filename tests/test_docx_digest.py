"""A DOCX tartalmi ujjlenyomata a szerkezetet is látja (10-22, ügynök-jelentés): a valódi lábjegyzet és az oldaltörés-
beállítás után a check „Nothing changed”-et mondott, és a szerkesztő az ilyen mentést változatlannak vette (eldobta).
Ami most számít: lábjegyzet, megjegyzés, követett beszúrás, törés, a bekezdés tördelési beállításai, élőláb, stílus.
A semmit nem változtató újramentés (python-docx, WordCraft) továbbra is „ugyanaz”."""
from __future__ import annotations

import asyncio
import re
import shutil
import zipfile
from pathlib import Path

import pytest


def _base(path: Path) -> Path:
    from docx import Document
    d = Document()
    d.add_heading("1. Előzmények", 1)
    d.add_paragraph("A munkaerőhiány nem zárja ki a munkanélküliséget.")
    d.add_paragraph("Második bekezdés, árvíztűrő tükörfúrógép.")
    d.save(path)
    return path


def _patched(src: Path, dst: Path, fn, extra: dict | None = None) -> Path:
    """Egy rész XML-je módosítva (a többi bájtra azonos)."""
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "word/document.xml":
                data = fn(data.decode("utf-8")).encode("utf-8")
            zo.writestr(it, data)
        for name, body in (extra or {}).items():
            zo.writestr(name, body)
    return dst


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def test_structure_changes_are_changes(tmp_path):
    from app.tools.content_digest import same_content
    a = _base(tmp_path / "a.docx")
    first_ppr = lambda x, add: re.sub(r"<w:pPr>", "<w:pPr>" + add, x, count=1)        # noqa: E731
    keep = _patched(a, tmp_path / "keep.docx", lambda x: first_ppr(x, "<w:keepNext/>"))
    brk = _patched(a, tmp_path / "brk.docx", lambda x: first_ppr(x, "<w:pageBreakBefore/>"))
    fn = _patched(a, tmp_path / "fn.docx", lambda x: x.replace("munkanélküliséget.</w:t></w:r>",
                                                                 "munkanélküliséget.</w:t></w:r><w:r><w:footnoteReference w:id=\"1\"/></w:r>", 1),
                  {"word/footnotes.xml": f'<w:footnotes {W}><w:footnote w:id="1"><w:p><w:r><w:t>KSH, 2025.</w:t></w:r></w:p></w:footnote></w:footnotes>'})
    ins = _patched(a, tmp_path / "ins.docx", lambda x: x.replace("Második bekezdés,</w:t>", "Második bekezdés,</w:t>", 1)
                   .replace("</w:p></w:body>", "<w:ins w:id=\"9\" w:author=\"X\"><w:r><w:t>ÚJ</w:t></w:r></w:ins></w:p></w:body>", 1)
                   if "</w:p></w:body>" in x else re.sub(r"(</w:r>)(</w:p>)(?!.*</w:p>)", r"\1<w:ins w:id=\"9\" w:author=\"X\"><w:r><w:t>ÚJ</w:t></w:r></w:ins>\2", x, flags=re.S))
    pg = _patched(a, tmp_path / "pg.docx", lambda x: x.replace("tükörfúrógép.</w:t></w:r>", "tükörfúrógép.</w:t></w:r><w:r><w:br w:type=\"page\"/></w:r>", 1))
    for changed in (keep, brk, fn, ins, pg):
        assert not same_content(a, changed), changed.name
    from docx import Document
    again = tmp_path / "again.docx"
    Document(str(a)).save(str(again))                      # újramentés változtatás nélkül: ugyanaz
    assert same_content(a, again)


def _wc_ok() -> bool:
    from app.tools import wordcraft
    return wordcraft.available()


@pytest.mark.engines
@pytest.mark.skipif(not _wc_ok(), reason="WordCraft CLI needed")
def test_wordcraft_footnote_and_keep_with_next_are_seen_and_a_plain_save_is_not(tmp_path, put_file):
    from app.storage import output_path
    from app.tools import checks, wordcraft
    from app.tools.content_digest import same_content
    fid = put_file(_base(tmp_path / "t.docx"))
    src = output_path(fid) if output_path(fid).exists() else Path(checks._path(fid))
    run = lambda cmds: asyncio.run(wordcraft.run(file_id=fid, commands=cmds))                # noqa: E731
    plain = run([{"command": "caret.docEnd", "params": {}}, {"command": "text.insert", "params": {"text": ""}}])
    fn = run([{"command": "select.text", "params": {"text": "munkanélküliséget"}}, {"command": "caret.end", "params": {}},
              {"command": "references.footnote", "params": {"text": "KSH, 2025."}}])
    keep = run([{"command": "edit.goto", "params": {"paragraph": 1}}, {"command": "para.keepNext", "params": {}}])
    path = lambda r: Path(checks._path(r["file_id"]))                                       # noqa: E731
    assert not same_content(src, path(fn)) and not same_content(src, path(keep))
    if plain["file_id"] != fid:                                    # a puszta mentés nem változás
        assert same_content(src, path(plain))
    chk = checks.check_docx(path(fn), src, "wordcraft", {})
    assert "Nothing changed" not in " ".join(chk["findings"]), chk
