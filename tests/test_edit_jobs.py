"""edit_with_instruction háttérben (10-22, egy MCP-ügynök tesztjelentése: „a nagyobb szerkesztési kérés időtúllépéssel
megszakadt, az új változat sem jelent meg”): a válasz ~40 s-on belül jön, a hosszabb szerkesztés job_id-t ad, amit
le lehet kérdezni; a változat a kliens nélkül is elkészül névvel és ellenőrzéssel; az elvágott AI-terv érthető hiba."""
from __future__ import annotations

import asyncio
import json

import pytest


def _doc(tmp_path, put_file):
    from docx import Document
    d = Document()
    d.add_heading("Bevezetés", 1)
    d.add_paragraph("A munkaerőhiány miatt nem lehet munkanélküliség.")
    d.add_paragraph("Második bekezdés.")
    p = tmp_path / "tanulmany.docx"
    d.save(p)
    return put_file(p)


@pytest.fixture()
def planner(monkeypatch):
    from app import mcp_server
    from app.tools import edit_ai
    state = {"delay": 0.0, "plan": None, "error": None}

    async def plan(text, instruction, format_hint=""):
        await asyncio.sleep(state["delay"])
        if state["error"]:
            raise state["error"]
        return state["plan"] or {"operations": [{"type": "find_replace", "pattern": "nem lehet munkanélküliség",
                                                 "replacement": "a munkanélküliség is jelen lehet"}],
                                 "confidence": 0.95, "warnings": []}, "fake-model"
    monkeypatch.setattr(edit_ai, "plan_edits", plan)
    monkeypatch.setattr(mcp_server, "_pool", lambda: None)
    return state


async def test_a_quick_edit_answers_at_once_with_a_named_version(planner, tmp_path, put_file):
    from app.mcp_server import edit_with_instruction
    from app.storage import read_meta
    fid = _doc(tmp_path, put_file)
    res = await edit_with_instruction(file_id=fid, instruction="Javítsd a logikai hibát.")
    assert res["executed"] and res["file_id"] != fid and res["check"]["status"] in ("ok", "needs_check")
    m = read_meta(res["file_id"])
    assert m["operation"] == "edit_with_instruction" and m["source_file_id"] == fid and m["label"].endswith(".docx")


async def test_a_long_edit_returns_a_job_and_finishes_without_the_client(planner, tmp_path, put_file, monkeypatch):
    from app import mcp_server
    from app.mcp_server import edit_with_instruction
    from app.storage import read_meta
    monkeypatch.setattr(mcp_server, "EDIT_INLINE_WAIT_S", 0.1)
    planner["delay"] = 0.6
    fid = _doc(tmp_path, put_file)
    first = await edit_with_instruction(file_id=fid, instruction="Írd át az egészet.")
    assert first["status"] == "running" and first["job_id"] and "planning the edit" in first["stage"]
    assert f'job_id=\\"{first["job_id"]}\\"' in json.dumps(first["hint"]) or first["job_id"] in first["hint"]
    await asyncio.sleep(1.0)                                   # a kliens közben nem kérdez — a munka akkor is kész
    done = await edit_with_instruction(job_id=first["job_id"], wait=5)
    assert done["executed"] and read_meta(done["file_id"])["operation"] == "edit_with_instruction" and "check" in done


async def test_a_cut_off_plan_says_what_to_do_and_changes_nothing(planner, tmp_path, put_file, monkeypatch):
    from app import mcp_server
    from app.llm.siliconflow import LLMTruncated
    from app.mcp_server import edit_with_instruction
    monkeypatch.setattr(mcp_server, "EDIT_INLINE_WAIT_S", 0.1)
    planner["delay"], planner["error"] = 0.3, LLMTruncated("m", 12000)
    fid = _doc(tmp_path, put_file)
    job = await edit_with_instruction(file_id=fid, instruction="Írd át az egészet feszesebbre.")
    with pytest.raises(ValueError, match="section by section"):
        await edit_with_instruction(job_id=job["job_id"], wait=5)
    with pytest.raises(ValueError, match="needs file_id and instruction"):
        await edit_with_instruction(file_id=fid)


async def test_the_llm_says_when_its_answer_was_cut(monkeypatch):
    import httpx
    from app.config import settings
    from app.llm import siliconflow

    class R:
        status_code = 200

        def json(self):
            return {"choices": [{"finish_reason": "length", "message": {"content": '{"operations": [{"type": "find'}}]}

    class C:
        def __init__(self, *a, **k):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *a):
            return False

        async def post(self, *a, **k):
            return R()
    monkeypatch.setattr(httpx, "AsyncClient", C)
    monkeypatch.setattr(settings, "SILICONFLOW_API_KEY", "x")
    with pytest.raises(siliconflow.LLMTruncated):
        await siliconflow._call_model("m", [{"role": "user", "content": "x"}], max_tokens=10)
