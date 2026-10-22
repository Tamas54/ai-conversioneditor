"""Szerkesztési módok határokkal (10-22): előbb kérdez („mit akar a költő”), aztán követett változásként javasol, gépi
őrökkel (számok, megnevezett szavak, nyelvi javításnál csak kis változás, új gondolat csak engedéllyel, a hossz-határ) —
semmi nincs elfogadva. Az AI itt utánzat."""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pytest

P_CUT = ("A munkaerőhiány 2025-ben 4,5 százalékos volt, ezért a vállalatok többsége bért emelt, és a bérek így tovább "
         "nőttek, ahogy azt már korábban is láttuk és sokszor elmondtuk.")
P_NUM = "A KSH adatai szerint 120 ezer betöltetlen állás volt az év végén."
P_TERM = "A Magyar Nemzeti Bank szerint a helyzet lassan javul az idén."
P_IDEA = "Ez a bekezdés rövid marad, a szerző álláspontja itt világos és egyértelmű."
P_NOTE = "Ehhez a mondathoz lábjegyzet tartozik a forrással együtt."


def _doc(path: Path) -> Path:
    from docx import Document
    from docx.oxml import parse_xml
    d = Document()
    d.add_heading("1. Előzmények", 1)
    for t in (P_CUT, P_NUM, P_TERM, P_IDEA):
        d.add_paragraph(t)
    note = d.add_paragraph(P_NOTE)
    note._p.append(parse_xml('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                             '<w:footnoteReference w:id="1"/></w:r>'))
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Táblázat szövege, amit nem szabad átírni", "2025"
    d.save(path)
    return path


PROPOSALS = {
    P_CUT: {"new": "A munkaerőhiány 2025-ben 4,5 százalékos volt, ezért a vállalatok többsége bért emelt.", "why": "ismétlés ki",
            "kind": "cut"},
    P_NUM: {"new": "A KSH adatai szerint sok betöltetlen állás volt az év végén.", "why": "egyszerűbb", "kind": "clarity"},
    P_TERM: {"new": "A jegybank szerint a helyzet lassan javul az idén.", "why": "rövidebb", "kind": "cut"},
    P_IDEA: {"new": "Ez a bekezdés rövid marad. Érdemes lenne egy nemzetközi összehasonlítást is beemelni a skandináv "
                    "modellekről és azok tanulságairól.", "why": "kitekintés", "kind": "new_idea"},
}


@pytest.fixture()
def ai(monkeypatch):
    from app.llm import siliconflow
    seen: list[str] = []

    async def fake(model, messages, thinking=True, json_mode=True, max_tokens=4096):
        user = messages[1]["content"]
        seen.append(messages[0]["content"] + "\n" + user)
        out = []
        for n, text in re.findall(r"\[(\d+)\] (.+)", user):
            if text in PROPOSALS:
                out.append({"n": int(n), **PROPOSALS[text]})
        return json.dumps({"changes": out}, ensure_ascii=False)
    monkeypatch.setattr(siliconflow, "_call_model", fake)
    return seen


async def test_first_the_questions_and_nothing_changes(tmp_path, put_file):
    from app.tools import edit_modes
    fid = put_file(_doc(tmp_path / "t.docx"))
    b = await edit_modes.brief(fid, "rövidítés")
    ids = [q["id"] for q in b["questions"]]
    assert b["mode"] == "shorten" and ids[:2] == ["goal", "limit"] and "new_ideas" in ids and "scope" in ids
    limit = next(q for q in b["questions"] if q["id"] == "limit")
    assert limit["default"] == 10 and "KSH" in next(q for q in b["questions"] if q["id"] == "keep_terms")["default"]
    assert b["document"]["tables"] == 1 and b["document"]["untouchable"] == 1 and "NOTHING CHANGED" in b["how"]
    with pytest.raises(ValueError, match="proofread"):
        edit_modes.mode_of("varázslat")


async def test_proposals_are_tracked_changes_within_the_limits(ai, tmp_path, put_file):
    from app.storage import output_path
    from app.tools import edit_modes
    from app.tools import review_loop as L
    fid = put_file(_doc(tmp_path / "t.docx"))
    res = await edit_modes.propose(fid, "shorten", {"limit": 20, "keep_terms": ["Magyar Nemzeti Bank"]})
    assert "Táblázat szövege" not in " ".join(ai) and P_NOTE not in " ".join(ai)      # a tábla és a jegyzetes bekezdés nem megy ki
    assert "Cut about 16%" in ai[0] and "NOT allowed" in ai[0]
    held = {h["paragraph"]: h["why"] for h in res["preservation"]["held_back"]}
    assert any("120" in w for w in held.values())                                     # a szám nem veszhet el
    assert any("Magyar Nemzeti Bank" in w for w in held.values())                     # a megnevezett szó marad
    assert any("new thoughts weren't allowed" in w for w in held.values())
    assert res["changes"] == 1 and res["preservation"]["not_touched"][0]["why"].startswith("it has a footnote")
    out = output_path(res["file_id"])
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "<w:del " in xml and 'w:author="Studio (shorten)"' in xml                   # követett változás, nem elfogadva
    texts = L._scan(L._body(L._Pkg(out)))["texts"]
    assert PROPOSALS[P_CUT]["new"] in texts and P_NUM in texts                         # végleges nézet = a javaslat
    com = zipfile.ZipFile(out).read("word/comments.xml").decode("utf-8")
    assert "Rövidítés: ismétlés ki" in com
    assert "NOTHING IS ACCEPTED" in res["how"] and res["check"]["status"] == "needs_check"


async def test_the_length_limit_holds_and_new_ideas_are_marked_when_allowed(ai, tmp_path, put_file):
    from app.storage import output_path
    from app.tools import edit_modes
    fid = put_file(_doc(tmp_path / "t.docx"))
    tight = await edit_modes.propose(fid, "shorten", {"limit": 2, "keep_numbers": False, "keep_terms": []})
    assert any("past the agreed 2%" in h["why"] for h in tight["preservation"]["held_back"]) or not tight["changed"]
    free = await edit_modes.propose(fid, "expand", {"limit": 50, "new_ideas": True, "keep_terms": []})
    com = zipfile.ZipFile(output_path(free["file_id"])).read("word/comments.xml").decode("utf-8")
    assert "Új gondolat — vállalod? kitekintés" in com and free["preservation"]["new_ideas"] == 1


async def test_proofreading_does_not_reword(ai, tmp_path, put_file, monkeypatch):
    from app.tools import edit_modes
    fid = put_file(_doc(tmp_path / "t.docx"))
    res = await edit_modes.propose(fid, "proofread", {"keep_terms": []})
    reasons = " ".join(h["why"] for h in res["preservation"]["held_back"]) if res.get("preservation") else \
        " ".join(h["why"] for h in res["held_back"])
    assert "rewords more than a proofreading should" in reasons


async def test_the_mcp_tool_asks_then_proposes(ai, tmp_path, put_file, monkeypatch):
    from app import mcp_server
    from app.mcp_server import edit_with_instruction
    monkeypatch.setattr(mcp_server, "_pool", lambda: None)
    fid = put_file(_doc(tmp_path / "t.docx"))
    asked = await edit_with_instruction(file_id=fid, mode="shorten")
    assert asked["questions"] and "NOTHING CHANGED" in asked["how"]
    done = await edit_with_instruction(file_id=fid, mode="shorten", brief={"limit": 20})
    assert done["changed"] and done["file_id"] != fid and "NOTHING IS ACCEPTED" in done["how"]
