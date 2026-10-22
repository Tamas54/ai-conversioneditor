"""A 10-22-i szerkesztői próbaüzem hibái (robot-tanulmány, tényellenőrzés + követett átírás MCP-n, WordCraft-tal):

 1. link végén gépelt szöveg a linkhez tapadt (100 bekezdés aikaveszunet.hu-ra mutatott) — a validátor nem szólt
 2. address new_comment megjegyzés-rész nélküli fájlon: 'NoneType' object has no attribute 'iter'
 3. review.compare {"text"} eldobta a lábjegyzetet, a szerző „Compare” lett
 4. address final: a lábjegyzeten belüli változás bent maradt, az árva jegyzet miatt a LibreOffice rossz szöveget mutatott
 5. elfogadás után a törölt bekezdés utáni címsor elvesztette a stílusát
 6. a szószám a követett beszúrást nem látta
 7. validate: rövid követett cím álhibás „missing”
 8. facts: gyenge forrás „supported”, PDF / 403 → „unverified”, rossz országhoz kötött szám
 9. csak olvasó parancs új verziót mentett
10. select.text csak a törzsben keresett
+1. „Mindent elfogad + megjegyzések törlése” gomb
"""
from __future__ import annotations

import asyncio
import re
import shutil
import zipfile
from pathlib import Path

import pytest

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
D = 'w:author="Ed" w:date="2026-10-22T10:00:00Z"'


def _rewrite(path: Path, fn=None, extra: dict | None = None, ct: str = "", rels: str = "") -> Path:
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    if fn:
        parts["word/document.xml"] = fn(parts["word/document.xml"].decode()).encode()
    for k, v in (extra or {}).items():
        parts[k] = v.encode()
    if ct:
        parts["[Content_Types].xml"] = parts["[Content_Types].xml"].decode().replace("</Types>", ct + "</Types>").encode()
    if rels:
        r = "word/_rels/document.xml.rels"
        parts[r] = parts[r].decode().replace("</Relationships>", rels + "</Relationships>").encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    return path


def _ref(i: int) -> str:
    return f'<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteReference w:id="{i}"/></w:r>'


def _note(i: int, body: str) -> str:
    return (f'<w:footnote w:id="{i}"><w:p><w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteRef/></w:r>'
            f'{body}</w:p></w:footnote>')


def _notes(*notes: str) -> str:
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes {W}>'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
            + "".join(notes) + "</w:footnotes>")


_FN_CT = ('<Override PartName="/word/footnotes.xml" '
          'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>')
_FN_REL = ('<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
           'Target="footnotes.xml"/>')


def tracked_doc(path: Path) -> Path:
    """Bevezető · IFR-mondat [1] · egészében törölt bekezdés [2] (a jegyzetjelével, a bekezdésjelével) · „4 Magyarország”
    címsor · záró. Az 1. jegyzetben követett csere (2023 → 2025)."""
    from docx import Document
    d = Document()
    d.add_paragraph("Bevezető bekezdés.")
    d.add_paragraph("Az IFR szerint a sűrűség nőtt.")
    d.add_paragraph("Ez a régi MI-rendeletes mondat törlendő.")
    d.add_heading("4 Magyarország helyzete", 1)
    d.add_paragraph("Záró bekezdés.")
    d.save(path)

    def doc(x: str) -> str:
        x = x.replace("<w:t>Az IFR szerint a sűrűség nőtt.</w:t></w:r>", "<w:t>Az IFR szerint a sűrűség nőtt.</w:t></w:r>" + _ref(1))
        return re.sub(r"<w:p>((?:(?!</w:p>).)*?)<w:r><w:t>Ez a régi MI-rendeletes mondat törlendő.</w:t></w:r></w:p>",
                      f'<w:p><w:pPr><w:rPr><w:del w:id="91" {D}/></w:rPr></w:pPr><w:del w:id="92" {D}><w:r>'
                      '<w:delText>Ez a régi MI-rendeletes mondat törlendő.</w:delText></w:r>' + _ref(2) + "</w:del></w:p>",
                      x, flags=re.S)
    notes = _notes(_note(1, f'<w:r><w:t xml:space="preserve"> IFR World Robotics </w:t></w:r><w:del w:id="93" {D}><w:r>'
                            f'<w:delText>2023</w:delText></w:r></w:del><w:ins w:id="94" {D}><w:r><w:t>2025</w:t></w:r></w:ins>'),
                   _note(2, '<w:r><w:t xml:space="preserve"> Régi MI-rendelet jegyzet.</w:t></w:r>'))
    return _rewrite(path, doc, {"word/footnotes.xml": notes}, _FN_CT, _FN_REL)


def _paras(path: Path) -> list[tuple[str, str, list[str]]]:
    """(stílus, szöveg a törölttel együtt, jegyzet-hivatkozások) bekezdésenként."""
    x = zipfile.ZipFile(path).read("word/document.xml").decode()
    out = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", x, re.S):
        st = re.search(r'<w:pStyle w:val="([^"]+)"', p)
        out.append((st.group(1) if st else "", "".join(re.findall(r"<w:(?:t|delText)[^>]*>([^<]*)<", p)),
                    re.findall(r'footnoteReference w:id="(\d+)"', p)))
    return out


def _note_texts(path: Path) -> dict[str, str]:
    f = zipfile.ZipFile(path).read("word/footnotes.xml").decode()
    return {m.group(1): m.group(2) for m in re.finditer(r'<w:footnote w:id="(\d+)">(.*?)</w:footnote>', f, re.S)}


# ---------- 4 + 5: tiszta példány / elfogadás ----------

def test_clean_copy_accepts_in_notes_keeps_the_heading_and_drops_orphan_notes(tmp_path):
    from app.tools import review_loop as L, review_ops
    pkg = L._Pkg(tracked_doc(tmp_path / "t.docx"))
    r = review_ops.final(pkg, L._comments(pkg))
    pkg.save(tmp_path / "f.docx")
    assert r["changes_accepted"] == 4 and r["in_notes_headers_footers"] == 2 and r["notes_removed"] == 1
    paras = _paras(tmp_path / "f.docx")
    assert [t for _, t, _ in paras] == ["Bevezető bekezdés.", "Az IFR szerint a sűrűség nőtt.", "4 Magyarország helyzete",
                                        "Záró bekezdés."]
    assert paras[2][0] == "Heading1"                                  # a címsor a SAJÁT stílusával marad
    notes = _note_texts(tmp_path / "f.docx")
    assert list(notes) == ["1"] and "<w:ins" not in notes["1"] and "<w:del" not in notes["1"] and "2025" in notes["1"]


def test_reject_restores_and_a_merge_keeps_the_first_paragraphs_form(tmp_path):
    from app.tools import review_loop as L, review_ops
    pkg = L._Pkg(tracked_doc(tmp_path / "t.docx"))
    review_ops.decide(pkg, "all", False)
    pkg.save(tmp_path / "r.docx")
    paras = _paras(tmp_path / "r.docx")
    assert [t for _, t, _ in paras][2] == "Ez a régi MI-rendeletes mondat törlendő." and paras[2][2] == ["2"]
    assert set(_note_texts(tmp_path / "r.docx")) == {"1", "2"} and "2023" in _note_texts(tmp_path / "r.docx")["1"]


def test_notes_are_put_in_reference_order(tmp_path):
    """A LibreOffice a jegyzetet a fájlbeli SORRENDJE szerint rendeli a hivatkozáshoz (mérve 10-22): rendezés kell."""
    from docx import Document
    from app.tools import review_loop as L, review_ops
    d = Document()
    d.add_paragraph("Első.")
    d.add_paragraph("Második.")
    d.save(tmp_path / "o.docx")
    _rewrite(tmp_path / "o.docx", lambda x: x.replace("<w:t>Első.</w:t></w:r>", "<w:t>Első.</w:t></w:r>" + _ref(7))
             .replace("<w:t>Második.</w:t></w:r>", "<w:t>Második.</w:t></w:r>" + _ref(3)),
             {"word/footnotes.xml": _notes(_note(3, "<w:r><w:t>harmadik</w:t></w:r>"), _note(5, "<w:r><w:t>árva</w:t></w:r>"),
                                           _note(7, "<w:r><w:t>hetedik</w:t></w:r>"))}, _FN_CT, _FN_REL)
    pkg = L._Pkg(tmp_path / "o.docx")
    assert review_ops.notes_report(pkg) == {"footnote": {"orphans": ["5"], "out_of_order": True}}
    assert review_ops.tidy_notes(pkg) == {"removed": 1, "reordered": 1}
    pkg.save(tmp_path / "o2.docx")
    assert list(_note_texts(tmp_path / "o2.docx")) == ["7", "3"]
    assert review_ops.notes_report(L._Pkg(tmp_path / "o2.docx")) == {}


# ---------- 2: új megjegyzés megjegyzés-rész nélkül ----------

async def test_new_comment_on_a_file_without_comments(tmp_path, put_file):
    from docx import Document
    from app.tools import review_loop
    d = Document()
    d.add_paragraph("A robotsűrűség Magyarországon nőtt.")
    d.save(tmp_path / "n.docx")
    fid = put_file(tmp_path / "n.docx")
    res = await review_loop.address(fid, [{"new_comment": "Forrás?", "paragraph": 1}], "Pupák")
    assert res["file_id"] and res["results"][0]["status"] == "commented", res
    from app.storage import output_path
    assert "Forrás?" in zipfile.ZipFile(output_path(res["file_id"])).read("word/comments.xml").decode()


# ---------- 3: összevetés a Studio saját gépezetével ----------

def test_text_compare_keeps_footnotes_and_the_author(tmp_path):
    from app.tools import review_loop as L, review_ops, text_compare
    src = tmp_path / "s.docx"
    pkg = L._Pkg(tracked_doc(tmp_path / "t.docx"))
    review_ops.final(pkg, {})
    pkg.save(src)                                                          # Bevezető · IFR [1] · cím · záró
    new = ("Elé írt sor.\nBevezető bekezdés, kicsit hosszabban.\nAz IFR szerint a robotsűrűség gyorsan nőtt.\n"
           "4 Magyarország helyzete\nEgy új gondolat a cím alatt.\nZáró bekezdés.\nUtóirat.")
    r = text_compare.compare(src, tmp_path / "c.docx", new, "Pupák")
    assert r["author"] == "Pupák" and r["inserted"] == 3 and r["changed"] == 2 and not r["not_done"]
    x = zipfile.ZipFile(tmp_path / "c.docx").read("word/document.xml").decode()
    assert set(re.findall(r'w:author="([^"]+)"', x)) == {"Pupák"}
    ifr = next(p for p in _paras(tmp_path / "c.docx") if "IFR" in p[1])
    assert ifr[2] == ["1"] and "1" in _note_texts(tmp_path / "c.docx")       # a jegyzet a helyén
    for accept, want in ((True, new.split("\n")), (False, ["Bevezető bekezdés.", "Az IFR szerint a sűrűség nőtt.",
                                                           "4 Magyarország helyzete", "Záró bekezdés."])):
        pkg = L._Pkg(tmp_path / "c.docx")
        review_ops.decide(pkg, "all", accept)
        pkg.save(tmp_path / "d.docx")
        assert [t for _, t, _ in _paras(tmp_path / "d.docx")] == want, accept
        styles = {t: s for s, t, _ in _paras(tmp_path / "d.docx")}
        assert styles["4 Magyarország helyzete"] == "Heading1"
        if accept:
            assert styles["Egy új gondolat a cím alatt."] != "Heading1"     # cím után szövegbekezdés formája


def test_text_compare_deletes_a_paragraph_with_its_note_and_rewrites_in_place(tmp_path):
    from app.tools import review_loop as L, review_ops, text_compare
    src = tmp_path / "s.docx"
    pkg = L._Pkg(tracked_doc(tmp_path / "t.docx"))
    review_ops.final(pkg, {})
    pkg.save(src)
    text_compare.compare(src, tmp_path / "del.docx", "Bevezető bekezdés.\n4 Magyarország helyzete\nZáró bekezdés.", "P")
    pkg = L._Pkg(tmp_path / "del.docx")
    assert review_ops.final(pkg, {})["notes_removed"] == 1                  # a törölt bekezdés jegyzete is megy
    text_compare.compare(src, tmp_path / "re.docx", "Bevezető bekezdés.\nTeljesen más szöveg, semmi köze.\n"
                         "4 Magyarország helyzete\nZáró bekezdés.", "P")
    assert next(p for p in _paras(tmp_path / "re.docx") if "Teljesen" in p[1])[2] == ["1"]   # helyben átírva: marad


# ---------- 6: szószám a végleges nézetből ----------

def test_word_count_sees_tracked_insertions(tmp_path):
    from app.tools import review_loop as L, text_compare, wordcraft
    from docx import Document
    d = Document()
    d.add_paragraph(" ".join(["régi"] * 40))
    d.save(tmp_path / "w.docx")
    text_compare.compare(tmp_path / "w.docx", tmp_path / "w2.docx", " ".join(["új"] * 50), "P")
    s = wordcraft._doc_summary(tmp_path / "w2.docx")
    assert s["words"] == 50 and s["words_before_tracked_changes"] == 40 and s["tracked_changes"] >= 2
    final, before = L.view_texts(tmp_path / "w2.docx")
    assert final[0].split() == ["új"] * 50


def test_an_edit_inside_a_tracked_insertion_is_a_change(tmp_path):
    from app.tools import content_digest
    from docx import Document
    d = Document()
    d.add_paragraph("Első.")
    d.add_paragraph("Záró bekezdés.")
    d.save(tmp_path / "a.docx")
    ins = lambda t: lambda x: x.replace("<w:r><w:t>Záró bekezdés.</w:t></w:r>",                 # noqa: E731
                                        f'<w:ins w:id="5" {D}><w:r><w:t>{t}</w:t></w:r></w:ins>')
    shutil.copy(tmp_path / "a.docx", tmp_path / "b.docx")
    _rewrite(tmp_path / "a.docx", ins("Záró bekezdés."))
    _rewrite(tmp_path / "b.docx", ins("Záró bekezdés, javítva."))
    assert not content_digest.same_content(tmp_path / "a.docx", tmp_path / "b.docx")


def test_digest_colour_tolerance():
    from app.tools.content_digest import _coarse
    assert _coarse("365F91") == _coarse("376092")                          # téma-szín két motor feloldásában
    assert _coarse("365F91") != _coarse("C00000") and _coarse("auto") == "auto"


# ---------- 1 + 7: validátor ----------

def _linked(path: Path, n: int) -> Path:
    from docx import Document
    d = Document()
    d.add_paragraph("Forrás: aikaveszunet.hu")
    for i in range(n):
        d.add_paragraph(f"A {i + 1}. bekezdés hosszú szövege a robotikáról, a munkaerőpiacról és a felzárkózásról szól.")
    d.add_paragraph("Valódi forrás: https://www.oecd.org/report")
    d.save(path)
    rid = 'r:id="rIdL"'

    def wrap(x: str) -> str:
        # az első bekezdés linkje jogos (a webhely nevét mondja), utána minden bekezdés „belegépelve” a linkbe
        x = re.sub(r"(<w:p>(?:(?!</w:p>).)*?)(<w:r>(?:(?!</w:p>).)*?</w:r>)(</w:p>)",
                   lambda m: m.group(1) + f"<w:hyperlink {rid}>" + m.group(2) + "</w:hyperlink>" + m.group(3)
                   if "oecd" not in m.group(2) else m.group(0), x, flags=re.S)
        return x
    return _rewrite(path, wrap, rels='<Relationship Id="rIdL" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                                     'relationships/hyperlink" Target="https://aikaveszunet.hu/hirek/" TargetMode="External"/>')


def test_validator_finds_text_that_spilled_into_a_link_and_the_fix_takes_it_off(tmp_path):
    from app.tools import layout_fix
    from app.tools import validate as v
    src = _linked(tmp_path / "l.docx", 4)
    issues, dc = v._analyze_doc(src, [])
    spill = [i for i in issues if i["kind"] == "link_spill"]
    assert len(spill) == 1 and spill[0]["severity"] == v.ERROR and "4 paragraphs" in spill[0]["message"]
    assert "aikaveszunet.hu" in spill[0]["message"]
    plan = [p for p in layout_fix._plan(issues, {b.pos: b for b in dc.blocks}) if p["kind"] == "link_spill"]
    done = layout_fix._apply(src, tmp_path / "l2.docx", plan)
    assert "4 paragraph(s)" in done[0]["how"]
    x = zipfile.ZipFile(tmp_path / "l2.docx").read("word/document.xml").decode()
    assert x.count("<w:hyperlink") == 1 and "Forrás: aikaveszunet.hu" in re.search(r"<w:hyperlink.*?</w:hyperlink>", x, re.S).group(0)
    assert not [i for i in v._analyze_doc(tmp_path / "l2.docx", [])[0] if i["kind"] == "link_spill"]
    few = _linked(tmp_path / "f.docx", 1)                                   # egy rövid bekezdés: nem riaszt
    assert not [i for i in v._analyze_doc(few, [])[0] if i["kind"] == "link_spill"]


def test_validator_flags_orphan_footnotes_and_proposes_the_tidy(tmp_path):
    from app.tools import layout_fix
    from app.tools import validate as v
    from app.tools import review_loop as L
    src = tracked_doc(tmp_path / "t.docx")
    pkg = L._Pkg(src)
    doc = pkg.xml("word/document.xml")                                      # a törölt bekezdés eltűnik, a jegyzete marad
    for p in list(doc.iter(L.P)):
        if "MI-rendeletes" in "".join(t.text or "" for t in p.iter(L.DT)):
            p.getparent().remove(p)
    pkg.put("word/document.xml", doc)
    pkg.save(tmp_path / "o.docx")
    issues, dc = v._analyze_doc(tmp_path / "o.docx", [])
    orphan = [i for i in issues if i["kind"] == "notes_orphan"]
    assert orphan and orphan[0]["severity"] == v.ERROR and "LibreOffice" in orphan[0]["message"]
    plan = [p for p in layout_fix._plan(issues, {b.pos: b for b in dc.blocks}) if p["kind"] == "notes_orphan"]
    done = layout_fix._apply(tmp_path / "o.docx", tmp_path / "o2.docx", plan)
    assert done and "1 note(s)" in done[-1]["how"] and list(_note_texts(tmp_path / "o2.docx")) == ["1"]


def test_short_tracked_title_is_found_in_a_layout_that_shows_the_deletion():
    """A lap a törölt szöveget is mutatja, egybeírva a beszúrással („A rossz forgatókönyvA magyar…”): a jelölt nézet
    alakja is illeszthető, nincs álhibás „missing”."""
    from app.tools.validate import Tok, _toks, align
    page = [Tok(t, 1, (0, 0, 1, 1)) for t in _toks("Előző bekezdés vége. A rossz forgatókönyvA magyar modernizáció "
                                                    "régi mondata törölve. Következő bekezdés szövege.")]
    final, markup = _toks("A rossz forgatókönyv"), _toks("A rossz forgatókönyv" + "A magyar modernizáció régi mondata törölve.")
    assert ("t",) not in align([(("t",), final)], page)
    assert ("t",) in align([(("t",), final, markup)], page)


# ---------- 8: tényellenőrzés ----------

def test_source_tiers_and_quotes():
    from app.tools import source_quality as sq
    assert sq.tier("https://ifr.org/downloads/press.pdf", "IFR") == "named"
    assert sq.tier("https://ifr.org/x", "International Federation of Robotics") == "named"
    assert sq.tier("https://www.pwc.com/hu/en/x.pdf", "PwC Magyarország") == "named"
    assert sq.tier("https://www.oecd.org/en/publications/x", None) == "official"
    assert sq.tier("https://digital-strategy.ec.europa.eu/en", None) == "official"
    assert sq.tier("https://math.stackexchange.com/q/1", "IFR") == "weak"
    assert sq.tier("https://gcn.com/robots", "IFR") == "secondary"
    assert sq.key_numbers("922,2 ezer munkahely 2025-ben") == ["922,2", "922.2"]
    txt = "Table 3. Exposure to generative AI: Hungary 27.1%, Korea 31.4%, OECD average 26.3%."
    assert sq.passages(txt, "a magyar kitettség 27,1%")[0].startswith("Table 3")
    assert sq.quoted("Hungary 27.1%", txt) and not sq.quoted("Korea 27.1%", txt)


@pytest.fixture()
def facts_world(monkeypatch):
    from app.tools import content_check as cc, datasources as ds
    pages = {"https://www.pwc.com/hu/report.pdf": "…Total employment exposed: 922,2 thousand jobs in Hungary (2025)…",
             "https://www.oecd.org/eo-2025.pdf": "Exposure to generative AI. Hungary 27.1%; Korea 31.4%."}
    results = {
        "pwc": [{"title": "PwC report", "url": "https://www.pwc.com/hu/report.pdf", "snippet": "How will AI impact…", "date": None}],
        "oecd": [{"title": "Employment Outlook", "url": "https://www.oecd.org/eo-2025.pdf", "snippet": "Generative AI…", "date": None}],
        "ifr": [{"title": "Robots", "url": "https://gcn.com/robots", "snippet": "IFR: robot density 4,1 per 1000", "date": None},
                {"title": "Q", "url": "https://math.stackexchange.com/q/9", "snippet": "density 4,1", "date": None}],
    }

    async def search(q, limit=5):
        return next((v for k, v in results.items() if k in q.lower()), [])

    async def scrape(url, max_chars=0):
        if url not in pages:
            raise ds.SourceError("reading the page answered HTTP 403")
        return pages[url]
    monkeypatch.setattr(ds, "search_enabled", lambda: True)
    monkeypatch.setattr(ds, "web_search", search)
    monkeypatch.setattr(ds, "scrape", scrape)
    return cc


async def test_facts_read_the_pdf_and_weigh_the_sources(tmp_path, facts_world, monkeypatch):
    from docx import Document
    cc = facts_world
    d = Document()
    for t in ("A PwC szerint 922,2 ezer munkahely érintett.", "Az OECD szerint a magyar kitettség 27,1%.",
              "Az IFR szerint a robotsűrűség 4,1 ezrelék."):
        d.add_paragraph(t)
    d.save(tmp_path / "f.docx")
    claims = {"claims": [
        {"paragraph": 1, "quote": "922,2 ezer munkahely érintett", "claim": "PwC: 922,2 ezer", "query": "pwc ai jobs hungary",
         "source_named": "PwC"},
        {"paragraph": 2, "quote": "a magyar kitettség 27,1%", "claim": "OECD: HU 27,1%", "query": "oecd generative ai exposure",
         "source_named": "OECD"},
        {"paragraph": 3, "quote": "robotsűrűség 4,1 ezrelék", "claim": "IFR: 4,1", "query": "ifr robot density", "source_named": "IFR"}]}
    seen: list[str] = []

    async def fake_ask(prompt, thinking=False):
        if prompt.startswith("List the checkable"):
            return claims
        seen.append(prompt)
        ev = {cid: blk for cid, blk in re.findall(r"Claim (F\d): (.*?)(?=\n\nClaim F|\Z)", prompt, re.S)}
        sid = lambda cid, host: (re.search(r"(S\d+) \[page[^\]]*" + re.escape(host), ev[cid])          # noqa: E731
                                 or re.search(r"(S\d+) \[[^\]]*" + re.escape(host), ev[cid])).group(1)
        return {"verdicts": [
            {"id": "F1", "verdict": "supported", "source": sid("F1", "pwc.com"), "explanation": "PwC: 922,2 ezer.",
             "evidence_quote": "Total employment exposed: 922,2 thousand jobs in Hungary"},
            {"id": "F2", "verdict": "contradicted", "source": sid("F2", "oecd.org"), "explanation": "Korea 27,1%.",
             "claimed": "27,1%", "correct": "31,4%", "replacement": "a magyar kitettség 31,4%", "evidence_quote": "Korea 27.1%"},
            {"id": "F3", "verdict": "supported", "source": sid("F3", "gcn.com"), "explanation": "gcn.com: 4,1.",
             "evidence_quote": "IFR: robot density 4,1 per 1000"}]}
    monkeypatch.setattr(cc, "_ask", fake_ask)
    r = await cc.facts(tmp_path / "f.docx")
    by = {c["claim"]: c for c in r["claims"]}
    assert "read from the page" in seen[0] and "922,2 thousand" in seen[0]          # a PDF-ből olvasott részlet az ítéletnél
    assert "named source: pwc.com" in seen[0] and "weak: math.stackexchange.com" in seen[0]
    pwc = by["PwC: 922,2 ezer"]
    assert pwc["verdict"] == "supported" and pwc["source"]["type"] == "page" and pwc["source"]["tier"] == "named"
    oecd = by["OECD: HU 27,1%"]                                               # a „Korea 27.1%” idézet nincs a forrásban
    assert oecd["verdict"] == "unverified" and oecd["replacement"] is None and "nem áll" in oecd["explanation"]
    ifr = by["IFR: 4,1"]                                                       # csak másodlagos oldal ismétli
    assert ifr["verdict"] == "unverified" and ifr["weak_support"] == "https://gcn.com/robots"
    assert "IFR saját kiadványát" in ifr["explanation"]


# ---------- 9 + 10 + 1 + 3 a motorral ----------

def _wc_version() -> tuple:
    import subprocess
    from app.tools import wordcraft
    exe = wordcraft.cli_path()
    if not exe:
        return ()
    out = subprocess.run([exe, "--version"], capture_output=True, text=True, env={"WORDCRAFT_NO_PREFS": "1"}).stdout
    m = re.search(r"(\d+)\.(\d+)\.(\d+)", out)
    return tuple(int(x) for x in m.groups()) if m else ()


@pytest.mark.engines
@pytest.mark.skipif(_wc_version() < (0, 9, 0), reason="WordCraft CLI 0.9+ needed")
async def test_engine_reads_dont_save_notes_are_found_compare_keeps_the_note(tmp_path, put_file):
    from app.storage import output_path
    from app.tools import review_loop as L, review_ops, wordcraft
    pkg = L._Pkg(tracked_doc(tmp_path / "t.docx"))
    review_ops.final(pkg, {})
    pkg.save(tmp_path / "s.docx")
    fid = put_file(tmp_path / "s.docx")
    ro = await wordcraft.run(file_id=fid, commands=[
        {"command": "caret.set", "params": {"pos": {"story": "body", "path": [0], "off": 0}}},
        {"command": "select.paragraph", "params": {}}, {"command": "document.selection", "params": {}}])
    assert ro["file_id"] == fid and ro.get("read_only", True)                  # nincs új változat
    sel = await wordcraft.run(file_id=fid, author="Pupák", commands=[
        {"command": "review.trackChanges", "params": {"value": True}},
        {"command": "select.text", "params": {"text": "World Robotics"}},
        {"command": "text.insert", "params": {"text": "World Robotics Report"}}])
    assert "World Robotics Report" in _note_texts(output_path(sel["file_id"]))["1"]
    cmp = await wordcraft.run(file_id=fid, author="Pupák", commands=[{"command": "review.compare", "params": {
        "text": "Bevezető bekezdés.\nAz IFR szerint a robotsűrűség gyorsan nőtt.\n4 Magyarország helyzete\nZáró bekezdés."}}])
    out = output_path(cmp["file_id"])
    assert cmp["compare"]["changed"] == 1 and next(p for p in _paras(out) if "IFR" in p[1])[2] == ["1"]
    assert set(re.findall(r'w:author="([^"]+)"', zipfile.ZipFile(out).read("word/document.xml").decode())) == {"Pupák"}


@pytest.mark.engines
@pytest.mark.skipif(_wc_version() < (0, 9, 0), reason="WordCraft CLI 0.9+ needed")
async def test_engine_text_typed_after_a_link_is_not_linked_and_accept_all_keeps_the_heading(tmp_path, put_file):
    from app.storage import output_path
    from app.tools import wordcraft
    src = _linked(tmp_path / "l.docx", 0)
    fid = put_file(src)
    r = await wordcraft.run(file_id=fid, commands=[
        {"command": "edit.goto", "params": {"paragraph": 1}}, {"command": "caret.end", "params": {}},
        {"command": "text.insert", "params": {"text": " és utána gépelt szöveg"}}])
    x = zipfile.ZipFile(output_path(r["file_id"])).read("word/document.xml").decode()
    assert "utána gépelt" not in "".join(re.findall(r"<w:hyperlink.*?</w:hyperlink>", x, re.S))
    acc = await wordcraft.run(file_id=put_file(tracked_doc(tmp_path / "t.docx")), commands=[{"command": "review.acceptAll"}])
    assert next(p for p in _paras(output_path(acc["file_id"])) if p[1].startswith("4 Magyar"))[0] == "Heading1"


# ---------- +1: tiszta példány a webről ----------

def test_clean_copy_endpoint(client, tmp_path, put_file):
    from app.storage import output_path
    fid = put_file(tracked_doc(tmp_path / "t.docx"))
    r = client.post("/api/make/clean_copy", json={"file_id": fid})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["file_id"] != fid and "4 changes accepted" in body["note"] and "original stays" in body["note"]
    assert "<w:ins" not in zipfile.ZipFile(output_path(body["file_id"])).read("word/footnotes.xml").decode()
    again = client.post("/api/make/clean_copy", json={"file_id": body["file_id"]})
    assert again.status_code == 400 and "Nothing to accept" in again.json()["detail"]


def test_only_reads_covers_caret_and_selection():
    from app.tools.run_report import only_reads
    assert only_reads([("caret.set", {}), ("select.paragraph", {}), ("document.selection", {})])
    assert not only_reads([("caret.set", {}), ("text.insert", {"text": "x"})])


def test_a_save_in_progress_does_not_break_the_file_list(client, tmp_path, put_file):
    """A WordCraft 0.9 mentés közben ideiglenes fájlt tart a kimeneti mappában (.wordcraft-save-<pid>-0.tmp): ha épp
    akkor listáz valaki, a lista nem dőlhet el („Invalid file_id”)."""
    from app.storage import output_path
    from docx import Document
    d = Document()
    d.add_paragraph("x")
    d.save(tmp_path / "x.docx")
    fid = put_file(tmp_path / "x.docx")
    tmp = output_path("probe.docx").with_name(".wordcraft-save-101-0.tmp")
    tmp.parent.mkdir(parents=True, exist_ok=True)
    tmp.write_bytes(b"PK")
    try:
        r = client.get("/api/files")
        assert r.status_code == 200, r.text
        names = [f["file_id"] for k in ("inputs", "outputs") for f in r.json().get(k, [])]
        assert fid in names and not any(n.startswith(".") for n in names)
    finally:
        tmp.unlink(missing_ok=True)
