"""A 10-22-i v5-ös szerkesztés (robot-tanulmány: katonai drón + humanoid fejezet, 759 követett változás) tanulságai —
amit a Studión kívül, saját eszközzel kellett megcsinálni:

 1. feltöltés az ügynök gépéről base64 nélkül: files(action="upload_link") + curl, a meglévő fájl ÚJ VÁLTOZATAKÉNT
 2. ábraszámozás: 5 új ábra után a régieket és a „lásd a 6. ábrát” hivatkozást kézzel kellett átszámozni
 3. a javaslat épsége: mindent elutasítva pontosan a kiinduló változat jön vissza (kézzel: pandoc reject = v4)
 4. házi stílusszabály a FELHASZNÁLÓHOZ kötve (-hat/-het, „korábbi vitairatunk”)
"""
from __future__ import annotations

import re
import time
import zipfile
from pathlib import Path

import pytest

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
D = 'w:author="Ed" w:date="2026-10-22T10:00:00Z"'


def _doc(path: Path, paras: list[str], caption_style: bool = False) -> Path:
    from docx import Document
    d = Document()
    for t in paras:
        if caption_style and re.match(r"\d+\. ábra", t):
            d.add_paragraph(t, style="Caption")
        else:
            d.add_paragraph(t)
    d.save(path)
    return path


def _rewrite(path: Path, fn, out: Path | None = None, notes: str | None = None) -> Path:
    out = out or path
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["word/document.xml"] = fn(parts["word/document.xml"].decode()).encode()
    if notes is not None:
        parts["word/footnotes.xml"] = notes.encode()
        parts["[Content_Types].xml"] = parts["[Content_Types].xml"].decode().replace("</Types>", (
            '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.'
            'wordprocessingml.footnotes+xml"/></Types>')).encode()
        r = "word/_rels/document.xml.rels"
        parts[r] = parts[r].decode().replace("</Relationships>", (
            '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
            'footnotes" Target="footnotes.xml"/></Relationships>')).encode()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    return out


def _track(x: str, old: str, new: str) -> str:
    """A bekezdés egy futásában `old` → követett törlés + beszúrás `new`."""
    i = x.index(old)
    a = x.rfind("<w:r>", 0, i)
    b = x.index("</w:r>", i) + len("</w:r>")
    run = x[a:b]
    t = re.search(r"<w:t[^>]*>(.*?)</w:t>", run).group(1)
    pre, post = t.split(old, 1)
    keep = lambda s: f'<w:r><w:t xml:space="preserve">{s}</w:t></w:r>' if s else ""
    return (x[:a] + keep(pre) + f'<w:del w:id="{abs(hash(old)) % 9000 + 100}" {D}><w:r><w:delText xml:space="preserve">'
            f'{old}</w:delText></w:r></w:del><w:ins w:id="{abs(hash(new)) % 9000 + 100}" {D}><w:r><w:t xml:space="preserve">'
            f'{new}</w:t></w:r></w:ins>' + keep(post) + x[b:])


def _structural(path: Path):
    from app.tools import validate as v
    return v._analyze_doc(path, [])[0]


def _kinds(issues, *kinds):
    return [i for i in issues if i["kind"] in kinds]


# ---------- 2: ábraszámozás ----------

def test_new_figures_without_renumbering_the_old_ones_are_caught(tmp_path):
    src = _doc(tmp_path / "a.docx", [
        "Bevezető, lásd az 1. ábrát.", "1. ábra. Robotsűrűség", "2. ábra. Drónok", "3. ábra. Humanoid szállítások",
        "Közben szöveg.", "3. ábra. Az ipar aránya", "4. ábra. OECD régiók", "Ahogy a 4. ábrán látszik."])
    caps = _kinds(_structural(src), "caption_numbering")
    assert [c["severity"] for c in caps] == ["error", "error"]
    assert "should be “4. ábra”" in caps[0]["message"] and "should be “5. ábra”" in caps[1]["message"]
    assert "the same number is used twice" in caps[0]["message"]
    skip = _kinds(_structural(_doc(tmp_path / "s.docx", ["1. ábra. A", "3. ábra. B"])), "caption_numbering")
    assert len(skip) == 1 and skip[0]["severity"] == "warning" and "skipped" in skip[0]["message"]


def test_references_to_a_missing_figure_but_not_to_other_works_or_robot_names(tmp_path):
    src = _doc(tmp_path / "r.docx", [
        "A Figure 03 robot dolgozik Spartanburgban, a Figure 02 flottát leállították.",
        "1. ábra. Robotsűrűség. Forrás: IFR (2025), 13–14. ábra.",
        "2. ábra. Drónok", "Lásd a 2. ábrát és a 7. ábrát.", "1. táblázat. Modell", "Az 1. táblázat szerint …"])
    refs = _kinds(_structural(src), "caption_reference")
    assert len(refs) == 1 and "“7. ábrát”" in refs[0]["message"] and "2. ábrát" not in refs[0]["message"]


def _new_para(text: str, n: int) -> str:
    """Egészében követett beszúrású bekezdés (a bekezdésjele is beszúrt)."""
    return (f'<w:p><w:pPr><w:rPr><w:ins w:id="{n}" {D}/></w:rPr></w:pPr><w:ins w:id="{n + 1}" {D}><w:r>'
            f'<w:t>{text}</w:t></w:r></w:ins></w:p>')


def test_a_tracked_renumbering_flags_the_reference_left_behind(tmp_path):
    """A javaslat új 2. ábrát szúr be a régi elé (követett), a régi 2 → 3; a régi „lásd a 2. ábrát” bent maradt,
    az új bekezdés már a 3-ra mutat."""
    base = _doc(tmp_path / "b.docx", ["1. ábra. A", "2. ábra. Régi ábra", "Lásd a 2. ábrát."])

    def fn(x):
        x = _track(x, "2. ábra. Régi", "3. ábra. Régi")
        x = x.replace("<w:p><w:r><w:t>Lásd", _new_para("Az új hivatkozás: lásd a 3. ábrát.", 7001) + "<w:p><w:r><w:t>Lásd", 1)
        return x.replace("<w:t>1. ábra. A</w:t></w:r></w:p>",
                         "<w:t>1. ábra. A</w:t></w:r></w:p>" + _new_para("2. ábra. Új ábra", 7003), 1)
    issues = _structural(_rewrite(base, fn, tmp_path / "p.docx"))
    assert not _kinds(issues, "caption_numbering")
    refs = _kinds(issues, "caption_reference")
    assert len(refs) == 1 and "renumbered" in refs[0]["message"] and "“2. ábrát”" in refs[0]["message"]
    assert "now “3. ábra”" in refs[0]["message"]


def test_captions_are_not_typed_list_numbers_and_caption_style_needs_no_punctuation(tmp_path):
    src = _doc(tmp_path / "c.docx", ["1. Első lépés", "2. Második lépés", "5. ábra Robotsűrűség", "1. ábra Drónok"],
               caption_style=True)
    issues = _structural(src)
    assert not _kinds(issues, "typed_numbering")
    caps = _kinds(issues, "caption_numbering")
    assert len(caps) == 2 and "should be “1. ábra”" in caps[0]["message"]


# ---------- 3: a javaslat épsége ----------

def _base_and_proposal(tmp_path, untracked: bool = False, note_untracked: bool = False):
    note = lambda body: (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes {W}>'
                         '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                         '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r>'
                         f'</w:p></w:footnote><w:footnote w:id="1"><w:p><w:r><w:t>{body}</w:t></w:r></w:p></w:footnote>'
                         '</w:footnotes>')
    ref = '<w:r><w:footnoteReference w:id="1"/></w:r>'
    base = _doc(tmp_path / "base.docx", ["A robotsűrűség nőtt.", "Ez okozhat gondot.", "Záró bekezdés hosszabb szöveggel."])
    base = _rewrite(base, lambda x: x.replace("<w:t>A robotsűrűség nőtt.</w:t></w:r>", "<w:t>A robotsűrűség nőtt.</w:t></w:r>" + ref),
                    notes=note("IFR World Robotics 2025."))

    def fn(x):
        x = _track(x, "okozhat", "okoz")
        # beszúrt bekezdésjel: a záró bekezdés két részre vágva (elutasítva újra egy bekezdés)
        x = x.replace("<w:p><w:r><w:t>Záró bekezdés hosszabb szöveggel.</w:t></w:r></w:p>",
                      f'<w:p><w:pPr><w:rPr><w:ins w:id="8001" {D}/></w:rPr></w:pPr><w:r><w:t xml:space="preserve">'
                      'Záró bekezdés </w:t></w:r></w:p><w:p><w:r><w:t>hosszabb szöveggel.</w:t></w:r></w:p>', 1)
        if untracked:
            x = x.replace("A robotsűrűség nőtt.", "A robotsűrűség meredeken nőtt.", 1)
        return x
    prop = _rewrite(base, fn, tmp_path / "prop.docx",
                    notes=note("IFR World Robotics 2026." if note_untracked else "IFR World Robotics 2025."))
    return base, prop


def test_a_clean_proposal_rejects_back_to_its_base_exactly(tmp_path):
    from app.tools import proposal_check as pc
    base, prop = _base_and_proposal(tmp_path)
    res = pc.compare(prop, base)
    assert res["tracked"] > 0 and res["changed"] == [] and res["notes"] == []


def test_an_untracked_edit_and_an_untracked_note_edit_are_errors(tmp_path, put_file):
    from app.storage import write_meta
    from app.tools import proposal_check as pc
    base, prop = _base_and_proposal(tmp_path, untracked=True, note_untracked=True)
    b_id, p_id = put_file(base), put_file(prop)
    write_meta(p_id, source_file_id=b_id, operation="wordcraft")
    from app.tools.crafts import existing
    issues, line = pc.issues(p_id, existing(p_id))
    assert line is None and len(issues) == 2 and {i["severity"] for i in issues} == {"error"}
    assert "Paragraph 1 was changed WITHOUT a tracked change" in issues[0]["message"]
    assert "[∅ → meredeken]" in issues[0]["message"]
    assert "footnote" in issues[1]["message"] and "2025" in issues[1]["message"]


def test_the_base_is_found_back_along_the_chain_and_decisions_soften_it(tmp_path, put_file):
    from app.storage import write_meta
    from app.tools import proposal_check as pc
    from app.tools.crafts import existing
    base, prop = _base_and_proposal(tmp_path, untracked=True)
    b_id, mid, p_id = put_file(base), put_file(prop), put_file(prop)
    write_meta(mid, source_file_id=b_id, operation="wordcraft")
    write_meta(p_id, source_file_id=mid, operation="wordcraft")
    assert pc.find_base(p_id) == (b_id, False)
    issues, _ = pc.issues(p_id, existing(p_id))
    assert issues and issues[0]["severity"] == "error"
    clean = put_file(_base_and_proposal(tmp_path)[1])
    write_meta(clean, source_file_id=b_id, operation="wordcraft")
    issues, line = pc.issues(clean, existing(clean))
    assert issues == [] and line.startswith("Proposal integrity: rejecting all")
    # közben elfogadás (kevesebb függő változás): figyelmeztetés
    more = put_file(_rewrite(prop, lambda x: _track(x, "Záró", "Utolsó"), tmp_path / "more.docx"))
    write_meta(more, source_file_id=b_id, operation="wordcraft")
    write_meta(p_id, source_file_id=more, operation="wordcraft")
    assert pc.find_base(p_id) == (b_id, True)
    assert pc.issues(p_id, existing(p_id))[0][0]["severity"] == "warning"


def test_no_base_no_check_and_a_clean_file_is_not_a_proposal(tmp_path, put_file):
    from app.tools import proposal_check as pc
    from app.tools.crafts import existing
    base, prop = _base_and_proposal(tmp_path)
    p_id = put_file(prop)
    issues, line = pc.issues(p_id, existing(p_id))
    assert issues == [] and "not checked" in line
    b_id = put_file(base)
    assert pc.issues(b_id, existing(b_id)) == ([], None)


# ---------- 4: házi szabályok ----------

@pytest.mark.parametrize("word,yes", [
    ("okozhat", True), ("lehet", True), ("mondhatjuk", True), ("kérhetnénk", True), ("tehessen", True),
    ("csökkentheti", True), ("lehetett", True), ("hat", False), ("hatás", False), ("hatékony", False),
    ("kihat", False), ("visszahatott", False), ("tizenhat", False), ("huszonhat", False), ("fenntartható", False),
    ("lehetőség", False), ("lehetetlen", False), ("tehetség", False), ("határ", False)])
def test_hu_modal_detection(word, yes):
    from app.tools import house_rules as hr
    assert hr.modal(word) is yes


def test_hu_modal_skips_conditions_and_quotes_but_not_hedge_words():
    from app.tools import house_rules as hr
    rule = hr.normalize([{"preset": "hu_modal"}])[0]
    assert hr.hits("Ha a kormány nem lép, a cégek elveszíthetik a piacot.", rule) == []
    assert hr.hits("A jelentés szerint „a robot okozhat balesetet”.", rule) == []
    assert hr.hits("Ez valószínűleg csökkentheti a költséget.", rule) == ["valószínűleg", "csökkentheti"]
    assert hr.hits("Ha sikerül, talán így lesz.", rule) == ["talán"]


def test_rules_normalize_and_check(tmp_path):
    from app.tools import house_rules as hr
    rules = hr.normalize([{"preset": "hu_modal"}, {"prefer": "korábbi vitairatunk", "instead_of": ["a 2025-ös vitairat"],
                                                   "why": "saját anyag"}, {"note": "Hibát nem írunk le."}])
    assert [r["id"] for r in rules] == ["R1", "R2", "R3"]
    for bad in ([{"preset": "x"}], [{"avoid": []}], [{"foo": 1}], "nem lista", [{"avoid": ["x" * 100]}]):
        with pytest.raises(ValueError):
            hr.normalize(bad)
    from app.tools import validate as v
    src = _doc(tmp_path / "h.docx", ["A 2025-ös vitairat szerint ez okozhat gondot.", "Ez rendben van."])
    blocks, _ = v.read_docx(src)
    issues = hr.check(blocks, rules)
    assert len(issues) == 1 and "R1" in issues[0]["message"] and "R2" in issues[0]["message"]
    assert "“okozhat”" in issues[0]["message"] and "“A 2025-ös vitairat”" in issues[0]["message"]      # ahogy a szövegben áll
    issues = v._analyze_doc(src, [], rules=rules)[0]
    assert len(_kinds(issues, "house_rule")) == 1


def test_house_rules_are_saved_per_user(db_client):
    from app import tenancy
    from app.tools import house_rules as hr
    tok = tenancy.set_user("aaaa1111bbbb2222")
    try:
        saved = db_client.portal.call(hr.save, [{"preset": "hu_modal"}])
        assert saved[0]["id"] == "R1" and db_client.portal.call(hr.load)[0]["preset"] == "hu_modal"
    finally:
        tenancy.reset(tok)
    assert db_client.portal.call(hr.load, "cccc3333dddd4444") == []        # másik felhasználó: nincs szabálya
    assert db_client.portal.call(hr.load, "aaaa1111bbbb2222")[0]["preset"] == "hu_modal"
    db_client.portal.call(hr.save, [], "aaaa1111bbbb2222")
    assert db_client.portal.call(hr.load, "aaaa1111bbbb2222") == []


# ---------- 1: feltöltési link ----------

def _link(filename=None, version_of=None):
    import asyncio
    from app import mcp_server
    return asyncio.run(mcp_server._upload_link(filename, version_of))


def test_upload_link_puts_a_file_once_as_a_new_version(client, tmp_path, put_file):
    from app.storage import read_meta
    from app.tools.revise import history
    v4 = put_file(_doc(tmp_path / "v4.docx", ["Régi szöveg."]))
    link = _link("robotos3 – v5.docx", v4)
    assert link["method"] == "PUT" and link["version_of"] == v4 and 1700 < link["expires_in_s"] <= 1800
    path = "/" + link["url"].split("/", 3)[3]
    data = _doc(tmp_path / "v5.docx", ["Új szöveg."]).read_bytes()
    r = client.put(path, content=data)
    assert r.status_code == 201, r.text
    got = r.json()
    assert got["name"] == "robotos3 – v5.docx" and got["version_of"] == v4 and got["edit_url"].endswith("/edit")
    meta = read_meta(got["file_id"])
    assert meta["source_file_id"] == v4 and meta["operation"] == "upload version"
    assert [s["file_id"] for s in history(got["file_id"])][-2:] == [v4, got["file_id"]]
    assert client.put(path, content=data).status_code == 409                     # egyszer
    assert client.put(path[:-3] + "abc", content=data).status_code == 403        # hamis aláírás


def test_upload_link_multipart_name_and_failures(client, tmp_path, monkeypatch):
    from app import upload_links
    link = _link()
    path = "/" + link["url"].split("/", 3)[3]
    r = client.post(path + "?name=jegyzet.docx", files={"file": ("x.docx", _doc(tmp_path / "m.docx", ["Szia."]).read_bytes())})
    assert r.status_code == 201 and r.json()["name"] == "jegyzet.docx" and "version_of" not in r.json()
    again = _link("ures.docx")
    p2 = "/" + again["url"].split("/", 3)[3]
    assert client.put(p2, content=b"").status_code == 400
    assert client.put(p2, content=_doc(tmp_path / "ok.docx", ["Most jó."]).read_bytes()).status_code == 201   # a hiba után újra
    tok, _ = upload_links.make("admin", "x.docx", None, ttl_s=-1)
    assert client.put(f"/upload-link/{tok}", content=b"abc").status_code == 403        # lejárt


def test_upload_link_refuses_a_version_of_someone_elses_file():
    with pytest.raises(ValueError, match="not one of this person's files"):
        _link("x.docx", "nincs-ilyen.docx")
