"""3. csoport: űrlap-felismerés bármiből, kitöltés az eredeti formátumban, aláírás (titkosítva tárolva, PDF/DOCX
elhelyezés), ügynöki kitöltés profilból és iratokból (kérdésekkel, validálással) — weben, REST API-n, MCP-n."""
from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest

pytestmark = pytest.mark.asyncio

VALUES = {"Név": "Kőváry Ödön", "lakcim": "9021 Győr, Fő utca 1.", "Születési dátum": "1980.05.17", "Hozzájárulok": "igen",
          "Adóazonosító jel": "8123456786", "Dátum": "2026-10-15", "IBAN": "HU42 1177 3016 1111 1018 0000 0000",
          "Telefonszám": "06301234567"}


def _out(fid: str) -> Path:
    from app.storage import output_path
    return output_path(fid)


async def run(key, fids, **opts):
    from app.ops import run_tool
    return await run_tool(key, fids, opts, charge=False)


def _engine(name: str):
    from app.convert import engines
    if not engines.available(name):
        pytest.skip(f"{name} not installed")


# ---------------------------------------------------------------- validátorok

def test_validators():
    from app.formkit import validate as V
    assert V.check_iban("hu42 1177 3016 1111 1018 0000 0000") == (True, "HU42 1177 3016 1111 1018 0000 0000", None)
    assert not V.check_iban("HU43117730161111101800000000")[0]
    assert not V.check_iban("HU4211773016111110180000")[0]                    # rövid
    assert V.check_hu_tax_id("8123456786")[0] and not V.check_hu_tax_id("8123456789")[0]
    assert V.hu_tax_id_birth("8123456786").year == 1867 + 12345 // 365
    assert V.check_taj("123 456 788") == (True, "123 456 788", None) and not V.check_taj("123456789")[0]
    assert V.check_hu_vat("10773381-2-44")[0] and not V.check_hu_vat("10773382-2-44")[0]
    assert V.check_hu_bank("11773016-11111018")[0] and not V.check_hu_bank("11773016-11111019")[0]
    assert V.check_date("2026. október 5.")[1] == "2026-10-05" and V.check_date("15/10/2026", "hu")[1] == "2026. 10. 15."
    assert not V.check_date("2026-02-30")[0]
    assert V.check_phone("06 1 234 5678")[1] == "+36 1 234 5678" and V.check_email("ödön@példa.hu")[0]
    from app.formkit.profile import hu_iban
    assert hu_iban("11773016-11111018-00000000") == "HU42 1177 3016 1111 1018 0000 0000"


# ---------------------------------------------------------------- felismerés

async def test_recognize_every_format(samples):
    from app.formkit import detect
    s = await detect.recognize(samples.fillable_pdf())
    assert s["kind"] == "acroform"
    assert [(f["name"], f["type"]) for f in s["fields"]] == [
        ("nev", "text"), ("lakcim", "text"), ("szuletesi_datum", "date"), ("hozzajarul", "checkbox"), ("alairas", "signature")]
    assert s["fields"][0]["label"] == "Név" and s["fields"][3]["label"] == "Hozzájárulok az adatkezeléshez"
    s = await detect.recognize(samples.flat_form_pdf())
    assert s["kind"] == "flat_pdf"
    assert [(f["label"], f["type"]) for f in s["fields"]] == [
        ("Név", "text"), ("Lakcím", "text"), ("Adóazonosító jel", "hu_tax_id"), ("Dátum", "date"),
        ("Hozzájárulok", "checkbox"), ("Aláírás", "signature")]
    assert s["fields"][0]["box"][0] == 200 and abs(s["fields"][0]["box"][3] - 152) < 1
    s = await detect.recognize(samples.form_docx())
    assert [(f["label"], f["type"], f["anchor"]["kind"]) for f in s["fields"]] == [
        ("Név", "text", "docx_text"), ("Születési dátum", "date", "docx_text"), ("Adóazonosító jel", "hu_tax_id", "docx_text"),
        ("Hozzájárulok", "checkbox", "docx_text"), ("Aláírás", "signature", "docx_text"), ("Lakcím", "text", "docx_cell"),
        ("Telefonszám", "phone", "docx_cell")]
    s = await detect.recognize(samples.form_xlsx())
    assert [(f["label"], f["anchor"]["cell"]) for f in s["fields"]] == [("Név", "B1"), ("Lakcím", "B2"), ("Születési dátum", "B3"), ("IBAN", "B4")]
    assert s["fields"][3]["type"] == "iban"


def _photo_of(pdf: Path, tmp_path: Path) -> Path:
    import pymupdf
    from PIL import Image, ImageFilter
    png = tmp_path / "urlap.png"
    with pymupdf.open(str(pdf)) as d:
        d[0].get_pixmap(dpi=150).save(str(png))
    photo = tmp_path / "fotó.jpg"
    Image.open(png).convert("RGB").filter(ImageFilter.GaussianBlur(0.7)).save(photo, quality=85)
    return photo


async def test_recognize_photo_offline_first(samples, tmp_path):
    """Fotó/szken: alapból ingyenes, helyi felismerés (Tesseract + OpenCV) — a látó modell nem fut."""
    _engine("tesseract")
    from app.formkit import detect
    s = await detect.recognize(_photo_of(samples.flat_form_pdf(), tmp_path))
    assert s["kind"] == "image" and s["engine"] == "offline" and s["quality"]["level"] == "good"
    assert [f["type"] for f in s["fields"]] == ["text", "text", "hu_tax_id", "date", "checkbox", "signature"]
    assert "deep" not in s


def _vision_mock(monkeypatch, answers: list[dict], seen: list):
    from app.config import settings
    from app.llm import vision
    monkeypatch.setattr(settings, "SILICONFLOW_API_KEY", "teszt")

    async def fake_call(msgs, max_tokens=4096, model=None, **kw):
        seen.append((model, msgs[0]["content"][0]["text"][:40], msgs[0]["content"][1]["image_url"]["url"][:22]))
        return json.dumps(answers.pop(0))
    monkeypatch.setattr(vision, "_call", fake_call)


async def test_deep_recognition_multi_step_for_paid_plans(samples, tmp_path, monkeypatch):
    """Kérve (deep=on) vagy gyenge eredménynél: a helyi mezők számozva az oldalképen → a modell javít, töröl,
    pótol (DeepSeek V4.1 Flash = settings.OCR_MODEL); sok változásnál második ellenőrző kör."""
    from app.config import settings
    from app.formkit import detect
    seen: list = []
    _vision_mock(monkeypatch, [
        {"fix": [{"n": 1, "label": "Teljes név", "kind": "text"}], "remove": [5],
         "add": [{"label": "Megjegyzés", "kind": "multiline", "box": [120, 480, 850, 560]}]},
        {"fix": [], "remove": [], "add": []}], seen)
    s = await detect.recognize(samples.flat_form_pdf(), deep="on")
    assert s["engine"] == "vision" and s["deep"]["used"] and s["deep"]["pages"] == 1
    labels = [f["label"] for f in s["fields"]]
    assert labels[0] == "Teljes név" and "Hozzájárulok" not in labels and labels[-1] == "Megjegyzés"
    assert s["fields"][-1]["type"] == "multiline" and s["fields"][-1]["box"][0] == pytest.approx(595 * 0.12, abs=0.5)
    assert seen[0][0] == settings.OCR_MODEL and seen[0][2].startswith("data:image/")
    assert s["deep"]["steps"] == len(seen) == 2
    # gyenge helyi eredmény (nincs vonal, csak feliratok) → automatikusan indul, ha jár (itt: admin = Pro)
    import pymupdf
    bare = tmp_path / "csak_feliratok.pdf"
    with pymupdf.open() as d:
        pg = d.new_page()
        for k, lab in enumerate(("Név:", "Lakcím:", "Telefon:", "E-mail:")):
            pg.insert_text((72, 100 + 30 * k), lab, fontsize=11)
        d.save(str(bare))
    from app.tools import forms as legacy

    async def fake_layout(page, hint=None):
        return {"items": [{"type": "field", "kind": "text", "name": "nev", "label": "Név", "box": [200, 100, 800, 125]}]}
    monkeypatch.setattr(legacy, "read_layout", fake_layout)
    seen.clear()
    _vision_mock(monkeypatch, [{"fix": [], "remove": [], "add": []}], seen)
    s = await detect.recognize(bare)
    assert s["deep"]["used"] and [f["name"] for f in s["fields"]] == ["nev"]


async def test_deep_recognition_gated_mcp_and_free_plan(samples, monkeypatch):
    from app import billing, tenancy
    from app.formkit import deep as D, detect
    from app.config import settings
    monkeypatch.setattr(settings, "SILICONFLOW_API_KEY", "teszt")
    tok = tenancy.client_makes_images.set(True)                      # MCP-hívás: a kliens ügynöke dolgozik
    try:
        s = await detect.recognize(samples.flat_form_pdf(), deep="auto")
        assert "deep" not in s                                       # jó eredménynél nincs teendő
        with pytest.raises(D.DeepUnavailable, match="own vision"):
            await detect.recognize(samples.flat_form_pdf(), deep="on")
    finally:
        tenancy.client_makes_images.reset(tok)

    async def free(uid):
        return "free"
    monkeypatch.setattr(billing, "plan_of", free)
    with pytest.raises(D.DeepUnavailable, match="Plus and Pro"):
        await detect.recognize(samples.flat_form_pdf(), deep="on")
    sch = {"kind": "image", "pages": [{"width": 10, "height": 10}], "fields": [], "label_hints": 4}
    assert D.quality(sch)["level"] == "low"
    out = await detect._maybe_deep(sch, Path("x.png"), "hun+eng", "auto")
    assert out["deep"] == {"used": False, "reason": "plan", "hint": D.UPGRADE}


async def test_deep_recognition_charges_credits(monkeypatch):
    from app import db, tenancy
    from app.formkit import deep as D
    calls = []

    async def spend(uid, n, reason, ref=None):
        calls.append((uid, n, reason))
        return 10
    monkeypatch.setattr(db, "enabled", lambda: True)
    monkeypatch.setattr(db, "credits_spend", spend)
    tok = tenancy.set_user("fizetos01")
    try:
        assert await D.charge(3) == 3 * D.CREDITS_PER_PAGE
    finally:
        tenancy.reset(tok)
    assert calls == [("fizetos01", 3 * D.CREDITS_PER_PAGE, "AI form recognition")]

    async def broke(uid, n, reason, ref=None):
        raise db.InsufficientCredits("no")
    monkeypatch.setattr(db, "credits_spend", broke)
    tok = tenancy.set_user("fizetos01")
    try:
        with pytest.raises(D.DeepUnavailable, match="credits"):
            await D.charge(1)
    finally:
        tenancy.reset(tok)


# ---------------------------------------------------------------- kitöltés az eredeti formátumban

async def test_fill_acroform_editable_and_flattened(samples, put_file):
    import pymupdf
    fid = put_file(samples.fillable_pdf())
    r = await run("fill-form", [fid], values=VALUES)
    assert r["data"]["filled"] == ["nev", "lakcim", "szuletesi_datum", "hozzajarul"] and not r["data"]["errors"]
    with pymupdf.open(str(_out(r["file_id"]))) as d:
        vals = {w.field_name: w.field_value for w in d[0].widgets()}
        assert vals["nev"] == "Kőváry Ödön" and vals["szuletesi_datum"] == "1980. 05. 17." and vals["hozzajarul"] not in ("Off", "")
        d.bake()                                         # a megjelenés (amit a néző lát) is ékezethelyes
        t = d[0].get_text()
        assert "Kőváry Ödön" in t and "9021 Győr, Fő utca 1." in t
    r = await run("fill-form", [fid], values=VALUES, flatten=True)
    with pymupdf.open(str(_out(r["file_id"]))) as d:
        assert not list(d[0].widgets()) and "Kőváry Ödön" in d[0].get_text()


async def test_fill_flat_docx_xlsx(samples, put_file):
    import pymupdf
    from docx import Document
    from openpyxl import load_workbook
    r = await run("fill-form", [put_file(samples.flat_form_pdf())], values=VALUES)
    with pymupdf.open(str(_out(r["file_id"]))) as d:
        t = d[0].get_text()
        assert all(x in t for x in ("Kőváry Ödön", "8123456786", "2026. 10. 15.", "✓"))
    r = await run("fill-form", [put_file(samples.form_docx())], values=VALUES)
    assert r["file_id"].endswith(".docx")
    d = Document(str(_out(r["file_id"])))
    paras = [p.text for p in d.paragraphs]
    assert "Név: Kőváry Ödön" in paras and "Hozzájárulok: ☒" in paras and "Születési dátum: 1980. 05. 17." in paras
    assert [c.text for c in d.tables[0].rows[1].cells] == ["Telefonszám", "+36 30 123 4567"]
    r = await run("fill-form", [put_file(samples.form_xlsx())], values=VALUES)
    ws = load_workbook(str(_out(r["file_id"]))).active
    assert ws["B1"].value == "Kőváry Ödön" and ws["B3"].value.year == 1980 and ws["B4"].value.startswith("HU42 1177")


async def test_fill_photo_keeps_look_or_becomes_pdf(samples, put_file, tmp_path, monkeypatch):
    import pymupdf
    from PIL import Image, ImageChops
    from app.formkit import detect, schema as S
    png = tmp_path / "lap.png"
    with pymupdf.open(str(samples.flat_form_pdf())) as d:
        d[0].get_pixmap(dpi=100).save(str(png))
    sch = S.make("image", "image", [{"page": 1, "width": 827, "height": 1170}],
                 [S.field("nev", "Név", "text", page=1, box=[278, 185, 694, 211], anchor={"kind": "xy", "page": 1, "rect": [278, 185, 694, 211]})])

    async def fake(path, language="hun+eng", deep="auto"):
        return json.loads(json.dumps(sch))
    monkeypatch.setattr(detect, "recognize", fake)
    fid = put_file(png)
    r = await run("fill-form", [fid], values={"nev": "Ősz Űrsula"})
    assert r["file_id"].endswith(".png")
    with Image.open(_out(r["file_id"])) as a, Image.open(png) as b:
        diff = ImageChops.difference(a.convert("RGB"), b.convert("RGB")).getbbox()
        assert diff and diff[0] >= 270 and diff[2] <= 700                     # csak a mező dobozába írt
    r = await run("fill-form", [fid], values={"nev": "Ősz Űrsula"}, keep_look=False)
    with pymupdf.open(str(_out(r["file_id"]))) as d:
        assert "Ősz Űrsula" in d[0].get_text()


async def test_manual_schema_and_validation_errors(samples, put_file):
    import pymupdf
    fid = put_file(samples.flat_form_pdf())
    sch = (await run("form-recognize", [fid]))["data"]
    assert len(sch["fields"]) == 6
    sch["fields"][0]["label"] = "Teljes név"
    sch["fields"].append({"name": "megjegyzes", "label": "Megjegyzés", "type": "text", "page": 1, "box": [72, 420, 500, 440]})
    sch["fields"].append({"name": "lakcim", "remove": True})
    r = await run("fill-form", [fid], schema=sch, values={"Teljes név": "Kőváry Ödön", "megjegyzes": "Őszi űrlap kézi mezővel",
                                                         "lakcim": "nem kell", "Adóazonosító jel": "8123456789", "IBAN": "x"})
    data = r["data"]
    assert data["filled"] == ["nev", "megjegyzes"] and data["skipped"] == ["lakcim", "IBAN"]
    assert "check digit" in data["errors"]["adoazonosito_jel"]
    with pymupdf.open(str(_out(r["file_id"]))) as d:
        t = d[0].get_text()
        assert "Őszi űrlap kézi mezővel" in t and "nem kell" not in t and "8123456789" not in t
    from app.convert.options import OptionError
    with pytest.raises(OptionError, match="needs a box"):
        await run("fill-form", [fid], schema={"fields": [{"name": "uj", "label": "Új"}]}, values={"uj": "x"})


async def test_bulk_fill_and_make_fillable(samples, put_file):
    import pymupdf
    from docx import Document
    data = put_file(samples.csv_rows())
    r = await run("bulk-fill-form", [put_file(samples.fillable_pdf())], data_file=data, flatten=True)
    assert r["data"]["copies"] == 2
    with pymupdf.open(str(_out(r["file_id"]))) as d:
        assert d.page_count == 2 and "Kőváry Ödön" in d[0].get_text() and "Űrmös Éva" in d[1].get_text()
    r = await run("bulk-fill-form", [put_file(samples.form_docx())], data_file=data)
    with zipfile.ZipFile(_out(r["file_id"])) as z:
        names = sorted(z.namelist())
        assert len(names) == 2 and names[0].startswith("001-")
        import io
        assert "Név: Űrmös Éva" in [p.text for p in Document(io.BytesIO(z.read(names[1]))).paragraphs]
    r = await run("make-fillable", [put_file(samples.flat_form_pdf())])
    fid = put_file(_out(r["file_id"]))
    with pymupdf.open(str(_out(r["file_id"]))) as d:
        ws = {w.field_name: w.field_type for w in d[0].widgets()}
        assert set(ws) == {"nev", "lakcim", "adoazonosito_jel", "datum", "hozzajarulok", "alairas"}
        assert ws["hozzajarulok"] == pymupdf.PDF_WIDGET_TYPE_CHECKBOX
    r = await run("fill-form", [fid], values={"Név": "Kőváry Ödön", "hozzajarulok": True})
    with pymupdf.open(str(_out(r["file_id"]))) as d:
        assert {w.field_name: w.field_value for w in d[0].widgets()}["nev"] == "Kőváry Ödön"


async def test_doc_and_odt_roundtrip_to_original_format(samples, put_file, tmp_path):
    _engine("libreoffice")
    from app.convert.service import convert_path
    outs, _ = await convert_path(samples.form_docx(), "odt", {}, tmp_path / "w")
    r = await run("fill-form", [put_file(outs[0])], values={"Név": "Kőváry Ödön"})
    assert r["file_id"].endswith(".odt")
    back, _ = await convert_path(_out(r["file_id"]), "txt", {}, tmp_path / "t")
    assert "Név: Kőváry Ödön" in back[0].read_text("utf-8")


# ---------------------------------------------------------------- aláírás

async def test_signature_clean_store_encrypted_and_isolated(samples, put_file):
    import numpy as np
    from PIL import Image
    from app import tenancy
    from app.formkit import signature as G
    r = await run("create-signature", [put_file(samples.signature_photo())], name="Ödön aláírása")
    sid = r["data"]["signature_id"]
    with Image.open(_out(r["file_id"])) as im:
        a = np.asarray(im.convert("RGBA"))
        assert im.mode == "RGBA" and a[0, 0, 3] == 0 and a[..., 3].max() == 255
        assert (a[..., 3] > 0).mean() < 0.2                             # a papír eltűnt, csak a vonás maradt
    raw = b"".join(p.read_bytes() for p in G._fs_dir(tenancy.current_user_id()).glob(f"{sid}.bin"))
    assert raw and b"\x89PNG" not in raw and b"IHDR" not in raw         # titkosítva a lemezen
    assert (await G.load(tenancy.current_user_id(), sid))[:4] == b"\x89PNG"
    with pytest.raises(KeyError):
        await G.load("masikfelhasznalo1", sid)
    other = tenancy.set_user("masikfelhasznalo1")
    try:
        assert all(s["id"] != sid for s in await G.list_("masikfelhasznalo1"))
    finally:
        tenancy.reset(other)
    import base64
    from PIL import ImageDraw
    c = Image.new("RGBA", (400, 150), (0, 0, 0, 0))
    ImageDraw.Draw(c).line([(20, 120), (150, 30), (260, 130), (380, 40)], fill=(0, 0, 0, 255), width=5)
    import io
    buf = io.BytesIO()
    c.save(buf, "PNG")
    r2 = await run("create-signature", [], image_base64="data:image/png;base64," + base64.b64encode(buf.getvalue()).decode(),
                   kind="initials", color="#1a237e")
    assert r2["data"]["kind"] == "initials"
    assert await G.delete(tenancy.current_user_id(), r2["data"]["signature_id"])


async def test_sign_pdf_and_docx(samples, put_file):
    import pymupdf
    from docx import Document
    sid = (await run("create-signature", [put_file(samples.signature_photo())]))["data"]["signature_id"]
    r = await run("sign-document", [put_file(samples.flat_form_pdf())], signature_id=sid)
    assert r["data"]["placement"] == "field"
    x0, y0, x1, y1 = r["data"]["rect"]
    assert 195 <= x0 <= 210 and 370 <= y1 <= 390                       # az „Aláírás:” vonalán
    with pymupdf.open(str(_out(r["file_id"]))) as d:
        assert d[0].get_images()
    r = await run("sign-document", [put_file(samples.pdf())], signature_id=sid, placement="xy", page=2, x=100, y=600, width=120)
    assert r["data"]["page"] == 2 and r["data"]["rect"][0] == pytest.approx(102, abs=1)
    r = await run("sign-document", [put_file(samples.pdf())], signature_id=sid)
    assert r["data"]["placement"] == "bottom-right"
    r = await run("sign-document", [put_file(samples.form_docx())], signature_id=sid)
    d = Document(str(_out(r["file_id"])))
    sig_p = [p for p in d.paragraphs if p.text.startswith("Aláírás")][0]
    assert "_" not in sig_p.text and sig_p._p.xpath(".//pic:pic") and len(d.inline_shapes) == 1
    from app.convert.options import OptionError
    with pytest.raises(OptionError, match="no saved signature"):
        await run("sign-document", [put_file(samples.pdf())], signature_id="sig_nincsilyenalairas")


# ---------------------------------------------------------------- ügynöki kitöltés

async def test_profile_encrypted_and_plan(samples):
    from app.formkit import detect, profile as P
    uid = "profil-teszt"
    saved = await P.save_profile(uid, {"full_name": "Kőváry Ödön", "tax_id": "8123456786", "bank_account": "11773016-11111018"})
    assert (await P.load_profile(uid)) == saved
    assert b"K\xc5\x91v" not in P._fs(uid).read_bytes()
    with pytest.raises(ValueError, match="unknown profile field"):
        await P.save_profile(uid, {"jelszo": "x"})
    sch = await detect.recognize(samples.flat_form_pdf())
    pl = await P.plan(sch, profile=await P.load_profile(uid), use_ai=False)
    assert pl["values"]["nev"] == "Kőváry Ödön" and pl["provenance"]["nev"] == "profile:full_name"
    assert pl["provenance"]["datum"] == "today" and [q["name"] for q in pl["questions"]] == ["lakcim"]
    d = P.derive(saved)
    assert d["iban"] == "HU42 1177 3016 1111 1018 0000 0000" and d["family_name"] == "Kőváry" and d["birth_date"].startswith("1900")
    bad = await P.plan(sch, profile={"full_name": "X", "tax_id": "8123456789"}, use_ai=False)
    q = {x["name"]: x for x in bad["questions"]}
    assert "check digit" in q["adoazonosito_jel"]["error"]
    src = "Személyi adatlap\nLakcím: 1051 Budapest, Ősz utca 5.\nSzámlaszám IBAN: HU42 1177 3016 1111 1018 0000 0000\n"
    pl = await P.plan(sch, profile={"full_name": "X"}, sources=src, use_ai=False)
    assert pl["values"]["lakcim"] == "1051 Budapest, Ősz utca 5." and pl["provenance"]["lakcim"] == "document"
    assert await P.delete_profile(uid)


async def test_agent_fill_asks_then_completes(samples, put_file, tmp_path):
    import pymupdf
    src = tmp_path / "korabbi_urlap.txt"
    src.write_text("Név: Kőváry Ödön\nAdóazonosító jel: 8123456786\n", "utf-8")
    sid = (await run("create-signature", [put_file(samples.signature_photo())]))["data"]["signature_id"]
    form = put_file(samples.flat_form_pdf())
    r = await run("ai-form-filler", [form], source_files=[put_file(src)], signature_id=sid)
    d = r["data"]
    assert d["status"] == "needs_input" and [q["name"] for q in d["questions"]] == ["lakcim"]
    assert d["provenance"] == {"nev": "document", "adoazonosito_jel": "document", "datum": "today", "alairas": "signature"}
    assert r["files"][1]["label"].endswith(".png")
    r = await run("ai-form-filler", [form], source_files=[put_file(src)], signature_id=sid, answers={"lakcim": "Győr, Ős tér 2."})
    assert r["data"]["status"] == "complete"
    with pymupdf.open(str(_out(r["file_id"]))) as doc:
        t = doc[0].get_text()
        assert "Győr, Ős tér 2." in t and "Kőváry Ödön" in t and doc[0].get_images()


async def test_agent_fill_uses_ai_when_configured(samples, put_file, monkeypatch):
    from app.config import settings
    from app.llm import siliconflow
    monkeypatch.setattr(settings, "SILICONFLOW_API_KEY", "teszt")
    seen = {}

    async def fake_call(model, messages, thinking=True, json_mode=True):
        seen["prompt"] = messages[0]["content"]
        return json.dumps({"lakcim": "Pécs, Ősz utca 9.", "nev": "Kitalált Név"})
    monkeypatch.setattr(siliconflow, "_call_model", fake_call)
    r = await run("ai-form-filler", [put_file(samples.flat_form_pdf())], profile={"full_name": "Kőváry Ödön"})
    d = r["data"]
    assert d["values"]["lakcim"] == "Pécs, Ősz utca 9." and d["provenance"]["lakcim"] == "ai"
    assert d["values"]["nev"] == "Kőváry Ödön"                          # a profil elsőbbséget kap az AI-val szemben
    assert "Kőváry Ödön" in seen["prompt"]


# ---------------------------------------------------------------- csatornák: web, REST API, MCP

def test_web_signatures_profile_and_pages(client, db_client, samples):
    from app import db
    from app.auth import COOKIE, new_session
    call = db_client.portal.call
    u = call(db.user_create, "alairo@example.com", "jelszó123", "Aláíró Ödön", 0)
    db_client.cookies.set(COOKIE, new_session(u["id"]))
    try:
        r = db_client.put("/api/profile", json={"full_name": "Aláíró Ödön", "taj": "123456788"})
        assert r.status_code == 200 and r.json()["profile"]["taj"] == "123456788"
        assert db_client.get("/api/profile").json()["profile"]["full_name"] == "Aláíró Ödön"
        assert db_client.put("/api/profile", json={"nincs": 1}).status_code == 400
        row = call(db.profile_get, u["id"])
        assert "Aláíró".encode() not in bytes(row["data"])
        up = db_client.post("/api/upload", files={"file": ("alairas.jpg", samples.signature_photo().read_bytes(), "image/jpeg")}).json()
        job = db_client.post("/api/tools/run", json={"tool": "create-signature", "file_ids": [up["file_id"]], "options": {"name": "Ödön"}}).json()
        import time
        for _ in range(200):
            job = db_client.get(f"/api/jobs/{job['id']}").json()
            if job["status"] in ("done", "failed"):
                break
            time.sleep(0.05)
        assert job["status"] == "done", job
        sid = job["result"]["data"]["signature_id"]
        sigs = db_client.get("/api/signatures").json()["signatures"]
        assert [s["id"] for s in sigs] == [sid] and sigs[0]["name"] == "Ödön"
        img = db_client.get(f"/api/signatures/{sid}/image")
        assert img.status_code == 200 and img.content[:4] == b"\x89PNG" and "no-store" in img.headers["cache-control"]
        assert db_client.delete(f"/api/signatures/{sid}").status_code == 200
        assert db_client.get(f"/api/signatures/{sid}/image").status_code == 404
        assert db_client.delete("/api/profile").json()["deleted"] is True
    finally:
        db_client.cookies.clear()
    from app.site import urls
    html = client.get(urls.tool("en", "create-signature")).text
    assert "<canvas" in html and "sig-list" in html
    html = client.get(urls.tool("hu", "ai-form-filler")).text
    assert "Mentett aláírásaid" in html


def test_web_profile_needs_account(client, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "DASHBOARD_PASSWORD", "pw")
    client.cookies.clear()
    assert client.get("/api/profile").status_code == 401
    assert client.get("/api/signatures").status_code == 200


def test_rest_api_forms_signatures_profile(client, samples, monkeypatch):
    import base64
    from app.config import settings
    monkeypatch.setattr(settings, "MCP_AUTH_TOKEN", "test-admin-token")
    client.headers["Authorization"] = "Bearer test-admin-token"
    try:
        doc = base64.b64encode(samples.form_docx().read_bytes()).decode()
        job = client.post("/v1/jobs", json={"tasks": {
            "be": {"operation": "import/base64", "file": doc, "filename": "meghatalmazás.docx"},
            "mezok": {"operation": "tool/form-recognize", "input": "be"},
            "kitolt": {"operation": "tool/fill-form", "input": "be", "values": {"Név": "Kőváry Ödön", "Hozzájárulok": True}},
            "ki": {"operation": "export/url", "input": "kitolt"}}}).json()
        job = client.get(f"/v1/jobs/{job['id']}", params={"wait": 30}).json()
        assert job["status"] == "finished", job
        f = [t for t in job["tasks"] if t["name"] == "ki"][0]["result"]["files"][0]
        assert f["filename"].endswith(".docx")
        from docx import Document
        import io
        assert "Név: Kőváry Ödön" in [p.text for p in Document(io.BytesIO(client.get(f["url"]).content)).paragraphs]
        r = client.put("/v1/profile", json={"full_name": "API Ödön", "email": "odon@example.com"})
        assert r.status_code == 200 and client.get("/v1/profile").json()["data"]["email"] == "odon@example.com"
        assert client.put("/v1/profile", json={"rossz": 1}).status_code == 422
        assert isinstance(client.get("/v1/signatures").json()["data"], list)
        assert client.delete("/v1/signatures/sig_nincsilyenalairas").status_code == 404
        assert client.delete("/v1/profile").json()["deleted"] is True
    finally:
        client.headers.pop("Authorization", None)


HDR = {"Accept": "application/json, text/event-stream", "Content-Type": "application/json"}


def _mcp(client, method, params=None, caps=None):
    client.post("/mcp/", headers=HDR, json={"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {
        "protocolVersion": "2025-06-18", "capabilities": caps or {}, "clientInfo": {"name": "t", "version": "1"}}})
    return client.post("/mcp/", headers=HDR, json={"jsonrpc": "2.0", "id": 2, "method": method, "params": params or {}}).json()


def test_mcp_forms_signature_assistant_and_ui(client, samples, put_file):
    tools = {t["name"]: t for t in _mcp(client, "tools/list")["result"]["tools"]}
    assert {"forms", "signature", "fill_assistant", "pdf_form"} <= set(tools)
    meta = tools["fill_assistant"]["_meta"]
    assert meta["ui"]["resourceUri"] == "ui://forms/form-review.html" and meta["openai/outputTemplate"].endswith("skybridge.html")
    res = {r["uri"]: r for r in _mcp(client, "resources/list")["result"]["resources"]}
    assert res["ui://forms/form-review.html"]["mimeType"] == "text/html;profile=mcp-app"
    assert res["ui://forms/signature-pad.skybridge.html"]["mimeType"] == "text/html+skybridge"
    html = _mcp(client, "resources/read", {"uri": "ui://forms/form-review.html"})["result"]["contents"][0]["text"]
    assert "const bridge" in html and 'src="bridge.js"' not in html and "ui/initialize" in html
    fid = put_file(samples.flat_form_pdf())
    r = _mcp(client, "tools/call", {"name": "forms", "arguments": {"action": "recognize", "file_id": fid}})["result"]
    data = json.loads(r["content"][0]["text"])
    assert [f["name"] for f in data["data"]["fields"]][:2] == ["nev", "lakcim"]
    r = _mcp(client, "tools/call", {"name": "fill_assistant", "arguments": {"file_id": fid, "profile": {"full_name": "Kőváry Ödön"}}})["result"]
    out = r.get("structuredContent") or json.loads(r["content"][0]["text"])
    assert out["status"] == "needs_input" and out["values"]["nev"] == "Kőváry Ödön" and out["form_file_id"] == fid
    assert {q["name"] for q in out["questions"]} == {"lakcim", "adoazonosito_jel"}
    r = _mcp(client, "tools/call", {"name": "fill_assistant", "arguments": {"file_id": fid, "profile": {"full_name": "Kőváry Ödön"},
                                                                            "answers": {"lakcim": "Győr", "adoazonosito_jel": "8123456786"}}})["result"]
    out = r.get("structuredContent") or json.loads(r["content"][0]["text"])
    assert out["status"] == "complete"


async def test_mcp_elicitation_asks_for_missing():
    from app.mcp_ext import _elicit_answers

    class Res:
        action = "accept"

        def __init__(self, data):
            self.data = data

    class Caps:
        elicitation = object()

    class Ctx:
        class session:
            class client_params:
                capabilities = Caps()

        async def elicit(self, message, schema):
            self.schema = schema
            return Res(schema(lakcim="Győr, Ős tér 2.", hozzajarulok=True))
    ctx = Ctx()
    got = await _elicit_answers(ctx, [{"name": "lakcim", "label": "Lakcím", "type": "text"},
                                      {"name": "hozzajarulok", "label": "Hozzájárulok", "type": "checkbox"}])
    assert got == {"lakcim": "Győr, Ős tér 2.", "hozzajarulok": True}
    assert set(ctx.schema.model_fields) == {"lakcim", "hozzajarulok"}

    class NoCaps:
        session = None
    assert await _elicit_answers(NoCaps(), [{"name": "x", "label": "X", "type": "text"}]) is None


def test_agent_has_form_assistant():
    from app.llm import agent
    names = {s["function"]["name"] for s in agent.TOOL_SCHEMAS}
    assert "form_assistant" in names and "form_assistant" in agent.TOOL_HANDLERS
