"""Megnézem → célzottan javítom → újraellenőrzöm (a ChatGPT-teszt 10-17 tapasztalatai).

- Csak olvasó parancs (inspect, *.get, *.list…) NEM ment új fájlt: ugyanaz a file_id jön vissza a teljes eredménnyel.
- A pakli-készítés végén ugyanaz a renderelt elemző fut, mint a validate_document-ben (egy kódút, egy ítélet).
- A hibajegyzék a pptx saját alakzat-azonosítóját adja és kész office-parancsot, ami csak azt az alakzatot javítja.
A motoros részek csak ott futnak, ahol a Craft-CLI-k (és a LibreOffice) telepítve vannak."""
from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import pytest



def _deckcraft():
    from app.tools import crafts
    if crafts.office_cli("deck") is None:
        pytest.skip("deckcraft-cli not installed")


def _outputs() -> set[str]:
    from app.storage import output_path
    return {p.name for p in output_path("x.pptx").parent.glob("*")}


def _deck(path: Path, overflow: bool = True) -> None:
    """Egy dia: fix méretű doboz (autofit nélkül) hosszú szöveggel, alatta egy másik szöveg."""
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    a = s.shapes.add_textbox(Pt(40), Pt(40), Pt(400), Pt(60))
    a.text_frame.word_wrap, a.text_frame.auto_size = True, MSO_AUTO_SIZE.NONE
    a.text_frame.text = ("Ez egy nagyon hosszú zárómondat, amely biztosan nem fér el ebben a kicsi dobozban, mert túl sok "
                         "a szó benne." if overflow else "Rövid mondat.")
    a.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    b = s.shapes.add_textbox(Pt(40), Pt(110), Pt(400), Pt(40))
    b.text_frame.text = "Második szöveg"
    b.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    prs.save(str(path))


# ---------------------------------------------------------------- motor nélkül

def test_read_only_commands_are_recognised():
    from app.tools import run_report as rr
    for c in ("slide.inspect", "shape.inspect", "document.inspect", "text.get", "cell.get", "styles.list", "format.state",
              "edit.find", "review.accessibility", "find.objects"):
        assert rr.read_only(c), c
    for c in ("text.set", "format.size", "slide.new", "shape.fill", "cell.set", "edit.delete", "file.save"):
        assert not rr.read_only(c), c
    assert rr.only_reads([("slide.go", {"index": 1}), ("slide.inspect", {"index": 1})])
    assert not rr.only_reads([("slide.go", {"index": 1})])                 # puszta navigáció nem kérdés
    assert not rr.only_reads([("slide.inspect", None), ("text.set", {"id": 3, "text": "x"})])


def test_printed_results_are_parsed_per_command():
    from app.tools import run_report as rr
    out = 'slide.go → {"index":1}\ntext.set → null\nWarning: font fallback\nstyles.list: [1, 2]\nwrote /tmp/x.pptx\n'
    assert rr.arrow_results(out) == [{"index": 1}, None, [1, 2]]
    rows = rr.full_results([("doc.info", None)], [{"path": "/srv/data/inputs/abc.pdf", "pages": 2}])
    assert rows[0]["result"] == {"path": "abc.pdf", "pages": 2}        # a szerver útvonala nem megy ki


def test_on_slide_goes_to_the_slide_first():
    from app.tools.crafts import _on_slide
    got = _on_slide([("text.set", {"on_slide": 3, "id": 7, "text": "a"}), ("text.get", {"on_slide": 3, "id": 7}),
                     ("edit.select", {"on_slide": 1, "ids": [2]}), ("format.size", {"size": 12})])
    assert got == [("slide.go", {"index": 2}), ("text.set", {"id": 7, "text": "a"}), ("text.get", {"id": 7}),
                   ("slide.go", {"index": 0}), ("edit.select", {"ids": [2]}), ("format.size", {"size": 12})]
    with pytest.raises(ValueError, match="1 = the first slide"):
        _on_slide([("text.set", {"on_slide": "x", "id": 1, "text": ""})])


# ---------------------------------------------------------------- DeckCraft

async def test_inspect_lists_shapes_without_a_new_file(put_file, tmp_path):
    _deckcraft()
    from pptx import Presentation
    from app.tools import crafts
    p = tmp_path / "d.pptx"
    _deck(p)
    fid = put_file(p)
    before = _outputs()
    r = await crafts.office_read("deck", "inspect", file_id=fid)
    assert r["read_only"] and r["slides_total"] == 1 and _outputs() == before
    ids = {sh.shape_id for sh in Presentation(str(p)).slides[0].shapes}
    assert {s["id"] for s in r["slides"][0]["shapes"]} == ids                  # DeckCraft id = a pptx alakzat-azonosítója
    one = await crafts.office_read("deck", "inspect", file_id=fid, slide=1, shape_id=min(ids))
    assert one["shape"]["id"] == min(ids)
    from app.tools import checks                                               # MCP-n: a nézés nem kap „nothing changed” ítéletet
    assert await checks.run("office", {"app": "deck", "file_id": fid, "read": "inspect", "slide": 1}, r) == (None, None)
    # ugyanez parancsként (ahogy a ChatGPT hívta): nincs új fájl, nincs „nothing changed” hiba, ott az objektumlista
    run = await crafts.office_run("deck", file_id=fid, commands=[{"command": "slide.inspect", "params": {"index": 0}}])
    assert run["file_id"] == fid and run["read_only"] and run["check"]["status"] == "ok" and _outputs() == before
    assert {s["id"] for s in run["results"][0]["result"]["shapes"]} == ids


async def test_deck_check_names_the_shape_and_the_fix_works(put_file, tmp_path):
    _deckcraft()
    if shutil.which("soffice") is None:
        pytest.skip("LibreOffice not installed")
    from app.storage import input_path
    from app.tools import checks, crafts, validate
    p = tmp_path / "d.pptx"
    _deck(p)
    fid = put_file(p)
    made, _ = await checks.check_deck(input_path(fid), None, "pptx_build", {})
    standalone, _ = await validate.validate(fid)
    assert made["status"] == standalone["status"] == "failed"                 # egy ítélet a két helyen
    issue = standalone["issues"][0]
    assert issue["slide"] == 1 and isinstance(issue["shape_id"], int) and "text.set" in issue["fix"]
    m = re.search(r'commands=\[(\{"command": "text.set".*?\}\})\]', made["findings"][0])
    cmd = json.loads(m.group(1))
    assert cmd["params"]["id"] == issue["shape_id"] and cmd["params"]["on_slide"] == 1
    cmd["params"]["text"] = "Rövid zárómondat."
    args = {"app": "deck", "file_id": fid, "commands": [cmd]}
    res = await crafts.office_run("deck", file_id=fid, commands=[cmd])
    after, _ = await checks.run("office", args, res)
    assert after["status"] == "ok", after
    assert any("changed slide 1" in v for v in after["verified"])


async def test_auto_growing_box_is_not_an_overflow(put_file, tmp_path):
    """„Méret a szöveghez” dobozt a PowerPoint a szöveggel együtt növeli — csak a ráfutás hiba."""
    _deckcraft()
    if shutil.which("soffice") is None:
        pytest.skip("LibreOffice not installed")
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Pt
    from app.storage import input_path
    from app.tools import checks
    p = tmp_path / "d.pptx"
    _deck(p)
    prs = Presentation(str(p))
    box = prs.slides[0].shapes[0]
    box.text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    prs.slides[0].shapes[1].top = Pt(400)                                      # semmire nem fut rá
    prs.save(str(p))
    c, _ = await checks.check_deck(input_path(put_file(p)), None, "pptx_build", {})
    assert not any("past its box" in f for f in c["findings"]), c


async def test_repeated_labels_are_measured_at_their_own_box_and_text_on_text_is_an_error(put_file, tmp_path):
    """10-17: a döntési fa „igen”/„nem” címkéit a gép a dia első „igen”-jéhez mérte (231 pt téves túllógás), a címlapon
    viszont a felső sor a címre íródott, és nem vette észre. Most: a saját dobozához legközelebbi előfordulás az övé;
    két egymásra írt szöveg hiba, kész paranccsal, ami a kisebb dobozt elteszi."""
    _deckcraft()
    if shutil.which("soffice") is None:
        pytest.skip("LibreOffice not installed")
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from app.tools import validate
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    def box(x, y, w, h, text, size):
        b = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        b.text_frame.word_wrap = True
        b.text_frame.text = text
        b.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        return b
    for x, y, t in ((600, 60, "igen"), (600, 160, "nem"), (760, 300, "igen"), (760, 420, "nem")):
        box(x, y, 60, 26, t, 14)
    title = box(40, 200, 480, 200, "Szeged energiaátállása 2027–2030", 44)
    kicker = box(40, 205, 400, 26, "KÖZGYŰLÉSI ELŐTERJESZTÉS", 14)
    p = tmp_path / "d.pptx"
    prs.save(str(p))
    info, _ = await validate.validate(put_file(p))
    kinds = [i["kind"] for i in info["issues"]]
    assert "text_overflow" not in kinds, info["issues"]
    hit = [i for i in info["issues"] if i["kind"] == "text_collision"]
    assert len(hit) == 1 and hit[0]["severity"] == "error" and hit[0]["shape_id"] == kicker.shape_id
    assert '"shape.setBounds"' in hit[0]["fix"] and title.shape_id != kicker.shape_id
