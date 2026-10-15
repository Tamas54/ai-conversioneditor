"""Minden kimenet iPhone-on is menjen: videó/hang (faststart, hvc1, H.264 4:2:0, AAC), Word/Excel sémahelyesen
(Pages/Numbers/Quick Look), a dián lévő videó, letöltés HEAD-del és byte-range-dzsel."""
from __future__ import annotations

import asyncio
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _encoders() -> str:
    if not shutil.which("ffmpeg"):
        return ""
    return subprocess.run(["ffmpeg", "-hide_banner", "-encoders"], capture_output=True, text=True).stdout


needs_ffmpeg = pytest.mark.skipif("libx264" not in _encoders(), reason="ffmpeg with libx264 needed")


def _video(path: Path, *args: str) -> Path:
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=320x240:rate=25",
                    "-f", "lavfi", "-i", "sine=f=440:sample_rate=48000", "-t", "1", *args, str(path)], check=True)
    return path


def _streams(path: Path) -> list[dict]:
    out = subprocess.run(["ffprobe", "-v", "error", "-print_format", "json", "-show_streams", str(path)],
                         capture_output=True, text=True).stdout
    return json.loads(out)["streams"]


@needs_ffmpeg
def test_media_gate_fixes_what_iphone_cannot_play(tmp_path):
    from app.tools import ios_safe
    bad = _video(tmp_path / "rgb.mp4", "-c:v", "libx264", "-pix_fmt", "yuv444p", "-c:a", "aac")      # moov a végén
    assert ios_safe.moov_first(bad) is False
    fixes = asyncio.run(ios_safe.finalize(bad))
    assert any("H.264 4:2:0" in f for f in fixes) and "faststart" in fixes
    v = next(s for s in _streams(bad) if s["codec_type"] == "video")
    assert v["codec_name"] == "h264" and v["pix_fmt"] == "yuv420p" and ios_safe.moov_first(bad) is True
    assert asyncio.run(ios_safe.finalize(bad)) == []                     # második futás: nincs teendő

    good = _video(tmp_path / "good.mp4", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-movflags", "+faststart")
    before = good.read_bytes()
    assert asyncio.run(ios_safe.finalize(good)) == [] and good.read_bytes() == before


@needs_ffmpeg
@pytest.mark.skipif("libopus" not in _encoders() or "libx265" not in _encoders(), reason="libopus + libx265 needed")
def test_media_gate_hvc1_tag_and_aac_audio(tmp_path):
    from app.tools import ios_safe
    hev = _video(tmp_path / "hev.mp4", "-c:v", "libx265", "-x265-params", "log-level=error", "-pix_fmt", "yuv420p", "-c:a", "aac")
    assert "HEVC tagged hvc1" in asyncio.run(ios_safe.finalize(hev))
    v = next(s for s in _streams(hev) if s["codec_type"] == "video")
    assert v["codec_name"] == "hevc" and v["codec_tag_string"] == "hvc1"
    op = _video(tmp_path / "opus.mp4", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "libopus")
    assert any("AAC" in f for f in asyncio.run(ios_safe.finalize(op)))
    assert next(s for s in _streams(op) if s["codec_type"] == "audio")["codec_name"] == "aac"
    # a kifejezetten kért kodeket nem kódoljuk újra (csak a veszteségmentes javítások)
    op2 = _video(tmp_path / "opus2.mp4", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "libopus")
    assert asyncio.run(ios_safe.finalize(op2, {"audio_codec": "opus"})) == ["faststart"]


@needs_ffmpeg
def test_converted_and_tool_videos_are_faststart(client, put_file, tmp_path):
    from app.convert import service
    from app.storage import output_path
    from app.tools import ios_safe
    fid = put_file(_video(tmp_path / "src.mp4", "-c:v", "libx264", "-pix_fmt", "yuv444p", "-c:a", "aac"))
    r = asyncio.run(service.convert_file(fid, "mov", {}))
    out = output_path(r["file_id"])
    assert ios_safe.moov_first(out) is True
    assert next(s for s in _streams(out) if s["codec_type"] == "video")["pix_fmt"] == "yuv420p"


@needs_ffmpeg
def test_video_on_a_slide_plays_on_iphone(put_file, tmp_path):
    """A „videó diára” a beágyazott videót is iPhone-biztossá teszi (4:4:4 → 4:2:0, moov elöl)."""
    from app.storage import output_path
    from app.tools import ios_safe, video
    fid = put_file(_video(tmp_path / "clip.mp4", "-c:v", "libx264", "-pix_fmt", "yuv444p", "-c:a", "aac"))
    r = asyncio.run(video.to_pptx(fid, None, None, "Árvíztűrő"))
    z = zipfile.ZipFile(output_path(r["file_id"]))
    media = [n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith((".mp4", ".mov"))]
    assert media
    emb = tmp_path / Path(media[0]).name
    emb.write_bytes(z.read(media[0]))
    assert ios_safe.moov_first(emb) is True
    assert next(s for s in _streams(emb) if s["codec_type"] == "video")["pix_fmt"] == "yuv420p"
    pres = z.read("ppt/presentation.xml")
    assert b'cx="12192000" cy="6858000"' in pres and b"screen4x3" not in pres
    sld = z.read("ppt/slides/slide1.xml").decode()
    # a videó magától indul a dia megjelenésekor, koppintásra szünet (a PowerPoint saját időzítése)
    assert "playFrom(0.0)" in sld and "togglePause" in sld and 'nodeType="mainSeq"' in sld and 'nodeType="interactiveSeq"' in sld
    assert sld.count("<p:timing>") == 1


@needs_ffmpeg
def test_keynote_pass_makes_python_pptx_videos_play(tmp_path):
    """Egy python-pptx-szel (horog nélkül) beszúrt videó csak „várakozó” médiacsomópontot kap — a javító a
    PowerPoint időzítésére cseréli, a videó valódi hosszával."""
    from pptx import Presentation
    from pptx.util import Inches
    from app.tools import keynote_safe
    clip = _video(tmp_path / "c.mp4", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac")
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_movie(str(clip), 0, 0, Inches(4), Inches(3), mime_type="video/mp4")
    p = tmp_path / "v.pptx"
    prs.part.package.save(str(p))
    assert "mediacall" not in zipfile.ZipFile(p).read("ppt/slides/slide1.xml").decode()
    assert any("video autoplay" in x for x in keynote_safe.fix(p))
    sld = zipfile.ZipFile(p).read("ppt/slides/slide1.xml").decode()
    assert "playFrom(0.0)" in sld and 'dur="1000"' in sld
    assert keynote_safe.fix(p) == []


def _assert_ordered(root):
    from app.tools.ooxml_safe import _ORDER_W
    for el in root.iter(f"{W}*"):
        local = etree.QName(el).localname
        if local not in _ORDER_W:
            continue
        order = _ORDER_W[local]
        idx = [order[etree.QName(c).localname] for c in el if isinstance(c.tag, str)
               and etree.QName(c).namespace == W[1:-1] and etree.QName(c).localname in order]
        assert idx == sorted(idx), f"{local}: children out of schema order"


def test_docx_and_xlsx_are_schema_ordered_and_clean(tmp_path):
    from docx import Document
    from docx.shared import Pt
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    import app.tools  # noqa: F401  — a mentés-horgokat a csomag telepíti
    doc = Document()
    run = doc.add_paragraph().add_run("Árvíztűrő tükörfúrógép")
    run.font.size = Pt(12)
    rpr = run._r.get_or_add_rPr()
    rpr.append(rpr.makeelement(f"{W}b", {}))                       # w:b a w:sz UTÁN — a séma szerint rossz
    tc = doc.add_table(rows=1, cols=1).cell(0, 0)._tc.get_or_add_tcPr()
    shd = tc.makeelement(f"{W}shd", {f"{W}fill": "#3366CC"})        # típus nélkül, „#” színnel (mint a pdf2docx)
    tc.append(shd)
    p = tmp_path / "x.docx"
    doc.save(str(p))
    z = zipfile.ZipFile(p)
    root = etree.fromstring(z.read("word/document.xml"))
    _assert_ordered(root)
    s = next(root.iter(f"{W}shd"))
    assert s.get(f"{W}val") == "clear" and s.get(f"{W}fill") == "3366CC"
    assert Document(str(p)).paragraphs[-1].text == "Árvíztűrő tükörfúrógép"

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ő ű"
    ws["A1"].font = Font(name="Arial", bold=True, color="FF0000", size=14)
    x = tmp_path / "x.xlsx"
    wb.save(str(x))
    styles = etree.fromstring(zipfile.ZipFile(x).read("xl/styles.xml"))
    X = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    from app.tools.ooxml_safe import _ORDER_X
    for f in styles.iter(f"{X}font"):
        idx = [_ORDER_X["font"][etree.QName(c).localname] for c in f]
        assert idx == sorted(idx)
    assert load_workbook(str(x)).active["A1"].value == "ő ű"


def test_ooxml_fix_repairs_engine_quirks_and_is_idempotent(tmp_path):
    """pdf2docx / pandoc-szerű hibák: „16.0” egész helyett, rövid nsid, kettőzött proofState, m:mathPr rossz helyen."""
    from docx import Document
    from app.tools import ooxml_safe
    p = tmp_path / "q.docx"
    Document().save(str(p))
    src = zipfile.ZipFile(p)
    data = {i.filename: src.read(i.filename) for i in src.infolist()}
    doc = data["word/document.xml"].decode()
    doc = doc.replace("<w:body>", '<w:body><w:p><w:pPr><w:jc w:val="left"/><w:pStyle w:val="Normal"/></w:pPr>'
                                  '<w:r><w:rPr><w:rFonts w:ascii="Aptos;Calibri;Segoe UI;Arial;sans-serif" '
                                  'w:hAnsi="\'Segoe UI\', Arial"/><w:sz w:val="22.0"/><w:b/></w:rPr><w:t>ő</w:t></w:r></w:p>', 1)
    data["word/document.xml"] = doc.encode()
    st = data["word/settings.xml"].decode()
    st = st.replace("</w:settings>", '<w:proofState w:spelling="clean"/><w:proofState w:grammar="clean"/>'
                    '<m:mathPr xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"/>'
                    '<w:zoom w:percent="100"/></w:settings>')
    data["word/settings.xml"] = st.encode()
    bad = tmp_path / "bad.docx"
    with zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as zo:
        for n, b in data.items():
            zo.writestr(n, b)
    assert ooxml_safe.fix(bad)
    assert ooxml_safe.fix(bad) == []
    z = zipfile.ZipFile(bad)
    assert z.testzip() is None
    root = etree.fromstring(z.read("word/document.xml"))
    _assert_ordered(root)
    assert next(root.iter(f"{W}sz")).get(f"{W}val") == "22"
    rf = next(root.iter(f"{W}rFonts"))                       # LibreOffice HTML-import: CSS-betűlista → első család
    assert (rf.get(f"{W}ascii"), rf.get(f"{W}hAnsi")) == ("Aptos", "Segoe UI")
    settings = etree.fromstring(z.read("word/settings.xml"))
    _assert_ordered(settings)
    names = [etree.QName(c).localname for c in settings]
    assert names.count("proofState") == 1 and names.index("zoom") < names.index("proofState")
    ps = next(settings.iter(f"{W}proofState"))
    assert ps.get(f"{W}spelling") == "clean" and ps.get(f"{W}grammar") == "clean"
    assert names.index("mathPr") > names.index("proofState")
    assert Document(str(bad)).paragraphs[0].text == "ő"


def test_download_head_and_range(client):
    from app.storage import output_path, public_url, write_meta
    from app.storage.files import new_file_id
    fid = new_file_id("mp4")
    output_path(fid).write_bytes(b"\x00" * 4096)
    write_meta(fid, operation="test", label="clip")
    url = public_url(fid).replace("http://testserver", "")
    h = client.head(url)
    assert h.status_code == 200 and h.headers["content-type"] == "video/mp4" and h.content == b""
    assert h.headers.get("accept-ranges") == "bytes"
    r = client.get(url, headers={"Range": "bytes=0-99"})
    assert r.status_code == 206 and len(r.content) == 100 and r.headers["content-range"] == "bytes 0-99/4096"
    assert "clip.mp4" in r.headers["content-disposition"]


def test_epub_mimetype_first_for_apple_books(tmp_path):
    from app.tools import ios_safe
    p = tmp_path / "book.epub"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:      # rossz: a mimetype nem elöl, tömörítve
        z.writestr("META-INF/container.xml", "<container/>")
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("OEBPS/ch1.xhtml", "<p>Árvíztűrő</p>")
    assert ios_safe.epub_ocf(p) == ["EPUB mimetype first"]
    z = zipfile.ZipFile(p)
    first = z.infolist()[0]
    assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED
    assert z.read("OEBPS/ch1.xhtml").decode() == "<p>Árvíztűrő</p>" and z.testzip() is None
    assert ios_safe.epub_ocf(p) == []
    if shutil.which("pandoc"):                                      # a pandoc EPUB-ja eleve jó: nem nyúlunk hozzá
        md = tmp_path / "a.md"
        md.write_text("# Fejezet\n\nŐsz ű", "utf-8")
        e = tmp_path / "a.epub"
        subprocess.run(["pandoc", str(md), "-o", str(e), "--metadata", "title=T"], check=True)
        before = e.read_bytes()
        assert ios_safe.epub_ocf(e) == [] and e.read_bytes() == before


APPLE_MIME = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
              "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
              "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
              "jpg": "image/jpeg", "png": "image/png", "heic": "image/heic", "gif": "image/gif", "webp": "image/webp",
              "mp4": "video/mp4", "mov": "video/quicktime", "m4a": "audio/mp4", "mp3": "audio/mpeg",
              "epub": "application/epub+zip", "zip": "application/zip", "csv": "text/csv", "txt": "text/plain",
              "html": "text/html", "svg": "image/svg+xml", "key": "application/vnd.apple.keynote"}


def test_every_format_downloads_with_its_real_extension_and_type():
    """Mind a 264 formátum: a letöltött név a valódi kiterjesztésre végződik (Safari nem fűz hozzá „.txt”-t),
    a típus nem „octet-stream”, és az Apple-appoknak fontosak a szabványos MIME-ot kapják."""
    from app.catalog import formats
    from app.routes.files import download_meta
    assert len(formats.FORMATS) >= 264
    for ext in formats.FORMATS:
        name, media = download_meta(f"abc.{ext}", "Árvíztűrő jelentés")
        assert name == f"Árvíztűrő jelentés.{ext}" and media != "application/octet-stream", (ext, name, media)
        assert download_meta(f"abc.{ext}", "jelentés.txt")[0] == f"jelentés.{ext}" or ext == "txt", ext
    for ext, mime in APPLE_MIME.items():
        assert download_meta(f"abc.{ext}", "x")[1] == mime, ext


def test_output_checker_catches_what_breaks_on_a_phone(tmp_path, samples):
    """A sweep --verify ellenőrzője: jó fájlra csend, rosszra pontos hiba."""
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
    import output_check as oc
    good = samples.docx()
    assert oc.check(good, "docx", text=True) == []
    fake = tmp_path / "fake.pdf"
    fake.write_bytes(b"<html>not a pdf</html>")
    assert "wrong signature" in oc.check(fake, "pdf")[0]
    latin = tmp_path / "latin.txt"
    latin.write_text("Árvíztûrõ tükörfúrógép", "utf-8")
    assert oc.check(latin, "txt", text=True) == ["accent ő missing", "accent ű missing", "Latin-1 substitute õ/û found (ő/ű lost)"]
    ep = tmp_path / "bad.epub"
    with zipfile.ZipFile(ep, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("META-INF/container.xml", "<container/>")
        z.writestr("mimetype", "application/epub+zip")
    assert oc.check(ep, "epub")[0] == "'mimetype' is not the first zip entry"
    bad_json = tmp_path / "x.json"
    bad_json.write_text('{"a": ', "utf-8")
    assert oc.check(bad_json, "json")[0].startswith("does not parse as json")


def test_chain_step_never_writes_over_its_own_input(tmp_path, monkeypatch):
    """md → pdf → PDF/A: a 2. lépés kimenete ugyanazt a nevet kapta, mint az 1.-é (s1/jegyzet.pdf), és a Ghostscript
    a saját bemenetére írt — üres lap lett. A lépésmappák számozása most a láncon át folytatódik."""
    from app.convert import graph, service, steps
    seen: list[tuple[Path, Path]] = []

    async def write(src, fmt, ctx):
        dst = ctx.out(src, fmt)
        seen.append((src, dst))
        assert dst != src, "a lépés a saját bemenetére írna"
        dst.write_bytes(src.read_bytes() + fmt.encode())
        return [dst]
    monkeypatch.setitem(steps.STEPS, "fake1", write)
    monkeypatch.setitem(steps.STEPS, "fake2", write)
    monkeypatch.setattr(service, "plan_or_raise", lambda *a, **k: [graph.Edge("md", "pdf", "fake1", 1),
                                                                    graph.Edge("pdf", "pdfa", "fake2", 1)])
    src = tmp_path / "jegyzet.md"
    src.write_text("ő ű", "utf-8")
    outs, used = asyncio.run(service.convert_path(src, "pdfa", {}, tmp_path / "w"))
    assert used == ["fake1", "fake2"] and outs[0].read_bytes() == "ő ű".encode() + b"pdfpdfa"
    assert len({d for _, d in seen}) == 2


def test_office_package_content_types_first(tmp_path):
    """A LibreOffice a [Content_Types].xml-t a csomag végére írja; a kapu előre teszi, a tartalom bájtra ugyanaz."""
    from app.tools import ios_safe
    from docx import Document
    good = tmp_path / "a.docx"
    d = Document()
    d.add_paragraph("Árvíztűrő ő ű")
    d.save(good)
    bad = tmp_path / "b.docx"
    with zipfile.ZipFile(good) as zi, zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as zo:
        names = [n for n in zi.namelist() if n != "[Content_Types].xml"] + ["[Content_Types].xml"]
        for n in names:
            zo.writestr(n, zi.read(n))
    assert asyncio.run(ios_safe.finalize(bad)) == ["[Content_Types].xml first"]
    with zipfile.ZipFile(bad) as z, zipfile.ZipFile(good) as g:
        assert z.namelist()[0] == "[Content_Types].xml" and z.testzip() is None
        assert {n: z.read(n) for n in z.namelist()} == {n: g.read(n) for n in g.namelist()}
    assert asyncio.run(ios_safe.finalize(bad)) == []


def test_fb2_and_xml_lose_illegal_control_characters(tmp_path):
    """A PDF-ből kinyert szöveg lapdobása (&#12;) tiltott XML 1.0-ban — a szigorú e-könyv-olvasó meg sem nyitja."""
    from app.tools import ios_safe
    p = tmp_path / "konyv.fb2"
    p.write_bytes('<?xml version="1.0" encoding="utf-8"?><FictionBook><body><p>ő&#12;ű\x0c&#xB; &#10;&#9; Ő</p></body></FictionBook>'.encode())
    assert asyncio.run(ios_safe.finalize(p)) == ["invalid XML characters removed"]
    t = p.read_text("utf-8")
    assert "<p>őű &#10;&#9; Ő</p>" in t
    etree.fromstring(p.read_bytes())
    assert asyncio.run(ios_safe.finalize(p)) == []
    u16 = tmp_path / "x.xml"
    u16.write_bytes("<a>\x0c</a>".encode("utf-16"))
    assert ios_safe.xml_clean(u16) == []                           # UTF-16-hoz nem nyúl


def test_cursor_is_a_cursor_not_an_icon(tmp_path):
    from app.convert.steps import _ico_as_cur
    from PIL import Image
    p = tmp_path / "k.cur"
    Image.new("RGBA", (32, 32), (200, 30, 30, 255)).save(p, format="ICO", sizes=[(32, 32)], bitmap_format="bmp")   # mint az ImageMagick: BMP-adat
    assert p.read_bytes()[:4] == b"\0\0\1\0"
    _ico_as_cur(p)
    b = p.read_bytes()
    assert b[:4] == b"\0\0\2\0" and b[10:14] == b"\0\0\0\0"        # típus 2, forrópont (0, 0)
    with Image.open(p) as im:
        assert im.format == "CUR" and im.size == (32, 32)
