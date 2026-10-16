"""Az MCP-n épített PPTX az Apple Keynote-ban is nyíljon meg (a PowerPoint elnéző, a Keynote nem).

Rögzített hibák: a jegyzetmester nincs beírva a presentation.xml-be, jegyzetdia színleképezés nélkül,
ellentmondó sldSz type, Windows nyomtató-csonk, negatív diagram-tengely-id; letöltésnél kiterjesztés nélküli
név / text/plain típus (a Safari .txt-t fűz a .pptx után).
"""
from __future__ import annotations

import asyncio
import io
import re
import zipfile

from lxml import etree

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def _deck(tmp_path):
    from PIL import Image
    from app.storage import output_path, write_meta
    from app.storage.files import new_file_id
    from app.tools import pptx_deck
    img = new_file_id("png")
    Image.new("RGB", (1200, 800), (40, 90, 160)).save(output_path(img))
    write_meta(img, operation="test", label="kép.png")
    spec = {"title": "Árvíztűrő tükörfúrógép", "footer": "ő ű", "slides": [
        {"layout": "cover", "title": "Árvíztűrő tükörfúrógép", "image": img, "notes": "Előadói jegyzet: ő, ű"},
        {"layout": "chart", "title": "Oszlop", "chart_type": "column", "categories": ["Q1", "Q2"],
         "series": [{"name": "A", "values": [1, 2]}, {"name": "B", "values": [3, 4]}], "notes": "második"},
        {"layout": "chart", "title": "Szórás", "chart_type": "scatter", "series": [{"name": "A", "points": [[1, 2], [3, 4]]}]},
        {"layout": "closing", "title": "Köszönöm"}]}
    out = asyncio.run(pptx_deck.build(spec))
    return output_path(out["file_id"])


def test_mcp_deck_is_keynote_safe(tmp_path):
    path = _deck(tmp_path)
    z = zipfile.ZipFile(path)
    names = z.namelist()
    assert names[0] == "[Content_Types].xml"
    assert not [n for n in names if "printerSettings" in n]
    assert b"printerSettings" not in z.read("[Content_Types].xml") + z.read("ppt/_rels/presentation.xml.rels")

    pres = etree.fromstring(z.read("ppt/presentation.xml"))
    rels = etree.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
    nm_rid = next(r.get("Id") for r in rels if r.get("Type", "").endswith("/notesMaster"))
    lst = pres.find(P + "notesMasterIdLst")
    assert lst is not None and lst[0].get(R_ID) == nm_rid
    order = [etree.QName(c).localname for c in pres]
    assert order.index("sldMasterIdLst") < order.index("notesMasterIdLst") < order.index("sldIdLst")
    sz = pres.find(P + "sldSz")
    assert (sz.get("cx"), sz.get("cy"), sz.get("type")) == ("12192000", "6858000", None)

    notes = [n for n in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)]
    assert len(notes) == 2
    for n in notes:
        root = etree.fromstring(z.read(n))
        kids = [etree.QName(c).localname for c in root]
        assert kids.index("clrMapOvr") == kids.index("cSld") + 1

    charts = [n for n in names if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)]
    assert len(charts) == 2
    for n in charts:
        root = etree.fromstring(z.read(n))
        ax = [int(e.get("val")) for e in root.iter(C + "axId")]
        cross = {int(e.get("val")) for e in root.iter(C + "crossAx")}
        assert ax and all(0 < v <= 2147483647 for v in ax)
        assert cross <= set(ax)          # a hivatkozások ugyanarra mutatnak

    from pptx import Presentation
    prs = Presentation(str(path))            # a python-pptx (újraépítés, szerkesztés) továbbra is megnyitja
    assert prs.slides[0].notes_slide.notes_text_frame.text == "Előadói jegyzet: ő, ű"
    assert sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart) == 2


def test_fix_is_idempotent_and_reports_nothing_on_clean_files(tmp_path):
    from app.tools import keynote_safe
    path = _deck(tmp_path)
    before = path.read_bytes()
    assert keynote_safe.fix(path) == []
    assert path.read_bytes() == before


def test_fix_repairs_foreign_packages(tmp_path):
    """Kívülről jött / más motor írta csomag: hiányzó clrMapOvr a jegyzetdián, image/jpg, negatív tengely."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    from app.tools import keynote_safe
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.notes_slide.notes_text_frame.text = "jegyzet"
    d = CategoryChartData(); d.categories = ["a"]; d.add_series("s", [1])
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), d)
    raw = io.BytesIO()
    prs.part.package.save(raw)          # nyers csomagírás, a mentés-horog nélkül (mint egy idegen motor)
    src = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    bad = tmp_path / "foreign.pptx"
    with zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in src.infolist():
            data = src.read(i)
            if i.filename.startswith("ppt/notesSlides/notesSlide"):
                data = re.sub(rb"<p:clrMapOvr>.*?</p:clrMapOvr>", b"", data)
            if i.filename == "[Content_Types].xml":
                data = data.replace(b'Extension="png" ContentType="image/png"', b'Extension="jpg" ContentType="image/jpg"')
                if b'Extension="jpg"' not in data:
                    data = data.replace(b"<Default ", b'<Default Extension="jpg" ContentType="image/jpg"/><Default ', 1)
            zo.writestr(i, data)
    assert b"-2068027336" in zipfile.ZipFile(bad).read("ppt/charts/chart1.xml")
    done = keynote_safe.fix(bad)
    assert any("notes slide" in x for x in done) and any("axis ids" in x for x in done)
    assert any("image/jpg" in x for x in done) and any("notesMasterIdLst" in x for x in done)
    z = zipfile.ZipFile(bad)
    assert b"clrMapOvr" in z.read("ppt/notesSlides/notesSlide1.xml")
    assert b"image/jpg" not in z.read("[Content_Types].xml")
    assert b'val="-' not in z.read("ppt/charts/chart1.xml")
    assert keynote_safe.fix(bad) == []
    Presentation(str(bad))


def test_save_hook_also_fixes_streams_and_keeps_valid_4x3():
    from pptx import Presentation
    import app.tools  # noqa: F401  — a horgot a csomag telepíti
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6]).notes_slide.notes_text_frame.text = "x"
    buf = io.BytesIO()
    prs.save(buf)
    z = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    pres = z.read("ppt/presentation.xml")
    assert b"notesMasterIdLst" in pres and not [n for n in z.namelist() if "printerSettings" in n]
    assert b'cx="9144000" cy="6858000" type="screen4x3"' in pres     # a valódi 4:3 maradhat


def test_fix_if_slides_ignores_other_formats(tmp_path):
    from app.tools import keynote_safe
    assert keynote_safe.fix_if_slides(tmp_path / "x.docx") == []          # más formátumhoz nem nyúl


def test_download_name_and_type(client):
    from app.routes.files import download_meta
    ppt = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert download_meta("a.pptx", "Árvíztűrő Q3/Q4") == ("Árvíztűrő Q3-Q4.pptx", ppt)
    assert download_meta("a.pptx", "deck.docx") == ("deck.pptx", ppt)
    assert download_meta("a.pptx", "v1.2 plan") == ("v1.2 plan.pptx", ppt)
    assert download_meta("b.tar.gz", "x.tar.gz") == ("x.tar.gz", "application/gzip")
    assert download_meta("c.webp", "c")[1] == "image/webp"
    assert download_meta("e.deckcraft", "e") == ("e.deckcraft", "application/octet-stream")

    from app.storage import output_path, public_url, write_meta
    from app.storage.files import new_file_id
    fid = new_file_id("pptx")
    from pptx import Presentation
    Presentation().save(str(output_path(fid)))
    write_meta(fid, operation="pptx_build", label="Negyedéves terv")       # kiterjesztés nélküli címke
    r = client.get(public_url(fid).replace("http://testserver", ""))
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert r.headers["content-type"] == ppt
    cd = r.headers["content-disposition"]
    assert "attachment" in cd and ".pptx" in cd and "text/plain" not in r.headers["content-type"]
