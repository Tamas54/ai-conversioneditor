"""Tördelési hibák gépi javítása visszaellenőrzéssel (10-22, ügynök-jelentés): cím a lap alján, képaláírás külön lapon,
két lapra törő táblázatsor, nem ismétlődő fejléc → javítás a DOCX-ben, újratördelés, legfeljebb 3 kör."""
from __future__ import annotations

import asyncio
import shutil
from pathlib import Path

import pytest


def _doc(path: Path) -> Path:
    from docx import Document
    d = Document()
    d.add_heading("1. Bevezetés", 1)                          # 0
    d.add_paragraph("Szöveg.")                               # 1
    d.add_paragraph("")                                      # 2 (üres)
    d.add_paragraph("1. ábra: a kép aláírása")                # 3
    t = d.add_table(rows=3, cols=2)                          # 4
    for r in range(3):
        for c in range(2):
            t.cell(r, c).text = f"{r}{c}"
    d.save(path)
    return path


def _ppr(path: Path, pos: int) -> set[str]:
    from docx import Document
    from app.tools.validate import _W, _body
    el = list(_body(Document(str(path)).element.body))[pos]
    ppr = el.find(_W + "pPr")
    return {c.tag.rsplit("}", 1)[1] for c in ppr} if ppr is not None else set()


def test_fixes_go_into_the_document(tmp_path):
    from docx import Document
    from app.tools import layout_fix
    from app.tools.validate import _W, _body
    src = _doc(tmp_path / "a.docx")
    issues = [{"kind": "orphan_heading", "_pos": 0, "message": "Heading … is the last line", "page": 1, "object": "heading"},
              {"kind": "orphan_heading", "_pos": 0, "message": "same in the other view", "view": "word"},
              {"kind": "caption_split", "_pos": 3, "message": "The picture is on page 1, its caption on page 2."},
              {"kind": "table_row_split", "_pos": 4, "row": 2, "message": "Row 2 splits."},
              {"kind": "table_header_not_repeated", "_pos": 4, "row": 1, "message": "no header"},
              {"kind": "blank_page", "_pos": None, "message": "Page 3 is empty."}]
    todo = layout_fix._plan(issues, {})
    assert [(i["kind"], how) for i, how in todo] == [("orphan_heading", "keep"), ("caption_split", "keep"),
                                                     ("table_row_split", "row"), ("table_header_not_repeated", "header")]
    out = tmp_path / "b.docx"
    done = layout_fix._apply(src, out, todo)
    assert len(done) == 4 and {"keepNext", "keepLines"} <= _ppr(out, 0)
    assert "keepNext" in _ppr(out, 1)                        # a kép UTÁNI aláírás: az előtte álló tartalom tart vele
    tbl = list(_body(Document(str(out)).element.body))[4]
    rows = tbl.findall(_W + "tr")
    assert rows[1].find(_W + "trPr").find(_W + "cantSplit") is not None
    assert rows[0].find(_W + "trPr").find(_W + "tblHeader") is not None
    # ismétlésnél erősebb javítás; a csak a szerkesztőben elhagyott cím (már keepNext) rögtön oldaltörést kap
    tried = {layout_fix._key(i): 1 for i, _ in todo}
    again = layout_fix._plan(issues, tried)
    assert [(i["kind"], how) for i, how in again] == [("orphan_heading", "break"), ("caption_split", "break")]
    editor_only = layout_fix._plan([{"kind": "orphan_heading", "_pos": 0, "message": "The heading is already set to keep "
                                     "with its text, so Word keeps them together"}], {})
    assert editor_only[0][1] == "break"
    out2 = tmp_path / "c.docx"
    layout_fix._apply(out, out2, again)
    assert "pageBreakBefore" in _ppr(out2, 0)
    tall = [{"kind": "table_row_split", "_pos": 4, "row": 2, "message": "splits although it is set not to break — …"}]
    assert layout_fix._plan(tall, {}) == []                  # a lapnál magasabb sor: nem gépi döntés


def _engines_ok() -> bool:
    from app.tools import wordcraft
    return wordcraft.available() and bool(shutil.which("soffice"))


def _orphan_doc(path: Path) -> Path:
    """Mért elrendezés (10-22): 50 rövid sor után a cím a Word-nézet (LibreOffice) 2. oldalának alján marad."""
    from docx import Document
    d = Document()
    d.styles["Heading 1"].paragraph_format.keep_with_next = False
    for i in range(50):
        d.add_paragraph(f"Sor {i + 1}: rövid bekezdés.")
    d.add_heading("2. Következtetések", 1)
    for _ in range(5):
        d.add_paragraph("A fejezet szövege. " * 8)
    t = d.add_table(rows=45, cols=2)
    for r in range(45):
        t.cell(r, 0).text, t.cell(r, 1).text = ("Év", "Érték") if r == 0 else (str(1980 + r), f"{r * 1.5:.1f}")
    d.save(path)
    return path


@pytest.mark.engines
@pytest.mark.skipif(not _engines_ok(), reason="WordCraft and LibreOffice needed")
async def test_validate_document_fixes_and_shows_what_is_left(tmp_path, put_file):
    from app.mcp_server import validate_document
    from app.storage import read_meta
    fid = put_file(_orphan_doc(tmp_path / "tanulmany.docx"))
    res = await validate_document(file_id=fid, fix=True)
    info = res.structuredContent
    kinds = {f["kind"] for f in info["fixed"]}
    assert "table_header_not_repeated" in kinds and info["file_id"] != fid and info["source_file_id"] == fid
    assert read_meta(info["file_id"])["operation"] == "layout fix"
    assert not [i for i in info.get("issues") or [] if i["kind"] in ("orphan_heading", "table_header_not_repeated")]
