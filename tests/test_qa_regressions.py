"""QA-regresszió (10-18): a hibajavítás-sorozat ELŐTTI működés rögzítése, MCP-n át — ahogy a kliensek hívják.

A tulajdonos „NE TÖRD EL” listája: md/docx/pdf konverzió, a WordCraft-út, a szerkesztők (find_replace, insert_after,
delete_section, replace_heading, delete_pages, merge), verziófa és összevetés (history, files), WordCraft követett
változások, táblázat-képletek, darabolt feltöltés UTF-8 határon, jobs-folyamat, időzóna DST-váltáson át, pptx_build a
saját check-jével, design (oklevél, auto), diagram (generate_chart_image, chart_redraw), képolvasás (ocr_image,
extract_table_from_image, image_to_xlsx), digitize, PDF-kitakarás, rövid videó, photo_edit, restore_photos.
Minden bemenet ékezetes (ő, ű), és a kimenetben is ezt nézzük.

Kulcs nincs a tesztben: ahol a szerver saját modellje dolgozik (látó modell, LLM, FLUX Kontext), a LEGALSÓ hívást
cseréljük valószerű kész válaszra, és a determinisztikus részt ellenőrizzük (fájl, cella, szöveg, réteg).
A motorok (LibreOffice, pandoc, WordCraft és a többi Craft-CLI, ffmpeg) a Docker-képben vannak; ahol hiányzik, kihagyás.
"""
from __future__ import annotations

import asyncio
import base64
import io
import json
import shutil
import subprocess
import time
from pathlib import Path

import pytest

from samples import HU

pytestmark = pytest.mark.engines

HDR = {"Accept": "application/json, text/event-stream", "Content-Type": "application/json"}


def _call(client, name, args):
    client.post("/mcp/", headers=HDR, json={"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {
        "protocolVersion": "2025-06-18", "capabilities": {}, "clientInfo": {"name": "t", "version": "1"}}})
    r = client.post("/mcp/", headers=HDR, json={"jsonrpc": "2.0", "id": 2, "method": "tools/call",
                                                "params": {"name": name, "arguments": args}})
    res = r.json()["result"]
    text = res["content"][0]["text"]
    return res.get("isError"), (json.loads(text) if text.startswith("{") else text)


def _ok(client, name, args):
    err, res = _call(client, name, args)
    assert not err, f"{name}: {res}"
    return res


def _wait(client, tool, res, limit=240):
    """Háttérmunka (digitize, photo_edit, restore_photos, video): újrahívás job_id-val, amíg kész nem lesz."""
    t0 = time.time()
    while res.get("status") == "running" and time.time() - t0 < limit:
        time.sleep(0.4)
        res = _ok(client, tool, {"action": "status", "job_id": res["job_id"]} if tool == "video" else {"job_id": res["job_id"]})
    assert res.get("status") == "done", res
    return res


def _out(fid: str) -> Path:
    from app.tools.crafts import existing
    return existing(fid)


def _pdf_text(fid: str) -> str:
    import pymupdf
    with pymupdf.open(str(_out(fid))) as d:
        return "".join(pg.get_text() for pg in d)


def _pdf_pages(fid: str) -> list[str]:
    import pymupdf
    with pymupdf.open(str(_out(fid))) as d:
        return [pg.get_text() for pg in d]


def _docx_paras(fid: str) -> list[tuple[str, str]]:
    from docx import Document
    return [(p.style.name if p.style else "", p.text) for p in Document(str(_out(fid))).paragraphs]


def _need(*engines: str) -> None:
    """A Craft-CLI-k és a többi motor csak a production képben vannak — máshol kihagyás (a Dockerben mind fut)."""
    from app.tools import crafts, wordcraft
    have = {"wordcraft": wordcraft.available, "sheet": lambda: bool(crafts.office_cli("sheet")),
            "deck": lambda: bool(crafts.office_cli("deck")), "pdf": lambda: bool(crafts.office_cli("pdf")),
            "photocraft": lambda: bool(crafts.photo_cli()), "vectorcraft": lambda: bool(crafts.vector_cli())}
    for e in engines:
        ok = have[e]() if e in have else bool(shutil.which(e))
        if not ok:
            pytest.skip(f"{e} is only in the production image")


@pytest.fixture()
def lo(client, monkeypatch):
    """A conftest LO_WARM=0-val indít: a régi konverziós út LibreOffice-szemafora így 0 helyes, és örökre vár.
    Élesben LO_WARM > 0 — itt egy helyet adunk neki, hogy a valódi (LibreOffice) út fusson."""
    _need("soffice")
    from app import mcp_server
    monkeypatch.setattr(mcp_server._pool(), "_lo_sem", asyncio.Semaphore(1))


def _vision(monkeypatch, answer):
    """A látó modell legalsó hívása (app.llm.vision._call) kész válasszal; a látott promptok listája vissza."""
    from app.config import settings
    from app.llm import vision
    monkeypatch.setattr(settings, "SILICONFLOW_API_KEY", "teszt")
    seen: list[str] = []

    async def fake_call(msgs, response_format=None, max_tokens=4096, model=None):
        prompt = msgs[0]["content"][0]["text"]
        seen.append(prompt)
        assert msgs[0]["content"][1]["image_url"]["url"].startswith("data:image/")    # a kép tényleg elmegy
        return answer(prompt) if callable(answer) else answer
    monkeypatch.setattr(vision, "_call", fake_call)
    return seen


# ---------------------------------------------------------------- konverzió: md / docx / pdf, WordCraft

def test_convert_md_docx_pdf_keep_accents(client, samples, put_file, lo):
    md, docx, pdf = put_file(samples.md()), put_file(samples.docx()), put_file(samples.pdf(pages=3))
    r = _ok(client, "convert", {"file_id": md, "target_format": "docx"})
    assert r["backend"] == "wordcraft" and r["check"]["status"] == "ok"      # 1. pont: md → docx alapból WordCraft
    paras = _docx_paras(r["file_id"])
    assert ("Heading 1", HU) in paras and any("félkövér ő ű" in t for _, t in paras)

    r = _ok(client, "convert", {"file_id": md, "target_format": "pdf"})
    flat = " ".join(_pdf_text(r["file_id"]).split())                          # a hosszú címsor két sorba törik
    assert r["backend"] == "md_to_pdf" and HU in flat and "Bekezdés félkövér ő ű." in flat

    r = _ok(client, "convert", {"file_id": docx, "target_format": "pdf"})
    assert r["backend"] == "libreoffice" and r["check"]["status"] == "ok"
    assert "It has a text layer." in r["check"]["verified"]
    t = _pdf_text(r["file_id"])
    assert "Címsor: " + HU in t and "Második bekezdés, ő és ű betűkkel." in t

    r = _ok(client, "convert", {"file_id": docx, "target_format": "md"})
    assert r["backend"] == "pandoc"
    text = _out(r["file_id"]).read_text("utf-8")                              # a pandoc a gondolatjelet „---”-ként írja
    assert "Első bekezdés: Árvíztűrő tükörfúrógép" in text and "Második bekezdés, ő és ű betűkkel." in text

    r = _ok(client, "convert", {"file_id": pdf, "target_format": "docx"})
    assert r["backend"] == "pdf2docx" and r["file_id"].endswith(".docx")
    body = " ".join(t for _, t in _docx_paras(r["file_id"]))
    assert all(f"{i}. oldal: {HU}" in body for i in (1, 2, 3))


def test_wordcraft_engine_is_fast_and_keeps_accents(client, samples, put_file):
    _need("wordcraft")
    docx = put_file(samples.docx())
    r = _ok(client, "convert", {"file_id": docx, "target_format": "pdf", "engine": "wordcraft"})
    assert r["backend"] == "wordcraft" and r["ms_elapsed"] < 1500          # ~20 ms; LibreOffice nélkül
    assert r["check"]["status"] == "ok"
    t = _pdf_text(r["file_id"])
    assert "Első bekezdés: " + HU in t and "ő és ű" in t
    r = _ok(client, "convert", {"file_id": docx, "target_format": "md", "engine": "wordcraft"})
    assert r["backend"] == "wordcraft" and "Második bekezdés, ő és ű" in _out(r["file_id"]).read_text("utf-8")


# ---------------------------------------------------------------- szerkesztők

def _chapters(samples) -> Path:
    from docx import Document
    d = Document()
    d.add_heading("Bevezető ő", 1)
    d.add_paragraph("Első rész szövege, árvíztűrő.")
    d.add_heading("Törlendő fejezet ű", 1)
    d.add_paragraph("Ez a szakasz eltűnik.")
    d.add_heading("Záró gondolatok", 1)
    d.add_paragraph("Vége, ő és ű.")
    p = samples.root / "fejezetek.docx"
    d.save(p)
    return p


def test_find_replace_docx_md_pdf(client, samples, put_file):
    r = _ok(client, "find_replace", {"file_id": put_file(samples.docx()), "pattern": "Második bekezdés",
                                     "replacement": "Őszinte űrlap"})
    assert r["matches"] == 1 and r["check"]["status"] == "ok"
    assert "“Második bekezdés” no longer appears." in r["check"]["verified"]
    assert r["name"] == "level – text replaced.docx"
    assert ("Normal", "Őszinte űrlap, ő és ű betűkkel.") in _docx_paras(r["file_id"])

    r = _ok(client, "find_replace", {"file_id": put_file(samples.md()), "pattern": "ő ű", "replacement": "Ő Ű"})
    assert r["matches"] == 1 and "Bekezdés **félkövér** Ő Ű." in _out(r["file_id"]).read_text("utf-8")

    pdf = put_file(samples.pdf(pages=2))
    r = _ok(client, "find_replace", {"file_id": pdf, "pattern": "tükörfúrógép", "replacement": "fűrészgép"})
    assert r["format"] == "pdf" and r["matches"] == 2
    t = _pdf_text(r["file_id"])
    assert "tükörfúrógép" not in t and t.count("fűrészgép") == 2 and "Árvíztűrő" in t


def test_insert_after_delete_section_replace_heading(client, samples, put_file):
    fid = put_file(_chapters(samples))
    r = _ok(client, "insert_after", {"file_id": fid, "anchor": "Bevezető ő", "content": "Új sor őszi ű\nMásodik új sor"})
    texts = [t for _, t in _docx_paras(r["file_id"])]
    i = texts.index("Bevezető ő")
    assert texts[i + 1:i + 4] == ["Új sor őszi ű", "Második új sor", "Első rész szövege, árvíztűrő."]
    assert r["name"] == "fejezetek – text added.docx" and r["check"]["status"] in ("ok", "needs_check")

    r = _ok(client, "delete_section", {"file_id": fid, "heading": "Törlendő fejezet ű"})
    assert [t for _, t in _docx_paras(r["file_id"])] == ["Bevezető ő", "Első rész szövege, árvíztűrő.",
                                                          "Záró gondolatok", "Vége, ő és ű."]

    r = _ok(client, "replace_heading", {"file_id": fid, "old_title": "Záró gondolatok", "new_title": "Összegzés űrhajóval"})
    paras = _docx_paras(r["file_id"])
    assert ("Heading 1", "Összegzés űrhajóval") in paras and ("Heading 1", "Záró gondolatok") not in paras
    assert r["check"]["status"] == "ok" and "Heading “Összegzés űrhajóval” is there." in r["check"]["verified"]

    md = put_file(samples.md())
    r = _ok(client, "replace_heading", {"file_id": md, "old_title": HU, "new_title": "Új cím: őszi űrlap"})
    assert _out(r["file_id"]).read_text("utf-8").startswith("# Új cím: őszi űrlap\n")
    err, msg = _call(client, "replace_heading", {"file_id": fid, "old_title": "Nincs ilyen ő", "new_title": "x"})
    assert err and "Heading not found" in msg


def test_delete_pages_and_merge(client, samples, put_file):
    pdf = put_file(samples.pdf(pages=3))
    r = _ok(client, "delete_pages", {"file_id": pdf, "page_range": "2"})
    assert r["pages_remaining"] == 2 and r["check"]["status"] == "ok"
    assert r["check"]["verified"] == ["1 of 3 pages removed: 2 pages, as expected."]
    pages = _pdf_pages(r["file_id"])
    assert pages[0].startswith(f"1. oldal: {HU}") and pages[1].startswith(f"3. oldal: {HU}")

    two = put_file(samples.pdf("masik.pdf", pages=2, text="Őszi űrlap, második fájl"))
    r = _ok(client, "merge", {"file_ids": [pdf, two]})
    assert r["check"]["status"] == "ok" and r["check"]["verified"] == ["2 files merged: 5 pages, as expected."]
    pages = _pdf_pages(r["file_id"])
    assert len(pages) == 5 and pages[2].startswith("3. oldal: " + HU) and pages[4].startswith("2. oldal: Őszi űrlap, második")
    assert r["name"] == "dok – merged.pdf"


# ---------------------------------------------------------------- verziófa (history) és összevetés (history, files)

def test_history_lineage(db_client, samples):
    client = db_client
    up = _ok(client, "upload_file", {"filename": "levél ő.docx", "content_base64": base64.b64encode(samples.docx().read_bytes()).decode()})
    assert up["original_name"] == up["name"] == "levél ő.docx"
    a = _ok(client, "find_replace", {"file_id": up["file_id"], "pattern": "Második", "replacement": "Harmadik ű"})
    b = _ok(client, "find_replace", {"file_id": a["file_id"], "pattern": "Első", "replacement": "Nulladik ő"})
    deadline = time.time() + 10                     # a verzió-napló háttérfeladat (write_meta → Postgres)
    while True:
        lin = _ok(client, "history", {"action": "lineage", "file_id": a["file_id"]})
        if (len(lin["ancestors"]), len(lin["descendants"])) == (2, 1) or time.time() > deadline:
            break
        time.sleep(0.2)
    assert [(x["file_id"], x["operation"], x["label"], x["depth"]) for x in lin["ancestors"]] == [
        (up["file_id"], "upload", "levél ő.docx", 1), (a["file_id"], "find_replace", "levél ő – text replaced.docx", 0)]
    assert lin["ancestors"][1]["source_file_id"] == up["file_id"]
    assert [(x["file_id"], x["source_file_id"], x["operation"]) for x in lin["descendants"]] == [
        (b["file_id"], a["file_id"], "find_replace")]
    err, msg = _call(client, "history", {"action": "lineage"})
    assert err and "lineage needs a file_id" in msg


def test_history_compare_and_files_compare(client, samples, put_file):
    base = put_file(samples.docx())
    other = _ok(client, "find_replace", {"file_id": base, "pattern": "Második", "replacement": "Harmadik ű"})["file_id"]
    cmp = _ok(client, "history", {"action": "compare", "file_id": base, "other_file_id": other})
    assert cmp["summary"] == "1 place where the two differ, 2 paragraphs the same."
    assert [(u["kind"], [x["text"] for x in u["a"]], [x["text"] for x in u["b"]]) for u in cmp["units"]] == [
        ("changed", ["Második bekezdés, ő és ű betűkkel."], ["Harmadik ű bekezdés, ő és ű betűkkel."])]

    pdf = put_file(samples.pdf(pages=3))
    cut = _ok(client, "delete_pages", {"file_id": pdf, "page_range": "2"})["file_id"]
    c = _ok(client, "files", {"action": "compare", "file_id": cut})            # alapból a forrásával (amiből készült)
    assert (c["a"]["file_id"], c["b"]["file_id"], c["pages_a"], c["pages_b"]) == (pdf, cut, 3, 2)
    assert c["changed_pages"] == [2] and c["changes"][0]["note"] == "gone from the second file"
    assert f"Page 2 is gone (“2. oldal: {HU}”)." in c["missing"]
    assert c["text"]["removed_sample"] == [f"2. oldal: {HU}"] and c["text"]["lines_added"] == 0

    c = _ok(client, "files", {"action": "compare", "file_id": base, "other_file_id": other})
    assert c["changes"][0]["text_removed"] == ["Második bekezdés, ő és ű betűkkel."]
    assert c["changes"][0]["text_added"] == ["Harmadik ű bekezdés, ő és ű betűkkel."]


# ---------------------------------------------------------------- WordCraft: követett változások szerzővel

def test_wordcraft_track_changes_by_author(client, samples, put_file):
    _need("wordcraft")
    docx = put_file(samples.docx())
    r = _ok(client, "wordcraft", {"file_id": docx, "author": "Kiss Ödön", "commands": [
        {"command": "review.trackChanges", "params": {"value": True}},
        {"command": "select.text", "params": {"text": "Második bekezdés"}},
        {"command": "text.insert", "params": {"text": "Őszinte űrlap"}}]})
    assert r["file_id"] != docx and [s["ok"] for s in r["steps"]] == [True, True, True]
    assert r["before_after"]["after"]["tracked_changes"] == 2 and r["before_after"]["after"]["authors"] == ["Kiss Ödön"]
    assert r["check"]["status"] == "ok" and "2 tracked changes by Kiss Ödön." in r["check"]["verified"]

    rev = _ok(client, "wordcraft", {"file_id": r["file_id"], "read": "review"})
    assert {(x["type"], x["author"], x["text"]) for x in rev["revisions"]} == {
        ("ins", "Kiss Ödön", "Őszinte űrlap"), ("del", "Kiss Ödön", "Második bekezdés")}
    assert rev["by_author"] == {"Kiss Ödön": {"ins": 1, "del": 1}}

    ch = _ok(client, "wordcraft", {"file_id": r["file_id"], "read": "changes", "base_file_id": docx})
    assert ch["summary"]["new_tracked_changes_by"] == {"Kiss Ödön": 2} and ch["summary"]["untracked_edits"] == 0
    p = ch["paragraphs"][0]
    assert (p["before"], p["after"]) == ("Második bekezdés, ő és ű betűkkel.", "Őszinte űrlap, ő és ű betűkkel.")
    assert p["diff"] == "[-Második bekezdés-]{+Őszinte űrlap+}, ő és ű betűkkel."


# ---------------------------------------------------------------- táblázat: képlet → kiszámolt érték

def test_sheet_formulas_computed(client):
    _need("sheet")
    r = _ok(client, "office", {"app": "sheet", "commands": [{"command": "range.setValues", "params": {"range": "A1", "values": [
        ["Tétel", "Összeg", "Felirat"], ["Őrs", 12, '=A2&" – "&B2'], ["Űr", 30, None], ["Összesen", "=SUM(B2:B3)", "=AVERAGE(B2:B3)"]]}}]})
    fid = r["file_id"]
    assert fid.endswith(".xlsx") and r["steps"][0]["ok"]
    got = _ok(client, "office", {"app": "sheet", "file_id": fid, "commands": [
        {"command": "cell.get", "params": {"cell": "B4"}}, {"command": "cell.get", "params": {"cell": "C2"}},
        {"command": "cell.get", "params": {"cell": "C4"}}]})
    assert got["read_only"] is True and got["file_id"] == fid                  # csak olvasás: nincs új változat
    b4, c2, c4 = (x["result"] for x in got["results"])
    assert (b4["formula"], b4["value"]) == ("=SUM(B2:B3)", 42.0)
    assert c2["value"] == "Őrs – 12" and c4["value"] == 21.0
    csv = _ok(client, "office", {"app": "sheet", "file_id": fid, "read": "text"})["text"]
    assert csv.splitlines()[3].startswith("Összesen,42")
    from openpyxl import load_workbook
    ws = load_workbook(_out(fid), data_only=True).active                     # a mentett fájlban is ott az érték
    assert [c.value for c in ws[4]][:2] == ["Összesen", 42] and ws["C2"].value == "Őrs – 12"
    assert load_workbook(_out(fid)).active["B4"].value == "=SUM(B2:B3)"


# ---------------------------------------------------------------- darabolt feltöltés UTF-8 határon

def test_upload_chunk_split_inside_multibyte_chars(client):
    text = "Árvíztűrő tükörfúrógép — őszi űrlap, ő és ű"
    data = text.encode("utf-8")
    c1 = data.index("ű".encode()) + 1                   # az „ű” két bájtja közé
    c2 = data.index("ő".encode(), c1) + 1               # és egy „ő” közepébe
    c3 = data.index("—".encode()) + 2                   # a hárombájtos gondolatjel közepébe
    cuts = sorted([c1, c2, c3])
    parts = [data[:cuts[0]], data[cuts[0]:cuts[1]], data[cuts[1]:cuts[2]], data[cuts[2]:]]
    for p in parts:                                     # egyik darab sem érvényes UTF-8 magában
        with pytest.raises(UnicodeDecodeError):
            p.decode("utf-8")
    uid = f"qa-{time.time_ns()}"
    order = [2, 0, 3, 1]                                # a sorrend tetszőleges
    for k, i in enumerate(order[:-1]):
        r = _ok(client, "upload_chunk", {"upload_id": uid, "chunk_index": i, "total_chunks": 4, "filename": "őszi jegyzet.txt",
                                         "content_base64": base64.b64encode(parts[i]).decode()})
        assert r["pending"] is True and r["chunks_received"] == k + 1
        assert r["received_bytes"] == sum(len(parts[j]) for j in order[:k + 1])
    r = _ok(client, "upload_chunk", {"upload_id": uid, "chunk_index": order[-1], "total_chunks": 4,
                                     "filename": "őszi jegyzet.txt", "content_base64": base64.b64encode(parts[order[-1]]).decode()})
    assert r["chunks_assembled"] == 4 and r["original_name"] == r["name"] == "őszi jegyzet.txt"
    assert _out(r["file_id"]).read_bytes() == data and _out(r["file_id"]).read_text("utf-8") == text
    err, msg = _call(client, "upload_chunk", {"upload_id": uid, "chunk_index": 4, "total_chunks": 4, "filename": "x.txt",
                                              "content_base64": "eA=="})
    assert err and "chunk_index out of range" in msg


# ---------------------------------------------------------------- jobs: import → convert → export

def test_jobs_pipeline_import_convert_export(client):
    src = "# Őszi űrlap\n\nÁrvíztűrő tükörfúrógép, ő és ű.\n"
    tasks = {"be": {"operation": "import/base64", "file": base64.b64encode(src.encode()).decode(), "filename": "őszi.md"},
             "word": {"operation": "convert", "input": "be", "output_format": "docx"},
             "ki": {"operation": "export/url", "input": "word"}}
    r = _ok(client, "jobs", {"action": "create", "tasks": tasks, "wait": 30})
    assert r["status"] == "finished"
    assert [(t["name"], t["operation"], t["status"]) for t in r["tasks"]] == [
        ("be", "import/base64", "finished"), ("word", "convert", "finished"), ("ki", "export/url", "finished")]
    fids = {f["file_id"] for f in r["files"]}
    assert len(fids) == 1 and next(iter(fids)).endswith(".docx")
    assert any(f.get("url") for f in r["files"]) and all(f["label"] == "őszi.docx" for f in r["files"])
    paras = _docx_paras(next(iter(fids)))
    assert ("Heading 1", "Őszi űrlap") in paras and any(t == "Árvíztűrő tükörfúrógép, ő és ű." for _, t in paras)
    st = _ok(client, "jobs", {"action": "status", "job_id": r["job_id"], "wait": 1})
    assert st["status"] == "finished" and st["job_id"] == r["job_id"]


# ---------------------------------------------------------------- units: időzóna a nyári időszámítás váltásán át

def test_units_timezone_across_dst_switch(client):
    def conv(t, src, dst):
        return _ok(client, "units", {"from_zone": src, "to_zones": dst, "time": t})
    # tavasz (03-29): 00:59Z még 01:59 CET, egy perccel később már 03:00 CEST
    a, b = conv("2026-03-29 00:59", "UTC", ["Europe/Budapest"]), conv("2026-03-29 01:00", "UTC", ["Europe/Budapest"])
    assert (a["to"][0]["time"], a["to"][0]["offset"], a["to"][0]["dst"]) == ("2026-03-29T01:59+01:00", "UTC+01:00", False)
    assert (b["to"][0]["time"], b["to"][0]["offset"], b["to"][0]["dst"]) == ("2026-03-29T03:00+02:00", "UTC+02:00", True)
    # ősz (10-25): az óra visszaáll — 00:59Z = 02:59 CEST, 01:00Z = 02:00 CET
    a, b = conv("2026-10-25 00:59", "UTC", ["Europe/Budapest"]), conv("2026-10-25 01:00", "UTC", ["Europe/Budapest"])
    assert (a["to"][0]["time"], a["to"][0]["abbr"]) == ("2026-10-25T02:59+02:00", "CEST")
    assert (b["to"][0]["time"], b["to"][0]["abbr"]) == ("2026-10-25T02:00+01:00", "CET")
    # helyi idő a váltás után: 03:30 CEST = 01:30 UTC; New York ekkor már nyári időben (EDT)
    r = conv("2026-03-29 03:30", "Europe/Budapest", ["UTC", "America/New_York"])
    assert r["from"]["time"] == "2026-03-29T03:30+02:00"
    assert [x["time"] for x in r["to"]] == ["2026-03-29T01:30+00:00", "2026-03-28T21:30-04:00"]
    # a két váltás közti héten Budapest–New York 5 óra (nem 6)
    r = conv("2026-03-20 12:00", "Europe/Budapest", ["America/New_York"])
    assert (r["to"][0]["time"], r["to"][0]["abbr"], r["to"][0]["dst"]) == ("2026-03-20T07:00-04:00", "EDT", True)


# ---------------------------------------------------------------- pptx_build és a saját check-je

def test_pptx_build_with_rendered_check(client):
    _need("deck")
    spec = {"title": "Őszi terv", "slides": [
        {"layout": "cover", "title": "Őszi űrlap-terv", "subtitle": HU},
        {"layout": "bullets", "title": "Fő pontok ő", "bullets": ["Első ű", "Második ő", "Harmadik"]},
        {"layout": "chart", "title": "Bevétel", "chart_type": "column", "categories": ["Ősz", "Tél"],
         "series": [{"name": "Összeg", "values": [3, 5]}]},
        {"layout": "closing", "title": "Köszönöm ű"}]}
    r = _ok(client, "pptx_build", {"spec": spec, "images": "none"})
    assert r["slides"] == 4 and r["name"] == "Őszi terv.pptx" and r["images"] == []
    chk = r["check"]
    assert chk["status"] == "ok" and chk["findings"] == []
    assert chk["verified"][0] == "4 slides, 0 pictures."
    assert any(v.startswith("Rendered and checked like validate_document") for v in chk["verified"])
    from pptx import Presentation
    prs = Presentation(str(_out(r["file_id"])))
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert "Őszi űrlap-terv" in texts[0] and HU in texts[0] and "Második ő" in texts[1] and "Köszönöm ű" in texts[3]
    chart = next(sh.chart for sh in prs.slides[2].shapes if sh.has_chart)
    assert list(chart.plots[0].categories) == ["Ősz", "Tél"] and list(chart.series[0].values) == [3.0, 5.0]


# ---------------------------------------------------------------- design: oklevél és auto

def test_design_certificate(client):
    _need("vectorcraft")
    r = _ok(client, "design", {"template": "certificate", "data": {
        "recipient": "Kőváry Ödön", "text": "az őszi űrhajós tanfolyam elvégzéséért", "date": "2026. október 18.",
        "signatures": [{"name": "Űrmös Éva", "role": "igazgató"}]}})
    exts = sorted(f["file_id"].rsplit(".", 1)[1] for f in r["files"])
    assert exts == ["pdf", "png"] and r["file_id"].endswith(".pdf") and r["template"] == "certificate"
    assert all(s["ok"] for s in r["steps"]) and r["check"]["status"] in ("ok", "needs_check")
    created = [s["params"]["text"] for s in r["steps"] if s["command"] == "text.create"]
    assert {"Kőváry Ödön", "az őszi űrhajós tanfolyam elvégzéséért", "Űrmös Éva", "igazgató"} <= set(created)
    t = _pdf_text(r["file_id"])
    assert "Kőváry Ödön" in t and "őszi űrhajós" in t and "Űrmös Éva" in t
    import pymupdf
    with pymupdf.open(str(_out(r["file_id"]))) as d:
        assert d.page_count == 1 and d[0].rect.width > d[0].rect.height             # fekvő
    from PIL import Image
    png = next(f["file_id"] for f in r["files"] if f["file_id"].endswith(".png"))
    with Image.open(_out(png)) as im:
        assert im.width > im.height


def test_design_auto_variants_from_the_planner(client, monkeypatch):
    _need("vectorcraft")
    from app.config import settings
    from app.llm import siliconflow
    monkeypatch.setattr(settings, "SILICONFLOW_API_KEY", "teszt")
    asked = []

    async def fake_model(model, messages, thinking=True, json_mode=True):       # az art director (LLM) kész terve
        asked.append(messages[-1]["content"])
        return json.dumps({"variants": [
            {"name": "Éjszakai fények", "archetype": "typographic", "fonts": "condensed",
             "palette": {"bg": "#14213D", "ink": "#FFFFFF", "accent": "#FCA311"}, "title_case": "upper", "title_scale": 1.1,
             "align": "left", "hero": "title", "decoration": "rule", "density": "balanced", "why": "Erős, éjszakai hangulat."},
            {"name": "Őszi keret", "archetype": "framed", "fonts": "editorial",
             "palette": {"bg": "#F7F4EE", "ink": "#1B1B1F", "accent": "#A62E38"}, "title_case": "as_is", "title_scale": 0.9,
             "align": "center", "hero": "date", "decoration": "frame", "density": "airy", "why": "Nyugodt, ünnepi."}]},
            ensure_ascii=False)
    monkeypatch.setattr(siliconflow, "_call_model", fake_model)
    r = _ok(client, "design", {"template": "auto", "variants": 2,
                               "data": {"title": "Őszi koncert a Főtéren", "date": "október 25.", "place": "Győr"}})
    assert len(asked) == 1 and "Őszi koncert a Főtéren" in asked[0]
    assert r["size"] == "ig_post" and r["ai"] is True
    v = r["variants"]
    assert [(x["name"], x["archetype"], x["fonts"]) for x in v] == [("Éjszakai fények", "typographic", "condensed"),
                                                                    ("Őszi keret", "framed", "editorial")]
    assert v[0]["palette"]["bg"] == "#14213D" and v[1]["why"] == "Nyugodt, ünnepi."
    for x in v:
        doc = json.loads(_out(x["file_id"]).read_text("utf-8"))
        assert x["file_id"].endswith(".vectorcraft") and doc["document"]["title"] == "Őszi koncert a Főtéren"
        assert _out(x["preview_file_id"]).read_bytes()[:4] == b"\x89PNG"
        assert {f["label"] for f in x["files"]} == {f"Őszi koncert a Főtéren – {x['name']} (Instagram post).vectorcraft",
                                                    f"Őszi koncert a Főtéren – {x['name']} (Instagram post).png"}


# ---------------------------------------------------------------- diagram: generate_chart_image, chart_redraw

def _pixels(fid: str, hex_color: str) -> int:
    from PIL import Image
    rgb = tuple(int(hex_color[i:i + 2], 16) for i in (1, 3, 5))
    with Image.open(_out(fid)) as im:
        return sum(1 for p in im.convert("RGB").getdata() if p == rgb)


def test_generate_chart_image_png_and_svg(client):
    r = _ok(client, "generate_chart_image", {"chart_type": "column", "categories": ["Ősz", "Tél", "Űrhét"],
                                             "series": [{"name": "Bevétel ű", "values": [3, 5, 4]}], "title": "Őszi bevétel",
                                             "palette": ["#C0392B"], "also_svg": True})
    assert r["chart_type"] == "column" and r["image_file_id"].endswith(".png") and r["svg_file_id"].endswith(".svg")
    assert _out(r["image_file_id"]).read_bytes()[:4] == b"\x89PNG" and _pixels(r["image_file_id"], "#C0392B") > 5000
    svg = _out(r["svg_file_id"]).read_text("utf-8")
    assert "<!-- Őszi bevétel -->" in svg and "<!-- Űrhét -->" in svg and "#c0392b" in svg.lower()
    err, msg = _call(client, "generate_chart_image", {"chart_type": "radar", "categories": ["a"], "series": [{"values": [1]}]})
    assert err and "unsupported chart_type" in msg


def test_chart_redraw_reads_then_redraws_with_palette(client, samples, put_file, monkeypatch):
    seen = _vision(monkeypatch, json.dumps({
        "chart_type": "column", "title": "Őszi bevétel", "x_label": "Évszak", "y_label": "Millió Ft", "categories": ["Ősz", "Tél"],
        "series": [{"name": "Bevétel ű", "values": [3, 5]}], "value_format": "integer", "notes": ""}, ensure_ascii=False))
    src = put_file(samples.png())
    r = _ok(client, "chart_redraw", {"image_file_id": src, "palette": ["#112233", "#AA5500"], "also_svg": True,
                                     "title_override": "Új cím: őszi űrlap"})
    assert len(seen) == 1 and "extract its structure as STRICT JSON" in seen[0]
    s = r["extracted_structure"]
    assert (s["title"], s["categories"], s["series"][0]["name"]) == ("Őszi bevétel", ["Ősz", "Tél"], "Bevétel ű")
    assert r["chart_type"] == "column" and r["image_file_id"] != src
    assert _pixels(r["image_file_id"], "#112233") > 5000                     # az új paletta színe
    svg = _out(r["svg_file_id"]).read_text("utf-8")
    assert "<!-- Új cím: őszi űrlap -->" in svg and "<!-- Évszak -->" in svg and "<!-- Millió Ft -->" in svg


# ---------------------------------------------------------------- képolvasás: ocr_image, extract_table_from_image, image_to_xlsx

TABLE = {"headers": ["Név", "Összeg"], "rows": [["Kőváry Ödön", "12"], ["Űrmös Éva", "30"]]}


def test_ocr_image(client, samples, put_file, monkeypatch):
    seen = _vision(monkeypatch, "Árvíztűrő tükörfúrógép\nŐszi űrlap")
    img = put_file(samples.jpg())
    r = _ok(client, "ocr_image", {"image_file_id": img, "language_hint": "Hungarian"})
    assert (r["image_file_id"], r["text"]) == (img, "Árvíztűrő tükörfúrógép\nŐszi űrlap")
    assert len(seen) == 1 and "The text appears to be in Hungarian." in seen[0]
    err, msg = _call(client, "ocr_image", {"image_file_id": put_file(samples.docx())})
    assert err and "not an image" in msg


def test_extract_table_and_image_to_xlsx(client, samples, put_file, monkeypatch):
    # a modell gyakran kódkerítésben adja a JSON-t — a feldolgozó leszedi
    _vision(monkeypatch, "```json\n" + json.dumps(TABLE, ensure_ascii=False) + "\n```")
    img = put_file(samples.png())
    r = _ok(client, "extract_table_from_image", {"image_file_id": img})
    assert (r["headers"], r["rows"], r["row_count"]) == (TABLE["headers"], TABLE["rows"], 2)

    r = _ok(client, "image_to_xlsx", {"image_file_id": img, "sheet_name": "Őszi tábla"})
    assert r["file_id"].endswith(".xlsx") and r["headers"] == ["Név", "Összeg"] and r["row_count"] == 2
    from openpyxl import load_workbook
    ws = load_workbook(_out(r["file_id"])).active
    assert ws.title == "Őszi tábla" and ws["A1"].font.bold
    assert [[c.value for c in row] for row in ws.iter_rows()] == [["Név", "Összeg"], ["Kőváry Ödön", "12"], ["Űrmös Éva", "30"]]


# ---------------------------------------------------------------- digitize

def test_digitize_page_photo_to_word_and_text(client, samples, put_file, monkeypatch, tmp_path):
    _need("pandoc")
    import tempfile
    from app.tools import digitize as dz
    # a qa-képben a /tmp/aice_book_reference.docx root tulajdonú (régi futásból) — más felhasználóként nem írható;
    # a könyv-referencia ezért a teszt saját ideiglenes mappájába készül
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    monkeypatch.setattr(dz, "_REF_DOCX", None)
    page = "# Őszi fejezet\n\nA tükörfúrógép árvíztűrő, az űrlap kész.\n\nMásodik bekezdés ő és ű."
    seen = _vision(monkeypatch, page)
    img = put_file(samples.jpg("régi lap ő.jpg"))
    job = _ok(client, "digitize", {"file_ids": [img], "outputs": ["docx", "txt"], "language": "Hungarian"})
    assert job["status"] == "running" and job["kind"] == "digitize"
    res = _wait(client, "digitize", job)
    out = res["result"]
    assert len(seen) == 1 and "Hungarian" in seen[0]
    assert (out["pages"], out["words"], out["readings"]["uncertain"], out["readings"]["illegible"]) == (1, 13, 0, 0)
    labels = {f["label"]: f["file_id"] for f in out["files"]}
    assert set(labels) == {"régi lap ő.docx", "régi lap ő.txt"}
    assert _docx_paras(labels["régi lap ő.docx"]) == [("Heading 1", "Őszi fejezet"),
                                                      ("First Paragraph", "A tükörfúrógép árvíztűrő, az űrlap kész."),
                                                      ("Body Text", "Második bekezdés ő és ű.")]
    assert _out(labels["régi lap ő.txt"]).read_text("utf-8") == \
        "Őszi fejezet\n\nA tükörfúrógép árvíztűrő, az űrlap kész.\n\nMásodik bekezdés ő és ű.\n"
    assert res["check"]["status"] == "ok" and "régi lap ő.docx: 13 words read." in res["check"]["verified"]


# ---------------------------------------------------------------- office pdf: kitakarás (a SZÖVEGRÉTEGBEN)

def test_office_pdf_redact_removes_text_layer(client, samples, put_file):
    _need("pdf")
    pdf = put_file(samples.pdf(pages=2))
    before = _pdf_pages(pdf)
    assert all("Árvíztűrő" in p for p in before)
    r = _ok(client, "office", {"app": "pdf", "file_id": pdf, "commands": [
        {"tool": "redact_mark", "args": {"find": "Árvíztűrő"}}, {"tool": "redact_apply", "args": {}}]})
    assert [s["result"] for s in r["steps"]] == [{"marked": 2, "pages_marked": [1, 2], "pages": 2}, {"applied": 2, "pages": 2}]
    after = _pdf_pages(r["file_id"])
    assert len(after) == 2 and not any("Árvíztűrő" in p for p in after)              # a szövegrétegből is eltűnt
    assert all("tükörfúrógép — őszi űrlap" in p for p in after)                     # a többi szöveg marad
    assert after[0].startswith("1. oldal") and after[1].startswith("2. oldal")
    assert r["check"]["status"] == "ok"
    assert "“Árvíztűrő”: 2 in the original → 0 left in the text layer." in r["check"]["verified"]
    assert r["name"] == "dok – redacted.pdf"


# ---------------------------------------------------------------- videó (rövid)

def test_video_make_short_from_scene_plan(client):
    _need("ffmpeg", "ffprobe")
    from app.tools import video
    if not video.available():
        pytest.skip("the video engine (EffectCraft) is only in the production image")
    scenes = [{"headline": "Őszi terv", "caption": "Győrben 12 új fát ültetnek ősszel."},
              {"headline": "Űrlap", "caption": "A jelentkezés az űrlapon történik."}]
    job = _ok(client, "video", {"action": "make", "text": "Győrben 12 új fát ültetnek ősszel. A jelentkezés az űrlapon történik.",
                                "scenes": scenes, "title": "Őszi faültetés", "format": "square", "duration": 12,
                                "pictures": "none", "narration": "none"})
    res = _wait(client, "video", job)["result"]
    assert res["duration"]["requested"] == 12 and res["duration"]["mode"] == "fixed" and abs(res["duration"]["actual"] - 12) <= 0.6
    assert res["check"]["status"] == "ok" and "Your scene plan: 2 scenes (s1, s2), used as given." in res["check"]["verified"]
    assert any(v.startswith("Length as asked: 12 s") for v in res["check"]["verified"])
    assert [(s["id"], s["headline"], s["caption"]) for s in res["story"]["scenes"]] == [
        ("s1", "Őszi terv", "Győrben 12 új fát ültetnek ősszel."), ("s2", "Űrlap", "A jelentkezés az űrlapon történik.")]
    assert res["project"]["label"] == "Őszi faültetés.ecproj" and res["video"]["label"] == "Őszi faültetés.mp4"
    probe = json.loads(subprocess.run(["ffprobe", "-v", "error", "-print_format", "json", "-show_format", "-show_streams",
                                       str(_out(res["video"]["file_id"]))], capture_output=True, text=True, check=True).stdout)
    v = next(s for s in probe["streams"] if s["codec_type"] == "video")
    assert (v["width"], v["height"]) == (1080, 1080) and abs(float(probe["format"]["duration"]) - 12) <= 0.6


# ---------------------------------------------------------------- photo_edit (saját terv, modell nélkül)

def test_photo_edit_with_explicit_plan(client, samples, put_file):
    _need("photocraft")
    img = put_file(samples.jpg("őszi fotó.jpg"))                # 640×480
    plan = [{"op": "crop", "aspect": "4:5"}, {"op": "color", "warmth": 30},
            {"op": "text", "text": "Őszi akció ű", "position": "top", "size": "large"}]
    job = _ok(client, "photo_edit", {"file_ids": [img], "plan": plan})
    res = _wait(client, "photo_edit", job)["result"]
    assert res["summary"] == "custom plan" and res["photos"] == 1
    assert [s["op"] for s in res["plan"]] == ["crop", "color", "text"] and res["plan"][2]["text"] == "Őszi akció ű"
    psd = next(f for f in res["files"] if f["file_id"].endswith(".psd"))
    jpg = next(f for f in res["files"] if f["file_id"].endswith(".jpg"))
    assert jpg["name"] == "őszi fotó (edited).jpg"
    from PIL import Image
    with Image.open(_out(jpg["file_id"])) as im:
        assert im.size == (384, 480)                               # 4:5 kivágás
    info = _ok(client, "photocraft", {"file_id": psd["file_id"], "read": "info"})
    layers = {(x["name"], x["kind"]) for x in info["layers"]}
    assert ("Őszi akció ű", "Type") in layers and ("Warmer", "Adjustment") in layers and (info["width"], info["height"]) == (384, 480)


# ---------------------------------------------------------------- restore_photos (az AI-lépés kész képpel)

def test_restore_photos_layers_and_jpg(client, samples, put_file, monkeypatch):
    _need("photocraft")
    from PIL import Image, ImageFilter
    from app.tools import restore
    prompts = []

    async def fake_kontext(src, prompt):                        # a FLUX Kontext hívás helyett: „helyreállított” kép
        prompts.append(prompt)
        with Image.open(src) as im:
            out = im.convert("RGB").filter(ImageFilter.SMOOTH_MORE)
        buf = io.BytesIO()
        out.save(buf, "PNG")
        return buf.getvalue()
    monkeypatch.setattr(restore, "_kontext", fake_kontext)
    img = put_file(samples.jpg("nagymama őszi fotója.jpg"))
    job = _ok(client, "restore_photos", {"file_ids": [img], "colorize": False, "split_prints": False})
    done = _wait(client, "restore_photos", job)
    res = done["result"]
    assert len(prompts) == 1 and res["photos"] == 1                 # színezés nélkül egy AI-lépés
    labels = {f["label"]: f for f in res["files"]}
    assert set(labels) == {"nagymama őszi fotója (restored).jpg", "nagymama őszi fotója (restored, layers).psd"}
    jpg = labels["nagymama őszi fotója (restored).jpg"]
    assert jpg["colorized"] is False and jpg["aligned_points"] > 0
    with Image.open(_out(jpg["file_id"])) as im:
        assert im.size == (640, 480)
    info = _ok(client, "photocraft", {"file_id": labels["nagymama őszi fotója (restored, layers).psd"]["file_id"], "read": "info"})
    names = [x["name"] for x in info["layers"]]
    assert "Restored (AI)" in names and len(names) == 2 and "Colour (AI)" not in names
    assert done["check"]["status"] in ("ok", "needs_check")
    assert "1 photo found and restored; each has a layered PSD (original, restored) and a JPG." in done["check"]["verified"]
