"""Rust-motoraink (DeckCraft, GridCraft, PhotoCraft, VectorCraft) a konverziós gráfban: ahol a mérés szerint
jobbak, ott ők futnak; ha a kért opciót nem tudják, a következő legjobb motor jön. A valódi átalakítás csak ott
fut, ahol a Craft-CLI-k telepítve vannak (a production képben)."""
from __future__ import annotations

import asyncio
import subprocess
import zipfile
from pathlib import Path

import pytest


@pytest.fixture()
def all_engines(monkeypatch):
    """Minden motor „telepítve” — a gráf döntéseit nézzük, nem a gépet."""
    from app.convert import engines, graph
    monkeypatch.setattr(engines, "available", lambda e: True)
    graph._adjacency.cache_clear()
    graph._reach.cache_clear()
    graph._targets.cache_clear()
    yield graph
    graph._adjacency.cache_clear()
    graph._reach.cache_clear()
    graph._targets.cache_clear()


def _route(graph, s, d, **kw):
    p = graph.plan(s, d, **kw)
    return [e.engine for e in p] if p else None


def test_rust_crafts_win_where_they_are_faster(all_engines):
    g = all_engines
    assert _route(g, "pptx", "pdf") == ["deckcraft"]                  # 0,08 s vs LibreOffice 1,6 s
    assert _route(g, "pptx", "png") == ["deckcraft"]
    assert _route(g, "xlsx", "csv") == ["gridcraft"] and _route(g, "csv", "xlsx") == ["gridcraft"]
    assert _route(g, "xlsx", "json") == ["python"]                     # a GridCraft JSON-ja a saját formátuma
    assert _route(g, "png", "webp") == ["photocraft"] and _route(g, "psd", "png") == ["photocraft"]
    assert _route(g, "png", "jpg") == ["pillow"]                      # ide a mérés szerint nem kell
    assert _route(g, "gif", "webp") == ["pillow"]                     # animált forrás: nem a PhotoCraft
    assert _route(g, "svg", "png") == ["vectorcraft"] and _route(g, "dxf", "svg") == ["ezdxf"]     # DXF: az ezdxf skálázik jól


def test_options_a_craft_cannot_honour_route_around_it(all_engines):
    from app.convert import service
    g = all_engines
    assert service.craft_avoid({}, "pdf") == set()
    assert service.craft_avoid({"page_range": "1-2"}, "pdf") == {"gridcraft", "photocraft", "vectorcraft"}
    assert "deckcraft" in service.craft_avoid({"page_range": "1-2"}, "png")
    assert "photocraft" in service.craft_avoid({"width": 640}, "webp")
    assert _route(g, "png", "webp", avoid=service.craft_avoid({"quality": 70}, "webp")) == ["pillow"]
    assert _route(g, "pptx", "pdf", avoid=service.craft_avoid({"page_range": "2"}, "pdf")) == ["deckcraft"]
    assert _route(g, "svg", "png", avoid=service.craft_avoid({"dpi": 300}, "png")) == ["rsvg"]
    assert _route(g, "xlsx", "csv", avoid=service.craft_avoid({"sheet": "2"}, "csv")) == ["python"]


def _has(cli: str) -> bool:
    from app.tools import crafts
    return bool({"deck": lambda: crafts.office_cli("deck"), "sheet": lambda: crafts.office_cli("sheet"),
                 "photo": crafts.photo_cli, "vector": crafts.vector_cli}[cli]())


needs_crafts = pytest.mark.skipif(not all(_has(c) for c in ("deck", "sheet", "photo", "vector")),
                                  reason="the Rust Craft CLIs are only in the production image")


@needs_crafts
def test_real_rust_conversions_keep_accents(samples, tmp_path):
    from app.convert import service
    from app.convert import graph
    graph.reset_cache()

    def conv(src, tgt):
        outs, used = asyncio.run(service.convert_path(src, tgt, {}, tmp_path / f"w-{src.suffix[1:]}-{tgt}"))
        return outs, used

    outs, used = conv(samples.pptx(), "pdf")
    assert used == ["deckcraft"]
    text = subprocess.run(["pdftotext", str(outs[0]), "-"], capture_output=True, text=True).stdout
    assert "Árvíztűrő" in text and "ű" in text
    outs, used = conv(samples.pptx(), "png")
    assert used == ["deckcraft"] and outs and all(p.read_bytes()[:4] == b"\x89PNG" for p in outs)
    outs, used = conv(samples.xlsx(), "csv")
    assert used == ["gridcraft"] and "Őrs" in outs[0].read_text("utf-8")
    outs, used = conv(samples.csv_rows(), "xlsx")
    from openpyxl import load_workbook
    rows = [[c.value for c in r] for r in load_workbook(outs[0]).active.iter_rows()]
    assert used == ["gridcraft"] and rows[1][0] == "Kőváry Ödön"
    outs, used = conv(samples.png(size=(320, 200)), "webp")
    assert used == ["photocraft"] and outs[0].read_bytes()[8:12] == b"WEBP"
    outs, used = conv(samples.svg(), "png")
    assert used == ["vectorcraft"] and outs[0].read_bytes()[:4] == b"\x89PNG"
    outs, used = asyncio.run(service.convert_path(samples.png(size=(320, 200)), "webp", {"width": 100},
                                                  tmp_path / "w-opt"))
    from PIL import Image
    assert used != ["photocraft"] and 98 <= Image.open(outs[0]).width <= 100   # a kért méretet a Pillow tudja
