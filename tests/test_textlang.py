"""A docx nyelve a szöveg nyelve: magyar levélen ne angol helyesírás-ellenőrzés fusson.

UI-próba (10-18): a briefből írt magyar levél w:lang="en-US" volt → a szerkesztőben szinte minden szó piros.
"""
from __future__ import annotations

import re
import zipfile

import docx

from app.tools import ooxml_safe
from app.tools.textlang import detect

HU = ("Tájékoztatjuk Önöket, hogy a társasház udvarán felújítási munkálatok kezdődnek. A felújítás november 3. és "
      "november 14. között tart. A munkálatok ideje alatt az udvaron nem lehet parkolni, kérjük, hogy az utcán parkoljanak.")
EN = ("We are writing to let you know that the courtyard of the building is being renovated. The work starts on "
      "3 November and it will take until 14 November. Please park your cars on the street during this time.")


def test_detect_is_sure_or_silent():
    assert detect(HU) == "hu-HU" and detect(EN) == "en-US"
    assert detect("Negyedév Árbevétel Változás Q1 420 +4%") is None          # kevés szó: nem tippel


def _lang(path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return re.findall(r'<w:lang w:val="([^"]+)"', z.read("word/styles.xml").decode())


def test_docx_default_language_follows_text(tmp_path):
    p = tmp_path / "level.docx"
    d = docx.Document()
    for _ in range(2):
        d.add_paragraph(HU)
    d.save(p)
    ooxml_safe.fix(p)                     # (a python-docx mentése az install() óta maga is átmegy ezen)
    assert "hu-HU" in _lang(p)
    assert "word/styles.xml" not in ooxml_safe.fix(p)                      # másodszor már nincs dolga


def test_english_and_explicit_languages_untouched(tmp_path):
    p = tmp_path / "en.docx"
    d = docx.Document()
    d.add_paragraph(EN)
    d.save(p)
    before = _lang(p)
    ooxml_safe.fix(p)
    assert _lang(p) == before


def test_pptx_language_follows_text(tmp_path):
    from pptx import Presentation
    from app.tools import keynote_safe
    p = tmp_path / "pakli.pptx"
    prs = Presentation()
    for _ in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Közösségi napelempark"
        s.placeholders[1].text = HU
    prs.save(p)
    keynote_safe.fix(p)                   # (a python-pptx mentése is átmehet rajta — a végeredmény számít)
    with zipfile.ZipFile(p) as z:
        xml = b"".join(z.read(n) for n in z.namelist() if n.startswith("ppt/") and n.endswith(".xml"))
    assert b'lang="hu-HU"' in xml and b'lang="en-US"' not in xml
    assert keynote_safe.fix(p) == []


def test_certificate_lead_in_the_content_language():
    """UI-próba 10-18: magyar oklevélen angol „This certificate is presented to”."""
    from app.tools.design import _certificate_lead
    assert _certificate_lead({"title": "Oklevél a könyvtári olvasóverseny első helyéért", "recipient": "Kovács Eszter"}) \
        == "Ezt az oklevelet kapja"
    assert _certificate_lead({"title": "Certificate of Completion", "recipient": "Jane Doe"}) == "This certificate is presented to"
