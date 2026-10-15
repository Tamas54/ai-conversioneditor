"""2. csoport: eszközök (tömörítés, PDF-csomag, kép, videó/hang, vízjel, web, fájlok, metaadat, haladó, segédek).

Minden eszköz az `app.ops.run_tool`-on át fut, ahogy a web, a REST API, az MCP és az ügynök is hívja; a kimenetet
megnyitjuk és ellenőrizzük (oldalszám, méret, időtartam, szöveg — ékezetekkel)."""
from __future__ import annotations

import json
import subprocess
import zipfile
from pathlib import Path

import pytest

from samples import HU

pytestmark = pytest.mark.asyncio


def _out(fid: str) -> Path:
    from app.storage import output_path
    return output_path(fid)


async def run(key: str, fids: list[str], **opts):
    from app.ops import run_tool
    return await run_tool(key, fids, opts, charge=False)


def _pdf(fid: str):
    import pymupdf
    return pymupdf.open(str(_out(fid)))


def _probe(p: Path) -> dict:
    out = subprocess.run(["ffprobe", "-v", "error", "-print_format", "json", "-show_format", "-show_streams", str(p)],
                         capture_output=True, text=True, check=True).stdout
    return json.loads(out)


def _dur(p: Path) -> float:
    return float(_probe(p)["format"]["duration"])


def _img(fid: str):
    from PIL import Image
    return Image.open(_out(fid))


def _engine(name: str):
    from app.convert import engines
    if not engines.available(name):
        pytest.skip(f"{name} not installed")


# ---------------------------------------------------------------- PDF

async def test_pdf_compress_merge_split(samples, put_file):
    _engine("ghostscript")
    pdf = put_file(samples.pdf(pages=3))
    r = await run("compress-pdf", [pdf], compression="extreme")
    with _pdf(r["file_id"]) as d:
        assert d.page_count == 3 and "tükörfúrógép" in d[0].get_text()
    r = await run("merge-pdf", [pdf, put_file(samples.png())], page_ranges=["2-3", "all"])
    with _pdf(r["file_id"]) as d:
        assert d.page_count == 3 and d[0].get_text().startswith("2. oldal")
    r = await run("split-pdf", [pdf], mode="each")
    assert r["file_id"].endswith(".zip") and len(r["files"]) == 4
    with zipfile.ZipFile(_out(r["file_id"])) as z:
        assert len(z.namelist()) == 3
    r = await run("split-pdf", [pdf], page_range="1,2-3", output="files")
    assert len(r["files"]) == 2


async def test_pdf_pages_rotate_organize_crop_resize(samples, put_file):
    pdf = put_file(samples.pdf(pages=3))
    with _pdf((await run("extract-pages", [pdf], page_range="2"))["file_id"]) as d:
        assert d.page_count == 1 and "2. oldal" in d[0].get_text()
    with _pdf((await run("delete-pdf-pages", [pdf], page_range="1,3"))["file_id"]) as d:
        assert d.page_count == 1 and "2. oldal" in d[0].get_text()
    with pytest.raises(Exception, match="every page"):
        await run("delete-pdf-pages", [pdf], page_range="1-3")
    with _pdf((await run("rotate-pdf", [pdf], angle=90, page_range="1"))["file_id"]) as d:
        assert d[0].rotation == 90 and d[1].rotation == 0
    with _pdf((await run("organize-pdf", [pdf], order="3,1,2"))["file_id"]) as d:
        assert [p.get_text()[:2] for p in d] == ["3.", "1.", "2."]
    with _pdf((await run("crop-pdf", [pdf], margins_mm=[10, 10, 10, 10]))["file_id"]) as d:
        assert d[0].rect.width < 595 - 50
    with _pdf((await run("resize-pdf", [pdf], page_size="A5"))["file_id"]) as d:
        assert round(d[0].rect.width) == 420 and "tükörfúrógép" in d[0].get_text()


async def test_pdf_protect_unlock_permissions(samples, put_file):
    import pikepdf
    pdf = put_file(samples.pdf())
    r = await run("protect-pdf", [pdf], password="titok-ő", allow_print=False, allow_copy=False)
    with pytest.raises(pikepdf.PasswordError):
        pikepdf.open(str(_out(r["file_id"])))
    with pikepdf.open(str(_out(r["file_id"])), password="titok-ő") as p:
        assert p.is_encrypted and not p.allow.print_highres and not p.allow.extract
    r2 = await run("unlock-pdf", [put_file(_out(r["file_id"]))], password="titok-ő")
    with pikepdf.open(str(_out(r2["file_id"]))) as p:
        assert not p.is_encrypted
    with pytest.raises(Exception, match="(?i)password"):
        await run("unlock-pdf", [put_file(_out(r["file_id"]))], password="rossz")


async def test_pdf_pdfa_flatten_images(samples, put_file):
    _engine("ghostscript")
    r = await run("pdf-to-pdfa", [put_file(samples.pdf())], pdfa_level="2b")
    with _pdf(r["file_id"]) as d:
        assert "pdfaid" in (d.get_xml_metadata() or "") and "tükörfúrógép" in d[0].get_text()
    form = put_file(samples.fillable_pdf())
    with _pdf((await run("flatten-pdf", [form]))["file_id"]) as d:
        assert not list(d[0].widgets()) and "őszi tábor" in d[0].get_text()
    merged = await run("merge-pdf", [put_file(samples.png(size=(400, 300))), put_file(samples.jpg())])
    r = await run("extract-pdf-images", [put_file(_out(merged["file_id"]))])
    assert len(r["files"]) >= 2


async def test_pdf_annotate_edit_redact_numbers_compare(samples, put_file):
    pdf = put_file(samples.pdf())
    r = await run("annotate-pdf", [pdf], annotations=[{"type": "highlight", "text": "tükörfúrógép"},
                                                       {"type": "note", "page": 1, "x": 50, "y": 50, "text": "Megjegyzés ő ű"}])
    with _pdf(r["file_id"]) as d:
        kinds = sorted(a.type[1] for a in d[0].annots())
        assert kinds == ["Highlight", "Text"]
    r = await run("edit-pdf", [pdf], edits=[{"type": "whiteout", "page": 1, "rect": [60, 80, 560, 110]},
                                            {"type": "text", "page": 1, "x": 72, "y": 100, "text": "Új szöveg: őrült űrhajó"}])
    with _pdf(r["file_id"]) as d:
        assert "Új szöveg: őrült űrhajó" in d[0].get_text()
    r = await run("redact-pdf", [pdf], words=["Árvíztűrő"])
    with _pdf(r["file_id"]) as d:
        t = d[0].get_text()
        assert "Árvíztűrő" not in t and "tükörfúrógép" in t
    r = await run("page-numbers", [pdf], position="bottom", text_format="{n}. oldal / {total}")
    with _pdf(r["file_id"]) as d:
        assert "2. oldal / 2" in d[1].get_text()
    other = put_file(samples.pdf("masik.pdf", text="Más szöveg ű"))
    with _pdf((await run("compare-pdf", [pdf, other]))["file_id"]) as d:
        assert d.page_count >= 1


async def test_pdf_ocr_layer(samples, put_file, tmp_path):
    _engine("ocrmypdf")
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 48)
    im = Image.new("RGB", (1240, 400), "white")
    ImageDraw.Draw(im).text((60, 150), "Tükörfúrógép és őszi űrlap", font=font, fill="black")
    scan = tmp_path / "szken.pdf"
    im.save(scan, "PDF", resolution=150)
    r = await run("ocr-pdf", [put_file(scan)], ocr_language="hun")
    with _pdf(r["file_id"]) as d:
        assert "őszi" in d[0].get_text()


# ---------------------------------------------------------------- kép

async def test_image_compress_family(samples, put_file, tmp_path):
    jpg = samples.jpg()
    r = await run("compress-image", [put_file(jpg)], quality=40, strip_metadata=True)
    assert _out(r["file_id"]).stat().st_size < jpg.stat().st_size
    with _img(r["file_id"]) as im:
        assert im.format == "JPEG" and not im.getexif()
    r = await run("compress-jpeg", [put_file(jpg)], quality=50, strip_metadata=False)
    from PIL import Image
    with _img(r["file_id"]) as im, Image.open(jpg) as orig:
        assert im.getexif().get(0x0110) == orig.getexif().get(0x0110)   # kérésre a metaadat marad
    r = await run("compress-png", [put_file(samples.png(alpha=True))], colors=16)
    with _img(r["file_id"]) as im:
        assert im.format == "PNG" and len(im.convert("RGBA").getcolors(4096) or []) <= 17
    from PIL import Image
    frames = [Image.new("RGB", (200, 120), (i * 40, 80, 200 - i * 30)) for i in range(5)]
    gif = tmp_path / "anim.gif"
    frames[0].save(gif, save_all=True, append_images=frames[1:], duration=100, loop=0)
    r = await run("compress-gif", [put_file(gif)], colors=32, max_width=100)
    with _img(r["file_id"]) as im:
        assert im.format == "GIF" and im.width == 100 and im.n_frames == 5
    batch = [put_file(samples.png(f"k{i}.png")) for i in range(3)]
    r = await run("compress-image", batch, quality=60)
    assert len(r["files"]) == 3 and all(f["label"].endswith(".png") for f in r["files"])


async def test_image_geometry(samples, put_file):
    png = put_file(samples.png(size=(320, 200)))
    with _img((await run("resize-image", [png], width=160))["file_id"]) as im:
        assert im.size == (160, 100)
    with _img((await run("resize-image", [png], percent=50))["file_id"]) as im:
        assert im.size == (160, 100)
    with _img((await run("crop-image", [png], box=[10, 10, 100, 50]))["file_id"]) as im:     # x, y, szélesség, magasság
        assert im.size == (100, 50)
    with _img((await run("crop-image", [png], aspect="1:1"))["file_id"]) as im:
        assert im.width == im.height == 200
    with _img((await run("rotate-image", [png], angle=90))["file_id"]) as im:
        assert im.size == (200, 320)
    from PIL import Image
    src = Image.open(samples.png(size=(320, 200)))
    with _img((await run("flip-image", [png], direction="vertical"))["file_id"]) as im:
        assert im.convert("RGB").getpixel((100, 199 - 100)) == src.convert("RGB").getpixel((100, 100))
    with _img((await run("upscale-image", [png], scale=2))["file_id"]) as im:
        assert im.size == (640, 400)
    r = await run("collage-maker", [png, put_file(samples.jpg())], columns=2, width=800)
    with _img(r["file_id"]) as im:
        assert im.format == "JPEG" and im.width == 800
    r = await run("color-picker", [png], colors=4)
    assert len(r["data"]["colors"]) <= 4 and all(c["hex"].startswith("#") for c in r["data"]["colors"])


# ---------------------------------------------------------------- videó, hang

async def test_video_tools(samples, put_file, tmp_path):
    _engine("ffmpeg")
    mp4 = put_file(samples.mp4(seconds=2))
    r = await run("compress-video", [mp4], crf=35, resolution="160x120")
    st = [s for s in _probe(_out(r["file_id"]))["streams"] if s["codec_type"] == "video"][0]
    assert (st["width"], st["height"]) == (160, 120)
    r = await run("trim-video", [mp4], start="0.5", end="1.5")
    assert 0.8 <= _dur(_out(r["file_id"])) <= 1.3
    r = await run("crop-video", [mp4], aspect="1:1")
    st = [s for s in _probe(_out(r["file_id"]))["streams"] if s["codec_type"] == "video"][0]
    assert st["width"] == st["height"]
    r = await run("rotate-video", [mp4], angle=90)
    st = [s for s in _probe(_out(r["file_id"]))["streams"] if s["codec_type"] == "video"][0]
    assert (st["width"], st["height"]) == (240, 320)
    r = await run("resize-video", [mp4], resolution="640x360", aspect_mode="pad")
    st = [s for s in _probe(_out(r["file_id"]))["streams"] if s["codec_type"] == "video"][0]
    assert (st["width"], st["height"]) == (640, 360)
    r = await run("mute-video", [mp4])
    assert {s["codec_type"] for s in _probe(_out(r["file_id"]))["streams"]} == {"video"}
    r = await run("video-to-gif", [mp4], fps=8, max_width=160)
    with _img(r["file_id"]) as im:
        assert im.format == "GIF" and im.width == 160 and im.n_frames > 4
    r2 = await run("gif-to-video", [put_file(_out(r["file_id"]))])
    assert r2["file_id"].endswith(".mp4") and _dur(_out(r2["file_id"])) > 1
    r = await run("extract-audio", [mp4], audio_format="mp3")
    assert r["file_id"].endswith(".mp3") and 1.5 <= _dur(_out(r["file_id"])) <= 2.5
    srt = tmp_path / "felirat.srt"
    srt.write_text("1\n00:00:00,000 --> 00:00:01,500\nŐszi űrlap — felirat\n", "utf-8")
    r = await run("add-subtitles", [mp4, put_file(srt)], burn=False, language="hun")
    subs = [s for s in _probe(_out(r["file_id"]))["streams"] if s["codec_type"] == "subtitle"]
    assert subs


async def test_audio_tools(samples, put_file):
    _engine("ffmpeg")
    wav = put_file(samples.wav(seconds=2))
    r = await run("trim-audio", [wav], start="0.5", fade_in=0.2, fade_out=0.2)
    assert 1.3 <= _dur(_out(r["file_id"])) <= 1.7
    r = await run("adjust-volume", [wav], volume="0.5")
    assert r["file_id"].endswith(".wav")
    r = await run("adjust-volume", [wav], normalize=True)
    assert _dur(_out(r["file_id"])) > 1.5
    r = await run("compress-mp3", [wav], audio_bitrate="64k")
    assert r["file_id"].endswith(".mp3")
    r = await run("compress-wav", [put_file(samples.wav("nagy.wav", seconds=1, rate=44100))], sample_rate=16000, bit_depth=16)
    a = [s for s in _probe(_out(r["file_id"]))["streams"] if s["codec_type"] == "audio"][0]
    assert a["sample_rate"] == "16000"


# ---------------------------------------------------------------- vízjel

async def test_watermarks(samples, put_file):
    r = await run("watermark-pdf", [put_file(samples.pdf())], text="BIZALMAS – ő ű", opacity=0.3, rotation=45)
    with _pdf(r["file_id"]) as d:
        assert "BIZALMAS – ő ű" in d[0].get_text() and "BIZALMAS" in d[1].get_text()
    logo = put_file(samples.png("logo.png", size=(80, 40), alpha=True))
    r = await run("watermark-pdf", [put_file(samples.pdf())], image_file=logo, position="bottom-right")
    with _pdf(r["file_id"]) as d:
        assert d[0].get_images()
    png = samples.png(size=(400, 300))
    r = await run("watermark-image", [put_file(png)], text="© Ősz Űr", position="tile", opacity=0.5)
    from PIL import Image, ImageChops
    with _img(r["file_id"]) as im, Image.open(png) as orig:
        assert im.size == orig.size and ImageChops.difference(im.convert("RGB"), orig.convert("RGB")).getbbox()
    _engine("ffmpeg")
    r = await run("watermark-video", [put_file(samples.mp4())], text="Minta ő", position="top-left")
    assert r["file_id"].endswith(".mp4") and _dur(_out(r["file_id"])) > 1.5


# ---------------------------------------------------------------- fájlok

async def test_files_merge_thumbnail_archives(samples, put_file, tmp_path):
    a, b = put_file(samples.pdf("a.pdf", pages=1)), put_file(samples.pdf("b.pdf", pages=2))
    with _pdf((await run("merge-files", [a, b]))["file_id"]) as d:
        assert d.page_count == 3
    _engine("ffmpeg")
    r = await run("merge-files", [put_file(samples.wav("x.wav")), put_file(samples.wav("y.wav"))])
    assert 1.8 <= _dur(_out(r["file_id"])) <= 2.2
    with _img((await run("thumbnail", [put_file(samples.pdf())], width=200))["file_id"]) as im:
        assert im.width <= 200 and im.format == "PNG"
    with _img((await run("thumbnail", [put_file(samples.mp4())], width=160, at="1"))["file_id"]) as im:
        assert im.width == 160
    files = [put_file(samples.txt("első.txt")), put_file(samples.png("kép.png"))]
    r = await run("create-archive", files)
    with zipfile.ZipFile(_out(r["file_id"])) as z:
        assert sorted(z.namelist()) == ["első.txt", "kép.png"]
        assert HU in z.read("első.txt").decode("utf-8")
    _engine("7zip")
    r = await run("create-archive", files, archive_format="7z", password="jelszó")
    assert r["file_id"].endswith(".7z")
    r2 = await run("extract-archive", [put_file(_out(r["file_id"]))], password="jelszó")
    with zipfile.ZipFile(_out(r2["file_id"])) as z:
        assert "első.txt" in z.namelist()
    r = await run("extract-archive", [put_file(samples.zip())])
    with zipfile.ZipFile(_out(r["file_id"])) as z:
        assert sorted(z.namelist()) == ["második.txt", "ő/első.txt"] and z.read("ő/első.txt").decode() == "első ő"


# ---------------------------------------------------------------- metaadat

async def test_metadata_image(samples, put_file):
    _engine("exiftool")
    jpg = put_file(samples.jpg())
    r = await run("read-metadata", [jpg])
    flat = json.dumps(r["data"], ensure_ascii=False)
    assert "TesztGy" in flat and "/tmp" not in flat and "Directory" not in flat
    assert r["file_id"].endswith(".json")
    r = await run("edit-metadata", [jpg], fields={"title": "Őszi kép", "author": "Kovács Ödön", "date": "2026-10-01"})
    back = (await run("read-metadata", [put_file(_out(r["file_id"]))]))["data"]
    flat = json.dumps(back, ensure_ascii=False)
    assert "Őszi kép" in flat and "Kovács Ödön" in flat and "2026:10:01" in flat
    r = await run("remove-metadata", [jpg])
    flat = json.dumps((await run("read-metadata", [put_file(_out(r["file_id"]))]))["data"], ensure_ascii=False)
    assert "TesztGy" not in flat and "Modell" not in flat
    with pytest.raises(Exception, match="unknown"):
        await run("edit-metadata", [jpg], fields={"gps": "1,2"})


async def test_metadata_pdf_office_media(samples, put_file):
    import pikepdf
    pdf = put_file(samples.pdf())
    r = await run("edit-metadata", [pdf], fields={"title": "Éves jelentés ű", "author": "Ő Kft."})
    with pikepdf.open(str(_out(r["file_id"]))) as p:
        assert str(p.docinfo["/Title"]) == "Éves jelentés ű" and str(p.docinfo["/Author"]) == "Ő Kft."
    r = await run("remove-metadata", [put_file(_out(r["file_id"]))])
    with pikepdf.open(str(_out(r["file_id"]))) as p:
        assert "/Title" not in p.docinfo and "/Metadata" not in p.Root
    docx = put_file(samples.docx())
    r = await run("edit-metadata", [docx], fields={"title": "Ügyirat ő", "author": "Űr Elek"})
    with zipfile.ZipFile(_out(r["file_id"])) as z:
        core = z.read("docProps/core.xml").decode()
        assert "<dc:title>Ügyirat ő</dc:title>" in core and "Űr Elek" in core
    from docx import Document
    d = Document(str(_out(r["file_id"])))
    assert d.core_properties.title == "Ügyirat ő" and HU in "\n".join(p.text for p in d.paragraphs)
    r = await run("remove-metadata", [put_file(_out(r["file_id"]))])
    assert Document(str(_out(r["file_id"]))).core_properties.author in ("", None)
    _engine("ffmpeg")
    r = await run("edit-metadata", [put_file(samples.mp4())], fields={"title": "Nyári videó ő"})
    assert _probe(_out(r["file_id"]))["format"]["tags"]["title"] == "Nyári videó ő"
    r = await run("remove-metadata", [put_file(_out(r["file_id"]))])
    assert "title" not in _probe(_out(r["file_id"]))["format"].get("tags", {})


# ---------------------------------------------------------------- haladó

async def test_advanced_allowlisted(samples, put_file):
    _engine("ffmpeg")
    mp4 = put_file(samples.mp4())
    r = await run("ffmpeg-advanced", [mp4], output_format="webm", args=["-vf", "scale=160:-2,hflip", "-an", "-crf", "40", "-b:v", "0"])
    p = _probe(_out(r["file_id"]))
    assert r["file_id"].endswith(".webm") and {s["codec_type"] for s in p["streams"]} == {"video"}
    for bad in (["-vf", "movie=/etc/passwd"], ["-i", "/etc/passwd"], ["-f", "lavfi"], ["-vf", "drawtext=textfile=/etc/passwd"]):
        with pytest.raises(Exception, match="args"):
            await run("ffmpeg-advanced", [mp4], output_format="mp4", args=bad)
    with pytest.raises(Exception, match="output_format"):
        await run("ffmpeg-advanced", [mp4], output_format="txt", args=["-an"])
    _engine("imagemagick")
    png = put_file(samples.png(size=(320, 200)))
    r = await run("imagemagick-advanced", [png], output_format="webp", args=["-resize", "50%", "-quality", "80"])
    with _img(r["file_id"]) as im:
        assert im.format == "WEBP" and im.size == (160, 100)
    for bad in (["-write", "/tmp/x.png"], ["-resize", "@/etc/passwd"], ["-script", "x"]):
        with pytest.raises(Exception, match="args"):
            await run("imagemagick-advanced", [png], args=bad)


# ---------------------------------------------------------------- web

async def test_website_capture_guarded(monkeypatch):
    _engine("chromium")
    import http.server
    import threading
    html = "<html><head><meta charset='utf-8'><title>Teszt</title></head><body><h1>Árvíztűrő tükörfúrógép</h1></body></html>"

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            body = html.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *a):
            pass
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{srv.server_port}/"
    try:
        with pytest.raises(Exception, match="(?i)not allowed|private|blocked"):
            await run("website-to-pdf", [], url=url)
        monkeypatch.setenv("ALLOW_PRIVATE_URLS", "1")
        try:
            r = await run("website-to-pdf", [], url=url, wait_ms=0)
        except Exception as e:  # pragma: no cover - a böngésző hiánya nem az eszköz hibája
            if "Executable doesn't exist" in str(e) or "browserType.launch" in str(e):
                pytest.skip(f"chromium can't start here: {str(e)[:80]}")
            raise
        with _pdf(r["file_id"]) as d:
            assert "Árvíztűrő tükörfúrógép" in d[0].get_text()
        r = await run("website-to-image", [], url=url, width=800, image_format="png", wait_ms=0)
        with _img(r["file_id"]) as im:
            assert im.width == 800
    finally:
        srv.shutdown()


# ---------------------------------------------------------------- segédek, és minden csatornán

async def test_utilities():
    r = await run("unit-converter", [], value=10, from_unit="km", to_unit="mi")
    assert abs(r["data"]["result"] - 6.21371) < 1e-4 and r["files"] == []
    r = await run("time-zone-converter", [], time="2026-10-15 09:00", from_zone="Europe/Budapest", to_zones=["UTC", "est"])
    assert r["data"]["to"][0]["time"].startswith("2026-10-15T07:00")


def test_every_tool_option_is_declared():
    from app.catalog.tools import TOOLS_LIST
    from app.convert.options import OPTIONS, TOOL_OPTIONS
    missing = [(t.key, o) for t in TOOLS_LIST for o in t.options if o not in OPTIONS and o not in TOOL_OPTIONS]
    assert not missing


def test_group2_handlers_all_live():
    from app.catalog import tools as T
    group2 = [t.key for t in T.TOOLS_LIST if t.group not in ("forms", "signature")]
    assert [k for k in group2 if T.handler(k) is None] == []


def _poll(client, job):
    import time
    for _ in range(400):
        if job["status"] in ("done", "failed"):
            return job
        time.sleep(0.05)
        job = client.get(f"/api/jobs/{job['id']}").json()
    raise AssertionError("job didn't finish")


def test_tools_run_on_web_route(client, samples, put_file):
    fid = put_file(samples.png(size=(300, 100)))
    job = _poll(client, client.post("/api/tools/run", json={"tool": "rotate-image", "file_ids": [fid], "options": {"angle": 90}}).json())
    assert job["status"] == "done", job
    from PIL import Image
    with Image.open(_out(job["result"]["file_id"])) as im:
        assert im.size == (100, 300)
    r = client.post("/api/tools/run", json={"tool": "rotate-image", "file_ids": [fid], "options": {"angle": 45}})
    assert r.status_code == 400
    r = client.post("/api/tools/run", json={"tool": "flip-image", "file_ids": [fid], "options": {"gps": 1}})
    assert r.status_code == 400 and "gps" in r.json()["detail"]


def test_tools_over_rest_api_and_mcp(client, samples, put_file, monkeypatch):
    """Ugyanaz az eszköz a REST API-n (tool/* feladat) és az MCP-n (file_tools) — ékezetes szöveggel."""
    import base64
    from app.config import settings
    monkeypatch.setattr(settings, "MCP_AUTH_TOKEN", "test-admin-token")
    client.headers["Authorization"] = "Bearer test-admin-token"
    try:
        pdf = base64.b64encode(samples.pdf(pages=2).read_bytes()).decode()
        job = client.post("/v1/jobs", json={"tasks": {
            "be": {"operation": "import/base64", "file": pdf, "filename": "jelentés.pdf"},
            "szamoz": {"operation": "tool/page-numbers", "input": "be", "text_format": "Oldal {n} – ő"},
            "jel": {"operation": "tool/watermark-pdf", "input": "szamoz", "text": "Tervezet ű"},
            "ki": {"operation": "export/url", "input": "jel"}}}).json()
        job = client.get(f"/v1/jobs/{job['id']}", params={"wait": 30}).json()
        assert job["status"] == "finished", job
        f = [t for t in job["tasks"] if t["name"] == "ki"][0]["result"]["files"][0]
        import pymupdf
        data = client.get(f["url"]).content
        with pymupdf.open(stream=data, filetype="pdf") as d:
            t = d[1].get_text()
            assert "Oldal 2 – ő" in t and "Tervezet ű" in t
        hdr = {"Accept": "application/json, text/event-stream", "Content-Type": "application/json"}
        client.post("/mcp/", headers=hdr, json={"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {
            "protocolVersion": "2025-06-18", "capabilities": {}, "clientInfo": {"name": "t", "version": "1"}}})
        lst = client.post("/mcp/", headers=hdr, json={"jsonrpc": "2.0", "id": 2, "method": "tools/list"}).json()
        doc = [t for t in lst["result"]["tools"] if t["name"] == "file_tools"][0]["description"]
        for key in ("compress-video", "remove-metadata", "website-to-pdf", "extract-archive", "redact-pdf"):
            assert key in doc
        fid = put_file(samples.jpg())
        r = client.post("/mcp/", headers=hdr, json={"jsonrpc": "2.0", "id": 3, "method": "tools/call", "params": {
            "name": "file_tools", "arguments": {"tool": "resize-image", "file_ids": [fid], "options": {"width": 64}}}}).json()
        res = json.loads(r["result"]["content"][0]["text"])
        assert not r["result"].get("isError"), res
        with _img(res["file_id"]) as im:
            assert im.width == 64
    finally:
        client.headers.pop("Authorization", None)


def test_web_import_url_and_chunked_upload(client, monkeypatch, samples):
    import http.server
    import threading
    body = samples.pdf().read_bytes()

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.send_header("Content-Type", "application/pdf")
            self.send_header("Content-Disposition", "attachment; filename*=UTF-8''%C5%91szi%20jelent%C3%A9s.pdf")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *a):
            pass
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{srv.server_port}/x"
    try:
        r = client.post("/api/import", json={"provider": "url", "url": url})
        assert r.status_code == 400 and "allowed" in r.json()["detail"]          # SSRF: belső cím tiltva
        monkeypatch.setenv("ALLOW_PRIVATE_URLS", "1")
        r = client.post("/api/import", json={"provider": "url", "url": url})
        assert r.status_code == 200, r.text
        assert r.json()["original_name"] == "őszi jelentés.pdf" and _out_or_in(r.json()["file_id"]).read_bytes() == body
    finally:
        srv.shutdown()
    assert client.post("/api/import", json={"provider": "ftp"}).status_code == 400
    assert client.get("/api/import/config").json()["dropbox"] is None
    data = bytes(range(256)) * 40000                                   # ~10 MB → 2 darab
    st = client.post("/api/upload/chunks", json={"filename": "nagy fájl.bin", "size": len(data)}).json()
    assert st["chunks"] == 2
    cs = st["chunk_size"]
    assert client.post(f"/api/upload/chunks/{st['upload_id']}/complete").status_code == 400     # hiányzó darab
    for i in range(st["chunks"]):
        assert client.put(f"/api/upload/chunks/{st['upload_id']}/{i}", content=data[i * cs:(i + 1) * cs]).status_code == 200
    r = client.post(f"/api/upload/chunks/{st['upload_id']}/complete").json()
    assert _out_or_in(r["file_id"]).read_bytes() == data and r["original_name"] == "nagy fájl.bin"
    assert client.put(f"/api/upload/chunks/{'x' * 20}/0", content=b"1").status_code == 404


def _out_or_in(fid: str) -> Path:
    from app.tools import crafts
    return crafts.existing(fid)
