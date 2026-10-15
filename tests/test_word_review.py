"""Beata munkafolyamata: a főnök Wordben kommentel, ő MCP-n dolgozik — minden Word-natív marad.

A minta (samples.word_review) úgy épül, ahogy a Word 2019+ ment: szál válasszal, lezárt megjegyzés, két bekezdésen
átívelő és táblázatcellába tett megjegyzés, a főnök saját követett változásai, és a Word-kiegészítő részek
(commentsExtended, commentsIds, commentsExtensible, people).
"""
from __future__ import annotations

import asyncio
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from lxml import etree

BOSS, BEATA = "Dr. Főnök Ferenc", "Kiss Beata"
W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"
W16CID = "{http://schemas.microsoft.com/office/word/2016/wordml/cid}"


@pytest.fixture()
def doc(samples, put_file):
    return put_file(samples.word_review())


def _part(fid: str, name: str):
    from app.tools.crafts import existing
    with zipfile.ZipFile(existing(fid)) as z:
        return etree.fromstring(z.read(name)) if name in z.namelist() else None


def _ex(fid: str) -> dict[str, tuple]:
    root = _part(fid, "word/commentsExtended.xml")
    return {e.get(W15 + "paraId"): (e.get(W15 + "paraIdParent"), e.get(W15 + "done")) for e in root.iter(W15 + "commentEx")}


def test_worklist_reads_word_threads(doc):
    from app.tools import review_loop
    w = review_loop.worklist(doc)
    assert (w["open"], w["resolved"]) == (3, 1)
    t = {x["id"]: x for x in w["threads"]}
    assert t["0"]["anchor"] == "2022-ben jött létre" and [r["text"] for r in t["0"]["replies"]] == ["Utánanézek."]
    assert t["4"]["paragraph"] == 3 and t["4"]["paragraph_to"] == 4                  # két bekezdésen átível
    assert t["3"]["anchor"] == "március 31."                                          # táblázatcellában
    ch = {(c["type"], c["text"]): c for c in w["tracked_changes"]}
    assert ch[("del", "4")]["id"] == "101" and ch[("ins", "5")]["author"] == BOSS


def test_address_keeps_everything_word_native(doc):
    from app.tools import review_loop
    res = asyncio.run(review_loop.address(doc, [
        {"comment": "0", "replace": {"find": "2022-ben jött létre", "with": "2022-ben, törvénnyel jött létre"},
         "reply": "Javítva.", "resolve": True},
        {"comment": "3", "replace": {"find": "március 31.", "with": "április 15."}, "reply": "Átírva."}], BEATA))
    assert [r["status"] for r in res["results"]] == ["fixed", "fixed"] and res["others_untouched"] is True
    out = res["file_id"]
    ex = _ex(out)
    assert ex["2B000001"] == (None, "1")                                             # a főnök szála lezárva
    replies = [pid for pid, (par, _) in ex.items() if par == "2B000001"]
    assert len(replies) == 2                                                          # Beata régi + új válasza a szálban
    ids = {e.get(W16CID + "paraId") for e in _part(out, "word/commentsIds.xml").iter(W16CID + "commentId")}
    assert set(replies) <= ids                                                        # tartós azonosító az új válasznak is
    people = [p.get(W15 + "author") for p in _part(out, "word/people.xml").iter(W15 + "person")]
    assert people.count(BEATA) == 1 and BOSS in people


def test_beata_review_moves(doc):
    from app.tools import review_loop
    res = asyncio.run(review_loop.address(doc, [
        {"new_comment": "Ide kell a pontos összeg?", "on": "milliárd forint"},
        {"new_comment": "Ezt átírnám.", "paragraph": 9},
        {"reject": "102"}, {"reject": "101"},                                       # a főnök 4 → 5 cseréje vissza
        {"comment": "2", "reopen": True},
        {"comment": "1", "edit": "Utánanéztem: 2022. évi XXVII. törvény."},
        {"comment": "4", "delete": True}], BEATA))
    st = [r["status"] for r in res["results"]]
    assert st == ["commented", "commented", "rejected", "rejected", "reopened", "edited", "deleted"], res["results"]
    out = res["file_id"]
    w = review_loop.worklist(out, include_resolved=True)
    t = {x["id"]: x for x in w["threads"]}
    assert t["5"]["author"] == BEATA and t["5"]["anchor"] == "milliárd forint"
    assert t["6"]["anchor"] == "Záró bekezdés, ő és ű betűkkel."
    assert t["2"]["resolved"] is False and "4" not in t
    assert t["0"]["replies"][0]["text"] == "Utánanéztem: 2022. évi XXVII. törvény."
    assert w["tracked_changes"] == []
    from app.tools.crafts import existing
    from docx import Document
    assert "A költségvetés évi 4 milliárd forint." in [p.text for p in Document(str(existing(out))).paragraphs]
    xml = zipfile.ZipFile(existing(out)).read("word/document.xml").decode()
    assert not re.search(r'<w:comment(?:RangeStart|RangeEnd|Reference) w:id="4"/>', xml)   # a törölt szál nyoma sem


def test_accept_by_author_then_final_clean_copy(doc):
    from app.tools import review_loop
    from app.tools.crafts import existing
    from docx import Document
    res = asyncio.run(review_loop.address(doc, [{"accept": {"author": BOSS}}], BEATA))
    assert res["results"][0]["changes"] == 2
    paras = [p.text for p in Document(str(existing(res["file_id"]))).paragraphs]
    assert "A költségvetés évi 5 milliárd forint." in paras
    fin = asyncio.run(review_loop.address(res["file_id"], [{"final": True}], BEATA))
    r = fin["results"][0]
    assert r["status"] == "final" and r["comments_removed"] == 5
    z = zipfile.ZipFile(existing(fin["file_id"]))
    xml = z.read("word/document.xml").decode()
    assert "commentRange" not in xml and "<w:ins " not in xml and "<w:del " not in xml
    assert not list(_part(fin["file_id"], "word/comments.xml").iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}comment"))
    assert z.testzip() is None


def test_bad_items_fail_loudly_not_silently(doc):
    from app.tools import review_loop
    res = asyncio.run(review_loop.address(doc, [
        {"new_comment": "x", "on": "nincs ilyen szöveg"}, {"accept": "999"}, {"comment": "0", "reply": "ok"}], BEATA))
    st = [(r.get("action"), r["status"]) for r in res["results"]]
    assert st == [("new_comment", "failed"), ("accept", "failed"), (None, "replied")]
    assert res["file_id"] and res["check"]["status"] == "failed"
    assert any("new_comment: FAILED" in f for f in res["check"]["findings"])


def test_word_extras_survive_an_engine_rewrite(samples, tmp_path):
    """A WordCraft eldobja a commentsIds / commentsExtensible / people részeket és átszámozza a bekezdés-azonosítókat —
    a helyreállítás a kimenet saját azonosítóihoz igazítva teszi vissza."""
    from app.tools import docx_extras
    src = samples.word_review()
    out = tmp_path / "engine.docx"
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zi.infolist():
            if i.filename in ("word/commentsIds.xml", "word/commentsExtensible.xml", "word/people.xml"):
                continue
            data = zi.read(i.filename)
            if i.filename in ("word/comments.xml", "word/commentsExtended.xml"):
                data = re.sub(rb"2B00000(\d)", rb"1000000\1", data)
            if i.filename == "word/_rels/document.xml.rels":
                data = re.sub(rb'<Relationship [^>]*(?:commentsIds|commentsExtensible|people)[^>]*/>', b"", data)
            zo.writestr(i.filename, data)
    got = docx_extras.restore(src, out)
    assert got == ["word/commentsExtensible.xml", "word/commentsIds.xml", "word/people.xml"]
    z = zipfile.ZipFile(out)
    ids = etree.fromstring(z.read("word/commentsIds.xml"))
    assert sorted(e.get(W16CID + "paraId") for e in ids.iter(W16CID + "commentId")) == [f"1000000{k}" for k in range(1, 6)]
    rels = z.read("word/_rels/document.xml.rels").decode()
    assert "commentsIds" in rels and "people" in rels and "commentsIds+xml" in z.read("[Content_Types].xml").decode()
    assert docx_extras.restore(src, out) == []                                       # második futás: nincs teendő


def test_real_wordcraft_keeps_threads_and_extras(samples, put_file):
    from app.tools import review_loop, wordcraft
    if not wordcraft.available():
        pytest.skip("WordCraft CLI is only in the production image")
    fid = put_file(samples.word_review())
    r = asyncio.run(wordcraft.run(file_id=fid, commands=[{"command": "edit.goto", "params": {"paragraph": 9}},
                                                         {"command": "select.paragraph"}, {"command": "format.bold"}],
                                  author=BEATA))
    out = r["file_id"]
    for part in ("word/commentsExtended.xml", "word/commentsIds.xml", "word/commentsExtensible.xml", "word/people.xml"):
        assert _part(out, part) is not None, part
    w = review_loop.worklist(out, include_resolved=True)
    t = {x["id"]: x for x in w["threads"]}
    assert [x["text"] for x in t["0"]["replies"]] == ["Utánanézek."] and t["2"]["resolved"] is True
    assert {(c["type"], c["author"]) for c in w["tracked_changes"]} == {("del", BOSS), ("ins", BOSS)}
